"""korean-tax-calc-mcp — 한국어 세금 계산 MCP 서버(코드 계산, AI 없음).

세율·한도·가산세는 사업연도·과세기간·공급일 기준 연도표(2016~2026)로 계산하고, 표에 없는 연도는 추정하지 않고 멈춘다.
결과마다 근거 조문을 붙인다. 근거 원문은 korean-tax-mcp(law_article)로 확인한다.
"""
import asyncio
import calendar
import functools
import math
import os
import re
from datetime import date
from typing import Annotated, Any, Literal

from mcp.server.mcpserver import MCPServer
from mcp.types import ToolAnnotations
from pydantic import BaseModel, Field

from .engine import calc as K, calc_cit as C, calc_income as I, calc_vat as V, calc_capgains as CG
from .engine import calc_inhgift as IG
from .engine import calc_gold_gaps6 as GG
from .engine import calc_tax_priority as TP
from .engine import calc_stock as ST
from .engine import calc_stock_price as SP
from .engine import calc_stock_val as SV
from .engine import calc_official_price as OP
from .engine import calc_building as BL
from .engine import calc_donation_ind as DI
from . import vocab
from .forms_workbook import render_form_sheet, _load_form_templates, tool_value, nonbusiness_form_sources, nonbusiness_form_data

# ── 세무조사 판정 도구 (audit 패키지) ──
from .audit.agents import related as _R
from .audit.agents.related_intake import norm_name
from .audit.agents.tunnel_intake import upstage as _tunnel_upstage
from .audit.agents.return_intake import upstage as _return_upstage
from .audit.agents.precheck import corp as _precheck_corp, vat as _precheck_vat, income as _precheck_income, run as _precheck_run
from .audit.agents.precheck.needs import guide as _precheck_guide
from .audit.agents.precheck import REGISTRY as _AUDIT_REGISTRY, Data as _AuditData, Missing as _AuditMissing
from .audit.agents import calc as _audit_calc, related as _audit_related

AUDIT_INSTRUCTIONS = (
    "한국 세무조사 판정 도구. 특수관계인 판정은 반드시 기준일(사실관계 시점, YYYY-MM-DD)을 넣는다 — "
    "친족 범위(2023.3.1. 전 6촌·4촌, 이후 4촌·3촌)와 조문 위치가 시점마다 다르다. "
    "결과의 '근거'를 그대로 인용하고, 도구가 '확인 필요'라고 한 부분은 추측으로 채우지 않는다. "
    "PDF에서 표를 뽑는 도구만 Upstage(Document Parse + Solar Pro 4)로 전송하며, 나머지는 로컬 계산이다.")

AI_GENERATED = "이 결과의 표 추출·정리는 생성형 AI(Upstage Solar Pro 4)가 수행했습니다. 사람이 원본과 대조해 확인하세요."
AI_GENERATED_EN = "Table extraction and organization in this result was performed by generative AI (Upstage Solar Pro 4). Please verify against the original document."

def _upstage_key():
    import os
    return os.environ.get("UPSTAGE_API_KEY")

def _pdf_guard(fn):
    """PDF 도구용 가드: 키 없으면 mode=host_ai로 로컬 추출, 있으면 solar_cloud."""
    @functools.wraps(fn)
    def w(*a, **k):
        if not _upstage_key():
            return _host_ai_response(k.get("pdf_path", ""), fn.__name__)
        try:
            return fn(*a, **k)
        except Exception as e:
            return {"error": str(e), "안내": "PDF 처리 중 오류 — 문서 형식·내용을 확인하세요"}
    return w


HOST_AI_CLEANUP = (
    "정리 안내 — 추출한 텍스트를 아래 형식으로 다듬어 다음 도구에 넣으세요.\n"
    "■ 추출 텍스트: pypdf로 PDF에서 로컬 추출한 텍스트(위 '추출 텍스트' 필드). 페이지 순서로 붙어 있으며 표·그림은 포함되지 않을 수 있습니다.\n"
    "※ 이 결과는 설치한 컴퓨터의 pypdf가 추출한 것입니다 — 사용자에게 보여 줄 때 AI가 추출한 것임을 표시하세요.\n"
    "※ UPSTAGE_API_KEY를 설정하면 Document Parse + Solar Pro 4로 더 정교한 표 추출이 가능합니다(선택 사항).\n"
)

_SCAN_GUIDE = (
    "이 PDF는 텍스트가 거의 없는 스캔 이미지로 보입니다. "
    "지금 쓰는 AI에게 이 PDF를 직접 첨부해 읽게 한 뒤, 아래 정리 안내 형식으로 다음 도구에 넣으세요. "
    "다음 도구 호출 예시도 정리 안내 안에 함께 들어 있습니다."
)

_NEXT_TOOLS = {
    "extract_relations": "related_judge — 사람·가족·지분 관계를 판정으로 연결",
    "tunnelling_from_pdf": "tunnelling_gift_nts — 일감몰아주기 증여의제이익 계산",
    "return_precheck_pdf": "return_precheck — 신고서 사전검토(오류 의심 목록)",
}

_TOOL_CASE_DOC = (
    "people: {이름: {\"구분\": \"개인\"|\"법인\"}}\n"
    "family: [[a, b, \"부모자녀\"|\"배우자\"]] — 부모자녀는 [부모, 자녀]\n"
    "stakes: [[보유자, 법인, 지분율(0~1)]] — 직접보유비율\n"
    "officers/employees: {법인: [이름]}, livelihood: [[부양자, 피부양자]], "
    "influence: [[사실상 영향력자, 법인]], group: {기업집단: [계열회사]}\n"
    "livelihood_shared: [{친족, 상대방(임원·사용인·생계유지자), 주민등록_동일세대: true|false|모름, 동거_사실: 같은집_공유구조|같은집_독립구조|따로거주|모름, 생활비_공동부담: 공동|각자분담_증빙있음|각자분담_증빙없음|일방지원|모름, 독립생계_소득: 안정소득_있음|소액_일시|없음|모름, 경제적_지원: 정기송금_있음|차입후상환|용돈수준|제3자_지원|없음|모름, 별거_사유: 취학|질병요양|근무|사업|혼인파탄|독립|해당없음, 별도_임대차: bool, 별거_증빙_유형: [공과금·관리비_명의|금융_송금|출입국|전화·인터넷_가입|임대차_중개확인|진술서만|없음], 연말정산_부양가족_등록: bool, 수입_상대방계좌_입금: bool, 혼인_독립가구: bool, 친족관계?·촌수?}] — 생계를 함께하는 친족"
)

_TOOL_TUNNEL_GUIDE = (
    "size=중소|중견|일반, sales_by_corp={거래처 법인: 매출액}, "
    "excluded_sales=⑩항 과세제외매출액(입력). relations={관계명: 보유비율}(예: {'직접':0.2,'A경유':0.18}), "
    "holdings_in_related={특관법인: 지배주주등 보유비율}, "
    "indirect_corp_of={관계명: 그 경로 간접출자법인}, "
    "excluded_by_corp={특관법인: ⑩항 과세제외액}. "
    " 국세청 2026 신고안내 작성사례 원 단위 검증."
)

_TOOL_PRECCHECK_GUIDE = (
    "data={입력키: 값} — 입력키는 agents/precheck/{corp,vat,income}.py SCHEMA. "
    "tax=법인|부가|소득(생략 시 전체). 예: {\"신고서.수입금액\": 5000000000, \"조정계산서.산출세액\": 50000000}"
)


def _host_ai_response(pdf_path: str, tool_name: str | None = None) -> dict:
    """UPSTAGE_API_KEY 없이 PDF 도구 호출 → mode=host_ai 응답자."""
    from .audit.agents import upstage
    raw = upstage.extract_text_local(pdf_path)
    text = raw.strip()
    scanned = len(text) < 200  # 텍스트 거의 없음 → 스캔 PDF

    extraction_text = text[:20000] if text else ""
    next_tool_fn = _NEXT_TOOLS.get(tool_name, "해당 도구")

    if scanned:
        guidance = _SCAN_GUIDE + "\n\n" + _cleanup_guide(tool_name)
        extraction_text = "(스캔 PDF — pypdf 텍스트 추출 결과 없음. PDF를 지금 쓰는 AI에 첨부해 읽게 한 뒤 아래 형식으로 정리하세요.)"
    else:
        guidance = _cleanup_guide(tool_name)

    return {
        "mode": "host_ai",
        "추출 텍스트": extraction_text,
        "정리 안내": guidance + f"\n다음 도구: {next_tool_fn}",
        "다음 도구": next_tool_fn,
        "선택": ("UPSTAGE_API_KEY를 설정하면 Document Parse + Solar Pro 4로 더 정교한 표 추출이 가능합니다." if not scanned else
                 "UPSTAGE_API_KEY를 설정하면 Document Parse + Solar Pro 4로 OCR·표 추출이 가능합니다."),
    }


def _cleanup_guide(tool_name: str) -> str:
    """도구별 정리 안내 텍스트(다음 도구 입력 형식 + 예시)."""
    common = "※ 이 결과는 설치한 컴퓨터의 pypdf가 추출한 것입니다 — 사용자에게 보여 줄 때 AI가 추출한 것임을 표시하세요.\n"
    if tool_name == "extract_relations":
        return (common +
                "INPUT 포맷(CASE_DOC):\n" + _TOOL_CASE_DOC + "\n\n"
                "예시:\n"
                "people: {\"김갑\": {\"구분\": \"개인\"}, \"㈜대한\": {\"구분\": \"법인\"}}\n"
                "family: [[\"김갑\", \"김을\", \"부모자녀\"]]\n"
                "stakes: [[\"김갑\", \"㈜대한\", 0.30]]\n"
                "다음 도구: related_judge(a=\"김갑\", b=\"김을\", on=\"2024-01-01\", people=..., family=..., stakes=...)")
    if tool_name == "tunnelling_from_pdf":
        return (common +
                "INPUT 포맷:\n" + _TOOL_TUNNEL_GUIDE + "\n\n"
                "예시:\n"
                "tunnelling_gift_nts(size=\"중소\", op_income=1000000000, taxable_income=800000000, "
                "tax=100000000, sales_total=1000000000, sales_by_corp={\"A\": 300000000}, "
                "relations={\"직접\": 0.20}, base_excluded=200000000)")
    if tool_name == "return_precheck_pdf":
        return (common +
                "INPUT 포맷:\n" + _TOOL_PRECCHECK_GUIDE + "\n\n"
                "예시:\n"
                "return_precheck(data={\"신고서.수입금액\": 5000000000, \"조정계산서.산출세액\": 50000000}, tax=\"법인\")")
    return HOST_AI_CLEANUP

def _add_ai_mark(result):
    """PDF 도구 결과에 AI 생성 표시 필드 추가."""
    result["AI 생성 표시"] = AI_GENERATED
    result["AI generation notice"] = AI_GENERATED_EN
    return result

# audit 서버 server.py의 CASE_DOC (도구 설명문에 사용)
CASE_DOC = """people: {이름: {"구분": "개인"|"법인"}}
family: [[a, b, "부모자녀"|"배우자"]] — 부모자녀는 [부모, 자녀]
stakes: [[보유자, 법인, 지분율(0~1)]] — 직접보유비율
officers/employees: {법인: [이름]}, livelihood: [[부양자, 피부양자]], influence: [[사실상 영향력자, 법인]], group: {기업집단: [계열회사]}, group_exempt: {법인: 공정위 계열편입 유예·제외 통지 내용} — 이 법인은 계열회사로 보지 않음
livelihood_shared: [{친족, 상대방(임원·사용인·생계유지자), 주민등록_동일세대: true|false|모름, 동거_사실: 같은집_공유구조|같은집_독립구조|따로거주|모름, 생활비_공동부담: 공동|각자분담_증빙있음|각자분담_증빙없음|일방지원|모름, 독립생계_소득: 안정소득_있음|소액_일시|없음|모름, 경제적_지원: 정기송금_있음|차입후상환|용돈수준|제3자_지원|없음|모름, 별거_사유: 취학|질병요양|근무|사업|혼인파탄|독립|해당없음, 별도_임대차: bool, 별거_증빙_유형: [공과금·관리비_명의|금융_송금|출입국|전화·인터넷_가입|임대차_중개확인|진술서만|없음], 연말정산_부양가족_등록: bool, 수입_상대방계좌_입금: bool, 혼인_독립가구: bool, 친족관계?·촌수?}] — 생계를 함께하는 친족"""

mcp = MCPServer(
    "korean-tax-calc-mcp", title="Korea Tax Calculator (한국 세금 계산)",
    instructions="도구가 70개가 넘으므로 맞는 도구 이름을 모르면 먼저 find_tool(query=평소 말)로 찾고, '도구가 없다'고 판단하기 전에도 find_tool을 부를 것. "
                 "계산 문제는 먼저 해당 계산 도구가 있는지 확인하고, 문제 자료를 도구 입력으로 옮겨 계산할 것. "
                 "도구 결과의 단계별 값을 그대로 인용할 것. "
                 "한국 세금 계산은 반드시 이 도구 결과를 쓰고 직접 계산하지 않는다. 결과의 근거 조문을 함께 제시한다. "
                 "연도표가 없는 해는 오류로 멈추므로 추정하지 않는다. 계산 근거 원문·해석은 korean-tax-mcp로 확인한다. "
                 + AUDIT_INSTRUCTIONS)
CALC = ToolAnnotations(read_only_hint=True, destructive_hint=False, idempotent_hint=True, open_world_hint=False)
YEAR = Annotated[int, Field(description="사업연도·귀속연도(2016~2026)", ge=2016, le=2026)]
WON = Annotated[int, Field(description="금액(원)", ge=0)]


# list[dict] 인자용 Pydantic 모델 — 입력 검증 강화 (지시서 4)
class Movement(BaseModel):
    date: str       # YYYY-MM-DD
    amount: int     # 원 (대여 +, 회수 -)


class Loan(BaseModel):
    interest: int   # 지급이자(원)
    jeoksu: int | None = None     # 차입금 적수(원·일)
    kind: str | None = None       # 일반|채권자불분명|비실명|건설자금|기업구매자금|금융회사제외차입금
    rate: float | None = None     # 이자율(소수) — 적수 = 이자 ÷ 이자율 × 사업연도 일수
    principal: int | None = None  # 원금(원) — days와 함께 적수 = 원금 × 일수
    days: int | None = None
    withholding: int | None = None  # 채권자불분명·비실명 이자의 원천징수세액(소득처분 분리용)


class ReceiptItem(BaseModel):
    amount: int           # 1회 지출액(원)
    qualified: bool = False       # 적격증빙 여부
    congratulatory: bool = False  # 경조금 여부


# ── 세무조사 판정 도구 입력 스키마 ──
class PersonEntry(BaseModel):
    구분: Literal["개인", "법인"] = "개인"


class FamilyLink(BaseModel):
    a: str
    b: str
    관계: Literal["부모자녀", "배우자"]


class StakeEntry(BaseModel):
    보유자: str
    법인: str
    지분율: Annotated[float, Field(description="직접보유비율(소수, 0~1)")]


class OfficerEntry(BaseModel):
    법인: str
    이름: list[str]


class LivelihoodLink(BaseModel):
    부양자: str
    피부양자: str


class InfluenceEntry(BaseModel):
    영향력자: str
    법인: str


class GroupEntry(BaseModel):
    기업집단: str
    계열회사: list[str]


class GroupExemptEntry(BaseModel):
    법인: str
    통지내용: str


# ── 끝에서 끝까지 계산기 공용 입력 검증 (Claude 10-13): 모르는 kind·빠진 사실관계는 조용히 기본값으로 계산하지 않고 바로 돌려보낸다 ──
def _validate_items(items, field, specs, extra=None):
    """items: 목록, field: 종류 필드명, specs: {종류: {"필요": [필드...], "허용": [필드...], "설명": str}}.
    문제가 있으면 {"error", "문제", "허용값", "안내"} 반환, 없으면 None.
    항목의 키 중 kind·field_Name과 sp["필요"]·sp["허용"]에 없는 키가 있으면 오류로 반환."""
    problems = []
    if not isinstance(items, list) or not items:
        problems.append(f"{field} 목록이 비어 있음")
    for i, it in enumerate(items or []):
        k = (it or {}).get(field)
        if k not in specs:
            problems.append(f"항목 {i+1}: {field}='{k}' 은(는) 지원하지 않음")
            continue
        sp = specs[k]
        for need in sp.get("필요", []):
            if it.get(need) in (None, "", []):
                problems.append(f"항목 {i+1}({k}): '{need}' 입력 필요 — {sp.get('설명', '')}")
        # 항목에 모르는 키가 있으면 오류
        필요셋 = set(sp.get("필요", []))
        허용셋 = set(sp.get("허용", []))
        허용전체 = 필요셋 | 허용셋 if 허용셋 else 필요셋
        항목키 = set((it or {}).keys()) - {field}
        모르는키 = 항목키 - 허용전체
        if 모르는키:
            problems.append(f"항목 {i+1}({k}): 모르는 키 {sorted(모르는키)} — 허용 키: {sorted(허용전체)}")
    if extra:
        problems += extra
    if not problems:
        # 별칭(예: 파산 → 회수불능)으로 들어온 종류는 정식 이름으로 바꿔 계산 분기가 알아보게 함
        for it in items:
            k = it.get(field)
            try: canon = vocab.resolve(k)["이름"]
            except Exception: continue
            if canon != k and canon in specs:
                it[field] = canon
        return None
    return {"error": "입력 확인 필요 — 계산하지 않음", "문제": problems,
            "허용값": {k: {"필요": v.get("필요", []), "허용": v.get("허용", []), "설명": v.get("설명", "")} for k, v in specs.items()},
            "안내": "문제·장부에 적힌 사실관계를 위 허용값의 종류로 나눠 다시 호출하세요. 한 금액에 두 성격이 섞여 있으면 항목을 나누세요."}


# 한글 사실관계 필드명 → 도구 입력 필드명 매핑
# 도구별로 필수로 요구할 입력 필드 (vocab의 필요한사실 중 일부만 필수 처리)
_도구별_필수필드 = {
    "financial_income_aggregate": {
        "예금이자": {"필요": ["amount"], "허용": ["withheld"]},
        "비영업대금이익": {"필요": ["amount"], "허용": ["withheld"]},
        "소비대차전환이자": {"필요": ["amount"], "허용": ["withheld"]},
        "외상매출금지급기일연장이자": {"필요": ["amount"], "허용": ["withheld"]},
        "저축성보험차익": {"필요": ["premium", "payout", "contract_year"], "허용": ["withheld", "amount", "annuity", "exempt_conditions_met"]},
        "투자신탁환매이익": {"필요": ["amount", "components"], "허용": ["withheld"]},
        "파생결합사채이익": {"필요": ["amount"], "허용": ["withheld"]},
        "장내파생상품이익": {"필요": ["amount"], "허용": ["withheld"]},
        "상장법인현금배당": {"필요": ["amount"], "허용": ["withheld"]},
        "비상장법인현금배당": {"필요": ["amount"], "허용": ["withheld"]},
        "의제배당_잉여금자본전입": {"필요": ["amount", "source"], "허용": ["withheld", "self_stock_shift"]},
        "인정배당": {"필요": ["amount"], "허용": ["withheld"]},
        "손해배상법정이자": {"필요": ["amount"], "허용": ["withheld"]},
        "공익신탁이익": {"필요": ["amount"], "허용": ["withheld"]},
        "통합발행채권할인액": {"필요": ["amount"]},
        "외국법인배당": {"필요": ["amount"], "허용": ["withheld"]},
        "환매조건부매매차익": {"필요": ["amount"], "허용": ["withheld"]},
        "출자공동사업자배당": {"필요": ["amount"]},
        "법원보증금이자": {"필요": ["amount"]},
        "직장공제회초과반환금": {"필요": ["amount"], "허용": ["contribution_years"]},
    },
    "bad_debt_adjustment": {
        "소멸시효완성": {"필요": ["amount", "arose_in"]},
        "화해·확정판결": {"필요": ["amount", "arose_in"]},
        "부도6개월": {"필요": ["amount", "arose_in", "instrument", "months_since_default"], "허용": ["mortgage"]},
        "회수불능": {"필요": ["amount", "arose_in"]},
    },
    "entertainment_adjustment": {
        "없음": {"필요": ["amount"]},
        "영수증 등": {"필요": ["amount"]},
        "적격": {"필요": ["amount"]},
    },
    "other_income_aggregate": {
        "산업재산권대여": {"필요": ["amount"], "허용": ["actual_cost", "actual_expense", "receipt_date"]},
        "공익사업지상권대여": {"필요": ["amount"], "허용": ["actual_cost", "actual_expense", "receipt_date"]},
        "서화골동품양도": {"필요": ["amount"], "허용": ["actual_cost", "held_years", "museum", "living_domestic_artist", "business_setup"]},
        "일시적인적용역": {"필요": ["amount"], "허용": ["actual_cost", "actual_expense", "receipt_date"]},
        "뇌물": {"필요": ["amount"], "허용": ["forfeited"]},
        "위약금·배상금": {"필요": ["amount"], "허용": ["deposit_substituted"]},
        "주택입주지체상금": {"필요": ["amount"]},
        "연금외수령": {"필요": ["amount"]},
        "직무발명보상금_퇴직후": {"필요": ["amount", "employment_exempt_used"]},
        "승마투표권환급금": {"필요": ["amount"], "허용": ["unit_bet", "ticket_total", "per_bet_refund"]},
        "슬롯머신당첨금": {"필요": ["amount"], "허용": ["stake"]},
        "복권당첨금": {"필요": ["amount"], "허용": ["actual_cost"]},
    },
    "earned_income_gross_pay": {
        "급여": {"허용": ["facts", "service_year", "regular_monthly"]},
        "상여": {"허용": ["facts", "service_year"]},
        "잉여금처분상여": {"허용": ["facts", "resolution_date"]},
        "인정상여": {"허용": ["facts", "service_year"]},
        "식사대": {"허용": ["facts", "months", "meal_provided", "regular_monthly"]},
        "자가운전보조금": {"허용": ["facts", "months", "vehicle", "business_use", "actual_travel_reimbursed", "regular_monthly"]},
        "출산·보육수당": {"허용": ["facts", "purpose", "months", "child_age", "child_ages", "child_count", "child", "within_2y_of_birth", "payment_count", "related_party", "regular_monthly"]},
        "직무발명보상금": {"허용": ["facts", "retired", "related_party", "regular_monthly"]},
        "연장근로수당": {"허용": ["facts", "eligible_occupation", "monthly_regular_wage", "prior_year_gross", "mine_or_daily", "months"]},
        "사택제공이익": {"허용": ["facts", "position", "regular_monthly"]},
        "주택자금대여이익": {"허용": ["facts", "sme", "related_party", "regular_monthly"]},
        "단체보험료": {"허용": ["facts", "regular_monthly"]},
        "학자금": {"허용": ["facts", "beneficiary", "requirements_met", "regular_monthly"]},
        "실비변상여비": {"허용": ["facts", "actual_cost", "regular_monthly"]},
        "사회보험료_본인부담분_회사대납": {"허용": ["facts", "regular_monthly"]},
        "사회보험료_사용자부담분": {"허용": ["facts"]},
    },
}


def _specs_from_vocab(tool_name: str, field: str) -> dict:
    """vocab.items_for로 _validate_items용 specs를 생성한다.

    허용값 키: kind 필드는 항목 이름, evidence 필드는 매핑값(='evidence=값'의 값).
    별칭도 허용값에 포함하고, 설명에는 항목의 뜻을 쓴다.
    필수 필드(필요)는 도구별_필수필드에서 가져온다(어휘의 필요한사실보다 완화된 기준).
    허용 필드(허용)는 도구별_필수필드에 명시된 추가 허용 키 목록 — 미지수 키 검증용.
    """
    specs: dict = {}
    필수맵 = _도구별_필수필드.get(tool_name, {})

    for item in vocab.items_for(tool_name, field):
        의미 = item.get("뜻", "")
        aliases = [a[0] for a in item.get("별칭", [])]

        # 필수 필드는 도구별 맵에서 가져온다
        if "_매핑값" in item:
            key = item["_매핑값"]
        else:
            key = item["이름"]

        필드설정 = 필수맵.get(key, {"필요": ["amount"]})
        필요 = 필드설정.get("필요", ["amount"])
        허용 = 필드설정.get("허용", [])

        specs[key] = {"필요": 필요, "허용": 허용, "설명": 의미}
        # 별칭도 허용값에 추가
        for alias in aliases:
            if alias not in specs:
                specs[alias] = {"필요": 필요, "허용": 허용, "설명": 의미}
    return specs




def _build_tax_adjustment_result(settlements, result, calc_dict=None):
    """세무조정 출력 통일: 결과 첫 키에 세무조정·소득처분별_합계·손금불산입_합계 배치."""
    처분별 = {}
    for s in settlements:
        처분 = s["소득처분"]
        처분별[처분] = 처분별.get(처분, 0) + s["금액"]
    처분별 = {k: v for k, v in 처분별.items() if v != 0}
    손금불산입 = sum(v for k, v in 처분별.items() if k != "△유보")
    new_result = {"세무조정": settlements, "소득처분별_합계": 처분별, "손금불산입_합계": 손금불산입}
    new_result.update(result)
    return new_result


def _ok(r, basis, **extra):
    return {"결과": r, "근거": basis, **extra}


def _front(r):
    """결과 dict의 키 순서를 결과 → 계산 → 나머지로 (답이 되는 값을 먼저 보이게, Claude 10-14)"""
    if isinstance(r, dict) and "결과" in r:
        return {k: r[k] for k in ("결과", "계산") if k in r} | {k: v for k, v in r.items() if k not in ("결과", "계산")}
    return r


import json as _json

_CALL_MEMO: dict[str, int] = {}   # key → 호출 횟수 (최근 500개만 보관, 초과 시 오래된 것부터 삭제)


def _reset_call_memo() -> None:
    """테스트용: 호출 메모 초기화(모듈 수준, 도구 아님)."""
    _CALL_MEMO.clear()


def _call_key(fn_name: str, args: tuple, kwargs: dict) -> str:
    """프로세스 안에서 같은 호출을 식별하기 위한 정규화 키."""
    return _json.dumps({"args": args, "kwargs": kwargs}, sort_keys=True, ensure_ascii=False, default=str)


def _guard(fn):
    import functools

    @functools.wraps(fn)
    def w(*a, **k):
        key = (fn.__name__, _call_key(fn.__name__, a, k))
        count = _CALL_MEMO.get(key, 0) + 1
        _CALL_MEMO[key] = count
        # 최근 500개 초과분 정리(가장 오래된 키부터 제거)
        if len(_CALL_MEMO) > 500:
            for _ in range(len(_CALL_MEMO) - 500):
                _CALL_MEMO.pop(next(iter(_CALL_MEMO)))
        repeat = count >= 2
        try:
            inner = fn(*a, **k)
            r = _front(inner)
        except (C.NoYearTable, getattr(V, "NoYearTable", C.NoYearTable), getattr(I, "NoYearTable", C.NoYearTable)) as e:
            r = {"error": str(e), "주의": "연도표 없는 해는 추정 계산하지 않음"}
        except (ValueError, KeyError, TypeError) as e:
            r = {"error": f"입력 오류: {e}"}
        except Exception as e:
            # Pydantic ValidationError 등 프레임워크 검증 오류 — 문제 필드명을 결과에 포함
            msg = str(e)
            r = {"error": f"입력 확인 필요 — 계산하지 않음", "문제": [msg]}
        if repeat:
            r = {"안내": f"같은 입력으로 {count}번째 계산입니다. 결과는 처음과 같습니다 — 빠진 사실이나 잘못 넣은 값이 없다면 이 결과를 그대로 쓰세요", **r}
        return r
    return w


@mcp.tool(annotations=CALC)
@_guard
def corporate_tax(base: WON, year: YEAR,
                  months: Annotated[int, Field(description="사업연도 월수(1년 미만이면 1~11)")] = 12) -> dict:
    """법인세 산출세액(사업연도 과세표준 × 세율표, 1년 미만 사업연도 월할 환산). 언제: 법인세 신고·세무조정에서 과세표준이 정해졌을 때 산출세액을 구한다.
입력: 과세표준(과세표준, 원) → base, 사업연도(2016~2026) → year, 월수(1~12) → months.
결과: 산출세액·과세표준·사업연도·월수와 근거(법인세법 제55조①·②). Corporate income tax on a tax base for a fiscal year (rate table, prorated for short fiscal years)."""
    return _ok({"산출세액": C.corp_tax(base, year, months), "과세표준": base, "사업연도": year, "월수": months}, "법인세법 제55조(세율)·제55조②(1년 미만)")


@mcp.tool(annotations=CALC)
@_guard
def foreign_tax_credit(
    year: YEAR,
    base: Annotated[int, Field(description="각 사업연도 소득에 대한 과세표준(원). Gross-up 반영 여부는 base_includes_grossup으로 지정", ge=0)],
    base_includes_grossup: Annotated[bool, Field(description="base에 간접외국납부세액 익금산입(Gross-up)이 이미 들어 있으면 true, 아니면 false(도구가 더함)")],
    method: Annotated[Literal["세액공제", "손금산입"], Field(description="외국납부세액 처리 방법. 손금산입은 2020 사업연도까지만(2021~ 삭제)")],
    countries: Annotated[list, Field(description="국가별 [{country, foreign_income, direct_tax, deemed_tax?, related_costs?, reduction_excluded?, subsidiaries?}] — "
                                                 "foreign_income: 국외원천소득(Gross-up 전, 수입배당금은 원천징수 전 금액 포함), direct_tax: 직접외국납부세액(자회사 배당 원천징수세 제외), "
                                                 "deemed_tax: 간주외국납부세액, related_costs: 국외원천소득대응비용, reduction_excluded: 감면·면제 대상 국외원천소득×감면비율, "
                                                 "subsidiaries: [{name, dividend, withholding_tax?, received_on, ownership, holding_months, resource_dev, sub_income, sub_tax, participation_exempt?}]")],
    months: Annotated[int, Field(description="사업연도 월수(1~12)", ge=1, le=12)] = 12,
) -> dict:
    """법인세 외국납부세액공제(직접·간주·간접, 국가별 한도, 한도초과 이월). 언제: 국외원천소득에 외국 법인세를 낸 내국법인의 공제액·Gross-up·이월액을 구할 때.
입력:
  year — 사업연도(2016~2026). 이월기간·손금산입방법 허용 여부가 연도별.
  base — 과세표준(원). base_includes_grossup — base에 Gross-up 포함 여부.
  method — '세액공제'|'손금산입'(2020년까지).
  countries — 국가별 foreign_income(Gross-up 전 국외원천소득), direct_tax, deemed_tax, related_costs(대응비용), reduction_excluded,
    subsidiaries(외국자회사 배당: dividend 원천징수 전, withholding_tax, received_on, ownership 소수, holding_months 배당기준일 현재, resource_dev, sub_income·sub_tax 자회사 소득금액·법인세액, participation_exempt 2023~ 제18조의4 적용 여부).
  months — 사업연도 월수.
결과: 간접외국납부세액, Gross-up 익금산입액, 국가별 국외원천소득·공제한도·공제액·한도초과 이월액·이월배제액, 산출세액, 공제액 합계, 이월공제기간, 산식, 근거.
  지분율 요건은 배당 받은 날 기준(2022년까지 25%, 2023~ 10%, 해외자원 5%). 손자회사·이월액 공제·이월결손금 안분(영 §94⑥)은 반영 안 함 — 국외원천소득에 미리 반영."""
    r = C.foreign_tax_credit(year, base, base_includes_grossup, method, countries, months)
    if "error" in r:
        return r
    return _ok(r["결과"], r["근거"], 연도별_근거=r.get("연도별_근거", {}))


@mcp.tool(annotations=CALC)
@_guard
def entertainment_limit(year: YEAR, sme: Annotated[bool, Field(description="중소기업 여부")],
                        general_revenue: WON,
                        expensed: WON,
                        related_revenue: WON = 0,
                        culture: WON = 0,
                        months: Annotated[int, Field(description="사업연도 월수(1~12)")] = 12,
                        market: Annotated[int, Field(description="전통시장(및 조특법 제136조⑥의 대상)에서 지출한 기업업무추진비(원, 2023 사업연도부터) — 일반한도의 일정 비율 안에서 추가한도", ge=0)] = 0) -> dict:
    """기업업무추진비(접대비) 한도와 한도초과액. 언제: 법인세 세무조정에서 기업업무추진비 손금불산입액을 구할 때.
입력: 수입금액(일반, 중단사업부문 매출 포함 기업회계 매출액) → general_revenue, 특수관계인 수입 → related_revenue, 중소기업 여부 → sme, 지출액(접대비) → expensed, 문화비 → culture, 전통시장 지출 → market(2023~), 사업연도 월수 → months.
결과: 한도·한도초과액(손금불산입)·근거(법인세법 제25조, 시행령 제42조). 기업업무추진비 한도와 한도초과액 계산; 접대비가 한도를 넘는가 확인할 때 쓴다."""
    return _ok(C.entertain(year, sme, general_revenue, related_revenue, expensed, culture=culture, months=months, market=market),
               "법인세법 제25조, 시행령 제42조")


@mcp.tool(annotations=CALC)
@_guard
def deemed_interest(movements: Annotated[list[Movement], Field(description="가지급금 증감 목록 [{date:'YYYY-MM-DD', amount:원(대여 +, 회수 -)}]")],
                    year_end: Annotated[str, Field(description="사업연도 종료일(YYYY-MM-DD)")],
                    market_rate: Annotated[float, Field(description="시가 이자율(소수, 예: 0.046 = 4.6%). 가중평균차입이자율 원칙, 예외 시 당좌대출이자율 0.046")] = 0.046,
                    charged_rate: Annotated[float, Field(description="실제 받은 이자율(소수, 예: 0.0 = 무수익)")] = 0.0,
                    received_interest: Annotated[int | None, Field(description="실제 받은 이자 금액(원, ge=0). 주면 charged_rate 계산 대신 이 금액을 수령이자로 사용")] = None,
                    fiscal_start: Annotated[str | None, Field(description="사업연도 개시일(YYYY-MM-DD). 생략 시 year_end의 1년 전 다음 날")] = None) -> dict:
    """가지급금 인정이자(적수 × 시가 이자율)와 부당행위계산 여부. 언제: 특수관계인 가지급금 이자와 부당행위 3억·5% 판정을 함께 구할 때.
입력: 가지급금 증감 [{date, amount(원, 대여 +, 회수 -)}] → movements, 사업연도 종료일(YYYY-MM-DD) → year_end, 시가 이자율(소수) → market_rate, 실제 이자율(소수) → charged_rate, 실제 받은 이자 금액(원) → received_interest, 사업연도 개시일(YYYY-MM-DD, 생략 시 year_end 기준 자동 계산) → fiscal_start.
결과: 적수·시가이자·수령이자·차액·익금산입액·부당행위 여부, 근거(법인세법 제52조, 시행령 제88조③·제89조③, 시행규칙 제43조)."""
    from datetime import date as _date, timedelta
    if received_interest is not None and charged_rate != 0.0:
        return {"error": "입력 확인 필요 — 계산하지 않음",
                "문제": ["received_interest와 charged_rate를 함께 입력할 수 없음 — 둘 중 하나만 사용"]}
    mv = [(m.date, int(m.amount)) for m in movements]
    y_end = _date.fromisoformat(year_end)
    # fiscal_start 없으면 year_end의 1년 전 다음 날
    if fiscal_start is None:
        try:
            fs_dt = y_end.replace(year=y_end.year - 1) + timedelta(days=1)
        except ValueError:
            fs_dt = y_end.replace(year=y_end.year - 1, day=28) + timedelta(days=1)
        fiscal_start = fs_dt.isoformat()
    fs = _date.fromisoformat(fiscal_start)
    days = 366 if y_end.year % 4 == 0 else 365
    # 개시일 전 증감은 모두 개시일 잔액으로 합산
    opening = 0
    fy_movements = []
    for d_str, amt in mv:
        d = _date.fromisoformat(d_str)
        if d < fs:
            opening += amt
        else:
            fy_movements.append((d_str, amt))
    if opening != 0:
        fy_movements.insert(0, (fiscal_start, opening))
    j = C.jeoksu(fy_movements, year_end)
    if received_interest is not None:
        market = int(j * market_rate / days)
        charged = received_interest
        diff = market - charged
        hit = diff >= 300_000_000 or (market and diff >= market * 0.05)
        return _ok({"적수": j,
                    "시가이자": market,
                    "수령이자": charged,
                    "차액": diff if diff > 0 else 0,
                    "부당행위": bool(hit),
                    "익금산입": diff if hit else 0},
                   "법인세법 제52조, 시행령 제88조③·제89조③, 시행규칙 제43조(이자율)")
    return _ok({"적수": j, **C.deemed_interest_check(market_rate, charged_rate, j, days)},
               "법인세법 제52조, 시행령 제88조③·제89조③, 시행규칙 제43조(이자율)")


@mcp.tool(annotations=CALC)
@_guard
def unfair_transaction(market_price: Annotated[float, Field(description="시가(단가, 원)")], actual_price: Annotated[float, Field(description="실제 거래 단가(원)")],
                       quantity: Annotated[int, Field(description="수량")]) -> dict:
    """부당행위계산 기준(시가와 실제 거래가 차이 3억 원 이상 또는 5% 이상). 언제: 특수관계인 거래 prices가 시가와 크게 차이 나 부당행위계산 부인 대상인지 확인할 때.
입력: 시가(단가, 원) → market_price, 실제 거래 단가(원) → actual_price, 수량 → quantity.
결과: 시가총액·차액·차액비율·부당행위계산 여부, 근거(법인세법 제52조, 시행령 제88조③). 부당행위계산 기준(시가와 실제 거래가 차액이 3억 원 이상 또는 5% 이상이면 부인 대상)."""
    gap = abs(market_price - actual_price) * quantity; base = market_price * quantity
    return _ok({"시가총액": int(base), "차액": int(gap), "차액비율": round(gap / base, 4) if base else None,
                "부당행위계산": bool(gap >= 300_000_000 or (base and gap >= base * 0.05))}, "법인세법 제52조, 시행령 제88조③")


@mcp.tool(annotations=CALC)
@_guard
def loss_carryforward_limit(year: YEAR, sme: Annotated[bool, Field(description="중소기업 등 여부")]) -> dict:
    """이월결손금 공제 한도율(사업연도 개시일 기준). 언제: 이월결손금이 있을 때 해당 사업연도에 얼마까지 공제되는지 한도율을 확인한다.
입력: 사업연도(2016~2026) → year, 중소기업 등 여부 → sme.
결과: 한도율(소수)과 근거(법인세법 제13조① 단서). Limit on using loss carryforwards in a year; returns the deduction limit ratio by fiscal year start."""
    return _ok({"한도율": C.loss_limit_ratio(year, sme)}, "법인세법 제13조① 단서")


@mcp.tool(annotations=CALC)
@_guard
def minimum_tax(base_before_incentives: WON, year: YEAR,
                sme: Annotated[bool, Field(description="중소기업 여부")] = True,
                grace_year: Annotated[int, Field(description="중소기업 졸업 후 경과 연차(0=해당 없음, 0~5)")] = 0) -> dict:
    """최저한세(감면 전 과세표준 × 최저한세율). 언제: 세액감면·공제를 받은 뒤에도 납부해야 할 최저한세를 확인할 때(중소기업 졸업 유예 연차 포함).
입력: 감면 전 과세표준(원) → base_before_incentives, 사업연도(2016~2026) → year, 중소기업 여부 → sme, 졸업 후 경과 연차(0~5) → grace_year.
결과: 최저한세액, 근거(조세특례제한법 제132조①). Minimum corporate tax = base before incentives × minimum tax rate."""
    return _ok({"최저한세": C.min_tax(base_before_incentives, sme=sme, grace_year=grace_year, year=year)}, "조세특례제한법 제132조①")


@mcp.tool(annotations=CALC)
def underreporting_penalty(additional_tax: WON, fraud: Annotated[bool, Field(description="부정행위로 인한 과소신고 여부")]) -> dict:
    """과소신고가산세(일반 과소신고 10%, 부정행위 40%). 언제: 신고한 세액보다 실제 세액이 더 클 때 가산세 금액을 구한다.
입력: 과소신고한 세액(원) → additional_tax, 부정행위 여부 → fraud.
결과: 가산세·적용 세율, 근거(국세기본법 제47조의3). Under-reporting penalty: 10% general, 40% for fraud."""
    return _ok({"가산세": K.penalty_underreport(additional_tax, fraud), "세율": 0.4 if fraud else 0.1}, "국세기본법 제47조의3")


@mcp.tool(annotations=CALC)
@_guard
def late_payment_penalty(unpaid_tax: WON, due: Annotated[str, Field(description="법정납부기한(YYYY-MM-DD)")],
                         until: Annotated[str, Field(description="납부일·고지일(YYYY-MM-DD)")]) -> dict:
    """납부지연가산세(미납 세액 × 일수 × 이율, 이율 변경일 기준 기간 안분). 언제: 세금을 법정납부기한보다 늦게 납부했을 때 일수분 가산세를 계산한다.
입력: 미납 세액(원) → unpaid_tax, 법정납부기한(YYYY-MM-DD) → due, 납부일·고지일(YYYY-MM-DD) → until.
결과: 가산세 총액, 근거(국세기본법 제47조의4, 시행령 제27조의4). Late payment penalty with daily rate split at rate-change dates."""
    return _ok(C.late_payment_penalty(unpaid_tax, due, until), "국세기본법 제47조의4, 시행령 제27조의4")


@mcp.tool(annotations=CALC)
@_guard
def invoice_penalty(supply_amount: WON, kind: Annotated[Literal["가공", "지연발급", "미발급", "미발급_종이", "지연전송", "미전송", "지연수취",
                                                               "매출처별합계표_미제출", "매출처별합계표_예정분확정제출", "신용카드매입공제"],
                                                       Field(description="가산세 유형 — 가공=실물 없는 발급·수취")],
                    supplied_on: Annotated[str, Field(description="공급일(발급·수취일) YYYY-MM-DD — 세율 기준 연도")] = "") -> dict:
    """세금계산서 가산세(가공 2%·3%·4%, 미발급 2%, 지연발급 1% 등, 공급일 기준 세율 적용). 언제: 세금계산서 발급·수취 의무를 위반했을 때 가산세 금액을 구한다.
입력: 공급가액(원) → supply_amount, 가산세 유형 → kind, 공급일(YYYY-MM-DD, 세율 기준) → supplied_on.
결과: 가산세·적용 세율, 근거(부가가치세법 제60조②⑤⑥⑦, 가공은 제60조③). VAT invoice penalties by type and supply date."""
    if kind == "가공":
        return _ok({"가산세": K.penalty_fake_invoice(supply_amount, supplied_on), "세율": K.fake_invoice_rate(supplied_on)}, "부가가치세법 제60조③")
    return _ok({"가산세": V.invoice_penalties(supply_amount, kind, supplied_on or None), "세율": V.invoice_penalty_rate(kind, supplied_on or None)}, "부가가치세법 제60조②⑤⑥⑦")


@mcp.tool(annotations=CALC)
def assessment_limitation(due_date: Annotated[str, Field(description="법정신고기한(YYYY-MM-DD, 부과할 수 있는 날 = 그 다음 날)")],
                          case: Annotated[Literal["일반", "무신고", "부정행위"], Field(description="일반(과소신고 등) / 무신고 / 부정행위로 포탈·환급")] = "일반",
                          offshore: Annotated[bool, Field(description="역외거래(국제거래 또는 국외 자산·용역 관련 거래) 여부")] = False,
                          inheritance_gift: Annotated[bool, Field(description="상속세·증여세 여부")] = False) -> dict:
    """국세 부과제척기간과 만료일(2020년 이후 부과분 기준 현행 조문). 언제: 어느 과세기간의 세금이든 부과·제척기간이 언제까지인지 역산한다.
입력: 법정신고기한(YYYY-MM-DD) → due_date, 부과 성격(일반·무신고·부정행위) → case, 역외거래 여부 → offshore, 상속세·증여세 여부 → inheritance_gift.
결과: 제척기간(년)·기산일·만료일, 근거(국세기본법 제26조의2①②④). 일반 5년(역외 7년), 무신고 7년(역외 10년), 부정행위 10년(역외 15년), 상속·증여 10년(부정 등 15년)."""
    if inheritance_gift: n = 10 if case == "일반" else 15
    else: n = {"일반": (5, 7), "무신고": (7, 10), "부정행위": (10, 15)}[case][1 if offshore else 0]
    d = date.fromisoformat(due_date)
    try: end = d.replace(year=d.year + n)
    except ValueError: end = d.replace(year=d.year + n, day=28)
    return _ok({"제척기간(년)": n, "기산일": "법정신고기한 다음 날", "만료일": end.isoformat()},
               "국세기본법 제26조의2" + ("④" if inheritance_gift else ("②" if case != "일반" else "①")),
               주의="2019.12.31. 개정 전 기간·특례(이월결손금 공제 시 ③, 상증세 안 날부터 1년 ⑤, 쟁송 결과 ⑥)는 별도 확인 — 원문은 korean-tax-mcp law_article")


@mcp.tool(annotations=CALC)
@_guard
def income_tax(base: Annotated[int, Field(description="종합소득 과세표준(원)")], year: YEAR) -> dict:
    """종합소득세 산출세액(귀속연도 세율표). 언제: 종합소득 과세표준이 정해지면 그에 대한 산출세액을 구한다(근로·사업·기타·이자·배당 등 종합과세 대상).
입력: 종합소득 과세표준(원) → base, 귀속연도(2016~2026) → year.
결과: 산출세액, 근거(소득세법 제55조). Comprehensive income tax on a tax base (year-specific rate table)."""
    return _ok({"산출세액": I.income_tax(base, year)}, "소득세법 제55조")


@mcp.tool(annotations=CALC)
@_guard
def withholding_tax(kind: Annotated[str, Field(description="소득 종류: 이자·배당·비영업대금·사업·기타·봉사료·일용·비실명·외국인직업운동가·온투업이자")],
                    amount: WON, year: Annotated[int, Field(description="지급연도(2016~2026)")] = 2025) -> dict:
    """국내 원천징수세액(소득세분). 언제: 이자·배당·사업·기타소득 등을 지급할 때 원천징수할 소득세액을 구한다.
입력: 소득 종류(이자·배당·비영업대금·사업·기타·봉사료·일용·비실명·외국인직업운동가·온투업이자) → kind, 지급액(원) → amount, 지급연도(2016~2026) → year.
결과: 원천징수세액·적용 세율, 근거(소득세법 제129조). 지방소득세 10% 별도. 비거주자·외국법인은 nonresident_withholding 사용. 국내 withholding tax (소득세분) on interest, dividends, business income, etc."""
    return _ok({"원천징수세액": I.withholding_tax(kind, amount, year), "세율": I.withholding_rate(kind, year)}, "소득세법 제129조",
               주의="지방소득세(소득세의 10%) 별도. 비거주자는 조세조약 제한세율 확인(korean-tax-mcp treaty_withholding_rates)")


@mcp.tool(annotations=CALC)
@_guard
def nonresident_withholding(income_kind: Annotated[Literal["이자", "배당", "사용료", "인적용역", "기타"],
                                                   Field(description="소득 종류(비거주자·외국법인 국내원천) — 이자·배당·사용료·인적용역·기타")],
                            amount: WON,
                            recipient_type: Annotated[Literal["개인", "법인"], Field(description="수취인 구분 — 개인(비거주자) / 법인(외국법인)")],
                            residence_country: Annotated[str, Field(description="거주지국(ISO 2자 코드 또는 국명)")],
                            treaty_rate: Annotated[float | None, Field(description="조세조약 제한세율(소수, 예: 0.15 = 15%). 생략 시 국내세율 적용")] = None,
                            bond_interest: Annotated[bool, Field(description="채권 이자 여부 — 이자 소득에서만 의미, 특례 적용 필요 시 별도 확인")] = False,
                            related_party: Annotated[bool, Field(description="수취인이 국외지배주주일 가능성 — 이자 소득에서 과소자본 쟁점 표시 시 사용")] = False) -> dict:
    """비거주자·외국법인 국내원천소득 원천징수(소득세법 제156조 / 법인세법 제98조). 언제: 비거주자·외국법인에 이자·배당·사용료·인적용역·기타 소득을 지급할 때 원천징수세액을 구한다.
입력: 소득 종류(이자·배당·사용료·인적용역·기타) → income_kind, 지급액(원) → amount, 수취인 구분(개인/법인) → recipient_type, 거주지국(ISO 2자 또는 국명) → residence_country, 조세조약 제한세율(소수) → treaty_rate, 채권 이자 여부 → bond_interest, 국외지배주주 가능성 → related_party.
결과: 원천징수세액·지방소득세·적용 세율·국내세율, 근거(소득세법 제156조 / 법인세법 제98조). 적용세율 = min(국내세율, 조약 제한세율), 지방소득세 10% 별도. Non-resident/foreign corporation withholding tax on Korean-source income."""
    r = I.nonresident_withholding(income_kind, amount, recipient_type, residence_country, treaty_rate, bond_interest, related_party)
    return _ok({k: v for k, v in r.items() if k != "쟁점"}, r["근거"], 쟁점=r["쟁점"])


@mcp.tool(annotations=CALC)
@_guard
def retirement_income_tax(income: Annotated[int, Field(description="퇴직소득금액(퇴직급여, 비과세 제외, 원)")],
                          years_of_service: Annotated[int, Field(description="근속연수(1~60년)")],
                          retire_year: Annotated[int, Field(description="퇴직한 해(2020~2026). 근속연수공제표(2023 개정)·기본세율표가 이 해 기준", ge=2016, le=2026)] = 2025) -> dict:
    """퇴직소득세(근속연수공제·환산급여 방식). 언제: 퇴직금(퇴직급여)을 지급할 때 퇴직소득세를 계산한다.
입력: 퇴직소득금액(퇴직급여, 비과세 제외, 원) → income, 근속연수(1~60년) → years_of_service, 퇴직한 해 → retire_year(기본 2025).
연도: 근속연수공제는 2023.1.1. 이후 퇴직 개정표(100만·200만·250만·300만), 2022년까지는 구표(30만·50만·80만·120만). 기본세율은 퇴직한 해 표.
지원하지 않음: 2016~2019년 퇴직(종전·개정 방식 세액을 비율로 섞는 경과 부칙, 법률 제12852호).
결과: 퇴직소득세, 근거(소득세법 제48조·제55조②). Retirement income tax (근속연수공제·환산급여 방식)."""
    if retire_year < 2020:
        return {"error": "입력 확인 필요 — 계산하지 않음",
                "문제": [f"retire_year={retire_year}: 2016~2019년 퇴직은 종전 방식과 개정 방식 세액을 비율로 섞는 경과 부칙(법률 제12852호) 대상 — 이 도구는 지원하지 않음"]}
    return _ok({"퇴직소득세": I.retirement_tax(income, years_of_service, retire_year),
                "근속연수공제": min(I.retirement_service_deduction(years_of_service, retire_year), income)},
               "소득세법 제48조·제55조②" + (" (2022년 이전 퇴직 — 근속연수공제 구표)" if retire_year <= 2022 else ""))


@mcp.tool(annotations=CALC)
@_guard
def vat_deemed_rent(deposit: WON, days: Annotated[int, Field(description="과세기간 중 임대 일수(1~184)")],
                    period: Annotated[str, Field(description="과세기간('YYYY-1' 또는 'YYYY-2', 정기예금이자율 연도표 기준)")],
                    rate: Annotated[float | None, Field(description="적용 이자율(소수, 예 0.015 = 1.5%). 생략 시 연도표")]=None) -> dict:
    """간주임대료(보증금 × 정기예금이자율 × 일수/연일수, 윤년이면 연수 366). 언제: 상가·오피스텔 등 임대 보증금을 받을 때 부가가치세 간주임대료를 계산한다.
입력: 보증금(원) → deposit, 과세기간 중 임대 일수(1~184) → days, 과세기간('YYYY-1' 또는 'YYYY-2') → period, 적용 이자율(소수, 생략 시 연도표) → rate.
결과: 간주임대료·적용_이자율, 근거(부가가치세법 시행령 제65조). VAT deemed rent on lease deposits."""
    결과 = V.deemed_rent(deposit, days, period=period, rate=rate)
    적용이자율 = rate if rate is not None else V.DEEMED_RENT_RATE.get(V._period(period)[0], 0.031) if period else 0.031
    return _ok({"간주임대료": 결과, "적용_이자율": 적용이자율}, "부가가치세법 시행령 제65조")


@mcp.tool(annotations=CALC)
@_guard
def vat_input_tax_deduction(period: Annotated[str, Field(description="과세기간('YYYY-1' 또는 'YYYY-2')")],
                            filing: Annotated[Literal["예정", "확정"], Field(description="예정신고 또는 확정신고")],
                            items: Annotated[list, Field(description="[{항목, 증빙, 공급가액?|공급대가?, 성격, 공급자?, 공급시기?, 발급일?, 실제공급가액?}]")]) -> dict:
    """매입세액 공제 판정(부가가치세법 제38조·제39조·제46조, 시행령 제75조). 언제: 부가가치세 신고에서 매입세액 중 공제·불공제 금액을 항목별로 판정할 때.
    입력: period=과세기간('YYYY-1'|'YYYY-2'), filing='예정'|'확정', items=[{항목, 증빙, 공급가액?|공급대가?, 성격, 공급자?, 공급시기?, 발급일?, 실제공급가액?}].
      증빙: '세금계산서'|'신용카드'|'현금영수증'|'영수증'.
      성격: '일반'|'접대'|'비영업용승용차'|'토지관련'|'면세사업'|'사업무관'.
      공급자: '일반과세자'(기본)|'간이과세자_세금계산서'|'간이과세자_영수증'.
      공급가액·공급대가 중 하나 이상 필요. 공급가액만 있으면 ×10%, 공급대가만 있으면 ×10/110.
      실제공급가액이 있으면 필요적 기재사항 착오로 실제분 공제(영 제75조2호).
    결과: {결과: {공제매입세액, 불공제매입세액}, 계산: {항목별: [{항목, 세액, 공제세액, 불공제사유?}], 판단: [{항목, 결론, 근거}]}}.
    판단 순서: ①성격별 불공제(접대·비영업용승용차·토지관련·면세사업·사업무관) → ②증빙 '영수증' → ③신용카드·현금영수증 + 간이과세자_영수증 → ④세금계산서 발급시기(지연수취 가산세) → ⑤실제공급가액(착오 기재)."""
    r = V.vat_input_tax_deduction(period, filing, items)
    if "error" in r:
        return r
    calc = r.pop("계산", None)
    판단 = r.pop("판단", None)
    if calc and "판단" in calc:
        calc = {k: v for k, v in calc.items() if k != "판단"}
    return _ok(r, "부가가치세법 제38조·제39조·제46조, 시행령 제75조", 계산=calc, 판단=판단)


@mcp.tool(annotations=CALC)
@_guard
def vat_deemed_supply_items(period: Annotated[str, Field(description="과세기간('YYYY-1' 또는 'YYYY-2')")],
                            items: Annotated[list, Field(description="[{항목, 거래, 매입세액공제(bool), 금액?|시가?|감가상각?, 1인당_연간금액?, 견본_광고선전물?}]")]) -> dict:
    """간주공급·판매 과세표준 판정(부가가치세법 제10조, 시행령 제19조의2·제20조·제66조). 언제: 재화의 판매·간주공급에 대한 부가가치세 과세표준을 항목별로 판정할 때.
    입력: period=과세기간('YYYY-1'|'YYYY-2'), items=[{항목, 거래, 매입세액공제(bool), 금액?|시가?|감가상각?, 1인당_연간금액?, 견본_광고선전물?}].
      거래: '판매'|'면세전용'|'비영업용승용차전용'|'타사업장반출'|'개인적공급'|'사업상증여'|'폐업시잔존'.
      감가상각: {구분:'건물'|'구축물'|'기타', 취득가액, 취득일, 공급일} → 간주시가 = 취득가액 × (1 − 체감률 × 경과과세기간수), 체감률 건물·구축물 5%·기타 25%.
      개인적공급: 1인당 연간 10만원 이하분은 공급 아님(영 제19조의2), 초과액만 과세.
      사업상증여: 견본_광고선전물이 True면 0(법 제10조⑤ 단서).
    결과: {결과: {과세표준}, 계산: {항목별: [{항목, 거래, 과세표준, 제외사유?}], 판단: [{항목, 결론, 근거}]}}.
    규칙: '판매'·'타사업장반출'은 금액 그대로 과세(매입세액 공제 여부와 무관). 나머지 간주공급은 매입세액 공제 여부가 False면 0, True면 시가(또는 감가상각 간주시가) 과세."""
    r = V.vat_deemed_supply_items(period, items)
    if "error" in r:
        return r
    calc = r.pop("계산", None)
    판단 = r.pop("판단", None)
    if calc and "판단" in calc:
        calc = {k: v for k, v in calc.items() if k != "판단"}
    return _ok(r, "부가가치세법 제10조, 시행령 제19조의2·제20조·제66조", 계산=calc, 판단=판단)


# ───────── 0.2 추가: 법인세 세무조정 ─────────
@mcp.tool(annotations=CALC)
@_guard
def nonbusiness_interest_disallowance(loans: Annotated[list[Loan], Field(description="차입금별 [{interest: 지급이자(원), jeoksu?: 차입금 적수(원·일) | principal·days | rate, kind?: 일반|채권자불분명|비실명|건설자금|기업구매자금|금융회사제외차입금, withholding?: 원천징수세액}]")],
                                      nonbusiness_jeoksu: Annotated[int, Field(description="업무무관자산(영 제49조①) 적수 + 직접 계산한 가지급금 적수 합(원·일). advances로 가지급금을 넣으면 그 적수는 여기서 빼고 넣는다", ge=0)] = 0,
                                      advances: Annotated[list[dict] | None, Field(description="가지급금 등 [{대상: 대표이사|대표자|임원|직원|주주|특수관계법인|기타특수관계인|비특수관계인, 이름?(동일인 상계용), 사유: 업무무관대여|미지급소득세대납|국외투자법인비용대납|우리사주취득자금|퇴직금전환금|대표자상여소득세대납|급료가불|경조사비대여|학자금대여|주택자금대여|자산관리공사출자법인대여, 금액, 지급일, 회수일? | movements [{일자, 증감}], 총결정세액·미지급소득·종합소득금액?(미지급소득세대납), 월정급여?(급료가불), 해당법인주식?(우리사주), 중소기업·지배주주등?(주택자금)}]")] = None,
                                      deposits: Annotated[list[dict] | None, Field(description="가수금 [{이름|대상, 금액, 일자, 반환일? | movements [{일자, 증감}], 상계불가?(상환기간·이자율 약정으로 상계 불가)}]")] = None,
                                      year: Annotated[int | None, Field(description="사업연도(개시일이 속하는 연도, 2016~2026). advances·deposits를 넣으면 필요", ge=2016, le=2026)] = None,
                                      fiscal_start: Annotated[str | None, Field(description="사업연도 개시일(YYYY-MM-DD). 생략 시 year-01-01")] = None,
                                      fiscal_end: Annotated[str | None, Field(description="사업연도 종료일(YYYY-MM-DD). 생략 시 개시일 1년 뒤 전날")] = None,
                                      company: Annotated[str | None, Field(description="가지급금을 지급한 법인 이름(relations 안의 법인). advances[].상대방으로 특수관계인 판정할 때 필요")] = None,
                                      relations: Annotated[dict | None, Field(description='특수관계 판정 자료 — related_judge와 같은 형식 {people, family, stakes, officers?, employees?, livelihood?, livelihood_shared?, influence?, group?, group_exempt?}. 예: {"people": {"㈜가람": {"구분": "법인"}, "김대표": {"구분": "개인"}}, "family": [], "stakes": [["김대표", "㈜가람", 0.6]], "officers": {"㈜가람": ["김대표"]}}')] = None) -> dict:
    """업무무관자산·가지급금 관련 지급이자 손금불산입 — 가지급금 제외 판정·적수·가수금 상계·지급이자 손금불산입 순위까지. 언제: 업무무관자산이나 특수관계인 가지급금이 있을 때 지급이자 중 얼마를 손금불산입하는지, 어떤 대여금이 가지급금에서 빠지는지(시행규칙 제44조 각 호) 판정한다.
입력: 차입금별 [{interest, jeoksu | principal·days | rate, kind?, withholding?}] → loans, 업무무관자산 적수(원·일) → nonbusiness_jeoksu, 가지급금 [{대상 | 상대방, 사유, 금액, 지급일, 회수일? | movements, 기준일?, 관련 정보}] → advances, 법인 이름 → company, 관계 자료(related_judge 형식) → relations, 가수금 [{이름, 금액, 일자, 반환일?, 상계불가?}] → deposits, 사업연도(2016~2026) → year, 개시일·종료일 → fiscal_start·fiscal_end.
판정: ⓪특수관계인 — 가지급금에 상대방을 주면 company↔상대방을 법인세법 특수관계인(법 제2조12호, 영 제2조⑧ — 2019.2.12. 전 영 제87조①, 2019.2.12.~2023.12.31. 영 제2조⑤) 각 호로 판정해(기준일: 기준일 또는 사업연도 종료일) 대상을 정한다 — 해당 없으면 비특수관계인, 법인이면 특수관계법인, 개인이면 officers·employees·stakes로 임원·직원·주주, 그 밖은 기타특수관계인. 임원·직원의 배우자 등이 지분 없이 상대방이면 relations.livelihood_shared의 생계 사실관계로 ⑧3호(생계를 함께하는 친족)를 판정 — 사실이 모자라 입력 확인 필요가 나오면 비특수관계인으로 넘기지 않고 확인할 항목·필요 자료를 돌려준다. 대상도 주면 대조만 하고 특수관계 여부(또는 규칙 제44조 6·7·7의2호에서 직원 여부)가 다르면 입력 확인 필요. ①차입금 구분 — 채권자불분명(영 제51조①)·비실명(영 제51조②)·건설자금(영 제52조)은 영 제55조 순위로 먼저 손금불산입하고, 기업구매자금대출·금융회사 등 제외 차입금(영 제53조④)과 함께 업무무관 산식의 지급이자·차입금에서 뺀다. ②가지급금 — 비특수관계인 대여는 제외, 규칙 제44조 각 호(미지급소득 소득세 대납은 총결정세액×미지급소득÷종합소득금액 한도, 직원 급료 가불은 월정급여 범위, 우리사주는 해당 회사 주식만, 중소기업 직원 주택자금 7의2호는 2020.1.1. 이후 개시 사업연도부터) 제외. ③적수 — 대여일 산입·회수일 불산입·말일 산입, 동일인 가수금 적수 상계(영 제53조③, 약정으로 상계 불가면 안 함 — 규칙 제28조②). ④업무무관 = 지급이자 × min(자산 적수, 차입금 적수) ÷ 차입금 적수(원 미만 절사).
결과: 순위별 손금불산입·소득처분(채권자불분명·비실명 대표자 상여·원천세 기타사외유출, 건설자금 유보, 업무무관 기타사외유출), 가지급금별 제외 판단, 동일인별 적수, 인정이자는 deemed_interest로 이어 계산할 movements 안내. 지원 2016~2026 사업연도.
지원 안 함: 일반차입금 건설자금이자 자본화 선택(법 제28조②), 업무무관자산 취득가액·적수 산정(영 제49조, 직접 적수 입력), 상대방 없이 대상만 준 가지급금의 특수관계 여부(입력 그대로 씀), 비소액주주등의 직원과 생계를 함께하는 친족(영 제2조⑧3호가목 중 주주의 직원)·대표이사 구분(officers에 대표 표시 없음 — 임원으로 봄), 금융회사 주된 수익사업 판단 — 해당하면 입력 확인 필요 또는 직접 적수 입력. 근거(법인세법 제28조, 시행령 제51조~제53조·제55조·제106조, 시행규칙 제28조·제44조)."""
    simple = not (advances or deposits or year or fiscal_start or fiscal_end) and all(
        l.jeoksu is not None and l.kind is None and l.rate is None and l.principal is None and l.withholding is None for l in loans)
    if simple:
        r = C.nonbusiness_interest([(int(l.interest), int(l.jeoksu)) for l in loans])(nonbusiness_jeoksu)
        return _ok(r, "법인세법 제28조①4호, 시행령 제53조", 산식="지급이자 × min(1, 업무무관 적수 ÷ 차입금 적수)")
    try:
        case = None
        if relations is not None:
            from .audit.agents.related_intake import normalize, to_case
            rj = dict(relations)
            case = to_case(normalize(rj))
            case["group_exempt"] = {norm_name(k): v for k, v in (rj.get("group_exempt") or {}).items()}
        r, calc, 판단, pointer = C.nonbusiness_interest_judge([l.model_dump() for l in loans], advances, deposits,
                                                            nonbusiness_jeoksu, year, fiscal_start, fiscal_end,
                                                            company=norm_name(company) if company else None, relations=case)
    except (ValueError, KeyError, TypeError) as e:
        return {"error": "입력 확인 필요 — 계산하지 않음", "문제": [str(e)]}
    extra = {"계산": calc, "판단": 판단}
    if pointer:
        extra["인정이자_연결"] = pointer
    return _ok(r, "법인세법 제28조①·" + ("⑤" if (r and calc["사업연도"][:4] <= "2018") else "③") +
               ", 시행령 제51조·제53조·제55조·제106조, 시행규칙 제28조·제44조" +
               ("; 특수관계인: " + ", ".join(dict.fromkeys(x["특수관계_판정"]["근거"] for x in calc["가지급금"] if "특수관계_판정" in x))
                if any("특수관계_판정" in x for x in calc["가지급금"]) else ""),
               산식="순위(영 제55조) 채권자불분명 → 비실명 → 건설자금 → 지급이자 × min(업무무관 적수, 차입금 적수) ÷ 차입금 적수", **extra)


@mcp.tool(annotations=CALC)
@_guard
def business_car_expense(year: YEAR, upkeep: Annotated[int, Field(description="유지비(유류·보험·수선 등, 원)")],
                         depreciation: Annotated[int, Field(description="세법상 감가상각비(5년 정액, 원)")],
                         booked_depreciation: Annotated[int, Field(description="장부상 감가상각비(원)")],
                         log_kept: Annotated[bool, Field(description="운행기록부 작성 여부")],
                         business_ratio: Annotated[float | None, Field(description="운행기록상 업무사용비율(소수, 작성 시, 0~1)")] = None,
                         months: Annotated[int, Field(description="보유 월수(1~12)")] = 12,
                         rental_corp: Annotated[bool, Field(description="부동산임대업 주업 등 특정법인")] = False,
                         plate_ok: Annotated[bool, Field(description="법인업무용 자동차등록번호판 부착 여부(2024.1.1. 이후 취득·임차 차량 중 취득가액(임차는 기준가액) 8천만원 이상인 차량만 대상 — acquisition_cost를 주면 기준 미만 차량엔 적용하지 않음)")] = True,
                         acquisition_cost: Annotated[int | None, Field(description="차량 취득가액(원, 임차는 기준가액). 8천만원 미만이면 번호판 요건 대상이 아니므로 plate_ok=False여도 업무사용비율 그대로", ge=0)] = None) -> dict:
    """업무용승용차 관련비용 손금불산입(업무사용비율·사적사용·미작성 시 기준비율·감가상각 연 800만 원 한도). 언제: 임원·직원이 사용하는 승용차 관련비용을 어디까지 손금으로 인정할지 계산할 때.
입력: 사업연도(2016~2026) → year, 유지비(유류·보험·수선 등, 원) → upkeep, 세법상 감가상각비(5년 정액, 원) → depreciation, 장부상 감가상각비(원) → booked_depreciation, 운행기록부 작성 여부 → log_kept, 업무사용비율(소수, 작성 시) → business_ratio, 보유 월수(1~12) → months, 부동산임대업 주업 등 특정법인 여부 → rental_corp, 법인 전용번호판 부착 여부 → plate_ok.
결과: 업무사용비율·손금불산입·감가상각 한도 등, 근거(법인세법 제27조의2, 시행령 제50조의2)."""
    plate_applies = (acquisition_cost is None) or (acquisition_cost >= 80_000_000)
    r = C.business_car(year, upkeep, depreciation, booked_depreciation, log_kept, business_ratio, months, rental_corp, plate_ok or not plate_applies)
    if not plate_ok and not plate_applies:
        r["번호판 요건"] = f"취득가액 {acquisition_cost:,}원 < 80,000,000원 → 법인업무용 번호판 부착 대상 아님(영 제50조의2④ 단서·규칙 기준) — 미부착이어도 업무사용비율 그대로 적용"
    elif not plate_ok and acquisition_cost is None:
        r["번호판 요건"] = "취득가액을 주지 않아 번호판 요건 대상으로 보고 업무사용금액 0원 적용 — 8천만원 미만 차량이면 acquisition_cost를 넣어 다시 계산"
    return _ok(r, "법인세법 제27조의2, 시행령 제50조의2", 산식="미작성 시 업무사용비율 = min(1, 기준금액 ÷ 총비용), 감가상각 업무사용분 연 800만원 한도")


@mcp.tool(annotations=CALC)
@_guard
def donation_limit(year: YEAR, base_income: WON,
                   special: WON = 0, general: WON = 0,
                   carried_loss: WON = 0,
                   special_carry: WON = 0, general_carry: WON = 0,
                   sme: Annotated[bool, Field(description="중소기업 여부")] = True,
                   social_enterprise: Annotated[bool, Field(description="사회적기업 여부(일반한도 20% 적용)")] = False) -> dict:
    """기부금 손금산입 한도(특례기부금 50%, 일반기부금 10%)와 한도초과·이월. 언제: 기부금을 손금산입할 때 한도와 한도초과액을 구한다.
입력: 기준소득(과세표준 계산 전 소득금액, 원) → base_income, 특례기부금 지출액(원) → special, 일반기부금 지출액(원) → general, 결손금(원) → carried_loss, 특례 이월액(원) → special_carry, 일반 이월액(원) → general_carry, 중소기업 여부 → sme, 사회적기업 여부 → social_enterprise, 사업연도(2016~2026) → year.
결과: 한도·한도초과·이월 손금산입, 근거(법인세법 제24조)."""
    return _ok(C.donation(base_income, carried_loss, special, general, special_carry, general_carry, sme, social_enterprise, year=year),
               "법인세법 제24조", 산식="특례 = (기준소득 − 결손금) × 50%, 일반 = (기준소득 − 결손금 − 특례 손금산입) × 10%(사회적기업 20%)")


@mcp.tool(annotations=CALC)
def bad_debt_allowance(receivables: WON, loss_rate: Annotated[float, Field(description="대손실적률(소수, 예: 0.02 = 2%)", ge=0, le=1)],
                       set_amount: WON, prior_disallowed: WON = 0) -> dict:
    """대손충당금 한도(채권잔액 × max(1%, 대손실적률))와 한도초과. 언제: 매출채권 등에 대해 대손충당금을 설정할 때 손금산입 한도를 구한다.
입력: 채권잔액(원) → receivables, 대손실적률(소수, 예: 0.02 = 2%) → loss_rate, 설정할 대손충당금(원) → set_amount, 이전 disallowed(원) → prior_disallowed.
결과: 한도·한도초과, 근거(법인세법 제34조, 시행령 제61조②). Bad debt allowance limit."""
    return _ok(C.bad_debt_allowance(receivables, loss_rate, set_amount, prior_disallowed), "법인세법 제34조, 시행령 제61조②")


@mcp.tool(annotations=CALC)
@_guard
def thin_capitalization(equity: WON,
                        borrowings: WON,
                        total_interest: WON,
                        ratio: Annotated[float | int, Field(description="적용 배수(기본적으로 2, 금융업 등 업종 배수는 별도 확인)")] = 2,
                        year: YEAR = 2026,
                        recipient_residence_country: Annotated[str | None, Field(description="수취인 거주지국(ISO 2자 코드 또는 국명). 지정 시 배당 처분에 따른 원천징수 재계산 제공")] = None,
                        interest_treaty_rate: Annotated[float | None, Field(description="이자 조세조약 제한세율(소수, 예: 0.12 = 12%). 생략 시 국내세율 적용")] = None,
                        dividend_treaty_rate: Annotated[float | None, Field(description="배당 조세조약 제한세율(소수, 예: 0.15 = 15%). 생략 시 국내세율 적용")] = None,
                        recipient_type: Annotated[Literal['개인', '법인'], Field(description="수취인 구분 — 개인(비거주자) / 법인(외국법인)")] = '개인') -> dict:
    """과소자본 손금불산입(국제조세조정에 관한 법률 제22조). 언제: 국외지배주주로부터 차입한 금액이 출자금액의 기준 배수(기본 2배)를 초과할 때 초과분 지급이자를 손금불산입한다.
입력: 국외지배주주 출자금액(원) → equity, 차입금액(원) → borrowings, 지급이자 총액(원) → total_interest, 적용 배수(기본 2) → ratio, 사업연도(2016~2026) → year, 수취인 거주지국 → recipient_residence_country, 이자 조세조약 제한세율(소수) → interest_treaty_rate, 배당 조세조약 제한세율(소수) → dividend_treaty_rate, 수취인 구분(개인/법인) → recipient_type.
결과: 초과차입금·손금불산입 이자·처분(배당/기타사외유출), recipient_residence_country 지정 시 원천징수 재계산 결과 포함, 근거(국제조세조정에 관한 법률 제22조, 시행령 제34조). Thin capitalization disallowance."""
    base = C.thin_capitalization(equity, borrowings, total_interest, ratio, year)
    result = {"결과": base, "근거": "국제조세조정에 관한 법률 제22조, 시행령 제34조"}

    # 쟁점 부착 (issues.json에서 읽음, 하드코딩 금지)
    from .engine.calc_income import _load_issues
    result["쟁점"] = _load_issues().get("thin_capitalization", [])
    if recipient_residence_country:
        disallowed_interest = base["손금불산입 이자"]
        if disallowed_interest > 0:
            # 이자 기준 원천징수 (기존 이자소득 처분 가정)
            interest_wh = I.nonresident_withholding("이자", disallowed_interest, recipient_type,
                                                    recipient_residence_country, interest_treaty_rate)
            # 배당 기준 원천징수 (국조법 제22조② 배당 처분)
            dividend_wh = I.nonresident_withholding("배당", disallowed_interest, recipient_type,
                                                    recipient_residence_country, dividend_treaty_rate)
            diff_income = dividend_wh["원천징수세액"] - interest_wh["원천징수세액"]
            diff_local = dividend_wh["지방소득세"] - interest_wh["지방소득세"]
            result["결과"]["원천징수 재계산"] = {
                "대상 금액(손금불산입 이자)": disallowed_interest,
                "이자 기준 원천징수세액": interest_wh["원천징수세액"],
                "이자 기준 지방소득세": interest_wh["지방소득세"],
                "배당 기준 원천징수세액": dividend_wh["원천징수세액"],
                "배당 기준 지방소득세": dividend_wh["지방소득세"],
                "소득세 차액": diff_income,
                "지방소득세 차액": diff_local,
                "근거": "국조법 제22조 ② 배당 처분 → 소득세법 제119조 2호 / 법인세법 제93조 2호 배당소득",
            }
    return result


@mcp.tool(annotations=CALC)
@_guard
def missing_receipt_disallowance(year: YEAR, items: Annotated[list[ReceiptItem], Field(description="[{amount: 1회 지출액(원), qualified: 적격증빙 여부, congratulatory: 경조금 여부}]")]) -> dict:
    """적격증빙 미수취 기업업무추진비 손금불산입(건당 3만 원, 경조금 20만 원 초과 + 적격증빙 없음). 언제: 접대비 등을 지출했는데 적격증빙을 받지 못한 건이 있으면 손금불산입액을 구한다.
입력: 사업연도(2016~2026) → year, 지출 건 목록 [{amount: 1회 지출액(원), qualified: 적격증빙 여부, congratulatory: 경조금 여부}] → items.
결과: 손금불산입 합계액, 근거(법인세법 제25조②, 시행령 제41조①). 적격증빙 미수취 기업업무추진비 손금불산입."""
    return _ok({"손금불산입": C.receipt_disallowed(year, [(int(i.amount), bool(i.qualified), bool(i.congratulatory)) for i in items])},
               "법인세법 제25조②, 시행령 제41조①")


# ───────── 0.2 추가: 부가가치세 ─────────
def _normalize_usage(usage):
    """items.usage 허용 3形態 → {구분: 비율%} dict로 정규화.

    ① {'제조': 60, '재판매': 30}          → 그대로 반환
    ② [{'구분': '제조', '비율': 60}, ...] → dict로 변환
    ③ {'구분': '제조', '비율': 60}         → ②처럼 처리해 dict로 변환
    """
    if not usage:
        return {}
    if isinstance(usage, list):
        out = {}
        for row in usage:
            g = row.get("구분", "")
            b = row.get("비율", 0)
            out[g] = int(b)
        return out
    if isinstance(usage, dict):
        if "구분" in usage and "비율" in usage:
            return {usage["구분"]: int(usage["비율"])}
        return {k: int(v) for k, v in usage.items()}
    raise TypeError(f"usage 유형 불일치: {type(usage)}")


def _merge_items_with_same_identity(items):
    """같은 name이고 purchase_price·수입·관세·운임 값이 같은 여러 행을 한 품목으로 합쳐 usage를 더한다."""
    groups = {}
    for it in items:
        key = (
            it.get("name", ""),
            int(it.get("purchase_price", 0)),
            bool(it.get("imported", False)),
            int(it.get("customs_value", 0)),
            int(it.get("customs_duty", 0)),
            int(it.get("freight_included", 0)),
        )
        groups.setdefault(key, []).append(it)
    merged = []
    for key, rows in groups.items():
        name = key[0]
        base = rows[0]
        merged_usage = {}
        for r in rows:
            u = _normalize_usage(r.get("usage", {}))
            for g, b in u.items():
                merged_usage[g] = merged_usage.get(g, 0) + b
        merged.append({
            "name": name,
            "purchase_price": base.get("purchase_price", 0),
            "imported": base.get("imported", False),
            "customs_value": base.get("customs_value", 0),
            "customs_duty": base.get("customs_duty", 0),
            "freight_included": base.get("freight_included", 0),
            "taxable_goods": base.get("taxable_goods", False),
            "usage": merged_usage,
        })
    return merged


_USAGE_ERROR_FORMAT = (
    "items[].usage 형식 오류 — 허용 형태: "
    "{{'제조': 60, '재판매': 30, '재고': 10}} 또는 "
    "[{{'구분': '제조', '비율': 60}}, ...]. "
    "구분 허용값: 제조, 재판매, 재고, 종업원제공, 면세사업. "
    "비율 합 100."
)


@mcp.tool(annotations=CALC)
@_guard
def vat_deemed_input_credit(period: Annotated[str, Field(description="과세기간('YYYY-1' 또는 'YYYY-2')")],
                            industry: Annotated[Literal["음식점", "유흥", "제조", "제조_떡방앗간등", "기타"], Field(description="업종")],
                            purchase: Annotated[int, Field(description=(
                                "공제대상 매입가액(원) = 과세사업에 사용했거나 사용할 면세농산물 등의 매입가액. "
                                "다른 사업자에게 그대로 판매한 분은 제외, 기말 재고분은 포함. "
                                "수입분은 관세의 과세가격 + 운반비 포함 구입액이면 purchase_total·freight_included 사용. "
                                "품목별로 따로 부르지 말고 합산해서 1회 호출. "
                                "used_taxable·used_exempt·ending_inventory_common이 모두 0이면 purchase(또는 purchase_total−freight_included)가 공제대상 매입가액이 됨. "
                                "items와 purchase/purchase_total을 동시에 주면 문제로 거부됨"
                            ))] = 0,
                            items: Annotated[list | None, Field(description=(
                                "품목별 면세농산물 등 목록. 각 항목: {name, purchase_price, imported?, customs_value?, customs_duty?, freight_included?, taxable_goods?, usage: {구분: 비율%}}. "
                                "usage 허용 형태 3가지 — ① {'제조': 60, '재판매': 30, '재고': 10} ② [{'구분': '제조', '비율': 60}, {'구분': '재판매', '비율': 30}] ③ {'구분': '제조', '비율': 60}(단일). 전부 {'제조': 60, ...} 형태로 자동 정규화됨. "
                                "usage 구분 허용값: 제조, 재판매, 재고, 종업원제공, 면세사업. 같은 name이고 purchase_price·수입·관세·운임이 같으면 여러 행을 한 품목으로 합쳐 usage를 더함. 비율 합이 100이 아니면 문제로 거부. "
                                "문항에 품목별 표가 있으면 items에 그대로 넣고(수입품은 관세의 과세가격·관세, 운임 포함액, 사용 비율), 합산·제외는 도구가 함. "
                                "items와 purchase/purchase_total을 동시에 주면 문제로 거부"
                            ))] = None,
                            base: Annotated[int, Field(description="해당 과세기간 과세표준(원). 없으면 taxable_supply 사용")] = 0,
                            individual: Annotated[bool, Field(description="개인사업자 여부")] = True, sme: Annotated[bool, Field(description="중소기업 여부(제조 법인)")] = True,
                            purchase_total: Annotated[int, Field(description="면세농산물 등 총 구입액( 운반비 포함 가능, 원)")] = 0,
                            freight_included: Annotated[int, Field(description="구입액에 포함된 채취 운반비(원) — 있으면 매입가액에서 제외")] = 0,
                            used_taxable: Annotated[int, Field(description="과세사업 사용·소비액(원)")] = 0,
                            used_exempt: Annotated[int, Field(description="면세사업 사용·판매액(원)")] = 0,
                            ending_inventory_common: Annotated[int, Field(description="실지귀속 불분명 기말재고(원)")] = 0,
                            taxable_supply: Annotated[int, Field(description="해당 기간 과세 공급가액(원)")] = 0,
                            exempt_supply: Annotated[int, Field(description="해당 기간 면세 공급가액(원)")] = 0,
                            filing: Annotated[Literal["확정", "예정"], Field(description=(
                                "신고 구분(기본 '확정'). '예정'이면 공제 한도(과세표준 × 한도율)를 적용하지 않음 — "
                                "한도는 확정신고 때 과세기간 전체로 정산(부가가치세법 제42조②)"))] = "확정") -> dict:
    """의제매입세액공제(면세농산물 등 매입가액 × 공제율, 연장·운반비율 제외·공통분 안분). 언제: 음식점·제조업 등에서 면세농산물 등을 매입할 때 의제매입세액공제를 계산한다.
입력: 과세기간('YYYY-1' 또는 'YYYY-2') → period, 업종 → industry, 면세농산물 등 매입가액(원) → purchase, 해당 과세기간 과세표준(원) → base, 개인사업자 여부 → individual, 중소기업 여부(제조 법인) → sme, 면세농산물 등 총 구입액(원) → purchase_total, 구입액에 포함된 운반비(원) → freight_included, 과세사업 사용·소비액(원) → used_taxable, 면세사업 사용·판매액(원) → used_exempt, 실지귀속 불분명 기말재고(원) → ending_inventory_common, 해당 기간 과세 공급가액(원) → taxable_supply, 해당 기간 면세 공급가액(원) → exempt_supply.
결과: 공제대상 매입가액·의제매입세액·공제율·한도율·과세비율·단계별 계산, 근거(부가가치세법 제42조, 시행령 제81조·제84조)."""
    from .engine.calc_vat import deemed_input, deemed_input_rate, deemed_input_limit_ratio

    # items와 purchase/purchase_total 동시 입력 거부
    if items is not None and (purchase > 0 or purchase_total > 0):
        return {"error": "입력 확인 필요 — 계산하지 않음",
                "문제": ["items와 purchase/purchase_total을 동시에 입력할 수 없음 — 둘 중 하나만 사용"]}

    # items 기반 계산
    품목별_결과 = None
    if items is not None:
        try:
            items = _merge_items_with_same_identity(items)
            품목별_결과 = []
            총_공제대상 = 0
            for item in items:
                name = item.get("name", "")
                purchase_price = int(item.get("purchase_price", 0))
                imported = bool(item.get("imported", False))
                customs_value = int(item.get("customs_value", 0))
                customs_duty = int(item.get("customs_duty", 0))
                freight_included = int(item.get("freight_included", 0))
                taxable_goods = bool(item.get("taxable_goods", False))
                usage = _normalize_usage(item.get("usage", {}))

                # usage 비율 합 검증
                usage_sum = sum(usage.values()) if usage else 0
                if usage_sum != 100:
                    return {"error": "입력 확인 필요 — 계산하지 않음",
                            "문제": [f"품목 '{name}'의 usage 비율 합이 {usage_sum}%로 100이 아님"]}

                # 기준 매입가액 계산
                if imported:
                    기준매입가액 = customs_value  # 관세의 과세가격 (관세 미포함)
                else:
                    기준매입가액 = max(0, purchase_price - freight_included)

                # taxable_goods (세금계산서 수취 과세 물품)면 공제대상 0
                if taxable_goods:
                    공제대상 = 0
                    판단_문구 = "과세 물품은 의제매입세액 대상 아님"
                else:
                    # usage별 포함/제외 판정
                    포함_usage = {"제조", "재고"}
                    제외_usage = {"재판매", "종업원제공", "면세사업"}
                    공제대상 = 0
                    판단_부분 = []
                    for 구분, 비율 in usage.items():
                        비율_int = int(비율)
                        if 구분 in 포함_usage:
                            공제대상 += int(기준매입가액 * 비율_int / 100)
                            판단_부분.append(f"{구분} {비율_int}% 포함")
                        elif 구분 in 제외_usage:
                            판단_부분.append(f"{구분} {비율_int}% 제외")
                        else:
                            return {"error": "입력 확인 필요 — 계산하지 않음",
                                    "문제": [f"품목 '{name}'의 usage 구분 '{구분}'은(는) 허용되지 않음 — 허용: 제조, 재판매, 재고, 종업원제공, 면세사업"]}
                    판단_문구 = ", ".join(판단_부분) if 판단_부분 else ""

                품목별_결과.append({
                    "name": name,
                    "기준매입가액": 기준매입가액,
                    "공제대상": 공제대상,
                    "판단": 판단_문구,
                })
                총_공제대상 += 공제대상
            eligible_purchase = 총_공제대상
        except Exception as e:
            return {"error": "입력 확인 필요 — 계산하지 않음",
                    "문제": [_USAGE_ERROR_FORMAT]}
    else:
        품목별_결과 = None

    # base가 없으면 taxable_supply를 base로 사용 (명세서 검산 예시 호환)
    effective_base = base if base > 0 else taxable_supply

    # 운반비 제외 비율: purchase_total > 0이면 항상 계산 (freight_included 기본 0)
    if purchase_total > 0:
        freight_ratio = freight_included / purchase_total if purchase_total > 0 else 0
        non_freight_factor = 1 - freight_ratio
        used_taxable_adj = int(used_taxable * non_freight_factor)
        used_exempt_adj = int(used_exempt * non_freight_factor)
        inventory_adj = int(ending_inventory_common * non_freight_factor)
    else:
        used_taxable_adj = used_taxable
        used_exempt_adj = used_exempt
        inventory_adj = ending_inventory_common

    # 과세 공급가액 비율 계산
    if taxable_supply > 0 and exempt_supply > 0:
        taxable_ratio = taxable_supply / (taxable_supply + exempt_supply)
    else:
        taxable_ratio = 1.0  # 과세만 있는 경우

    # items가 제공되면 items 기반 계산 결과 사용, 아니면 기존 로직
    if items is None:
        # 공제대상 매입가액 = 과세 사용분 + 공통 재고 × 과세 공급가액 비율
        eligible_purchase = used_taxable_adj + int(inventory_adj * taxable_ratio)

        # 사용처 구분이 모두 0이면 purchase 또는 purchase_total − freight_included
        if used_taxable == 0 and used_exempt == 0 and ending_inventory_common == 0:
            if purchase > 0:
                eligible_purchase = purchase
            else:
                eligible_purchase = max(0, purchase_total - freight_included)

        # used_taxable이 매입가액 초과 시 오류
        매입가액 = purchase if purchase > 0 else max(0, purchase_total - freight_included)
        if used_taxable > 매입가액 and 매입가액 > 0:
            return {"error": "입력 확인 필요 — 계산하지 않음",
                    "문제": [f"used_taxable({used_taxable:,}원) > 매입가액({매입가액:,}원) — used_taxable이 매입가액보다 클 수 없음"]}

    # 기존 방식(purchase, base 사용)의 공제율·한도율
    num, den = deemed_input_rate(industry, period, individual, effective_base, sme)
    # effective_base가 0이면 한도 계산 생략 (한도 무한대)
    if filing == "예정":
        limit_ratio = float("inf")  # 예정신고는 한도 없이 공제, 확정신고 때 과세기간 전체로 정산
    elif effective_base > 0:
        limit_ratio = deemed_input_limit_ratio(period, individual, industry == "음식점", effective_base)
    else:
        limit_ratio = None  # 한도 없음

    # 의제매입세액 계산
    deemed_tax = deemed_input(
        eligible_purchase if eligible_purchase > 0 else purchase,
        effective_base,
        rate_num=num, rate_den=den, limit_ratio=limit_ratio,
        period=period, industry=industry, individual=individual, sme=sme
    )
    if filing == "예정":
        limit_ratio = None  # 출력용 — JSON에 Infinity를 내보내지 않음

    # 운반비율 적용 결과
    freight_exclusion = {
        "총_구입액": purchase_total,
        "운반비": freight_included,
        "운반비_비율": freight_ratio if purchase_total > 0 else 0,
        "운반비_제외_계수": non_freight_factor if purchase_total > 0 else 1,
        "과세_사용분_조정": used_taxable_adj,
        "면세_사용분_조정": used_exempt_adj,
        "공통_재고_조정": inventory_adj,
    } if purchase_total > 0 else {}

    # 과세 공급가액 비율
    allocation = {
        "과세_공급가액": taxable_supply,
        "면세_공급가액": exempt_supply,
        "과세_공급가액_비율": taxable_ratio,
    } if taxable_supply > 0 else {}

    steps = [
        {"항목": "총 구입액(운반비 포함)", "금액": purchase_total, "근거": "부가가치세법 제42조"},
        {"항목": "채취 운반비", "금액": freight_included, "근거": "부가가치세법 제42조"},
    ]
    if purchase_total > 0:
        steps.extend([
            {"항목": "운반비 제외 비율", "금액": freight_ratio, "근거": "부가가치세법 제42조"},
            {"항목": "운송비 제외 계수 (1 − 운반비 비율)", "금액": non_freight_factor, "근거": "부가가치세법 제42조"},
            {"항목": "과세 사용·소비액(운반비 제외)", "금액": used_taxable_adj, "근거": "부가가치세법 제42조"},
            {"항목": "면세 사용·판매액(운반비 제외)", "금액": used_exempt_adj, "근거": "부가가치세법 제42조"},
            {"항목": "공통 기말재고(실지귀속 불분명, 운반비 제외)", "금액": inventory_adj, "근거": "부가가치세법 제42조"},
        ])
    else:
        steps.append({"항목": "운반비 제외 비율", "금액": 0, "근거": "부가가치세법 제42조"})

    steps.extend([
        {"항목": "과세 공급가액", "금액": taxable_supply, "근거": "부가가치세법 시행령 제81조"},
        {"항목": "면세 공급가액", "금액": exempt_supply, "근거": "부가가치세법 시행령 제81조"},
        {"항목": "과세 공급가액 비율", "금액": taxable_ratio, "근거": "부가가치세법 시행령 제81조"},
    ])

    steps.extend([
        {"항목": "공제대상 매입가액 = 과세 사용분 + 공통 재고 × 과세비율", "금액": eligible_purchase, "근거": "부가가치세법 제42조, 시행령 제84조, 제81조"},
        {"항목": f"공제율 ({num}/{den})", "금액": f"{num}/{den}", "근거": "부가가치세법 제42조"},
        {"항목": f"한도율 ({limit_ratio})" if limit_ratio is not None else ("한도율 (예정신고 — 한도 없이 공제, 확정신고 때 정산)" if filing == "예정" else "한도율 (과세표준 없음 — 한도 미적용)"), "금액": limit_ratio, "근거": "부가가치세법 시행령 제84조"},
        {"항목": "의제매입세액 = min(공제대상 매입, 과세표준 × 한도율) × 공제율", "금액": deemed_tax, "근거": "부가가치세법 제42조"},
    ])

    result = {
        "공제대상_매입가액": eligible_purchase,
        "의제매입세액": deemed_tax,
        "공제율_분자": num,
        "공제율_분모": den,
        "한도율": limit_ratio,
        "과세_공급가액_비율": taxable_ratio,
        "운반비율_적용": freight_exclusion,
        "안분_정보": allocation,
        "단계별_계산": steps,
    }
    if 품목별_결과 is not None:
        return _ok(result, "부가가치세법 제42조, 시행령 제81조·제84조",
                   계산={"품목별": 품목별_결과},
                   산식="공제대상 매입 = 과세 사용분 + 공통 재고 × 과세비율(운반비 제외 후), 의제매입세액 = min(공제대상 매입, 과세표준 × 한도율) × 공제율")
    return _ok(result, "부가가치세법 제42조, 시행령 제81조·제84조",
               산식="공제대상 매입 = 과세 사용분 + 공통 재고 × 과세비율(운반비 제외 후), 의제매입세액 = min(공제대상 매입, 과세표준 × 한도율) × 공제율")


@mcp.tool(annotations=CALC)
@_guard
def vat_common_input_allocation(common_tax: WON, total_supply: WON, exempt_supply: WON) -> dict:
    """겸영사업자 공통매입세액 안분(불공제분 계산). 언제: 과세·면세 사업을 겸영할 때 공통 매입세액 중 얼마가 불공제인지 계산한다.
입력: 공통매입세액(원) → common_tax, 총공급가액(원) → total_supply, 면세공급가액(원) → exempt_supply.
결과: 불공제 매입세액, 근거(부가가치세법 시행령 제81조). 면세비율 5% 미만·공통매입세액 5만 원 미만이면 전액 공제. 불공제 = 공통매입세액 × 면세공급가액 ÷ 총공급가액. 겸영사업자 공통매입세액 안분(불공제분 계산)."""
    return _ok(V.common_input_allocation(common_tax, total_supply, exempt_supply), "부가가치세법 시행령 제81조", 산식="불공제 = 공통매입세액 × 면세공급가액 ÷ 총공급가액")


@mcp.tool(annotations=CALC)
@_guard
def vat_simplified_taxpayer(supply_incl_vat: Annotated[int, Field(description="공급대가(부가세 포함, 원)")],
                            value_added_rate: Annotated[float, Field(description="업종별 부가가치율(소수, 예: 0.15 = 15%, 0 초과 1 미만)")],
                            invoice_purchase_incl_vat: Annotated[int, Field(description="세금계산서 등 수취 매입 공급대가(원)")] = 0,
                            mileage_supply_incl_vat: Annotated[int, Field(description="자기적립마일리지 등 결제분 공급대가(원) — 과세표준에서 제외")] = 0,
                            card_cash_receipt_supply_incl_vat: Annotated[int, Field(description="신용카드매출전표·현금영수증 등 발급분 공급대가(원)")] = 0,
                            no_invoice_purchase_incl_vat: Annotated[int, Field(description="세금계산서 발급 의무자로부터 미수취한 매입 공급대가(원)")] = 0,
                            prior_year_supply_incl_vat: Annotated[int, Field(description="직전 연도 공급대가(원) — 납부의무 면제 판정 참조")] = 0,
                            prepaid_tax: Annotated[int, Field(description="예정부과 고지·수시부과세액(원)")] = 0,
                            year: Annotated[int, Field(description="적용 연도(2016~2026)")] = 2026) -> dict:
    """간이과세자 납부세액 전 과정(과세표준·납부세액·매입공제·카드 발행세액공제·가산세·납부의무 면제 판정·차가감 납부세액). 언제: 간이과세자 부가가치세 신고 시 납부세액을 단계별로 계산한다.
입력: 공급대가(부가세 포함, 원) → supply_incl_vat, 업종별 부가가치율(소수, 예: 0.15 = 15%) → value_added_rate, 세금계산서 등 수취 매입 공급대가(원) → invoice_purchase_incl_vat, 자기적립마일리지 등 결제분 공급대가(원) → mileage_supply_incl_vat, 신용카드매출전표·현금영수증 등 발급분 공급대가(원) → card_cash_receipt_supply_incl_vat, 세금계산서 미수취 매입 공급대가(원) → no_invoice_purchase_incl_vat, 직전 연도 공급대가(원) → prior_year_supply_incl_vat, 예정부과 고지·수시부과세액(원) → prepaid_tax, 적용 연도(2016~2026) → year.
결과: 과세표준·납부세액·매입공제·카드공제·가산세·공제합계·차가감 납부세액·납부의무 면제 여부·단계별 계산, 근거(부가가치세법 제46조·제63조·제68조의2·제69조)."""
    from .engine.calc_vat import simplified_exempt_threshold, card_sales_credit
    from .engine.calc_vat import simplified_exempt_threshold, card_sales_credit

    # 과세표준 = 공급대가 − 마일리지
    tax_base = supply_incl_vat - mileage_supply_incl_vat

    # 납부세액 = 과세표준 × 부가가치율 × 10%
    tax_due = int(tax_base * value_added_rate * 0.1)

    # 매입세금계산서 등 공제 = 수취 매입 × 0.5%
    input_credit = int(invoice_purchase_incl_vat * 0.005)

    # 신용카드매출전표등 발행세액공제 (제46조, 1.3%, 연 한도 1,000만원)
    card_credit = card_sales_credit(card_cash_receipt_supply_incl_vat, f"{year}-01-01")

    # 세금계산서 미수취 가산세 (제68조의2, 0.5%)
    no_invoice_penalty = int(no_invoice_purchase_incl_vat * 0.005)

    # 공제 합계 (납부세액 한도)
    total_credits = min(input_credit + card_credit, tax_due)

    # 가산세 합계
    total_penalties = no_invoice_penalty

    # 차가감 납부세액 = 납부세액 − 공제 합계 + 가산세 (공제는 납부세액 한도, 음수이면 0)
    net_tax = max(0, tax_due - total_credits) + total_penalties

    # 납부의무 면제 여부 (제69조, 과세기간 공급대가 4,800만원 미만)
    exempt_threshold = simplified_exempt_threshold(year)
    exempt = supply_incl_vat < exempt_threshold

    steps = [
        {"항목": "공급대가(부가세 포함)", "금액": supply_incl_vat, "근거": "부가법 제63조"},
        {"항목": "자기적립마일리지 등 차감", "금액": mileage_supply_incl_vat, "근거": "부가가치세법 제63조"},
        {"항목": "과세표준(공급대가 − 마일리지)", "금액": tax_base, "근거": "부가법 제63조"},
        {"항목": f"납부세액(과세표준 × 부가가치율 {value_added_rate} × 10%)", "금액": tax_due, "근거": "부가가치세법 제63조"},
        {"항목": "매입세금계산서 등 공제(수취 매입 × 0.5%)", "금액": input_credit, "근거": "부가가치세법 제63조"},
        {"항목": "신용카드매출전표등 발행세액공제(1.3%, 한도 1,000만원)", "금액": card_credit, "근거": "부가가치세법 제46조"},
        {"항목": "공제 합계(납부세액 한도)", "금액": total_credits, "근거": "부가가치세법 제63조"},
        {"항목": "세금계산서 미수취 가산세(0.5%)", "금액": no_invoice_penalty, "근거": "부가가치세법 제68조의2"},
        {"항목": "가산세 합계", "금액": total_penalties, "근거": "부가가치세법 제68조의2"},
        {"항목": "차가감 납부세액(납부세액 − 공제 + 가산세, 음수 0)", "금액": net_tax, "근거": "부가가치세법 제63조"},
        {"항목": f"납부의무 면제 여부(해당 과세기간 공급대가 {supply_incl_vat:,}원, 기준 {exempt_threshold:,}원 미만)", "금액": "면제" if exempt else "해당없음", "근거": "부가가치세법 제69조"},
    ]

    result = {
        "과세표준": tax_base,
        "납부세액": tax_due,
        "매입세금계산서등_공제": input_credit,
        "카드발행세액공제": card_credit,
        "가산세": total_penalties,
        "공제_합계": total_credits,
        "차가감_납부세액": net_tax,
        "납부의무_면제": exempt,
        "직전_연도_공급대가": prior_year_supply_incl_vat,
        "면제_기준": exempt_threshold,
        "단계별_계산": steps,
    }
    return _ok(result, "부가가치세법 제46조·제63조·제68조의2·제69조",
               산식="과세표준 = 공급대가 − 마일리지, 납부세액 = 과세표준 × 부가가치율 × 10%, 차가감 = max(0, 납부세액 − min(매입공제+카드공제, 납부세액)) + 가산세")


@mcp.tool(annotations=CALC)
@_guard
def vat_card_sales_credit(amount: Annotated[int, Field(description="신용카드·현금영수증 발급분 공급대가(원)")], supplied_on: Annotated[str, Field(description="공급일 YYYY-MM-DD")],
                          used_this_year: Annotated[int, Field(description="같은 해 이미 공제받은 금액(원)")] = 0,
                          individual: Annotated[bool, Field(description="개인사업자 여부(True면 공제 대상)")] = True,
                          prior_year_supply: Annotated[int | None, Field(description="직전 연도 공급가액 합계(원). 한도 초과 시 공제 제외")] = None) -> dict:
    """신용카드매출전표 등 발행세액공제(공제율 1.3%, 연 한도 1,000만 원). 언제: 개인사업자가 신용카드·현금영수증 매출에 대해 발행세액공제를 얼마나 받는지 계산한다.
입력: 신용카드·현금영수증 발급분 공급대가(원) → amount, 공급일(YYYY-MM-DD, 세율·한도 연도표 기준) → supplied_on, 같은 해 이미 공제받은 금액(원) → used_this_year, 개인사업자 여부 → individual(기본 True), 직전 연도 공급가액 합계(원) → prior_year_supply(직전연도_공급가액_한도 초과 시 공제 0).
결과: 공제액(법인·한도 초과 시 0, 판단 문구 포함), 근거(부가가치세법 제46조①). 신용카드매출전표 등 발행세액공제; 공제율 1.3%(간이 음식·숙박 2.6%는 2021.7.1. 전), 연 한도 1,000만 원(2018년 제2기 확정부터), 개인사업자만 해당·직전 연도 공급가액 한도 이하."""
    # 법인 또는 직전 연도 공급가액 한도 초과: 공제 0 (부가가치세법 제46조①)
    공급가액_한도 = vocab.rule("직전연도_공급가액_한도", int(supplied_on[:4]))["한도"]
    if not individual or (prior_year_supply is not None and prior_year_supply > 공급가액_한도):
        return _ok({"공제액": 0},
                   f"부가가치세법 제46조① — 개인사업자 아님 또는 직전 연도 공급가액 {공급가액_한도:,}원 초과로 발행세액공제 대상 아님")
    return _ok({"공제액": V.card_sales_credit(amount, supplied_on, used_this_year=used_this_year)},
               "부가가치세법 제46조①")


@mcp.tool(annotations=CALC)
def vat_bad_debt_credit(bad_debt_incl_vat: Annotated[int, Field(description="대손금액(부가세 포함, 원)")],
                        supply_receivable: Annotated[bool, Field(description="재화·용역 공급 대가 채권 여부(True면 대손세액 공제 대상)")] = True) -> dict:
    """대손세액공제(대손금액 × 10/110). 언제: 외상 매출이 대손되었을 때 부가가치세 대손세액공제를 계산한다.
입력: 대손금액(부가세 포함, 원) → bad_debt_incl_vat, 재화·용역 공급 대가 채권 여부 → supply_receivable(기본 True).
결과: 대손세액(공급 대가 아니면 0), 근거(부가가치세법 제45조①). VAT bad debt credit."""
    if not supply_receivable:
        return _ok({"대손세액": 0, "공급_대가_채권_여부": False},
                   "부가가치세법 제45조① — 재화·용역의 공급 대가가 아닌 채권은 대손세액공제 대상 아님")
    return _ok({"대손세액": V.bad_debt_vat(bad_debt_incl_vat), "공급_대가_채권_여부": True},
               "부가가치세법 제45조①")


# ── 부가가치세 공급시기 판정(법 제15~17조, 영 제28~30조, 규칙 제17~20조) ─────────────
@mcp.tool(annotations=CALC)
@_guard
def vat_supply_timing(
    transactions: Annotated[list, Field(description=(
        "거래 목록 [{이름?, 종류, 인도일?(재화 인도·이용가능일, 용역은 완료일), 대가?: [{약정일, 금액, 구분?('계약금'|'중도금'|'잔금')}], "
        "공급가액?(대가 목록이 없을 때), 선적일?(수출), 확정일?(조건부·기한부 판매), 국고_선금급?(bool), "
        "세금계산서?: [{발급일, 금액?, 수령일?(대가 받은 날), 청구지급시기_별도기재?(bool), 지급시기?(계약서상 지급일), 공급받는자_조기환급?(bool)}]}]. "
        "종류: 일반판매·할부·장기할부·중간지급조건부·완성도기준지급·계속적공급(전력 등)·조건부판매·수출, 용역·용역_장기할부·용역_중간지급조건부·용역_완성도기준지급·용역_계속적공급"))],
    business: Annotated[Literal["일반", "간이"], Field(description="사업자 구분 — '일반'(과세기간 1·2기, 예정신고기간 1.1~3.31·7.1~9.30) 또는 '간이'(과세기간 1년, 예정부과기간 1.1~6.30)")] = "일반",
    period: Annotated[str | None, Field(description="조회 기간(선택). 일반: 'YYYY-1기 예정'·'YYYY-1기 확정'(4.1~6.30)·'YYYY-2기'(과세기간 전체). 간이: 'YYYY' 또는 'YYYY 예정부과'")] = None,
) -> dict:
    """부가가치세 공급시기 판정 — 거래별 공급시기·판정 사유와 예정신고기간·과세기간별 과세표준(공급가액) 합계.
    언제: 장기할부·중간지급조건부·완성도기준지급·계속적 공급·용역 완료·수출·세금계산서 선발급 거래가 어느 과세기간(예정/확정)에 들어가는지 정할 때. 공급 2016-01-01~2026-12-31 지원.
    입력: 거래 목록(종류·인도일/완료일·대가 지급 약정·세금계산서 발급일과 대가 수령일) → transactions, 사업자 구분('일반'|'간이') → business, 조회 기간 → period(선택).
    판정: 장기할부 = 2회 이상 분할 + 인도(완료)일 다음 날부터 최종 지급기일까지 1년 이상(규칙 제17·19조) → 각 약정일.
      중간지급조건부 = 계약금 약정일 다음 날부터 인도(완료)일까지 6개월 이상 + 기간 안 계약금 외 분할(규칙 제18·20조) → 각 약정일, 인도(완료)일 이후 약정분은 인도(완료)일(영 제28조③·제29조① 단서).
      기간은 민법 제157조·제160조(초일 불산입, 해당일 전일 만료, 해당일 없으면 말일). 요건 미충족이면 인도(완료)일 일괄.
      선발급 특례(법 제17조) 공급분 연도별 — 2017까지 ①은 '동시에' 발급, ③은 청구·지급시기 별도기재+30일 모두 / 2018 ③ 별도기재+(30일 또는 과세기간 내 지급 확인) /
      2019~2021 ③ 30일 약정 또는 과세기간 내 공급시기 도래+지급 확인 / 2022~ 지급 확인 요건 삭제(부칙 제6조). 장기할부·계속적 공급은 ④(영 제30조) 발급일.
    미지원(입력 확인 필요): 상품권·가공·무인판매기·보세구역·위탁판매·간주공급, 선적일 외 수출(중계무역·외국인도 등), 부동산 임대·선불 연회비(영 제29조②), 외국항행용역, 폐업일 의제, 공급시기 후 발급 가산세.
    결과: 기간별_과세표준·조회기간_과세표준, 거래별 공급시기·금액·신고기간·사유·근거, 판단(요건 경계 계산)."""
    from .engine import calc_vat_timing as VT
    try:
        r, rows, 판단 = VT.judge(transactions, business, period)
    except VT.InputError as e:
        return {"error": "입력 확인 필요 — 계산하지 않음", "문제": [str(e)]}
    # ── 계산식 단계: 거래별·공급시기별 금액(공급시기 판정 결과)을 입력으로, 기간별 합계는 그 행들의 합 ──
    계산식_입력: list[dict] = []
    기간행: dict[str, list[str]] = {}
    조회행: list[str] = []
    q = VT.parse_query(period, business.startswith("간이"))
    for n, row in enumerate(rows, 1):
        for m, it in enumerate(row["공급시기별"], 1):
            nm = f"거래{n}_공급{m}"
            계산식_입력.append({"이름": nm, "값": it["금액"],
                             "설명": f"{row['거래']} — 공급시기 {it['공급시기']} · {it['신고기간']} ({it['사유']}; {it['근거']})"})
            기간행.setdefault(it["신고기간"], []).append(nm)
            if q and q[0] <= date.fromisoformat(it["공급시기"]) <= q[1]:
                조회행.append(nm)
    계산식 = []
    기간이름: list[str] = []
    for k, (기간, 이름들) in enumerate(sorted(기간행.items()), 1):
        nm = "과세표준_" + re.sub(r"_+", "_", re.sub(r"[^가-힣A-Za-z0-9]", "_", 기간.split("(")[0])).strip("_")
        if nm in 기간이름:
            nm = f"{nm}_{k}"
        기간이름.append(nm)
        계산식.append({"이름": nm, "식": "+".join(이름들), "값": r["기간별_과세표준"][기간],
                     "설명": f"{기간} 과세표준(공급가액) 합계 — 이 기간에 공급시기가 든 거래 행의 합"})
    if 계산식:
        계산식.append({"이름": "거래_합계", "식": "+".join(기간이름), "값": r["거래_합계"], "설명": "모든 기간 합계"})
    if q:
        if 조회행:
            계산식.append({"이름": "조회기간_과세표준", "식": "+".join(조회행), "값": r["조회기간_과세표준"], "설명": f"조회기간 {r['조회기간']} 귀속 합계"})
        else:
            계산식_입력.append({"이름": "조회기간_과세표준", "값": 0, "설명": f"조회기간 {r['조회기간']} 귀속 공급 없음"})
    return _ok(r, "부가가치세법 제15조·제16조·제17조, 시행령 제28조·제29조·제30조, 시행규칙 제17조~제20조, 제5조①·제48조①·제66조①",
               거래별=rows, 판단=판단, 계산식_입력=계산식_입력, 계산식=계산식)


# ───────── 0.2 추가: 근로·원천 ─────────
@mcp.tool(annotations=CALC)
@_guard
def wage_income_tax(gross_pay: Annotated[int, Field(description="총급여(원, 비과세 제외) — 연봉에서 비과세 수당을 뺀 금액")],
                    year: YEAR,
                    tax_base: Annotated[int | None, Field(description="종합소득 과세표준(원). 모르면 생략 — 근로소득금액에서 본인 기본공제 150만 원만 뺀 값으로 계산하고 그 사실을 판단에 적음")] = None) -> dict:
    """근로소득공제·산출세액·근로소득세액공제·결정세액(참고). 언제: 근로소득만 있는 근로자의 연말정산에서, 또는 '연봉 얼마면 세금 얼마'를 어림할 때.
입력: 총급여(원, 비과세 제외) → gross_pay, 귀속연도(2016~2026) → year, 종합소득 과세표준(원, 모르면 생략) → tax_base.
결과: 근로소득공제·근로소득금액·과세표준·산출세액·근로소득세액공제·결정세액(참고: 표준세액공제 13만 원만 적용)·지방소득세, 근거(소득세법 제47조①·제55조·제59조·제59조의4⑨). Wage income: employment income deduction, computed tax and wage tax credit."""
    ded = I.earned_deduction(gross_pay, year)
    income_amt = max(0, gross_pay - ded)
    assumed = tax_base is None
    if assumed:
        tax_base = max(0, income_amt - 1_500_000)
    t = I.income_tax(tax_base, year)
    credit = I.earned_tax_credit(t, gross_pay, year=year)
    std = min(130_000, max(0, t - credit)) if assumed else 0
    decided = max(0, t - credit - std)
    판단 = [("과세표준을 주지 않아 근로소득금액 − 본인 기본공제 150만 원(소득세법 제50조①1호)으로 계산 — 부양가족·보험료·주택자금 등 다른 소득공제가 있으면 tax_base를 직접 주세요" if assumed
             else "과세표준은 입력값 그대로"),
          ("결정세액(참고)은 특별소득공제·특별세액공제·월세액 세액공제를 신청하지 않은 근로자의 표준세액공제 13만 원(제59조의4⑨)만 뺀 값 — 자녀·연금계좌·의료비 등 세액공제가 있으면 그만큼 줄어듦" if assumed
           else "과세표준을 직접 준 경우 표준세액공제는 적용하지 않음(특별소득공제·특별세액공제를 신청했는지 알 수 없음) — 결정세액(참고)은 산출세액 − 근로소득세액공제이며 표준세액공제 13만 원(제59조의4⑨) 또는 특별·자녀·연금계좌 세액공제는 별도로 빼세요"),
          "지방소득세는 소득세 결정세액의 10%(지방세법 제103조의3)를 참고로 붙임"]
    return _ok({"근로소득공제": ded, "근로소득금액": income_amt, "과세표준": tax_base, "산출세액": t, "근로소득세액공제": credit,
                "표준세액공제": std, "결정세액(참고)": decided, "지방소득세(참고)": decided // 10},
               "소득세법 제47조①·제55조·제59조·제59조의4⑨", 판단=판단, 주의="그 밖의 소득공제·세액공제는 별도")


@mcp.tool(annotations=CALC)
@_guard
def daily_worker_withholding(daily_wage: Annotated[int, Field(description="일급(원)")], year: Annotated[int, Field(description="지급연도(2016~2026)")] = 2025) -> dict:
    """일용근로자 원천징수세액((일급 − 15만 원) × 6% × (1 − 55%), 1천 원 미만 소액부징수). 언제: 일용근로자에게 일급을 지급할 때 원천징수할 세액을 계산한다.
입력: 일급(원) → daily_wage, 지급연도(2016~2026) → year.
결과: 원천징수세액, 근거(소득세법 제134조③·제129조①4호·제86조). Withholding on daily workers' wages."""
    return _ok({"원천징수세액": I.daily_worker_tax(daily_wage, year)}, "소득세법 제134조③·제129조①4호·제86조",
               산식="(일급 − 15만원) × 6% × (1 − 55%), 1천원 미만 부징수")


@mcp.tool(annotations=CALC)
@_guard
def deemed_bonus_resettlement(gross_pay: Annotated[int, Field(description="귀속연도 당초 총급여(원)")], tax_base: Annotated[int, Field(description="당초 과세표준(원)")],
                              bonus: Annotated[int, Field(description="소득처분 상여(원)")], year: YEAR,
                              other_credits: Annotated[int, Field(description="근로소득세액공제 외 세액공제 합계(원)")] = 0,
                              prev_decided: Annotated[int | None, Field(description="당초 결정세액(원). 모르면 생략")] = None) -> dict:
    """소득처분 상여 연말정산 재정산(추가 원천징수세액). 언제: 법인세 소득금액 변동으로 상여 처분이 발생했을 때 해당 근로자의 연말정산을 재정산해 추가 원천징수세액을 구한다.
입력: 귀속연도 당초 총급여(원) → gross_pay, 당초 과세표준(원) → tax_base, 소득처분 상여(원) → bonus, 귀속연도(2016~2026) → year, 근로소득세액공제 외 세액공제 합계(원) → other_credits, 당초 결정세액(원, 모르면 생략) → prev_decided.
결과: 추가 원천징수세액, 근거(법인세법 제67조, 소득세법 시행령 제192조, 소득세법 제137조). Year-end resettlement after a deemed bonus (income disposition) is added to wages."""
    return _ok(I.deemed_bonus_resettlement(gross_pay, tax_base, bonus, other_credits, prev_decided, year=year),
               "법인세법 제67조, 소득세법 시행령 제192조, 소득세법 제137조",
               주의="총급여 연동 소득·세액공제(신용카드 공제 등)는 수기 보정 필요")


# ───────── 0.2 추가: 근로소득 총급여 ─────────
# 근거 조문: 소득세법 제20조(근로소득)·제12조3호(비과세)·제47조(근로소득공제), 시행령 제11조·제12조·제17조·제17조의2~제17조의4·제38조·제49조
_EGP_KIND_필요사실 = {
    "급여":             ["amount"],
    "상여":             ["amount"],
    "잉여금처분상여":   ["amount", "resolution_date"],
    "인정상여":         ["amount"],
    "식사대":           ["amount", "months", "meal_provided"],
    "자가운전보조금":   ["amount", "months", "vehicle", "business_use", "actual_travel_reimbursed"],
    "출산·보육수당":    ["amount", "purpose"],
    "직무발명보상금":   ["amount", "retired"],
    "연장근로수당":     ["amount", "eligible_occupation"],
    "사택제공이익":     ["amount", "position"],
    "주택자금대여이익": ["amount", "sme"],
    "단체보험료":       ["amount"],
    "학자금":           ["amount", "beneficiary"],
    "실비변상여비":     ["amount", "actual_cost"],
    "사회보험료_본인부담분_회사대납": ["amount"],
    "사회보험료_사용자부담분":        ["amount"],
}

_EGP_월한도_종류 = {"식사대", "자가운전보조금", "출산·보육수당"}


_EGP_근무월수_종류 = ("식사대", "자가운전보조금", "출산·보육수당")


def _egp_월정액급여_산정(year: int, items: list, 판정: dict, 제외_idx: dict, 본인_idx: int) -> dict:
    """영 제17조① 후단의 월정액급여를 같은 호출의 다른 항목으로 산정한다.
    월정액급여 = 매월 직급별로 받는 봉급·급료·보수·임금·수당, 그 밖에 이와 유사한 성질의 급여
      (상여 등 부정기적인 급여, 영 제12조 실비변상적 급여, [2021~] 영 제17조의4 복리후생적 급여 제외)의 총액
      − 「근로기준법」 연장·야간·휴일근로 급여(·「선원법」 생산수당).
    판정: {원래 index: 항목별 dict(과세액·비과세액·비과세_성격)}, 제외_idx: {원래 index: 제외 사유}.
    반환: {"월정액급여", "근무월수", "포함", "제외", "산식"} 또는 {"error", ...}."""
    복리후생_제외 = year >= 2021   # 영 제17조① 개정 2021.2.17 — 제17조의4 복리후생적 급여 제외 문구 추가
    포함: list[dict] = []
    포함_i: list[int] = []      # 포함 항목의 원래 index(계산식에서 항목 입력을 가리킬 때)
    제외: list[dict] = []

    def _포함(i, d):
        포함.append(d)
        포함_i.append(i)
    모호: list[str] = []
    필요: dict[str, list[str]] = {}

    def _모호(k, 문구, 키="regular_monthly"):
        모호.append(f"{k}: {문구}")
        필요.setdefault(k, [])
        if 키 not in 필요[k]:
            필요[k].append(키)

    # 근무월수 — 연장근로수당 facts.months, 없으면 월 한도 항목의 months(모두 같을 때), 둘 다 없으면 12
    본인_facts = (items[본인_idx].get("facts") or {})
    if 본인_facts.get("months") not in (None, ""):
        근무월수 = int(본인_facts["months"]); 월수_출처 = "연장근로수당 facts.months"
    else:
        월수들 = {int((it.get("facts") or {})["months"]) for it in items
                  if it.get("kind") in _EGP_근무월수_종류 and (it.get("facts") or {}).get("months") not in (None, "")}
        if len(월수들) > 1:
            return {"error": "입력 확인 필요 — 계산하지 않음",
                    "문제": [f"연장근로수당: 월정액급여 산정용 근무월수가 항목마다 다름 {sorted(월수들)} — "
                             "연장근로수당 facts.months(근무월수)를 넣거나 monthly_regular_wage를 직접 입력"],
                    "필요 입력": {"연장근로수당": ["months", "monthly_regular_wage"]}}
        if 월수들:
            근무월수 = 월수들.pop(); 월수_출처 = "식사대·자가운전보조금 등 facts.months"
        else:
            근무월수 = 12; 월수_출처 = "근무월수 입력 없음 — 12개월로 봄"
    if 근무월수 <= 0:
        return {"error": "입력 확인 필요 — 계산하지 않음",
                "문제": ["연장근로수당: 근무월수(months)가 0 이하"],
                "필요 입력": {"연장근로수당": ["months"]}}

    if not any(it.get("kind") == "급여" and i not in 제외_idx for i, it in enumerate(items)):
        return {"error": "입력 확인 필요 — 계산하지 않음",
                "문제": ["연장근로수당: monthly_regular_wage가 없고, 월정액급여를 산정할 해당 연도 급여(봉급·임금) 항목도 없음"],
                "필요 입력": {"연장근로수당": ["monthly_regular_wage"], "급여": ["amount"]}}

    for i, it in enumerate(items):
        k = it.get("kind", "")
        amount = int(it.get("amount", 0) or 0)
        facts = it.get("facts") or {}
        정기 = facts.get("regular_monthly")   # 매월 정기 지급 여부(선택)
        if i in 제외_idx:
            제외.append({"kind": k, "금액": amount, "사유": f"{year}년 근로소득 아님 — {제외_idx[i]}"})
            continue
        p = 판정.get(i) or {"과세액": amount, "비과세액": 0}   # 연장근로수당(아직 판정 전)은 금액 그대로
        과세, 비과세 = p["과세액"], p["비과세액"]
        if k == "연장근로수당":
            제외.append({"kind": k, "금액": amount, "사유": "「근로기준법」 연장·야간·휴일근로로 통상임금에 더하여 받는 급여 — 총액에서 뺌(영 제17조① 후단)"})
        elif k in ("상여", "잉여금처분상여", "인정상여"):
            제외.append({"kind": k, "금액": amount, "사유": "상여 등 부정기적인 급여 — 제외(영 제17조① 후단 괄호)"})
        elif k == "사회보험료_사용자부담분":
            제외.append({"kind": k, "금액": amount, "사유": "사용자 부담분 — 근로자가 받는 급여 아님(법 제12조3호너목)"})
        elif k in ("급여", "식사대", "사회보험료_본인부담분_회사대납") or (k == "출산·보육수당" and facts.get("purpose") == "보육"):
            if 정기 is False:
                제외.append({"kind": k, "금액": amount, "사유": "매월 정기 지급 아님(regular_monthly=false) — 부정기적인 급여로 제외"})
            else:
                사유 = {"급여": "매월 받는 봉급·급료·임금 — 포함",
                        "식사대": "매월 받는 수당 — 포함(비과세 식사대도 영 제17조의2 근거로 제12조 실비변상·제17조의4 복리후생 제외 대상 아님)",
                        "사회보험료_본인부담분_회사대납": "근로자 부담분 회사 대납 — 매월 급여와 유사한 성질의 급여로 포함",
                        "출산·보육수당": "매월 받는 보육수당 — 포함(비과세 근거 법 제12조3호머목, 제외 대상 아님)"}[k]
                _포함(i, {"kind": k, "금액": amount, "사유": 사유})
        elif k == "자가운전보조금":
            if 비과세 > 0:
                제외.append({"kind": k, "금액": 비과세, "사유": "비과세 자가운전보조금 — 영 제12조 실비변상적 급여로 제외"})
            if 과세 > 0:
                if 정기 is False:
                    제외.append({"kind": k, "금액": 과세, "사유": "과세분 — 매월 정기 지급 아님(regular_monthly=false)"})
                else:
                    _포함(i, {"kind": k, "금액": 과세, "사유": "비과세 한도 초과 과세분 — 영 제12조 실비변상적 급여 아님, 매월 받는 수당으로 포함"})
        elif k == "실비변상여비":
            if 비과세 > 0:
                제외.append({"kind": k, "금액": 비과세, "사유": "실비변상 여비 — 영 제12조 실비변상적 급여로 제외"})
            if 과세 > 0:
                if 정기 is None:
                    _모호(k, "실비 정산이 아닌 여비(과세분) — 매월 정기 지급 수당인지(regular_monthly) 입력 필요")
                elif 정기:
                    _포함(i, {"kind": k, "금액": 과세, "사유": "실비 정산 아닌 정액 여비, 매월 정기 지급 — 포함"})
                else:
                    제외.append({"kind": k, "금액": 과세, "사유": "부정기 지급 — 제외"})
        elif k in ("사택제공이익", "주택자금대여이익", "단체보험료"):
            나머지 = amount
            if 비과세 > 0 and 복리후생_제외:
                제외.append({"kind": k, "금액": 비과세, "사유": "영 제17조의4 복리후생적 급여 — 제외(영 제17조① 후단, 2021.2.17 개정분)"})
                나머지 -= 비과세
            if p.get("비과세_성격") == "근로소득 아님":
                제외.append({"kind": k, "금액": 나머지, "사유": "근로소득 아님(영 제38조① 단서) — 급여 총액에 넣지 않음"})
                나머지 = 0
            if 나머지 > 0:
                if 정기 is None:
                    근거문 = "" if 복리후생_제외 else f"({year}년 영 제17조①에는 복리후생적 급여 제외 문구 없음 — 2021.2.17 개정으로 추가) "
                    _모호(k, f"{나머지:,}원 {근거문}— 매월 정기적으로 받는 급여인지(regular_monthly) 입력 필요")
                elif 정기:
                    _포함(i, {"kind": k, "금액": 나머지, "사유": "매월 정기적으로 받는 급여 — 포함"})
                else:
                    제외.append({"kind": k, "금액": 나머지, "사유": "부정기 — 제외"})
        else:
            # 직무발명보상금·학자금·출산 급여 등 — 지급 형태가 항목 이름만으로 정해지지 않음
            if 정기 is None:
                _모호(k, "매월 정기 지급(월정액급여 포함)인지 부정기 지급(제외)인지 알 수 없음 — regular_monthly(true/false) 입력 필요")
            elif 정기:
                _포함(i, {"kind": k, "금액": amount, "사유": "매월 정기 지급 — 포함"})
            else:
                제외.append({"kind": k, "금액": amount, "사유": "부정기적인 급여 — 제외(영 제17조① 후단 괄호)"})

    if 모호:
        필요.setdefault("연장근로수당", []).append("monthly_regular_wage")
        return {"error": "입력 확인 필요 — 계산하지 않음",
                "문제": [f"연장근로수당: monthly_regular_wage가 없어 월정액급여를 항목으로 산정하려 했으나 판단할 수 없는 항목이 있음 — {m}"
                         for m in 모호],
                "필요 입력": 필요}

    포함합 = sum(x["금액"] for x in 포함)
    월정액 = 포함합 // 근무월수
    산식 = (f"월정액급여 = (" + " + ".join(f"{x['kind']} {x['금액']:,}" for x in 포함) +
            f") ÷ {근무월수}개월 = {포함합:,} ÷ {근무월수} = {월정액:,}원 ({월수_출처})")
    return {"월정액급여": 월정액, "포함_합계": 포함합, "근무월수": 근무월수, "포함": 포함, "제외": 제외, "산식": 산식,
            "포함_i": 포함_i, "월수_출처": 월수_출처}


@mcp.tool(annotations=CALC)
@_guard
def earned_income_gross_pay(year: YEAR,
                            items: Annotated[list, Field(description="근로소득 항목 [{kind, amount, facts}] — kind: 급여·상여·잉여금처분상여·인정상여·식사대·자가운전보조금·출산·보육수당·직무발명보상금·연장근로수당·사택제공이익·주택자금대여이익·단체보험료·학자금·실비변상여비·사회보험료_본인부담분_회사대납·사회보험료_사용자부담분. facts는 종류별 사실관계 dict(빠지면 필요한 입력을 돌려줌)")]) -> dict:
    """근로소득 항목별 과세·비과세를 판정하고 귀속연도를 확인해 총급여액·근로소득공제·근로소득금액을 계산한다(소득세법 제20조·제12조3호·제47조, 시행령 제11조·제12조·제17조·제17조의2~4·제38조·제49조).
    언제: 연말정산·종합소득세 신고에서 급여·상여·수당·복리후생 급여 자료로 총급여액과 근로소득금액을 구할 때.
    입력: 귀속연도(2016~2026) → year, 항목 목록 [{kind, amount, facts}] → items.
      kind와 facts(필요한 사실):
        · 급여·상여 — 다른 해 근로분이면 service_year
        · 잉여금처분상여 — resolution_date(결의일 연도 귀속) / 인정상여 — service_year
        · 식사대 — months, meal_provided / 자가운전보조금 — months, vehicle(본인소유|본인명의임차|기타), business_use, actual_travel_reimbursed
        · 출산·보육수당 — purpose(보육: months, child_age(과세기간 개시일 현재 만 나이, 6세 이하) / 출산: 2024~ within_2y_of_birth·payment_count·related_party, 그 전 months).
          보육 한도: 2016~2023 월 10만원(출산 포함), 2024~2025 월 20만원(근로자 1인당), 2026~ 자녀 1명당 월 20만원 —
          2026~는 child_ages(자녀별 개시일 나이 목록) 또는 child_count, 또는 자녀별 항목으로 나누고 facts.child(자녀 이름)를 다르게
        · 직무발명보상금 — retired(퇴직 후면 기타소득), 2024~ related_party
        · 연장근로수당 — eligible_occupation, 해당 시 prior_year_gross(필수)·monthly_regular_wage(광산·일용은 mine_or_daily).
          monthly_regular_wage가 없으면 같은 호출 항목으로 영 제17조① 월정액급여를 산정(연간 합계 ÷ months, 기본 12) —
          지급 형태가 이름만으로 정해지지 않는 항목(직무발명보상금·학자금·출산 급여 등)은 facts.regular_monthly(매월 정기 지급 여부)
        · 사택제공이익 — position / 주택자금대여이익 — sme(2024~ related_party)
        · 학자금 — beneficiary(본인|가족), 본인이면 requirements_met / 실비변상여비 — actual_cost
        · 단체보험료(연 한도) / 사회보험료_본인부담분_회사대납(과세) / 사회보험료_사용자부담분 — insurance
    결과: 항목별(과세액·비과세액·근거·사유)·제외 항목(다른 해·다른 소득)·총급여액·비과세 합계·근로소득공제·근로소득금액."""

    # ── 1. vocab 기반 허용값 검증 ──
    _v = _validate_items(items, "kind", _specs_from_vocab("earned_income_gross_pay", "kind"))
    if _v: return _v

    # ── 2. 전적 월한도 항목 중복 체크 ──
    월한도_건수: dict[str, int] = {}
    for it in items:
        k = it.get("kind", "")
        if k in _EGP_월한도_종류:
            월한도_건수[k] = 월한도_건수.get(k, 0) + 1
    for k, cnt in 월한도_건수.items():
        if cnt > 1 and k == "출산·보육수당" and vocab.rule(k, year).get("보육_자녀1명당"):
            # 2026~ 보육수당은 자녀 1명당 월 20만원(제12조3호머목2) — 자녀별 항목(facts.child 서로 다름)이면 허용
            _보육 = [it for it in items if it.get("kind") == k]
            _자녀 = [(it.get("facts") or {}).get("child") for it in _보육]
            if all((it.get("facts") or {}).get("purpose") == "보육" for it in _보육) \
                    and all(c not in (None, "") for c in _자녀) and len(set(_자녀)) == len(_자녀):
                continue
            return {"error": "입력 확인 필요 — 계산하지 않음",
                    "문제": [f"같은 kind '{k}'가 {cnt}번 입력됨 — {year}년 보육수당은 자녀 1명당 월 한도이므로 "
                             "자녀별로 나눌 때는 모든 항목을 purpose='보육'으로 하고 facts.child(자녀 이름)를 서로 다르게 넣거나, "
                             "한 항목에 child_ages(과세기간 개시일 기준 자녀 나이 목록)를 넣으세요"],
                    "필요 입력": {k: ["amount", "months", "purpose", "child", "child_age"]}}
        if cnt > 1:
            return {"error": "입력 확인 필요 — 계산하지 않음",
                    "문제": [f"같은 kind '{k}'가 {cnt}번 입력됨 — 월 한도는 한 항목으로 합쳐 입력하세요"],
                    "필요 입력": {k: ["amount", "months", "meal_provided" if k == "식사대" else
                                       "vehicle" if k == "자가운전보조금" else "purpose"]}}

    # ── 3. 필요 사실 누락 체크 ──
    # facts 자체가 없는 항목({kind, amount}만 있는 경우)은 허용:
    #   · 급여·상여·인정상여: service_year가 없어도 귀속연도=year로 계산
    #   · 단체보험료·사회보험료_본인부담분_회사대납·사회보험료_사용자부담분: facts 불필요
    #   · 그 외 kinds는 필요한 facts가 빠지면 필요 입력 error
    누락: dict[str, list[str]] = {}
    for it in items:
        k = it.get("kind", "")
        사실과있음 = bool(it.get("facts"))   # facts 키 자체가 없거나 빈 dict면 False
        facts = it.get("facts", {}) or {}
        필요 = _EGP_KIND_필요사실.get(k, ["amount"])
        빠진 = []
        for f in 필요:
            if f == "amount":
                if it.get("amount") in (None, ""):
                    빠진.append("amount")
            elif 사실과있음 and (f not in facts or facts.get(f) in (None, "", [])):
                빠진.append(f)
        if 빠진:
            누락[k] = 빠진
    if 누락:
        return {"error": "입력 확인 필요 — 계산하지 않음",
                "문제": [f"항목 kind='{k}': 필요 사실 {빠진} 누락" for k, 빠진 in 누락.items()],
                "필요 입력": 누락}

    # ── 4. 항목별 판정 ──
    항목별: list[dict] = []
    월정액_식: dict | None = None   # 연장근로수당 비과세 판정에 월정액급여를 산정했으면 계산식용 정보
    제외_항목: list[dict] = []
    총급여액 = 0
    비과세_합계 = 0
    근로소득_아님_합계 = 0
    직무발명보상금_비과세_사용액 = 0
    산식_목록: list[str] = []
    근거_목록: list[list[str]] = []
    판단_목록: list[dict] = []
    # 연장근로수당은 월정액급여를 다른 항목 판정 결과로 산정할 수 있게 맨 뒤에 판정(출력 순서는 입력 순서 유지)
    처리순서 = [i for i, it in enumerate(items) if it.get("kind") != "연장근로수당"] + \
              [i for i, it in enumerate(items) if it.get("kind") == "연장근로수당"]
    판정_idx: dict[int, dict] = {}     # 원래 index → 항목별 dict
    제외_idx: dict[int, str] = {}      # 원래 index → 제외 사유

    for _idx in 처리순서:
        it = items[_idx]
        k = it["kind"]
        amount = int(it.get("amount", 0))
        facts = it.get("facts", {}) or {}
        r = vocab.rule(k, year)
        b = vocab.basis(k, year)
        구분 = r.get("구분", "과세")
        과세액 = amount
        비과세액 = 0
        비과세_성격 = ""
        사유 = ""

        # ── 귀속시기 확인 (제외 판정) ──
        if k in ("급여", "상여"):
            sy = facts.get("service_year")
            if sy is not None and int(sy) != year:
                제외_항목.append({"kind": k, "금액": amount, "사유": f"{int(sy)}년 귀속(근로 제공 연도) — {year}년 총급여에서 제외",
                                  "근거": b})
                제외_idx[_idx] = 제외_항목[-1]["사유"]; continue
        elif k == "잉여금처분상여":
            rd = facts.get("resolution_date")
            if rd:
                결의연도 = int(str(rd)[:4])
                if 결의연도 != year:
                    제외_항목.append({"kind": k, "금액": amount, "사유": f"{결의연도}년 귀속(잉여금처분결의일) — {year}년 총급여에서 제외",
                                      "근거": b})
                    제외_idx[_idx] = 제외_항목[-1]["사유"]; continue
        elif k == "인정상여":
            sy = facts.get("service_year")
            if sy is not None and int(sy) != year:
                제외_항목.append({"kind": k, "금액": amount, "사유": f"{int(sy)}년 귀속(근로 제공 사업연도) — {year}년 총급여에서 제외",
                                  "근거": b})
                제외_idx[_idx] = 제외_항목[-1]["사유"]; continue

        # ── 항목별 과세·비과세 판정 ──
        if k == "식사대":
            months = int(facts.get("months", 0))
            meal_provided = facts.get("meal_provided", False)
            월한도 = r["월한도"]
            if meal_provided:
                비과세액 = 0
                과세액 = amount
                사유 = "현물 식사 별도 제공 — 식사대 전액 과세(영 제17조의2 2호 단서)"
                비과세_성격 = "과세(현물 식사 제공)"
            else:
                비과세액 = min(amount, 월한도 * months)
                과세액 = amount - 비과세액
                사유 = f"월 {월한도:,}원 × {months}개월 = {월한도 * months:,}원 이내 비과세"
                비과세_성격 = "월한도비과세"
            if 비과세액 > 0:
                산식_목록.append(f"{k}: min({amount:,}, {월한도:,}×{months}) = {비과세액:,}원 비과세, 과세 {과세액:,}원")

        elif k == "자가운전보조금":
            months = int(facts.get("months", 0))
            vehicle = facts.get("vehicle", "기타")
            business_use = facts.get("business_use", False)
            actual_travel = facts.get("actual_travel_reimbursed", False)
            월한도 = r["월한도"]
            차량요건 = r.get("차량요건", [])
            비과세_대상 = (vehicle in 차량요건) and bool(business_use) and not actual_travel
            if 비과세_대상:
                비과세액 = min(amount, 월한도 * months)
                과세액 = amount - 비과세액
                사유 = f"본인 소유·임차 차량 업무 사용, 실제 여비 별도 수령 없음 — 월 {월한도:,}원 × {months}개월 이내 비과세(영 제12조 3호)"
                비과세_성격 = "월한도비과세"
            else:
                과세액 = amount
                비과세액 = 0
                if vehicle not in 차량요건:
                    사유 = f"차량 요건 미충족(차량: {vehicle}, 허용: {차량요건}) — 전액 과세"
                elif not business_use:
                    사유 = "업무 사용 아님 — 전액 과세"
                else:
                    사유 = "실제 여비 별도 수령 — 전액 과세"
                비과세_성격 = "과세(요건 미충족)"
            if 비과세액 > 0:
                산식_목록.append(f"{k}: min({amount:,}, {월한도:,}×{months}) = {비과세액:,}원 비과세, 과세 {과세액:,}원")

        elif k == "출산·보육수당":
            purpose = facts.get("purpose", "")
            months = int(facts.get("months", 0))
            child_age = facts.get("child_age")
            within_2y = facts.get("within_2y_of_birth", False)
            payment_count = facts.get("payment_count", 0)
            related_party = facts.get("related_party", False)
            월한도 = r["월한도"]
            출산전액 = r.get("출산전액비과세", False)

            if purpose == "보육":
                # 자녀 나이는 과세기간 개시일(1/1) 현재 만 나이 — "6세 이하"(제12조3호머목2)
                자녀1명당 = bool(r.get("보육_자녀1명당"))   # 2026~ 자녀 1명당 월 20만원, 2025년까지 근로자 1인당
                child_ages = facts.get("child_ages")
                if child_ages:
                    대상자녀수 = sum(1 for a in child_ages if int(a) <= 6)
                elif facts.get("child_count") is not None:
                    대상자녀수 = int(facts["child_count"])
                else:
                    대상자녀수 = 0 if (child_age is not None and int(child_age) > 6) else 1
                if 대상자녀수 == 0:
                    과세액 = amount; 비과세액 = 0
                    _나이 = child_ages if child_ages else f"{int(child_age)}세"
                    사유 = f"자녀 나이 {_나이}(과세기간 개시일 기준) — 6세 이하 아님, 전액 과세"
                    비과세_성격 = "과세(연령 초과)"
                else:
                    배수 = 대상자녀수 if 자녀1명당 else 1
                    월한도 = 월한도 * 배수
                    비과세액 = min(amount, 월한도 * months)
                    과세액 = amount - 비과세액
                    if 자녀1명당:
                        사유 = (f"6세 이하 자녀 {대상자녀수}명 보육 — 자녀 1명당 월 {r['월한도']:,}원 × {대상자녀수}명 × "
                                f"{months}개월 이내 비과세({year}년~ 제12조3호머목2)")
                    else:
                        사유 = f"6세 이하 자녀 보육 — 월 {월한도:,}원 × {months}개월 이내 비과세"
                        if 대상자녀수 > 1:
                            사유 += f"({year}년은 자녀 수와 관계없이 근로자 1인당 월 한도 — 자녀 1명당 한도는 2026년부터)"
                    비과세_성격 = "월한도비과세"
            elif purpose == "출산":
                if year >= 2024 and 출산전액 and within_2y and payment_count <= r.get("출산지급횟수한도", 2) and not related_party:
                    비과세액 = amount; 과세액 = 0
                    사유 = f"출생 후 2년 이내 {int(payment_count)}번째 출산 급여 — 전액 비과세(특수관계자 제외, 영 제17조의2 2호)"
                    비과세_성격 = "출산전액비과세"
                else:
                    비과세액 = min(amount, 월한도 * months)
                    과세액 = amount - 비과세액
                    사유 = f"출산 급여 — 월 {월한도:,}원 × {months}개월 이내 비과세(2024년 전 또는 요건 미충족 시 월한도 적용)"
                    비과세_성격 = "월한도비과세"
            else:
                비과세액 = min(amount, 월한도 * months)
                과세액 = amount - 비과세액
                사유 = f"출산·보육수당 — 월 {월한도:,}원 × {months}개월 이내 비과세"
                비과세_성격 = "월한도비과세"
            if 비과세액 > 0:
                산식_목록.append(f"{k}: min({amount:,}, {월한도:,}×{months}) = {비과세액:,}원 비과세, 과세 {과세액:,}원")

        elif k == "직무발명보상금":
            retired = facts.get("retired", False)
            related_party = facts.get("related_party", False)
            연한도 = r.get("연한도", 0)
            근로해당 = r.get("근로소득해당", True)
            특수관계제외 = r.get("특수관계제외", False)

            if year == 2016:
                제외_항목.append({"kind": k, "금액": amount, "사유": "기타소득(2016년까지는 근로소득 범위 밖 — 소득세법 제20조 2016 시행본에 직무발명보상금 없음)",
                                  "근거": b})
                제외_idx[_idx] = 제외_항목[-1]["사유"]; continue
            if retired:
                제외_항목.append({"kind": k, "금액": amount, "사유": "기타소득(퇴직 후 수령 — 발명진흥법 직무발명보상금, 소득세법 제21조 22호의2)",
                                  "근거": b})
                제외_idx[_idx] = 제외_항목[-1]["사유"]; continue
            if year >= 2024 and 특수관계제외 and related_party:
                과세액 = amount; 비과세액 = 0
                사유 = "2024년부터 사용자등과 특수관계자는 비과세 제외(영 제17조의3 ②)"
                비과세_성격 = "과세(특수관계)"
            else:
                비과세액 = min(amount, 연한도)
                과세액 = amount - 비과세액
                직무발명보상금_비과세_사용액 += 비과세액
                사유 = f"재직 중 직무발명보상금 — 연 {연한도:,}원 이내 비과세(영 제17조의3)"
                비과세_성격 = "연한도비과세"

        elif k == "연장근로수당":
            eligible = facts.get("eligible_occupation", False)
            mine_or_daily = facts.get("mine_or_daily", False)
            if eligible:
                mrw = facts.get("monthly_regular_wage")
                pYG = facts.get("prior_year_gross")
                if pYG is None:
                    # 직전 과세기간 총급여액은 이번 호출 항목으로 알 수 없음 — 항상 입력 필요
                    return {"error": "입력 확인 필요 — 계산하지 않음",
                            "문제": ["연장근로수당: eligible_occupation=True이나 직전연도 총급여(prior_year_gross) 누락"
                                     + ("" if mrw is not None else " — 월정액급여는 monthly_regular_wage가 없으면 같은 호출의 항목으로 산정")],
                            "필요 입력": {"연장근로수당": ["prior_year_gross"] + (
                                ["monthly_regular_wage"] if mrw is None and not any(x.get("kind") == "급여" for x in items) else [])}}
                월정액한도 = r["월정액급여한도"]
                총급여한도 = r["직전연도총급여한도"]
                월정액_초과 = None
                월정액_산정됨 = mrw is None
                if mrw is None:
                    # 영 제17조① 후단 정의대로 같은 호출의 다른 항목으로 월정액급여 산정
                    산정 = _egp_월정액급여_산정(year, items, 판정_idx, 제외_idx, _idx)
                    if "error" in 산정:
                        return 산정
                    mrw = 산정["월정액급여"]
                    월정액_초과 = 산정["포함_합계"] > 월정액한도 * 산정["근무월수"]   # 원 단위 절사 전 비교
                    월정액_식 = {"idx": _idx, "산정": 산정, "월정액한도": 월정액한도, "총급여한도": 총급여한도,
                               "직전": int(pYG), "광산일용": bool(mine_or_daily), "연한도": r.get("연한도", 0), "지급액": amount}
                    산식_목록.append(산정["산식"])
                    판단_목록.append({"항목": "월정액급여 산정(영 제17조① 후단)",
                                     "결론": f"월정액급여 {mrw:,}원 — {'한도 ' + format(월정액한도, ',') + '원 초과' if 월정액_초과 else '한도 ' + format(월정액한도, ',') + '원 이하'}",
                                     "포함": 산정["포함"], "제외": 산정["제외"], "산식": 산정["산식"],
                                     "근거": "소득세법 시행령 제17조① 후단(상여 등 부정기 급여·영 제12조 실비변상적 급여"
                                             + ("·영 제17조의4 복리후생적 급여" if year >= 2021 else "")
                                             + " 제외, 연장·야간·휴일근로 급여 차감)",
                                     "주의": "연간 금액을 근무월수로 나눈 평균 — 달마다 월정액급여가 달라 한도를 넘는 달이 있으면 monthly_regular_wage로 그 달 금액을 직접 입력"})
                if 월정액_초과 is None:
                    월정액_초과 = int(mrw) > 월정액한도
                if 월정액_초과 or int(pYG) > 총급여한도:
                    과세액 = amount; 비과세액 = 0
                    사유 = f"월정액급여 {int(mrw):,}원 > {월정액한도:,}원 또는 직전연도 총급여 {int(pYG):,}원 > {총급여한도:,}원 — 비과세 요건 미충족(영 제17조①)"
                    비과세_성격 = "과세(요건 미충족)"
                elif mine_or_daily:
                    비과세액 = amount; 과세액 = 0
                    사유 = "광산·일용근로자 — 연장근로수당 전액 비과세(영 제17조② 1호)"
                    비과세_성격 = "연한도비과세(전액)"
                else:
                    연한도 = r["연한도"]
                    비과세액 = min(amount, 연한도)
                    과세액 = amount - 비과세액
                    사유 = f"생산직 등 요건 충족 — 연 {연한도:,}원 이내 비과세(영 제17조② 1호)"
                    비과세_성격 = "연한도비과세"
                if 월정액_산정됨:
                    사유 += f" [월정액급여 {int(mrw):,}원은 같은 입력 항목으로 산정 — 판단 참고]"
            else:
                과세액 = amount; 비과세액 = 0
                사유 = "생산직 등 요건 미해당 — 전액 과세"
                비과세_성격 = "과세(비대상)"

        elif k == "사택제공이익":
            position = facts.get("position", "")
            비과세대상 = r.get("비과세대상", [])
            비과세성격_v = r.get("비과세성격", "비과세")
            if position in 비과세대상:
                if 비과세성격_v == "근로소득 아님":
                    # 2016~2020년: 사택제공이익은 비과세 금액이나 비과세_합계가 아닌 근로소득_아님_합계에 산입
                    과세액 = 0; 비과세액 = 0
                    사유 = f"비과세 대상({position}) — 사택 제공 이익 근로소득에서 제외(비과세 성격: 근로소득 아님, 영 제38조①6호 단서)"
                    비과세_성격 = 비과세성격_v
                else:
                    비과세액 = amount; 과세액 = 0
                    사유 = f"비과세 대상({position}) — 사택 제공 이익 비과세(영 제38조①6호, 영 제17조의4 1호)"
                    비과세_성격 = 비과세성격_v
            else:
                과세액 = amount; 비과세액 = 0
                사유 = f"비과세 대상 아님({position}) — 전액 과세"
                비과세_성격 = "과세(비대상)"

        elif k == "주택자금대여이익":
            sme = facts.get("sme", False)
            related_party = facts.get("related_party", False)
            중소기업제외 = r.get("중소기업제외", False)
            특수관계제외 = r.get("특수관계제외", False)
            비과세성격_v = r.get("비과세성격", "근로소득 아님")
            if not sme or not 중소기업제외:
                # sme=False: 중소기업 종업원 아님 → 전액 과세
                # 중소기업제외=False(2016~2019): 중소기업 종업원이라도 예외 없음 → 전액 과세
                과세액 = amount; 비과세액 = 0
                사유 = "중소기업 종업원 아님 또는 해당 연도 중소기업 예외 없음 — 주택자금 대여이익 전액 과세(영 제38조①7호)"
                비과세_성격 = "과세(비대상)"
            elif year >= 2024 and 특수관계제외 and related_party:
                과세액 = amount; 비과세액 = 0
                사유 = "2024년부터 중소기업 종업원이라도 특수관계자(지배주주등·친족)는 비과세 제외 — 전액 과세(영 제17조의4 2호 단서, 부칙 2024.2.29)"
                비과세_성격 = "과세(특수관계)"
            else:
                if 비과세성격_v == "근로소득 아님":
                    # 2020년: 중소기업 종업원 주택자금 대여이익은 근로소득에서 제외(비과세 금액 아님)
                    과세액 = 0; 비과세액 = 0
                    사유 = "중소기업 종업원 주택자금 대여이익 — 근로소득에서 제외(비과세 성격: 근로소득 아님, 영 제38조①7호 단서)"
                    비과세_성격 = 비과세성격_v
                else:
                    비과세액 = amount; 과세액 = 0
                    사유 = "중소기업 종업원 주택자금 대여이익 — 비과세(영 제38조①7호 단서, 영 제17조의4 2호)"
                    비과세_성격 = 비과세성격_v

        elif k == "단체보험료":
            연한도 = r["연한도"]
            비과세액 = min(amount, 연한도)
            과세액 = amount - 비과세액
            사유 = f"단체순수보장성·환급부보장성 보험료 — 연 {연한도:,}원 이내 비과세(영 제38조①12호, 영 제17조의4 3호가목)"
            비과세_성격 = "연한도비과세"

        elif k == "학자금":
            beneficiary = facts.get("beneficiary", "")
            if beneficiary == "본인":
                req = facts.get("requirements_met", False)
                if req:
                    비과세액 = amount; 과세액 = 0
                    사유 = "근로자 본인 학자금 — 업무관련·지급기준·반납조건 3요건 충족, 비과세(영 제11조, 법 제12조3호아목)"
                    비과세_성격 = "요건비과세"
                else:
                    과세액 = amount; 비과세액 = 0
                    사유 = "근로자 본인 학자금이나 3요건(업무관련·지급기준·반납조건) 미충족 — 전액 과세"
                    비과세_성격 = "과세(요건 미충족)"
            else:
                과세액 = amount; 비과세액 = 0
                사유 = "가족(자녀 등) 학자금 — 전액 과세(영 제11조, 법 제12조3호아목)"
                비과세_성격 = "과세(가족 학자금)"

        elif k == "실비변상여비":
            actual_cost = facts.get("actual_cost")
            if actual_cost is None:
                return {"error": "입력 확인 필요 — 계산하지 않음",
                        "문제": ["실비변상여비: actual_cost(실비 정산 여부, true/false) 누락"],
                        "필요 입력": {"실비변상여비": ["actual_cost"]}}
            if not isinstance(actual_cost, bool):
                return {"error": "입력 확인 필요 — 계산하지 않음",
                        "문제": ["실비변상여비: actual_cost는 실비 정산 여부(true/false)여야 함 — 숫자가 입력됨"],
                        "필요 입력": {"실비변상여비": ["actual_cost"]}}
            if actual_cost:
                비과세액 = amount; 과세액 = 0
                사유 = "실비변상 실제 비용 정산 지급 — 전액 비과세(영 제12조 3호, 법 제12조3호자목)"
                비과세_성격 = "요건비과세"
            else:
                과세액 = amount; 비과세액 = 0
                사유 = "실비변상 실제 비용 정산 미지급 — 전액 과세(영 제12조 3호, 법 제12조3호자목)"
                비과세_성격 = "과세"

        elif k == "사회보험료_본인부담분_회사대납":
            과세액 = amount; 비과세액 = 0
            사유 = "근로자 본인 부담 사회보험료 회사 대납 — 근로소득 과세(법 제20조①1호, 법 제12조3호너목 단서)"
            비과세_성격 = "과세"

        elif k == "사회보험료_사용자부담분":
            비과세액 = amount; 과세액 = 0
            사유 = "사용자 부담 사회보험료 — 근로소득 아님(법 제12조3호너목)"
            비과세_성격 = "비과세"

        else:
            # 기타 과세 항목 (급여, 상여, 잉여금처분상여, 인정상여)
            과세액 = amount; 비과세액 = 0
            사유 = "과세 근로소득"
            비과세_성격 = "과세"

        # 항목별 kind는 vocab 정식 이름(별칭이 입력됐어도 정식 이름으로 출력)
        정식_kind = vocab.resolve(k)["이름"]

        항목별.append({
            "kind": 정식_kind,
            "금액": amount,
            "과세액": 과세액,
            "비과세액": 비과세액,
            "비과세_성격": 비과세_성격,
            "근거": b,
            "사유": 사유,
        })
        판정_idx[_idx] = 항목별[-1]
        총급여액 += 과세액
        비과세_합계 += 비과세액

        # 근로소득 아님 항목 집계: vocab 판단.비과세성격이 "근로소득 아님"인 항목만
        # 사택제공이익(2016~2020)·주택자금대여이익(2020) 등. 사회보험료_사용자부담분은
        #   vocab 판단.비과세성격 필드가 없고 판단.구분="비과세" → 비과세_합계에만 포함.
        if r.get("비과세성격") == "근로소득 아님":
            근로소득_아님_합계 += amount

    # 출력 순서는 입력 순서대로(연장근로수당을 뒤에서 판정했어도)
    _순번 = {id(v): i for i, v in 판정_idx.items()}
    항목별.sort(key=lambda x: _순번.get(id(x), 0))

    # ── 5. 근로소득공제·근로소득금액 ──
    근로소득공제 = I.earned_deduction(총급여액, year)
    근로소득금액 = 총급여액 - 근로소득공제

    if 직무발명보상금_비과세_사용액 > 0:
        산식_목록.append(f"직무발명보상금 비과세 사용액: {직무발명보상금_비과세_사용액:,}원 (기타소득 비과세 한도 공유용, 영 제18조②)")
    산식_목록.append(f"총급여액 = Σ 과세액 = {총급여액:,}원")
    산식_목록.append(f"근로소득공제 = earned_deduction({총급여액:,}원, {year}) = {근로소득공제:,}원 (소득세법 제47조①)")
    산식_목록.append(f"근로소득금액 = 총급여액 {총급여액:,}원 − 근로소득공제 {근로소득공제:,}원 = {근로소득금액:,}원")

    근거_목록.extend(b for _, _, b in [(it["kind"], it["금액"], it["근거"]) for it in 항목별])

    # ── 계산식 단계 (SPEC_export_workbook_f) ──────────────────────────────
    # 항목별 과세액을 계산식_입력으로 추가
    계산식_입력 = []
    항목식_parts = []
    월정액_단계: list[dict] = []
    연장_entry = 판정_idx.get(월정액_식["idx"]) if 월정액_식 else None
    위치 = {id(item): idx for idx, item in enumerate(항목별, 1)}
    for idx, item in enumerate(항목별, 1):
        과세액 = item["과세액"]
        항목식_parts.append(f"항목{idx}")
        if item is not 연장_entry:
            계산식_입력.append({"이름": f"항목{idx}", "값": 과세액, "설명": f"{item['kind']} 과세액 = {과세액:,}원"})
            continue
        # 연장근로수당: 월정액급여(영 제17조① 후단)를 같은 입력 항목으로 산정 → 비과세 요건 → 과세액까지 식으로
        산정 = 월정액_식["산정"]
        refs = []
        for m, (j, x) in enumerate(zip(산정["포함_i"], 산정["포함"]), 1):
            e = 판정_idx.get(j)
            if e is not None and e is not 연장_entry and e["과세액"] == x["금액"]:
                refs.append(f"항목{위치[id(e)]}")
            else:
                계산식_입력.append({"이름": f"월정액_포함{m}", "값": x["금액"], "설명": f"월정액급여 포함 — {x['kind']}: {x['사유']}"})
                refs.append(f"월정액_포함{m}")
        계산식_입력 += [
            {"이름": "월정액_근무월수", "값": 산정["근무월수"], "설명": f"근무월수 ({산정['월수_출처']})"},
            {"이름": "월정액급여한도", "값": 월정액_식["월정액한도"], "설명": "생산직 등 연장근로수당 비과세 월정액급여 한도(영 제17조①)"},
            {"이름": "직전연도총급여", "값": 월정액_식["직전"], "설명": "직전 과세기간 총급여액(입력)"},
            {"이름": "직전연도총급여한도", "값": 월정액_식["총급여한도"], "설명": "직전 과세기간 총급여액 한도(영 제17조①)"},
            {"이름": "연장근로수당_지급액", "값": 월정액_식["지급액"], "설명": "연장·야간·휴일근로수당 지급액"},
            {"이름": "광산일용여부", "값": 월정액_식["광산일용"], "설명": "광산근로자·일용근로자면 TRUE — 전액 비과세(영 제17조②1호)"},
            {"이름": "연장근로_연한도", "값": 월정액_식["연한도"] or 0, "설명": "연 비과세 한도(영 제17조②1호)"},
        ]
        월정액_단계 += [
            {"이름": "월정액_포함합계", "식": "+".join(refs) if refs else "0*월정액_근무월수", "값": 산정["포함_합계"],
             "설명": "매월 받는 봉급·수당 합계(상여 등 부정기 급여·실비변상적 급여" + ("·복리후생적 급여" if year >= 2021 else "") + " 제외, 연장근로 급여 차감)"},
            {"이름": "월정액급여", "식": "ROUNDDOWN(월정액_포함합계/월정액_근무월수,0)", "값": 산정["월정액급여"], "설명": "포함 합계 ÷ 근무월수(영 제17조① 후단)"},
            {"이름": "연장근로_비과세", "식": "IF(월정액_포함합계>월정액급여한도*월정액_근무월수,0,IF(직전연도총급여>직전연도총급여한도,0,"
                                           "IF(광산일용여부,연장근로수당_지급액,MIN(연장근로수당_지급액,연장근로_연한도))))",
             "값": item["비과세액"], "설명": "월정액급여·직전연도 총급여 요건 충족 시 연 한도 안에서 비과세(영 제17조①②)"},
            {"이름": f"항목{idx}", "식": "연장근로수당_지급액-연장근로_비과세", "값": 과세액, "설명": f"{item['kind']} 과세액"},
        ]

    # 근로소득공제 구간표 (2016~2026 동일)
    # 500만 이하 70%, 1500만 이하 40%, 4500만 이하 15%, 1억 이하 5%, 초과 2%
    # 2020년부터 한도 2,000만
    한도적용 = year >= 2020

    if len(항목식_parts) > 0:
        총급여_식 = "+".join(항목식_parts)
    else:
        총급여_식 = "0"

    # 근로소득공제 식 (중첩 IF)
    공제_식 = (
        f"IF(총급여액<=5000000,총급여액*0.7,"
        f"IF(총급여액<=15000000,5000000*0.7+(총급여액-5000000)*0.4,"
        f"IF(총급여액<=45000000,5000000*0.7+10000000*0.4+(총급여액-15000000)*0.15,"
        f"IF(총급여액<=100000000,5000000*0.7+10000000*0.4+30000000*0.15+(총급여액-45000000)*0.05,"
        f"5000000*0.7+10000000*0.4+30000000*0.15+55000000*0.05+(총급여액-100000000)*0.02))))"
    )
    if 한도적용:
        공제_식 = f"MIN({공제_식},20000000)"

    계산식 = 월정액_단계 + [
        {"이름": "총급여액", "식": 총급여_식, "값": 총급여액,
         "설명": f"총급여액 = Σ 항목별 과세액 = {총급여액:,}원 (비과세 판정 후 금액 합계)"},
        {"이름": "근로소득공제", "식": 공제_식, "값": 근로소득공제,
         "설명": f"근로소득공제 = 구간별 공제율 적용 (500만↓70%, 1500만↓40%, 4500만↓15%, 1억↓5%, 초과2%) {'+ 한도 2,000만(2020~)' if 한도적용 else '(한도 없음)'} = {근로소득공제:,}원 (소득세법 제47조①)"},
        {"이름": "근로소득금액", "식": "총급여액-근로소득공제", "값": 근로소득금액,
         "설명": f"근로소득금액 = 총급여액 − 근로소득공제 = {총급여액:,} − {근로소득공제:,} = {근로소득금액:,}원"},
    ]

    return _ok({
        "총급여액": 총급여액,
        "비과세_합계": 비과세_합계,
        "근로소득_아님_합계": 근로소득_아님_합계,
        "근로소득공제": 근로소득공제,
        "근로소득금액": 근로소득금액,
        "직무발명보상금_비과세_사용액": 직무발명보상금_비과세_사용액,
    }, 근거_목록,
       항목별=항목별,
       제외_항목=제외_항목,
       산식=산식_목록,
       계산식_입력=계산식_입력,
       계산식=계산식,
       **({"판단": 판단_목록} if 판단_목록 else {}))


# ───────── 세무조사 판정 도구 11개 ─────────
@mcp.tool(description="""두 당사자의 특수관계 여부·호수를 국세기본법·법인세법·상증법 기준으로 판정한다. 언제: 세무조사·증여의제 검토에서 두 사람·법인 간 특수관계와 적용 호수를 확인할 때.
입력: 당사자 a(이름) → a, 당사자 b(이름) → b, 사실관계 시점(YYYY-MM-DD, 필수) → on, people={이름: {구분: "개인"|"법인"}} → people, family=[[a, b, "부모자녀"|"배우자"]] → family, stakes=[[보유자, 법인, 지분율(0~1)]] → stakes, officers={법인: [이름]} → officers, employees={법인: [이름]} → employees, livelihood=[[부양자, 피부양자]] → livelihood, 생계를 함께하는 친족 사실관계 [{친족, 상대방, 주민등록_동일세대, 동거_사실, 생활비_공동부담, 독립생계_소득, 경제적_지원, 별거_사유, 별도_임대차, 별거_증빙_유형, 연말정산_부양가족_등록, 수입_상대방계좌_입금, 혼인_독립가구}] → livelihood_shared, influence=[[사실상 영향력자, 법인]] → influence, group={기업집단: [계열회사]} → group, group_exempt={법인: 공정위 통지 내용} → group_exempt, 법령 인용 포함 여부 → with_citation.
생계를 함께하는 친족(국기령 제1조의2②3호, 법인령 제2조⑧3호 — 2019.2.12.~2023.12.31. 제2조⑤3호, 그 전 제87조①3호): 전제는 상대방이 다른 당사자 본인의 임원·사용인(같은 법인 동료는 아님) 또는 생계유지자, 친족이 기준일 친족 범위(2023.3.1.부터 4촌 혈족·3촌 인척) 안. 1차 기준은 동거 불요·생활비 공동 부담: 생활비 공동부담, 무소득+정기송금, 수입이 상대방 계좌로 입금이면 생계 공동 / 안정 소득+따로 거주 또는 별도 세대+생활비 각자+지원 없음, 제3자 지원(금융자료), 혼인 독립 가구+주소 다름(소득 없으면 제외)이면 생계 별도 / 주민등록 동일세대는 객관 자료(진술서만은 안 됨)로 반증해야 하고, 동일세대+소득만·동거만·부양가족 공제 이력이면 생계 공동 추정(인정되는 증빙 안내) / 그 밖은 입력 확인 필요(확인할 항목·필요 자료). 상증법에는 이 항목 없음.
결과: 특수관계 여부·적용 호수와 근거 조문(기준일 시행 법령 위치), 근거, 생계를_함께하는_친족(결론·사유·결정·해석 번호). 두 당사자의 특수관계 여부와 적용 호수를 국세기본법·법인세법·상증법 기준으로 판정한다(두 사람·법인 간 특수관계와 적용 호수 확인).""")
def related_judge(a: Annotated[str, Field(description='당사자 a(이름)')], b: Annotated[str, Field(description='당사자 b(이름)')],
                  on: Annotated[str, Field(description='사실관계 시점 YYYY-MM-DD(필수, 친족 범위·조문 위치가 시점마다 다름)')],
                  people: Annotated[dict, Field(description='사람·법인 목록 {이름: {"구분": "개인"|"법인"}. 예: {"김갑": {"구분": "개인"}, "㈜대한": {"구분": "법인"}}')],
                  family: Annotated[list, Field(description='친족 관계 목록 [[a, b, "부모자녀"|"배우자"]]. 예: [["김갑", "김을", "부모자녀"]] — 부모자녀는 [부모, 자녀] 순서')],
                  stakes: Annotated[list, Field(description='직접보유비율 목록 [[보유자, 법인, 지분율(0~1)]]. 예: [["김갑", "㈜대한", 0.30]]')],
                  officers: Annotated[dict | None, Field(description='임원 목록 {법인: [이름]}. 예: {"㈜대한": ["김갑"]}')] = None,
                  employees: Annotated[dict | None, Field(description='직원 목록 {법인: [이름]}. 예: {"㈜대한": ["김을"]}')] = None,
                  livelihood: Annotated[list | None, Field(description='부양 관계 [[부양자, 피부양자]]. 예: [["김갑", "김이쁜"]]')] = None,
                  livelihood_shared: Annotated[list | None, Field(description='생계를 함께하는 친족 사실관계 [{친족, 상대방(임원·사용인·생계유지자), 주민등록_동일세대: true|false|모름, 동거_사실: 같은집_공유구조|같은집_독립구조|따로거주|모름, 생활비_공동부담: 공동|각자분담_증빙있음|각자분담_증빙없음|일방지원|모름, 독립생계_소득: 안정소득_있음|소액_일시|없음|모름, 경제적_지원: 정기송금_있음|차입후상환|용돈수준|제3자_지원|없음|모름, 별거_사유: 취학|질병요양|근무|사업|혼인파탄|독립|해당없음, 별도_임대차: bool, 별거_증빙_유형: [공과금·관리비_명의|금융_송금|출입국|전화·인터넷_가입|임대차_중개확인|진술서만|없음], 연말정산_부양가족_등록: bool, 수입_상대방계좌_입금: bool, 혼인_독립가구: bool, 친족관계?: 배우자|혈족|인척, 촌수?}]. 예: [{"친족": "이배우", "상대방": "김대표", "생활비_공동부담": "공동"}]')] = None,
                  influence: Annotated[list | None, Field(description='사실상 영향력 [[사실상 영향력자, 법인]]. 예: [["김갑", "㈜대한"]]')] = None,
                  group: Annotated[dict | None, Field(description='기업집단 {기업집단: [계열회사]}. 예: {"XX그룹": ["㈜대한", "㈜민국"]}')] = None,
                  group_exempt: Annotated[dict | None, Field(description='공정위 계열편입 유예·제외 통지 {법인: 통지 내용}. 예: {"㈜대한": "유예 통지"}')] = None,
                  with_citation: Annotated[bool, Field(description='법령 인용 포함 여부(기본 True)')] = True) -> dict:
    from .audit.agents.related_intake import normalize, to_case
    try:
        case = to_case(normalize({"people": people, "family": family, "stakes": stakes, "officers": officers, "employees": employees,
                                  "livelihood": livelihood, "livelihood_shared": livelihood_shared, "influence": influence, "group": group}))
        case["group_exempt"] = {norm_name(k): v for k, v in (group_exempt or {}).items()}
        a, b = norm_name(a), norm_name(b)
        return _R.judge_all_at(case, a, b, on) if with_citation else _R.judge_all(case, a, b, on)
    except ValueError as e:
        return {"error": "입력 확인 필요 — 판정하지 않음", "문제": [str(e)]}


@mcp.tool(description="""두 개인의 친족관계(혈족·인척 촌수, 배우자)와 기준일 친족 범위 해당 여부. 언제: 두 사람이 친족인지, 국세기본법·상증법상 친족 범위에 들어가는지 확인할 때.
입력: 당사자 a(이름) → a, 당사자 b(이름) → b, 기준일(YYYY-MM-DD) → on, family=[[a, b, "부모자녀"|"배우자"]] → family(related_judge와 동일 형식).
결과: 관계(촌수·배우자)·기준일·친족 범위·국기법 친족 여부·상증법 친족(사돈 포함) 여부. 두 개인의 친족관계(혈족·인척 촌수, 배우자)와 기준일 친족 범위 해당 여부를 확인한다.""")
def kinship_check(a: Annotated[str, Field(description='당사자 a(이름)')], b: Annotated[str, Field(description='당사자 b(이름)')],
                  on: Annotated[str, Field(description='기준일 YYYY-MM-DD')],
                  family: Annotated[list, Field(description='친족 관계 목록 [[a, b, "부모자녀"|"배우자"]]. 예: [["김갑", "김을", "부모자녀"]] — related_judge와 동일 형식')]) -> dict:
    fam = [tuple(x) for x in family]
    blood, inlaw = _R.kin_limits(on)
    k = _R.kinship(fam, a, b)
    return {"관계": k, "기준일": on, "범위": f"혈족 {blood}촌·인척 {inlaw}촌·배우자",
            "국기법_친족": _R.is_kin(fam, a, b, "국기", on), "상증법_친족(사돈 포함)": _R.is_kin(fam, a, b, "상증", on)}


@mcp.tool(description="""보유자→법인 직접·간접 보유비율(경로별 곱 합산, 순환출자 반영). 언제: 일감몰아주기·지배주주 판정에서 특정인이 법인에 얼마나 간접 보유하는지 계산할 때.
입력: 보유자(이름) → holder, 대상 법인(이름) → target, stakes=[[보유자, 법인, 지분율(0~1)]] → stakes(직접보유비율, related_judge와 동일).
결과: 직접비율·경로별 비율·합계·순환출자 여부, 근거(상증법 시행령 제34조의3②). 보유자→법인 직접·간접 보유비율을 계산한다(간접은 단계별 직접보유비율의 곱, 경로별 합산, 순환출자는 단계별 모두 반영).""")
def ownership_ratio(holder: Annotated[str, Field(description='보유자(이름)')], target: Annotated[str, Field(description='대상 법인(이름)')],
                    stakes: Annotated[list, Field(description='직접보유비율 목록 [[보유자, 법인, 지분율(0~1)]]. 예: [["김갑", "㈜대한", 0.30]]')]) -> dict:
    st = [(h, c, float(r)) for h, c, r in stakes]
    ps = _R.paths(st, holder, target)
    return {"직접": _R.direct(st, holder, target), "경로": [{"경로": " → ".join(p), "비율": r} for p, r in ps],
            "합계": sum(r for _, r in ps), "순환출자": _R.has_cycle(st)}


@mcp.tool(description="""수혜법인의 지배주주 판정(상증법 시행령 제34조의3①). 언제: 일감몰아주기 증여의제이익 계산에서 수혜법인의 지배주주가 누구인지 판정할 때.
입력: 수혜법인(이름) → corp, people={이름: {구분: "개인"|"법인"}} → people, family=[[a, b, "부모자녀"|"배우자"]] → family, stakes=[[보유자, 법인, 지분율(0~1)]] → stakes(related_judge와 동일 형식).
결과: 지배주주(이름)·판정 이유. 수혜법인의 지배주주를 판정한다Amid(중견기업 중 직접보유 최고자 / 법인이면 직접+간접 최고 개인).""")
def dominant_shareholder(corp: Annotated[str, Field(description='수혜법인(이름)')],
                         people: Annotated[dict, Field(description='사람·법인 목록 {이름: {"구분": "개인"|"법인"}}. 예: {"김갑": {"구분": "개인"}}')],
                         family: Annotated[list, Field(description='친족 관계 [[a, b, "부모자녀"|"배우자"]]. 예: [["김갑", "김을", "부모자녀"]]')],
                         stakes: Annotated[list, Field(description='직접보유비율 [[보유자, 법인, 지분율(0~1)]]. 예: [["김갑", "㈜대한", 0.30]]')]) -> dict:
    who, why = _R.dominant_shareholder(people, [tuple(x) for x in family], [(h, c, float(r)) for h, c, r in stakes], corp)
    return {"지배주주": who, "판정": why}


@mcp.tool(description="""일감몰아주기 증여의제이익(상증법 제45조의3, ⑩항 과세제외매출액 직접 입력, ⑭·⑮ 미반영). 언제: 일감몰아주기 검토에서 증여의제이익을 간단히 계산할 때(국세청 신고안내 전체 산식은 tunnelling_gift_nts 사용).
입력: 수혜자(이름) → beneficiary, 기업 규모(중소|중견|일반) → size, 총 매출액 → sales_total, sales_by_corp={거래처 법인: 매출액} → sales_by_corp, 세후영업이익 → after_tax_op_profit, people={이름: {구분: "개인"|"법인"}} → people, family=[[a, b, "부모자녀"|"배우자"]] → family, stakes=[[보유자, 법인, 지분율(0~1)]] → stakes, ⑩·⑭항 과세제외매출액 → excluded_sales.
결과: 증여의제이익, 근거(상증법 제45조의3, 시행령 제34조의3). 일감몰아주기 증여의제이익(상증법 제45조의3).""")
def tunnelling_gift(beneficiary: Annotated[str, Field(description='수혜자(이름)')],
                    size: Annotated[str, Field(description="기업 규모('중소'|'중견'|'일반')")],
                    sales_total: Annotated[float, Field(description="총 매출액(원)")],
                    sales_by_corp: Annotated[dict, Field(description='거래처 법인별 매출액 {법인: 매출액(원)}. 예: {"A": 300000000}')],
                    after_tax_op_profit: Annotated[float, Field(description="세후영업이익(원)")],
                    people: Annotated[dict, Field(description='사람·법인 목록 {이름: {"구분": "개인"|"법인"}}')],
                    family: Annotated[list, Field(description='친족 관계 [[a, b, "부모자녀"|"배우자"]]')],
                    stakes: Annotated[list, Field(description='직접보유비율 [[보유자, 법인, 지분율(0~1)]]')],
                    excluded_sales: Annotated[float, Field(description="⑩·⑭항 과세제외매출액(원)")] = 0) -> dict:
    return _R.tunnelling(people, [tuple(x) for x in family], [(h, c, float(r)) for h, c, r in stakes], beneficiary,
                        sales_total, sales_by_corp, after_tax_op_profit, size, excluded_sales)


@mcp.tool(description="""일감몰아주기 증여의제이익 — 국세청 2026 신고안내 산식(세후영업이익 ⑫, 출자관계별 추가 과세제외 ⑭1·3호, 한계보유 ⑬, 배당공제 ⑮). 언제: 국세청 신고안내 공식에 따라 일감몰아주기 증여의제이익을 정확히 계산할 때.
입력: 기업 규모(중소|중견|일반) → size, 수혜법인(이름, 고정) → '수혜법인', 세후영업이익 ⑫(원) → op_income, 과세표준(원) → taxable_income, 법인세 산출세액(원) → tax, 총 매출액(원) → sales_total, sales_by_corp={거래처 법인: 매출액(원)} → sales_by_corp, relations={관계명: 보유비율(소수)}(예: {'직접':0.2,'A경유':0.18}) → relations, 기본 과세제외매출액(⑩항, 원) → base_excluded, holdings_in_related={특관법인: 지배주주등 보유비율(소수)} → holdings_in_related, indirect_corp_of={관계명: 그 경로 간접출자법인} → indirect_corp_of, excluded_by_corp={특관법인: ⑩항 과세제외액(원)} → excluded_by_corp, dividends={관계명: [배당소득(원), 분모]}(⑮항) → dividends, 토지등 양도소득 법인세(차감, 원) → land_tax, 공제·감면세액(차감, 원) → credits, 미환류소득 법인세(원, tax에 포함이 맞음) → unreturned_tax.
결과: 증여의제이익·단계별 계산, 근거(상증법 제45조의3, 시행령 제34조의3, 국세청 2026 신고안내). 일감몰아주기 증여의제이익 — 국세청 2026 신고안내 산식.""")
def tunnelling_gift_nts(size: Annotated[str, Field(description="기업 규모('중소'|'중견'|'일반')")],
                        op_income: Annotated[float, Field(description="세후영업이익 ⑫(원)")],
                        taxable_income: Annotated[float, Field(description="과세표준(원)")],
                        tax: Annotated[float, Field(description="법인세법 제55조 산출세액(원)")],
                        sales_total: Annotated[float, Field(description="총 매출액(원)")],
                        sales_by_corp: Annotated[dict, Field(description='거래처 법인별 매출액 {법인: 매출액(원)}. 예: {"A": 300000000}')],
                        relations: Annotated[dict, Field(description="관계명별 보유비율(소수). 예: {'직접': 0.2, 'A경유': 0.18}")],
                        base_excluded: Annotated[float, Field(description="기본 ⑩항 과세제외매출액(원)")] = 0,
                        holdings_in_related: Annotated[dict | None, Field(description='특관법인별 지배주주등 보유비율(소수). 예: {"A": 0.3}')] = None,
                        indirect_corp_of: Annotated[dict | None, Field(description='관계명별 간접출자법인. 예: {"A경유": "B법인"}')] = None,
                        excluded_by_corp: Annotated[dict | None, Field(description='특관법인별 ⑩항 과세제외액(원). 예: {"A": 10000000}')] = None,
                        dividends: Annotated[dict | None, Field(description='관계명별 [배당소득(원), 분모]. 예: {"직접": [10000000, 300000000]}')] = None,
                        land_tax: Annotated[float, Field(description="토지등 양도소득 법인세(차감, 원)")] = 0,
                        credits: Annotated[float, Field(description="공제·감면세액(차감, 원)")] = 0,
                        unreturned_tax: Annotated[float, Field(description="미환류소득 법인세(원, tax에 포함이 맞음)")] = 0) -> dict:
    return _R.tunnelling_nts("수혜법인", size, op_income, taxable_income, tax, sales_total, sales_by_corp, relations,
                            base_excluded=base_excluded, holdings_in_related=holdings_in_related, indirect_corp_of=indirect_corp_of,
                            dividends={k: tuple(v) for k, v in (dividends or {}).items()},
                            excluded_by_corp=excluded_by_corp,
                            land_tax=land_tax, credits=credits, unreturned_tax=unreturned_tax)


@mcp.tool(description="""특수관계인 범위 조문이 기준일에 어느 법령에 있었는지 확인. 언제: 특수관계 판정 결과를 근거로 인용할 때 해당 기준일 시행 조문의 위치를 확인한다.
입력: 법령 구분(국기|법인|상증) → law, 기준일(YYYY-MM-DD) → on. 예: law='법인', on='2019-02-11' → 법인세 시행령 제87조①.
결과: 법령명·조항·원문·기준일. 특수관계인 범위 조문이 기준일에 어디 있었는지 확인한다(law=국기|법인|상증).""")
def related_provision_at(law: Annotated[str, Field(description="법령 구분('국기'|'법인'|'상증')")],
                         on: Annotated[str, Field(description="기준일 YYYY-MM-DD")]) -> dict:
    name, jo, head = _R.cite_at(law, on)
    return {"법령": name, "조항": jo, "원문": head, "기준일": on}


@mcp.tool(description="""주주명부·가족관계·임원명단 PDF에서 관계표(people/family/stakes/officers)를 추출한다. 언제: 세무조사·의료법인 친족확인 등에서 PDF 문서에 있는 관계표를 도구 입력으로 정리할 때.
입력: PDF 파일 경로 → pdf_path.
결과: people·family·stakes·오 officers 추출 결과(초안) + AI 생성 표시(필요 시). related_judge에 넣기 전에 사람이 확인한다. UPSTAGE_API_KEY 없이도 pypdf 로컬 추출 후 정리 가능(기본), 키 설정 시 Document Parse + Solar Pro 4로 더 정교한 표 추출. 주주명부·가족관계·임원명단 PDF에서 관계표를 뽑는다.""")
@_pdf_guard
def extract_relations(pdf_path: Annotated[str, Field(description="PDF 파일 경로")]) -> dict:
    from .audit.agents import related_intake
    return _add_ai_mark(related_intake.from_pdf(pdf_path))


@mcp.tool(description="""일감몰아주기 검토 자료 PDF → 표 추출 → 코드가 지배주주·출자관계·증여의제이익 계산. 언제: PDF로 된 일감몰아주기 검토 자료에서 관계표를 추출해 tunnelling_gift_nts 입력으로 정리할 때.
입력: PDF 파일 경로 → pdf_path.
결과: 지배주주·출자관계·증여의제이익 계산 결과(초안) + AI 생성 표시(필요 시). UPSTAGE_API_KEY 없이도 pypdf 로컬 추출 후 정리 가능(기본), 키 설정 시 Document Parse + Solar Pro 4로 더 정교한 표 추출. 결과는 사람 확인. 일감몰아주기 검토 자료 PDF에서 표를 뽑는다.""")
@_pdf_guard
def tunnelling_from_pdf(pdf_path: Annotated[str, Field(description="PDF 파일 경로")]) -> dict:
    from .audit.agents import tunnel_intake
    return _add_ai_mark(tunnel_intake.from_pdf(pdf_path))


@mcp.tool(description="""신고서 사전검토(조사 착수 전 오류 의심 목록). 언제: 세무조사 착수 전에 신고서 값으로 오류 의심 항목을 미리 점검할 때. 입력: data={입력키: 값}, tax=법인|부가|소득(생략 시 전체) — 입력키는 precheck_schema로 먼저 확인. 결과: 오류 의심·이상 없음·자료 부족·사실 확인 필요.""")
def return_precheck(data: Annotated[dict, Field(description='신고서 입력키·값 사전(dictionary). 입력키는 precheck_schema로 확인')], tax: Annotated[str | None, Field(description="세목('법인'|'부가'|'소득'), 생략 시 전체")] = None) -> dict:
    return _precheck_run(data, tax)


@mcp.tool(description="""신고서 사전검토 입력스키마. 언제: return_precheck·영어로 된 신고서 PDF 도구 입력 전에 어떤 키를 넣어야 하는지 확인할 때. 입력: tax=법인|부가|소득. 결과: 각 입력키의 의미·단위·필수 여부를 목록으로 반환한다.""")
def precheck_schema(tax: Annotated[Literal["법인", "부가", "소득"], Field(description="사전검토 세목('법인'|'부가'|'소득')")] ) -> dict:
    from .audit.agents.precheck import REGISTRY
    from .audit.agents.precheck import corp as _c, vat as _v, income as _i

    schema_map = {"법인": _c.SCHEMA, "부가": _v.SCHEMA, "소득": _i.SCHEMA}
    schema = schema_map[tax]

    # 해당 세목 규칙의 needs에 등장하는 키 → '요건별 필수'로 표시
    needed = set()
    prefix = {"법인": "법인", "부가": "부가", "소득": "소득"}[tax]
    for r in REGISTRY:
        if r.세목 != prefix: continue
        for k in r.needs:
            needed.add(k)

    rows = []
    for k, desc in schema.items():
        unit = _infer_unit(desc)
        required = "요건별 필수" if k in needed else "선택"
        rows.append({"키": k, "의미": desc, "단위": unit, "필수": required})
    return {"세목": tax, "입력스키마": rows, "총_입력키": len(rows)}


def _infer_unit(desc: str) -> str:
    """SCHEMA 설명에서 단서를 읽어 단위를 추론한다."""
    if "bool" in desc or "여부" in desc: return "bool"
    if "YYYY-MM-DD" in desc or "제출일" in desc or "개시일" in desc or "종료일" in desc or "일" in desc: return "YYYY-MM-DD/string"
    if "리스트" in desc or "[" in desc: return "list[dict]/list"
    if "원" in desc: return "원 (int)"
    if "소수" in desc or "비율" in desc: return "소수 (float)"
    if "string" in desc or "구분" in desc or "코드" in desc: return "문자열"
    if "명" in desc or "수" in desc: return "int (명/수)"
    return "문맥 참조"


@mcp.tool(description="""법인세 신고서 PDF(별지 1·3·50호) → 표 추출 → 사전검토(오류 의심 목록). 언제: 법인세 신고서 PDF에서 값을 추출해 return_precheck 입력으로 정리할 때.
입력: PDF 파일 경로 → pdf_path.
결과: 오류 의심·이상 없음·자료 부족·사실 확인 필요 목록 + 서식 간 대사·세액 체인 재계산 + AI 생성 표시(필요 시). UPSTAGE_API_KEY 없이도 pypdf 로컬 추출 후 return_precheck 입력으로 정리 가능(기본), 키 설정 시 Document Parse + Solar Pro 4로 더 정교한 표 추출. 법인세 신고서 PDF에서 표를 추출해 사전검토한다.""")
@_pdf_guard
def return_precheck_pdf(pdf_path: Annotated[str, Field(description="PDF 파일 경로")]) -> dict:
    from .audit.agents import return_intake
    return _add_ai_mark(return_intake.precheck_pdf(pdf_path))


# ───────── 끝에서 끝까지 계산기 시범 3종 ─────────

def _loss_rate_floor(rate: float, digits: int = 3) -> float:
    """대손실적률 소수 자릿수 절사(넷째 자리 이하 버림)."""
    return math.floor(rate * 10 ** digits) / 10 ** digits


@mcp.tool(annotations=CALC)
@_guard
def bad_debt_adjustment(year: YEAR, sme: Annotated[bool, Field(description="중소기업 여부")],
                        book: Annotated[dict, Field(description='장부상 대손충당금 {opening: 기초, offset: 당기상계, provision: 당기설정, closing: 기말}(원)')],
                        receivables_prev_end_book: WON,
                        receivables_end_book: WON,
                        writeoffs: Annotated[list, Field(description="대손금 목록 [{kind, amount, arose_in:'전기'|'당기', instrument?:'외상매출금'|'받을어음'|'기타', months_since_default?, mortgage?:bool}]")],
                        prior_excess: Annotated[int, Field(description="전기 대손충당금 한도초과액(원)")] = 0,
                        loss_rate_digits: Annotated[int, Field(description="대손실적률 소수 자릿수(기본 3 — 넷째 자리 이하 절사)")] = 3) -> dict:
    """대손금·대손충당금 세무조정(법인세법 제19조의2·제34조, 시행령 제19조의2·제61조).
    언제: 법인세 세무조정에서 대손금 인정 여부와 대손충당금 한도·국회초과를 함께 계산할 때.
    입력: 사업연도(2016~2026) → year, 중소기업 여부 → sme, 전기 대손충당금 한도초과액(원) → prior_excess,
    장부상 대손충당금 {opening, offset, provision, closing}(원) → book,
    장부상 설정대상 채권 전기말·당기말(원) → receivables_prev_end_book·receivables_end_book,
    대손금 목록 [{kind, amount, arose_in, instrument?, months_since_default?, mortgage?}] → writeoffs,
    대손실적률 소수 자릿수(기본 3) → loss_rate_digits.
    결과: 판단(항목별·근거)·계산(세법상 당기 대손금·전기말·당기말 채권·실적률)·결과(한도·한도초과)·세무조정. 대손금·대손충당금 세무조정."""

    _v = _validate_items(writeoffs, "kind", _specs_from_vocab("bad_debt_adjustment", "kind"))
    if _v: return _v
    # ── 대손금 인정 판단 ──
    judgments = []  # 항목별 판단
    tax_bad_debt_recognized = 0   # 세법상 당기 대손금(인정분)
    prior_year_reduced = 0        # 전기 사유 신고조정 대손 → 전기말 채권 차감
    denied_bad_debt = 0           # 당기 부인 대손 합계
    current_disallowed_entries = []  # 당기 손금불산입(유보) 항목
    memo_total = 0                   # 비망가액 합계: 인정 채권별 min(금액, 1,000원) (CodeSolar 리뷰 반영)

    for w in writeoffs:
        kind = w["kind"]; amount = int(w["amount"]); arose_in = w["arose_in"]
        instrument = w.get("instrument"); months = w.get("months_since_default"); mortgage = w.get("mortgage", False)

        if kind == "화해·확정판결":
            # 영 제19조의2①10호(재판상 화해 등 확정판결과 같은 효력)는 ③1호 열거 밖 → 결산조정: 손비로 계상한 사업연도 손금 (Claude 10-13 정정)
            tax_bad_debt_recognized += amount
            judgments.append({"항목": f"{kind} {amount:,}원", "결론": "인정(대손금, 결산조정)",
                              "금액": amount, "근거": "법인세법 시행령 제19조의2①10호·③2호",
                              "이유": "확정판결과 같은 효력으로 회수불능 확정 — 손비로 계상한 사업연도에 손금"})
            continue
        if kind == "소멸시효완성":
            # 신고조정 대손: 사유 발생 사업연도 손금. arose_in='전기'인데 당기 장부 처리 → 당기 손금불산입(유보), 전기말 채권 차감
            if arose_in == "전기":
                # 전기 손금이었어야 할 것을 당기에 장부 처리 → 당기 손금불산입(유보), 전기 신고조정 추심
                prior_year_reduced += amount
                judgments.append({"항목": f"{kind} {amount:,}원(전기 사유, 당기 장부 처리)",
                                  "결론": "손금불산입(유보)·전기 신고조정 추심",
                                  "금액": amount, "근거": "법인세법 제19조의2·시행령 제19조의2(소멸시효완성 등 신고조정 대손)",
                                  "이유": f"{arose_in} 사유 발생, 전기 손금 귀속분이나 당기 장부 처리 → 당기 손금불산입(유보)"})
                current_disallowed_entries.append({"구분": "전기 사유분", "금액": amount, "소득처분": "유보"})
            else:
                # 당기 사유 → 세법상 인정
                tax_bad_debt_recognized += amount
                judgments.append({"항목": f"{kind} {amount:,}원(당기 사유)", "결론": "인정(대손금)",
                                  "금액": amount, "근거": "법인세법 제19조의2·시행령 제19조의2",
                                  "이유": "사유 발생 사업연도 손금 귀속, 당기 대손금 인정"})

        elif kind == "회수불능":
            # 영 제19조의2①8호(파산·강제집행·사업폐지·사망·실종 등) — 결산조정, 손비 계상 사업연도 손금
            tax_bad_debt_recognized += amount
            judgments.append({"항목": f"{kind} {amount:,}원", "결론": "인정(대손금, 결산조정)",
                              "금액": amount, "근거": "법인세법 시행령 제19조의2①8호·③2호",
                              "이유": "채무자 파산·강제집행·사업폐지·사망·실종 등으로 회수할 수 없는 채권 — 손비로 계상한 사업연도에 손금"})
        elif kind == "부도6개월" and months is not None and int(months) < 6:
            # 부도발생일부터 6개월이 지나지 않음 → 요건 미충족, 부인(유보) (CodeSolar 리뷰 반영)
            denied_bad_debt += amount
            judgments.append({"항목": f"부도6개월 {instrument or '기타'} {amount:,}원(부도 후 {months}개월)",
                              "결론": "부인(손금불산입·유보) — 6개월 미경과",
                              "금액": amount, "근거": "법인세법 시행령 제19조의2①9호",
                              "이유": "부도발생일부터 6개월 이상 지나야 대손금 — 요건 미충족"})
            current_disallowed_entries.append({"구분": "부인 대손(6개월 미경과)", "금액": amount, "소득처분": "유보"})
        elif kind == "부도6개월":
            if mortgage:
                # 저당권 설정 시 불인정
                denied_bad_debt += amount
                judgments.append({"항목": f"부도6개월 {instrument or '기타'} {amount:,}원(저당권 설정)",
                                  "결론": "불인정(저당권으로 인한 부인, 손금불산입·유보)",
                                  "금액": amount, "근거": "법인세법 제19조의2·시행령 제19조의2(저당권 설정 부도6개월분 불인정)",
                                  "이유": "저당권 설정 사유로 부도6개월 대손금 불인정"})
                current_disallowed_entries.append({"구분": "부인 대손(저당권)", "금액": amount, "소득처분": "유보"})
            elif instrument == "받을어음":
                # 받을어음: 중소기업·일반 모두 인정 (비망 1,000원 공제)
                recognized = max(0, amount - 1000); memo_total += min(amount, 1000)
                tax_bad_debt_recognized += recognized
                judgments.append({"항목": f"부도6개월 받을어음 {amount:,}원",
                                  "결론": f"인정(대손금, 비망 1,000원 공제 → {recognized:,}원)",
                                  "금액": recognized, "근거": "법인세법 제19조의2·시행령 제19조의2(부도발생 후 6개월 경과 받을어음)",
                                  "이유": "받을어음 부도6개월분 인정, 비망 1,000원 공제"})
            elif instrument == "외상매출금":
                if sme:
                    # 중소기업: 부도 이전 발생 외상매출금 인정 (비망 1,000원 공제)
                    recognized = max(0, amount - 1000); memo_total += min(amount, 1000)
                    tax_bad_debt_recognized += recognized
                    judgments.append({"항목": f"부도6개월 외상매출금 {amount:,}원(중소기업)",
                                      "결론": f"인정(대손금, 비망 1,000원 공제 → {recognized:,}원)",
                                      "금액": recognized, "근거": "법인세법 제19조의2·시행령 제19조의2(중소기업 외상매출금 부도6개월분)",
                                      "이유": "중소기업 외상매출금 부도6개월분 인정, 비망 1,000원 공제"})
                else:
                    # 일반법인: 외상매출금 부도6개월분 불인정
                    denied_bad_debt += amount
                    judgments.append({"항목": f"부인 대손(외상매출금) {amount:,}원(일반법인)",
                                      "결론": "부인(손금불산입·유보)",
                                      "금액": amount, "근거": "법인세법 제19조의2·시행령 제19조의2(중소기업 외 외상매출금 부도6개월분 불인정)",
                                      "이유": "중소기업 아닌 일반법인, 외상매출금 부도6개월분 손금 불인정 — 부인"})
                    current_disallowed_entries.append({"구분": "부인 대손(외상매출금)", "금액": amount, "소득처분": "유보"})
            else:
                # instrument 미지정 기타
                if sme:
                    recognized = max(0, amount - 1000); memo_total += min(amount, 1000)
                    tax_bad_debt_recognized += recognized
                    judgments.append({"항목": f"부도6개월 기타 {amount:,}원(중소기업)", "결론": f"인정(비망 1,000원 공제 → {recognized:,}원)",
                                      "금액": recognized, "근거": "법인세법 제19조의2·시행령 제19조의2"})
                else:
                    denied_bad_debt += amount
                    judgments.append({"항목": f"부도6개월 기타 {amount:,}원(일반법인)", "결론": "불인정(손금불산입·유보)",
                                      "금액": amount, "근거": "법인세법 제19조의2·시행령 제19조의2"})
                    current_disallowed_entries.append({"구분": "부인 대손(기타)", "금액": amount, "소득처분": "유보"})

        else:
            judgments.append({"항목": f"{kind} {amount:,}원", "결론": "검토 필요", "금액": amount,
                              "근거": "종류 미정", "이유": f"kind={kind} — 세무 판단 기준 확인 필요"})

    # 세법상 전기말 채권 = 장부 전기말 − 전기 사유 신고조정 대손
    tax_receivables_prev = receivables_prev_end_book - prior_year_reduced
    # 세법상 당기말 채권 = 장부 당기말 + 부인 채권 + 비망가액
    # (부인 채권 = 인정되지 않은 대손금, 비망가액은 인식된 부도6개월 채권의 1,000원)
    # 비망가액 합계 계산 필요 — 인정된 부도6개월 채권별 1,000원
    reservation_total = memo_total   # 인정된 부도6개월 채권별 비망가액 min(금액, 1,000원)
    # 부인 채권: 인정되지 않은 writeoff
    tax_receivables_end = receivables_end_book + denied_bad_debt + reservation_total

    # 대손실적률 = 세법상 당기 대손금 ÷ 세법상 전기말 채권 (절사)
    if tax_receivables_prev > 0:
        raw_rate = tax_bad_debt_recognized / tax_receivables_prev
        loss_rate = _loss_rate_floor(raw_rate, loss_rate_digits)
    else:
        loss_rate = 0.0

    # 한도 = 세법상 당기말 채권 × max(1%, 실적률)
    limit = int(tax_receivables_end * max(0.01, loss_rate))
    # 한도초과 = 장부 기말잔액 − 한도
    book_closing = int(book["closing"])
    over = max(0, book_closing - limit)

    # 세무조정 조립
    settlements = []
    # 1. 부인 대손(손금불산입·유보)
    if denied_bad_debt > 0:
        settlements.append({"구분": "부인 대손", "금액": denied_bad_debt, "소득처분": "유보"})
    # 2. 전기 사유분(손금불산입·유보)
    for e in current_disallowed_entries:
        if e["구분"] == "전기 사유분":
            settlements.append({"구분": "전기 사유분", "금액": e["금액"], "소득처분": "유보"})
    # 3. 한도초과(손금불산입·유보)
    if over > 0:
        settlements.append({"구분": "한도초과", "금액": over, "소득처분": "유보"})
    # 4. 전기 한도초과(손금산입·△유보)
    if prior_excess > 0:
        settlements.append({"구분": "전기 한도초과", "금액": prior_excess, "소득처분": "△유보"})

    result = {
        "한도": limit,
        "한도초과": over,
        "전기한도초과_손금산입": prior_excess,
        "세법상_당기_대손금_인정합계": tax_bad_debt_recognized,
        "비망가액_합계": reservation_total,
    }
    calc_dict = {
        "세법상 당기 대손금": tax_bad_debt_recognized,
        "세법상 전기말 채권": tax_receivables_prev,
        "세법상 당기말 채권": tax_receivables_end,
        "대손실적률": loss_rate,
        "한도": limit,
        "한도초과": over,
        "비망가액 합계": reservation_total,
        "장부 당기말": book_closing,
    }
    unified = _build_tax_adjustment_result(settlements, result, calc_dict)
    return _ok(unified, "법인세법 제19조의2·제34조, 시행령 제19조의2·제61조",
               세무조정=settlements, 판단=judgments, 계산=calc_dict,
               산식="세법상 당기 대손금 = 인정 대손금 합계, 세법상 전기말 채권 = 장부 전기말 − 전기 신고조정 대손, 세법상 당기말 채권 = 장부 당기말 + 부인 채권 + 비망가액, 실적률 = 당기 대손금 ÷ 전기말 채권(절사), 한도 = 당기말 채권 × max(1%, 실적률), 한도초과 = 장부 기말잔액 − 한도")


@mcp.tool(annotations=CALC)
@_guard
def entertainment_adjustment(year: YEAR, sme: Annotated[bool, Field(description="중소기업 여부")],
                             months: Annotated[int, Field(description="사업연도 월수(1~12)")] = 12,
                             부가세별도: Annotated[bool, Field(description="현물 접대비 평가 시 부가가치세 포함 여부(기본 True — 사업상 증여는 매출세액 포함, 영 제42조⑥·법 제36조①3호 준용)")] = True,
                             ledger: Annotated[list, Field(description="원장형 입력 [{항목, 금액, 증빙:'없음'|'영수증 등'|'적격'(기본 '적격'), 성격?:'교제목적'|'직원단체_법인'|'직원단체_비법인'|'임직원부담'|'현물'|'경조금'|'문화'|'사적사용', 시가?, 장부가액?, 건당기준초과?}] — 성격별 포함·제외는 어휘 기업업무추진비_포함항목으로 판단. '사적사용'은 시부인 대상 제외, 직접 손금불산입(상여") ] = ...,
                             revenue: Annotated[dict | None, Field(description="수입금액 원장형 입력 {매출액, 특수관계매출?(기본 0), 매출할인등:{일반?, 특수관계?}(기본 0)} — 기업회계기준 매출액 기준(영 제42조①). 없으면 수입금액 0")] = None) -> dict:
    """기업업무추진비 세무조정 — 2023년 이전 사업연도의 "접대비"와 같은 항목(2024.1.1. 명칭 변경, 법률 19193). 접대비 문제·장부도 이 도구로 계산한다(법인세법 제25조, 시행령 제40조·제41조·제42조).
    언제: 법인세 신고에서 기업업무추진비 손금불산입을 항목별로 계산할 때.
    입력:
      · 원장: ledger=[{항목, 금액, 증빙:'없음'|'영수증 등'|'적격'(기본 '적격'), 성격?:'교제목적'|'직원단체_법인'|'직원단체_비법인'|'임직원부담'|'현물'|'경조금'|'문화', 시가?, 장부가액?, 건당기준초과?:(기본 true)}],
        revenue={매출액, 특수관계매출?(기본 0), 매출할인등:{일반?, 특수관계?}(기본 0)}.
      증빙 허용값(어휘 기반):
        · 없음 — 지출 증빙이 전혀 없는 건(손금불산입·대표자상여)
        · 영수증 등 — 영수증 등은 있으나 적격증명서류가 아닌 건(건당 기준 초과 시 손금불산입·기타사외유출)
        · 적격 — 적격증명서류(세금계산서·계산서·신용카드·현금영수증) 수취 건(한도 시부인 대상)
      성격 판단(어휘 기업업무추진비_포함항목.판단.성격별):
        · 없음·교제목적 → 포함(법 제25조①) / 직원단체_법인 → 포함(영 제40조②) / 직원단체_비법인 → 제외(복리후생비)
        · 임직원부담 → 제외(영 제40조①) / 현물 → 포함(max(시가, 장부가액) + 부가세, 부가세별도 기본 True, 영 제42조⑥·법 제36조①3호) / 경조금·문화·전통시장 → 기존과 동일
        · 사적사용 → 시부인 대상 제외, 직접 손금불산입(상여, 법인세법 제19조②·시행령 제40조①)
        · 전통시장 → 포함(「전통시장 및 상점가 육성을 위한 특별법」상 전통시장에서 쓴 적격 기업업무추진비, 조특법 제136조⑥, 전통시장등 추가한도 대상)
        · 지역사랑상품권 → 포함(지역사랑상품권으로 지급한 적격 기업업무추진비, 조특법 제136조⑥, 2025 사업연도부터 전통시장등 추가한도 대상)
        · 온누리상품권 → 포함(온누리상품권으로 지급한 적격 기업업무추진비, 조특법 제136조⑥, 2026 사업연도 지출분부터 전통시장등 추가한도 대상)
    결과: 판단(항목별·근거)·결과(한도·한도초과)·세무조정(구분·금액·소득처분)·소득처분별_합계·손금불산입_합계·계산. 기업업무추진비 세액공제."""

    # ── 빈 원장 거부 ──
    if not ledger:
        return {"error": "입력 확인 필요 — 원장 항목이 없음",
                "허용값": {"증빙": _specs_from_vocab("entertainment_adjustment", "evidence")}}

    # ── 수입금액 원장형 입력(revenue) ──
    if revenue is not None:
        매출액 = int(revenue.get("매출액", 0))
        특수관계매출 = int(revenue.get("특수관계매출", 0))
        매출할인등 = revenue.get("매출할인등", {}) or {}
        할인_일반 = int(매출할인등.get("일반", 0))
        할인_특수관계 = int(매출할인등.get("특수관계", 0))
        revenue_related_calc = 특수관계매출 - 할인_특수관계
        revenue_general_calc = 매출액 - 특수관계매출 - 할인_일반
        if revenue_related_calc < 0 or revenue_general_calc < 0:
            return {"error": "입력 확인 필요 — 수입금액이 음수: 매출액·특수관계매출·매출할인등 확인"}
        revenue_general = revenue_general_calc
        revenue_related = revenue_related_calc
    else:
        revenue_general = 0
        revenue_related = 0

    # ── 증빙 허용값 검증(어휘 기반) ──
    증빙_허용값 = _specs_from_vocab("entertainment_adjustment", "evidence")

    # ── 원장형 입력 처리 ──
    포함항목_규칙 = vocab.rule("기업업무추진비_포함항목", year)
    성격별_판단 = 포함항목_규칙.get("성격별", {})

    judgments = []
    items_for_limit = []
    items_for_limit_names = []
    no_receipt_total = 0
    no_qual_total = 0
    private_use_total = 0
    settlements = []

    for item in ledger:
        항목명 = item.get("항목", "")
        금액 = int(item.get("금액", 0))
        증빙 = item.get("증빙", "적격")
        성격 = item.get("성격", None)
        시가 = item.get("시가", None)
        장부가액 = item.get("장부가액", None)
        건당기준초과 = item.get("건당기준초과", True)

        # 증빙 허용값 밖이면 계산 거부(0.8.x의 '검토 필요' 대신 통일)
        if 증빙 not in 증빙_허용값:
            return {"error": "입력 확인 필요 — 증빙 허용값 아님",
                    "허용값": {"증빙": 증빙_허용값},
                    "문제": [f"증빙='{증빙}' 은(는) {list(증빙_허용값)} 중 하나여야 함"]}

        _현물_판단_이유 = None
        if 성격 == "현물":
            if 시가 is not None and 장부가액 is not None:
                평가액 = max(int(시가), int(장부가액))
            elif 시가 is not None:
                평가액 = int(시가)
            elif 장부가액 is not None:
                평가액 = int(장부가액)
            elif 금액:
                평가액 = int(금액); 시가 = 금액   # 시가·장부가액이 없으면 입력 금액을 평가액으로(영 제42조⑥ — 시가·장부가액 중 큰 금액을 알면 따로 입력)
            else:
                return {"error": "입력 확인 필요 — 현물 항목인데 금액·시가·장부가액 중 하나 필요"}
            if 부가세별도:
                금액 = int(평가액 * 11 // 10)
                if 시가 is not None and 장부가액 is not None:
                    _현물_판단_이유 = f"현물 max(시가 {int(시가):,}, 장부가액 {int(장부가액):,}) + 부가세 {int(평가액 * 11 // 10 - 평가액):,} = {금액:,}원 (영 제42조⑥·법 제36조①3호)"
                elif 시가 is not None:
                    _현물_판단_이유 = f"현물 시가 {int(시가):,} + 부가세 {int(평가액 * 11 // 10 - 평가액):,} = {금액:,}원 (영 제42조⑥·법 제36조①3호)"
                else:
                    _현물_판단_이유 = f"현물 장부가액 {int(장부가액):,} + 부가세 {int(평가액 * 11 // 10 - 평가액):,} = {금액:,}원 (영 제42조⑥·법 제36조①3호)"
            else:
                금액 = 평가액
                if 시가 is not None and 장부가액 is not None:
                    _현물_판단_이유 = f"현물 max(시가 {int(시가):,}, 장부가액 {int(장부가액):,}) = {금액:,}원 (영 제42조⑥·법 제36조①3호)"
                elif 시가 is not None:
                    _현물_판단_이유 = f"현물 시가 {int(시가):,} = {금액:,}원 (영 제42조⑥·법 제36조①3호)"
                else:
                    _현물_판단_이유 = f"현물 장부가액 {int(장부가액):,} = {금액:,}원 (영 제42조⑥·법 제36조①3호)"
            if int(item.get("금액", 0)) > 0 and int(item.get("금액", 0)) != 금액:
                _현물_판단_이유 += f" — 계상액 {int(item.get('금액', 0)):,}원 → {금액:,}원으로 평가"

        if 성격 is None:
            결론 = "포함"
            근거 = "법인세법 제25조①"
            이유 = "명목 불문, 기업업무추진비 포함"
        elif 성격 in 성격별_판단:
            성급_판단 = 성격별_판단[성격]
            결론 = 성급_판단.get("결론", "포함")
            근거 = 성급_판단.get("근거", "")
            이유 = 성급_판단.get("이유", "")
        else:
            결론 = "포함"
            근거 = "법인세법 제25조①"
            이유 = f"성격 '{성격}' — 기본 포함"

        if 결론 == "제외":
            judgments.append({"항목": 항목명, "성격": 성격, "결론": f"제외·근거: {근거}", "금액": 0, "근거": 근거, "이유": 이유})
            continue

        if 성격 == "사적사용":
            judgments.append({"항목": 항목명, "성격": "사적사용", "결론": "제외·손금불산입(상여)", "금액": 금액, "근거": "법인세법 제19조②·시행령 제40조①", "이유": "업무와 관련 없는 개인 용도 지출 — 시부인 대상 제외, 직접 손금불산입(상여)"})
            settlements.append({"구분": "사적사용 손금불산입", "금액": 금액, "소득처분": "대표자상여"})
            private_use_total += 금액
            continue

        if 증빙 == "없음":
            no_receipt_total += 금액
            judgments.append({"항목": 항목명, "성격": 성격, "결론": "손금불산입·대표자상여", "금액": 금액, "근거": "법인세법 제25조②(증빙서류 없는 접대비 손금불산입)", "이유": ("증빙서류 없음 — 전액 손금불산입, 대표자상여 소득처분" + (f" — {_현물_판단_이유}" if _현물_판단_이유 else ""))})
        elif 증빙 == "영수증 등":
            기준_일반 = vocab.rule("기업업무추진비_적격증빙미수취", year).get("기준", {}).get("일반", 30000)
            기준_경조금 = vocab.rule("기업업무추진비_적격증빙미수취", year).get("기준", {}).get("경조금", 200000)
            limit_amount = 기준_경조금 if 성격 == "경조금" else 기준_일반
            if 건당기준초과 and 금액 > limit_amount:
                no_qual_total += 금액
                judgments.append({"항목": 항목명, "성격": 성격, "결론": "적격증명서류 미수취 → 손금불산입·기타사외유출", "금액": 금액, "근거": "법인세법 제25조②·시행령 제41조①(적격증빙 미수취)", "이유": (f"{성격 or '일반'} {limit_amount:,}원 초과, 적격증명서류 미수취 → 기타사외유출" + (f" — {_현물_판단_이유}" if _현물_판단_이유 else ""))})
            else:
                items_for_limit.append(금액)
                items_for_limit_names.append(항목명)
                judgments.append({"항목": 항목명, "성격": 성격, "결론": "인정", "금액": 금액, "근거": "법인세법 제25조②", "이유": (f"{성격 or '일반'} {limit_amount:,}원 이하 → 적격증빙 요구 해당 없음" + (f" — {_현물_판단_이유}" if _현물_판단_이유 else ""))})
        elif 증빙 == "적격":
            items_for_limit.append(금액)
            items_for_limit_names.append(항목명)
            if 성격 == "문화":
                judgments.append({"항목": 항목명, "성격": 성격, "결론": "시부인 대상(문화접대비)", "금액": 금액, "근거": "법인세법 제25조④(문화접대비)", "이유": ("적격증빙·문화비 지출 → 한도 계산 시 문화접대비 포함" + (f" — {_현물_판단_이유}" if _현물_판단_이유 else ""))})
            elif 성격 == "경조금":
                judgments.append({"항목": 항목명, "성격": 성격, "결론": "시부인 대상(경조금)", "금액": 금액, "근거": "법인세법 제25조①", "이유": ("적격증빙·경조금 지출 → 한도 계산 대상" + (f" — {_현물_판단_이유}" if _현물_판단_이유 else ""))})
            elif 성격 == "전통시장":
                judgments.append({"항목": 항목명, "성격": 성격, "결론": "시부인 대상(전통시장 기업업무추진비)", "금액": 금액, "근거": "조세특례제한법 제136조⑥", "이유": ("적격증빙·전통시장 지출 → 한도 계산 시 전통시장 추가한도 대상" + (f" — {_현물_판단_이유}" if _현물_판단_이유 else ""))})
            else:
                judgments.append({"항목": 항목명, "성격": 성격, "결론": 결론, "금액": 금액, "근거": 근거, "이유": ("적격증명서류 수취 → 한도 계산 대상" + (f" — {_현물_판단_이유}" if _현물_판단_이유 else ""))})

    si_bu_shin = sum(items_for_limit)
    culture_total = sum(int(item.get("금액", 0)) for item in ledger if item.get("성격") == "문화" and item.get("증빙", "적격") == "적격")
    전통시장_총액 = sum(int(item.get("금액", 0)) for item in ledger if item.get("성격") == "전통시장" and item.get("증빙", "적격") == "적격")
    지역사랑상품권_총액 = sum(int(item.get("금액", 0)) for item in ledger if item.get("성격") == "지역사랑상품권" and item.get("증빙", "적격") == "적격")
    온누리상품권_총액 = sum(int(item.get("금액", 0)) for item in ledger if item.get("성격") == "온누리상품권" and item.get("증빙", "적격") == "적격")
    기본한도_중소 = vocab.rule("기업업무추진비_적격수취", year).get("기본한도", {}).get("중소", 36000000)
    기본한도_일반 = vocab.rule("기업업무추진비_적격수취", year).get("기본한도", {}).get("일반", 12000000)

    # ── 전통시장등 추가한도 대상 성격 합계 (연도별) ──
    # 2023·2024: 전통시장만 (법률 19936 부칙§35, 2024.1.1. 이후 신고분 → 2023 사업연도부터)
    # 2025: 전통시장 + 지역사랑상품권 (법률 21223 부칙§31, 2026.1.1. 이후 신고분 → 2025 사업연도부터)
    # 2026~2028: 전통시장 + 지역사랑상품권 + 온누리상품권 (법률 21612 부칙§4, 2026.5.12 시행)
    # 2029 이후: 대상 없음
    if year <= 2022 or year >= 2029:
        전통시장등_총액 = 0
    elif year in (2023, 2024):
        전통시장등_총액 = 전통시장_총액
    elif year == 2025:
        전통시장등_총액 = 전통시장_총액 + 지역사랑상품권_총액
    else:  # 2026~2028
        전통시장등_총액 = 전통시장_총액 + 지역사랑상품권_총액 + 온누리상품권_총액

    # ── 한도 계산 (계산식 단계보다 먼저 실행하여 ent 결과를 계산식 값으로 사용) ──
    ent = C.entertain(year, sme, revenue_general, revenue_related, si_bu_shin, culture=culture_total,
                      market=전통시장등_총액, months=months)
    limit = ent["한도"]
    over_limit = ent["한도초과(기타사외유출)"]

    # ── 계산식 단계 (SPEC_export_workbook_c) ──────────────────────────────
    ent_일반한도 = ent["일반한도"]
    ent_문화한도 = ent["문화한도"]
    ent_한도 = ent["한도"]
    ent_한도초과 = ent["한도초과(기타사외유출)"]

    # 수입금액 구간 tiers (식 안의 숫자로 표현) — 연도표 원문 전부(3구간) 사용
    tiers = C.YEAR[year]["접대비_수입"]
    t1_cap, t1_r = tiers[0]
    t2_cap, t2_r = tiers[1]
    t3_cap, t3_r = tiers[2]  # math.inf 가능 (식에서는 그대로 숫자 처리)

    # 일반수입금액 티어 값: t1_r × min(일반, t1) + t2_r × max(0, min(일반, t2) − t1) + t3_r × max(0, 일반 − t2)
    수입금액일반한도 = int(t1_r * min(revenue_general, t1_cap) +
                             t2_r * max(0, min(revenue_general, t2_cap) - t1_cap) +
                             t3_r * max(0, revenue_general - t2_cap))
    # 특수관계 포함 합계 티어 값
    수입금액합계한도 = int(t1_r * min(revenue_general + revenue_related, t1_cap) +
                              t2_r * max(0, min(revenue_general + revenue_related, t2_cap) - t1_cap) +
                              t3_r * max(0, revenue_general + revenue_related - t2_cap))

    계산식_입력: list[dict] = [
        {"이름": "중소기업여부", "값": True if sme else False, "설명": "중소기업 여부 (한도 기준 차등)"},
        {"이름": "기본한도_중소", "값": 기본한도_중소, "설명": "중소기업 기본한도(원) — 법인세법 제25조①, 시행령 제42조"},
        {"이름": "기본한도_일반", "값": 기본한도_일반, "설명": "일반기업 기본한도(원) — 법인세법 제25조①, 시행령 제42조"},
        {"이름": "월수", "값": months, "설명": "사업연도 월수"},
        {"이름": "문화금액", "값": culture_total, "설명": "문화접대비 지출 합계(적격증빙·문화 성격)"},
        {"이름": "일반수입금액", "값": revenue_general, "설명": "일반 수입금액 (매출액 − 특수관계매출 − 매출할인등)"},
        {"이름": "특수관계수입금액", "값": revenue_related, "설명": "특수관계 수입금액 (특수관계매출 − 매출할인등)"},
        {"이름": "수입구간1_한도", "값": t1_cap, "설명": f"수입금액 구간1 한도(원) — {t1_cap:,}원 이하"},
        {"이름": "수입구간1_비율", "값": t1_r, "설명": f"수입금액 구간1 비율 — {t1_cap:,}원 이하"},
        {"이름": "수입구간2_한도", "값": t2_cap, "설명": f"수입금액 구간2 한도(원) — {t2_cap:,}원 이하"},
        {"이름": "수입구간2_비율", "값": t2_r, "설명": f"수입금액 구간2 비율 — {t2_cap:,}원 이하"},
        {"이름": "수입구간3_한도", "값": t3_cap, "설명": f"수입금액 구간3 한도(원) — {t2_cap:,}원 초과"},
        {"이름": "수입구간3_비율", "값": t3_r, "설명": f"수입금액 구간3 비율 — {t2_cap:,}원 초과"},
    ]
    for idx, item in enumerate(items_for_limit, 1):
        설명 = f"기업업무추진비 시부인 대상 항목 {idx} (적격증빙)"
        # 원장 항목 원래 이름과 판단 요지 포함
        항목명 = items_for_limit_names[idx - 1] if idx - 1 < len(items_for_limit_names) else ""
        if 항목명:
            설명 = f"{항목명} — " + 설명
        계산식_입력.append({"이름": f"항목{idx}", "값": item,
                             "설명": 설명})

    # 시부인대상 식: 항목1+항목2+... (항목 수만큼 동적 생성)
    항목식parts = [f"항목{i}" for i in range(1, len(items_for_limit) + 1)]
    시부인대상식 = "+".join(항목식parts) if 항목식parts else "0"

    # 전통시장등 추가한도 입력·단계 (연도별 대상 성격에 따라)
    # 2023·2024: 전통시장만, 비율 10% (법률 19936 부칙§35)
    # 2025: 전통시장+지역사랑상품권, 비율 20% (법률 21223 부칙§31)
    # 2026~2028: 전통시장+지역사랑+온누리, 비율 20% (법률 21612 부칙§4)
    # 지출이 0이어도 비율·금액을 넘긴다 — 서식 엑셀에서 지출액을 나중에 넣어도 한도 수식이 살아 있게
    if 2023 <= year <= 2028:
        비율 = C.YEAR[year].get("전통시장등_비율", 0)
        계산식_입력.append({"이름": "전통시장등비율", "값": 비율,
                             "설명": f"전통시장등 추가한도 비율 — 조특법 제136조⑥, {'법률 19936 부칙§35(2023·2024)' if year in (2023, 2024) else '법률 21223 부칙§31(2025)' if year == 2025 else '법률 21612 부칙§4(2026~2028)'}"})
        계산식_입력.append({"이름": "전통시장등금액", "값": 전통시장등_총액,
                             "설명": f"전통시장등(적격 기업업무추진비 중 해당 연도 대상 성격) 지출 합계 — 조특법 제136조⑥"})

    실제기본한도 = 기본한도_중소 if sme else 기본한도_일반
    계산식: list[dict] = [
        {"이름": "기본한도", "식": "IF(중소기업여부,기본한도_중소,기본한도_일반)*월수/12",
         "값": int(실제기본한도 * months / 12),
         "설명": "IF(중소기업여부, 중소, 일반) 기본한도 × 월수/12"},
        {"이름": "수입금액일반한도", "식": f"수입구간1_비율*MIN(일반수입금액,수입구간1_한도)+수입구간2_비율*MAX(0,MIN(일반수입금액,수입구간2_한도)-수입구간1_한도)+수입구간3_비율*MAX(0,일반수입금액-수입구간2_한도)",
         "값": 수입금액일반한도, "설명": "일반 수입금액 구간별 비율 적용액 (3구간)"},
        {"이름": "수입금액합계한도", "식": f"수입구간1_비율*MIN(일반수입금액+특수관계수입금액,수입구간1_한도)+수입구간2_비율*MAX(0,MIN(일반수입금액+특수관계수입금액,수입구간2_한도)-수입구간1_한도)+수입구간3_비율*MAX(0,일반수입금액+특수관계수입금액-수입구간2_한도)",
         "값": 수입금액합계한도, "설명": "일반+특수관계 수입금액 구간별 비율 적용액 (3구간)"},
        {"이름": "수입금액한도", "식": "수입금액일반한도+(수입금액합계한도-수입금액일반한도)*0.1",
         "값": int(ent_일반한도 - 실제기본한도 * months / 12), "설명": "일반 한도 + (특수관계 − 일반) × 10% (영 제42조①)"},
        {"이름": "문화한도", "식": "MIN(문화금액,(기본한도+수입금액한도)*0.2)",
         "값": ent_문화한도, "설명": "문화접대비 min(문화비, (기본한도+수입금액한도) × 20%) — 조특법 제136조③"},
    ]
    if 2023 <= year <= 2028 and 전통시장등_총액 > 0:
        비율 = C.YEAR[year].get("전통시장등_비율", 0)
        계산식.append({"이름": "전통시장한도", "식": f"MIN(전통시장등금액,(기본한도+수입금액한도)*{비율})",
                        "값": ent["전통시장한도"],
                        "설명": f"전통시장등 추가한도 min(전통시장등금액, (기본한도+수입금액한도) × {비율}) — 조특법 제136조⑥"})
    한도_식 = "기본한도+수입금액한도+문화한도" + ("+전통시장한도" if 2023 <= year <= 2028 and 전통시장등_총액 > 0 else "")
    한도_설명 = "기본한도 + 수입금액한도 + 문화한도" + (" + 전통시장한도" if 2023 <= year <= 2028 and 전통시장등_총액 > 0 else "")
    계산식.extend([
        {"이름": "한도", "식": 한도_식,
         "값": ent_한도, "설명": 한도_설명},
        {"이름": "시부인대상", "식": 시부인대상식,
         "값": si_bu_shin, "설명": "적격증빙 기업업무추진비 합계 (시부인 대상)"},
        {"이름": "한도초과", "식": "MAX(0,시부인대상-한도)",
         "값": ent_한도초과, "설명": "한도초과액 (시부인대상 − 한도, 0 미만 절사) — 기타사외유출"},
    ])

    # ── 세무조정 조립 ──
    # 루프의 settlements(사적사용 등)는 이미 collect 돼 있음 — 덮어쓰지 않고 추가만 한다.
    if no_receipt_total > 0:
        settlements.append({"구분": "증빙 없음", "금액": no_receipt_total, "소득처분": "대표자상여"})
    if no_qual_total > 0:
        settlements.append({"구분": "적격증빙 미수취", "금액": no_qual_total, "소득처분": "기타사외유출"})
    if over_limit > 0:
        settlements.append({"구분": "한도초과", "금액": over_limit, "소득처분": "기타사외유출"})

    # 전통시장등 추가한도 판단 (연도별)
    if 2023 <= year <= 2028 and 전통시장등_총액 > 0:
        비율 = C.YEAR[year].get("전통시장등_비율", 0)
        대상설명 = "전통시장" if year in (2023, 2024) else ("전통시장+지역사랑상품권" if year == 2025 else "전통시장+지역사랑상품권+온누리상품권")
        부칙근거 = "법률 19936 부칙§35" if year in (2023, 2024) else ("법률 21223 부칙§31" if year == 2025 else "법률 21612 부칙§4")
        if year == 2025:
            judgments.append({"항목": "전통시장등 비율(2025 사업연도)", "성격": "전통시장", "결론": "20% 적용 — 2025.12.31 시행 조문 본문은 100분의 10이지만, 2025.12.23 개정 법률 제21223호 부칙 제31조로 2026.1.1 이후 신고분부터 20%. 개정 전 기준(10%)이 필요하면 따로 계산",
                              "금액": 0, "근거": "조세특례제한법 제136조⑥, 법률 제21223호 부칙 제31조", "이유": "본문 시행본과 부칙 적용 시점이 달라 결과가 다르게 보일 수 있음"})
        judgments.append({"항목": "전통시장등 추가한도", "성격": "전통시장", "결론": f"{대상설명} 지출 {전통시장등_총액:,}원 → 추가한도 min({전통시장등_총액:,}, (기본한도+수입금액한도) {ent['일반한도']:,} × {비율}) = {ent['전통시장한도']:,}원 (조특법 제136조⑥, {부칙근거})",
                          "금액": 전통시장등_총액, "근거": "조세특례제한법 제136조⑥",
                          "이유": f"전통시장등 추가한도 대상 성격 합계 {전통시장등_총액:,}원 — (기본한도+수입금액한도) × {비율}과 지출액 중 작은 금액이 추가한도"})
    elif 2023 <= year <= 2028:
        judgments.append({"항목": "전통시장등 추가한도", "성격": "전통시장", "결론": f"{year} 사업연도는 추가한도 적용 기간이나 전통시장등 지출이 없어 0원", "금액": 0,
                          "근거": "조세특례제한법 제136조⑥", "이유": "추가한도 대상 성격(전통시장·지역사랑상품권·온누리상품권) 지출 0"})
    else:
        judgments.append({"항목": "전통시장 추가한도", "성격": "전통시장", "결론": "전통시장 추가한도는 2024.1.1 이후 신고분(2023 사업연도)부터 — 일반 한도로 처리", "금액": 0,
                          "근거": "조세특례제한법 제136조⑥",
                          "이유": "전통시장등 추가한도는 2024.1.1 이후 신고분부터 적용(2023 사업연도~) — 해당 연도 이전 또는 2029 이후는 일반 기업업무추진비로 시부인"})

    result = {"한도": limit, "한도초과": over_limit, "직부인_대표자상여": no_receipt_total + private_use_total, "직부인_기타사외유출": no_qual_total, "시부인_대상": si_bu_shin}
    # 지출이 0이어도 그 해 비율을 넘긴다 — 서식 엑셀에서 지출액을 나중에 넣어도 한도 수식이 살아 있게
    전통시장등비율 = C.YEAR[year].get("전통시장등_비율", 0) if 2023 <= year <= 2028 else 0
    calc_dict = {"기본한도": ent["일반한도"], "문화한도": ent["문화한도"], "전통시장한도": ent["전통시장한도"],
               "시부인 대상": si_bu_shin, "직부인(증빙없음·심사비)": no_receipt_total + no_qual_total,
               "전통시장_지출": (전통시장_총액 if 2023 <= year <= 2028 else 0),
               "지역사랑상품권_지출": (지역사랑상품권_총액 if 2025 <= year <= 2028 else 0),
               "온누리상품권_지출": (온누리상품권_총액 if 2026 <= year <= 2028 else 0),
               "전통시장등비율": 전통시장등비율}
    근거_목록 = vocab.basis("기업업무추진비_적격수취", year)
    근거_문자열 = ", ".join(f"{조문[0]} {조문[1]} {조문[2]} {조문[3]}" if 조문[3] else f"{조문[0]} {조문[1]} {조문[2]}" for 조문 in 근거_목록)
    unified = _build_tax_adjustment_result(settlements, result, calc_dict)
    return _ok(unified, 근거_문자열, 세무조정=settlements, 판단=judgments, 계산=calc_dict,
               산식=f"기본한도 = 중소 {기본한도_중소:,}원·일반 {기본한도_일반:,}원 × 월수/12, 수입금액 한도 = 수입금액 구간별 비율을 적용한 금액(+ 특수관계 10%), 문화한도 = min(문화비, 일반한도 × 20%), 전통시장등한도 = min(전통시장등 지출, (기본한도+수입금액한도) × {C.YEAR[year].get('전통시장등_비율', 0)}) [2023~2028, 연도별 대상·비율 다름], 한도 = 기본 + 수입금액 + 문화 + 전통시장등, 한도초과 = 시부인 대상 − 한도",
               계산식_입력=계산식_입력, 계산식=계산식)


# ───────── 끝에서 끝까지: 의제배당(잉여금 자본전입) ─────────
# 재원별 허용값·판단 규칙은 facts_vocab.json '의제배당_잉여금자본전입' 항목에서 함께 관리한다.
def _bonus_share_sources():
    """어휘 '의제배당_잉여금자본전입.판단.재원별'에서 (과세 제외 재원, 과세 대상 재원, Gross-up 비대상 과세 재원) — 코드 상수 없음 (Claude 10-13)"""
    by = vocab.rule("의제배당_잉여금자본전입", date.today().year)["재원별"]
    non = tuple(k for k, v in by.items() if v == "의제배당 아님")
    tax = tuple(k for k, v in by.items() if "의제배당" in v and "Gross-up 대상" in v and "지분율" not in k)
    tax_gu_false = tuple(k for k, v in by.items() if v == "의제배당·Gross-up 비대상" and "지분율" not in k)
    return non, tax, tax_gu_false


_NON_TAXABLE_SOURCES, _TAXABLE_SOURCES, _TAXABLE_SOURCES_GU_FALSE = _bonus_share_sources()


def _corporate_basis(source, has_self_stock_shift):
    """법인세 납세 주주 대상 근거 조문(법인세법 제16조①2호·3호, 영 제14조①1호가목)."""
    if source in _TAXABLE_SOURCES_GU_FALSE:
        b = "법인세법 제16조①2호 나목" + ("·제16조①3호" if has_self_stock_shift else "")
    elif source in _TAXABLE_SOURCES:
        b = "법인세법 제16조①2호" + ("·제16조①3호" if has_self_stock_shift else "")
    elif source in _NON_TAXABLE_SOURCES:
        if has_self_stock_shift:
            b = "법인세법 제16조①3호, 영 제14조①1호가목"
        else:
            b = ("법인세법 제16조①2호 단서 나목(재평가적립금 — 의제배당 아님)" if source == "재평가적립금"
                 else "법인세법 제16조①2호 단서 가목, 영 제12조①(자본준비금 — 의제배당 아님)")
    else:
        b = ""
    return b


def _income_basis(source, shareholder_type, has_self_stock_shift):
    """소득세법 납세 주주 대상 근거 조문(소득세법 제17조②2호·5호)."""
    if shareholder_type != "개인":
        return _corporate_basis(source, has_self_stock_shift)
    if source in _TAXABLE_SOURCES_GU_FALSE:
        b = "소득세법 제17조②2호 나목·제56조" + ("·제17조②5호" if has_self_stock_shift else "")
    elif source in _TAXABLE_SOURCES:
        b = "소득세법 제17조②2호·제56조" + ("·제17조②5호" if has_self_stock_shift else "")
    elif source in _NON_TAXABLE_SOURCES:
        if has_self_stock_shift:
            b = "소득세법 제17조②5호"
        else:
            b = "소득세법 제17조②2호 단서(과세 제외 재원 — 의제배당 아님)"
    else:
        b = ""
    return b


@mcp.tool(annotations=CALC)
@_guard
def deemed_dividend_bonus_shares(
    year: YEAR,
    source: Annotated[str, Field(description="재원 — 어휘 '의제배당_잉여금자본전입.판단.재원별' 키: 주식발행초과금·자본준비금·재평가적립금·자기주식처분이익·이익잉여금·재평가적립금_토지1%·상환주식_주식발행초과금·자기주식소각이익_2년내·자기주식소각이익")],
    capitalized_amount: WON,
    par_value: Annotated[int, Field(description="액면가(원/주), 신주 1주의 액면금액")],
    shareholders: Annotated[list, Field(description="[{name, shares, self_stock?:bool}] — 무상증자 직전 보유주식수, 발행법인 자기주식은 self_stock=true")],
    self_stock_allotment: Annotated[Literal["다른주주에게교부", "미교부"], Field(description="자기주식 몫의 무상주를 다른 주주에게 교부했는지")],
    shareholder_type: Annotated[Literal["법인", "개인"], Field(description="주주 구분 — 법인세법 제16조(법인) / 소득세법 제17조(개인) 근거 조문이 다름")],
) -> dict:
    """잉여금을 자본에 전입(무상증자)할 때 주주별 의제배당 금액. 언제: 법인이 잉여금을 자본에 전입해 무상주를 발행할 때 주주별 의제배당을 계산한다.
입력:
  year — 사업연도(2016~2026).
  source — 재원(어휘 '의제배당_잉여금자본전입.판단.재원별' 키): 주식발행초과금·자본준비금·재평가적립금(과세 제외), 이익잉여금·자기주식처분이익·상환주식_주식발행초과금·자기주식소각이익_2년내(과세 대상), 재평가적립금_토지1%(과세 대상·Gross-up 비대상), 자기주식소각이익(과세 제외).
  capitalized_amount — 자본전입액(원).
  par_value — 액면가(원/주).
  shareholders — [{name, shares, self_stock?:bool}]: 무상증자 직전 보유주식수, 발행법인 자기주식은 self_stock=true.
  self_stock_allotment — '다른주주에게교부' | '미교부': 자기주식 몫 무상주를 다른 주주에게 교부했는지.
  shareholder_type — '법인' | '개인': 근거 조문(법인세법 제16조 / 소득세법 제17조)이 다름.
결과: 주주별 {배정주식수, 전입전지분율, 전입후지분율, 의제배당, 과세여부, 근거}, 전입후총주식수, 재원구분, 과세대상합계, 산식.
  재원 어휘에 없으면 'error' 반환(계산하지 않음).
  재평가적립금을 일부만 전입하면 1% 토지분은 전입액 × (1% 토지 재평가차액 ÷ 재평가차액 전체)로 나눠 넣을 것(법인세법 시행령 제12조④, 소득세법 시행령 제27조⑤)."""
    source = source.strip()
    if source not in _NON_TAXABLE_SOURCES and source not in _TAXABLE_SOURCES and source not in _TAXABLE_SOURCES_GU_FALSE:
        return {"error": f"지원하지 않는 재원 '{source}' — facts_vocab.json '의제배당_잉여금자본전입.판단.재원별' 허용값 확인 필요",
                "허용값": {"과세 제외(의제배당 아님, 단 자기주식 미배정 시 지분율 상승분은 과세)": list(_NON_TAXABLE_SOURCES),
                           "과세 대상·Gross-up 대상": list(_TAXABLE_SOURCES),
                           "과세 대상·Gross-up 비대상": list(_TAXABLE_SOURCES_GU_FALSE)}}

    from fractions import Fraction

    taxable_source = source in _TAXABLE_SOURCES or source in _TAXABLE_SOURCES_GU_FALSE
    total_pre = sum(int(h["shares"]) for h in shareholders)
    n_new = int(capitalized_amount / par_value)  # 신주 총수

    # 자기주식·배정 (Claude 10-13 재작성 — CodeSolar 지적: 미교부 시 자기주식 배정 0, 교부 시 단주 누락)
    self_idx = {i for i, h in enumerate(shareholders) if h.get("self_stock")}
    self_stock_shares = sum(int(shareholders[i]["shares"]) for i in self_idx)
    others = [i for i in range(len(shareholders)) if i not in self_idx]
    other_total = sum(int(shareholders[i]["shares"]) for i in others)

    def _largest_remainder(pool, idxs, denom):
        """idxs 주주에게 pool 주를 보유주식 비례로 정수 배정(큰 나머지 순, 합계 = pool)."""
        exact = {i: Fraction(int(shareholders[i]["shares"]) * pool, denom) for i in idxs}
        base = {i: int(exact[i]) for i in idxs}
        left = pool - sum(base.values())
        for i in sorted(idxs, key=lambda i: (-(exact[i] - base[i]), shareholders[i]["name"]))[:left]:
            base[i] += 1
        return base

    if self_stock_shares > 0 and self_stock_allotment == "미교부":
        # 자기주식 몫은 발행하지 않음: 다른 주주는 전체 주식 기준 비례분만 (합계 = 신주 − 자기주식 몫)
        allot = _largest_remainder(n_new - (self_stock_shares * n_new) // total_pre, others, other_total) if others else {}
    elif self_stock_shares > 0:
        # 자기주식 몫까지 다른 주주에게 교부: 신주 전부를 다른 주주끼리 보유 비례로
        allot = _largest_remainder(n_new, others, other_total) if others else {}
    else:
        allot = _largest_remainder(n_new, list(range(len(shareholders))), total_pre)
    for i in self_idx:
        allot[i] = 0
    n_issued_to_others = sum(allot.values())
    total_post = total_pre + n_issued_to_others
    allotments = [{"index": i, "name": h["name"], "shares": int(h["shares"]), "proportion": int(h["shares"]) / total_pre,
                   "BaseAllot": allot.get(i, 0), "self_stock": i in self_idx} for i, h in enumerate(shareholders)]

    # 전입 전·후 지분율 및 의제배당 계산
    result_list = []
    taxable_total = 0
    for a in allotments:
        name = a["name"]
        allot = a["BaseAllot"]
        prop_pre = a["proportion"]
        prop_post = (a["shares"] + allot) / total_post if total_post > 0 else 0

        deemed_div = 0
        taxable = False
        if taxable_source:
            deemed_div = allot * par_value  # 받은 주식 수 × 액면가
            taxable = True
        else:
            # 과세 제외 재원 + 자기주식 존재 + 다른주주에게교부/미교부 →
            # 다른 주주의 지분율 상승분에 대해 의제배당 과세 (법인세법 제16조①3호 / 소득세법 제17조②5호)
            if self_stock_shares > 0 and not a["self_stock"]:
                # share_val = (전입후지분율 − 전입전지분율) × total_post × par_value
                #           = [(shares+allot)/total_post − shares/total_pre] × total_post × par_value
                #           = par_value × [(shares+allot)×total_pre − shares×total_post] ÷ total_pre  (정수 계산)
                shift_num = (a["shares"] + allot) * total_pre - a["shares"] * total_post
                share_val = (par_value * shift_num) // total_pre
                if share_val > 0:
                    deemed_div += share_val
                    taxable = True

        taxable_total += deemed_div

        basis = _income_basis(source, shareholder_type, taxable)
        if shareholder_type == "개인" and source in _TAXABLE_SOURCES_GU_FALSE:
            basis = basis + " — Gross-up 비대상(소득세법 제17조③)"
        result_list.append({
            "name": name,
            "의제배당": deemed_div,                # 답은 이 값. 지분율은 참고(지분율 차이로 다시 계산하지 않도록)
            "배정주식수": allot,
            "전입전지분율": round(prop_pre * 100, 4),
            "전입후지분율": round(prop_post * 100, 4),
            "과세여부": "과세" if taxable else "의제배당 아님",
            "근거": basis,
        })

    # 산식 문자열
    if taxable_source:
        calc_formula = (f"신주 총수 = 자본전입액 ÷ 액면가 = {capitalized_amount:,} ÷ {par_value} = {n_new:,}주; "
                        f"각 주주 배정 = 신주 총수 × 지분율(자기주식 포함 총 {total_pre:,}주 기준); "
                        f"자기주식 몫 {'발행 후 다른 주주에게 교부' if self_stock_allotment == '다른주주에게교부' else '미발행'}; "
                        f"의제배당 = 배정주식수 × 액면가 (재원 {source} = 과세 대상, 법인세법 제16조①2호 / 소득세법 제17조②2호)")
    else:
        if self_stock_shares > 0:
            calc_formula = (f"신주 총수 = {capitalized_amount:,} ÷ {par_value} = {n_new:,}주; "
                            f"자기주식 {self_stock_shares:,}주 {'미발행(미교부)' if self_stock_allotment == '미교부' else '발행 후 다른 주주에게 교부(다른주주에게교부)'}; "
                            f"발행 총수 = {n_issued_to_others:,}주; "
                            f"전입 후 총 주식수(자기주식 {'제외' if self_stock_allotment == '미교부' else '포함'}) = {total_post:,}주; "
                            f"각 주주 의제배당 = (전입 후 지분율 − 전입 전 지분율) × 전입 후 총 주식수{('(자기주식 제외)' if self_stock_allotment == '미교부' else '')} × 액면가 "
                            f"(법인세법 제16조①3호 / 소득세법 제17조②5호, 영 제14조①1호가목 — 액면가액)")
        else:
            if source in ("합병차익", "분할차익"):
                calc_formula = (f"신주 총수 = {capitalized_amount:,} ÷ {par_value} = {n_new:,}주; "
                                f"재원 {source} = 과세 제외(자본준비금과 같은 과세 제외 재원) → 의제배당 없음; "
                                f"자기주식 없음 → 지분율 상승분 없음 → 전원 의제배당 0; "
                                f"※ 합병·분할차익의 과세 부분은 별도 확인(상법·합병 세법에 따라 과세 여부가 달라질 수 있음)")
            elif source in ("주식의포괄적교환차익", "주식의포괄적이전차익"):
                calc_formula = (f"신주 총수 = {capitalized_amount:,} ÷ {par_value} = {n_new:,}주; "
                                f"재원 {source} = 과세 제외(자본준비금과 같은 과세 제외 재원) → 의제배당 없음; "
                                f"자기주식 없음 → 지분율 상승분 없음 → 전원 의제배당 0")
            else:
                calc_formula = (f"신주 총수 = {capitalized_amount:,} ÷ {par_value} = {n_new:,}주; "
                                f"재원 {source} = 과세 제외(주식발행초과금·자본준비금·재평가적립금) → 의제배당 없음; "
                                f"자기주식 없음 → 지분율 상승분 없음 → 전원 의제배당 0")

    result = {
        "주주별_의제배당": {r["name"]: r["의제배당"] for r in result_list},   # 요약: 이 숫자가 각 주주의 의제배당(원). 두 사례 비교는 이 값끼리 뺄 것
        "주주별": result_list,
        "전입후총주식수": total_post,
        "재원구분": "과세 대상" if taxable_source else "과세 제외",
        "과세대상합계": taxable_total,
        "산식": calc_formula,
    }
    return _ok(result, _income_basis(source, shareholder_type, taxable_source))


def _직장공제회_원천(amt: int, 납입연수, year: int) -> str:
    """직장공제회 초과반환금 원천징수 설명 — 기본세율 연분연승(소득세법 제129조①1호다목, 제63조①).
    납입연수가 없으면 금액을 내지 않고 계산 방법만 적는다."""
    방법 = "분리과세(제14조③3호), 원천징수는 기본세율(제129조①1호다목) 연분연승(제63조①): (초과반환금 − 40% − 납입연수공제) ÷ 납입연수 × 기본세율 × 납입연수"
    if 납입연수 is None:
        return 방법 + " — 납입연수(contribution_years) 입력 시 세액 계산, 금융소득 합계·합산액 모두 제외"
    if isinstance(납입연수, bool) or not isinstance(납입연수, (int, float)) or 납입연수 < 0 or (납입연수 >= 1 and 납입연수 != int(납입연수)):
        return 방법 + f" — contribution_years={납입연수!r}: 납입연수는 정수 연수로 입력(1년 미만은 1년), 세액 미계산, 금융소득 합계·합산액 모두 제외"
    n = max(1, int(납입연수))   # 1년 미만은 1년(제63조①)
    if n <= 5:
        연수공제 = 300_000 * n
    elif n <= 10:
        연수공제 = 1_500_000 + 500_000 * (n - 5)
    elif n <= 20:
        연수공제 = 4_000_000 + 800_000 * (n - 10)
    else:
        연수공제 = 12_000_000 + 1_200_000 * (n - 20)
    과표 = max(0, amt - int(amt * 0.4) - 연수공제)
    연평균 = 과표 // n
    세액 = I.income_tax(연평균, year) * n
    return (방법 + f" — 납입연수 {n}년: ({amt:,} − {int(amt * 0.4):,} − {연수공제:,}) ÷ {n} = {연평균:,}원,"
            f" 기본세율 세액 {I.income_tax(연평균, year):,}원 × {n} = 산출세액 {세액:,}원, 금융소득 합계·합산액 모두 제외")


@mcp.tool(annotations=CALC)
@_guard
def financial_income_aggregate(year: YEAR,
                               items: Annotated[list, Field(description="금융소득 항목 [{kind, amount, withheld?, ...}]. "
                                   "kind: 예금이자·비영업대금이익·소비대차전환이자·외상매출금지급기일연장이자·저축성보험차익·투자신탁환매이익·파생결합사채이익·장내파생상품이익·"
                                   "상장법인현금배당·비상장법인현금배당·의제배당_잉여금자본전입·인정배당·출자공동사업자배당·외국법인배당·환매조건부매매차익·법원보증금이자·직장공제회초과반환금·통합발행채권할인액·손해배상법정이자·공익신탁이익. "
                                   "withheld=false: 국내에서 원천징수되지 않은 이자·배당(개인 간 대여 이자, 국외 배당)은 2천만원 이하여도 종합과세(제14조③6호), 기본 true. "
                                   "의제배당_잉여금자본전입은 source 필수 — 이익잉여금·자기주식처분이익·상환주식_주식발행초과금·자기주식소각이익_2년내(과세) / 주식발행초과금·자본준비금·재평가적립금·합병차익·분할차익·주식의포괄적교환차익·주식의포괄적이전차익·자기주식소각이익(과세 제외) / 재평가적립금_토지1%. "
                                   "투자신탁환매이익은 components 필수 — [{asset, gain}] 목록, asset이 '상장주식'인 손익만 빼고 나머지(채권·비상장주식·배당·이자 등 아무 이름)는 합산, 손실은 음수. "
                                   "예: {kind:'투자신탁환매이익', amount:18000000, components:[{asset:'상장주식', gain:3000000}, {asset:'배당', gain:15000000}]}. "
                                   "직장공제회초과반환금은 contribution_years(납입연수)를 주면 연분연승 세액까지")]) -> dict:
    """종합소득 합산 금융소득(이자·배당)과 Gross-up 계산(소득세법 제14조·제16조·제17조·제56조).
    언제: 종합소득세 신고에서 금융소득 종합과세 대상 여부와 배당가산(Gross-up)액을 계산할 때.
    입력: 귀속연도(2016~2026) → year, 금융소득 항목 목록 [{kind, amount, ...}] → items.
      kind 허용값(어휘 기반, 별칭 포함):
        · 예금이자 — 예금·적금·부금의 이자
        · 비영업대금이익 — 사업이 아닌 금전 대여에서 받는 이자
        · 소비대차전환이자 — 외상매출금 등 거래 대금이 소비대차로 바뀐 뒤 그 원금에서 생긴 이자
        · 외상매출금지급기일연장이자 — 물품·용역 대금의 지급기일을 연장해 주고 추가로 받는 금액(사업소득)
        · 저축성보험차익 — 보험 만기·해지 시 받는 보험금에서 납입보험료를 뺀 차익
        · 투자신탁환매이익 — 집합투자기구(투자신탁 등)로부터 받는 이익
        · 파생결합사채이익 — 상법상 파생결합사채(ELB 등)에서 받는 이익
        · 장내파생상품이익 — KOSPI200 선물·옵션 등 장내파생상품 거래 이익(양도소득)
        · 상장법인현금배당 — 상장 내국법인에서 받는 이익배당(Gross-up 대상)
        · 비상장법인현금배당 — 비상장 내국법인에서 받는 이익배당(Gross-up 대상)
        · 의제배당_잉여금자본전입 — 잉여금 자본전입 무상주(재평가적립금 일부 전입 시 1% 토지분 안분, 법인세법 시행령 제12조④, 소득세법 시행령 제27조⑤)
        · 인정배당 — 법인세법에 따라 배당으로 소득처분된 금액(Gross-up 대상)
        · 출자공동사업자배당 — 출자공동사업자 배당(무조건 종합과세, 2천만원 판정 제외, Gross-up 비대상, 원천 25%)
        · 외국법인배당 — 외국법인으로부터 받는 배당(조건부 종합과세, Gross-up 비대상)
        · 환매조건부매매차익 — RP(환매조건부매매) 차익(이자소득, 조건부 종합과세)
        · 법원보증금이자 — 법원 보증금·경락대금 이자(14% 분리과세 제129조②1호, 합계·합산액 제외)
        · 직장공제회초과반환금 — 분리과세, 합계 제외, 연분연승 제63조(납입연수 contribution_years)
        · 손해배상법정이자 — 계약 위약·해약에 따른 손해배상금의 법정이자(이자소득 아님, 기타소득 소득세법 제21조①10호, 금융소득 합계 제외)
        · 공익신탁이익 — 공익신탁 이익(비과세 제12조1호, 합계 제외)
        · 통합발행채권할인액 — 공개시장 통합발행 채권의 매각·액면 차액(이자 아님, 영 제22조의2②, 합계 제외)
    결과: 판단(항목별 이자/배당/제외·Gross-up 대상 여부)·결과(이자소득금액·배당소득금액·합산액)·계산(중간값). 금융소득 종합과세·배당가산."""

    # 어휘의 항목 이름과 다른 흔한 표기를 정식 이름으로 해결(예: 장내파생상품→장내파생상품이익)
    _kind_별칭 = {"장내파생상품": "장내파생상품이익",
                 # 평소 말(처음 쓰는 사람 점검 2026-10-30): 내국법인 현금배당은 상장·비상장 계산이 같아 상장으로 둠, 이자는 예금이자로
                 "배당": "상장법인현금배당", "배당금": "상장법인현금배당", "현금배당": "상장법인현금배당", "배당소득": "상장법인현금배당", "주식배당금": "상장법인현금배당",
                 "이자": "예금이자", "이자소득": "예금이자", "은행이자": "예금이자", "적금이자": "예금이자", "예금 이자": "예금이자"}
    _해석 = []
    for it in items:
        k = it.get("kind", "")
        if k in _kind_별칭:
            it["kind"] = _kind_별칭[k]
            _해석.append(f"kind '{k}' → '{it['kind']}'" + ("(내국법인 현금배당은 상장·비상장 계산이 같음)" if it["kind"] == "상장법인현금배당" else ""))

    _extra = [f"항목 {i+1}(저축성보험차익): 계약 {year - int(it.get('contract_year'))}년 유지 — 비과세 요건(시행령 제25조: 납입 한도·월적립식 등) 충족 여부를 'exempt_conditions_met'(true/false)로 입력 필요"
              for i, it in enumerate(items or []) if isinstance(it, dict) and it.get("kind") == "저축성보험차익" and it.get("contract_year")
              and year - int(it.get("contract_year")) >= 10 and it.get("exempt_conditions_met") is None]
    _excl = tuple(k for k, v in vocab.rule("의제배당_잉여금자본전입", year)["재원별"].items() if v == "의제배당 아님")
    _extra += [f"항목 {i+1}(의제배당_잉여금자본전입, {it.get('source')}): 과세 제외 재원이라 무상주 자체는 의제배당이 아니지만, 법인이 자기주식을 보유해 다른 주주의 지분율이 오른 부분은 의제배당(소득세법 제17조②5호, Gross-up 비대상) — 그 금액을 'self_stock_shift'로 입력 필요(없으면 0)"
               for i, it in enumerate(items or []) if isinstance(it, dict) and it.get("kind") == "의제배당_잉여금자본전입"
               and it.get("source") in _excl and it.get("self_stock_shift") is None]   # 과세 제외 재원이면 자기주식 미배정 상승분이 빠지지 않도록
    _v = _validate_items(items, "kind", _specs_from_vocab("financial_income_aggregate", "kind"), extra=_extra)
    if _v: return _v

    # Gross-up 가산율은 어휘에서 가져온다(연도별)
    배당_규칙 = vocab.rule("상장법인현금배당", year)
    rate = 배당_규칙.get("Gross-up가산율", 0.1)
    판단, inc, div_n, div_g, exc = [], 0, 0, 0, 0
    항목별: list[dict] = []  # 계산.항목별용: [{항목, 금액, 구분}]

    미원천 = 0          # 국내에서 원천징수되지 않은 이자·배당 — 2천만원 이하여도 종합과세(제14조③6호)
    미원천_이자 = 0
    cur_withheld = True

    def add(항목, 구분, 금액, 근거, 이유, gu=None):
        nonlocal inc, div_n, div_g, exc, 미원천, 미원천_이자
        if 구분 in ("이자", "배당") and not cur_withheld and 항목 != "출자공동사업자배당":
            미원천 += 금액
            if 구분 == "이자":
                미원천_이자 += 금액
            이유 = 이유 + " / 원천징수되지 않음 — 2천만원 이하여도 종합과세(소득세법 제14조③6호)"
        if 구분 == "이자": inc += 금액
        elif 구분 == "배당" and gu: div_g += 금액
        elif 구분 == "배당": div_n += 금액
        else: exc += 금액
        결론 = {"이자": "이자소득", "배당": "배당소득(Gross-up 대상)" if gu else "배당소득(Gross-up 비대상)"}.get(구분, "금융소득 제외")
        판단.append({"항목": 항목, "결론": 결론, "금액": 금액, "근거": 근거, "이유": 이유})
        # 계산.항목별용 세부 구분 결정
        if 구분 == "이자":
            세부구분 = "이자"
        elif 구분 == "배당":
            if 항목 == "출자공동사업자배당":
                세부구분 = "출자공동"
            elif gu:
                세부구분 = "GU대상"
            else:
                세부구분 = "GU비대상"
        else:
            세부구분 = "제외"
        항목별.append({"항목": 항목, "금액": 금액, "구분": 세부구분})

    for it in items:
        k, amt = it.get("kind", ""), int(it.get("amount", 0))
        cur_withheld = it.get("withheld", True) is not False
        if k == "예금이자":
            add(k, "이자", amt, "소득세법 제16조①1호", "국내 예금 이자")
        elif k in ("비영업대금이익", "소비대차전환이자"):
            add(k, "이자", amt, "소득세법 제16조①11호, 시행령 제26조③", "금전 대여 이익 — 외상매출금이 소비대차로 전환된 이자 포함")
        elif k == "외상매출금지급기일연장이자":
            add(k, "제외", amt, "소득세법 시행령 제26조③", "물품 대금 지급기일 연장 이자는 사업소득(금융소득 아님)")
        elif k == "장내파생상품이익":
            add(k, "제외", amt, "소득세법 제94조①5호", "파생상품 거래 이익은 양도소득")
        elif k == "저축성보험차익":
            gain = int(it.get("payout", amt)) - int(it.get("premium", 0))
            held = year - int(it.get("contract_year", year))
            if held >= 10 and it.get("exempt_conditions_met") is True:
                add(k, "제외", max(gain, 0), "소득세법 제16조①9호, 시행령 제25조①", f"계약 {held}년 유지 + 비과세 요건(납입 한도 등) 충족 — 비과세")
            elif it.get("annuity"):
                add(k, "제외", max(gain, 0), "소득세법 제16조①9호, 시행령 제25조", "비과세 요건 연금형 — 과세 제외(요건 충족 전제)")
            else:
                add(k, "이자", max(gain, 0), "소득세법 제16조①9호, 시행령 제25조", f"만기보험금 − 납입보험료 = {gain:,}원(10년 미만 등 비과세 요건 미충족)")
        elif k == "투자신탁환매이익":
            comps = it.get("components") or []
            if comps and not (isinstance(comps, list) and all(isinstance(c, dict) for c in comps)):   # dict 형태 components는 크래시 대신 형식 안내
                return {"error": "입력 확인 필요 — 계산하지 않음",
                        "문제": [f"항목({k}): 'components'는 [{{'asset': '상장주식'|'채권'|'비상장주식'|…, 'gain': 원(손실은 음수)}}] 목록이어야 함 — 집합투자기구 이익 중 상장주식 양도차익은 제외(소득세법 시행령 제26조의2④)"],
                        "예시": {"kind": "투자신탁환매이익", "amount": 12000000, "components": [{"asset": "상장주식", "gain": 10000000}, {"asset": "채권", "gain": -3000000}, {"asset": "비상장주식", "gain": 5000000}]}}
            net = sum(int(c.get("gain", 0)) for c in comps if c.get("asset") != "상장주식") if comps else amt
            add(k, "배당", max(net, 0), "소득세법 제17조①5호, 시행령 제26조의2", "상장주식 양도·평가 손익 제외, 나머지 손익 통산", gu=False)
        elif k == "파생결합사채이익":
            add(k, "배당", amt, "소득세법 제17조①5호의2", "파생결합사채 이익 — 배당소득, Gross-up 비대상", gu=False)
        elif k in ("상장법인현금배당", "비상장법인현금배당"):
            add(k, "배당", amt, "소득세법 제17조①1호, 제56조", "내국법인 이익배당 — 법인세 과세 재원, Gross-up 대상", gu=True)
        elif k == "인정배당":
            add(k, "배당", amt, "소득세법 제17조①4호, 제56조", "법인세법상 배당 소득처분 — Gross-up 대상", gu=True)
        elif k == "의제배당_잉여금자본전입":
            src, shift = it.get("source", ""), int(it.get("self_stock_shift", 0))
            _non, _tax, _gu_false = _bonus_share_sources()
            if src in _non:
                add(f"{k}({src})", "제외", amt - shift, "소득세법 제17조②2호 단서, 시행령 제27조④", "과세되지 않는 자본잉여금 전입 — 의제배당 아님")
                if shift:
                    add(f"{k}({src}) 자기주식 미배정 지분율 상승분", "배당", shift, "소득세법 제17조②5호", "자기주식 몫 미배정으로 다른 주주 지분 증가 — 의제배당, Gross-up 비대상", gu=False)
            elif src in _gu_false and src != "재평가적립금_토지1%":
                add(f"{k}({src})", "배당", amt - shift, "소득세법 제17조②2호, 제17조③", f"{src} 자본전입 — 의제배당, Gross-up 제외(제17조③)", gu=False)
                if shift:
                    add(f"{k}({src}) 자기주식 미배정 지분율 상승분", "배당", shift, "소득세법 제17조②5호", "자기주식 몫 미배정으로 다른 주주 지분 증가 — 의제배당, Gross-up 비대상", gu=False)
            elif src == "재평가적립금_토지1%":
                add(f"{k}({src})", "배당", amt - shift, "소득세법 제17조②2호 나목, 제56조", "토지 재평가차액(재평가세 1%) 자본전입 — 의제배당, Gross-up 제외(소득세법 제17조③)", gu=False)
                if shift:
                    add(f"{k}({src}) 자기주식 미배정 지분율 상승분", "배당", shift, "소득세법 제17조②5호", "자기주식 몫 미배정으로 다른 주주 지분 증가 — 의제배당, Gross-up 비대상", gu=False)
            else:
                add(f"{k}({src})", "배당", amt - shift, "소득세법 제17조②2호, 제56조", f"{src} 자본전입 — 의제배당, 법인세 과세 재원이라 Gross-up 대상", gu=True)
                if shift:
                    add(f"{k}({src}) 자기주식 미배정 지분율 상승분", "배당", shift, "소득세법 제17조②5호", "자기주식 몫 미배정으로 다른 주주 지분 증가 — 의제배당, Gross-up 비대상", gu=False)
        elif k == "출자공동사업자배당":
            wr = vocab.rule("출자공동사업자배당", year)
            wr_rate = wr.get("원천징수세율", 0.25)
            add(k, "배당", amt, "소득세법 제17조①8호, 제14조③6호 괄호, 제129조①2호가목", f"무조건 종합과세 — 2천만원 이하라도 합산액에 포함, 원천징수 {wr_rate * 100:.0f}%(제129조①2호가목), Gross-up 비대상", gu=False)
        elif k == "외국법인배당":
            add(k, "배당", amt, "소득세법 제17조①6호, 제14조③6호", "외국법인 배당 — 조건부 종합과세, Gross-up 비대상(법인세 과세 재원 아님)")
        elif k == "환매조건부매매차익":
            add(k, "이자", amt, "소득세법 제16조①8호", "환매조건부매매(RP) 차익 — 이자소득, 조건부 종합과세")
        elif k == "법원보증금이자":
            원천율 = vocab.rule("법원보증금이자", year).get("원천징수세율", 0.14)
            원천징수세액 = int(amt * 원천율)
            add(k, "제외", amt, "소득세법 제129조②1호, 제14조③3호", f"민사집행법상 법원 보증금·경락대금 이자 — 원천징수 {원천율 * 100:.0f}%(제129조②1호) {원천징수세액:,}원으로 분리과세, 금융소득 합계·합산액 모두 제외")
        elif k == "직장공제회초과반환금":
            add(k, "제외", amt, "소득세법 제16조①10호, 제14조③3호, 제129조①1호다목, 제63조①", _직장공제회_원천(amt, it.get("contribution_years"), year))
        elif k == "손해배상법정이자":
            add(k, "제외", amt, "소득세법 제21조①10호, 제16조", "이자소득 아님 — 기타소득(소득세법 제21조①10호)")
        elif k == "공익신탁이익":
            add(k, "제외", amt, "소득세법 제12조1호", "비과세 — 공익신탁의 이익(소득세법 제12조1호)")
        elif k == "통합발행채권할인액":
            add(k, "제외", amt, "소득세법 시행령 제22조의2②, 제16조①1호·2호", "공개시장 통합발행 채권의 매각가액−액면가액 차액 — 이자·할인액에 포함되지 않음, 금융소득 합계 제외")
        else:
            판단.append({"항목": k, "결론": "확인 필요", "금액": amt, "근거": "", "이유": "지원하지 않는 항목 — 직접 판단 필요"})

    # 출자공동사업자배당은 2천만원 판정 합계에서 제외하되 합산액에는 항상 포함
    partner_div = next((h["금액"] for h in 판단 if h["항목"] == "출자공동사업자배당"), 0)
    cond_total = inc + div_n + div_g - partner_div   # 출자공동사업자배당 제외한 조건부 합계
    threshold = 20_000_000
    if cond_total <= threshold:
        gu_base = 0
        grossup = 0
        합산 = partner_div + 미원천
        분리과세 = True
    else:
        filled = min(threshold, inc + (div_n - partner_div))   # 이자 → 비대상 배당(비출자공동) 순으로 2천만 원 채움
        gu_base = div_g - max(0, threshold - filled)
        gu_base = max(0, min(div_g, gu_base))
        grossup = int(gu_base * rate)
        합산 = cond_total + grossup + partner_div
        분리과세 = False
    # ── 계산식 단계 (SPEC_export_workbook_f) ──────────────────────────────
    결과_배당소득금액 = (div_n + div_g + grossup) if not 분리과세 else partner_div + (미원천 - 미원천_이자)
    결과_합산액 = 합산
    결과_Grossup = grossup

    계산식_입력 = [
        {"이름": "종합과세기준금액", "값": 20000000, "설명": "금융소득 종합과세 기준액 = 2,000만 원 (소득세법 제14조③6호)"},
        {"이름": "가산율", "값": rate, "설명": f"Gross-up 가산율 = {rate} ({year}년 기준, 소득세법 제56조)"},
        {"이름": "이자소득금액", "값": inc, "설명": f"이자소득금액 = {inc:,}원"},
        {"이름": "배당비대상소득", "값": div_n, "설명": f"배당소득 중 Gross-up 비대상 = {div_n:,}원 (출자공동사업자배당 등)"},
        {"이름": "배당대상소득", "값": div_g, "설명": f"배당소득 중 Gross-up 대상 = {div_g:,}원 (상장법인·비상장법인 배당, 의제배당 등)"},
        {"이름": "출자공동사업자배당", "값": partner_div, "설명": f"출자공동사업자배당(always 합산 포함, 2천만원 판정 제외) = {partner_div:,}원"},
        {"이름": "제외합계", "값": exc, "설명": f"분리과세·제외 금융소득 합계 = {exc:,}원 (법원보증금이자, 직장공제회초과반환금 등)"},
    ]

    # 판정합계 = 이자 + 배당비대상 + 배당대상 − 출자공동사업자배당 (출자공동사업자배당은 2천만원 판정에서 제외)
    판정합계_식 = f"이자소득금액+배당비대상소득+배당대상소득-출자공동사업자배당"
    # Gross-up 기준액: 2천만원을 이자→배당비대상(출자공동 제외)→배당대상 순으로 채운 뒤 남은 배당대상 × 가산율
    grossup_기준액_식 = f"MAX(0,MIN(배당대상소득,배당대상소득-MAX(0,종합과세기준금액-MIN(종합과세기준금액,이자소득금액+배당비대상소득-출자공동사업자배당))))"
    grossup_식 = f"ROUNDDOWN({grossup_기준액_식}*가산율,0)"
    # 배당소득금액: 판정합계(이자+배당비대상+배당대상−출자공동) > 기준금액이면 종합과세
    #   → 배당비대상+배당대상+grossup, 아니면 분리과세 → max(0, 출자공동사업자배당)
    배당소득금액_식 = f"IF({판정합계_식}>종합과세기준금액,배당비대상소득+배당대상소득+{grossup_식},MAX(0,출자공동사업자배당))"
    # 합산액: 판정합계 > 기준금액이면 판정합계+grossup+출자공동사업자배당, 아니면 MAX(0, 출자공동사업자배당)
    합산액_식 = f"IF({판정합계_식}>종합과세기준금액,{판정합계_식}+{grossup_식}+출자공동사업자배당,MAX(0,출자공동사업자배당))"

    배당소득금액_설명 = f"배당비대상 + 배당대상 + Gross-up = {div_n:,} + {div_g:,} + {grossup:,} = {결과_배당소득금액:,}원 (분리과세 시 출자공동사업자배당만 배당소득)"
    합산액_설명 = f"판정합계 + Gross-up + 출자공동사업자배당 = {cond_total:,} + {grossup:,} + {partner_div:,} = {결과_합산액:,}원 (분리과세 시 출자공동사업자배당만 합산)"

    계산식 = [
        {"이름": "배당소득금액",
         "식": 배당소득금액_식,
         "값": 결과_배당소득금액,
         "설명": 배당소득금액_설명},
        {"이름": "합산액",
         "식": 합산액_식,
         "값": 결과_합산액,
         "설명": 합산액_설명},
        {"이름": "GrossUp",
         "식": grossup_식,
         "값": 결과_Grossup,
         "설명": f"Gross-up = MAX(0, MIN(배당대상소득, 배당대상소득 − MAX(0, 종합과세기준금액 − MIN(종합과세기준금액, 이자소득금액 + 배당비대상소득 − 출자공동사업자배당)))) × 가산율 = {grossup:,}원"},
    ]

    return _ok({"이자소득금액": inc if not 분리과세 else 미원천_이자,
                "배당소득금액": (div_n + div_g + grossup) if not 분리과세 else partner_div + (미원천 - 미원천_이자),
                "합산액": 합산, "분리과세": 분리과세, "Gross-up": grossup},
               "소득세법 제14조·제16조·제17조·제56조",
               계산={"이자": inc, "배당(Gross-up 비대상)": div_n, "배당(Gross-up 대상)": div_g,
                     "금융소득 합계(조건부)": cond_total, "출자공동사업자배당(always 포함)": partner_div,
                     "원천징수되지 않은 금융소득(2천만원 이하여도 종합과세)": 미원천,
                     "제외 합계": exc, "Gross-up 기준액": gu_base, "가산율": rate,
                     "항목별": 항목별},
               판단=판단,
               산식="판정합계(이자+배당비대상+배당대상−출자공동사업자배당) ≤ 2천만 원이면 분리과세(배당소득금액=출자공동사업자배당, 합산액=출자공동사업자배당). 초과 시 이자 → 배당비대상(출자공동 제외) → 배당대상 순으로 2천만 원을 채우고, 2천만 원을 넘는 배당대상 × 가산율 = Gross-up. 출자공동사업자배당은 항상 합산액에 포함되고 2천만원 판정 합계에서는 제외.",
               계산식_입력=계산식_입력,
               계산식=계산식, **({"입력 해석": _해석} if _해석 else {}))


# ── 배당세액공제(소득세법 제56조)와 비교과세(제62조) ─────────────────────────────
# 연도표(조문·부칙 대조, Claude 10-28):
#   · 제56조①④(배당세액공제 = Gross-up 금액, 대상은 2천만원 초과 배당), 제15조2호 후단(공제 후 세액이
#     제62조2호 금액보다 작아지지 않게 — 곧 공제 한도 = 산출세액 − 제62조2호 금액): 2016~2026 문언 같음.
#   · Gross-up 가산율: 11%(2016~2023), 10%(2024~) — financial_income_aggregate와 같은 어휘 값을 쓴다.
#   · 제62조2호 가목 단서(원천징수되지 않은 이자소득등의 세율):
#       2016 — "제127조에 따라 원천징수되지 아니하는 소득에 대해서는 제129조제1항제1호라목의 세율"(14%) 하나.
#       2017~ — 2)·3)으로 나뉨: 원천징수되지 않은 비영업대금의 이익(제16조①11호)은 제129조①1호나목(25%),
#               그 밖은 라목(14%). 법률 제14389호(2016.12.20) 부칙 제1조(2017.1.1. 시행)·
#               제2조①("이 법 시행 이후 발생하는 소득분부터 적용") → 2017 귀속부터.
#   · 비영업대금 중 온라인투자연계금융업자 경유 이자 14%: 엔진 I.withholding_rate("온투업이자", year) 연도표를 그대로 쓴다.
#   · 2016~2018 제62조1호 나목 단서·2호 가목 1)(조세특례제한법 제104조의27 배당)은 2018.12.31. 삭제 — 지원하지 않음.
_DTC_비영업대금 = ("비영업대금이익", "소비대차전환이자")   # 제16조①11호 (소비대차전환이자도 같은 호, 시행령 제26조③)


@mcp.tool(annotations=CALC)
@_guard
def dividend_tax_credit(year: YEAR,
                        items: Annotated[list | None, Field(description="금융소득 항목 — financial_income_aggregate의 items와 같은 형식 [{kind, amount, withheld?, ...}]. 비영업대금이익·소비대차전환이자는 온라인투자연계금융업자 경유면 p2p=true. items를 주면 아래 숫자 입력(grossup·financial_14·financial_25·partner_dividend)은 비움")] = None,
                        other_base: Annotated[int | None, Field(description="이자소득등(금융소득, Gross-up 포함)을 뺀 나머지 종합소득 과세표준(원) — 출자공동사업자배당은 여기에 포함. taxable_base와 둘 중 하나만", ge=0)] = None,
                        taxable_base: Annotated[int | None, Field(description="종합소득 과세표준 전체(원, 금융소득·Gross-up 포함). other_base 대신 줄 때만", ge=0)] = None,
                        grossup: Annotated[int | None, Field(description="숫자 입력: Gross-up 금액(원, financial_income_aggregate 결과 'Gross-up')", ge=0)] = None,
                        financial_14: Annotated[int | None, Field(description="숫자 입력: 종합과세되는 이자소득등 중 14% 적용분(원, Gross-up 제외, 출자공동사업자배당 제외)", ge=0)] = None,
                        financial_25: Annotated[int | None, Field(description="숫자 입력: 종합과세되는 이자소득등 중 25% 적용분(원, 비영업대금의 이익)", ge=0)] = None,
                        partner_dividend: Annotated[int | None, Field(description="숫자 입력: 출자공동사업자 배당소득금액(원, other_base 안에 들어 있는 금액)", ge=0)] = None) -> dict:
    """배당세액공제와 금융소득 종합과세 비교산출세액(소득세법 제56조·제62조·제15조2호).
언제: 금융소득이 2천만원을 넘어 종합과세될 때 일반산출세액(제62조1호)·비교산출세액(제62조2호)·종합소득 산출세액·배당세액공제액(Gross-up과 한도 중 적은 금액)을 구할 때.
입력: 귀속연도(2016~2026) → year, 금융소득 항목(financial_income_aggregate 형식) → items 또는 그 결과 숫자(grossup, financial_14, financial_25, partner_dividend), 금융소득을 뺀 종합소득 과세표준 → other_base(또는 과세표준 전체 → taxable_base).
결과: 일반산출세액(1호: (이자소득등 − 2천만 + 다른 과세표준) × 기본세율 + 2천만 × 14%), 비교산출세액(2호: 이자소득등(Gross-up 제외) × 원천징수세율 + 다른 과세표준 × 기본세율, 출자공동사업자배당 비교세액 포함), 종합소득 산출세액(큰 금액), 배당세액공제 한도(산출세액 − 비교산출세액), 배당세액공제액 = min(Gross-up, 한도), 공제 후 세액.
연도: 2016~2026. Gross-up 가산율 11%(2016~2023)·10%(2024~). 원천징수되지 않은 비영업대금 이익은 2016년 14%, 2017년부터 25%(제62조2호 가목, 법률 제14389호 부칙 제2조①).
지원하지 않음(입력 확인 필요로 멈춤): 종합소득공제가 금융소득 외 소득보다 커서 다른 과세표준이 음수(또는 출자공동사업자배당보다 작음)인 경우, 조세특례제한법 제104조의27 배당(2016~2018), 2025년 이후 Gross-up 제외 사유 개정분 판정(Gross-up은 financial_income_aggregate 판정 그대로), 외국납부세액공제 등 다른 세액공제·감면(제15조2호 — 이 도구 결과 뒤에 적용), 비거주자."""
    M = 20_000_000
    숫자입력 = {"grossup": grossup, "financial_14": financial_14, "financial_25": financial_25, "partner_dividend": partner_dividend}
    문제 = []
    if items is not None and any(v is not None for v in 숫자입력.values()):
        문제.append("items와 숫자 입력(grossup·financial_14·financial_25·partner_dividend)을 함께 줌 — 한쪽만")
    if items is None and all(v is None for v in 숫자입력.values()):
        문제.append("금융소득 입력 없음 — items 또는 grossup·financial_14·financial_25·partner_dividend 필요")
    if (other_base is None) == (taxable_base is None):
        문제.append("other_base(금융소득 뺀 과세표준)와 taxable_base(과세표준 전체) 중 하나만 입력")
    if 문제:
        return {"error": "입력 확인 필요 — 계산하지 않음", "문제": 문제,
                "예시": {"year": 2024, "items": [{"kind": "예금이자", "amount": 10_000_000}, {"kind": "상장법인현금배당", "amount": 40_000_000}], "other_base": 50_000_000}}

    rate_14 = I.withholding_rate("이자", year)          # 제129조①1호라목·2호나목
    rate_25 = I.withholding_rate("비영업대금", year)     # 제129조①1호나목
    판단 = []
    if items is not None:
        import copy
        its = copy.deepcopy(items)
        p2p_flags = []
        for i, it in enumerate(its):
            if not isinstance(it, dict):
                return {"error": "입력 확인 필요 — 계산하지 않음", "문제": [f"항목 {i+1}: {{kind, amount, ...}} 형식이어야 함"]}
            p = it.pop("p2p", None)
            if p is not None and it.get("kind") not in _DTC_비영업대금:
                return {"error": "입력 확인 필요 — 계산하지 않음", "문제": [f"항목 {i+1}({it.get('kind')}): p2p는 비영업대금이익·소비대차전환이자에만"]}
            p2p_flags.append(bool(p))
        agg = financial_income_aggregate.__wrapped__(year, copy.deepcopy(its))
        if "error" in agg:
            return agg
        종합과세 = not agg["결과"]["분리과세"]
        gu = agg["결과"]["Gross-up"]
        partner = agg["계산"]["출자공동사업자배당(always 포함)"]
        f14 = f25 = 0
        for i, it in enumerate(its):
            k = it.get("kind")
            if k == "출자공동사업자배당":
                continue
            withheld = it.get("withheld", True) is not False
            if not 종합과세 and withheld:
                continue          # 2천만원 이하 분리과세 — 원천징수된 것은 과세표준에 없음
            one = financial_income_aggregate.__wrapped__(year, [copy.deepcopy(it)])["계산"]
            amt = one["이자"] + one["배당(Gross-up 비대상)"] + one["배당(Gross-up 대상)"]
            if not amt:
                continue
            if k in _DTC_비영업대금:
                if p2p_flags[i]:
                    r, why = I.withholding_rate("온투업이자", year), "온라인투자연계금융업자 경유 이자 — 제129조①1호나목 단서"
                elif not withheld and year <= 2016:
                    r, why = rate_14, "원천징수되지 않은 소득 — 2016년 제62조2호 가목 단서: 제129조①1호라목 14%"
                else:
                    r, why = rate_25, "비영업대금의 이익 — 제129조①1호나목" + (" (원천징수되지 않음: 제62조2호 가목 2), 2017년부터)" if not withheld else "")
            else:
                r, why = rate_14, "제129조①1호라목·2호나목" + (" (원천징수되지 않음: 제62조2호 가목 단서)" if not withheld else "")
            if r == rate_25:
                f25 += amt
            else:
                f14 += amt
            판단.append({"항목": k, "금액": amt, "세율": r, "이유": why})
        F = agg["결과"]["이자소득금액"] + agg["결과"]["배당소득금액"] - partner   # 과세표준에 들어간 이자소득등(Gross-up 포함)
        출처 = "items → financial_income_aggregate"
    else:
        gu, f14, f25, partner = grossup or 0, financial_14 or 0, financial_25 or 0, partner_dividend or 0
        종합과세 = f14 + f25 > M
        if gu and not 종합과세:
            return {"error": "입력 확인 필요 — 계산하지 않음",
                    "문제": [f"이자소득등(Gross-up 제외) {f14 + f25:,}원이 2천만원 이하인데 Gross-up {gu:,}원 — Gross-up은 2천만원 초과 배당에만 생김(제56조④, 제17조③)"]}
        F = f14 + f25 + gu
        출처 = "숫자 입력"

    if other_base is None:
        other_base = taxable_base - F
        if other_base < 0:
            return {"error": "입력 확인 필요 — 계산하지 않음",
                    "문제": [f"과세표준 {taxable_base:,}원 < 이자소득등(Gross-up 포함) {F:,}원 — 종합소득공제가 금융소득 외 소득보다 큰 경우는 지원하지 않음(다른 과세표준을 음수로 둘지 판단 필요)"]}
    if partner > other_base:
        return {"error": "입력 확인 필요 — 계산하지 않음",
                "문제": [f"출자공동사업자배당 {partner:,}원이 다른 과세표준 {other_base:,}원보다 큼 — 종합소득공제를 어디서 뺐는지 판단이 필요해 지원하지 않음(제62조2호 나목 단서)"]}

    # 제62조2호: 가목(이자소득등 × 원천징수세율) + 나목(다른 종합소득 산출세액, 출자공동사업자배당 비교세액 하한)
    가목 = int(f14 * rate_14) + int(f25 * rate_25)
    나목_본문 = I.income_tax(other_base, year)
    나목_비교 = int(partner * rate_14) + I.income_tax(other_base - partner, year) if partner else 0   # 제129조①1호라목 14%
    나목 = max(나목_본문, 나목_비교)
    비교산출 = 가목 + 나목
    if 종합과세:
        # 제62조1호: (이자소득등 − 2천만 + 다른 종합소득) × 기본세율 + 2천만 × 14%
        일반산출 = I.income_tax(F - M + other_base, year) + int(M * rate_14)
        산출 = max(일반산출, 비교산출)
        방식 = "1호(종합과세)" if 일반산출 >= 비교산출 else "2호(비교과세)"
    else:
        일반산출 = None
        산출 = 비교산출
        방식 = "2호(2천만원 이하 — 제62조 본문 후단)"
    한도 = max(0, 산출 - 비교산출)
    공제 = min(gu, 한도)
    rate_gu = vocab.rule("상장법인현금배당", year).get("Gross-up가산율", 0.1)

    # ── 계산식 단계 (SPEC_export_workbook_c 틀) — 기본세율은 그 연도 구간표 전부를 식에 ──
    계산식_입력 = [
        {"이름": "이자소득등_14", "값": f14, "설명": f"종합과세 이자소득등 중 원천징수세율 {rate_14:.0%} 적용분(Gross-up·출자공동사업자배당 제외) — {출처}"},
        {"이름": "이자소득등_25", "값": f25, "설명": f"비영업대금의 이익 {rate_25:.0%} 적용분(제129조①1호나목)"},
        {"이름": "GrossUp", "값": gu, "설명": "Gross-up 금액(제17조③ 단서, financial_income_aggregate 판정)"},
        {"이름": "출자공동사업자배당", "값": partner, "설명": "출자공동사업자 배당소득금액 — 다른 과세표준 안에 포함(제62조2호 나목 단서)"},
        {"이름": "원천징수세율_14", "값": rate_14, "설명": "제129조①1호라목·2호나목"},
        {"이름": "원천징수세율_25", "값": rate_25, "설명": "제129조①1호나목"},
        {"이름": "종합과세기준금액", "값": M, "설명": "이자소득등 2천만원(제14조③6호, 제62조)"},
    ]
    계산식 = []
    if F == f14 + f25 + gu:
        계산식.append({"이름": "이자소득등", "식": "이자소득등_14+이자소득등_25+GrossUp", "값": F,
                     "설명": "과세표준에 들어간 이자소득등(Gross-up 포함, 출자공동사업자배당 제외)"})
    else:
        계산식_입력.append({"이름": "이자소득등", "값": F, "설명": "과세표준에 들어간 이자소득등(Gross-up 포함) — financial_income_aggregate 판정 결과"})
    if taxable_base is not None:
        계산식_입력.append({"이름": "과세표준전체", "값": taxable_base, "설명": "종합소득 과세표준 전체(금융소득·Gross-up 포함)"})
        계산식.append({"이름": "다른과세표준", "식": "과세표준전체-이자소득등", "값": other_base, "설명": "금융소득을 뺀 다른 종합소득 과세표준"})
    else:
        계산식_입력.append({"이름": "다른과세표준", "값": other_base, "설명": "금융소득을 뺀 다른 종합소득 과세표준(입력)"})
    if 종합과세:
        계산식 += [
            {"이름": "일반_과세표준", "식": "이자소득등-종합과세기준금액+다른과세표준", "값": F - M + other_base,
             "설명": "제62조1호: 이자소득등 − 2천만원 + 다른 과세표준"},
            {"이름": "일반_기본세율세액", "식": I.income_tax_expr("일반_과세표준", year), "값": I.income_tax(F - M + other_base, year),
             "설명": f"기본세율(제55조①, {year} 귀속 구간표)"},
            {"이름": "일반산출세액", "식": "일반_기본세율세액+ROUNDDOWN(종합과세기준금액*원천징수세율_14,0)", "값": 일반산출,
             "설명": "제62조1호: 기본세율 세액 + 2천만원 × 14%"},
        ]
    계산식 += [
        {"이름": "비교_가목", "식": "ROUNDDOWN(이자소득등_14*원천징수세율_14,0)+ROUNDDOWN(이자소득등_25*원천징수세율_25,0)", "값": 가목,
         "설명": "제62조2호 가목: 이자소득등(Gross-up 제외) × 원천징수세율"},
        {"이름": "비교_나목_본문", "식": I.income_tax_expr("다른과세표준", year), "값": 나목_본문,
         "설명": "제62조2호 나목: 다른 과세표준 × 기본세율"},
    ]
    if partner:
        계산식 += [
            {"이름": "다른과세표준_출자공동제외", "식": "다른과세표준-출자공동사업자배당", "값": other_base - partner, "설명": "다른 과세표준 − 출자공동사업자배당"},
            {"이름": "출자공동_기본세율세액", "식": I.income_tax_expr("다른과세표준_출자공동제외", year),
             "값": I.income_tax(other_base - partner, year), "설명": "기본세율(제55조①)"},
            {"이름": "비교_나목_출자공동", "식": "ROUNDDOWN(출자공동사업자배당*원천징수세율_14,0)+출자공동_기본세율세액", "값": 나목_비교,
             "설명": "제62조2호 나목 단서: 출자공동사업자배당 × 14% + 나머지 × 기본세율"},
            {"이름": "비교_나목", "식": "MAX(비교_나목_본문,비교_나목_출자공동)", "값": 나목, "설명": "제62조2호 나목(큰 금액)"},
        ]
    else:
        계산식.append({"이름": "비교_나목", "식": "비교_나목_본문", "값": 나목, "설명": "출자공동사업자배당 없음 — 본문 그대로"})
    계산식 += [
        {"이름": "비교산출세액", "식": "비교_가목+비교_나목", "값": 비교산출, "설명": "제62조2호"},
        {"이름": "종합소득산출세액", "식": "MAX(일반산출세액,비교산출세액)" if 종합과세 else "비교산출세액", "값": 산출,
         "설명": "제62조 각 호 중 큰 금액" if 종합과세 else "이자소득등 2천만원 이하 — 비교산출세액(제62조 본문 후단)"},
        {"이름": "배당세액공제한도", "식": "MAX(0,종합소득산출세액-비교산출세액)", "값": 한도, "설명": "제56조④·제15조2호: 산출세액 − 비교산출세액"},
        {"이름": "배당세액공제액", "식": "MIN(GrossUp,배당세액공제한도)", "값": 공제, "설명": "제56조①: Gross-up과 한도 중 적은 금액"},
        {"이름": "공제후세액", "식": "종합소득산출세액-배당세액공제액", "값": 산출 - 공제, "설명": "산출세액 − 배당세액공제"},
    ]
    return _ok({"일반산출세액": 일반산출, "비교산출세액": 비교산출, "종합소득산출세액": 산출, "적용": 방식,
                "Gross-up": gu, "배당세액공제한도": 한도, "배당세액공제액": 공제, "공제후세액": 산출 - 공제},
               "소득세법 제56조①④·제62조·제15조2호, 시행령 제116조의2",
               계산={"입력": 출처, "종합과세": 종합과세, "이자소득등(Gross-up 포함)": F, "다른 종합소득 과세표준": other_base,
                     "종합소득 과세표준": F + other_base, "Gross-up 가산율": rate_gu,
                     "2호가목_14%대상": f14, "2호가목_25%대상": f25, "2호가목": 가목,
                     "2호나목_본문": 나목_본문, "2호나목_출자공동비교세액": 나목_비교, "2호나목": 나목,
                     "출자공동사업자배당": partner},
               판단=판단, 계산식_입력=계산식_입력, 계산식=계산식,
               산식="일반산출세액 = 기본세율(이자소득등 − 2천만 + 다른 과세표준) + 2천만 × 14%(제62조1호). "
                    "비교산출세액 = 이자소득등(Gross-up 제외) × 원천징수세율 + MAX(기본세율(다른 과세표준), 출자공동배당 × 14% + 기본세율(다른 과세표준 − 출자공동배당))(제62조2호). "
                    "산출세액 = MAX(일반, 비교)(2천만원 이하면 비교). 배당세액공제 = MIN(Gross-up, 산출세액 − 비교산출세액)(제56조, 제15조2호 후단).",
               주의="외국납부세액공제 등 다른 세액공제·감면은 공제후세액에서 이어서 적용(제15조2호). 2025년 이후 Gross-up 제외 사유 개정은 이 도구에서 따로 판정하지 않음.")


def _기타소득_분리_원천징수(소득: int, rule: dict) -> tuple:
    """복권·승마투표권·슬롯머신 등(제14조③8호라목·마목) 원천징수세액: 소득금액 3억원 이하 20%, 초과분 30%(제129조①6호가목·라목)."""
    기준 = rule.get("고율기준금액")
    세율 = rule.get("원천징수세율", 0.2)
    if 기준 is not None and 소득 > 기준:
        고율 = rule.get("고율세율", 0.3)
        return int(기준 * 세율 + (소득 - 기준) * 고율), f"{기준:,} × {세율 * 100:.0f}% + 초과분 {소득 - 기준:,} × {고율 * 100:.0f}%"
    return int(소득 * 세율), f"{소득:,} × {세율 * 100:.0f}%"


@mcp.tool(annotations=CALC)
@_guard
def other_income_aggregate(year: YEAR,
                           items: Annotated[list, Field(description="기타소득 항목 [{kind, amount, ...}] — kind: 산업재산권대여·공익사업지상권대여·서화골동품양도·일시적인적용역·뇌물·위약금·배상금·주택입주지체상금·연금외수령·직무발명보상금_퇴직후·승마투표권환급금·슬롯머신당첨금")]) -> dict:
    """거주자의 기타소득 항목을 원래 사실관계로 받아 ① 항목별 과세 구분(종합/분리/비과세/제외) ② 비과세·필요경비 ③ 소득금액·원천징수세액 ④ 종합소득 합산액과 기납부세액까지 계산한다(소득세법 제21조·제12조·제14조③·제37조·제84조·제127조·제129조, 시행령 제18조②·제21조·제87조).
    언제: 종합소득세 신고에서 기타소득의 과세 구분·필요경비·원천징수세액·종합소득 합산액을 계산할 때.
    입력: 귀속연도(2016~2026) → year, 기타소득 항목 목록 [{kind, amount, ...}] → items.
      kind 허용값(어휘 기반, 별칭 포함):
        · 산업재산권대여 — 광업권·어업권·산업재산권·산업정보 등을 양도·대여하고 받는 금품
        · 공익사업지상권대여 — 공익사업 관련 지역권·지상권 설정·대여 대가(필요경비율은 산업재산권대여와 같음, 공익사업과 무관한 지상권 대여는 사업소득)
        · 서화골동품양도 — 개당·점당 6천만원 이상 서화·골동품 양도(필요경비: 2020 과세기간부터 1억 이하 90%·초과분 80%, 2019년까지 80%, held_years≥10이면 90% — 2020.2.11. 개정 부칙 제36조, business_setup=true(2021~ 사업장 갖춤)면 사업소득이라 제외, museum=true 박물관·미술관 양도 비과세, living_domestic_artist=true 생존 국내 작가 작품 과세 제외, 무조건 분리과세 20%)
        · 일시적인적용역 — 고용관계 없이 일시적으로 제공한 강연·방송 해설·자문 등 인적용역의 대가
          (산업재산권대여·공익사업지상권대여·일시적인적용역 필요경비율: ~2017 80%, 2018은 receipt_date(받은 날) 필수 — 2018.3.31.까지 80%·2018.4.1.부터 70%, 2019~ 60%, 시행령 제87조)
        · 뇌물 — 뇌물로 받은 금품(몰수·추징 확정액은 forfeited로 입력, 기본 0)
        · 위약금·배상금 — 계약의 위약·해약으로 받는 위약금·배상금(계약금 대체분은 deposit_substituted로 입력, 기본 0, 주택입주 지체상금 제외. 계약금 대체분은 원천징수되지 않아 2019년 이전에는 선택적 분리과세 대상 아님 — 제14조③8호가목 2020.1.1. 개정)
        · 주택입주지체상금 — 주택 입주 지연으로 받는 지체상금
        · 연금외수령 — 연금계좌의 세액공제분·운용수익을 연금 외의 방식으로 수령한 금액(무조건 분리과세)
        · 직무발명보상금_퇴직후 — 퇴직 후 사용자등으로부터 받는 직무발명보상금(제21조①22호의2, 같은 해 재직 중 받은 비과세액은 employment_exempt_used로 필수 입력, 비과세 연 한도 2017~2018 300만·2019~2023 500만·2024~ 700만(시행령 제18조②), 2016은 22호의2가 없어 계산하지 않음, 선택적 분리과세는 2020~)
        · 배임수재 — 배임수재로 받은 금품(무조건 종합과세, 원천징수 제외, 소득세법 제21조①24호)
        · 알선수재 — 알선수재로 받은 금품(무조건 종합과세, 원천징수 제외, 소득세법 제21조①24호)
        · 복권당첨금 — 복권 당첨금(제21조①2호, 제14조③8호라목 무조건 분리과세, 필요경비는 actual_cost(구입비)만, 원천징수 20%·소득금액 3억원 초과분 30%(제129조①6호가목). 과세최저한: year≥2023은 당첨금(분할 지급이면 합계) 건별 200만원 이하(제84조2호), 그 전은 소득금액 건별 5만원 이하)
        · 승마투표권환급금 — 승마·경륜·경정·소싸움·체육진흥투표권 환급금(제21조①4호, 제14조③8호마목 무조건 분리과세, 원천징수 20%·소득금액 3억원 초과분 30%). unit_bet=적중된 투표권의 단위투표금액(필요경비, 제37조①1호, 기본 0), per_bet_refund=적중한 개별투표당 환급금(기본 amount), ticket_total=건별 권면 표시 금액 합계. 과세최저한(제84조1호): ticket_total≤10만원이고 (개별투표당 환급금≤10만원 또는 (≤unit_bet×100이면서 ≤200만원)) — 가·나목에 해당하면 ticket_total 필수
        · 슬롯머신당첨금 — 슬롯머신·투전기 등 당첨금품(제21조①14호, 제14조③8호마목 무조건 분리과세, 원천징수 20%·소득금액 3억원 초과분 30%). stake=당첨 당시 투입한 금액(필요경비, 제37조①2호, 기본 0). 당첨금 건별 200만원 이하 과세최저한(제84조2호)
    결과: 판단(항목별 과세 구분·근거)·결과(종합소득 합산액·기납부세액·분리과세·비과세·선택적 분리과세 가능)·계산(항목별 소득금액·원천징수세액)."""

    _v = _validate_items(items, "kind", _specs_from_vocab("other_income_aggregate", "kind"))
    if _v: return _v

    종합소득_합산_기타소득금액 = 0
    기납부_원천징수세액 = 0
    분리과세_기타소득금액 = 0
    분리과세_원천징수세액 = 0
    비과세_합계 = 0
    항목별 = []
    산식_목록 = []
    근거_목록 = []
    뇌물_여부 = False
    # 제14조③8호가목 열거에서 빠져 선택적 분리과세를 못 하는 종합과세분(2019년 이전: 22호의2·계약금 대체 위약금)
    선택불가_소득 = 0
    선택불가_설명 = []

    # 금액 칸에 bool이 오면 오류 (True가 1원으로 계산되던 문제 방지)
    _금액필드 = ["amount", "deposit_substituted", "forfeited", "actual_cost", "actual_expense",
               "unit_bet", "ticket_total", "per_bet_refund", "stake", "employment_exempt_used"]
    for it in items:
        for f in _금액필드:
            v = it.get(f)
            if isinstance(v, bool):
                return {"error": f"입력 확인 필요 — 계산하지 않음",
                        "문제": [f"항목 '{it.get('kind', '')}': {f}에 bool이 입력됨 — 숫자여야 함"]}

    for it in items:
        k = it.get("kind", "")
        amount = int(it.get("amount", 0))
        rule = vocab.rule(k, year)
        basis = vocab.basis(k, year)

        if k in ("뇌물", "배임수재", "알선수재"):
            forfeited = int(it.get("forfeited", 0)) if k == "뇌물" else 0
            if k == "뇌물" and forfeited > amount:
                return {"error": f"뇌물 몰수·추징 확정액 {forfeited:,}원이 받은 금액 {amount:,}원을 초과 — 입력 확인 필요"}
            소득 = amount - forfeited
            원천징수세액 = 0
            항목별.append({"kind": k, "총수입금액": amount, "비과세": 0, "필요경비": 0,
                           "소득금액": 소득, "과세구분": "종합", "원천징수세액": 원천징수세액, "근거": basis})
            종합소득_합산_기타소득금액 += 소득
            if k == "뇌물":
                산식_목록.append(f"{k}: {amount:,}원 − 몰수·추징 {forfeited:,}원 = 소득 {소득:,}원 (제127조①6호다목 원천징수 제외)")
            else:
                산식_목록.append(f"{k}: {amount:,}원 = 소득 {소득:,}원 (제127조①6호다목 원천징수 제외, 무조건 종합과세)")
            근거_목록.extend(basis)

        elif k == "위약금·배상금":
            deposit_substituted = int(it.get("deposit_substituted", 0))
            if deposit_substituted > amount:
                return {"error": f"계약금 대체 위약금 {deposit_substituted:,}원이 위약금 총액 {amount:,}원을 초과 — 입력 확인 필요"}
            소득 = amount
            원천징수세액 = int((amount - deposit_substituted) * rule["원천징수세율"])
            # 과세최저한(소득세법 제84조3호): 건별 기타소득금액이 5만원 이하이면 과세하지 않음
            과세최저한 = rule.get("과세최저한", 50000)
            if 소득 <= 과세최저한:
                소득 = 0
                원천징수세액 = 0
                과세구분표시 = "과세최저한"
            else:
                과세구분표시 = "종합"
            항목별.append({"kind": k, "총수입금액": amount, "비과세": 0, "필요경비": 0,
                           "소득금액": 소득, "과세구분": 과세구분표시, "원천징수세액": 원천징수세액, "근거": basis})
            종합소득_합산_기타소득금액 += 소득
            기납부_원천징수세액 += 원천징수세액
            if year < 2020 and 소득 > 0 and deposit_substituted > 0:
                # 2019 시행본까지 제14조③8호가목은 "제127조에 따라 원천징수된 소득"만 — 계약금 대체분(제127조①6호나목 원천징수 제외)은 빠짐
                선택불가_소득 += min(deposit_substituted, 소득)
                선택불가_설명.append(f"계약금 대체 위약금 {min(deposit_substituted, 소득):,}원(원천징수되지 않은 소득, 2020.1.1. 개정 전 제14조③8호가목)")
            산식_목록.append(f"{k}: 총액 {amount:,}원(계약금 대체 {deposit_substituted:,}원 포함), 원천징수 ({amount:,} − {deposit_substituted:,}) × 20% = {원천징수세액:,}원")
            근거_목록.extend(basis)

        elif k == "연금외수령":
            소득 = amount
            원천징수세액 = int(amount * rule["원천징수세율"])
            분리과세_기타소득금액 += 소득
            분리과세_원천징수세액 += 원천징수세액
            항목별.append({"kind": k, "총수입금액": amount, "비과세": 0, "필요경비": 0,
                           "소득금액": 소득, "과세구분": "분리과세", "원천징수세액": 원천징수세액, "근거": basis})
            산식_목록.append(f"{k}: {amount:,}원 × 15% = {원천징수세액:,}원 (제14조③8호나목 무조건 분리과세, 종합소득 합산액·기납부세액에서 제외)")
            근거_목록.extend(basis)

        elif k == "직무발명보상금_퇴직후":
            if rule.get("미지원"):
                return {"error": "입력 확인 필요 — 계산하지 않음",
                        "문제": [f"직무발명보상금_퇴직후({year}년): {rule['미지원']} — 이 종류로 계산하지 않음"]}
            employment_exempt_used = int(it["employment_exempt_used"])
            비과세한도 = rule["비과세한도"]
            비과세 = max(0, min(amount, 비과세한도 - employment_exempt_used))
            소득 = amount - 비과세
            원천징수세율 = rule.get("원천징수세율", 0.2)
            원천징수세액 = int(소득 * 원천징수세율) if 소득 > 0 else 0
            비과세소득 = 비과세
            비과세_합계 += 비과세
            if year < 2020 and 소득 > 0:
                선택불가_소득 += 소득
                선택불가_설명.append(f"직무발명보상금_퇴직후 {소득:,}원(제21조①22호의2는 2020.1.1. 개정 전 제14조③8호가목 열거에 없음)")
            종합소득_합산_기타소득금액 += 소득
            기납부_원천징수세액 += 원천징수세액
            항목별.append({"kind": k, "총수입금액": amount, "비과세": 비과세, "필요경비": 0,
                           "소득금액": 소득, "과세구분": "종합", "원천징수세액": 원천징수세액, "근거": basis})
            산식_목록.append(f"{k}: {amount:,}원 − 비과세 {비과세:,}원(한도 {비과세한도:,}원 − 재직 중 비과세 사용 {employment_exempt_used:,}원) = 소득 {소득:,}원, 원천징수 {소득:,} × {원천징수세율 * 100:.0f}% = {원천징수세액:,}원")
            근거_목록.extend(basis)

        elif k == "복권당첨금":
            actual_cost = int(it.get("actual_cost", 0))
            if actual_cost > amount:
                return {"error": "입력 확인 필요 — 계산하지 않음",
                        "문제": [f"복권당첨금: 구입비 {actual_cost:,}원이 당첨금 {amount:,}원을 초과"]}
            소득 = amount - actual_cost
            필요경비 = actual_cost
            과세최저한 = rule.get("과세최저한", 50000)
            # 제84조: 2023.1.1. 이후 지급분은 당첨금(분할 지급이면 합계) 건별 200만원 이하(2호), 그 전은 기타소득금액 건별 5만원 이하(3호)
            당첨금기준 = rule.get("과세최저한기준") == "당첨금"
            if (amount if 당첨금기준 else 소득) <= 과세최저한:
                최저한설명 = (f"당첨금 {amount:,}원 ≤ {과세최저한:,}원(제84조2호, 2023.1.1. 이후 지급분)" if 당첨금기준
                          else f"소득금액 {소득:,}원 ≤ {과세최저한:,}원(제84조3호)")
                소득, 원천징수세액, 과세구분표시 = 0, 0, "과세최저한"
            else:
                원천징수세액, 세율설명 = _기타소득_분리_원천징수(소득, rule)
                과세구분표시 = "분리과세"
            분리과세_기타소득금액 += 소득
            분리과세_원천징수세액 += 원천징수세액
            항목별.append({"kind": k, "총수입금액": amount, "비과세": 0, "필요경비": 필요경비,
                           "소득금액": 소득, "과세구분": 과세구분표시, "원천징수세액": 원천징수세액, "근거": basis})
            if 과세구분표시 == "과세최저한":
                산식_목록.append(f"{k}: {amount:,}원 − 구입비 {actual_cost:,}원 → {최저한설명} 과세최저한으로 과세 제외, 원천징수 0")
            else:
                산식_목록.append(f"{k}: {amount:,}원 − 구입비 {actual_cost:,}원 = 소득금액 {소득:,}원, 원천징수 {세율설명} = {원천징수세액:,}원 (제14조③8호라목 무조건 분리과세, 제129조①6호)")
            근거_목록.extend(basis)

        elif k == "승마투표권환급금":
            unit_bet = int(it.get("unit_bet", 0))
            # 적중한 개별투표당 환급금(기본: 이 항목의 환급금), 건별 권면 표시 금액 합계(제84조1호)
            개별환급금 = int(it.get("per_bet_refund", amount))
            if unit_bet > amount or 개별환급금 > amount:
                return {"error": "입력 확인 필요 — 계산하지 않음",
                        "문제": [f"승마투표권환급금: 단위투표금액 {unit_bet:,}원·개별투표당 환급금 {개별환급금:,}원은 환급금 {amount:,}원을 넘을 수 없음"]}
            소득 = amount - unit_bet
            # 제84조1호 가·나목: 권면 합계 요건을 빼고 먼저 본다
            가목 = 개별환급금 <= 100000
            나목 = unit_bet > 0 and 개별환급금 <= unit_bet * 100 and 개별환급금 <= 2000000
            과세최저한_조건 = False
            if 가목 or 나목:
                if it.get("ticket_total") in (None, ""):
                    return {"error": "입력 확인 필요 — 계산하지 않음",
                            "문제": [f"승마투표권환급금 {amount:,}원: 과세최저한(소득세법 제84조1호)은 건별 권면 표시 금액 합계가 10만원 이하여야 함 — ticket_total(권면 금액 합계) 입력 필요"]}
                ticket_total = int(it["ticket_total"])
                과세최저한_조건 = ticket_total <= 100000
            if 과세최저한_조건:
                소득, 원천징수세액, 과세구분표시 = 0, 0, "과세최저한"
            else:
                원천징수세액, 세율설명 = _기타소득_분리_원천징수(소득, rule)
                과세구분표시 = "분리과세"
            분리과세_기타소득금액 += 소득
            분리과세_원천징수세액 += 원천징수세액
            항목별.append({"kind": k, "총수입금액": amount, "비과세": 0, "필요경비": unit_bet,
                           "소득금액": 소득, "과세구분": 과세구분표시, "원천징수세액": 원천징수세액, "근거": basis})
            if 과세최저한_조건:
                사유 = "개별투표당 환급금 10만원 이하(가목)" if 가목 else "단위투표금액의 100배 이하이면서 200만원 이하(나목)"
                산식_목록.append(f"{k}: 권면 합계 {ticket_total:,}원 ≤ 100,000원이고 {사유} → 과세최저한(소득세법 제84조1호)으로 과세 제외, 원천징수 0")
            else:
                산식_목록.append(f"{k}: 환급금 {amount:,}원 − 단위투표금액 {unit_bet:,}원 = 소득금액 {소득:,}원, 원천징수 {세율설명} = {원천징수세액:,}원 (소득세법 제21조①4호·제37조①1호, 제14조③8호마목 무조건 분리과세)")
            근거_목록.extend(basis)

        elif k == "슬롯머신당첨금":
            stake = int(it.get("stake", 0))
            if stake > amount:
                return {"error": "입력 확인 필요 — 계산하지 않음",
                        "문제": [f"슬롯머신당첨금: 투입금액 {stake:,}원이 당첨금 {amount:,}원을 초과"]}
            과세최저한 = 2000000
            if amount <= 과세최저한:
                소득, 원천징수세액, 과세구분표시 = 0, 0, "과세최저한"
            else:
                소득 = amount - stake      # 제37조①2호: 당첨 당시 슬롯머신등에 투입한 금액이 필요경비
                원천징수세액, 세율설명 = _기타소득_분리_원천징수(소득, rule)
                과세구분표시 = "분리과세"
            분리과세_기타소득금액 += 소득
            분리과세_원천징수세액 += 원천징수세액
            항목별.append({"kind": k, "총수입금액": amount, "비과세": 0, "필요경비": stake,
                           "소득금액": 소득, "과세구분": 과세구분표시, "원천징수세액": 원천징수세액, "근거": basis})
            if 과세구분표시 == "과세최저한":
                산식_목록.append(f"{k}: 당첨금 {amount:,}원 ≤ 2,000,000원 → 과세최저한(소득세법 제84조2호)으로 과세 제외, 원천징수 0")
            else:
                산식_목록.append(f"{k}: 당첨금 {amount:,}원 − 투입금액 {stake:,}원 = 소득금액 {소득:,}원, 원천징수 {세율설명} = {원천징수세액:,}원 (소득세법 제21조①14호·제37조①2호, 제14조③8호마목 무조건 분리과세)")
            근거_목록.extend(basis)

        elif k == "서화골동품양도":
            held = int(it.get("held_years", 0) or 0)
            actual = int(it.get("actual_cost", 0) or 0)
            # 2020.2.11. 개정 시행령 부칙 제36조: 시행일이 속하는 과세기간 전 발생분만 종전 규정 → 2020 과세기간 전체가 구간제
            구간제 = year >= 2020
            if it.get("business_setup") and year >= 2021:
                과세구분표시, 소득, 원천징수세액, 필요경비, 비과세 = "제외", 0, 0, 0, 0
                산식_목록.append(f"{k}: 사업장을 갖추는 등 사업 형태 거래 → 기타소득이 아니라 사업소득(소득세법 제21조②, 2021.1.1. 이후)")
            elif it.get("museum"):
                과세구분표시, 소득, 원천징수세액, 필요경비, 비과세 = "비과세", 0, 0, 0, amount
                산식_목록.append(f"{k}: 박물관·미술관 양도 {amount:,}원 → 비과세(소득세법 제12조5호)")
            elif it.get("living_domestic_artist") or amount < rule.get("과세기준가액", 60000000):
                과세구분표시, 소득, 원천징수세액, 필요경비, 비과세 = "제외", 0, 0, 0, 0
                사유 = "양도일 현재 생존한 국내 원작자 작품" if it.get("living_domestic_artist") else "개당·점당 6천만원 미만"
                산식_목록.append(f"{k}: {amount:,}원 — {사유} → 과세 대상 아님(소득세법 시행령 제41조⑭)")
            else:
                if 구간제:   # 2020 과세기간부터: 1억 이하 90%, 초과분 80%(보유 10년 이상 90%)
                    율경비 = int(amount * 0.9) if (held >= 10 or amount <= 100000000) else int(100000000 * 0.9 + (amount - 100000000) * 0.8)
                    율설명 = '보유 10년 이상 90%' if held >= 10 else ('1억 이하 90%' if amount <= 100000000 else '1억까지 90%·초과분 80%')
                else:        # 2019 과세기간까지: 받은 금액의 80%(보유 10년 이상 90%)
                    율경비 = int(amount * (0.9 if held >= 10 else 0.8))
                    율설명 = '2019년까지 — ' + ('보유 10년 이상 90%' if held >= 10 else '80%')
                필요경비 = max(율경비, actual)
                소득 = amount - 필요경비
                원천징수세액 = int(소득 * rule.get("원천징수세율", 0.2))
                과세구분표시, 비과세 = "분리과세", 0
                분리과세_기타소득금액 += 소득
                분리과세_원천징수세액 += 원천징수세액
                산식_목록.append(f"{k}: {amount:,}원 − 필요경비 max(율 {율경비:,}, 실제 {actual:,}) = 소득금액 {소득:,}원"
                                 f"({율설명}), 원천징수 20% = {원천징수세액:,}원 (무조건 분리과세, 시행령 제87조2호)")
            비과세_합계 += 비과세
            항목별.append({"kind": k, "총수입금액": amount, "비과세": 비과세, "필요경비": 필요경비,
                           "소득금액": 소득, "과세구분": 과세구분표시, "원천징수세액": 원천징수세액, "근거": basis})
            근거_목록.extend(basis)

        elif k in ("산업재산권대여", "공익사업지상권대여", "일시적인적용역", "주택입주지체상금"):
            수입일 = str(it.get("receipt_date") or "")
            if 수입일:
                if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", 수입일) or int(수입일[:4]) != year:
                    return {"error": "입력 확인 필요 — 계산하지 않음",
                            "문제": [f"{k}: receipt_date '{수입일}'는 {year}년 안의 YYYY-MM-DD여야 함"]}
                rule = vocab.rule_on(k, int(수입일.replace("-", "")))
                basis = vocab.basis_on(k, int(수입일.replace("-", "")))
            elif "필요경비율" not in rule:
                # 시행령 제87조: 2018.3.31.까지 발생분 1호나목 100분의 80, 2018.4.1.부터 1호의2 100분의 70(2018.2.13. 개정)
                return {"error": "입력 확인 필요 — 계산하지 않음",
                        "문제": [f"{k}({year}년): 필요경비율이 해당 연도 중에 바뀜(2018.3.31.까지 발생분 80%, 2018.4.1.부터 70% — 소득세법 시행령 제87조1호나목·1호의2) — receipt_date(YYYY-MM-DD, 받은 날) 입력 필요"]}
            필요경비율 = rule.get("필요경비율", 0)
            actual_expense = int(it.get("actual_cost", it.get("actual_expense", 0)))
            과세대상금액 = amount
            필요경비 = min(과세대상금액, max(과세대상금액 * 필요경비율, actual_expense))
            소득 = amount - 필요경비
            원천징수세율 = rule.get("원천징수세율", 0.2)
            원천징수세액 = int(소득 * 원천징수세율) if 소득 > 0 else 0
            # 과세최저한(소득세법 제84조3호): 건별 기타소득금액이 5만원 이하이면 과세하지 않음
            과세최저한 = rule.get("과세최저한", 50000)
            if 소득 <= 과세최저한:
                소득 = 0
                원천징수세액 = 0
                과세구분표시 = "과세최저한"
            else:
                과세구분표시 = "종합"
            종합소득_합산_기타소득금액 += 소득
            기납부_원천징수세액 += 원천징수세액
            항목별.append({"kind": k, "총수입금액": amount, "비과세": 0, "필요경비": int(필요경비),
                           "소득금액": int(소득), "과세구분": 과세구분표시, "원천징수세액": 원천징수세액, "근거": basis})
            if rule.get("주의"): 산식_목록.append(f"{k} 주의: {rule['주의']}")
            if 수입일 and year == 2018:
                산식_목록.append(f"{k}: 받은 날 {수입일} → 2018.3.31.까지 발생분 80%, 2018.4.1.부터 70%(시행령 제87조1호나목·1호의2) 중 {필요경비율 * 100:.0f}% 적용")
            if 소득 == 0:
                산식_목록.append(f"{k}: {amount:,}원 − 필요경비 {int(필요경비):,}원 = 소득금액 {int(amount - 필요경비):,}원 → 건별 5만원 이하 과세최저한(제84조3호)으로 과세 제외, 원천징수 0")
            else:
                산식_목록.append(f"{k}: {amount:,}원 × 필요경비율 {필요경비율 * 100:.0f}% = {int(amount * 필요경비율):,}원, 실제경비 {actual_expense:,}원 → 필요경비 max({int(amount * 필요경비율):,}, {actual_expense:,}) = {int(필요경비):,}원, 소득 {int(소득):,}원, 원천징수 {int(소득):,} × {원천징수세율 * 100:.0f}% = {원천징수세액:,}원")
            근거_목록.extend(basis)

    # 선택적 분리과세 판정(제14조③8호가목): 열거된 호의 원천징수된 기타소득금액 합계가 300만원 이하
    # 뇌물·배임수재·알선수재(23·24호)는 열거 밖. 2019년 이전은 22호의2와 원천징수되지 않은 계약금 대체 위약금도 열거 밖(2019.12.31. 개정, 2020.1.1. 시행)
    _뇌물등_소득 = sum(h["소득금액"] for h in 항목별 if h["kind"] in ("뇌물", "배임수재", "알선수재"))
    _종합대상_소득 = sum(h["소득금액"] for h in 항목별 if h["과세구분"] == "종합" and h["kind"] not in ("뇌물", "배임수재", "알선수재")) - 선택불가_소득

    선택적_분리과세_가능 = _종합대상_소득 > 0 and _종합대상_소득 <= 3_000_000
    선택적_분리과세_적용시_종합합산_기타소득금액 = (_뇌물등_소득 + 선택불가_소득) if 선택적_분리과세_가능 else 종합소득_합산_기타소득금액
    if 선택불가_설명:
        산식_목록.append("선택적 분리과세 대상 아님(반드시 종합과세): " + ", ".join(선택불가_설명))

    return _ok({
        "종합소득_합산_기타소득금액": int(종합소득_합산_기타소득금액),
        "기납부_원천징수세액": int(기납부_원천징수세액),
        "분리과세_기타소득금액": int(분리과세_기타소득금액),
        "분리과세_원천징수세액": int(분리과세_원천징수세액),
        "비과세_합계": int(비과세_합계),
        "선택적_분리과세_가능": 선택적_분리과세_가능,
        "선택적_분리과세_적용시_종합합산_기타소득금액": int(선택적_분리과세_적용시_종합합산_기타소득금액),
    }, 근거_목록, 항목별=항목별, 산식=산식_목록)


# ── 종교인소득(기타소득 제21조①26호) ──────────────────
# 2018.1.1. 시행 — 부칙 <제13558호,2015.12.15> 제1조 단서: "제12조, 제14조, 제21조, 제73조, 제80조, 제145조의3,
# 제155조의6, 제164조 및 제170조의 개정규정은 2018년 1월 1일부터 시행한다." → 2016·2017 과세기간은 과세 전.
# 연도표: (부터, 까지, 값, 근거). 까지=None이면 2026까지.
_종교인_식사대_월한도 = [
    (2018, 2022, 100_000, "시행령 제19조②2호 월 10만원 이하 식사대"),
    # 부칙 <대통령령 제33267호,2023.2.28> 제2조: 이 영 시행일이 속하는 과세기간에 발생하는 소득부터 적용
    (2023, None, 200_000, "시행령 제19조②2호 월 20만원 이하 식사대(개정 2023.2.28, 부칙 제2조 — 2023 과세기간 발생분부터)"),
]
_종교인_보육_월한도 = [
    (2018, 2023, 100_000, False, "소득세법 제12조5호아목4) 출산·6세 이하 자녀 보육 월 10만원 이내"),
    # 부칙 <제19933호,2023.12.31> 제1조 2024.1.1. 시행, 제2조③ 시행 전 지급분은 종전 한도
    (2024, 2025, 200_000, False, "소득세법 제12조5호아목4) 월 20만원 이내(개정 2023.12.31, 2024.1.1. 이후 지급분 — 부칙 제2조③)"),
    # 부칙 <제21221호,2025.12.23> 제1조 2026.1.1. 시행, 제3조 시행 이후 지급받는 소득분부터
    (2026, None, 200_000, True, "소득세법 제12조5호아목4) 해당 자녀 1명당 월 20만원 이내(개정 2025.12.23, 부칙 제3조 — 2026.1.1. 이후 지급분)"),
]
_종교인_자가운전_임차포함 = [
    (2018, 2024, False, "시행령 제19조③2호 본인 소유 차량 — 월 20만원 이내"),
    # 부칙 <대통령령 제35349호,2025.2.28> 제4조: 이 영 시행일이 속하는 과세기간에 지급받는 소득분부터 적용
    (2025, None, True, "시행령 제19조③2호 소유하거나 본인 명의로 임차한 차량 — 월 20만원 이내(개정 2025.2.28, 부칙 제4조 — 2025 과세기간 지급분부터)"),
]
# 시행령 제87조3호 표(2018~2026 같은 표 — 기준일별 law_article로 대조): (구간 상한, 기본액, 초과분 비율, 구간 하한)
_종교인_필요경비표 = [(20_000_000, 0, 0.8, 0), (40_000_000, 16_000_000, 0.5, 20_000_000),
                     (60_000_000, 26_000_000, 0.3, 40_000_000), (None, 32_000_000, 0.2, 60_000_000)]


def _연도값(표, year):
    for row in 표:
        if row[0] <= year and (row[1] is None or year <= row[1]):
            return row[2:]
    raise C.NoYearTable(f"{year}년 종교인소득 연도표 없음")


def _종교인_필요경비_표(받은금액: int) -> tuple[int, str]:
    for 상한, 기본, 율, 하한 in _종교인_필요경비표:
        if 상한 is None or 받은금액 <= 상한:
            금액 = int(기본 + (받은금액 - 하한) * 율)
            if 기본 == 0:
                return 금액, f"{받은금액:,} × {율 * 100:.0f}% = {금액:,}"
            return 금액, f"{기본:,} + ({받은금액:,} − {하한:,}) × {율 * 100:.0f}% = {금액:,}"
    raise ValueError("필요경비 표 구간 없음")


@mcp.tool(annotations=CALC)
@_guard
def religious_income(year: YEAR,
                     amount: Annotated[int, Field(strict=True, description="해당 과세기간에 종교단체로부터 받은 금액 총액(원, 아래 비과세 지급액 포함)", ge=0)],
                     scholarship: Annotated[int, Field(strict=True, description="학자금(원) — 종교 활동 관련 교육·훈련을 위한 학교·평생교육시설 입학금·수업료 등(시행령 제19조①), 전액 비과세", ge=0)] = 0,
                     meals_in_kind: Annotated[int, Field(strict=True, description="종교단체가 현물로 제공한 식사·음식물(원), 전액 비과세(시행령 제19조②1호)", ge=0)] = 0,
                     meal_allowance: Annotated[int, Field(strict=True, description="식사대(원)", ge=0)] = 0,
                     meal_months: Annotated[int | None, Field(description="식사대를 받은 개월 수(1~12)", ge=1, le=12)] = None,
                     meal_provided: Annotated[bool | None, Field(description="식사(현물)를 제공받았는지 — true면 식사대는 과세(시행령 제19조②2호)")] = None,
                     vehicle_allowance: Annotated[int, Field(strict=True, description="자가운전보조금(원) — 실제 여비 대신 종교단체 지급기준에 따라 받은 금액", ge=0)] = 0,
                     vehicle_months: Annotated[int | None, Field(description="자가운전보조금을 받은 개월 수(1~12)", ge=1, le=12)] = None,
                     vehicle: Annotated[Literal["본인소유", "본인명의임차", "기타"] | None, Field(description="차량: 본인소유|본인명의임차(2025~ 비과세)|기타")] = None,
                     expense_reimbursement: Annotated[int, Field(strict=True, description="일직·숙직료, 실비 정도의 여비, 재해로 받은 지급액(원), 전액 비과세(시행령 제19조③1·2·4호)", ge=0)] = 0,
                     religious_activity_fund: Annotated[int, Field(strict=True, description="종교활동비(원) — 규약·의결기구 결정 지급기준에 따라 종교 활동에 쓰도록 받은 금액·물품, 전액 비과세(시행령 제19조③3호·제12조18호)", ge=0)] = 0,
                     childcare_allowance: Annotated[int, Field(strict=True, description="출산·보육 관련 지급액(원, 본인·배우자 출산 또는 6세 이하 자녀 보육)", ge=0)] = 0,
                     childcare_months: Annotated[int | None, Field(description="출산·보육 지급액을 받은 개월 수(1~12)", ge=1, le=12)] = None,
                     childcare_children: Annotated[int | None, Field(description="출산 자녀 또는 과세기간 개시일 기준 6세 이하 자녀 수(0이면 전액 과세)", ge=0)] = None,
                     housing_benefit: Annotated[int, Field(strict=True, description="사택 제공 이익(원, 시행규칙 사택 요건 충족분), 전액 비과세", ge=0)] = 0,
                     actual_expense: Annotated[int, Field(strict=True, description="실제 소요된 필요경비(원) — 표 금액보다 크면 이 금액(시행령 제87조3호 단서)", ge=0)] = 0,
                     withheld: Annotated[bool | None, Field(description="종교단체가 원천징수(간이세액표·연말정산)를 했는지 — 선택적 분리과세·신고의무 판단용, 모르면 생략")] = None,
                     route: Annotated[Literal["기타소득", "근로소득"], Field(description="기타소득(기본) | 근로소득(근로소득으로 원천징수·신고를 선택한 경우 — 이 도구는 계산하지 않고 안내)")] = "기타소득") -> dict:
    """종교인소득(소득세법 제21조①26호, 2018.1.1. 시행) 기타소득 계산 — 비과세(제12조5호아목) → 필요경비 표(시행령 제87조3호) → 기타소득금액 → 선택적 분리과세 가능 여부·원천징수·신고 안내.
    언제: 목사·신부·승려 등 종교관련종사자가 종교단체로부터 받은 사례비 등의 종교인소득금액을 기타소득으로 계산할 때(2018~2026 귀속).
    입력: 귀속연도(2018~2026, 2016·2017은 과세 전이라 오류) → year, 받은 금액 총액 → amount,
      비과세 항목 → scholarship(학자금)·meals_in_kind(현물 식사)·meal_allowance+meal_months+meal_provided(식사대 월 10만원, 2023~ 20만원)·
      vehicle_allowance+vehicle_months+vehicle(자가운전보조금 월 20만원, 본인명의임차는 2025~)·expense_reimbursement(일직·숙직료·실비 여비·재해 지급액)·
      religious_activity_fund(종교활동비)·childcare_allowance+childcare_months+childcare_children(출산·보육 월 10만원, 2024~ 20만원, 2026~ 자녀 1명당 20만원)·housing_benefit(사택),
      실제 필요경비 → actual_expense, 원천징수 여부 → withheld, 계산 경로 → route(기타소득|근로소득).
    결과: 비과세 항목별·과세대상 받은 금액·필요경비(표 금액 vs 실제)·기타소득금액·선택적 분리과세 가능 여부(제14조③8호가목 — 26호 열거, 300만원 이하·원천징수된 소득)·원천징수·신고 안내·근거·연도표.
    지원하지 않음: 근로소득 선택 경로 계산(earned_income_gross_pay + wage_income_tax로 계산), 간이세액표 월별 원천징수세액·연말정산 세액, 비과세 요건(학자금 대상 학교, 사택 요건) 자체의 판정, 퇴직소득 해당 여부(시행령 제42조의2④4호), 다른 기타소득과의 300만원 합산 판정."""
    # 금액 칸은 strict int — bool(True=1원)·문자열은 검증 단계에서 거부
    문제: list[str] = []
    if year < 2018:
        문제.append(f"{year}년은 종교인소득 과세 전 — 소득세법 제21조①26호는 2018.1.1. 시행(부칙 <제13558호,2015.12.15> 제1조 단서)")
    근로_안내 = ("종교인소득을 근로소득으로 원천징수하거나 확정신고하면 근로소득으로 본다(소득세법 제21조"
               + ("③" if year <= 2020 else "④") + " — 2018~2020 ③, 2021~ ④). 그 경우 비과세는 제12조3호(근로소득 비과세)로 보고 "
               "earned_income_gross_pay(총급여) → wage_income_tax(근로소득공제·산출세액)로 계산한다.")
    if route == "근로소득":
        문제.append("route='근로소득' — 이 도구는 기타소득 경로만 계산. " + 근로_안내)
    if meal_allowance > 0 and (meal_months is None or meal_provided is None):
        문제.append("식사대가 있으면 meal_months(개월 수)와 meal_provided(현물 식사 제공 여부)가 필요")
    if vehicle_allowance > 0 and (vehicle_months is None or vehicle is None):
        문제.append("자가운전보조금이 있으면 vehicle_months(개월 수)와 vehicle(본인소유|본인명의임차|기타)이 필요")
    if childcare_allowance > 0 and (childcare_months is None or childcare_children is None):
        문제.append("출산·보육 지급액이 있으면 childcare_months(개월 수)와 childcare_children(출산 자녀 또는 개시일 기준 6세 이하 자녀 수)이 필요")
    비과세_입력합 = (scholarship + meals_in_kind + meal_allowance + vehicle_allowance + expense_reimbursement
                + religious_activity_fund + childcare_allowance + housing_benefit)
    if not 문제 and 비과세_입력합 > amount:
        문제.append(f"비과세 항목 입력 합계 {비과세_입력합:,}원이 받은 금액 총액 {amount:,}원보다 큼 — amount는 비과세 지급액을 포함한 총액")
    if 문제:
        return {"error": "입력 확인 필요 — 계산하지 않음", "문제": 문제}

    식사월한도, 식사근거 = _연도값(_종교인_식사대_월한도, year)
    보육월한도, 보육자녀별, 보육근거 = _연도값(_종교인_보육_월한도, year)
    임차포함, 차량근거 = _연도값(_종교인_자가운전_임차포함, year)

    비과세: list[dict] = []
    산식: list[str] = []

    def _add(항목, 받은, 비과세액, 근거, 사유=""):
        if 받은 > 0:
            비과세.append({"항목": 항목, "받은금액": 받은, "비과세": 비과세액, "과세": 받은 - 비과세액, "근거": 근거, **({"사유": 사유} if 사유 else {})})

    _add("학자금", scholarship, scholarship, "소득세법 제12조5호아목1), 시행령 제19조①")
    _add("현물 식사", meals_in_kind, meals_in_kind, "소득세법 제12조5호아목2), 시행령 제19조②1호")
    if meal_allowance:
        if meal_provided:
            _add("식사대", meal_allowance, 0, 식사근거, "현물 식사를 제공받아 식사대는 과세(시행령 제19조②2호)")
        else:
            nt = min(meal_allowance, 식사월한도 * meal_months)
            _add("식사대", meal_allowance, nt, 식사근거, f"min({meal_allowance:,}, 월 {식사월한도:,} × {meal_months}개월)")
    if vehicle_allowance:
        if vehicle == "본인소유" or (vehicle == "본인명의임차" and 임차포함):
            nt = min(vehicle_allowance, 200_000 * vehicle_months)
            _add("자가운전보조금", vehicle_allowance, nt, 차량근거, f"min({vehicle_allowance:,}, 월 200,000 × {vehicle_months}개월)")
        else:
            사유 = ("본인 명의 임차 차량은 2025 과세기간부터 비과세(시행령 제19조③2호 개정 2025.2.28, 부칙 제4조) — 전액 과세"
                  if vehicle == "본인명의임차" else "본인 소유(2025~ 본인 명의 임차 포함) 차량이 아님 — 전액 과세")
            _add("자가운전보조금", vehicle_allowance, 0, 차량근거, 사유)
    _add("일직·숙직료·실비 여비·재해 지급액", expense_reimbursement, expense_reimbursement, "소득세법 제12조5호아목3), 시행령 제19조③1·2·4호")
    _add("종교활동비", religious_activity_fund, religious_activity_fund, "소득세법 제12조5호아목3), 시행령 제19조③3호·제12조18호")
    if childcare_allowance:
        if childcare_children == 0:
            _add("출산·보육", childcare_allowance, 0, 보육근거, "출산 자녀·개시일 기준 6세 이하 자녀 없음 — 전액 과세")
        else:
            배수 = childcare_children if 보육자녀별 else 1
            nt = min(childcare_allowance, 보육월한도 * 배수 * childcare_months)
            사유 = (f"min({childcare_allowance:,}, 월 {보육월한도:,} × 자녀 {배수}명 × {childcare_months}개월)" if 보육자녀별
                  else f"min({childcare_allowance:,}, 월 {보육월한도:,} × {childcare_months}개월) — {year}년은 자녀 수와 관계없이 1인당 월 한도")
            _add("출산·보육", childcare_allowance, nt, 보육근거, 사유)
    _add("사택 제공 이익", housing_benefit, housing_benefit, "소득세법 제12조5호아목5)")

    비과세_합계 = sum(b["비과세"] for b in 비과세)
    과세대상 = amount - 비과세_합계
    표금액, 표산식 = _종교인_필요경비_표(과세대상)
    필요경비 = max(표금액, actual_expense)
    소득금액 = 과세대상 - 필요경비 if 과세대상 > 필요경비 else 0
    산식.append(f"과세대상 받은 금액 = {amount:,} − 비과세 {비과세_합계:,} = {과세대상:,}원")
    산식.append(f"필요경비 표(시행령 제87조3호): {표산식}원")
    산식.append(f"필요경비 = max(표 {표금액:,}, 실제 {actual_expense:,}) = {필요경비:,}원")
    산식.append(f"기타소득금액 = {과세대상:,} − {필요경비:,} = {소득금액:,}원")

    # ── 계산식 단계 — 비과세 한도 판정(현물 식사 제공·차량 요건·자녀 수)은 판단 결과(월 한도·배수)를 입력으로 ──
    계산식_입력 = [{"이름": "받은금액총액", "값": amount, "설명": "종교단체로부터 받은 금액 총액(비과세 지급액 포함)"}]
    비과세식: list[str] = []
    for 이름, 값, 설명 in (("학자금", scholarship, "시행령 제19조① — 전액 비과세"),
                         ("현물식사", meals_in_kind, "시행령 제19조②1호 — 전액 비과세"),
                         ("실비변상", expense_reimbursement, "일직·숙직료·실비 여비·재해 지급액, 시행령 제19조③1·2·4호 — 전액 비과세"),
                         ("종교활동비", religious_activity_fund, "시행령 제19조③3호 — 전액 비과세"),
                         ("사택제공이익", housing_benefit, "제12조5호아목5) — 전액 비과세")):
        if 값:
            계산식_입력.append({"이름": 이름, "값": 값, "설명": 설명}); 비과세식.append(이름)
    계산식 = []

    def _한도식(이름, 받은, 월한도, 월수, 설명, 배수=None):
        계산식_입력.append({"이름": 이름, "값": 받은, "설명": "받은 금액"})
        계산식_입력.append({"이름": f"{이름}_월한도", "값": 월한도, "설명": 설명})
        계산식_입력.append({"이름": f"{이름}_개월수", "값": 월수, "설명": "받은 개월 수"})
        식 = f"MIN({이름},{이름}_월한도*{이름}_개월수" + (f"*{이름}_자녀수)" if 배수 is not None else ")")
        if 배수 is not None:
            계산식_입력.append({"이름": f"{이름}_자녀수", "값": 배수, "설명": "자녀 1명당 한도면 자녀 수, 아니면 1(자녀 없으면 0)"})
        nt = next(b["비과세"] for b in 비과세 if b["항목"] == {"식사대": "식사대", "자가운전보조금": "자가운전보조금", "출산보육": "출산·보육"}[이름])
        계산식.append({"이름": f"{이름}_비과세", "식": 식, "값": nt, "설명": f"MIN(받은 금액, 월 한도 × 개월 수) — {설명}"})
        비과세식.append(f"{이름}_비과세")

    if meal_allowance:
        _한도식("식사대", meal_allowance, 0 if meal_provided else 식사월한도, meal_months,
               "현물 식사를 제공받아 식사대 과세 — 한도 0(시행령 제19조②2호)" if meal_provided else 식사근거)
    if vehicle_allowance:
        ok = vehicle == "본인소유" or (vehicle == "본인명의임차" and 임차포함)
        _한도식("자가운전보조금", vehicle_allowance, 200_000 if ok else 0, vehicle_months,
               차량근거 if ok else f"차량 요건 미충족({vehicle}) — 한도 0, 전액 과세")
    if childcare_allowance:
        배수 = 0 if childcare_children == 0 else (childcare_children if 보육자녀별 else 1)
        _한도식("출산보육", childcare_allowance, 보육월한도, childcare_months, 보육근거, 배수=배수)
    if not 비과세식:
        계산식_입력.append({"이름": "비과세없음", "값": 0, "설명": "비과세 지급액 없음"}); 비과세식.append("비과세없음")
    계산식_입력.append({"이름": "실제필요경비", "값": actual_expense, "설명": "실제 소요된 필요경비(시행령 제87조3호 단서)"})
    계산식 += [
        {"이름": "비과세_합계", "식": "+".join(비과세식), "값": 비과세_합계, "설명": "제12조5호아목·시행령 제19조"},
        {"이름": "과세대상_받은금액", "식": "받은금액총액-비과세_합계", "값": 과세대상, "설명": "받은 금액 − 비과세"},
        {"이름": "필요경비_표금액",
         "식": "ROUNDDOWN(IF(과세대상_받은금액<=20000000,과세대상_받은금액*0.8,IF(과세대상_받은금액<=40000000,16000000+(과세대상_받은금액-20000000)*0.5,"
               "IF(과세대상_받은금액<=60000000,26000000+(과세대상_받은금액-40000000)*0.3,32000000+(과세대상_받은금액-60000000)*0.2))),0)",
         "값": 표금액, "설명": "시행령 제87조3호 표(2천만 이하 80%, 4천만 이하 1,600만+50%, 6천만 이하 2,600만+30%, 초과 3,200만+20%)"},
        {"이름": "필요경비", "식": "MAX(필요경비_표금액,실제필요경비)", "값": 필요경비, "설명": "표 금액과 실제 필요경비 중 큰 금액"},
        {"이름": "기타소득금액", "식": "MAX(과세대상_받은금액-필요경비,0)", "값": 소득금액, "설명": "과세대상 받은 금액 − 필요경비"},
    ]

    if 소득금액 == 0:
        선택 = False; 선택판단 = "기타소득금액 0 — 분리과세 판단 불필요"
    elif 소득금액 > 3_000_000:
        선택 = False; 선택판단 = f"기타소득금액 {소득금액:,}원 > 300만원 — 종합과세(제14조③8호가목)"
    elif withheld is False:
        선택 = False; 선택판단 = "원천징수되지 않은 종교인소득 — 선택적 분리과세 대상 아님, 종합소득 확정신고 필요(제14조③8호가목·제155조의6 후단)"
    elif withheld is None:
        선택 = None; 선택판단 = "기타소득금액 300만원 이하 — 원천징수됐다면 선택적 분리과세 가능(withheld로 확정). 다른 기타소득이 있으면 합계로 판단"
    else:
        선택 = True; 선택판단 = "기타소득금액 300만원 이하·원천징수된 소득 — 종합과세와 분리과세(원천징수로 종결) 중 선택 가능. 다른 기타소득이 있으면 합계로 판단"

    return _ok({
        "비과세_합계": 비과세_합계,
        "과세대상_받은금액": 과세대상,
        "필요경비_표금액": 표금액,
        "필요경비": 필요경비,
        "기타소득금액": 소득금액,
        "선택적_분리과세_가능": 선택,
    }, ["소득세법 제21조①26호(종교인소득)", "소득세법 제12조5호아목·시행령 제19조(비과세)",
        "소득세법 제37조②2호·시행령 제87조3호(필요경비)", "소득세법 제14조③8호가목(26호 열거 — 300만원 이하 선택적 분리과세)",
        "부칙 <제13558호,2015.12.15> 제1조(2018.1.1. 시행)"],
        비과세=비과세, 산식=산식, 계산식_입력=계산식_입력, 계산식=계산식,
        판단={"과세구분": "기타소득(종합과세, 조건부 선택적 분리과세)", "선택적_분리과세": 선택판단,
              "근로소득_선택": 근로_안내},
        원천징수_신고={
            "원천징수": "종교단체는 매월 종교인소득 간이세액표(시행령 제202조④·별표 3의4)로 원천징수하고, 다음 연도 2월분 지급 시(또는 소속관계 종료 시) 연말정산한다(소득세법 제145조의3, 시행령 제202조의4).",
            "원천징수_예외": "종교단체는 원천징수를 하지 않을 수 있고, 그 경우 받은 사람이 종합소득 확정신고를 해야 한다(소득세법 제155조의6).",
            "확정신고_예외": "원천징수(연말정산)된 종교인소득만 있으면 확정신고를 하지 않을 수 있다(소득세법 제73조①4호의2).",
            "이 도구": "간이세액표 월 세액·연말정산 세액은 계산하지 않음 — 종합소득 합산은 income_tax에 기타소득금액을 더해 계산",
        },
        연도표={"식사대_월한도": 식사월한도, "출산·보육_월한도": 보육월한도, "출산·보육_자녀1명당": 보육자녀별,
               "자가운전보조금_본인명의임차_포함": 임차포함, "근거": [식사근거, 보육근거, 차량근거,
               "시행령 제87조3호 필요경비 표 2018~2026 동일(2천만 이하 80%, 4천만 이하 1,600만+50%, 6천만 이하 2,600만+30%, 초과 3,200만+20%)"]})


# ── 끝에서 끝까지: 간이과세자 부가가치세 확정신고 ──────────────────
# 어휘 항목 5개(간이과세자_납부세액·세금계산서등수취세액공제·신용카드매출전표등발급세액공제·
# 간이과세자_납부의무면제·간이과세자_세금계산서미수취가산세)는 Claude가 검증해 넣었음 — 수정 금지.
# 업종 목록은 vocab.rule("간이과세자_납부세액", year)["부가가치율"]에서 읽는다(코드 상수 금지).



@mcp.tool(annotations=CALC)
@_guard
def simplified_vat_return(year: YEAR,
                         business_type: Annotated[str, Field(description="업종 — 어휘 간이과세자_납부세액.적용[].판단.부가가치율 키(소매업·음식점업·제조업·숙박업·건설업·부동산임대업·그 밖의 서비스업 등)")],
                         prior_year_supplies: Annotated[int | None, Field(description="직전연도 공급대가 합계(원) — 세금계산서 미수취 가산세 판정에만 씀. 미수취 매입이 없으면 생략 가능")] = None,
                         bad_debt_confirmed: Annotated[int, Field(description="대손 확정액(기본 0, 납부세액에 영향 없음 — 판단에 명시)")] = 0,
                         prepaid_assessed: Annotated[int, Field(description="예정부과 고지세액")] = 0,
                         ledger: Annotated[dict, Field(description="원장형 입력 {공급: [{항목, 금액}], 매입: [{항목, 금액}]} — 금액은 공급대가(부가세 포함). 공급 항목이 없으면 error. 항목명은 어휘 간이과세자_원장항목의 키워드로 역할을 정함")] = ...) -> dict:
    """간이과세자 부가가치세 확정신고 차가감납부세액 계산(부가가치세법 제63조·제68조의2·제69조·제46조).
    언제: 간이과세자의 과세기간(1.1~12.31) 부가가치세 확정신고 시 차가감납부세액을 원래 사실관계에서 계산한다.
    입력:
      · 원장: ledger={공급:[{항목, 금액}], 매입:[{항목, 금액}]}.
        항목명 → 역할 매핑은 어휘 간이과세자_원장항목.판단.공급/매입에서 읽음(항목명에 키워드가 포함되면 해당 역할).
        공급 금액 전부 합(total), 공제 대상 합(카드발급), 마일리지 합(자기마일리지)으로 변환.
        매입도 같은 식으로 변환.
      대손 확정액(기본 0) → bad_debt_confirmed, 직전연도 합계 → prior_year_supplies,
      예정부과 고지세액 → prepaid_assessed.
    결과: 판단(항목·결론·근거)·계산(과세표준·부가가치율·납부세액·공제·가산세)·결과(납부의무면제·차가감납부세액)·산식·근거."""

    # 2021-07-01 이전 체계 미지원
    if year < 2021:
        return {"error": "2021-07-01 이전 체계 미지원"}

    # ── 빈 원장 거부 ──
    if not ledger.get("공급"):
        return {"error": "입력 확인 필요 — 원장 항목이 없음",
                "허용값": {"공급": {"필수": "공급 항목 목록 [{항목, 금액}]"},
                           "매입": {"선택": "매입 항목 목록 [{항목, 금액}]"}}}

    # ── 원장형 입력 처리 ──
    원장항목_규칙 = vocab.rule("간이과세자_원장항목", year)
    공급_역할목록 = 원장항목_규칙.get("공급", {})
    매입_역할목록 = 원장항목_규칙.get("매입", {})

    총공급대가 = 0
    카드발급 = 0
    자기마일리지 = 0
    세금계산서수취 = 0
    계산서수취 = 0
    미수취매입 = 0
    면세농산물 = 0
    문제_항목 = []

    for 공급항목 in ledger.get("공급", []):
        항목명 = 공급항목.get("항목", "")
        금액 = int(공급항목.get("금액", 0))
        총공급대가 += 금액
        역할_찾음 = False
        for 역할명, 역할정보 in 공급_역할목록.items():
            for kw in 역할정보.get("키워드", []):
                if kw in 항목명:
                    if 역할명 == "발급세액공제대상":
                        카드발급 += 금액
                    elif 역할명 == "과세표준제외":
                        자기마일리지 += 금액
                    역할_찾음 = True
                    break
            if 역할_찾음:
                break
        if not 역할_찾음:
            문제_항목.append(항목명)

    for 매입항목 in ledger.get("매입", []):
        항목명 = 매입항목.get("항목", "")
        금액 = int(매입항목.get("금액", 0))
        역할_찾음 = False
        for 역할명, 역할정보 in 매입_역할목록.items():
            for kw in 역할정보.get("키워드", []):
                if kw in 항목명:
                    if 역할명 == "세금계산서등수취":
                        세금계산서수취 += 금액
                    elif 역할명 == "세금계산서미수취":
                        미수취매입 += 금액
                    elif 역할명 == "계산서수취":
                        계산서수취 += 금액
                    elif 역할명 == "면세농산물":
                        면세농산물 += 금액
                    역할_찾음 = True
                    break
            if 역할_찾음:
                break
        if not 역할_찾음:
            문제_항목.append(항목명)

    if 문제_항목:
        허용값 = {역할명: {"키워드": 역할정보.get("키워드", [])} for 역할명, 역할정보 in {**공급_역할목록, **매입_역할목록}.items()}
        return {"error": "입력 확인 필요 — 계산하지 않음", "문제": [f"어휘에 없는 항목명: {', '.join(문제_항목)}"], "허용값": 허용값}

    # 업종 유효성 검사 — vocab.rule에서 연도별로 부가가치율 표 읽기
    vat_rule = vocab.rule("간이과세자_납부세액", year)
    부가가치율표 = vat_rule["부가가치율"]
    if not 부가가치율표:
        return {"error": "입력 확인 필요 — 계산하지 않음",
                "문제": [f"year={year}의 부가가치율 표를 vocab에서 읽을 수 없음"],
                "허용값": {}, "안내": "연도표 없는 해는 추정하지 않음"}
    if business_type not in 부가가치율표:
        return {"error": "입력 확인 필요 — 계산하지 않음",
                "문제": [f"business_type='{business_type}' 은(는) 지원하지 않음"],
                "허용값": {k: {"설명": f"부가가치율 {v}"} for k, v in 부가가치율표.items()},
                "안내": "허용 업종 목록에서 선택하세요."}
    대손확정 = int(bad_debt_confirmed)
    if prior_year_supplies is None:
        if 미수취매입 > 0:
            return {"error": "입력 확인 필요 — 계산하지 않음",
                    "문제": [f"세금계산서 미수취 매입 {미수취매입:,}원이 있어 가산세 판정에 prior_year_supplies(직전연도 공급대가)가 필요"]}
        prior_year_supplies = 0   # 미수취 매입이 없으면 가산세 0이라 값과 무관
    직전연도 = int(prior_year_supplies)
    예정고지 = int(prepaid_assessed)

    부가가치율 = 부가가치율표[business_type]

    # 1. 과세표준 = total − self_mileage (영 제61조②9호)
    과세표준 = 총공급대가 - 자기마일리지

    # 2. 납부세액 = 과세표준 × 부가가치율 × 10% (제63조②). 원 미만 절사
    납부세액 = int(과세표준 * 부가가치율 * 0.1)

    # 3. 납부의무 면제: 과세표준(공급대가 합계) < 기준(그해) → 모두 0 (제69조①)
    면제기준_rule = vocab.rule("간이과세자_납부의무면제", year)
    면제기준 = 면제기준_rule["기준"]
    # 기본 판단 기준과 적용 기간 기준 중 연도별로 vocab.basis에서 정함 — rule이 이미 적용 블록 기준으로 덮어쓴 값을 반환
    납부의무면제 = 총공급대가 < 면제기준

    # 공제·가산에 필요한 규칙(면제 여부와 무관하게 연도별로 vocab에서 읽기)
    세금계산서수취공제_rule = vocab.rule("세금계산서등수취세액공제", year)
    세금계산서수취공제율 = 세금계산서수취공제_rule["공제율"]
    카드공제_rule = vocab.rule("신용카드매출전표등발급세액공제", year)
    카드공제율 = 카드공제_rule["공제율"]
    카드공제한도 = 카드공제_rule["한도"]
    미수취_rule = vocab.rule("간이과세자_세금계산서미수취가산세", year)
    미수취가산세율 = 미수취_rule["가산세율"]
    영수증발급기준 = 미수취_rule["영수증발급기준"]

    if 납부의무면제:
        납부세액 = 0
        세금계산서수취공제 = 0
        카드발급공제 = 0
        공제합계 = 0
        공제후세액 = 0
    else:
        # 4. 세금계산서등수취세액공제 = tax_invoice_received × 0.5% (제63조③1호)
        세금계산서수취공제 = int(세금계산서수취 * 세금계산서수취공제율)

        # 5. 신용카드매출전표등발급세액공제 = min(card × 공제율, 한도) (제46조①3호)
        카드발급공제 = int(min(카드발급 * 카드공제율, 카드공제한도))

        # 6. 공제 합계가 납부세액을 넘으면 넘는 부분은 없음 (제63조⑥)
        공제합계 = min(세금계산서수취공제 + 카드발급공제, 납부세액)
        공제후세액 = max(0, 납부세액 - 공제합계)

    # 7. 세금계산서미수취가산세 = no_tax_invoice_from_issuer × 0.5% — 납부의무 면제와 무관하게 적용
    #   (제69조①은 제63조② 납부의무만 면제, 제69조②는 제60조① 미등록 가산세만 배제. CodeSolar 지적 후 Claude 확인 10-13)
    #   다만 prior_year_supplies < 영수증발급기준(4,800만)이면 0 (제68조의2②1호, 제36조①2호가목)
    if 직전연도 < 영수증발급기준:
        미수취가산세 = 0
    else:
        미수취가산세 = int(미수취매입 * 미수취가산세율)

    # 8. 의제매입세액공제 0 (2021-07-01부터 적용 없음), 대손세액공제 적용 없음 — 판단에 명시
    의제매입세액공제 = 0

    # 9. 차가감납부세액 = 공제후세액 + 가산세 − 예정고지 (음수면 0, 환급 없음)
    차가감납부세액 = max(0, 공제후세액 + 미수취가산세 - 예정고지)

    # 판단 목록
    판단 = [
        {"항목": "과세표준", "결론": f"{과세표준:,}원 (공급대가 합계 {총공급대가:,}원 − 자기적립마일리지 {자기마일리지:,}원)",
         "근거": "부가가치세법 제63조①·⑤, 시행령 제61조②9호"},
        {"항목": "부가가치율", "결론": f"{부가가치율*100:.0f}% ({business_type})",
         "근거": vocab.basis("간이과세자_납부세액", year)},
        {"항목": "납부세액", "결론": f"{납부세액:,}원 (과세표준 × 부가가치율 × 10%)",
         "근거": "부가가치세법 제63조②"},
        {"항목": "세금계산서등수취세액공제", "결론": f"{세금계산서수취공제:,}원 (수취 공급대가 {세금계산서수취:,}원 × {세금계산서수취공제율*100:.1f}%)",
         "근거": "부가가치세법 제63조③1호"},
        {"항목": "신용카드매출전표등발급세액공제", "결론": f"{카드발급공제:,}원 (발급분 {카드발급:,}원 × {카드공제율*100:.1f}%, 한도 {카드공제한도:,}원)",
         "근거": "부가가치세법 제46조①3호"},
        {"항목": "의제매입세액공제", "결론": "0원 (2021-07-01부터 간이과세자 적용 없음, 제65조 삭제)",
         "근거": "부가가치세법 제65조 삭제(2021.2.17. 개정)"},
        {"항목": "대손세액공제", "결론": "적용 없음 (제45조 대손세액공제는 일반과세자 납부세액 공제, 간이과세자 미적용)",
         "근거": "부가가치세법 제45조① — 간이과세자 납부세액 계산 범위 밖"},
        {"항목": "공제합계", "결론": f"{공제합계:,}원 (납부세액 {납부세액:,}원 한도)",
         "근거": "부가가치세법 제63조⑥"},
        {"항목": "공제후세액", "결론": f"{공제후세액:,}원",
         "근거": "부가가치세법 제63조⑥"},
        {"항목": "세금계산서미수취가산세", "결론": (
            f"{미수취가산세:,}원 (미수취 매입 {미수취매입:,}원 × {미수취가산세율*100:.1f}%)" if 미수취가산세 > 0 else
            (f"0원 (직전연도 공급대가 {직전연도:,}원 < 영수증발급기준 {영수증발급기준:,}원 → 영수증 발급 기간)" if 직전연도 < 영수증발급기준 else
             f"0원 (직전연도 공급대가 {직전연도:,}원 ≥ {영수증발급기준:,}원 — 세금계산서 발급 대상 간이과세자, 미수취 매입 {미수취매입:,}원)"),
            ),
         "근거": "부가가치세법 제68조의2②1호, 제36조①2호가목"},
        {"항목": "차가감납부세액", "결론": f"{차가감납부세액:,}원 (공제후세액 {공제후세액:,}원 + 가산세 {미수취가산세:,}원 − 예정고지 {예정고지:,}원)",
         "근거": "부가가치세법 제63조⑥, 제66조(예정부과세액 차감)"},
    ]
    if 대손확정 > 0:
        판단.append({"항목": "대손 확정액", "결론": f"{대손확정:,}원 — 간이과세자 납부세액에 영향 없음(제45조 일반과세자 대상)",
                     "근거": "부가가치세법 제45조① — 간이과세자 과세표준·납부세액 계산 범위 밖"})

    산식 = [
        f"과세표준 = total − self_mileage = {총공급대가:,} − {자기마일리지:,} = {과세표준:,}원 (영 제61조②9호)",
        f"납부세액 = 과세표준 × 부가가치율 × 10% = {과세표준:,} × {부가가치율} × 0.1 = {납부세액:,}원 (제63조②)",
        f"세금계산서등수취세액공제 = tax_invoice_received × 0.5% = {세금계산서수취:,} × 0.005 = {세금계산서수취공제:,}원 (제63조③1호)",
        f"신용카드매출전표등발급세액공제 = min(card × {카드공제율}, {카드공제한도:,}) = min({카드발급:,} × {카드공제율}, {카드공제한도:,}) = {카드발급공제:,}원 (제46조①3호)",
        f"공제합계 = min(매입공제 + 카드공제, 납부세액) = min({세금계산서수취공제:,} + {카드발급공제:,}, {납부세액:,}) = {공제합계:,}원 (제63조⑥)",
        f"공제후세액 = max(0, 납부세액 − 공제합계) = max(0, {납부세액:,} − {공제합계:,}) = {공제후세액:,}원",
        f"세금계산서미수취가산세 = {('0원 (영수증 발급 기간)' if 직전연도 < 영수증발급기준 else '0원 (미수취 매입 없음)') if 미수취가산세 == 0 else f'no_tax_invoice × 0.5% = {미수취매입:,} × 0.005 = {미수취가산세:,}원 (제68조의2②1호)'}",
        f"차가감납부세액 = max(0, 공제후세액 + 가산세 − 예정고지) = max(0, {공제후세액:,} + {미수취가산세:,} − {예정고지:,}) = {차가감납부세액:,}원",
    ]

    return {
        "판단": 판단,
        "계산": {
            "과세표준": 과세표준,
            "부가가치율": 부가가치율,
            "납부세액": 납부세액,
            "세금계산서등수취세액공제": 세금계산서수취공제,
            "신용카드매출전표등발급세액공제": 카드발급공제,
            "의제매입세액공제": 의제매입세액공제,
            "공제합계": 공제합계,
            "공제후세액": 공제후세액,
            "세금계산서미수취가산세": 미수취가산세,
        },
        "결과": {
            "납부의무면제": 납부의무면제,
            "차가감납부세액": 차가감납부세액,
        },
        "산식": 산식,
        "근거": vocab.basis("간이과세자_납부세액", year) + vocab.basis("세금계산서등수취세액공제", year) +
                vocab.basis("신용카드매출전표등발급세액공제", year) + vocab.basis("간이과세자_납부의무면제", year) +
                vocab.basis("간이과세자_세금계산서미수취가산세", year),
    }


# ── 끝에서 끝까지: 겸영사업자 공통매입세액 안분·정산 + 공통사용재화 공급 ──
# 어휘 항목 4개(공통매입세액_안분·공통사용재화_공급·공통매입세액_재계산·토지건물_일괄공급_안분)는 Claude 검증본 — 수정 금지.
# 기준값·판단 규칙은 모두 vocab.rule/vocab.basis에서 읽는다(코드 상수 금지).



@mcp.tool(annotations=CALC)
@_guard
def common_input_tax(year: YEAR,
                     period: Annotated[str, Field(description="과세기간('1기'|'2기')")],
                     filing: Annotated[str, Field(description="신고 구분('예정'|'확정')")],
                     supplies: Annotated[dict, Field(description="공급가액 {예정:{taxable,exempt}, 확정:{taxable,exempt}} — 부가세 제외 금액, 확정신고는 전체 비율 사용. 둘 다 null이면 error")],
                     new_business: Annotated[bool, Field(description="신규 사업 여부(기본 False)")] = False,
                     prior_nondeductible: Annotated[int | None, Field(description="예정신고 때 불공제한 공통매입세액. 없으면(null) 예정 기간 비율로 계산")] = None,
                     prior_period_supplies: Annotated[dict | None, Field(description="직전 과세기간 {taxable, exempt} — 공통사용재화 공급 과세표준용")] = None,
                     common_asset_sales: Annotated[list | None, Field(description="공통사용재화 공급 목록 [{sale_price(공급가액), input_tax(그 재화의 매입세액 — ledger에 넣지 말고 여기만 입력, 신규 사업 아니면 직전 과세기간 비율로 안분 공제)}]")] = None,
                     ledger: Annotated[list, Field(description="원장형 입력 당기 매입 항목 [{항목, 매입세액 | 매입가액(×10%, 둘 중 하나), 사용처:'과세'|'면세'|'공통'|'실지귀속'|'불공제', 과세비율(0~1) | 횟수:{과세,면세} | 면적:{과세,면세} | 신축건물(bool, 기본 False) — 둘 중 하나: 매입세액 또는 매입가액. 공통+신축건물+면적 있으면 예정사용면적비율 우선 적용")] = ...,
                     prior_common_input_tax: Annotated[int | None, Field(description="ledger로 확정신고할 때 예정신고 기간의 공통매입세액")] = None) -> dict:
    """겸영사업자 공통매입세액 안분·정산 + 공통사용재화 공급(부가법 제40조, 영 제81조~제83조·제63조).
    언제: 과세·면세 겸영 사업자의 예정·확정신고에서 공통매입세액 공제액과 공통사용재화 공급 시 추가 매출세액을 계산한다.
    입력:
      · 공급가액(예정·확정별 과세·면세), 예정 불공제액, 직전 과세기간 공급가액, 공통사용재화 공급 목록.
      · 원장: ledger=[{항목, 매입세액 | 매입가액(→×10%, 둘 중 하나), 사용처:'과세'|'면세'|'공통'|'실지귀속'|'불공제',
        과세비율(0~1) | 횟수:{과세,면세} | 면적:{과세,면세}}].
        사용처별 처리는 어휘 공통매입세액_귀속.판단.사용처별에서 읽음.
        prior_common_input_tax: 확정신고에서 ledger 사용 시 예정신고 기간의 공통매입세액. 기존 prior_nondeductible는 그대로.
    결과: 판단·계산·결과(공제매입세액·추가_매출세액)·산식·근거. 계산의 실지귀속_공제는 기존 계산 키 유지."""

    # 기간 유효성 검사
    if period not in ("1기", "2기"):
        return {"error": f"period='{period}' 은(는) 지원하지 않음 — '1기' 또는 '2기'여야 함"}

    # 빈 원장 거부
    if not ledger:
        return {"error": "입력 확인 필요 — 원장 항목이 없음",
                "허용값": {"사용처": ["과세", "면세", "공통", "실지귀속", "불공제"]}}

    # 공급가액 검증
    예정 = supplies.get("예정")
    확정 = supplies.get("확정")
    if (예정 is None or 예정.get("taxable") is None or 예정.get("exempt") is None) and \
       (확정 is None or 확정.get("taxable") is None or 확정.get("exempt") is None):
        return {"error": "공급가액 입력 필요 — 예정·확정 중 하나는 과세·면세 공급가액을 모두 입력해야 함"}

    # 어휘에서 기준값 읽기
    안분_규칙 = vocab.rule("공통매입세액_안분", year)
    공급_규칙 = vocab.rule("공통사용재화_공급", year)

    면세비율미만_기준 = 안분_규칙["전액공제_면세비율미만"]
    공통매입세액_상한 = 안분_규칙["전액공제_면세비율미만_공통매입세액상한"]
    공통매입세액_전액기준 = 안분_규칙["전액공제_공통매입세액미만"]

    직전면세비율미만_기준 = 공급_규칙["전액과세_직전면세비율미만"]
    직전면세비율미만_상한 = 공급_규칙["전액과세_직전면세비율미만_공급가액상한"]
    공급가액_전액기준 = 공급_규칙["전액과세_공급가액미만"]

    # 공급가액 파싱
    def _get_supplies(s):
        if s is None:
            return 0, 0
        return int(s.get("taxable", 0)), int(s.get("exempt", 0))

    예정_과세, 예정_면세 = _get_supplies(예정)
    확정_과세, 확정_면세 = _get_supplies(확정)

    # ── 원장형 입력 처리 ──
    taxable_only = 0
    taxable_only_nondeductible = 0
    exempt_only = 0
    attributed = []
    common_예정 = 0
    common_확정 = 0

    판단_목록 = []
    산식_목록 = []

    _ledger_허용키 = {"항목", "매입세액", "매입가액", "사용처", "과세비율", "횟수", "면적", "신축건물"}

    for item in ledger:
        # ledger 키 끝의 물음표 제거 (모델이 매입세액 물음표를 그대로 입력한 경우 대응)
        item = {k[:-1] if k.endswith("?") else k: v for k, v in item.items()}

        항목명 = item.get("항목", "")
        매입세액 = item.get("매입세액", None)
        매입가액 = item.get("매입가액", None)
        사용처 = item.get("사용처", "")
        과세비율 = item.get("과세비율", None)
        횟수 = item.get("횟수", None)
        면적 = item.get("면적", None)

        # 모르는 키 검사
        모르는키 = sorted(set(item.keys()) - _ledger_허용키)
        if 모르는키:
            return {"error": f"입력 확인 필요 — 계산하지 않음",
                    "문제": [f"항목 '{항목명}': 모르는 키 {모르는키} — 허용 키: {_ledger_허용키}"],
                    "모르는 키": 모르는키, "허용 키": sorted(_ledger_허용키)}

        if 매입세액 is not None:
            금액 = int(매입세액)
        elif 매입가액 is not None:
            금액 = int(매입가액) * 10 // 100
        else:
            return {"error": f"입력 확인 필요 — 항목 '{항목명}': 매입세액 또는 매입가액 필요"}

        if 사용처 == "과세":
            taxable_only += 금액
        elif 사용처 == "면세":
            exempt_only += 금액
        elif 사용처 == "불공제":
            taxable_only += 금액                      # 과세사업분이지만 제39조로 불공제 — 기존 필드 의미(taxable_only 중 불공제분)에 맞춰 둘 다에 더해 순효과 0
            taxable_only_nondeductible += 금액
        elif 사용처 == "공통":
            신축 = item.get("신축건물", False)
            if 신축 and 면적 is not None:
                과세 = 면적.get("과세", 0)
                총면적 = 과세 + 면적.get("면세", 0)
                if 총면적 == 0:
                    return {"error": f"입력 확인 필요 — 항목 '{항목명}': 신축건물인데 면적 합계 0"}
                비율 = 과세 / 총면적
                attributed.append({"input_tax": 금액, "taxable_ratio": 비율})
                판단_목록.append({"항목": f"{항목명}", "결론": f"건물 신축·취득 겸용 예정 → 예정사용면적비율 {비율:.4f} 우선 적용 (영 제81조④ 단서)", "근거": "부가가치세법 시행령 제81조④"})
                산식_목록.append(f"{항목명}: 건물 신축·취득 겸용 예정 → 예정사용면적비율 {비율:.4f} = {과세}/{총면적} 우선 적용, 공제 = {금액:,} × {비율:.4f} = {int(금액 * 비율):,}원 (영 제81조④ 단서)")
            elif 신축 and 면적 is None:
                return {"error": f"입력 확인 필요 — 항목 '{항목명}': 신축건물인데 면적 없음"}
            else:
                if filing == "확정":
                    common_확정 += 금액
                else:
                    common_예정 += 금액
        elif 사용처 == "실지귀속":
            if 과세비율 is not None:
                비율 = float(과세비율)
            elif 횟수 is not None:
                과세 = 횟수.get("과세", 0)
                면세 = 횟수.get("면세", 0)
                합계 = 과세 + 면세
                if 합계 == 0:
                    return {"error": f"입력 확인 필요 — 항목 '{항목명}': 실지귀속 횟수 합계 0"}
                비율 = 과세 / 합계
            elif 면적 is not None:
                과세 = 면적.get("과세", 0)
                총면적 = 과세 + 면적.get("면세", 0)
                if 총면적 == 0:
                    return {"error": f"입력 확인 필요 — 항목 '{항목명}': 실지귀속 면적 합계 0"}
                비율 = 과세 / 총면적
            else:
                return {"error": f"입력 확인 필요 — 항목 '{항목명}': 실지귀속인데 과세비율·횟수·면적 없음"}

            attributed.append({"input_tax": 금액, "taxable_ratio": 비율})
        else:
            return {"error": f"입력 확인 필요 — 항목 '{항목명}': 사용처='{사용처}' 은(는) 지원하지 않음 — '과세'|'면세'|'공통'|'실지귀속'|'불공제'"}

    if prior_common_input_tax is not None:
        common_예정 = int(prior_common_input_tax)

    # ── 1. 면세비율 계산 ──
    # 예정신고: 예정 기간 비율
    if filing == "예정":
        총예정 = 예정_과세 + 예정_면세
        if 총예정 == 0:
            return {"error": "예정신고 기간 공급가액이 없음 — 안분 불가"}
        면세비율_예정 = 예정_면세 / 총예정
        면세비율_과세기간 = 면세비율_예정  # 예정신고는 과세기간 개념 없음
    else:  # 확정
        총예정 = 예정_과세 + 예정_면세
        총확정 = 확정_과세 + 확정_면세
        총과세기간 = 총예정 + 총확정
        if 총과세기간 == 0:
            return {"error": "확정신고 과세기간 전체 공급가액이 없음 — 안분 불가"}
        면세비율_과세기간 = (예정_면세 + 확정_면세) / 총과세기간
        if 총예정 > 0:
            면세비율_예정 = 예정_면세 / 총예정
        else:
            면세비율_예정 = 0.0  # 예정 공급 없음

    # ── 2. 공통매입세액 계산 ──

    if filing == "예정":
        # 예정신고: 공제 = common.예정 × (1 − 예정 면세비율)
        공통매입세액_합계 = common_예정
        불공제_예정 = int(common_예정 * 면세비율_예정)
        공통매입세액_공제 = common_예정 - 불공제_예정
        판단_목록.append({"항목": "공통매입세액 예정 공제", "결론": f"안분 공제 {공통매입세액_공제:,}원 (불공제 {불공제_예정:,}원 = {common_예정:,} × {면세비율_예정:.4f})",
                          "근거": "부가가치세법 시행령 제81조①"})
        산식_목록.append(f"불공제 = 공통매입세액 × 면세비율 = {common_예정:,} × {면세비율_예정:.4f} = {불공제_예정:,}원 (영 제81조①)")
        산식_목록.append(f"공제 = {common_예정:,} − {불공제_예정:,} = {공통매입세액_공제:,}원")
    else:  # 확정
        공통매입세액_합계 = common_예정 + common_확정
        불공제_과세기간합계 = 0
        불공제_예정반영 = 0
        불공제_확정 = 0

        # 전액 공제 예외 판단 (영 제81조②) — new_business는 공통사용재화 매입세액에 적용, 공통매입세액은 정산 대상
        if 공통매입세액_합계 < 공통매입세액_전액기준:
            공통매입세액_공제 = 공통매입세액_합계
            판단_목록.append({"항목": "공통매입세액 확정 공제", "결론": f"전액 공제(공통매입세액 합계 {공통매입세액_합계:,}원 < {공통매입세액_전액기준:,}원, 영 제81조②2호)",
                              "근거": "부가가치세법 시행령 제81조②2호"})
            산식_목록.append(f"공통매입세액 합계 {공통매입세액_합계:,}원 < {공통매입세액_전액기준:,}원 → 전액 공제 (영 제81조②2호)")
        elif 면세비율_과세기간 < 면세비율미만_기준 and 공통매입세액_합계 < 공통매입세액_상한:
            공통매입세액_공제 = 공통매입세액_합계
            판단_목록.append({"항목": "공통매입세액 확정 공제", "결론": f"전액 공제(면세비율 {면세비율_과세기간:.1%} < 5%, 공통매입세액 합계 {공통매입세액_합계:,}원 < 500만원, 영 제81조②1호)",
                              "근거": "부가가치세법 시행령 제81조②1호"})
            산식_목록.append(f"면세비율 {면세비율_과세기간:.1%} < 5% & 공통매입세액 합계 {공통매입세액_합계:,}원 < 500만원 → 전액 공제 (영 제81조②1호)")
        else:
            # 정산 계산 (영 제82조)
            불공제_과세기간합계 = int(공통매입세액_합계 * 면세비율_과세기간)

            # 불공제_예정반영: prior_nondeductible이 있으면 그대로, 없으면 예정 기간 비율로 계산
            if prior_nondeductible is not None:
                불공제_예정반영 = prior_nondeductible
            else:
                # 예정 기간 비율로 계산 = common.예정 × 예정 면세비율
                불공제_예정반영 = int(common_예정 * 면세비율_예정) if common_예정 > 0 else 0

            불공제_확정 = 불공제_과세기간합계 - 불공제_예정반영
            공통매입세액_공제 = common_확정 - 불공제_확정

            판단_목록.append({"항목": "공통매입세액 확정 정산", "결론": (
                f"과세기간 합계 {공통매입세액_합계:,}원, 면세비율 {면세비율_과세기간:.4f}, "
                f"불공제 총액 {불공제_과세기간합계:,}원, 기불공제 {불공제_예정반영:,}원, "
                f"확정 불공제 {불공제_확정:,}원, 확정 공제 {공통매입세액_공제:,}원"
                ),
                "근거": "부가가치세법 시행령 제81조①·제82조"})

            산식_목록.append(f"공통매입세액 과세기간 합계 = {common_예정:,} + {common_확정:,} = {공통매입세액_합계:,}원")
            산식_목록.append(f"불공제 총액 = {공통매입세액_합계:,} × {면세비율_과세기간:.4f} = {불공제_과세기간합계:,}원 (영 제81조①)")
            산식_목록.append(f"기불공제 = {불공제_예정반영:,}원 (예정신고 때 불공제)")
            산식_목록.append(f"확정 불공제 = {불공제_과세기간합계:,} − {불공제_예정반영:,} = {불공제_확정:,}원 (영 제82조)")
            산식_목록.append(f"공통매입세액 공제 = {common_확정:,} − {불공제_확정:,} = {공통매입세액_공제:,}원")

    # ── 3. 실지귀속 공제 ──
    실지귀속_공제 = 0
    for attr in attributed:
        attr_tax = int(attr.get("input_tax", 0))
        attr_ratio = float(attr.get("taxable_ratio", 1.0))
        실지귀속_공제 += int(attr_tax * attr_ratio)
    if attributed:
        산식_목록.append(f"실지귀속 공제 = Σ(input_tax × taxable_ratio) = {실지귀속_공제:,}원")

    # ── 4. 공통사용재화 공급 계산 ──
    공통사용재화_공급_과세표준 = 0
    공통사용재화_매입세액_공제 = 0
    추가_매출세액 = 0

    if common_asset_sales:
        직전_과세 = int((prior_period_supplies or {}).get("taxable", 0))
        직전_면세 = int((prior_period_supplies or {}).get("exempt", 0))
        직전_총계 = 직전_과세 + 직전_면세
        직전_과세비율 = 직전_과세 / 직전_총계 if 직전_총계 > 0 else 0.0
        직전_면세비율 = 직전_면세 / 직전_총계 if 직전_총계 > 0 else 0.0

        for sale in common_asset_sales:
            sale_price = int(sale.get("sale_price", 0))
            sale_input_tax = int(sale.get("input_tax", 0))

            # 전액 과세표준 여부 판단 (영 제63조③)
            if new_business:
                # ③3호: 신규 사업은 전액 과세표준
                과세표준 = sale_price
                판단_목록.append({"항목": f"공통사용재화 공급(공급가액 {sale_price:,}원)", "결론": f"전액 과세표준(신규 사업, 영 제63조③3호)",
                                  "근거": "부가가치세법 시행령 제63조③3호"})
                산식_목록.append(f"신규 사업 → 공급가액 {sale_price:,}원 전액 과세표준 (영 제63조③3호)")
            elif sale_price < 공급가액_전액기준:
                # ③2호: 공급가액 50만 원 미만 전액
                과세표준 = sale_price
                판단_목록.append({"항목": f"공통사용재화 공급(공급가액 {sale_price:,}원)", "결론": f"전액 과세표준(50만원 미만, 영 제63조③2호)",
                                  "근거": "부가가치세법 시행령 제63조③2호"})
                산식_목록.append(f"공급가액 {sale_price:,}원 < 50만원 → 전액 과세표준 (영 제63조③2호)")
            elif 직전_총계 > 0 and 직전_면세비율 < 직전면세비율미만_기준 and sale_price < 직전면세비율미만_상한:
                # ③1호: 직전 면세비율 5% 미만이고 5천만 원 미만
                과세표준 = sale_price
                판단_목록.append({"항목": f"공통사용재화 공급(공급가액 {sale_price:,}원)", "결론": f"전액 과세표준(직전 면세비율 {직전_면세비율:.1%} < 5%, 5천만원 미만, 영 제63조③1호)",
                                  "근거": "부가가치세법 시행령 제63조③1호"})
                산식_목록.append(f"직전 면세비율 {직전_면세비율:.1%} < 5% & 공급가액 {sale_price:,}원 < 5천만원 → 전액 과세표준 (영 제63조③1호)")
            elif prior_period_supplies is not None and 직전_총계 > 0:
                # 직전 과세기간 과세비율로 안분
                과세표준 = int(sale_price * 직전_과세비율)
                판단_목록.append({"항목": f"공통사용재화 공급(공급가액 {sale_price:,}원)", "결론": f"과세표준 {과세표준:,}원 = {sale_price:,} × 직전 과세비율 {직전_과세비율:.4f}",
                                  "근거": "부가가치세법 시행령 제63조①"})
                산식_목록.append(f"과세표준 = {sale_price:,} × 직전 과세비율 {직전_과세비율:.4f} = {과세표준:,}원 (영 제63조①)")
            else:
                return {"error": f"공통사용재화 공급가액 {sale_price:,}원 — 신규 사업이 아니고 직전 과세기간 자료가 없음 → 안분 불가"}

            공통사용재화_공급_과세표준 += 과세표준
            추가_매출세액 += int(과세표준 * 0.1)

            # 그 재화의 매입세액 공제 판단
            if new_business:
                # ③3호 적용 재화: 매입세액 전액 공제 (영 제81조②3호)
                공통사용재화_매입세액_공제 += sale_input_tax
                판단_목록.append({"항목": f"공통사용재화 매입세액(입력세액 {sale_input_tax:,}원)", "결론": f"전액 공제(신규 사업 재화, 영 제81조②3호)",
                                  "근거": "부가가치세법 시행령 제81조②3호"})
                산식_목록.append(f"신규 사업 재화 매입세액 {sale_input_tax:,}원 전액 공제 (영 제81조②3호)")
            else:
                # 신규 사업이 아니면 그 재화의 매입세액은 ledger(공통)에 넣지 말고
                # common_asset_sales.input_tax에만 입력 — 직전 과세기간 과세비율로 안분 공제 (영 제81조④)
                if prior_period_supplies is not None and 직전_총계 > 0:
                    공제액 = int(sale_input_tax * 직전_과세비율)
                    공통사용재화_매입세액_공제 += 공제액
                    판단_목록.append({"항목": f"공통사용재화 매입세액(입력세액 {sale_input_tax:,}원)", "결론": f"공통사용재화 매입세액 {sale_input_tax:,}원 × 직전 과세비율 {직전_과세비율:.4f} = {공제액:,}원 공제 (영 제81조④)",
                                      "근거": "부가가치세법 시행령 제81조④"})
                    산식_목록.append(f"공통사용재화 매입세액 {sale_input_tax:,}원 × 직전 과세비율 {직전_과세비율:.4f} = {공제액:,}원 공제 (영 제81조④)")

    # ── 5. 공제매입세액 합계 ──
    공제매입세액 = (taxable_only - taxable_only_nondeductible) + 실지귀속_공제 + 공통매입세액_공제 + 공통사용재화_매입세액_공제

    # 판단 추가
    판단_목록.append({"항목": "공제매입세액 합계", "결론": (
        f"과세전용({taxable_only:,} − 불공제 {taxable_only_nondeductible:,}) + "
        f"실지귀속 {실지귀속_공제:,} + "
        f"공통매입세액 {공통매입세액_공제:,} + "
        f"공통사용재화 {공통사용재화_매입세액_공제:,} = {공제매입세액:,}원"
        ),
        "근거": "부가가치세법 제40조, 시행령 제81조"})

    # ── 계산식 단계 (SPEC_export_workbook_f) ──────────────────────────────
    # 공통사용재화 공급 계산 관련 입력/식
    계산식_입력 = [
        {"이름": "과세공급가액", "값": 확정_과세 + 예정_과세, "설명": f"당기 과세 공급가액 합계 = 확정 {확정_과세:,} + 예정 {예정_과세:,} = {확정_과세 + 예정_과세:,}원"},
        {"이름": "면세공급가액", "값": 확정_면세 + 예정_면세, "설명": f"당기 면세 공급가액 합계 = 확정 {확정_면세:,} + 예정 {예정_면세:,} = {확정_면세 + 예정_면세:,}원"},
        {"이름": "직전과세공급가액", "값": int((prior_period_supplies or {}).get("taxable", 0)), "설명": f"직전 과세기간 과세 공급가액 = {int((prior_period_supplies or {}).get('taxable', 0)):,}원"},
        {"이름": "직전면세공급가액", "값": int((prior_period_supplies or {}).get("exempt", 0)), "설명": f"직전 과세기간 면세 공급가액 = {int((prior_period_supplies or {}).get('exempt', 0)):,}원"},
        {"이름": "공급가액_전액기준", "값": 500000, "설명": "공급가액 50만 원 미만 전액 과세표준 기준 (영 제63조③2호)"},
    ]

    # 공통사용재화 공급 과세표준 계산 (공통사용재화_공급_과세표준을 구하는 식)
    계산식 = []

    if common_asset_sales:
        # 직전 과세기간 과세비율은 입력이 아니라 단계에서 계산한다 (SPEC h).
        # 여러 공통사용재화가 같은 직전 과세기간 비율을 공유하므로 한 번만 추가.
        직전_총계 = int((prior_period_supplies or {}).get("taxable", 0)) + int((prior_period_supplies or {}).get("exempt", 0))
        직전_면세비율 = int((prior_period_supplies or {}).get("exempt", 0)) / 직전_총계 if 직전_총계 > 0 else 0
        직전_과세비율 = int((prior_period_supplies or {}).get("taxable", 0)) / 직전_총계 if 직전_총계 > 0 else 0
        안분_사용 = prior_period_supplies is not None and 직전_총계 > 0  # 직전 비율로 안분하는 경우에만 단계 추가

        # 각 공통사용재화 공급에 대한 식 생성
        for idx, sale in enumerate(common_asset_sales, 1):
            sale_price = int(sale.get("sale_price", 0))
            sale_input_tax = int(sale.get("input_tax", 0))

            if new_business:
                # 신규 사업: 전액 과세표준
                과세표준_식 = f"공급가액{idx}"
                설명_이유 = "신규 사업 (영 제63조③3호)"
            elif sale_price < 500000:
                # 50만 원 미만: 전액 과세표준
                과세표준_식 = f"공급가액{idx}"
                설명_이유 = "50만 원 미만 (영 제63조③2호)"
            elif 직전_총계 > 0 and 직전_면세비율 < 0.05 and sale_price < 50000000:
                # 직전 면세비율 5% 미만 & 5천만 원 미만: 전액 과세표준
                과세표준_식 = f"공급가액{idx}"
                설명_이유 = f"직전 면세비율 {직전_면세비율:.1%} < 5% & 5천만 원 미만 (영 제63조③1호)"
            elif 안분_사용:
                # 직전 과세기간 과세비율로 안분 — 단계 직전과세비율 참조
                과세표준_식 = f"ROUNDDOWN(공급가액{idx}*직전과세비율,0)"
                설명_이유 = f"직전 과세비율 {직전_과세비율:.4f} 안분 (영 제63조①)"
            else:
                과세표준_식 = "0"
                설명_이유 = "안분 불가"

            계산식_입력.extend([
                {"이름": f"공급가액{idx}", "값": sale_price, "설명": f"공통사용재화 공급가액 {idx} = {sale_price:,}원"},
            ])

            계산식.append({
                "이름": f"과세표준{idx}" if len(common_asset_sales) > 1 else "공통사용재화_공급_과세표준",
                "식": 과세표준_식,
                "값": int(sale_price * 직전_과세비율) if (안분_사용 and not (new_business or sale_price < 500000 or (직전_총계 > 0 and 직전_면세비율 < 0.05 and sale_price < 50000000))) else sale_price,
                "설명": f"공통사용재화 {idx} 과세표준: {설명_이유}",
            })

        # 직전 과세기간 과세비율 단계 (안분에 사용한 경우에만)
        if 안분_사용:
            계산식.insert(0, {
                "이름": "직전과세비율",
                "식": "직전과세공급가액 / (직전과세공급가액 + 직전면세공급가액)",
                "값": 직전_과세비율,
                "설명": f"직전 과세기간 과세비율 = 직전과세공급가액 / (직전과세공급가액 + 직전면세공급가액) = {직전_과세비율:.4f}",
            })

        # 공통사용재화_공급_과세표준 합계 식 (여러 개인 경우)
        if len(common_asset_sales) > 1:
            과세표준_parts = "+".join([f"과세표준{n}" for n in range(1, len(common_asset_sales) + 1)])
            계산식.append({
                "이름": "공통사용재화_공급_과세표준",
                "식": 과세표준_parts,
                "값": 공통사용재화_공급_과세표준,
                "설명": f"공통사용재화 공급 과세표준 합계 = Σ 과세표준(1~{len(common_asset_sales)}) = {공통사용재화_공급_과세표준:,}원",
            })
            # 도구 코드와 같은 위치에서 절사: 각 재화별 과세표준에 10% 곱한 후 합산 (개별 ROUNDDOWN 후 합)
            추가_매출세액_식 = "+".join([f"ROUNDDOWN(과세표준{n}*0.1,0)" for n in range(1, len(common_asset_sales) + 1)])
        else:
            추가_매출세액_식 = "ROUNDDOWN(공통사용재화_공급_과세표준*0.1,0)"

        계산식.append({
            "이름": "추가_매출세액",
            "식": 추가_매출세액_식,
            "값": 추가_매출세액,
            "설명": f"추가 매출세액 = Σ ROUNDDOWN(과세표준_i × 10%, 0) = {추가_매출세액:,}원 (영 제63조)",
        })
    else:
        # 공통사용재화 공급이 없는 경우
        계산식_입력.append({"이름": "공통사용재화_공급_과세표준", "값": 0, "설명": "공통사용재화 공급 없음"})
        계산식.append({
            "이름": "공통사용재화_공급_과세표준",
            "식": "0",
            "값": 0,
            "설명": "공통사용재화 공급 없음 → 과세표준 0",
        })
        계산식.append({
            "이름": "추가_매출세액",
            "식": "0",
            "값": 0,
            "설명": "공통사용재화 공급 없음 → 추가 매출세액 0",
        })

    # 공제매입세액 식
    계산식.append({
        "이름": "공제매입세액",
        "식": f"ROUNDDOWN((과세전용매입세액-불공제매입세액)+실지귀속공제+공통매입세액공제+공통사용재화매입세액공제,0)",
        "값": 공제매입세액,
        "설명": f"공제매입세액 = (과세전용 − 불공제) + 실지귀속 + 공통매입세액 + 공통사용재화매입세액 = {공제매입세액:,}원 (법 제40조, 영 제81조)",
    })

    계산식_입력.extend([
        {"이름": "과세전용매입세액", "값": taxable_only - taxable_only_nondeductible, "설명": f"과세전용 매입세액 = {taxable_only:,} − 불공제 {taxable_only_nondeductible:,} = {taxable_only - taxable_only_nondeductible:,}원"},
        {"이름": "불공제매입세액", "값": taxable_only_nondeductible, "설명": f"불공제 매입세액 = {taxable_only_nondeductible:,}원"},
        {"이름": "실지귀속공제", "값": 실지귀속_공제, "설명": f"실지귀속 공제 = {실지귀속_공제:,}원"},
        {"이름": "공통매입세액공제", "값": 공통매입세액_공제, "설명": f"공통매입세액 공제 = {공통매입세액_공제:,}원"},
        {"이름": "공통사용재화매입세액공제", "값": 공통사용재화_매입세액_공제, "설명": f"공통사용재화 매입세액 공제 = {공통사용재화_매입세액_공제:,}원"},
    ])

    return {
        "판단": 판단_목록,
        "계산": {
            "면세비율_예정": 면세비율_예정,
            "면세비율_과세기간": 면세비율_과세기간,
            "공통매입세액_과세기간합계": 공통매입세액_합계 if filing == "확정" else 공통매입세액_합계,
            "불공제_과세기간합계": int(공통매입세액_합계 * 면세비율_과세기간) if filing == "확정" and 공통매입세액_합계 > 0 else 0,
            "불공제_예정반영": 불공제_예정반영 if filing == "확정" else 0,
            "불공제_확정": 불공제_확정 if filing == "확정" else 0,
            "공통매입세액_공제": 공통매입세액_공제,
            "실지귀속_공제": 실지귀속_공제,
            "공통사용재화_공급_과세표준": 공통사용재화_공급_과세표준,
            "공통사용재화_매입세액_공제": 공통사용재화_매입세액_공제,
        },
        "결과": {
            "공제매입세액": 공제매입세액,
            "추가_매출세액": 추가_매출세액,
        },
        "산식": 산식_목록,
        "근거": vocab.basis("공통매입세액_안분", year) + vocab.basis("공통사용재화_공급", year),
        "계산식_입력": 계산식_입력,
        "계산식": 계산식,
    }


# ── 끝에서 끝까지: 감가상각자산 공통매입세액 재계산 ──


@mcp.tool(annotations=CALC)
@_guard
def input_tax_recalculation(asset_type: Annotated[str, Field(description="자산 종류('건물·구축물'|'기타감가상각자산')")],
                            acquired: Annotated[dict, Field(description="{year, period} — 취득 과세기간")],
                            ratios: Annotated[list, Field(description="[{year, period, exempt_ratio}] — 취득 과세기간부터 당기까지 순서대로(마지막이 당기)")],
                            input_tax: Annotated[int | None, Field(description="해당 재화 매입세액 총액(원). bundle을 주면 생략")] = None,
                            bundle: Annotated[dict | None, Field(description="토지·건물 일괄공급 안분 입력(input_tax 대안) {total_price, vat_included, appraisal?, standard_value?, book_value?} — bundled_supply_split과 같은 규칙으로 건물 매입세액을 구함")] = None) -> dict:
    """감가상각자산 공통매입세액 재계산(부가법 제41조, 영 제83조).
    언제: 감가상각자산 매입 후 면세비율이 취득 과세기간 대비 5%p 이상 변동하면 확정신고 때 납부·환급세액을 재계산한다.
    입력:
      · 자산 종류 → asset_type, 취득 과세기간 {year, period} → acquired,
        과세기간별 면세비율 목록(취득~당기, 마지막이 당기) → ratios.
      · 기존: 해당 재화 매입세액 총액 → input_tax.
      · 일괄공급 안분 연결(신규 대안): bundle={total_price, vat_included, appraisal?:{building,land}, standard_value?:{building,land}, book_value?:{building,land}}.
        bundle이 주어지면 내부에서 bundled_supply_split과 같은 규칙으로 건물 공급가액을 안분해 부가가치세를 input_tax로 쓴다.
        input_tax와 bundle 둘 다 주면 error.
    결과: 이력(기간별 비율·기준비율·변동·재계산 여부·경과과세기간수·세액)·결과(당기_납부세액_가산·매입세액, 건물_매입세액은 옛 이름)·판단(안분 기준)·근거."""

    # ── 입력 충돌 검사 ──
    if input_tax is not None and bundle is not None:
        return {"error": "입력 확인 필요 — input_tax와 bundle을 함께 입력할 수 없음. 한 가지 방식으로 입력하세요"}

    # ── bundle 처리 ──
    if bundle is not None:
        total_price = int(bundle.get("total_price", 0))
        vat_included = bool(bundle.get("vat_included", False))
        appraisal = bundle.get("appraisal", None)
        standard_value = bundle.get("standard_value", None)
        book_value = bundle.get("book_value", None)

        안분_규칙 = vocab.rule("토지건물_일괄공급_안분", 2026)
        판단_목록 = []

        if appraisal and appraisal.get("building") and appraisal.get("land"):
            기준_건물 = float(appraisal["building"])
            기준_토지 = float(appraisal["land"])
            판단_목록.append({"항목": "안분 기준", "결론": f"감정평가가액(건물 {기준_건물:,.0f}원, 토지 {기준_토지:,.0f}원) — 우선순위 1순위", "근거": "부가가치세법 시행령 제64조①1호"})
        elif standard_value and standard_value.get("building") and standard_value.get("land"):
            기준_건물 = float(standard_value["building"])
            기준_토지 = float(standard_value["land"])
            판단_목록.append({"항목": "안분 기준", "결론": f"기준시가(건물 {기준_건물:,.0f}원, 토지 {기준_토지:,.0f}원) — 우선순위 2순위", "근거": "부가가치세법 시행령 제64조①1호"})
        elif book_value and book_value.get("building") and book_value.get("land"):
            기준_건물 = float(book_value["building"])
            기준_토지 = float(book_value["land"])
            판단_목록.append({"항목": "안분 기준", "결론": f"장부가액(건물 {기준_건물:,.0f}원, 토지 {기준_토지:,.0f}원) — 우선순위 3순위", "근거": "부가가치세법 시행령 제64조①2호"})
        else:
            return {"error": "안분 기준 없음 — bundle에 appraisal·standard_value·book_value 중 하나 필요"}

        기준_총계 = 기준_건물 + 기준_토지
        건물비율 = 기준_건물 / 기준_총계

        if vat_included:
            L_B비 = 기준_토지 / 기준_건물
            건물_공급가액 = int(total_price / (1 + L_B비 + 0.1))
            부가가치세 = int(건물_공급가액 * 0.1)
        else:
            건물_공급가액 = int(total_price * 건물비율)
            부가가치세 = int(건물_공급가액 * 0.1)

        input_tax = 부가가치세
        판단 = 판단_목록
    else:
        판단 = []

    recal_year = int(acquired["year"])
    재계산_규칙 = vocab.rule("공통매입세액_재계산", recal_year)

    변동기준 = 재계산_규칙["변동기준"]
    체감률표 = 재계산_규칙["체감률"]
    if asset_type not in 체감률표:
        return {"error": f"지원하지 않는 자산 종류 '{asset_type}'", "허용값": list(체감률표)}
    체감률 = 체감률표[asset_type]

    if asset_type not in ("건물·구축물", "기타감가상각자산"):
        return {"error": f"asset_type='{asset_type}' 은(는) 지원하지 않음 — '건물·구축물' 또는 '기타감가상각자산'"}

    acq_year = int(acquired["year"])
    acq_period = acquired["period"]

    def _period_key(y, p):
        return int(y) * 2 + (0 if p == "1기" else 1)

    acq_pk = _period_key(acq_year, acq_period)

    # 기준비율 = 취득 과세기간 비율
    기준비율 = None
    for r in ratios:
        if r["year"] == acq_year and r["period"] == acq_period:
            기준비율 = float(r["exempt_ratio"])
            break
    if 기준비율 is None:
        return {"error": f"취득 과세기간({acq_year}-{acq_period}) 비율이 ratios에 없음"}

    이력 = []
    당기_세액 = 0

    for i, r in enumerate(ratios):
        y = int(r["year"])
        p = r["period"]
        비율 = float(r["exempt_ratio"])
        pk = _period_key(y, p)

        변동 = round(비율 - 기준비율, 10)  # 부동소수 오차 제거

        # 경과과세기간수: 취득 기간은 0, 그다음부터 1씩 증가
        경과과세기간수 = pk - acq_pk
        if 경과과세기간수 < 0:
            경과과세기간수 = 0

        재계산 = abs(변동) >= 변동기준
        현재기준비율 = 기준비율  # 이 시점의 기준비율 (재계산 전)

        if 재계산:
            경감 = max(0.0, 1.0 - 체감률 * 경과과세기간수)
            if 경감 <= 0:
                세액 = 0
            else:
                세액 = int(round(input_tax * 경감 * 변동, 6))
            기준비율 = 비율  # 새 기준비율로 갱신
        else:
            세액 = 0

        이력.append({
            "year": y,
            "period": p,
            "비율": 비율,
            "기준비율": 현재기준비율,
            "변동": 변동,
            "재계산": 재계산,
            "경과과세기간수": 경과과세기간수,
            "세액": 세액,
        })
        당기_세액 = 세액

    return {
        "이력": 이력,
        "결과": {
            "당기_납부세액_가산": 당기_세액,
            "매입세액": input_tax,
            "건물_매입세액": input_tax,   # 옛 이름(호환용) — 자산 종류와 무관하게 같은 값
        },
        "판단": 판단,
        "근거": vocab.basis("공통매입세액_재계산", recal_year),
    }


# ── 끝에서 끝까지: 토지·건물 일괄공급 안분 ──


@mcp.tool(annotations=CALC)
@_guard
def bundled_supply_split(total_price: WON,
                         vat_included: Annotated[bool, Field(description="부가세 포함 여부")],
                         appraisal: Annotated[dict | None, Field(description="{building, land} — 감정평가가액(선택)")] = None,
                         standard_value: Annotated[dict | None, Field(description="{building, land} — 기준시가(선택)")] = None,
                         book_value: Annotated[dict | None, Field(description="{building, land} — 장부가액(선택)")] = None) -> dict:
    """토지·건물 일괄공급 안분(부가법 제29조⑨, 영 제64조①).
    언제: 토지와 건물을 함께 공급하며 가액 구분이 불분명할 때 감정평가가액 → 기준시가 → 장부가액 순으로 건물 공급가액을 안분한다.
    입력: 일괄 거래가액(원) → total_price, 부가세 포함 여부 → vat_included,
      감정평가가액 {building, land} → appraisal, 기준시가 {building, land} → standard_value,
      장부가액 {building, land} → book_value.
    결과: 판단·결과(건물_공급가액·토지_가액·부가가치세)·산식·근거."""

    안분_규칙 = vocab.rule("토지건물_일괄공급_안분", 2026)  # 적용 블록 없음 → 기본 판단 반환

    우선순위 = 안분_규칙["우선순위"]

    판단_목록 = []
    산식_목록 = []

    # 기준 선택: 우선순위 따라
    기준_건물 = 0
    기준_토지 = 0
    기준_명칭 = ""

    if appraisal and appraisal.get("building") and appraisal.get("land"):
        기준_건물 = float(appraisal["building"])
        기준_토지 = float(appraisal["land"])
        기준_명칭 = "감정평가가액"
        판단_목록.append({"항목": "안분 기준", "결론": f"감정평가가액(건물 {기준_건물:,.0f}원, 토지 {기준_토지:,.0f}원) — 우선순위 1순위",
                          "근거": "부가가치세법 시행령 제64조①1호"})
        산식_목록.append(f"안분 기준: 감정평가가액 (건물 {기준_건물:,.0f}원, 토지 {기준_토지:,.0f}원) — 영 제64조①1호")
    elif standard_value and standard_value.get("building") and standard_value.get("land"):
        기준_건물 = float(standard_value["building"])
        기준_토지 = float(standard_value["land"])
        기준_명칭 = "기준시가"
        판단_목록.append({"항목": "안분 기준", "결론": f"기준시가(건물 {기준_건물:,.0f}원, 토지 {기준_토지:,.0f}원) — 우선순위 2순위",
                          "근거": "부가가치세법 시행령 제64조①1호"})
        산식_목록.append(f"안분 기준: 기준시가 (건물 {기준_건물:,.0f}원, 토지 {기준_토지:,.0f}원) — 영 제64조①1호")
    elif book_value and book_value.get("building") and book_value.get("land"):
        기준_건물 = float(book_value["building"])
        기준_토지 = float(book_value["land"])
        기준_명칭 = "장부가액"
        판단_목록.append({"항목": "안분 기준", "결론": f"장부가액(건물 {기준_건물:,.0f}원, 토지 {기준_토지:,.0f}원) — 우선순위 3순위",
                          "근거": "부가가치세법 시행령 제64조①2호"})
        산식_목록.append(f"안분 기준: 장부가액 (건물 {기준_건물:,.0f}원, 토지 {기준_토지:,.0f}원) — 영 제64조①2호")
    else:
        return {"error": "안분 기준 없음 — 감정평가가액·기준시가·장부가액 중 하나를 입력하세요"}

    기준_총계 = 기준_건물 + 기준_토지
    건물비율 = 기준_건물 / 기준_총계

    if vat_included:
        # B + L + 0.1B = total, B:L = 기준 비율
        # B = total ÷ (1 + L/B비 + 0.1)
        L_B비 = 기준_토지 / 기준_건물
        B = int(total_price / (1 + L_B비 + 0.1))
        부가가치세 = int(B * 0.1)
        토지_가액 = total_price - B - 부가가치세
        판단_목록.append({"항목": "건물 공급가액(안분)", "결론": f"{B:,}원 (부가세 포함 총 {total_price:,}원, 건물·토지 비율 {건물비율:.4f}, 부가세 별도)",
                          "근거": "부가가치세법 시행령 제64조①"})
        산식_목록.append(f"부가세 포함: B + L + 0.1B = {total_price:,}, B:L = {기준_건물:,.0f}:{기준_토지:,.0f}")
        산식_목록.append(f"B = {total_price:,} ÷ (1 + {L_B비:.4f} + 0.1) = {B:,}원")
    else:
        B = int(total_price * 건물비율)
        부가가치세 = int(B * 0.1)
        토지_가액 = total_price - B
        판단_목록.append({"항목": "건물 공급가액(안분)", "결론": f"{B:,}원 (총 {total_price:,}원 × 건물비율 {건물비율:.4f})",
                          "근거": "부가가치세법 시행령 제64조①"})
        산식_목록.append(f"부가세 별도: B = {total_price:,} × {건물비율:.4f} = {B:,}원")

    산식_목록.append(f"부가가치세 = {B:,} × 10% = {부가가치세:,}원")
    산식_목록.append(f"토지 가액 = {total_price:,} − {B:,} − {부가가치세:,} = {토지_가액:,}원" if vat_included else
                     f"토지 가액 = {total_price:,} − {B:,} = {토지_가액:,}원")

    # ── 계산식 단계 (SPEC_export_workbook_f · h) ─────────────────────────
    # 기준총계·건물비율·토지건물비율은 입력이 아니라 단계에서 계산한다.
    토지건물비율 = 기준_토지 / 기준_건물 if 기준_건물 > 0 else 0

    계산식_입력 = [
        {"이름": "총거래가액", "값": total_price, "설명": f"일괄 거래가액 = {total_price:,}원"},
        {"이름": "부가세포함여부", "값": vat_included, "설명": f"부가세 포함 여부 = {vat_included}"},
        {"이름": "기준건물", "값": int(기준_건물), "설명": f"건물 기준가액({기준_명칭}) = {기준_건물:,.0f}원"},
        {"이름": "기준토지", "값": int(기준_토지), "설명": f"토지 기준가액({기준_명칭}) = {기준_토지:,.0f}원"},
    ]

    계산식 = [
        {
            "이름": "기준총계",
            "식": "기준건물 + 기준토지",
            "값": int(기준_총계),
            "설명": f"기준가액 합계 = 기준건물 + 기준토지 = {기준_건물:,.0f} + {기준_토지:,.0f} = {기준_총계:,.0f}원",
        },
        {
            "이름": "건물비율",
            "식": "기준건물 / 기준총계",
            "값": 건물비율,
            "설명": f"건물비율 = 기준건물 / 기준총계 = {기준_건물:,.0f} / {기준_총계:,.0f} = {건물비율:.4f}",
        },
        {
            "이름": "건물_공급가액",
            "식": f"IF(부가세포함여부,ROUNDDOWN(총거래가액*기준건물/(기준건물*1.1+기준토지),0),ROUNDDOWN(총거래가액*건물비율,0))",
            "값": B,
            "설명": f"부가세 포함: B = 총거래가액 × 기준건물 / (기준건물×1.1 + 기준토지) = {B:,}원, 부가세 별도: B = 총거래가액 × 건물비율 = {B:,}원 (영 제64조①)",
        },
        {
            "이름": "부가가치세",
            "식": "ROUNDDOWN(건물_공급가액*0.1,0)",
            "값": 부가가치세,
            "설명": f"부가가치세 = 건물_공급가액 × 10% = {B:,} × 0.1 = {부가가치세:,}원",
        },
        {
            "이름": "토지_가액",
            "식": "IF(부가세포함여부,총거래가액-건물_공급가액-부가가치세,총거래가액-건물_공급가액)",
            "값": 토지_가액,
            "설명": f"토지 가액 = 총거래가액 − 건물_공급가액{' − 부가가치세' if vat_included else ''} = {토지_가액:,}원",
        },
    ]

    return {
        "판단": 판단_목록,
        "결과": {
            "건물_공급가액": B,
            "토지_가액": 토지_가액,
            "부가가치세": 부가가치세,
        },
        "산식": 산식_목록,
        "근거": vocab.basis("토지건물_일괄공급_안분", 2026),
        "계산식_입력": 계산식_입력,
        "계산식": 계산식,
    }


@mcp.tool(annotations=CALC)
@_guard
def depreciation_adjustment(
    year: YEAR,
    method: Annotated[Literal["정액법", "정률법", "무신고"], Field(description="신고한 상각방법. 신고하지 않았으면 '무신고'(asset_class로 법정 방법 결정)")],
    acquisition_cost: Annotated[int, Field(description="기말 재무상태표상 취득가액(원, 자산으로 처리한 자본적 지출 포함)", ge=0)],
    booked_depreciation: Annotated[int, Field(description="당기 손익계산서 감가상각비(회사계상, 원)", ge=0)],
    useful_life: Annotated[int | None, Field(description="신고내용연수(무신고면 기준내용연수, 2~60년) — 시행규칙 별표4 상각률을 고름")] = None,
    rate: Annotated[float | None, Field(description="상각률을 직접 줄 때(소수, 예: 0.259). useful_life와 함께 주면 별표4와 같아야 함", gt=0, lt=1)] = None,
    asset_class: Annotated[Literal["건축물", "건축물외_유형자산", "무형자산"] | None, Field(description="자산 구분 — 무신고 법정 방법(영 제26조④)과 정률법 허용 여부 판단")] = None,
    accumulated_dep_begin: Annotated[int | None, Field(description="기초(전기말) 재무상태표상 감가상각누계액(원)", ge=0)] = None,
    accumulated_dep_end: Annotated[int | None, Field(description="기말 재무상태표상 감가상각누계액(원, 당기 감가상각비 포함)", ge=0)] = None,
    expensed_capex: Annotated[int, Field(description="당기 수선비 등 비용으로 처리한 자본적 지출(원)", ge=0)] = 0,
    repair_total: Annotated[int | None, Field(description="이 자산의 당기 수선비 지출 총액(원, 자본적+수익적). 생략하면 expensed_capex", ge=0)] = None,
    periodic_repair: Annotated[bool, Field(description="3년 미만 주기적 수선 지출 여부(영 제31조③3호)")] = False,
    prior_book_value: Annotated[int | None, Field(description="직전 사업연도말 재무상태표상 장부가액(취득가액 − 감가상각누계액, 원) — 수선비 5% 기준용", ge=0)] = None,
    prior_deemed_capex: Annotated[int, Field(description="전기 이전 즉시상각의제액 누계(원) — 정액법 취득가액에 가산", ge=0)] = 0,
    prior_disallowed: Annotated[int, Field(description="전기말 상각부인누계액(유보 잔액, 원). 감가상각의제 △유보 잔액이면 음수")] = 0,
    in_service_date: Annotated[str | None, Field(description="사업연도 중 취득해 사업에 사용한 날(YYYY-MM-DD) — 월할 상각")] = None,
    fy_start: Annotated[str | None, Field(description="사업연도 개시일(YYYY-MM-DD), 생략하면 year-01-01")] = None,
    fy_end: Annotated[str | None, Field(description="사업연도 종료일(YYYY-MM-DD), 생략하면 year-12-31")] = None,
    tax_exempt: Annotated[bool, Field(description="당기 법인세를 면제·감면받았는지(영 제30조 감가상각의제)")] = False,
) -> dict:
    """법인세 감가상각 시부인 — 개별 자산별 상각범위액과 회사계상액 비교(법인세법 제23조, 시행령 제26조·제28조·제30조·제31조·제32조).
언제: 감가상각비 세무조정(상각부인액·시인부족액·전기 부인액 추인·감면법인 의제상각)과 세무상 미상각잔액을 구할 때.
입력:
  year — 사업연도(2016~2026), method — 정액법|정률법|무신고(asset_class 필요).
  useful_life 또는 rate — 내용연수(별표4 상각률) 또는 상각률.
  acquisition_cost·booked_depreciation — 기말 취득가액·당기 감가상각비.
  accumulated_dep_begin 또는 accumulated_dep_end — 기초 또는 기말 감가상각누계액.
  expensed_capex·repair_total·periodic_repair·prior_book_value — 비용 처리한 자본적 지출과 즉시상각의제 예외 판단(600만원·2019년까지 300만원 미만, 직전 장부가액 5% 미만, 3년 미만 주기 수선).
  prior_deemed_capex — 전기 이전 즉시상각의제 누계(정액법).
  prior_disallowed — 전기말 상각부인누계액(△유보면 음수).
  in_service_date — 기중 취득·사용개시일(월할, 1개월 미만은 1개월).
  fy_start·fy_end — 사업연도(12개월만 지원).
  tax_exempt — 법인세 면제·감면 여부.
결과: 세무조정(상각부인액 유보, 추인·감가상각의제 △유보)·상각범위액·회사계상액·기말 상각부인누계액·세무상 미상각잔액·판단·근거. 사실이 빠지거나 맞지 않으면 error와 필요 입력을 돌려주고 계산하지 않음."""
    r = C.depreciation_adjustment(year, method, acquisition_cost, booked_depreciation,
                                   useful_life, rate, asset_class,
                                   accumulated_dep_begin, accumulated_dep_end,
                                   expensed_capex, repair_total, periodic_repair,
                                   prior_book_value, prior_deemed_capex, prior_disallowed,
                                   in_service_date, fy_start, fy_end, tax_exempt)
    if "error" in r:
        return r
    # r = {"결과": {...}, "계산": {...}, "판단": [...], "산식": [...]}
    result_data = r.pop("결과")
    calc_data = r.pop("계산")
    판단_data = r.pop("판단")
    산식_data = r.pop("산식")

    # ── 계산식 단계 (SPEC_export_workbook_f · h) ──────────────────────────
    # 회사계상액은 입력이 아니라 단계에서 계산(장부상각비 + 즉시상각의제).
    # 월수12 → 연간월수로 이름 변경.
    methodName = calc_data["상각방법"]
    상각률 = calc_data["상각률"]
    사용월수 = calc_data["사용월수"]
    세무상취득가액 = calc_data["세무상 취득가액"]
    기초미상각잔액 = calc_data["기초 미상각잔액(즉시상각의제 포함)"]
    회사계상액 = result_data["회사계상액"]
    기초감가상각누계액 = calc_data["기초 감가상각누계액"]
    즉시상각의제 = calc_data["즉시상각의제액"]
    is_정액법 = (methodName == "정액법")

    계산식_입력 = [
        {"이름": "정액법여부", "값": is_정액법, "설명": f"상각방법이 정액법이면 TRUE, 정률법·무신고(정률법)이면 FALSE"},
        {"이름": "취득가액", "값": acquisition_cost, "설명": "기말 재무상태표상 취득가액(원, 자본적 지출 포함)"},
        {"이름": "상각률", "값": 상각률, "설명": f"시행규칙 별표4 상각률 ({상각률}) — 내용연수 {useful_life or '직접입력'}년"},
        {"이름": "사용월수", "값": 사용월수, "설명": f"월할 상각 월수 ({사용월수}개월) — 사업연도 12개월"},
        {"이름": "연간월수", "값": 12, "설명": "연간 월수(기준)"},
        {"이름": "기초감가상각누계", "값": 기초감가상각누계액, "설명": f"기초 감가상각누계액 = {기초감가상각누계액:,}원"},
        {"이름": "전기부인누계", "값": prior_disallowed, "설명": f"전기말 상각부인누계액(유보 잔액) = {prior_disallowed:,}원"},
        {"이름": "즉시상각의제", "값": 즉시상각의제, "설명": f"당기 즉시상각의제액 = {즉시상각의제:,}원"},
        {"이름": "전기즉시상각의제누계", "값": prior_deemed_capex, "설명": f"전기 이전 즉시상각의제 누계 = {prior_deemed_capex:,}원"},
        {"이름": "장부상각비", "값": booked_depreciation, "설명": f"당기 손익계산서 감가상각비(회사계상, 장부상각비) = {booked_depreciation:,}원"},
    ]

    계산식 = [
        {"이름": "정액법_취득가액",
         "식": "취득가액 + 전기즉시상각의제누계 + 즉시상각의제",
         "값": 세무상취득가액,
         "설명": f"정액법용 세무상 취득가액 = 취득가액 + 전기즉시상각의제누계 + 당기즉시상각의제 = {세무상취득가액:,}원"},
        {"이름": "정률법_기초미상각잔액",
         "식": "취득가액 - 기초감가상각누계 + 전기부인누계 + 즉시상각의제",
         "값": 기초미상각잔액,
         "설명": f"정률법용 기초 미상각잔액 = 취득가액 − 기초감가상각누계 + 전기부인누계 + 당기즉시상각의제 = {기초미상각잔액:,}원"},
        {"이름": "상각범위액",
         "식": f"IF(정액법여부,ROUNDDOWN(정액법_취득가액*상각률*사용월수/연간월수,0),ROUNDDOWN(정률법_기초미상각잔액*상각률*사용월수/연간월수,0))",
         "값": result_data["상각범위액"],
         "설명": f"정액법: 정액법_취득가액 × 상각률 × 사용월수/연간월수, 정률 법: 정률법_기초미상각잔액 × 상각률 × 사용월수/연간월수 (법인세법 시행령 제26조⑥⑦)"},
        {"이름": "회사계상액",
         "식": "장부상각비 + 즉시상각의제",
         "값": 회사계상액,
         "설명": f"회사계상 감가상각비 = 장부상각비 + 당기즉시상각의제 = {booked_depreciation:,} + {즉시상각의제:,} = {회사계상액:,}원"},
        {"이름": "상각부인액",
         "식": "MAX(회사계상액-상각범위액,0)",
         "값": result_data["상각부인액"],
         "설명": "MAX(회사계상액 − 상각범위액, 0) — 손금산입 부인액 (유보)"},
        {"이름": "시인부족액",
         "식": "MAX(상각범위액-회사계상액,0)",
         "값": result_data["시인부족액"],
         "설명": "MAX(상각범위액 − 회사계상액, 0) — 시인부족액"},
        {"이름": "전기부인액_추인",
         "식": "MIN(시인부족액,전기부인누계)",
         "값": result_data["전기부인액_추인"],
         "설명": "MIN(시인부족액, 전기부인누계) — 전기 상각부인액 추인 (△유보, 영 제32조①)"},
    ]

    return {"결과": result_data, "계산": calc_data, "판단": 판단_data,
            "근거": "법인세법 제23조, 시행령 제26조·제28조·제30조·제31조·제32조, 시행규칙 별표 4",
            "산식": 산식_data,
            "계산식_입력": 계산식_입력, "계산식": 계산식}


# ── 끝에서 끝까지: 퇴직급여충당금·확정급여형 퇴직연금충당금 세무조정 ──
# 어휘 항목: 퇴직급여충당금_한도(vocab.rule에서 총급여_비율·추계액_적립비율·처리·지급·퇴직연금_지급 읽음)
# facts_vocab.json 수정 금지 — 연도표 적립비율은 vocab.rule로 읽는다.


@mcp.tool(annotations=CALC)
@_guard
def retirement_allowance_adjustment(
    year: YEAR,
    allowance: Annotated[dict, Field(description="장부상 퇴직급여충당금(원) {opening: 기초, closing: 기말, provision?: 당기 설정(손비 계상), paid_cash?: 회사가 직접 지급하고 충당금과 상계, paid_pension?: 퇴직연금운용자산에서 지급하고 충당금과 상계, other_increase?: 합병·사업양수 승계 등 설정 아닌 증가} — 생략 키는 0, 기초+증가−감소=기말이어야 함")],
    prior_disallowed: Annotated[int, Field(description="기초 퇴직급여충당금 부인누계(유보 잔액, 원)", ge=0)],
    estimate: Annotated[dict | None, Field(description="기말 퇴직급여추계액(원) {lump_sum: 일시퇴직기준, actuarial?: 보험수리기준(영 제44조의2④1호의2 금액)} — 설정액이나 퇴직연금이 있으면 필수")] = None,
    salary: Annotated[dict | None, Field(description="당기 총급여액(원) {total: 전체 임원·직원 총급여, dc_members?: 확정기여형 퇴직연금 설정자분, not_eligible?: 퇴직급여 지급대상 아닌 자(예: 규정상 1년 미만 근속자)분} — 당기 설정액이 있으면 필수")] = None,
    succeeded_disallowed: Annotated[int, Field(description="other_increase 중 합병 등으로 승계한 부인누계(유보, 원)", ge=0)] = 0,
    conversion_fund: Annotated[int, Field(description="국민연금 퇴직금전환금 기말잔액(원)", ge=0)] = 0,
    pension: Annotated[dict | None, Field(description="확정급여형 퇴직연금(원) {asset_closing: 퇴직연금운용자산 기말잔액, prior_deducted: 세무상 퇴직연금충당금 기초잔액(이미 손금산입한 부담금 누계, △유보), book_provision?: 결산에 계상한 퇴직연금충당금(신고조정이면 0), book_paid?: 당기 퇴직연금운용자산에서 지급하고 장부 퇴직연금충당금과 상계한 금액(퇴직급여충당금과 상계한 몫은 allowance.paid_pension), dc_transferred_allowance?: 확정기여형 전환자의 설정 전 기계상 퇴직급여충당금} — 없으면 퇴직연금 조정 생략")] = None,
) -> dict:
    """퇴직급여충당금·확정급여형 퇴직연금충당금 세무조정(법인세법 제33조, 시행령 제44조의2·제60조, 시행규칙 제24조).
    언제: 법인세 세무조정에서 퇴직급여충당금 한도초과, 퇴직금 지급 시 상계·부인액 추인, 퇴직연금 부담금 손금산입 한도를 함께 계산할 때.
    입력:
      year — 사업연도(2016~2026), 추계액 적립비율은 연도표(2016년 이후 0%).
      allowance — 장부상 퇴직급여충당금 {opening, closing, provision?, paid_cash?, paid_pension?, other_increase?}.
      prior_disallowed — 기초 부인누계(유보). succeeded_disallowed — 합병 등으로 승계한 부인누계.
      estimate — 기말 추계액 {lump_sum, actuarial?}, 큰 금액을 씀.
      salary — 총급여 {total, dc_members?, not_eligible?}: 확정기여형 설정자·지급대상 아닌 자 제외.
      conversion_fund — 퇴직금전환금 기말잔액.
      pension — 확정급여형 {asset_closing, prior_deducted, book_provision?, book_paid?(운용자산 지급 중 장부 퇴직연금충당금 상계분), dc_transferred_allowance?}.
      운용자산에서 퇴직금을 지급한 해에는 장부 퇴직연금충당금 상계분(book_paid)과 퇴직급여충당금 상계분(paid_pension)을 나눠 넣어야 세무상 연금충당금 잔액이 맞는다.
    결과: 세무조정(과목·조정 익금산입/손금불산입/손금산입·금액·소득처분 유보/△유보), 소득금액_순영향, 계산(세무상 잔액·한도). 입력이 모자라거나 장부 증감이 맞지 않으면 error."""

    # ── 입력 정리 ──
    opening = int(allowance.get("opening", 0))
    closing = int(allowance.get("closing", 0))
    provision = int(allowance.get("provision", 0) or 0)
    paid_cash = int(allowance.get("paid_cash", 0) or 0)
    paid_pension = int(allowance.get("paid_pension", 0) or 0)
    other_increase = int(allowance.get("other_increase", 0) or 0)

    has_prov = provision > 0
    has_pension = pension is not None
    needs_estimate = has_prov or has_pension
    needs_salary = has_prov

    # ── 입력 누락 검사 ──
    missing = []
    if needs_estimate and (estimate is None or int(estimate.get("lump_sum", 0) or 0) == 0):
        missing.append("estimate.lump_sum — 기말 일시퇴직기준 추계액")
    if needs_salary and (salary is None or int(salary.get("total", 0) or 0) == 0):
        missing.append("salary.total — 당기 총급여액")
    if paid_pension > 0 and pension is None:
        missing.append("pension.prior_deducted — 세무상 퇴직연금충당금 기초잔액")
    if missing:
        return {"error": "입력 확인 필요 — 계산하지 않음",
                "문제": [f"input 누락: {', '.join(missing)}"],
                "필요 입력": missing}

    # ── 장부 증감 검증 ──
    calc_closing = opening + provision + other_increase - paid_cash - paid_pension
    if calc_closing != closing:
        return {"error": "입력 확인 필요 — 계산하지 않음",
                "문제": [f"장부 증감 불일치: opening({opening:,}) + provision({provision:,}) + other_increase({other_increase:,}) − paid_cash({paid_cash:,}) − paid_pension({paid_pension:,}) = {calc_closing:,} ≠ closing({closing:,})"],
                "필요 입력": ["opening + provision + other_increase − paid_cash − paid_pension = closing 확인"]}

    # ── 어휘 규칙 ──
    rule_h = vocab.rule("퇴직급여충당금_한도", year)
    급여_비율 = rule_h.get("총급여_비율")
    적립비율 = rule_h.get("추계액_적립비율")
    if 급여_비율 is None:
        return {"error": "입력 확인 필요 — 계산하지 않음",
                "문제": ["vocab '퇴직급여충당금_한도'에 총급여_비율이 없음"],
                "필요 입력": ["vocab 확인"]}

    # ── 추정 추계액 ──
    lump_sum = int(estimate.get("lump_sum", 0) or 0) if estimate else 0
    actuarial = int(estimate.get("actuarial", 0) or 0) if estimate else 0
    추계액 = max(lump_sum, actuarial)

    # ── 대상 총급여 ──
    총급여_전체 = int(salary.get("total", 0) or 0) if salary else 0
    dc_members = int(salary.get("dc_members", 0) or 0) if salary else 0
    not_eligible = int(salary.get("not_eligible", 0) or 0) if salary else 0
    대상_총급여 = 총급여_전체 - dc_members - not_eligible

    # ── 총급여 기준 한도 ──
    총급여기준_한도 = int(대상_총급여 * 급여_비율)

    # ── 추계액 기준 누적한도 ──
    # 누적한도 = max(일시퇴직, 보험수리) × 적립비율 + 퇴직금전환금 − 세무상 설정전 잔액(음수 0 처리 전)
    # 적립비율 적용 추계액 부분
    적립_추계액 = int(추계액 * 적립비율)
    # 세무상 설정전 잔액(아래에서 계산) 차감 전 임시값
    # 실제 계산시 세무상_설정전잔액을 먼저 구해야 함 — 순환 참조는 아님(세무상 설정전잔액은 당기 설정과 별개)
    # "세무상 설정전 잔액 = 기초 − 기초 부인누계 − (상계 중 세무상 잔액 사용분) + (other_increase − succeeded_disallowed)… 구체 산식은 어휘 처리"
    # 어휘 '처리'에 상세 산식이 없으므로 계산 규칙대로:
    # 세무상 설정전 잔액 = 기초 − 기초부인누계 − min(지급총액, max(0, 기초 − 기초부인누계)) + other_increase − succeeded_disallowed

    # 세무상 기초 잔액 = opening − prior_disallowed
    세무_기초_잔액 = opening - prior_disallowed

    # 지급액이 세무상 잔액 초과 여부
    # 운용자산 지급분 중 세무상 퇴직연금충당금과 상계되는 몫은 퇴직급여충당금을 쓰지 않는다(2026-10-30 수정: 종전엔 지급총액 전부를 세무상 잔액에서 뺐음)
    _pd_early = int(pension.get("prior_deducted", 0) or 0) if (pension is not None) else 0
    _book_paid_early = int(pension.get("book_paid", 0) or 0) if (pension is not None) else 0
    _연금감소총_early = min(_book_paid_early + paid_pension, _pd_early) if pension is not None else 0
    _연금상계_early = max(0, _연금감소총_early - min(_book_paid_early, _pd_early)) if pension is not None else 0
    지급총액 = paid_cash + paid_pension
    세무상_잔액_사용액 = min(paid_cash + (paid_pension - _연금상계_early), max(0, 세무_기초_잔액))

    # 세무상 설정전 잔액 (당기 설정 전)
    세무상_설정전잔액 = 세무_기초_잔액 - 세무상_잔액_사용액 + other_increase - succeeded_disallowed
    # 음수면 0 (한도 계산 시)
    세무_설정전잔액_하한 = max(0, 세무상_설정전잔액)

    # 누적한도 = 적립_추계액 + conversion_fund − 세무상 설정전잔액(음수 0 처리)
    누적한도 = 적립_추계액 + conversion_fund - 세무_설정전잔액_하한
    누적한도 = max(0, 누적한도)

    추계액기준_한도 = 누적한도
    누계액기준_한도 = 추계액기준_한도   # 판단·산식 문자열용 동의어

    # ── 퇴직급여충당금 한도 ──
    퇴직급여충당금_한도 = min(총급여기준_한도, 추계액기준_한도)

    # ── 한도초과액 ──
    한도초과액 = max(0, provision - 퇴직급여충당금_한도)

    # prior_deducted는 pension 블록 앞에서 미리 추출(pension 없으면 0) — 부인액_추인 계산에 필요
    _prior_deducted = int(pension.get("prior_deducted", 0) or 0) if has_pension else 0

    # ── 지급 상계 & 부인액 추인 ──
    # 1. 세무상 충당금(기초 − 기초 부인누계)과 지급 상계
    세무상_충당금_기초 = max(0, 세무_기초_잔액)
    # 2. 연금상계분을 뺀 과세 지급액만 세무상 충당금과 상계
    #    연금상계 = min(paid_pension, prior_deducted) (pension 있으면), 없으면 0
    # 2026-10-30: 위 _연금상계_early와 같은 식(장부 퇴직연금충당금 상계분 book_paid를 먼저 세무상 연금충당금에서 뺀 뒤 남는 몫) — CodeSolar 지적 반영
    연금상계 = _연금상계_early if has_pension else 0
    과세_지급액 = paid_cash + (paid_pension - 연금상계)
    # 3. 과세 지급액이 세무상 잔액을 초과하는 부분 = 부인액과 상계 → 손금산입△유보
    초과지급액 = max(0, 과세_지급액 - 세무상_충당금_기초)
    부인액_추인 = min(초과지급액, prior_disallowed)  # 부인누계(유보)와 상계 → 추인

    # ── 지급 후 세무상 잔액(설정 전) 재계산 ──
    # 실제 세무상 사용 후 잔액 = 세무상_설정전잔액(이미 computed) − 한도초과 반영 전
    # 한도초과는 provision에서 발생하므로 별도

    # ── 퇴직연금 조정일 때 ──
    # 연금충당금 상계액은 pension 유무와 무관하게 정의(pension 없으면 0)
    # 운용자산 지급분(paid_pension)은 먼저 세무상 연금충당금(prior_deducted)과 상계 →
    #     그 상계액만큼 퇴직연금충당금 익금산입(유보) + 퇴직급여충당금 손금산입(△유보)
    #     초과분은 퇴직급여충당금 직접 지급으로 보아 부인누계(유보)와 상계 → 손금산입(△유보) 추인
    연금충당금_상계액 = 0
    연금초과 = max(0, paid_pension - 연금충당금_상계액)   # pension 없으면 paid_pension 전액
    dc_transferred = 0   # DC 전환자 설정 전 기계상 퇴직급여충당금 (pension 없을 때 0)
    퇴직연금_결과 = None
    if has_pension:
        asset_closing = int(pension.get("asset_closing", 0) or 0)
        prior_deducted = int(pension.get("prior_deducted", 0) or 0)
        book_provision = int(pension.get("book_provision", 0) or 0)
        dc_transferred = int(pension.get("dc_transferred_allowance", 0) or 0)

        # 운용자산 지급분 처리(2026-10-30 수정): 운용자산에서 지급된 금액은 장부에서 퇴직연금충당금과 상계한 몫(book_paid)과
        # 퇴직급여충당금과 상계한 몫(paid_pension)으로 나뉘지만, 세무상으로는 둘 다 먼저 세무상 퇴직연금충당금(prior_deducted)에서 빠진다.
        # book_paid 몫은 장부·세무가 같아 조정 없음, paid_pension 몫은 익금산입(연금충당금)·손금산입(퇴직급여충당금)으로 바꾼다.
        book_paid = int(pension.get("book_paid", 0) or 0)
        세무_연금_감소총 = min(book_paid + paid_pension, prior_deducted)
        연금충당금_상계액 = max(0, 세무_연금_감소총 - min(book_paid, prior_deducted))
        연금초과 = max(0, paid_pension - 연금충당금_상계액)

        # 이미 손금산입한 부담금 = prior_deducted − 당기 세무상 감소 총액(장부 상계분 + 퇴직급여충당금 상계분)
        # (시행규칙 제24조②: 이미 손금산입한 부담금 = 누계 − 당기 수령·상계액)
        이미_손금산입한_부담금 = prior_deducted - 세무_연금_감소총

        # 기말 세무상 퇴직급여충당금(퇴직연금 반영 전)
        # 설정 후 세무상 잔액 = 세무상_설정전잔액 + (provision − 한도초과액) = 세무_설정전잔액_하한 + min(provision, 한도)
        세무상_퇴직급여충당금_기말 = 세무_설정전잔액_하한 + (provision - 한도초과액)

        # DC 전환자 설정 전 기계상 퇴직급여충당금이 있으면 그만큼 차감(퇴직연금 한도 계산 시)
        # "추계액기준 연금한도 = 추계액 − 세무상 퇴직급여충당금 기말잔액(− dc_transferred_allowance 반영)"
        dc_조정_기말 = max(0, 세무상_퇴직급여충당금_기말 - dc_transferred)
        추계액기준_연금한도 = max(0, 추계액 - dc_조정_기말 - 이미_손금산입한_부담금)

        # 운용자산기준 연금한도 = max(0, asset_closing − 이미_손금산입한_부담금)
        운용자산기준_연금한도 = max(0, asset_closing - 이미_손금산입한_부담금)

        # 손금산입범위 = min(두 한도)
        손금산입범위 = min(추계액기준_연금한도, 운용자산기준_연금한도)
        손금산입범위 = max(0, 손금산입범위)

        # 결산 계상(book_provision)과 비교해 부족분 손금산입·△유보, 초과분 손금불산입·유보
        퇴직연금_납입_부족 = max(0, 손금산입범위 - book_provision)  # 손금산입·△유보
        퇴직연금_납입_초과 = max(0, book_provision - 손금산입범위)  # 손금불산입·유보

        # 세무상_퇴직연금충당금_기말 = 이미_손금산입한_부담금 + 당기 손금산입액
        # 당기 손금산입액 = max(0, 손금산입범위 − book_provision) = 퇴직연금_납입_부족
        세무상_퇴직연금충당금_기말 = 이미_손금산입한_부담금 + 퇴직연금_납입_부족

        퇴직연금_결과 = {
            "연금충당금_상계액": 연금충당금_상계액,
            "장부_연금충당금_상계(book_paid)": book_paid,
            "세무상_연금충당금_당기감소": 세무_연금_감소총,
            "이미_손금산입한_부담금": 이미_손금산입한_부담금,
            "추계액기준_연금한도": 추계액기준_연금한도,
            "운용자산기준_연금한도": 운용자산기준_연금한도,
            "퇴직연금_손금산입범위액": 손금산입범위,
            "퇴직연금_납입_부족": 퇴직연금_납입_부족,
            "퇴직연금_납입_초과": 퇴직연금_납입_초과,
            "세무상_퇴직연금충당금_기말": max(0, 세무상_퇴직연금충당금_기말),
        }

    # ── 세무조정 조립 ──
    settlements = []
    # 1. 한도초과(손금불산입·유보)
    if 한도초과액 > 0:
        settlements.append({"구분": "퇴직급여충당금 한도초과",
                            "조정": "손금불산입",
                            "금액": 한도초과액,
                            "소득처분": "유보",
                            "근거": "법인세법 제33조①·시행령 제60조①·②",
                            "이유": f"당기 설정액 {provision:,}원 > 한도 {퇴직급여충당금_한도:,}원 → 초과 {한도초과액:,}원 손금불산입(유보)"})
    # 2. 부인액 추인(손금산입·△유적)
    if 부인액_추인 > 0:
        settlements.append({"구분": "퇴직급여충당금(부인액 추인)",
                            "조정": "손금산입",
                            "금액": 부인액_추인,
                            "소득처분": "△유보",
                            "근거": "법인세법 제33조②·기본통칙 33-60…5",
                            "이유": f"지급 총액 {지급총액:,}원 중 세무상 충당금 {세무상_충당금_기초:,}원 초과분 {초과지급액:,}원 → 부인누계(유보)와 상계, {부인액_추인:,}원 추인(△유보)"})
    # 3. 퇴직연금 부족(손금산입·△유적) — 있으면
    if has_pension and 퇴직연금_결과:
        r = 퇴직연금_결과
        if r["퇴직연금_납입_부족"] > 0:
            settlements.append({"구분": "퇴직연금충당금",
                                "조정": "손금산입",
                                "금액": r["퇴직연금_납입_부족"],
                                "소득처분": "△유보",
                                "근거": "법인세법 제33조①·시행령 제44조의2④·시행규칙 제24조②",
                                "이유": f"손금산입범위 {r['퇴직연금_손금산입범위액']:,}원 − 결산 계상 {book_provision:,}원 = 부족 {r['퇴직연금_납입_부족']:,}원 → 신고조정 손금산입(△유보)"})
        if r["퇴직연금_납입_초과"] > 0:
            settlements.append({"구분": "퇴직연금충당금 한도초과",
                                "조정": "손금불산입",
                                "금액": r["퇴직연금_납입_초과"],
                                "소득처분": "유보",
                                "근거": "법인세법 제33조①·시행령 제44조의2④·시행규칙 제24조②",
                                "이유": f"결산 계상 {book_provision:,}원 > 손금산입범위 {r['퇴직연금_손금산입범위액']:,}원 → 초과 {r['퇴직연금_납입_초과']:,}원 손금불산입(유보)"})
        # 퇴직연금 지급 시 상계: 운용자산 지급분은 우선 연금충당금(prior_deducted)과 상계
        #   → 상계액만큼 퇴직연금충당금 익금산입(유보) + 퇴직급여충당금 손금산입(△유보)
        #   → 남은 초과분은 퇴직급여충당금 직접 지급으로 보아 부인누계(유보)와 상계 → 추인
        if paid_pension > 0:
            if 연금충당금_상계액 > 0:
                settlements.append({"구분": "퇴직연금충당금(지급 상계)",
                                    "조정": "익금산입",
                                    "금액": 연금충당금_상계액,
                                    "소득처분": "유보",
                                    "근거": "법인세법 제33조②·어휘 퇴직연금_지급(기본통칙 33-60…5)",
                                    "이유": f"퇴직연금운용자산에서 {paid_pension:,}원 지급, 그중 {연금충당금_상계액:,}원(min(지급액, prior_deducted {prior_deducted:,}))은 세무상 연금충당금과 상계 → 퇴직연금충당금 익금산입(유보)"})
                settlements.append({"구분": "퇴직급여충당금(퇴직연금 지급 대체)",
                                    "조정": "손금산입",
                                    "금액": 연금충당금_상계액,
                                    "소득처분": "△유보",
                                    "근거": "법인세법 제33조②·어휘 퇴직연금_지급(기본통칙 33-60…5)",
                                    "이유": f"연금충당금 상계액 {연금충당금_상계액:,}원만큼 퇴직급여충당금과 상계 → 손금산입(△유보)"})


    # ── 소득처분별 합계 ──
    유보_합계 = sum(s["금액"] for s in settlements if s["소득처분"] == "유보")
    델타유보_합계 = sum(s["금액"] for s in settlements if s["소득처분"] == "△유보")
    손금불산입_합계 = sum(s["금액"] for s in settlements if s["조정"] == "손금불산입")
    손금산입_합계 = sum(s["금액"] for s in settlements if s["조정"] == "손금산입")
    익금산입_합계 = sum(s["금액"] for s in settlements if s["조정"] == "익금산입")
    소득금액_순영향 = 손금불산입_합계 + 익금산입_합계 - 손금산입_합계
    가산조정 = 손금불산입_합계 + 익금산입_합계
    차감조정 = 손금산입_합계

    처분별 = {}
    if 유보_합계 != 0:
        처분별["유보"] = 유보_합계
    if 델타유보_합계 != 0:
        처분별["△유보"] = 델타유보_합계

    result = {
        "소득처분별_합계": 처분별,
        "손금불산입_합계": 손금불산입_합계,
        "소득금액_순영향": 소득금액_순영향,
        "가산조정": 가산조정,
        "차감조정": 차감조정,
    }

    # ── 계산 키 재정의 ──
    # 퇴직급여충당금_상계 = (paid_cash + 연금초과) 중 세무상 충당금(기초 − 부인누계)과 실제 상계한 금액
    # 연금초과 = pension 없으면 paid_pension 전액, pension 있으면 paid_pension − 연금충당금_상계액
    과세_지급액 = paid_cash + 연금초과
    퇴직급여충당금_상계 = min(과세_지급액, max(0, 세무_기초_잔액))

    # 누적한도 = 추계액 × 적립비율 + 퇴직금전환금 (세무_설정전잔액 차감 전)
    누적한도_raw = 적립_추계액 + conversion_fund
    # 추계액을_기준으로_한도 = max(0, 누적한도 − 세무_설정전잔액)
    추계액을_기준으로_한도 = max(0, 누적한도_raw - 세무_설정전잔액_하한)
    # 퇴직급여충당금_한도 = min(총급여기준_한도, 추계액을_기준으로_한도)
    퇴직급여_한도 = min(총급여기준_한도, 추계액을_기준으로_한도)
    한도초과액 = max(0, provision - 퇴직급여_한도)

    calc = {
        "세무상_퇴직급여충당금_지급전잔액": 세무_설정전잔액_하한,
        "퇴직연금충당금_상계": 연금충당금_상계액,
        "퇴직급여충당금_상계": 퇴직급여충당금_상계,
        "부인액_추인": 부인액_추인,
        "세무상_설정전잔액": 세무_설정전잔액_하한,
        "대상_총급여": 대상_총급여,
        "총급여기준_한도": 총급여기준_한도,
        "추계액": 추계액,
        "추계액_적립비율": 적립비율,
        "누적한도": 누적한도_raw,
        "추계액기준_한도": 추계액을_기준으로_한도,
        "퇴직급여충당금_한도": 퇴직급여_한도,
        "세무상_퇴직급여충당금_기말": max(0, 세무_설정전잔액_하한 + min(provision, 퇴직급여_한도)),
        "퇴직급여충당금_유보_기말": prior_disallowed + succeeded_disallowed + max(0, provision - 퇴직급여_한도) - 부인액_추인 - 연금충당금_상계액,
    }

    if 퇴직연금_결과:
        calc.update({
            "이미_손금산입한_부담금": 퇴직연금_결과["이미_손금산입한_부담금"],
            "추계액기준_연금한도": 퇴직연금_결과["추계액기준_연금한도"],
            "운용자산기준_연금한도": 퇴직연금_결과["운용자산기준_연금한도"],
            "퇴직연금_손금산입범위액": 퇴직연금_결과["퇴직연금_손금산입범위액"],
            "세무상_퇴직연금충당금_기말": 퇴직연금_결과["세무상_퇴직연금충당금_기말"],
        })

    근거_목록 = vocab.basis("퇴직급여충당금_한도", year)
    근거_문자열 = ", ".join(
        f"{b[0]} {b[1]} {b[2]}" + (f" {b[3]}" if b[3] else "")
        for b in 근거_목록
    )

    판단_목록 = [
        {"항목": "장부 증감 검증", "결론": "일치" if calc_closing == closing else "불일치",
         "근거": "기본통칙 33-60…5"},
        {"항목": "대상 총급여", "결론": f"{대상_총급여:,}원 (전체 {총급여_전체:,} − DC 설정자 {dc_members:,} − 지급대상 외 {not_eligible:,})",
         "근거": "법인세법 시행령 제60조①"},
        {"항목": "총급여 기준 한도", "결론": f"{총급여기준_한도:,}원 (= {대상_총급여:,} × {급여_비율*100:.0f}%)",
         "근거": "법인세법 시행령 제60조①"},
        {"항목": "추계액", "결론": f"{추계액:,}원 (일시퇴직 {lump_sum:,}원, 보험수리 {actuarial:,}원 중 큰 금액)",
         "근거": "법인세법 시행령 제60조②·제44조의2④"},
        {"항목": "추계액 적립비율", "결론": f"{적립비율*100:.0f}%",
         "근거": "법인세법 시행령 제60조②(연도별)"},
        {"항목": "추계액 기준 누적한도", "결론": f"{누계액기준_한도:,}원 (= {추계액:,} × {적립비율*100:.0f}% + 전환금 {conversion_fund:,} − 세무상 설정전 {세무_설정전잔액_하한:,})",
         "근거": "법인세법 시행령 제60조②·④"},
        {"항목": "퇴직급여충당금 한도", "결론": f"{퇴직급여충당금_한도:,}원 (= min(총급여기준 {총급여기준_한도:,}, 추계액기준 {누계액기준_한도:,}))",
         "근거": "법인세법 제33조①·시행령 제60조①·②"},
        {"항목": "한도초과", "결론": f"{한도초과액:,}원" if 한도초과액 > 0 else "없음",
         "근거": "법인세법 제33조①·시행령 제60조③"},
        {"항목": "퇴직금 지급 상계", "결론": f"지급 {지급총액:,}원 → 세무상 충당금 {세무상_충당금_기초:,}원 상계, 초과 {초과지급액:,}원 → 부인누계 추인 {부인액_추인:,}원",
         "근거": "법인세법 제33조②·기본통칙 33-60…5"},
    ]

    if has_pension and 퇴직연금_결과:
        r = 퇴직연금_결과
        판단_목록.extend([
            {"항목": "이미 손금산입한 부담금", "결론": f"{r['이미_손금산입한_부담금']:,}원 (누계 {prior_deducted:,} − 연금충당금 상계액 {r['연금충당금_상계액']:,})",
             "근거": "법인세법 시행규칙 제24조②"},
            {"항목": "추계액 기준 연금한도", "결론": f"{r['추계액기준_연금한도']:,}원 (= 추계액 {추계액:,} − (기말 세무상 퇴직급여충당금 {calc['세무상_퇴직급여충당금_기말']:,} − DC전환 {dc_transferred:,}))",
             "근거": "법인세법 시행령 제44조의2④"},
            {"항목": "운용자산 기준 연금한도", "결론": f"{r['운용자산기준_연금한도']:,}원 (= 운용자산 기말잔액 {asset_closing:,} − 이미 손금산입한 부담금 {r['이미_손금산입한_부담금']:,})",
             "근거": "법인세법 시행규칙 별지 제33호서식 ⑨"},
            {"항목": "퇴직연금 손금산입범위", "결론": f"{r['퇴직연금_손금산입범위액']:,}원 (= min(추계액기준 {r['추계액기준_연금한도']:,}, 운용자산기준 {r['운용자산기준_연금한도']:,}))",
             "근거": "법인세법 시행규칙 별지 제33호서식 ⑩"},
            {"항목": "퇴직연금 결산 계상", "결론": f"{book_provision:,}원",
             "근거": "법인세법 시행령 제44조의2④"},
            {"항목": "퇴직연금 부족/초과", "결론": (f"부족 {r['퇴직연금_납입_부족']:,}원(손금산입·△유적)" if r['퇴직연금_납입_부족'] > 0 else
                                                      f"초과 {r['퇴직연금_납입_초과']:,}원(손금불산입·유적)" if r['퇴직연금_납입_초과'] > 0 else "일치"),
             "근거": "법인세법 시행령 제44조의2④·시행규칙 제24조②"},
        ])

    result_data = {"세무조정": settlements, **result}
    산식_str = (
        f"① 장부 증감: opening({opening:,}) + provision({provision:,}) + other_increase({other_increase:,}) "
        f"− paid_cash({paid_cash:,}) − paid_pension({paid_pension:,}) = closing({closing:,}) [일치]\n"
        f"② 세무상 기초 잔액 = opening − prior_disallowed = {opening:,} − {prior_disallowed:,} = {세무_기초_잔액:,}\n"
        f"③ 세무상 충당금 지급 상계: min(지급액 {지급총액:,}, 세무상 잔액 {세무상_충당금_기초:,}) = {min(지급총액,세무상_충당금_기초):,}원 사용\n"
        f"④ 부인액 추인: max(0, 지급액 {지급총액:,} − 세무상 잔액 {세무상_충당금_기초:,}) × min(…, prior_disallowed {prior_disallowed:,}) = {부인액_추인:,}원\n"
        f"⑤ 대상 총급여 = {총급여_전체:,} − DC {dc_members:,} − 지급대상 외 {not_eligible:,} = {대상_총급여:,}원\n"
        f"⑥ 총급여 기준 한도 = {대상_총급여:,} × {급여_비율*100:.0f}% = {총급여기준_한도:,}원\n"
        f"⑦ 추계 추계액 = max(일시퇴직 {lump_sum:,}, 보험수리 {actuarial:,}) = {추계액:,}원\n"
        f"⑧ 추계액 기준 누적한도 = {추계액:,} × {적립비율*100:.0f}% + 전환금 {conversion_fund:,} = {calc['누적한도']:,}원 (세무상 설정전 잔액 차감 전)\n"
        f"⑨ 추계액을_기준으로_한도 = max(0, ⑧ {calc['누적한도']:,} − 세무상 설정전 {세무_설정전잔액_하한:,}) = {calc['추계액기준_한도']:,}원\n"
        f"⑩ 퇴직급여충당금 한도 = min(⑥ {총급여기준_한도:,}, ⑨ {calc['추계액기준_한도']:,}) = {calc['퇴직급여충당금_한도']:,}원\n"
        f"⑪ 한도초과 = max(0, provision {provision:,} − 한도 {calc['퇴직급여충당금_한도']:,}) = {한도초과액:,}원 → 손금불산입(유보)\n"
        + (f"⑫ 퇴직연금: 이미 손금산입 부담금 = {prior_deducted:,} − 연금충당금 상계액 {r['연금충당금_상계액']:,} = {r['이미_손금산입한_부담금']:,}원\n"
           f"    추계액을_기준으로 연금한도 = {추계액:,} − (기말 충당금 {calc['세무상_퇴직급여충당금_기말']:,} − DC전환 {dc_transferred:,}) = {r['추계액기준_연금한도']:,}원\n"
           f"    운용자산 기준 연금한도 = {asset_closing:,} − 이미 손금산입 {r['이미_손금산입한_부담금']:,} = {r['운용자산기준_연금한도']:,}원\n"
           f"    손금산입범위 = min({r['추계액기준_연금한도']:,}, {r['운용자산기준_연금한도']:,}) = {r['퇴직연금_손금산입범위액']:,}원\n"
           f"    부족 {r['퇴직연금_납입_부족']:,}원 → 손금산입(△유적), 초과 {r['퇴직연금_납입_초과']:,}원 → 손금불산입(유보)\n"
           f"    퇴직연금 지급 상계: {r['연금충당금_상계액']:,}원 → 퇴직연금충당금 익금산입(유보) + 퇴직급여충당금 손금산입(△유적)\n"
           f"    퇴직연금 지급 초과분: {연금초과:,}원 → 퇴직급여충당금 직접 지급(손금산입·△유적)\n"
           if has_pension and r else "")
    )
    return _ok(result_data, 근거_문자열, 계산=calc, 판단=판단_목록, 산식=산식_str)


def main():
    import argparse
    p = argparse.ArgumentParser(prog="korean-tax-calc-mcp", description="한국 세금 계산 MCP 서버")
    p.add_argument("--http", action="store_true"); p.add_argument("--host", default="127.0.0.1")
    p.add_argument("--port", type=int, default=int(os.environ.get("PORT", 8001)))
    a = p.parse_args()
    if a.http:
        import uvicorn
        uvicorn.run(mcp.streamable_http_app(streamable_http_path="/mcp", host=a.host), host=a.host, port=a.port)
    else:
        mcp.run("stdio")


# ── 끝에서 끝까지: 의제배당 — 감자·주식소각·퇴사·탈퇴·해산
# 어휘 항목: 의제배당_감자 (vocab.rule_on / vocab.basis 사용, facts_vocab.json 수정 금지)
# 기준일: 결의·퇴사·탈퇴일(YYYY-MM-DD), 해산은 잔여재산가액 확정일


@mcp.tool(annotations=CALC)
@_guard
def deemed_dividend_capital_reduction(
    shareholder_type: Annotated[Literal["법인", "개인"], Field(description="주주 구분 — 법인세법 제16조·영 제14조(법인) / 소득세법 제17조·영 제27조(개인)")],
    event: Annotated[Literal["감자", "주식소각", "퇴사·탈퇴", "해산"], Field(description="사유 — 감자(유상감자)·주식소각은 일부 소멸 가능, 퇴사·탈퇴·해산은 보유 주식 전부 소멸")],
    resolution_date: Annotated[str, Field(description="감자·소각 결의일(퇴사·탈퇴는 그 날, 해산은 잔여재산가액 확정일), YYYY-MM-DD. 단기소각 2년 기간의 기준")],
    acquisitions: Annotated[list, Field(description="주식 취득·처분 이력 [{date:'YYYY-MM-DD', shares, kind:'유상취득'|'무상주_과세'|'무상주_비과세'|'처분', unit_cost?(유상취득 1주당 취득가액), source?(무상주_비과세 재원: 주식발행초과금·자본준비금·재평가적립금)}] — 무상주 date는 자본전입 결의일")],
    cash_received: Annotated[int, Field(description="받은 금전(원)", ge=0)],
    property_value: Annotated[int, Field(description="받은 금전 외 재산의 가액(원) — 취득 당시 시가(주식이면 영 제14조①·소득세법 시행령 제27조① 평가액). 없으면 0", ge=0)] = 0,
    cancelled_shares: Annotated[int | None, Field(description="소멸(소각·감자)된 주식 수. 퇴사·탈퇴·해산은 생략하면 보유 주식 전부", ge=1)] = None,
    par_value: Annotated[int | None, Field(description="1주당 액면가액(원) — 무상주_과세가 있으면 필수(취득가액 = 액면가액)", ge=1)] = None,
) -> dict:
    """감자·주식소각·퇴사·탈퇴·해산 시 의제배당 = 받은 금전·재산가액 − 소멸 주식 취득가액. 언제: 유상감자·주식소각·퇴사·탈퇴·해산으로 주주가 대가를 받았을 때 의제배당을 계산한다. 무상주 과세 여부는 deemed_dividend_bonus_shares로 먼저 판단해 kind에 넣는다.
    입력:
      shareholder_type — '법인' | '개인'.
      event — '감자' | '주식소각' | '퇴사·탈퇴' | '해산'.
      resolution_date — 결의일 등(YYYY-MM-DD).
      acquisitions — [{date, shares, kind, unit_cost?, source?}]: 유상취득은 unit_cost, 무상주_과세는 액면가액, 무상주_비과세는 0. 처분은 보유 감소.
      cash_received — 받은 금전(원). property_value — 금전 외 재산 시가(원).
      cancelled_shares — 소멸 주식 수(퇴사·탈퇴·해산은 생략 시 전부).
      par_value — 1주당 액면가액(무상주_과세가 있으면 필수).
    결과: 의제배당금액, 단기소각주식수(결의일 전 2년 이내 비과세 무상주 — 먼저 소각, 취득가액 0; 개인은 주식발행초과금 전입분 제외), 평균단가(나머지 주식 총평균), 소멸주식 취득가액, 소각 후 1주당 장부가액, 산식, 근거.
      사실이 빠지거나 맞지 않으면 'error'와 '필요 입력'을 돌려주고 계산하지 않음."""
    # ── dates ──
    try:
        res_dt = date.fromisoformat(resolution_date)
    except (ValueError, TypeError):
        return {"error": "입력 확인 필요 — 계산하지 않음",
                "문제": [f"resolution_date='{resolution_date}' YYYY-MM-DD 형식 아님"],
                "필요 입력": ["resolution_date (YYYY-MM-DD)"]}

    res_yyyymmdd = res_dt.year * 10000 + res_dt.month * 100 + res_dt.day
    year = res_dt.year

    # ── vocab 규칙 ──
    rule = vocab.rule_on("의제배당_감자", res_yyyymmdd)
    basis_list = vocab.basis("의제배당_감자", year)

    단기소각_기간_년 = rule.get("단기소각_기간_년", 2)
    단기소각_적용사유 = rule.get("단기소각_적용사유", [])
    취득구분 = rule.get("취득구분", {})
    단기소각_재원 = rule.get("단기소각_재원", {})
    사유별근거 = rule.get("사유별_근거", {})
    취득가액근거 = rule.get("취득가액_근거", {})

    # ── event·shareholder_type 검증 ──
    if event not in 사유별근거:
        허용 = list(사유별근거.keys())
        return {"error": f"지원하지 않는 사유 '{event}' — 계산 안 함",
                "허용값": {"event": 허용}}
    if shareholder_type not in ("법인", "개인"):
        return {"error": "입력 확인 필요 — 계산하지 않음",
                "문제": [f"shareholder_type='{shareholder_type}' 은(는) 지원하지 않음"],
                "허용값": {"shareholder_type": ["법인", "개인"]}}

    # ── acquisitions 검증 ──
    problems: list[str] = []
    for i, a in enumerate(acquisitions or []):
        if not isinstance(a, dict):
            problems.append(f"항목 {i+1}: dict 아님")
            continue
        d = str(a.get("date", ""))
        if len(d) != 10:
            problems.append(f"항목 {i+1}({a.get('kind','-')}): date='{d}' — YYYY-MM-DD 필요")
            continue
        try:
            adt = date.fromisoformat(d)
        except (ValueError, TypeError):
            problems.append(f"항목 {i+1}({a.get('kind','-')}): date='{d}' 파싱 오류")
            continue
        k = a.get("kind", "")
        if k not in ("유상취득", "무상주_과세", "무상주_비과세", "처분"):
            problems.append(f"항목 {i+1}: kind='{k}' 지원하지 않음")
        elif k in ("무상주_과세", "무상주_비과세") and d > resolution_date:
            problems.append(f"항목 {i+1}({k}): 결의일({resolution_date}) 이후 취득")
        s = a.get("shares")
        if not isinstance(s, int) or s <= 0:
            problems.append(f"항목 {i+1}({k}): shares 양수 정수 필요 (현재 {s!r})")
        if k == "유상취득" and (a.get("unit_cost") is None or not isinstance(a.get("unit_cost"), int) or a["unit_cost"] <= 0):
            problems.append(f"항목 {i+1}(유상취득): unit_cost(1주당 취득가액, 양수 정수) 필요")

    # ── par_value 검증 ──
    has_taxable_bonus = any((a.get("kind") == "무상주_과세") for a in (acquisitions or []))
    if has_taxable_bonus and (par_value is None or not isinstance(par_value, int) or par_value <= 0):
        problems.append("무상주_과세가 있는데 par_value(액면가, 양수 정수) 필요")

    # ── cancelled_shares 검증 ──
    short_term_event = event in 단기소각_적용사유
    if short_term_event and (cancelled_shares is None or not isinstance(cancelled_shares, int) or cancelled_shares <= 0):
        problems.append(f"{event}은(는) cancelled_shares(소멸 주식 수, 양수 정수) 필수")

    has_missing_unit_cost = any(a.get("kind") == "유상취득" and a.get("unit_cost") is None for a in (acquisitions or []))

    if problems:
        return {"error": "입력 확인 필요 — 계산하지 않음",
                "문제": problems,
                "필요 입력": _capital_reduction_필요입력(has_taxable_bonus, short_term_event, problems, event, has_missing_unit_cost),
                "허용값": {"kind": {"유상취득": "주식 유상 매입", "무상주_과세": "과세 무상주(액면가)", "무상주_비과세": "비과세 무상주(취득가액 0)", "처분": "보유 주식 처분"},
                           "단기소각_재원": {"주식발행초과금": "주식발행초과금 전입", "자본준비금": "자본준비금 전입", "재평가적립금": "재평가적립금 전입"}}}

    # ── cutoff 계산 + short_term_event 판정 ──
    cutoff = date(res_dt.year - 단기소각_기간_년, res_dt.month, res_dt.day)
    try:
        cutoff = date(res_dt.year - 단기소각_기간_년, res_dt.month, res_dt.day)
    except ValueError:
        cutoff = date(res_dt.year - 단기소각_기간_년, 12, 31)
    short_term_event = event in 단기소각_적용사유

    # ── 날짜순 정렬 후 fra 구축 + 기타 주식 집계 + 처분 처리 ──
    # (1) 단기소각 대상 비과세 무상주 프로필 추출 (단기소각 이벤트인 경우만)
    단기소각_대상_프로필 = []  # (date, shares, source)
    if short_term_event:
        for a in sorted(acquisitions, key=lambda a: a["date"]):
            if a.get("kind") != "무상주_비과세":
                continue
            adt = date.fromisoformat(a["date"])
            if adt >= cutoff and adt <= res_dt:
                src = a.get("source", "")
                if shareholder_type == "개인" and src == "주식발행초과금":
                    continue  # 개인: 주식발행초과금 전입분 제외
                # 개인: 단기소각기간 내 비과세 무상주는 재원(source) 필수 — 재원 따라 단기소각 포함/제외 결정
                if shareholder_type == "개인" and not src:
                    return {"error": "입력 확인 필요 — 계산하지 않음",
                            "문제": [f"항목 ({a['date']}, 무상주_비과세): 개인 주주의 단기소각기간 내 비과세 무상주는 source(재원) 필요"],
                            "필요 입력": ["acquisitions[i].source — 무상주 재원(주식발행초과금·자본준비금·재평가적립금)"],
                            "허용값": {"kind": {"유상취득": "주식 유상 매입", "무상주_과세": "과세 무상주(액면가)", "무상주_비과세": "비과세 무상주(취득가액 0)", "처분": "보유 주식 처분"},
                                       "단기소각_재원": {"주식발행초과금": "주식발행초과금 전입", "자본준비금": "자본준비금 전입", "재평가적립금": "재평가적립금 전입"}}}
                단기소각_대상_프로필.append((a["date"], int(a["shares"]), src))

    # (2) 처분 목록 (단기소각 기간 중 여부 표시)
    처분_목록 = []
    for a in sorted(acquisitions, key=lambda a: a["date"]):
        if a.get("kind") != "처분":
            continue
        adt = date.fromisoformat(a["date"])
        is_within = short_term_event and adt >= cutoff and adt <= res_dt
        처분_목록.append((a["date"], int(a["shares"]), is_within))

    # (3) fra 구축 + 기타 주식 집계 + 처분 비례 감소 처리
    acqs = sorted(acquisitions, key=lambda a: a["date"])
    held = 0
    fra = []
    기타_총취득원가 = 0  # 유상취득 + 무상주_과세 + 비과세(단기 대상 제외)
    기타_총주식수 = 0
    # 단기소각 대상 프로필의 현재 주식 수 (처분 시 감소)
    단기_프로필_상태 = {dp[0]: {"shares": dp[1], "source": dp[2]} for dp in 단기소각_대상_프로필}

    for a in acqs:
        k = a.get("kind", "")
        s = int(a["shares"])
        if k == "유상취득":
            uc = int(a["unit_cost"])
            held += s
            기타_총취득원가 += s * uc
            기타_총주식수 += s
            fra.append({"date": a["date"], "kind": k, "shares": s, "unit_cost": uc,
                        "source": None, "held_after": held})
        elif k == "무상주_과세":
            held += s
            기타_총취득원가 += s * par_value
            기타_총주식수 += s
            fra.append({"date": a["date"], "kind": k, "shares": s, "unit_cost": par_value,
                        "source": None, "held_after": held})
        elif k == "무상주_비과세":
            src = a.get("source", "")
            held += s
            is_short = (a["date"], src) in {(dp[0], dp[2]) for dp in 단기소각_대상_프로필}
            if not is_short:
                기타_총주식수 += s
                # 취득가액 0이므로 기타_총취득원가 += 0 (변화 없음)
            fra.append({"date": a["date"], "kind": k, "shares": s, "unit_cost": 0,
                        "source": src, "held_after": held, "_short_eligible": is_short})
        elif k == "처분":
            adt = date.fromisoformat(a["date"])
            held -= s
            is_within = short_term_event and adt >= cutoff and adt <= res_dt
            dispose_ratio = 0.0
            if is_within:
                # 단기소각 기간 중 처분 → 비례 처분 (영 제14조③ 후단 / 소득세법 시행령 제27조③ 후단)
                prev_held = held + s
                if prev_held > 0:
                    dispose_ratio = s / prev_held
                    # 단기소각 대상 주식 수 비례 감소
                    총_단기 = sum(dp["shares"] for dp in 단기_프로필_상태.values())
                    if 총_단기 > 0:
                        감소_총량 = int(총_단기 * dispose_ratio)  # 정수 내림
                        # 각 항목에 비례 배분 (나머지는 주식 수가 큰 항목에 할당)
                        배분 = []
                        for dp_key, dp_val in 단기_프로필_상태.items():
                            if dp_val["shares"] > 0:
                                몫 = int(dp_val["shares"] * dispose_ratio)
                                배분.append((dp_key, 몫))
                        남은_감소 = 감소_총량 - sum(b[1] for b in 배분)
                        if 남은_감소 > 0:
                            배분.sort(key=lambda x: 단기_프로필_상태[x[0]]["shares"], reverse=True)
                            for i in range(min(남은_감소, len(배분))):
                                배분[i] = (배분[i][0], 배분[i][1] + 1)
                        for dp_key, 감소 in 배분:
                            단기_프로필_상태[dp_key]["shares"] = max(0, 단기_프로필_상태[dp_key]["shares"] - 감소)
                    # 기타 주식 총취득원가·총주식수 비례 감소
                    기타_총취득원가 = int(기타_총취득원가 * (1 - dispose_ratio))
                    기타_총주식수 = int(기타_총주식수 * (1 - dispose_ratio))
            fra.append({"date": a["date"], "kind": k, "shares": s, "unit_cost": None,
                        "source": None, "held_after": held, "_dispose_ratio": dispose_ratio,
                        "_is_within_period": is_within})

    if held < 0:
        return {"error": "입력 확인 필요 — 계산하지 않음",
                "문제": [f"취득이력상 보유 주식 음수({held}) — 처분의 합계 > 취득 합계"],
                "필요 입력": ["취득·처분 이력 합계 확인"]}

    # 처분 후 단기소각 대상 프로필 재구성
    단기소각_대상_프로필 = [(dp_key, dp_val["shares"], dp_val["source"]) for dp_key, dp_val in 단기_프로필_상태.items() if dp_val["shares"] > 0]
    단기소각_대상_주식수 = sum(dp[1] for dp in 단기소각_대상_프로필)

    # ── 소멸 주식 수 결정 ──
    if event in ("퇴사·탈퇴", "해산"):
        cs = held if cancelled_shares is None else int(cancelled_shares)
    else:
        cs = int(cancelled_shares)

    if cs > held:
        return {"error": "입력 확인 필요 — 계산하지 않음", "문제": [f"cancelled_shares({cs}) > 결의일 보유 주식({held})"],
                "필요 입력": [f"cancelled_shares ≤ 결의일 보유 주식({held})"]}

    if cs == 0:
        return {"error": "입력 확인 필요 — 계산하지 않음",
                "문제": ["소멸 주식 수 0 — 계산할 대상 없음"],
                "필요 입력": ["cancelled_shares ≥ 1"]}

    # ── 소멸 주식 배분: 단기소각 먼저, 나머지를 일반소각 ──
    D = min(단기소각_대상_주식수, cs)
    G = cs - D

    # ── 평균단가 계산 (총평균법) ──
    # 수정 1: 단기소각 대상 비과세 주식은 총평균 계산에서 제외
    #   평균단가 = (단기소각 제외 기타 주식 총취득원가) ÷ (기타 주식 수)
    총취득원가 = 기타_총취득원가
    총주식수 = 기타_총주식수
    평균단가 = (총취득원가 // 총주식수) if 총주식수 > 0 else 0

    # ── 소멸주식 취득가액 ──
    # 수정 2: 해산·퇴사·탈퇴는 단기소각 특례 없음 → 평균 없이 전체 취득가액
    if event in ("퇴사·탈퇴", "해산"):
        소멸주식_취득가액 = 기타_총취득원가  # 처분 비례 반영 후 기타 총취득원가
    else:
        일반취득가액 = G * 평균단가
        소멸주식_취득가액 = D * 0 + 일반취득가액

    # ── 의제배당 ──
    받은대가 = cash_received + property_value
    의제배당 = max(0, 받은대가 - 소멸주식_취득가액)

    # ── 소각 후 잔여 보유주식 + 장부가액 ──
    소각후_보유주식수 = held - cs
    소각전_장부가액 = 기타_총취득원가  # 처분 비례 반영 후
    소각후_장부가액 = 소각전_장부가액 - 소멸주식_취득가액
    소각후_1주당_장부가액 = (소각후_장부가액 // 소각후_보유주식수) if 소각후_보유주식수 > 0 else 0.0

    # ── 근거 조성 ──
    # 개인 주주는 단기소각(D) 유무와 무관하게 소득세법 시행령 제27조(취득가액 계산 근거)를 항상 포함
    # 법인 주주는 D>0일 때만 법인세법 시행령 제14조(취득가액 근거) 포함 (기존 동작 유지)
    근거_기본 = 사유별근거.get(event, {}).get(shareholder_type, "")
    if shareholder_type == "개인":
        근거_전체 = 근거_기본 + f" {취득가액근거.get('개인', '')}"
    elif D > 0:
        근거_전체 = 근거_기본 + f" {취득가액근거.get('법인', '')}"
    else:
        근거_전체 = 근거_기본

    # 비고/주의 — 최상위 키로 반환 (받은대가 ≤ 소멸주식 취득가액이면 "의제배당 없음"이 첫 원소)
    주의: list[str] = []
    if event in ("감자", "주식소각") and 단기소각_재원.get(shareholder_type, {}).get("재평가적립금") == "확인 필요":
        주의.append(f"개인 주주의 재평가적립금 재원 비과세 무상주는 단기소각 해당 여부를 별도 확인 필요")
    if G > 0 and 평균단가 == 0:
        주의.append("평균단가 0 — 소멸주식 취득가액 0, 의제배당 = 받은 대가 전액")
    if property_value > 0:
        주의.append(f"받은 재산가액 {property_value:,}원 포함 — 주식이면 결의일 시가(영 제14조①·소득세법 시행령 제27조①) 기준")
    if 받은대가 <= 소멸주식_취득가액 and 의제배당 == 0:
        주의.insert(0, f"의제배당 없음 — 받은 대가 {받은대가:,}원이 소멸주식 취득가액 {소멸주식_취득가액:,}원 이하(손실은 의제배당 아님, 주식 처분손실로 별도 처리)")

    # ── 취득이력 상세 (소각분 기록) ──
    취득이력: list[dict] = []
    rem_D = D
    rem_G = G
    단기소각_대상_날짜소스_현재 = {(dp[0], dp[2]) for dp in 단기소각_대상_프로필}

    for fa in fra:
        k = fa["kind"]
        s = fa["shares"]
        date_str = fa["date"]
        source = fa.get("source", "") or ""
        is_short_eligible = (k == "무상주_비과세" and (date_str, source) in 단기소각_대상_날짜소스_현재)
        is_disposal = (k == "처분")

        if is_short_eligible and not is_disposal:
            short_take = min(s, rem_D)
            gen_take = 0
            if short_take < s:
                gen_take = min(s - short_take, rem_G)
            short_take = min(short_take, rem_D)
            gen_take = min(gen_take, rem_G)
            rem_D -= short_take
            rem_G -= gen_take
            short_reduction = short_take
            gen_reduction = gen_take
            value_reduction = gen_take * 평균단가
            short_flag = True
        elif is_disposal:
            short_reduction = 0
            gen_reduction = 0
            value_reduction = 0
            short_flag = False
        else:
            gen_take = min(s, rem_G)
            gen_reduction = gen_take
            value_reduction = gen_take * 평균단가
            rem_G -= gen_take
            short_reduction = 0
            short_flag = False

        취득이력.append({
            "일자": date_str,
            "구분": k,
            "주식수": s,
            "취득가액": (fa["unit_cost"] * s) if fa["unit_cost"] else 0,
            "단기소각": short_flag,
            "재원": source if k == "무상주_비과세" else "",
            "단기소각_감소": short_reduction,
            "일반_감소": gen_reduction,
            "취득가액_감소": value_reduction,
        })

    # 산식 문자열
    산식_parts = [
        f"의제배당 = 받은 금전 {cash_received:,}원 + 재산가액 {property_value:,}원 − 소멸주식 취득가액 {소멸주식_취득가액:,}원",
        f"받은 대가 합계 = {받은대가:,}원",
        f"소멸주식수 = 단기소각 {D}주 + 일반소각 {G}주 = {cs}주",
    ]
    if D > 0:
        산식_parts.append(f"단기소각 {D}주(결의일 전 {단기소각_기간_년}년 이내 비과세 무상주, 먼저 소각, 취득가액 0)")
    if G > 0:
        산식_parts.append(f"일반소각 {G}주 × 평균단가 {평균단가:,}원 = 취득가액 {일반취득가액 if event not in ('퇴사·탈퇴', '해산') else 소멸주식_취득가액:,}원")
        산식_parts.append(f"평균단가 = 총취득원가 {총취득원가:,}원 ÷ 총주식수 {총주식수}주 = {평균단가:,}원 (총평균법: 유상취득 + 무상주_과세 액면가 + 무상주_비과세 0, 단기소각 대상 제외)")
    산식_parts.append(f"소멸주식 취득가액 = {소멸주식_취득가액:,}원")
    산식_parts.append(f"의제배당 = {받은대가:,} − {소멸주식_취득가액:,} = {의제배당:,}원 (음수면 0)")
    산식_parts.append(f"소각후 보유주식수 = 결의일 보유 {held}주 − 소멸 {cs}주 = {소각후_보유주식수}주")
    산식_parts.append(f"소각후 1주당 장부가액 = (총취득원가 {총취득원가:,} − 소멸 취득가액 {소멸주식_취득가액:,}) ÷ {소각후_보유주식수}주 = {소각후_1주당_장부가액:,}원" if 소각후_보유주식수 > 0 else "소각후 1주당 장부가액 = 0 (보유 주식 없음)")

    return _ok({
        "의제배당금액": 의제배당,
        "받은대가": 받은대가,
        "소멸주식_취득가액": 소멸주식_취득가액,
        "소멸주식수": cs,
        "단기소각주식수": D,
        "단기소각_보유주식수": 단기소각_대상_주식수,
        "일반소각주식수": G,
        "평균단가": float(평균단가),
        "결의일_보유주식수": held,
        "소각후_보유주식수": 소각후_보유주식수,
        "소각후_1주당_장부가액": float(소각후_1주당_장부가액),
        "주주구분": shareholder_type,
        "사유": event,
        "근거": 근거_전체,
        "산식": "\n".join(산식_parts),
        "취득이력": 취득이력,
    }, 근거_전체, 주의=주의)


def _capital_reduction_필요입력(has_taxable_bonus: bool, short_term_event: bool, problems: list[str], event: str, has_missing_unit_cost: bool) -> list[str]:
    필요 = []
    if has_taxable_bonus:
        필요.append("par_value — 1주당 액면가액(무상주_과세가 있으면 필수)")
    if short_term_event:
        필요.append(f"{event}은(는) cancelled_shares(소멸 주식 수, 양수 정수) 필수")
    if has_missing_unit_cost or any("unit_cost" in p for p in problems):
        필요.append("acquisitions[].unit_cost — 1주당 취득가액(유상취득)")
    if not any("date" in p for p in problems):
        필요.append("date — YYYY-MM-DD")
    if not any("shares" in p for p in problems):
        필요.append("shares")
    if not 필요:
        필요.append("취득이력(date·shares·kind)·resolution_date·cancelled_shares·par_value 확인")
    return 필요


@mcp.tool(annotations=CALC)
@_guard
def capital_gains_income(
        asset_type: Annotated[str, Field(description="양도자산 종류 — '토지'|'건물'(주택 아닌 건물)|'주택'(딸린 토지 포함). 그 밖의 자산은 지원하지 않음")],
        acquisition_date: Annotated[str, Field(description="취득일 YYYY-MM-DD(대금청산일, 등기접수일이 앞서면 그날)")],
        transfer_date: Annotated[str, Field(description="양도일 YYYY-MM-DD(대금청산일, 2017-01-01 ~ 2026-12-31)")],
        transfer_price: Annotated[int, Field(description="양도가액 — 양도 당시 실지거래가액(원). burdened_gift 입력 시 생략 가능(생략 시 0)", ge=0)] = 0,
        acquisition_price: Annotated[int, Field(description="취득가액 — 취득 당시 실지거래가액(원, 취득세·취득 중개수수료 등 포함, 재산세 제외, 환산취득가액 사용 시 0). burdened_gift 입력 시 생략 가능(생략 시 0)", ge=0)] = 0,
        acquisition_tax: Annotated[int, Field(description="취득세(원) — 문항 표의 '취득세'는 이 칸에 넣고, 취득가액 = acquisition_price + 취득세 + 취득 중개수수료로 도구가 합산함")] = 0,
        acquisition_brokerage: Annotated[int, Field(description="취득 시 중개수수료(원) — 문항 표의 '취득시 중개수수료'는 이 칸에 넣고, 취득가액 = acquisition_price + 취득세 + 취득 중개수수료로 도구가 합산함")] = 0,
        holding_tax_paid: Annotated[int, Field(description="보유 중 납부한 재산세·종합부동산세(원) — 필요경비에 넣지 않고 판단에서 필요경비가 아님을 명시함")] = 0,
        capital_expenditure: Annotated[int, Field(description="자본적지출액(원, 증명서류로 확인되는 것, 없으면 0)", ge=0)] = 0,
        transfer_expenses: Annotated[int, Field(description="양도비(원, 양도 시 중개수수료 등, 없으면 0)", ge=0)] = 0,
        registered: Annotated[bool, Field(description="등기 여부(False = 미등기양도자산: 비과세·장기보유특별공제·기본공제 배제)")] = True,
        non_business_land: Annotated[bool, Field(description="토지만: 비사업용 토지 여부(True면 세율 기본세율+10%p, 제104조①8호). apply_basic_deduction=True일 때 산출세액에 반영")] = False,
        acquisition_standard_price: Annotated[int | None, Field(description="환산취득가액 계산용 취득 당시 기준시가(원). acquisition_price=0일 때 transfer_standard_price와 함께 입력")] = None,
        transfer_standard_price: Annotated[int | None, Field(description="환산취득가액 계산용 양도 당시 기준시가(원). acquisition_price=0일 때 acquisition_standard_price와 함께 입력")] = None,
        one_house: Annotated[bool | None, Field(description="주택만: 양도일 현재 1세대 1주택(일시적 2주택 등 1주택으로 보는 경우 포함) 여부")] = None,
        exempt_met: Annotated[bool | None, Field(description="one_house=True일 때: 1세대 1주택 비과세 요건(보유 2년 이상, 취득 당시 조정대상지역이면 거주 2년 이상) 충족 여부")] = None,
        residence_years: Annotated[int | None, Field(description="one_house=True일 때: 보유기간 중 거주기간(만 년수, 1년 미만 버림)", ge=0)] = None,
        residence_start: Annotated[str | None, Field(description="거주 시작일 YYYY-MM-DD — residence_end와 함께 입력하면 거주기간(만 년수)을 도구가 계산함. residence_years와 다르면 문제로 거부")] = None,
        residence_end: Annotated[str | None, Field(description="거주 종료일 YYYY-MM-DD — 생략 시 transfer_date를 사용")] = None,
        heavy_multi_home: Annotated[bool | None, Field(description="one_house=False인 주택: 다주택 중과(제104조⑦) 대상 여부(이전 입력). False=중과 대상 아님. True만으로는 계산하지 않음 — houses_count·adjusted_area로 입력")] = None,
        houses_count: Annotated[int | None, Field(description="one_house=False인 주택: 양도일 현재 1세대 주택 수(조합원입주권·분양권 포함, 분양권은 2021.1.1. 이후 취득분부터). 영 제167조의3①·제167조의10① 각 호 외의 부분 괄호(지방 기준시가 3억원 이하 주택 등)로 주택 수에서 빼는 주택은 사용자가 판단해 뺀 수. 2=1세대 2주택(제104조⑦1호·2호), 3 이상=3주택 이상(3호·4호)", ge=1)] = None,
        adjusted_area: Annotated[bool | None, Field(description="one_house=False인 주택: 양도일 현재 양도 주택이 조정대상지역(주택법 제63조의2①1호)에 있는지")] = None,
        exclusion_reason: Annotated[str | None, Field(description="중과 배제 사유(사용자 주장, 요건 충족은 사용자 확인): '지방저가주택'|'장기임대주택'|'감면대상장기임대주택'|'장기사원용주택'|'조특법감면주택'|'문화재주택'|'상속주택'|'저당권실행취득주택'|'어린이집주택'|'취학근무요양'(2주택만)|'수도권밖주택'(2주택만)|'소송주택'(2주택만)|'기준시가1억이하'(2주택만)|'일반주택'|'조정지역공고전계약'|'1세대1주택특례주택'|'한시배제_경과'(2026.5.10.~11.9. 양도, 영 12호의2 나·다목)|'기타'(exclusion_basis 필수). 한시 배제(2019.12.17.~2020.6.30. 보유 10년 이상, 2022.5.10.~2026.5.9. 보유 2년 이상)는 도구가 날짜로 판정")] = None,
        exclusion_basis: Annotated[str | None, Field(description="exclusion_reason='기타'일 때 근거 조문·호와 요건 충족 사실")] = None,
        designated_area: Annotated[bool, Field(description="non_business_land=True인 토지: 지정지역(제104조의2②) 소재 여부. True면 2018.1.1. 이후 양도분 기본세율+20%p(제104조④3호). apply_basic_deduction=True일 때 산출세액에 반영")] = False,
        apply_basic_deduction: Annotated[bool, Field(description="True면 양도소득 기본공제(연 250만원)와 과세표준까지 계산(같은 해 같은 그룹에서 이미 공제받았으면 False)")] = False,
        burdened_gift: Annotated[dict | None, Field(description="부담부증여(소득세법 시행령 제159조) — {증여가액, 채무액, 평가:'시가'|'기준시가', 취득_기준시가?(평가='기준시가'면 필수), 취득가액_실지?(평가='시가'면 필수)}. 증여가액·채무액으로 양도가액과 취득가액을 자동 계산. transfer_price나 acquisition_price와 함께 입력 불가")] = None) -> dict:
    """토지·건물·주택의 양도차익 → 1세대1주택 비과세·고가주택 안분 → 장기보유특별공제(표1·표2) → 양도소득금액(소득세법 제89조·제92조·제95조·제97조·제103조, 시행령 제154조·제156조·제159조의3(4)·제160조·제176조의2·제163조).
    언제: 부동산 양도소득금액·장기보유특별공제를 양도일 기준 공제율·고가주택 기준(2021.12.8. 전 9억원, 이후 12억원)으로 계산할 때.
    입력:
      자산 종류 → asset_type, 취득일·양도일 → acquisition_date·transfer_date,
      양도가액·취득가액(실지거래가액) → transfer_price·acquisition_price,
      자본적지출·양도비 → capital_expenditure·transfer_expenses, 등기 여부 → registered,
      (환산취득가액) acquisition_price=0이고 취득·양도 당시 기준시가 둘 다 있으면 환산취득가액 자동 계산(소득세법 시행령 제176조의2②2호, 제163조⑥).
      (부담부증여) burdened_gift={증여가액, 채무액, 평가:'시가'|'기준시가', 취득_기준시가?, 취득가액_실지?} — 비율=채무액÷증여가액, 양도가액=증여가액×비율(=채무액), 취득가액: 평가='기준시가'면 취득_기준시가×비율+개산공제(토지·건물 3%, 미등기 0.3%), '시가'면 취득가액_실지×비율. transfer_price·acquisition_price와 함께 입력 불가.
      (주택) 1세대1주택 여부 → one_house, 비과세 요건 충족 → exempt_met, 거주기간(년) → residence_years,
      (주택, 1주택 아님) 다주택 중과(제104조⑦) → houses_count(1세대 주택 수)·adjusted_area(양도일 현재 조정대상지역)·exclusion_reason(배제 사유, 선택). 중과 대상 아님이 확실하면 heavy_multi_home=False.
        중과: 2018.4.1.~2021.5.31. 양도 2주택 +10%p·3주택 이상 +20%p, 2021.6.1.~ +20%p·+30%p, 장기보유특별공제 배제(제95조②), 단기 보유는 제104조①2호·3호 세액과 비교해 큰 세액.
        한시 배제: 2019.12.17.~2020.6.30. 보유 10년 이상(영 12호), 2022.5.10.~2026.5.9. 보유 2년 이상(영 12호의2) — 날짜로 도구가 판정. 2026.5.10. 이후 나·다목 경과 요건은 exclusion_reason='한시배제_경과'(사용자 주장).
      기본공제·과세표준·산출세액까지 → apply_basic_deduction(True면 제104조 세율로 산출세액·지방소득세도 계산: 보유 1년 미만·2년 미만 단기세율, 미등기 70%, 비사업용 토지 → non_business_land, 지정지역 비사업용 토지 → designated_area).
      지원하지 않음(error): 분양권·조합원입주권·주식, 2018.4.1. 전 양도 지정지역 3주택 이상 가산(제104조④1호·2호), 주택 수 계산(사용자 판단), 비과세 보유기간 예외, 주거용 전환, 고가주택 일부 양도.
      토지는 one_house=False, exempt_met=False 입력해도 오류 아님(주택 요건 무관).
    결과: 판단(비과세·고가주택·적용 표, 환산취득가액 사용 시 필요경비 선택 근거)·결과(양도차익·과세대상 양도차익·보유·거주기간·공제율·장기보유특별공제액·양도소득금액, 선택 시 기본공제·과세표준·산출세액·지방소득세·적용 세율)·산식·근거."""
    양도일_dt = date.fromisoformat(transfer_date)
    양도일_int = 양도일_dt.year * 10_000 + 양도일_dt.month * 100 + 양도일_dt.day

    def _vocab_on(name: str, yyyymmdd: int) -> dict:
        return vocab.rule_on(name, yyyymmdd)

    # 부담부증여 처리 (소득세법 시행령 제159조) — 환산취득가액 처리보다 먼저 실행하여
    # acquisition_price가 0인 상태에서의 환산취득가액 오류 체크를 우회
    if burdened_gift is not None:
        if transfer_price > 0 or acquisition_price > 0:
            return {"error": "입력 확인 필요 — 계산하지 않음",
                    "문제": ["burdened_gift와 transfer_price·acquisition_price를 함께 입력할 수 없음 — 둘 중 하나만 사용"]}
        증여가액 = int(burdened_gift["증여가액"])
        채무액 = int(burdened_gift["채무액"])
        평가 = burdened_gift["평가"]
        비율 = 채무액 / 증여가액
        양도_가액 = int(증여가액 * 비율)  # = 채무액
        if 평가 == "기준시가":
            취득_기준시가 = int(burdened_gift["취득_기준시가"])
            취득_기준시가_비율 = int(취득_기준시가 * 비율)
            필요경비개산공제 = int(취득_기준시가_비율 * (0.003 if not registered else 0.03))  # 토지·건물 3%, 미등기 0.3% (영 제159조①·제163조⑥)
            취득_가액 = 취득_기준시가_비율 + 필요경비개산공제
            부담부판단 = (f"비율 = 채무액 {채무액:,}원 ÷ 증여가액 {증여가액:,}원 = {비율:.4f}; "
                          f"양도가액 = 증여가액 {증여가액:,}원 × 비율 = {양도_가액:,}원(= 채무액); "
                          f"취득가액 = 취득기준시가 {취득_기준시가:,}원 × 비율 = {취득_기준시가_비율:,}원 + 개산공제 {필요경비개산공제:,}원(취득기준시가×비율 × {'0.3%' if not registered else '3%'}) = {취득_가액:,}원 (영 제159조①·제163조⑥)")
        elif 평가 == "시가":
            취득_실지 = int(burdened_gift["취득가액_실지"])
            취득_가액 = int(취득_실지 * 비율)
            부담부판단 = (f"비율 = 채무액 {채무액:,}원 ÷ 증여가액 {증여가액:,}원 = {비율:.4f}; "
                          f"양도가액 = 증여가액 {증여가액:,}원 × 비율 = {양도_가액:,}원(= 채무액); "
                          f"취득가액 = 실지취득가액 {취득_실지:,}원 × 비율 = {취득_가액:,}원 (영 제159조①)")
        else:
            return {"error": "입력 확인 필요 — 계산하지 않음",
                    "문제": [f"burdened_gift 평가='{평가}' 은(는) '시가'|'기준시가' 중 하나여야 함"]}
        # burden_gift 계산값을 이후 calc_cg 호출에 사용
        _bp = 양도_가액
        _ba = 취득_가액
        _bt = 0
        _bb = 0
        _bce = 0
        _be = 0
        _burdensome = True
    else:
        _bp = transfer_price
        _ba = acquisition_price
        _bt = acquisition_tax
        _bb = acquisition_brokerage
        _bce = capital_expenditure
        _be = transfer_expenses
        _burdensome = False

    # 환산취득가액 처리 (acquisition_price=0, 기준시가 둘 다 있으면 환산)
    사용환산취득가액 = False
    개산공제 = 0  # 실지거래가액 경로에서도 calc_cg에 전달하기 위한 기본값
    if not _burdensome and _ba == 0:
        if acquisition_standard_price is not None and transfer_standard_price is not None:
            사용환산취득가액 = True
            환산취득가액 = int(_bp * acquisition_standard_price / transfer_standard_price)  # 원 미만 버림 (시행령 제176조의2②2호)
            개산공제 = int(acquisition_standard_price * (0.003 if not registered else 0.03))   # 미등기 0.3%, 등기 3% (시행령 제163조⑥)
            필요경비_환산 = 환산취득가액 + 개산공제
            실제필요경비 = _bce + _be
            if 필요경비_환산 >= 실제필요경비:
                최종취득가격 = 0  # 환산취득가액 사용
                최종필요경비 = 필요경비_환산
                취사용근거 = f"환산취득가액 {환산취득가액:,}원 + 개산공제 {개산공제:,}원 = {필요경비_환산:,}원 (환산취득가액+개산공제가 실제 필요경비 {실제필요경비:,}원 이상 → 제97조②2호 단서로 환산 side 사용)"
            else:
                최종취득가격 = 0
                최종필요경비 = 실제필요경비
                취사용근거 = f"환산취득가액 {환산취득가액:,}원 + 개산공제 {개산공제:,}원 = {필요경비_환산:,}원 < 실제 필요경비(자본적지출 {_bce:,} + 양도비 {_be:,} = {실제필요경비:,}원) → 제97조②2호 단서로 실제 필요경비 사용"
        else:
            return {"error": "입력 확인 필요 — 계산하지 않음",
                    "문제": [f"acquisition_price=0: 환산취득가액 계산 필요 — acquisition_standard_price와 transfer_standard_price 둘 다 입력 필요 (소득세법 시행령 제176조의2②2호)"]}

    # 토지·건물에 one_house/exempt_met=False는 무시 (주택 요건 무관)
    is_주택 = asset_type == "주택"
    if not is_주택 and one_house is False:
        one_house = None
    if not is_주택 and exempt_met is False:
        exempt_met = None

    # residence_start/end로 residence_years 계산 (residence_years가 이미 있으면 검증만)
    if residence_start is not None:
        start_dt = date.fromisoformat(residence_start)
        end_dt = date.fromisoformat(residence_end) if residence_end else 양도일_dt
        calc_residence_years = CG.full_years(start_dt, end_dt)
        if residence_years is not None and residence_years != calc_residence_years:
            return {"error": "입력 확인 필요 — 계산하지 않음",
                    "문제": [f"residence_years({residence_years})와 residence_start/end로 계산한 거주기간({calc_residence_years}년)이 다름"]}
        effective_residence_years = calc_residence_years
    else:
        effective_residence_years = residence_years

    if _burdensome:
        r = CG.calc_cg({
            "asset_type": asset_type,
            "acquisition_date": acquisition_date,
            "transfer_date": transfer_date,
            "transfer_price": _bp,
            "acquisition_price": _ba,
            "acquisition_tax": _bt,
            "acquisition_brokerage": _bb,
            "holding_tax_paid": holding_tax_paid,
            "capital_expenditure": _bce,
            "transfer_expenses": _be,
            "registered": registered,
            "one_house": one_house,
            "exempt_met": exempt_met,
            "residence_years": effective_residence_years,
            "heavy_multi_home": heavy_multi_home,
            "houses_count": houses_count, "adjusted_area": adjusted_area,
            "exclusion_reason": exclusion_reason, "exclusion_basis": exclusion_basis,
            "apply_basic_deduction": apply_basic_deduction,
            "환산취득가액_사용": False,
            "개산공제": 0,
        }, _vocab_on)
    else:
        r = CG.calc_cg({
            "asset_type": asset_type,
            "acquisition_date": acquisition_date,
            "transfer_date": transfer_date,
            "transfer_price": _bp,
            "acquisition_price": _ba if _ba != 0 else 환산취득가액,
            "acquisition_tax": _bt,
            "acquisition_brokerage": _bb,
            "holding_tax_paid": holding_tax_paid,
            "capital_expenditure": _bce if _ba != 0 else 0,
            "transfer_expenses": _be if _ba != 0 else 0,
            "registered": registered,
            "one_house": one_house,
            "exempt_met": exempt_met,
            "residence_years": effective_residence_years,
            "heavy_multi_home": heavy_multi_home,
            "houses_count": houses_count, "adjusted_area": adjusted_area,
            "exclusion_reason": exclusion_reason, "exclusion_basis": exclusion_basis,
            "apply_basic_deduction": apply_basic_deduction,
            "환산취득가액_사용": 사용환산취득가액,
            "개산공제": 개산공제,
        }, _vocab_on)

    if "error" in r:
        return r

    # 부담부증여 판단 추가
    if burdened_gift is not None:
        r["판단"] = r.get("판단", []) + [{"항목": "부담부증여(양도가액·취득가액)", "결론": 부담부판단, "근거": "소득세법 시행령 제159조"}]

    # 환산취득가액 사용 시 필요경비 재계산 및 결과 덮어쓰기
    if 사용환산취득가액:
        필요경비_환산 = 환산취득가액 + 개산공제
        실제필요경비 = capital_expenditure + transfer_expenses
        if 필요경비_환산 >= 실제필요경비:
            최종필요경비 = 필요경비_환산
            최종사용 = "환산취득가액 + 개산공제"
        else:
            최종필요경비 = 실제필요경비
            최종사용 = "실제 필요경비(자본적지출 + 양도비)"
        새양도차익 = max(transfer_price - 최종필요경비, 0)
        # 덮어쓰기 전 원본 값 저장 (고가주택 안분비율 계산용)
        old_양도차익 = r["결과"]["양도차익"]
        old_과세대상 = r["결과"].get("과세대상_양도차익", old_양도차익)
        안분비율 = old_과세대상 / old_양도차익 if old_양도차익 > 0 else 1.0
        r["결과"]["양도차익"] = 새양도차익
        공제율 = r["결과"].get("공제율", 0)
        새장기공제 = int(새양도차익 * 공제율 * 안분비율)
        r["결과"]["장기보유특별공제액"] = 새장기공제
        r["결과"]["과세대상_양도차익"] = int(새양도차익 * 안분비율)
        r["결과"]["양도소득금액"] = max(r["결과"]["과세대상_양도차익"] - 새장기공제, 0)
        r["판단"] = r.get("판단", []) + [{"항목": "필요경비(환산취득가액)", "결론": f"필요경비 {최종필요경비:,}원 = {최종사용} (소득세법 제97조②2호 단서)", "근거": "소득세법 제97조②2호 단서, 시행령 제176조의2②2호·제163조⑥"}]
        if acquisition_tax > 0 or acquisition_brokerage > 0:
            r["판단"] = r.get("판단", []) + [{"항목": "취득가액 구성(환산취득가액)", "결론": "환산취득가액 사용 시 취득세·중개수수료는 합산하지 않음(개산공제로 대체, 소득세법 시행령 제163조)", "근거": "소득세법 시행령 제163조"}]
        r["산식"] = r.get("산식", []) + [f"환산취득가액 {환산취득가액:,}원(취득당시 기준시가 {acquisition_standard_price:,}원 ÷ 양도당시 기준시가 {transfer_standard_price:,}원 × 양도가액 {transfer_price:,}원, 원 미만 버림) + 개산공제 {개산공제:,}원(취득기준시가 × {'0.3%' if not registered else '3%'}) = {필요경비_환산:,}원"]

    # vocab 근거 추가 (양도일 기준 적용 블록 근거 — vocab.rule_on으로 날짜 단위 조회)
    장기보유_적용근거 = vocab.rule_on("양도소득_장기보유특별공제", 양도일_int)
    고가주택_적용근거 = vocab.rule_on("1세대1주택_고가주택", 양도일_int)
    기본공제_적용근거 = vocab.rule_on("양도소득_기본공제", 양도일_int)

    def _근거문자열(rule_dict):
        items = rule_dict.get("근거", [])
        return ", ".join(str(x) for x in items) if items else ""

    r["근거"] = r.get("근거", []) + [
        ["소득세법", "제95조", "②", f"장기보유특별공제 표1·표2 — {_근거문자열(장기보유_적용근거)}"],
    ]
    if "고가주택 안분" in str(r.get("판단", [])):
        r["근거"] = r["근거"] + [
            ["소득세법", "제89조", "①", f"1세대 1주택 비과세·고가주택 — {_근거문자열(고가주택_적용근거)}"],
        ]
    if r["결과"].get("양도소득_기본공제", 0) > 0:
        r["근거"] = r["근거"] + [
            ["소득세법", "제103조", "①", f"양도소득 기본공제 연 250만원 — {_근거문자열(기본공제_적용근거)}"],
        ]

    # 토지는 주택 요건 무관 안내 (one_house=False, exempt_met=False가 무시된 경우)
    if not is_주택:
        r["판단"] = r.get("판단", []) + [{"항목": "자산 종류", "결론": "토지는 주택 요건(1세대1주택·비과세 요건)과 무관", "근거": "소득세법 제89조①·제95조"}]

    if apply_basic_deduction:
        # 환산취득가액 경로는 양도소득금액을 덮어쓰므로 과세표준도 다시 맞춘다
        r["결과"]["양도소득_과세표준"] = max(0, r["결과"]["양도소득금액"] - r["결과"].get("양도소득_기본공제", 0))
        _중과 = r["결과"].get("다주택_중과") or {}
        세 = CG.capital_gains_tax_amount(r["결과"]["양도소득_과세표준"], asset_type, acquisition_date, transfer_date,
                                      registered, non_business_land, I.income_tax, designated_area=designated_area,
                                      heavy_surcharge=_중과["가산세율"] if _중과.get("적용") else 0.0,
                                      heavy_label=_중과.get("구분"))
        식입력, 식단계 = 세.pop("_계산식")
        # 과세표준 입력 대신 양도소득금액 − 기본공제 단계로(제92조②)
        식입력 = [{"이름": "양도소득금액", "값": r["결과"]["양도소득금액"],
                  "설명": "양도차익 − 장기보유특별공제(비과세·고가주택 안분·환산취득가액 판단 반영 결과)"},
                 {"이름": "양도소득_기본공제", "값": r["결과"].get("양도소득_기본공제", 0), "설명": "연 250만원(제103조①, 미등기 제외)"}]
        식단계 = [{"이름": "과세표준", "식": "MAX(0,양도소득금액-양도소득_기본공제)", "값": r["결과"]["양도소득_과세표준"],
                  "설명": "양도소득 과세표준(제92조②)"}] + 식단계
        r["결과"]["산출세액"] = 세["산출세액"]
        r["결과"]["지방소득세"] = 세["지방소득세"]
        r["결과"]["적용_세율"] = 세["적용_세율"]
        서식_세율 = 세["_서식"]
        후보문 = ", ".join("{} {:,}".format(c["세율"], c["세액"]) for c in 세["세율_후보"])
        결론문 = "{} → 산출세액 {:,}원, 지방소득세 {:,}원".format(세["적용_세율"], 세["산출세액"], 세["지방소득세"])
        if len(세["세율_후보"]) > 1:
            결론문 += " (후보: " + 후보문 + ")"
        r["판단"] = r.get("판단", []) + [{"항목": "양도소득 세율", "결론": 결론문, "근거": 세["근거"]}]
        for 주의문 in 세.get("주의", []):
            r["판단"].append({"항목": "지정지역 비사업용 토지", "결론": 주의문, "근거": "소득세법 제104조④3호"})
        r["근거"] = r["근거"] + [["소득세법", "제104조", "⑦" if "⑦" in 세["근거"] else ("④" if "④" in 세["근거"] else "①"), 세["적용_세율"]]]
    elif non_business_land or designated_area:
        r["판단"] = r.get("판단", []) + [{"항목": "양도소득 세율", "결론": "non_business_land·designated_area는 apply_basic_deduction=True일 때 산출세액에 반영", "근거": "소득세법 제104조①8호·④3호"}]

    # ── 서식(별지 제84호 부표1·신고서) 칸 값: 계산·계산식_입력(결과 키는 그대로) ──
    if 사용환산취득가액:
        if 환산취득가액 + 개산공제 >= capital_expenditure + transfer_expenses:
            서식_취득가액, 서식_기타필요경비, 취득가액_종류 = 환산취득가액, 개산공제, "환산취득가액"
        else:
            서식_취득가액, 서식_기타필요경비 = 0, capital_expenditure + transfer_expenses
            취득가액_종류 = "환산취득가액(자본적지출액+양도비 선택, 제97조②2호 단서)"
        서식_양도가액 = transfer_price
    else:
        서식_양도가액 = _bp
        서식_취득가액 = _ba + _bt + _bb
        서식_기타필요경비 = _bce + _be
        취득가액_종류 = "부담부증여 안분" if _burdensome else "실지거래가액"
    결 = r["결과"]
    _고가판단 = next((j for j in r.get("판단", []) if j.get("항목") == "고가주택 안분"), None)
    고가주택_기준 = int(vocab.rule_on("1세대1주택_고가주택", 양도일_int).get("기준금액", 900_000_000)) if _고가판단 else 0
    _중과적용 = bool((결.get("다주택_중과") or {}).get("적용"))
    _표 = 결.get("장기보유특별공제_적용표")
    장기_코드 = "01" if _표 == "표2" else ("02" if _표 == "표1" else ("" if 결.get("비과세") else "03"))
    if asset_type == "토지":
        자산_코드 = "1"
    elif asset_type == "주택":
        자산_코드 = "2" if _고가판단 else "3"
    else:
        자산_코드 = "4"
    계산: dict[str, Any] = {
        "양도가액": 서식_양도가액, "취득가액": 서식_취득가액, "기타필요경비": 서식_기타필요경비,
        "취득가액_종류": 취득가액_종류, "고가주택_기준": 고가주택_기준,
        "비과세_양도차익": 결["양도차익"] - 결["과세대상_양도차익"],
        "자산종류_코드": 자산_코드, "장기보유특별공제_코드": 장기_코드,
        "양도일": transfer_date.replace("-", "."), "취득일": acquisition_date.replace("-", "."),
        "귀속연도": 양도일_dt.year,
    }
    if is_주택 and one_house:
        보유 = 결["보유기간_년"]
        계산["비과세_보유기간_이상"], 계산["비과세_보유기간_미만"] = 보유, 보유 + 1
        if 결.get("거주기간_년") is not None:
            계산["비과세_거주기간_이상"], 계산["비과세_거주기간_미만"] = 결["거주기간_년"], 결["거주기간_년"] + 1
    if _표 == "표2" and 결.get("거주기간_년") is not None:
        계산["장기공제_거주기간_이상"], 계산["장기공제_거주기간_미만"] = 결["거주기간_년"], 결["거주기간_년"] + 1
    if _중과적용:
        계산["장기보유특별공제_배제"] = "다주택 중과(제104조⑦) — 제95조② 괄호로 배제"
    계산식_입력 = [
        {"이름": "양도가액", "값": 서식_양도가액, "설명": "⑪ 양도가액(실지거래가액)"},
        {"이름": "취득가액", "값": 서식_취득가액, "설명": f"⑫ 취득가액({취득가액_종류})"},
        {"이름": "기타필요경비", "값": 서식_기타필요경비, "설명": "⑭ 기타 필요경비(자본적지출액+양도비, 환산취득가액이면 개산공제)"},
        {"이름": "비과세_전액", "값": bool(결.get("비과세")), "설명": "1세대 1주택 비과세(양도가액이 고가주택 기준 이하)"},
        {"이름": "고가주택_기준", "값": 고가주택_기준, "설명": "고가주택 기준금액(고가주택 안분 아니면 0, 시행령 제160조①)"},
        {"이름": "공제율", "값": 결["공제율"], "설명": "장기보유특별공제율(표1·표2, 배제면 0)"},
    ]
    if apply_basic_deduction:
        계산["세율구분_코드"] = 서식_세율["세율구분_코드"]
        계산["세율"], 계산["누진공제"], 계산["가산세율"] = 서식_세율["세율"], 서식_세율["누진공제"], 서식_세율["가산세율"]
        계산식_입력 += [
            {"이름": "기본공제", "값": 결["양도소득_기본공제"], "설명": "⑦ 양도소득기본공제(미등기 0)"},
            {"이름": "세율", "값": 서식_세율["세율"], "설명": "적용 세율(기본세율이면 제55조① 구간 세율)"},
            {"이름": "누진공제", "값": 서식_세율["누진공제"], "설명": "제55조① 기본세율 누진공제액(단일세율이면 0)"},
            {"이름": "가산세율", "값": 서식_세율["가산세율"], "설명": "기본세율에 더하는 세율(비사업용 토지·다주택 중과, 없으면 0)"},
        ]

    if apply_basic_deduction:
        # 서식 칸 입력 + 세액 계산식 단계(엑셀 살아 있는 수식) — 같은 이름은 앞의 것을 쓴다
        _이름 = {x["이름"] for x in 계산식_입력}
        계산식_입력 += [x for x in 식입력 if x["이름"] not in _이름]
        return _ok(r["결과"], r["근거"], 판단=r.get("판단", []), 산식=r.get("산식", []), 계산=계산, 계산식_입력=계산식_입력, 계산식=식단계)
    return _ok(r["결과"], r["근거"], 판단=r.get("판단", []), 산식=r.get("산식", []), 계산=계산, 계산식_입력=계산식_입력)



# ── 상속세 ──────────────────────────────────────────────────────────
@mcp.tool(annotations=CALC)
@_guard
def inheritance_tax(
    year: YEAR,
    estate_value: WON,
    deemed_estate: WON = 0,
    presumed_items: Annotated[list, Field(description="[{kind, window, amount, unexplained}]")] = None,
    nontaxable: WON = 0,
    excluded: WON = 0,
    public_charges: WON = 0,
    funeral_cost: WON = 0,
    enshrinement_cost: WON = 0,
    debts: WON = 0,
    prior_gifts_heirs: WON = 0,
    prior_gifts_others: WON = 0,
    prior_gift_deductions: WON = 0,
    bequest_to_non_heirs: WON = 0,
    renounced_to_next: WON = 0,
    spouse: Annotated[dict | None, Field(description="실제 상속받은 금액·법정상속분 등. {actual(실제 상속액), legal_share('3/5' 같은 분수 문자열 또는 0.6 같은 0 초과 1 이하 소수)}")] = None,
    sole_spouse_heir: Annotated[bool, Field(description="배우자 단독상속")] = False,
    children: Annotated[int, Field(description="자녀 수")] = 0,
    minors: Annotated[list, Field(description="[미성년자별 19세까지 남은 연수]")] = None,
    seniors: Annotated[int, Field(description="65세 이상 인원")] = 0,
    disabled: Annotated[list, Field(description="[장애인별 기대여명 연수]")] = None,
    net_financial_assets: WON = 0,
    cohabiting_house: Annotated[dict | None, Field(description="{value, secured_debt?}")] = None,
    business_farm_deduction: WON = 0,
    appraisal_fee: WON = 0,
    share_appraisal_fee: WON = 0,
    share_appraisal_count: Annotated[int, Field(description="평가기관 수")] = 1,
    generation_skip_surcharge: WON = 0,
    gift_tax_credit: WON = 0,
    other_credits: WON = 0,
    filed: Annotated[bool, Field(description="법정신고기한 내 신고")] = True,
) -> dict:
    """거주자 피상속인 상속세 과세가액·상속공제·과세표준·산출세액·납부할세액. 언제: 상속세 신고·계산 시.
입력: 본래 상속재산가액 → estate_value, 상속개시연도 → year, 간주상속재산 → deemed_estate, 추정상속재산 목록 → presumed_items,
비과세·불산입·공과금·장례비용·봉안비용·채무·사전증여·특별공제 등.
spouse.legal_share는 '3/5' 같은 분수 문자열 또는 0.6 같은 0 초과 1 이하 소수.
결과: 과세가액·상속공제합계·과세표준·산출세액·신고세액공제·납부할세액, 공제 내역, 근거. 상속세 과세가액·공제·세액 계산."""
    return IG.inheritance_tax(
        year=year, estate_value=estate_value, deemed_estate=deemed_estate,
        presumed_items=presumed_items, nontaxable=nontaxable, excluded=excluded,
        public_charges=public_charges, funeral_cost=funeral_cost, enshrinement_cost=enshrinement_cost,
        debts=debts, prior_gifts_heirs=prior_gifts_heirs, prior_gifts_others=prior_gifts_others,
        prior_gift_deductions=prior_gift_deductions, bequest_to_non_heirs=bequest_to_non_heirs,
        renounced_to_next=renounced_to_next, spouse=spouse, sole_spouse_heir=sole_spouse_heir,
        children=children, minors=minors, seniors=seniors, disabled=disabled,
        net_financial_assets=net_financial_assets, cohabiting_house=cohabiting_house,
        business_farm_deduction=business_farm_deduction, appraisal_fee=appraisal_fee,
        share_appraisal_fee=share_appraisal_fee, share_appraisal_count=share_appraisal_count,
        generation_skip_surcharge=generation_skip_surcharge, gift_tax_credit=gift_tax_credit,
        other_credits=other_credits, filed=filed,
    )


# ── 증여세 ──────────────────────────────────────────────────────────
@mcp.tool(annotations=CALC)
@_guard
def gift_tax(
    gift_date: Annotated[str, Field(description="증여일(YYYY-MM-DD, 2016-01-01~2026-12-31)")],
    gift_value: WON,
    relation: Annotated[str, Field(description="수증자 기준 증여자와의 관계 — 배우자|직계존속|직계비속|기타친족|기타. 평소 말(부모·자녀·남편·형제·장인 등)도 됨")],
    recipient_minor: Annotated[bool, Field(description="수증자가 미성년자")] = False,
    prior_gifts_same_donor: WON = 0,
    nontaxable: WON = 0,
    excluded: WON = 0,
    assumed_debt: WON = 0,
    prior_deduction_used: WON = 0,
    marriage_birth: Annotated[bool, Field(description="혼인일 전후 2년·자녀 출생일부터 2년 내 직계존속 증여")] = False,
    marriage_birth_used: WON = 0,
    appraisal_fee: WON = 0,
    share_appraisal_fee: WON = 0,
    share_appraisal_count: Annotated[int, Field(description="평가기관 수")] = 1,
    generation_skip: Annotated[bool, Field(description="자녀 아닌 직계비속 수증")] = False,
    prior_gift_tax_paid: WON = 0,
    other_credits: WON = 0,
    filed: Annotated[bool, Field(description="법정신고기한 내 신고")] = True,
) -> dict:
    """거주자 수증자 증여세 과세가액·증여재산공제·과세표준·산출세액·납부할세액. 언제: 증여세 신고·계산 시.
입력: 증여일 → gift_date, 증여재산가액 → gift_value, 증여자 관계 → relation, 10년 내 동일 증여자 증여 → prior_gifts_same_donor,
부담부증여 인수채무 → assumed_debt, 사전증여 공제 사용량 → prior_deduction_used, 혼인·출산 공제 → marriage_birth.
결과: 과세가액·증여재산공제·과세표준·산출세액·할증세액·신고세액공제·납부할세액, 근거. 증여세 과세가액·공제·세액 계산."""
    rel_in = relation
    relation = _GIFT_RELATION_ALIAS.get(relation.replace(" ", ""), relation)
    if relation not in ("배우자", "직계존속", "직계비속", "기타친족", "기타"):
        return {"error": "입력 확인 필요 — 계산하지 않음", "문제": [f"relation='{rel_in}'은(는) 알 수 없는 관계"],
                "허용값": {"배우자": "남편·아내", "직계존속": "부모·조부모(계부모 포함)", "직계비속": "자녀·손자녀", "기타친족": "6촌 이내 혈족·4촌 이내 인척(형제자매·장인장모·시부모·사위·며느리 등)", "기타": "친족 아닌 사람"},
                "근거": "상속세 및 증여세법 제53조"}
    r = IG.gift_tax(
        gift_date=gift_date, gift_value=gift_value, relation=relation,
        recipient_minor=recipient_minor, prior_gifts_same_donor=prior_gifts_same_donor,
        nontaxable=nontaxable, excluded=excluded, assumed_debt=assumed_debt,
        prior_deduction_used=prior_deduction_used, marriage_birth=marriage_birth,
        marriage_birth_used=marriage_birth_used, appraisal_fee=appraisal_fee,
        share_appraisal_fee=share_appraisal_fee, share_appraisal_count=share_appraisal_count,
        generation_skip=generation_skip, prior_gift_tax_paid=prior_gift_tax_paid,
        other_credits=other_credits, filed=filed,
    )
    if isinstance(r, dict) and relation != rel_in:
        r["입력 해석"] = f"relation '{rel_in}' → '{relation}'(상증법 제53조 증여재산공제 구분)"
    return r


# 증여자 관계 평소 말 → 상증법 제53조 구분(처음 쓰는 사람 점검 2026-10-30). 장인·장모·시부모는 배우자의 직계존속이라 기타친족
_GIFT_RELATION_ALIAS = {
    **{k: "직계존속" for k in ("부모", "부모님", "아버지", "어머니", "아빠", "엄마", "부친", "모친", "조부모", "할아버지", "할머니", "외할아버지", "외할머니", "조부", "조모", "계부", "계모", "직계존속")},
    **{k: "직계비속" for k in ("자녀", "자식", "아들", "딸", "손자", "손녀", "손자녀", "증손", "증손자", "증손녀", "직계비속")},
    **{k: "배우자" for k in ("배우자", "남편", "아내", "처", "부인", "와이프")},
    **{k: "기타친족" for k in ("형제", "자매", "형제자매", "형", "누나", "오빠", "언니", "동생", "삼촌", "고모", "이모", "외삼촌", "큰아버지", "작은아버지", "조카",
                              "사위", "며느리", "장인", "장모", "장인어른", "장모님", "시부모", "시아버지", "시어머니", "친족", "기타친족", "사촌", "친척")},
    **{k: "기타" for k in ("기타", "타인", "남", "지인", "친구", "제3자", "법인")},
}


# ── 연금계좌 인출(연금수령한도·인출순서) ─────────────────────────────
@mcp.tool(annotations=CALC)
@_guard
def pension_account_withdrawal(
    year: Annotated[int, Field(description="귀속연도(과세기간)")],
    birth_date: Annotated[str, Field(description="출생일(YYYY-MM-DD)")],
    join_date: Annotated[str, Field(description="연금계좌 가입일(YYYY-MM-DD)")],
    start_request_date: Annotated[str, Field(description="연금수령 개시 신청일(YYYY-MM-DD)")],
    balance_at_period_start: WON,
    withdrawal: WON,
    contributions_total: WON,
    contributions_not_deducted: WON,
    deferred_retirement_income: WON = 0,
    investment_income: WON = 0,
) -> dict:
    """연금계좌 연금수령한도·인출순서 계산(소득세법 시행령 제40조의2·제40조의3, 2024.1.1. 시행본).
    언제: 연금계좌에서 인출할 때 연금수령분(연금소득)과 연금외수령분(기타소득)을 구분한다.
    입력:
      year — 귀속연도(과세기간)
      birth_date — 출생일
      join_date — 연금계좌 가입일
      start_request_date — 연금수령 개시 신청일
      balance_at_period_start — 과세기간 개시일 현재 평가액(개시 신청 연도는 신청일 현재)
      withdrawal — 당해 인출액
      contributions_total — 납입액 합계
      contributions_not_deducted — 세액공제 받지 못한 납입액(과세제외금액)
      deferred_retirement_income — 이연퇴직소득(있으면 가입 5년 요건 생략)
      investment_income — 운용수익
    결과: 연금수령분·연금외수령분 각 {과세제외금액, 이연퇴직소득, 그밖}, 사적연금_총연금액(종합과세 대상 사적연금),
      이연퇴직_연금수령(연금소득·분리과세), 기타소득_연금외수령, 계산.연금수령연차·연금수령한도, 판단.
    주의: contributions_not_deducted > contributions_total, 음수, 날짜 형식 오류 → error."""
    return GG.pension_account_withdrawal(
        year=year, birth_date=birth_date, join_date=join_date,
        start_request_date=start_request_date, balance_at_period_start=balance_at_period_start,
        withdrawal=withdrawal, contributions_total=contributions_total,
        contributions_not_deducted=contributions_not_deducted,
        deferred_retirement_income=deferred_retirement_income,
        investment_income=investment_income,
    )


# ── 상속·증여재산 평가 ────────────────────────────────────────────────
@mcp.tool(annotations=CALC)
@_guard
def inheritance_gift_valuation(
    kind: Annotated[Literal["상속", "증여"], Field(description="'상속' 또는 '증여'")],
    valuation_date: Annotated[str, Field(description="상속개시일 또는 증여일(YYYY-MM-DD)")],
    supplementary_value: Annotated[int, Field(description="보충적 평가액(개별공시지가 등, 원)")],
    market_values: Annotated[list | None, Field(description="[{대상:'해당재산'|'유사재산', 종류:'매매'|'감정'|'수용'|'경매'|'공매', 가액, 날짜}]")] = None,
    secured_debt: Annotated[dict | None, Field(description="{채권최고액?, 채무액} — 저당권·근저당 등 담보")] = None,
) -> dict:
    """상속·증여재산 평가: 시가 우선 → 보충적 평가 → 담보권 특례(상증법 제60조·제66조, 영 제49조·제63조).
    언제: 상속·증여재산의 시가를 구할 때. 해당재산 가액이 있으면 시가, 없으면 유사재산, 둘 다 없으면 보충적 평가액(법 제61조).
    입력:
      kind — '상속' 또는 '증여'
      valuation_date — 상속개시일·증여일
      supplementary_value — 보충적 평가액(개별공시지가 등)
      market_values — [{대상, 종류, 가액, 날짜}] 선택
      secured_debt — {채권최고액?, 채무액} 선택
    결과: 평가액, 계산.시가·적용·담보채권액, 판단.
    주의: 평가기간 밖 가액은 시가에서 제외(상속=전후 6개월, 증여=전 6개월~후 3개월)."""
    return GG.inheritance_gift_valuation(
        kind=kind, valuation_date=valuation_date, supplementary_value=supplementary_value,
        market_values=market_values, secured_debt=secured_debt,
    )


# ── 수입배당금 익금불산입 ──────────────────────────────────────────
@mcp.tool(annotations=CALC)
@_guard
def dividend_received_deduction(
    year: YEAR,
    investees: Annotated[list, Field(description="[{name, listed?, ratio|shares|lots, dividend, book_value_jeoksu?, exclusion?}]")],
    interest_paid: WON = 0,
    interest_excluded: WON = 0,
    total_assets_jeoksu: Annotated[int, Field(description="자산총계 적수(원·일)")] = 0,
    fy_start: Annotated[str | None, Field(description="사업연도 개시일")] = None,
    fy_end: Annotated[str | None, Field(description="사업연도 종료일")] = None,
    company_type: Annotated[Literal["일반", "지주회사"], Field(description="피지배기업 구분")] = "일반",
    holding_legacy_election: Annotated[bool, Field(description="지주회사가 종전 제18조의3 비율 선택")] = False,
) -> dict:
    """내국법인 수입배당금액 익금불산입(피출자법인별 익금불산입률·지급이자 차감·세무조정). 언제: 법인세 세무조정에서 수입배당금 익금을 부인할 때.
입력: 사업연도 → year, 피출자법인 목록 → investees, 지급이자 → interest_paid, 차감 이자 → interest_excluded, 자산총액 적수 → total_assets_jeoksu.
결과: 익금불산입액의 피출자법인별 내역·세무조정·근거. 수입배당금 익금불산입 계산."""
    r = C.dividend_received_deduction(
        year=year, investees=investees, interest_paid=interest_paid,
        interest_excluded=interest_excluded, total_assets_jeoksu=total_assets_jeoksu,
        fy_start=fy_start, fy_end=fy_end, company_type=company_type,
        holding_legacy_election=holding_legacy_election,
    )
    if "error" in r:
        return r
    result_data = r.pop("결과")
    calc_data = r.pop("계산")
    판단_data = r.pop("판단")
    근거_data = r.pop("근거")

    # ── 계산식 단계 (SPEC_export_workbook_f) ──────────────────────────────
    A = calc_data["차입금이자"]  # max(0, interest_paid - interest_excluded)
    total_assets = calc_data["자산총액_적수"]  # total_assets_jeoksu if A > 0 else None

    # 원본 입력값(지급이자, 제외이자)을 계산식_입력에 추가 → 차입금이자는 단계로 계산
    계산식_입력 = [
        {"이름": "지급이자", "값": interest_paid, "설명": f"지급이자(원금 상환액 포함) = {interest_paid:,}원"},
        {"이름": "제외이자", "값": interest_excluded, "설명": f"차감 제외 이자 = {interest_excluded:,}원 (법 제18조의2①2호)"},
        {"이름": "자산총액적수", "값": total_assets if total_assets else 0, "설명": f"자산총계 적수(원·일) = {total_assets:,}" if total_assets else "자산총계 적수 없음 (차입금이자 0)"},
    ]

    # 입력 investees에서 book_value_jeoksu 등 추가 정보 추출 (결과에는 없는 값)
    inv_input_map = {inv.get("name"): inv for inv in investees}

    계산식 = []
    # 차입금이자 단계 (A > 0이든 0이든 같은 식)
    계산식.append({
        "이름": "차입금이자",
        "식": "MAX(0,지급이자-제외이자)",
        "값": A,
        "설명": f"차감 대상 지급이자 = MAX(0, 지급이자 − 제외이자) = MAX(0, {interest_paid:,} − {interest_excluded:,}) = {A:,}원 (법 제18조의2①2호)",
    })
    대상액_합 = 0
    차감액_합 = 0

    for idx, inv_result in enumerate(result_data.get("피출자법인별", []), 1):
        name = inv_result.get("이름", f"피출자법인{idx}")
        대상액 = inv_result.get("익금불산입_대상액", 0)
        차감액 = inv_result.get("차감액", 0)
        대상배당금 = inv_result.get("대상_배당금", 0)
        익금불산입률 = inv_result.get("익금불산입률", 0)
        수입배당금액 = inv_result.get("수입배당금액", 0)
        제외배당금 = inv_result.get("제외_배당금", 0)

        # 입력에서 book_value_jeoksu 가져오기
        inv_input = inv_input_map.get(name, {})
        장부가액적수 = inv_input.get("book_value_jeoksu", 0)

        대상액_합 += 대상액
        차감액_합 += 차감액

        계산식_입력.extend([
            {"이름": f"수입배당금액{idx}", "값": 수입배당금액, "설명": f"{name} 수입배당금액 = {수입배당금액:,}원"},
            {"이름": f"제외배당금{idx}", "값": 제외배당금, "설명": f"{name} 3개월 이내 취득 주식 배당 제외분 = {제외배당금:,}원 (법 제18조의2②)" if 제외배당금 else f"{name} 3개월 이내 취득분 없음"},
            {"이름": f"익금불산입률{idx}", "값": 익금불산입률, "설명": f"{name} 적용 익금불산입률 = {익금불산입률} (출자비율·상장 여부·연도 기준, 법 제18조의2①)"},
            {"이름": f"장부가액적수{idx}", "값": 장부가액적수, "설명": f"{name} 익금불산입 대상 주식 장부가액 적수 = {장부가액적수:,}원·일 (법 제18조의2①2호·영 제17조의2③)"},
        ])

        # 대상배당금{idx}는 단계 (값이 0이어도 같은 식)
        계산식.append({
            "이름": f"대상배당금{idx}",
            "식": f"수입배당금액{idx}-제외배당금{idx}",
            "값": 대상배당금,
            "설명": f"{name} 익금불산입 대상 배당 = 수입배당금액 − 제외배당금 = {수입배당금액:,} − {제외배당금:,} = {대상배당금:,}원",
        })
        계산식.append({
            "이름": f"대상액{idx}",
            "식": f"ROUNDDOWN(대상배당금{idx}*익금불산입률{idx},0)",
            "값": 대상액,
            "설명": f"{name} 대상액 = 대상배당금 × 익금불산입률 = {대상배당금:,} × {익금불산입률} = {대상액:,}원",
        })
        # 차감액{idx}: 차입금이자 0이어도 같은 식 (MAX(1, 자산총액적수)로 0 나누기 방지)
        계산식.append({
            "이름": f"차감액{idx}",
            "식": f"ROUNDDOWN(차입금이자*장부가액적수{idx}*익금불산입률{idx}/MAX(1,자산총액적수),0)",
            "값": 차감액,
            "설명": f"{name} 차감액 = ROUNDDOWN(차입금이자 × 장부가액적수 × 익금불산입률 ÷ MAX(1, 자산총액적수), 0) = {A:,} × {장부가액적수:,} × {익금불산입률} ÷ {max(1, total_assets) if total_assets else 1:,} = {차감액:,}원",
        })

    # 익금불산입액 식 구성
    n = len(result_data.get('피출자법인별', []))
    if n > 1:
        대상액_parts = "+".join([f"대상액{i}" for i in range(1, n + 1)])
        차감액_parts = "+".join([f"차감액{i}" for i in range(1, n + 1)])
        익금불산입액_식 = f"MAX(0,{대상액_parts}-({차감액_parts}))"
    else:
        익금불산입액_식 = "MAX(0,대상액1-차감액1)"

    계산식.append({
        "이름": "익금불산입액",
        "식": 익금불산입액_식,
        "값": result_data["익금불산입액"],
        "설명": f"익금불산입액 = max(0, Σ대상액 − Σ차감액) = max(0, {대상액_합:,} − {차감액_합:,}) = {result_data['익금불산입액']:,}원 (법 제18조의2①)",
    })

    return {"결과": result_data, "계산": calc_data, "판단": 판단_data,
            "근거": 근거_data,
            "계산식_입력": 계산식_입력, "계산식": 계산식}


# ── 특별세액공제 ────────────────────────────────────────────────────
@mcp.tool(annotations=CALC)
@_guard
def special_tax_credit(
    year: YEAR,
    earner_type: Annotated[Literal["근로", "일용근로", "근로소득없음"], Field(description="근로소득 종류")],
    total_salary: Annotated[int, Field(description="총급여액(원). 문항에 총급여액이 없으면 0이 아니라 생략 — 0이면 의료비 문턱(total_salary×문턱비율)이 0이 되어 모든 의료비가 특정지출로 분류됨")],
    persons: Annotated[list, Field(description=(
        "[{name, relation, age, disabled?, income_ok?, severe_illness?, age_at_year_start?, birth_date?, school_grade?}] — "
        "name: 사람 표시 이름(문자열, 중복 불가). "
        "relation: 공제대상자와의 관계. 허용값: 본인·배우자·직계비속·입양자·형제자매·직계존속·위탁아동·수급자. "
        "문항 표현 그대로 '자녀'·'아들'·'딸'은 '직계비속', '부모'·'모친'·'부친'은 '직계존속'으로 적을 것. "
        "age: 과세기간 종료일(12/31) 현재 만 나이(정수). disabled: 장애인 여부(bool). "
        "age_at_year_start: 과세기간 개시일(1/1) 현재 만 나이 — 의료비 6세 이하 특정지출(2024~)에서 age=7이면 필요, "
        "birth_date(YYYY-MM-DD)로 대신 가능. school_grade: 초등학교 학년(예체능학원_초등저학년에서 age 9 이상일 때 필요). "
        "income_ok: 연 소득금액 100만원 이하(근로소득만 있으면 총급여 500만원 이하)면 true — 문항의 '소득 없음'은 true로 입력. "
        "severe_illness: 중증질환자·희귀난치성질환자·결핵환자 여부(bool) — 2018년부터 의료비 특정지출, 2016~2017은 일반(700만원 한도). "
        "예시: [{'name':'본인','relation':'본인','age':50,'income_ok':true}, {'name':'딸','relation':'직계비속','age':16,'income_ok':true}]"
    ))],
    expenses: Annotated[list, Field(description=(
        "[{kind, amount, person, reimbursed?}] — person은 persons의 name. "
        "1명당 연 한도는 같은 person의 지출을 합산해 적용: 안경 50만원, 교복 50만원, 체험학습 30만원(2017~), "
        "취학전·초중고 300만원(교복·체험학습 포함), 대학 900만원. 같은 사람의 지출은 같은 name으로 넣을 것. "
        "kind 예체능학원_초등저학년(2026~, 제59조의4③1호마목): 종료일 현재 9세 미만 또는 2학년 이하 초등학생의 "
        "예능학원·체육시설 월 단위·주 1회 이상 교습과정 수강료, 300만원 한도에 합산"))],
    wage_computed_tax: Annotated[int | None, Field(description="근로소득 종합소득산출세액")] = None,
    monthly_rent_credit: WON = 0,
    donation_credit: WON = 0,
    special_income_deduction: WON = 0,
) -> dict:
    """근로소득자 특별세액공제(보험료·의료비·교육비)와 표준세액공제 비교. 언제: 연말정산·종합소득세 신고에서 특별세액공제와 표준세액공제 중 선택할 때.
입력: 귀속연도 → year, 총급여액 → total_salary, 공제대상자 → persons, 지출 목록 → expenses, 근로소득 산출세액 → wage_computed_tax(선택).
결과: 보험료·의료비·교육비 공제액·특별세액공제 합계·표준세액공제 선택·적용 공제·근거. 특별세액공제 vs 표준세액공제 비교.
연도별: 의료비 특정(2호) 중증질환자 등 2018~, 과세기간 개시일 현재 6세 이하 2024~, 미숙아·선천성이상아 20% 2022~. 교육비 체험학습 2017~, 대입전형료 2023~."""
    return I.special_tax_credit(
        year=year, earner_type=earner_type, total_salary=total_salary,
        persons=persons, expenses=expenses, wage_computed_tax=wage_computed_tax,
        monthly_rent_credit=monthly_rent_credit, donation_credit=donation_credit,
        special_income_deduction=special_income_deduction,
    )


# ── 종합소득 인적공제(소득세법 제50조·제51조·제53조) ─────────────────────────────
@mcp.tool(annotations=CALC)
@_guard
def personal_deduction(
    year: YEAR,
    taxpayer: Annotated[dict, Field(description=(
        "{sex, birth_year, has_spouse, comprehensive_income, disabled?, household_head?} — "
        "sex: '남'|'여'. birth_year: 출생연도(정수). has_spouse: 배우자 유무(bool). "
        "comprehensive_income: 종합소득금액(원, 기본공제 한도 판단용). "
        "disabled: 장애인 여부(bool, 기본 무시). household_head: 세대주 여부(bool, 부녀자 공제 판단용)."
    ))],
    persons: Annotated[list, Field(description=(
        "[{name, relation, birth_year, income_items?, income?, wage_only_gross?, disabled?, cohabiting?}] — "
        "name: 사람 표시 이름(문자열, 중복 불가). "
        "relation: 공제대상자와의 관계. 허용값: 배우자·직계존속·직계비속·입양자·형제자매·수급자·위탁아동. "
        "문항 표현 그대로 '아들'·'딸'·'자녀'는 '직계비속', '부친'·'모친'·'부모'는 '직계존속'으로 적을 것. "
        "birth_year: 출생연도(정수, 나이 계산용). "
        "income_items: [{kind, amount}] — 부양가족 소득 종류·금액 목록. kind 허용값: 근로_총급여·사업소득금액·이자·배당·공적연금·사적연금·기타소득금액·양도소득금액·퇴직소득금액·전통주_제조소득_농어촌·비과세·공적연금_과세제외. "
        "income_items를 쓰면 income·wage_only_gross는 함께 입력할 수 없음. "
        "disabled: 장애인 여부(bool). cohabiting: 동거 여부(bool, 기본 True, 형제자매·수급자는 동거 필요). "
        "소득 없음 → income_items: []. 근로소득만 있음 → income_items: [{\"kind\": \"근로_총급여\", \"amount\": 총급여}]. "
        "예시: [{\"name\":\"부친\",\"relation\":\"부친\",\"birth_year\":1938,\"income_items\":[{\"kind\":\"공적연금_과세제외\",\"amount\":30000000}]}]"
    ))],
) -> dict:
    """종합소득 인적공제(기본공제·추가공제). 언제: 종합소득세 신고에서 기본공제·추가공제 금액을 구할 때.
    입력: 귀속연도 → year, 납세자 정보 → taxpayer, 부양가족 목록 → persons.
    결과: 기본공제·추가공제·인적공제·사람별 내역·근거. 기본공제 1인당 150만원, 추가공제(경로우대 100만·장애인 200만·부녀자 50만·한부모 100만)."""
    r = I.personal_deduction(year=year, taxpayer=taxpayer, persons=persons)
    if "error" in r:
        return r
    return _ok(r["결과"], "소득세법 제50조·제51조·제53조", 사람별=r["결과"]["사람별"])


# ── 자녀세액공제(소득세법 제59조의2) ────────────────────────────────────────
class ChildEntry(BaseModel):
    model_config = {"extra": "forbid"}
    name: str | None = None
    relation: str = "자녀"
    age: int | None = None
    birth_date: str | None = None
    birth_year: int | None = None
    basic_deduction: bool
    born_or_adopted_this_year: bool = False
    order: int | str | None = None
    school_enrolled: bool | None = None


@mcp.tool(annotations=CALC)
@_guard
def child_tax_credit(
    year: YEAR,
    children: Annotated[list[ChildEntry], Field(description=(
        "[{name?, relation?, age?|birth_date?|birth_year?, basic_deduction, born_or_adopted_this_year?, order?, school_enrolled?}] — "
        "relation: 자녀(기본, 입양자·위탁아동 포함)·손자녀. age: 과세기간 종료일 현재 만 나이(= 귀속연도 − 출생연도), "
        "birth_date(YYYY-MM-DD)·birth_year로 대신 입력 가능. basic_deduction: 이 납세자의 기본공제대상자인지(필수). "
        "born_or_adopted_this_year: 해당 과세기간에 출생·입양 신고(기본 false). order: 출산·입양 자녀의 순서 1·2·3(셋째 이상) 또는 '첫째'·'둘째'·'셋째' "
        "(2017 귀속부터 출산·입양 자녀에 필수). school_enrolled: 2019 귀속 6세 자녀의 취학 여부(그때만 필수). "
        "예: [{\"age\":10,\"basic_deduction\":true},{\"age\":0,\"basic_deduction\":true,\"born_or_adopted_this_year\":true,\"order\":2}]"
    ))],
    settlement_date: Annotated[str | None, Field(description=(
        "신고·결정·연말정산일(YYYY-MM-DD). 2023년 이전 귀속분에 손자녀가 있을 때만 필요 — "
        "손자녀 포함은 2024.1.1 이후 신고·결정·연말정산분부터(2023.12.31 개정 부칙 제3조①)"))] = None,
) -> dict:
    """자녀세액공제(소득세법 제59조의2) 귀속연도별 계산 — 기본공제대상 자녀 수별 금액·6세 이하 추가공제·출산·입양 공제. Child tax credit by tax year (2016~2026).
    언제: 종합소득세 신고·연말정산에서 자녀세액공제 금액을 구할 때(자녀 수·나이·출산·입양·손자녀).
    입력: 귀속연도(2016~2026) → year, 자녀 목록 → children(나이·기본공제대상 여부·출산·입양·순서), 손자녀가 있는 2023년 이전 귀속분만 신고·연말정산일 → settlement_date.
    결과: 자녀세액공제 합계와 ①인원별 금액·②6세 이하 추가(2016~2017)·③출산·입양, 자녀별 판단·산식·근거(개정 부칙 포함).
    연도별: ① 2016~2023 15만·30만·초과 1명당 30만 / 2024 15만·35만·30만 / 2025~ 25만·55만·40만. 연령 2019 7세 이상(7세 미만 취학아동 포함)·2020 7세·2023 8세·2026 9세(2017년생은 13세, 2026.4.21 부칙). 손자녀 2024~(2023 귀속은 2024 이후 연말정산·신고분). ② 6세 이하 추가 2016~2017만. ③ 2016 1명당 30만, 2017~ 첫째 30만·둘째 50만·셋째 이상 70만.
    지원 안 함: 2016 전·2027 이후 귀속, 산출세액 초과분 처리(제61조)·다른 공제와 적용순서, 2022년 이전 귀속 손자녀를 2024년 이후 신고·결정하는 경우(입력 확인 필요로 멈춤). 아동수당 수급 여부는 조문 요건이 아니어서 묻지 않는다."""
    return I.child_tax_credit(year=year, children=[c.model_dump(exclude_none=True) for c in children],
                              settlement_date=settlement_date)


# ── 연금소득금액(소득세법 제20조의3·제14조③9호·제47조의2·제64조의4·제40조) ────────
@mcp.tool(annotations=CALC)
@_guard
def pension_income(
    year: YEAR,
    public_pension: Annotated[int, Field(description="공적연금 과세분(2002년 이후 기여분 기준 금액, 원). public_pension_detail 사용 시 0이어야 함")] = 0,
    private_pension: Annotated[int, Field(description="연금계좌 연금수령분 중 퇴직소득 재원·의료목적 인출 외(원)")] = 0,
    retirement_source_pension: Annotated[int, Field(description="퇴직소득 재원 연금수령액(원, 항상 분리과세)")] = 0,
    medical_withdrawal: Annotated[int, Field(description="의료목적·부득이한 사유 인출액(원, 항상 분리과세)")] = 0,
    private_choice: Annotated[Literal["분리", "종합"], Field(description="사적연금 처리 선택(기본 '분리')")] = "분리",
    public_pension_detail: Annotated[dict | None, Field(description=(
        "{수령액, 환산소득_2002이후, 환산소득_전체, 과세제외기여금?} — "
        "공적연금 과세분 = max(0, int(수령액 × 환산소득_2002이후 ÷ 환산소득_전체) − 과세제외기여금) (소득세법 시행령 제40조①, 국민연금 기준). "
        "과세제외기여금: 연금보험료 소득공제를 받지 않은 기여분(소득세법 시행령 제40조 과세제외기여금), 기본 0. "
        "예: {\"수령액\":45000000, \"환산소득_2002이후\":380000000, \"환산소득_전체\":480000000, \"과세제외기여금\":5625000}"
    ))] = None,
) -> dict:
    """연금소득금액 계산. 언제: 연금소득의 종합과세·분리과세 구분과 연금소득금액을 구할 때.
    입력: 귀속연도 → year, 공적연금 과세분 → public_pension, 사적연금 → private_pension,
          퇴직소득 재원 연금 → retirement_source_pension, 의료목적 인출 → medical_withdrawal,
          사적연금 선택 → private_choice, 공적연금 상세 → public_pension_detail(수령액·환산소득·과세제외기여금).
    결과: 총연금액·분리과세연금소득·연금소득공제·연금소득금액·안내·근거. 공적연금 과세분은 public_pension_detail로 계산하거나 직접 입력."""
    return I.pension_income(year=year, public_pension=public_pension, private_pension=private_pension,
                            retirement_source_pension=retirement_source_pension, medical_withdrawal=medical_withdrawal,
                            private_choice=private_choice, public_pension_detail=public_pension_detail)


@mcp.tool(annotations=CALC)
@_guard
def capital_gains_tax(base: Annotated[int, Field(description="양도소득 과세표준(원)", ge=0)],
                      asset_type: Annotated[Literal["토지", "건물", "주택"], Field(description="양도자산 종류")],
                      acquisition_date: Annotated[str, Field(description="취득일 YYYY-MM-DD(세율 보유기간 계산용)")],
                      transfer_date: Annotated[str, Field(description="양도일 YYYY-MM-DD(2016-01-01 ~ 2026-12-31)")],
                      registered: Annotated[bool, Field(description="등기 여부(False = 미등기 70%)")] = True,
                      non_business_land: Annotated[bool, Field(description="토지만: 비사업용 토지(기본세율+10%p)")] = False,
                      designated_area: Annotated[bool, Field(description="비사업용 토지가 지정지역(제104조의2②)에 있으면 True — 2018.1.1. 이후 양도분 기본세율+20%p(제104조④3호)")] = False,
                      heavy_multi_home: Annotated[bool | None, Field(description="주택만: 조정대상지역 다주택 중과(제104조⑦) 대상 여부(이전 입력). True만 주면 2018.4.1. 전 양도·한시 배제(2022.5.10.~2026.5.9. 양도·보유 2년 이상)만 계산하고, 중과 기간은 houses_count·adjusted_area를 요구")] = None,
                      houses_count: Annotated[int | None, Field(description="주택만: 양도일 현재 1세대 주택 수(조합원입주권·분양권 포함, 분양권은 2021.1.1. 이후 취득분부터). 영 제167조의3①·제167조의10① 괄호로 주택 수에서 빼는 주택은 사용자가 판단해 뺀 수", ge=1)] = None,
                      adjusted_area: Annotated[bool | None, Field(description="주택만: 양도일 현재 조정대상지역 소재 여부")] = None,
                      exclusion_reason: Annotated[str | None, Field(description="중과 배제 사유(사용자 주장) — capital_gains_income의 exclusion_reason과 같은 값. 한시 배제는 도구가 날짜로 판정")] = None,
                      exclusion_basis: Annotated[str | None, Field(description="exclusion_reason='기타'일 때 근거 조문·호와 요건 충족 사실")] = None) -> dict:
    """양도소득 과세표준 → 산출세액·지방소득세(소득세법 제104조①·④·⑦). 언제: 양도소득 과세표준이 이미 주어졌거나 따로 계산했을 때 세율을 적용한다(가액부터 계산하면 capital_gains_income에 apply_basic_deduction=True).
    입력: 과세표준 → base, 자산 종류 → asset_type, 취득일·양도일 → acquisition_date·transfer_date, 등기 → registered, 비사업용 토지 → non_business_land, 지정지역 → designated_area, 다주택 중과(주택) → houses_count·adjusted_area·exclusion_reason(선택).
    세율: 2년 이상 기본세율(제55조) / 1년 이상 2년 미만 40%(주택 60%, 2021.6.1. 전 양도 주택은 기본세율) / 1년 미만 50%(주택 70%, 2021.6.1. 전 40%) / 미등기 70% / 비사업용 토지 기본세율+10%p(지정지역 2018~ +20%p) — 둘 이상이면 큰 세액.
    다주택 중과(제104조⑦, 조정대상지역): 2018.4.1.~2021.5.31. 양도 2주택 +10%p·3주택 이상 +20%p, 2021.6.1.~ +20%p·+30%p, 단기 보유는 제104조①2호·3호 세액과 비교해 큰 세액. 한시 배제(영 제167조의3①·제167조의10① 12호 2019.12.17.~2020.6.30. 보유 10년 이상, 12호의2 2022.5.10.~2026.5.9. 보유 2년 이상)는 날짜로 판정, 그 밖의 배제 사유와 2026.5.10. 이후 나·다목 경과 요건은 사용자 주장(exclusion_reason).
    지원하지 않음: 2018.4.1. 전 양도 지정지역 3주택 이상 가산(제104조④1호·2호), 주택 수 계산(사용자 판단), 분양권·조합원입주권·주식.
    결과: 산출세액·지방소득세(10%)·적용 세율·세율 후보·근거."""
    if not (2016 <= int(transfer_date[:4]) <= 2026):
        return {"error": "입력 확인 필요 — 계산하지 않음", "문제": [f"transfer_date={transfer_date} — 2016~2026 양도분만 지원"]}
    판단: list[dict] = []
    중과사실 = houses_count is not None or adjusted_area is not None or exclusion_reason is not None
    중과 = None
    if heavy_multi_home or 중과사실:
        if asset_type != "주택":
            return {"error": "입력 확인 필요 — 계산하지 않음",
                    "문제": [f"heavy_multi_home·houses_count·adjusted_area·exclusion_reason은 주택에만 입력(asset_type={asset_type})"],
                    "필요 입력": {"heavy_multi_home": "None|False(주택이 아닌 자산)"}}
        if heavy_multi_home and not 중과사실 and designated_area and date.fromisoformat(transfer_date) < date(2018, 4, 1):
            return {"error": "지원하지 않음 — 지정지역 다주택 가산(제104조④1호·2호, 2018.4.1. 전 양도)",
                    "문제": ["heavy_multi_home=True·designated_area=True, 2018.4.1. 전 양도: 지정지역 1세대 3주택 이상 등 +10%p(제104조④1호·2호) — 지원하지 않음"],
                    "지원하지 않음": "지정지역 다주택 가산(제104조④1호·2호)"}
        중과 = CG.다주택_중과_판정(acquisition_date, transfer_date, houses_count=houses_count, adjusted_area=adjusted_area,
                            exclusion_reason=exclusion_reason, exclusion_basis=exclusion_basis,
                            heavy_multi_home=heavy_multi_home, flag_only_ok=True)
        if "error" in 중과:
            return 중과
        판단.extend(중과["판단"])
    적용 = bool(중과 and 중과["적용"])
    r = CG.capital_gains_tax_amount(base, asset_type, acquisition_date, transfer_date, registered, non_business_land,
                                    I.income_tax, designated_area=designated_area,
                                    heavy_surcharge=중과["가산세율"] if 적용 else 0.0, heavy_label=중과["구분"] if 적용 else None)
    근거 = r.pop("근거")
    서식 = r.pop("_서식")
    계산 = {"세율구분_코드": 서식["세율구분_코드"], "세율": 서식["세율"], "누진공제": 서식["누진공제"], "가산세율": 서식["가산세율"],
            "귀속연도": int(transfer_date[:4])}
    계산식_입력 = [
        {"이름": "과세표준", "값": base, "설명": "양도소득 과세표준(입력)"},
        {"이름": "세율", "값": 서식["세율"], "설명": "적용 세율(기본세율이면 제55조① 구간 세율)"},
        {"이름": "누진공제", "값": 서식["누진공제"], "설명": "제55조① 기본세율 누진공제액(단일세율이면 0)"},
        {"이름": "가산세율", "값": 서식["가산세율"], "설명": "기본세율에 더하는 세율(비사업용 토지·다주택 중과, 없으면 0)"},
    ]
    식입력, 식단계 = r.pop("_계산식")
    _이름 = {x["이름"] for x in 계산식_입력}
    계산식_입력 += [x for x in 식입력 if x["이름"] not in _이름]
    if 적용:
        근거 = f"{근거}; {중과['근거']}"
        r["다주택_중과"] = {"적용": True, "가산세율": 중과["가산세율"], "구분": 중과["구분"]}
    elif 중과 is not None and 중과["사실입력"]:
        r["다주택_중과"] = {"적용": False, "가산세율": 0.0, "구분": 중과["구분"]}
    for 주의문 in r.pop("주의", []):
        판단.append({"항목": "지정지역 비사업용 토지", "결론": 주의문, "근거": "소득세법 제104조④3호"})
    if 판단:
        return _ok(r, 근거, 판단=판단, 계산=계산, 계산식_입력=계산식_입력, 계산식=식단계)
    return _ok(r, 근거, 계산=계산, 계산식_입력=계산식_입력, 계산식=식단계)


# ── 주택 임대 보증금 간주임대료(소득세법 제25조①, 시행령 제53조) ─────────────────
def _소형주택_제외기준(year: int) -> tuple[int, int] | None:
    """주택 수에서 빼는 소형주택 기준(면적㎡, 기준시가) — 제25조① 단서 괄호, 연도별."""
    if year <= 2016:
        return (85, 300_000_000)
    if year <= 2018:
        return (60, 300_000_000)
    if year <= 2026:
        return (40, 200_000_000)
    return None


def _주택간주임대료_계산식(rows: list, days: int, rate: float, 금융: int, estimated: bool,
                         적수합: int, 전: int, 산입액: int, 사유: str) -> tuple[list, list]:
    """간주임대료 계산식_입력·계산식 — 3억 차감은 적수가 큰 주택부터 차례로 MIN(보증금, 남은 공제)로(행 순서가 그 판단)."""
    입력 = [{"이름": "공제기준_3억", "값": 300_000_000, "설명": f"보증금 합계에서 빼는 3억원(제25조①, 시행령 제53조③) — 과세 판정: {사유}"}]
    단계 = []
    차감들, 적수들 = [], []
    for k, row in enumerate(rows, 1):
        입력.append({"이름": f"보증금{k}", "값": row["보증금"], "설명": f"{row['주택']} 보증금(적수 큰 순서 {k}번째)"})
        입력.append({"이름": f"일수{k}", "값": row["일수"], "설명": f"{row['주택']} 보증금 받은 일수"})
    for k, row in enumerate(rows, 1):
        남은 = "공제기준_3억" + "".join(f"-차감{j}" for j in range(1, k))
        단계.append({"이름": f"차감{k}", "식": f"MIN(보증금{k},MAX({남은},0))", "값": row["3억_차감"],
                   "설명": f"{row['주택']} — 3억 중 남은 공제액을 적수 큰 주택부터 차감"})
        단계.append({"이름": f"적수{k}", "식": f"(보증금{k}-차감{k})*일수{k}", "값": row["적수"], "설명": f"{row['주택']} (보증금 − 차감) × 일수"})
        차감들.append(f"차감{k}"); 적수들.append(f"적수{k}")
    입력 += [{"이름": "연일수", "값": days, "설명": "그 해 일수(윤년 366)"},
             {"이름": "정기예금이자율", "값": rate, "설명": "시행규칙 제23조 정기예금이자율"},
             {"이름": "금융수익", "값": 금융, "설명": "추계 — 빼지 않음(시행령 제53조④)" if estimated else "임대사업 금융수익(시행령 제53조③1호)"}]
    단계 += [{"이름": "적수_합계", "식": "+".join(적수들) or "0*연일수", "값": 적수합, "설명": "주택별 적수 합계"},
             {"이름": "금융수익_차감_전", "식": "ROUNDDOWN(적수_합계*0.6/연일수*정기예금이자율,0)", "값": 전, "설명": "적수 × 60% ÷ 연일수 × 정기예금이자율"},
             {"이름": "간주임대료", "식": "MAX(0,금융수익_차감_전-금융수익)", "값": 산입액, "설명": "금융수익 차감, 0 미만이면 0 — 총수입금액 산입액"}]
    return 입력, 단계


def _주택간주임대료(year: int, houses: list, financial_income: int = 0, estimated: bool = False,
                   rate: float | None = None) -> dict:
    """housing_deemed_rent_income 본체 — housing_rental_income도 같은 계산을 쓴다."""
    if not houses or not all(isinstance(h, dict) for h in houses):
        return {"error": "입력 확인 필요 — 계산하지 않음", "문제": ["houses는 [{기준시가, 면적, 보증금, 일수?}] 목록"]}
    days_in_year = 366 if calendar.isleap(year) else 365
    r = rate if rate is not None else V.DEEMED_RENT_RATE.get(year)
    if r is None:
        return {"error": "입력 확인 필요 — 계산하지 않음", "문제": [f"{year}년 정기예금이자율 연도표 없음 — rate 입력"]}
    소형 = _소형주택_제외기준(year)
    판단, 대상 = [], []
    for i, h in enumerate(houses):
        nm = h.get("이름") or f"주택{i + 1}"
        price, area, dep = int(h.get("기준시가", 0)), float(h.get("면적", 0) or 0), int(h.get("보증금", 0) or 0)
        days = int(h.get("일수", days_in_year) or days_in_year)
        if 소형 and area and area <= 소형[0] and price <= 소형[1]:
            판단.append({"항목": nm, "결론": f"소형주택 — 주택 수·보증금 합계에서 제외({소형[0]}㎡ 이하·기준시가 {소형[1]:,}원 이하)",
                       "근거": "소득세법 제25조① 단서 괄호"})
            continue
        대상.append({"이름": nm, "기준시가": price, "보증금": dep, "일수": days})
    주택수 = len(대상)
    보증금합계 = sum(h["보증금"] for h in 대상)
    과세, 사유 = False, ""
    if 주택수 >= 3 and 보증금합계 > 300_000_000:
        과세, 사유 = True, f"{주택수}주택, 보증금 합계 {보증금합계:,}원 > 3억 (제25조①1호)"
    elif year >= 2026:
        고가 = [h for h in 대상 if h["기준시가"] > 1_200_000_000]
        고가보증금 = sum(h["보증금"] for h in 고가)
        if len(고가) >= 2 and 고가보증금 > 1_200_000_000:
            과세, 사유 = True, f"기준시가 12억 초과 {len(고가)}주택, 보증금 합계 {고가보증금:,}원 > 12억 (제25조①2호, 시행령 제53조①, 2026 과세기간부터)"
            대상 = 고가
    if not 과세:
        why = f"주택 수 {주택수}(소형 제외), 보증금 합계 {보증금합계:,}원 — 3주택 이상·3억 초과" + (" 또는 기준시가 12억 초과 2주택·12억 초과" if year >= 2026 else "") + " 요건 미충족"
        return _ok({"과세": False, "총수입금액_산입액": 0, "주택_수": 주택수}, "소득세법 제25조①",
                   판단=판단 + [{"항목": "간주임대료 과세 여부", "결론": why, "근거": "소득세법 제25조①"}])
    남은공제 = 300_000_000
    적수합 = 0
    계산 = []
    for h in sorted(대상, key=lambda x: -x["보증금"] * x["일수"]):
        차감 = min(h["보증금"], 남은공제)
        남은공제 -= 차감
        적수 = (h["보증금"] - 차감) * h["일수"]
        적수합 += 적수
        계산.append({"주택": h["이름"], "보증금": h["보증금"], "3억_차감": 차감, "일수": h["일수"], "적수": 적수})
    전 = int(적수합 * 0.6 / days_in_year * r)
    금융 = 0 if estimated else financial_income
    산입액 = max(0, 전 - 금융)
    판단.append({"항목": "간주임대료 과세 여부", "결론": "과세 — " + 사유, "근거": "소득세법 제25조①"})
    판단.append({"항목": "금융수익 차감", "결론": "추계 — 빼지 않음(시행령 제53조④)" if estimated else f"{금융:,}원 차감(시행령 제53조③1호)",
               "근거": "소득세법 시행령 제53조③④"})
    계산식_입력, 계산식 = _주택간주임대료_계산식(계산, days_in_year, r, 금융, estimated, 적수합, 전, 산입액, 사유)
    return _ok({"과세": True, "총수입금액_산입액": 산입액, "주택_수": len(대상), "정기예금이자율": r},
               "소득세법 제25조①, 시행령 제53조③1호·④1호",
               계산={"주택별": 계산, "적수_합계": 적수합, "금융수익_차감_전": 전, "연_일수": days_in_year},
               판단=판단, 계산식_입력=계산식_입력, 계산식=계산식,
               산식=f"(보증금 − 3억, 적수 큰 주택부터)의 적수 {적수합:,} × 60% ÷ {days_in_year} × {r} = {전:,}원 − 금융수익 {금융:,}원 = {산입액:,}원")


@mcp.tool(annotations=CALC)
@_guard
def housing_deemed_rent_income(
    year: YEAR,
    houses: Annotated[list, Field(description=(
        "소유 주택 전부 [{이름?, 기준시가(원), 면적?(㎡, 주거 전용 — 소형주택 판정용, 모르면 생략하면 소형 아님으로 봄), 보증금(원, 임대 안 하면 0), 일수?(보증금 받은 일수, 기본 그 해 일수)}] — "
        "임대하지 않는 주택도 주택 수 판정에 넣는다(부부 합산 등 주택 수 계산은 시행령 제8조의2 — 합산한 목록을 넣을 것)"))],
    financial_income: Annotated[int, Field(description="임대사업 부분에서 생긴 수입이자·할인료·배당금(장부로 보증금 운용이 확인되는 것, 원). 추계면 무시", ge=0)] = 0,
    estimated: Annotated[bool, Field(description="추계신고·결정이면 True — 금융수익을 빼지 않음(시행령 제53조④)")] = False,
    rate: Annotated[float | None, Field(description="정기예금이자율(소수). 생략 시 연도표(소득세법 시행규칙 제23조, 부가·법인 간주임대료와 같은 율)")] = None,
) -> dict:
    """주택 임대 보증금 간주임대료 — 개인 사업소득 총수입금액 산입액(소득세법 제25조①, 시행령 제53조③④).
    언제: 개인이 주택을 전세·보증금으로 임대할 때 종합소득세(사업소득) 간주임대료를 계산한다. 부가가치세 간주임대료는 vat_deemed_rent, 법인은 corporate_deemed_rent_income.
    입력: 귀속연도 → year, 소유 주택 전부(기준시가·면적·보증금·일수) → houses, 임대 금융수익 → financial_income, 추계 여부 → estimated, 정기예금이자율 → rate(생략 시 연도표).
    판정(연도별): 3주택 이상 + 보증금 합계 3억 초과(전 기간) / 2026~: 2주택(기준시가 12억 이하 주택은 주택 수 제외) + 보증금 합계 12억 초과.
      소형주택(면적·기준시가 기준 이하)은 주택 수에서 뺀다 — 2016 85㎡·3억, 2017~2018 60㎡·3억, 2019~2026 40㎡·2억.
    산식: (보증금 − 3억, 적수가 큰 주택부터 차감)의 적수 × 60% × 1/365(윤년 366) × 정기예금이자율 − 금융수익(추계면 빼지 않음). 0보다 작으면 0.
    결과: 과세 여부·주택 수(제외 사유)·총수입금액 산입액·주택별 적수·근거."""
    return _주택간주임대료(year, houses, financial_income, estimated, rate)


# ── 개인 주택임대소득: 1주택 비과세·소액 비과세(~2018)·분리과세 선택(2019~) ─────────────
# 연도표 — 근거 조문과 부칙(적용 시점):
#  · 1주택 비과세 고가주택 기준(제12조2호나목 괄호): 기준시가 9억 초과 제외 → 12억 초과 제외.
#    2022.12.31. 개정(법률 제19196호, 2023.1.1. 시행) 부칙 제11조 "이 법 시행 전에 발생한 주택임대소득의 비과세기준에 관하여는
#    제12조제2호나목의 개정규정에도 불구하고 종전의 규정에 따른다" → 2022 귀속까지 9억, 2023 귀속부터 12억.
#    판정 시점: 과세기간 종료일 또는 양도일(시행령 제8조의2⑤).
#  · 총수입금액 2천만원 이하 비과세(제12조2호나목 후단): 2014.12.23. 개정 때 "2016년 12월 31일 이전에 끝나는 과세기간까지",
#    2016.12.20. 개정으로 "2018년 12월 31일 이전에 끝나는 과세기간까지" → 2014~2018 귀속 비과세(이 도구는 2016~2018).
#  · 분리과세 선택(제14조③7호, 제64조의2): 법률 제12852호(2014.12.23.) 부칙 제1조 단서·제4조 "2019년 1월 1일 이후 발생하는
#    소득분부터 적용"(2016.12.20. 개정) → 2019 귀속부터. 세율 14%, 필요경비율 50%(등록임대 60%), 다른 종합소득금액 2천만원
#    이하면 200만원(등록임대 400만원) 추가 차감 — 2019~2026 시행본 같은 문구.
#  · 등록임대주택 범위(시행령 제122조의2①): 2019 시행본은 공공지원·장기일반·단기 모두. 2020.10.7. 개정으로 아파트 매입
#    장기일반·단기는 2020.7.10. 이전 등록 신청분만, 2020.7.11. 이후 단기→장기 변경 신고분 제외(같은 조②).
#    이 개정의 적용 부칙은 조회 도구에서 확인하지 못함 → 2021 귀속부터 적용하고, 2020 귀속에서 결과가 갈리는 입력은 멈춤.
#  · 의무임대기간(제64조의2③2호): 4년 → 10년(2020.12.29. 개정) — 부칙 제13조② "2020년 8월 18일 이후 등록을 신청하는
#    민간임대주택부터 적용", 제36조 그 전 신청분은 종전(4년). ③ 단서(자동말소 등 추징 제외)는 부칙 제13조① 2020.8.18. 이후 말소분.
_주택임대_고가기준 = {**{y: 900_000_000 for y in range(2016, 2023)}, **{y: 1_200_000_000 for y in range(2023, 2027)}}


def _등록임대_판정(year: int, h: dict, nm: str) -> tuple[bool | None, str]:
    """등록임대주택(시행령 제122조의2①②) 해당 여부. (None, 사유)면 입력 확인 필요."""
    유형 = h.get("임대유형")
    if 유형 not in ("공공지원", "장기일반", "단기"):
        return None, f"{nm}: 등록임대=True면 임대유형(공공지원|장기일반|단기) 입력"
    신청일 = h.get("등록신청일")
    변경일 = h.get("단기에서_변경신고일")
    if year == 2019:
        return True, f"{유형}민간임대주택 — 2019 시행본은 공공지원·장기일반·단기 모두 등록임대주택(시행령 제122조의2①1호)"
    제한대상 = 유형 == "단기" or (유형 == "장기일반" and h.get("아파트"))
    if 제한대상 and not 신청일:
        return None, f"{nm}: {'단기' if 유형 == '단기' else '아파트 장기일반(매입)'} 임대 — 등록신청일(YYYY-MM-DD) 입력(2020.7.10. 이전 신청분만 등록임대주택)"
    늦은신청 = bool(제한대상 and date.fromisoformat(신청일) > date(2020, 7, 10))
    늦은변경 = bool(변경일 and date.fromisoformat(변경일) >= date(2020, 7, 11))
    if year == 2020 and (늦은신청 or 늦은변경):
        return None, (f"{nm}: 2020 귀속 — 2020.10.7. 개정 시행령 제122조의2①1호·②(2020.7.10. 이전 신청·2020.7.11. 이후 변경 신고 제외)의 "
                      "적용 시점을 확인하지 못해 계산하지 않음")
    if 늦은신청:
        return False, f"{유형} 등록 신청 {신청일} — 2020.7.10. 이전 신청분만 등록임대주택(시행령 제122조의2①1호 나·다목)"
    if 늦은변경:
        return False, f"단기→장기 변경 신고 {변경일} — 2020.7.11. 이후 변경 신고분은 등록임대주택 제외(시행령 제122조의2②)"
    return True, f"{유형}민간임대주택 — 등록임대주택(시행령 제122조의2①1호)"


@mcp.tool(annotations=CALC)
@_guard
def housing_rental_income(
    year: YEAR,
    houses: Annotated[list, Field(description=(
        "본인과 배우자가 소유한 주택 전부(주택 수는 부부 합산, 시행령 제8조의2③4호) — "
        "[{이름?, 소유자?('본인'|'배우자', 기본 본인), 소재?('국내'|'국외', 기본 국내), 기준시가(원, 국내 주택), 면적?(㎡, 주거 전용 — 보증금 있으면 필요), "
        "월세?(그 해 본인 귀속 월세·사용료 합계, 원), 보증금?(원), 일수?(보증금 받은 일수), "
        "등록임대?(bool), 임대유형?('공공지원'|'장기일반'|'단기'), 아파트?(bool), 등록신청일?(YYYY-MM-DD), 등록_월수?(그 해 등록임대로 임대한 월수 1~12, 기본 12), "
        "단기에서_변경신고일?(YYYY-MM-DD), 공동소유?(bool), 주택수_산입?(공동소유일 때 본인 주택 수에 넣는지 bool), 겸용?(bool)}]"))],
    other_income_amount: Annotated[int | None, Field(description="분리과세 주택임대소득을 뺀 그 해 종합소득금액(원) — 2019~ 분리과세 추가 차감(200만·400만) 요건 판정용", ge=0)] = None,
    other_taxable_base: Annotated[int | None, Field(description="종합과세 비교용: 주택임대소득을 뺀 종합소득 과세표준(원)", ge=0)] = None,
    comprehensive_rental_income: Annotated[int | None, Field(description="종합과세 비교용: 종합과세할 때의 주택임대 사업소득금액(장부 또는 추계, 원)", ge=0)] = None,
    financial_income: Annotated[int, Field(description="간주임대료에서 뺄 임대사업 금융수익(원, 시행령 제53조③)", ge=0)] = 0,
    estimated: Annotated[bool, Field(description="추계신고·결정이면 True — 간주임대료에서 금융수익을 빼지 않음")] = False,
    rate: Annotated[float | None, Field(description="간주임대료 정기예금이자율(소수). 생략 시 연도표")] = None,
) -> dict:
    """개인 주택임대소득 과세 판정 — 1주택 비과세, 총수입금액 2천만원 이하 비과세(~2018)·분리과세 선택(2019~), 분리과세 세액과 종합과세 참고값(소득세법 제12조2호나목, 제14조③7호, 제64조의2).
    언제: 개인이 주택을 월세·전세로 임대할 때 과세 여부와 분리과세(14%) 세액을 구하고, 분리과세와 종합과세 중 무엇을 고를지 비교한다. 간주임대료만 필요하면 housing_deemed_rent_income.
    입력: 귀속연도 → year(2016~2026), 부부 합산 소유 주택 전부(월세·보증금·기준시가·등록임대 여부) → houses, 다른 종합소득금액 → other_income_amount(2019~ 분리과세 계산 시 필요),
          종합과세 비교용 다른 과세표준 → other_taxable_base, 종합과세 시 주택임대 소득금액 → comprehensive_rental_income, 간주임대료 금융수익·추계·이자율 → financial_income·estimated·rate.
    판정(연도별): 주택 수 1개(부부 합산·국외 포함)이고 국내·기준시가 기준 이하면 전액 비과세 — 기준 2016~2022 9억, 2023~ 12억(2022.12.31. 개정 부칙 제11조).
      총수입금액(월세 + 간주임대료, 시행령 제8조의2⑥) 2천만원 이하: 2016~2018 비과세 / 2019~ 분리과세 선택 가능. 초과면 종합과세.
    분리과세 산식: (총수입금액 − 필요경비 50%, 등록임대 60% − 다른 종합소득금액 2천만원 이하면 200만원, 등록임대 400만원 — 섞이면 수입금액 비율 안분) × 14%.
      등록임대: 2019 공공지원·장기일반·단기 / 2021~ 아파트 매입 장기일반·단기는 2020.7.10. 이전 등록 신청분만. 의무임대기간 4년, 2020.8.18. 이후 등록 신청분 10년.
    지원하지 않음(입력 확인 필요로 멈춤): 겸용주택 주택 부분 판정(제8조의2④), 공동소유 주택의 주택 수 판정(제8조의2③2호 — 주택수_산입을 직접 입력하면 계산),
      국외주택 보증금 간주임대료, 조특법 제96조 소형주택 임대 감면, 2020 귀속 중 2020.10.7. 개정 시행령으로 결과가 갈리는 등록임대, 2016 전·2027 이후 연도.
    결과: 과세 여부·구분(비과세|분리과세 선택 가능|종합과세)·주택 수·총수입금액(월세·간주임대료)·분리과세 세액·종합과세 참고 세액·의무임대기간·근거."""
    if not houses or not all(isinstance(h, dict) for h in houses):
        return {"error": "입력 확인 필요 — 계산하지 않음", "문제": ["houses는 [{기준시가, 월세?, 보증금?, ...}] 목록"]}
    문제: list[str] = []
    판단: list[dict] = []
    정리 = []
    for i, h in enumerate(houses):
        nm = h.get("이름") or f"주택{i + 1}"
        소유자 = h.get("소유자", "본인")
        소재 = h.get("소재", "국내")
        if 소유자 not in ("본인", "배우자"):
            문제.append(f"{nm}: 소유자는 '본인' 또는 '배우자'")
        if 소재 not in ("국내", "국외"):
            문제.append(f"{nm}: 소재는 '국내' 또는 '국외'")
        if h.get("겸용"):
            문제.append(f"{nm}: 겸용주택 — 주택 부분 판정(시행령 제8조의2④: 주택 면적 > 사업용 면적이면 전부 주택)을 지원하지 않음. "
                      "판정한 주택 부분만 넣고 겸용=False로 다시 호출")
        if h.get("공동소유") and h.get("주택수_산입") is False and int(h.get("보증금", 0) or 0):
            문제.append(f"{nm}: 주택 수에 넣지 않는 공동소유 주택의 보증금 간주임대료는 지원하지 않음")
        if h.get("공동소유") and not isinstance(h.get("주택수_산입"), bool):
            문제.append(f"{nm}: 공동소유 주택 — 주택 수 판정(시행령 제8조의2③2호: 최대 지분자, 2020~ 수입 600만원 이상·고가주택 30% 초과 지분자 등)을 "
                      "지원하지 않음. 본인 주택 수에 넣는지 주택수_산입(true/false)과 본인 귀속 월세를 넣을 것")
        보증금 = int(h.get("보증금", 0) or 0)
        if 소재 == "국외" and 보증금 and 소유자 == "본인":
            문제.append(f"{nm}: 국외주택 보증금 간주임대료는 지원하지 않음")
        if 소재 == "국내" and h.get("기준시가") is None:
            문제.append(f"{nm}: 국내 주택은 기준시가(원) 입력 — 1주택 비과세 고가주택·소형주택 판정에 필요")
        if 소재 == "국내" and 보증금 and 소유자 == "본인" and not h.get("면적"):
            문제.append(f"{nm}: 보증금이 있으면 면적(㎡) 입력 — 간주임대료 소형주택 제외 판정(제25조①)에 필요")
        월수 = int(h.get("등록_월수", 12) or 12)
        if not (1 <= 월수 <= 12):
            문제.append(f"{nm}: 등록_월수는 1~12")
        정리.append({"이름": nm, "소유자": 소유자, "소재": 소재, "기준시가": int(h.get("기준시가", 0) or 0),
                   "면적": float(h.get("면적", 0) or 0), "월세": int(h.get("월세", 0) or 0), "보증금": 보증금,
                   "일수": h.get("일수"), "산입": (h.get("주택수_산입") if h.get("공동소유") else True),
                   "등록임대": bool(h.get("등록임대")), "등록_월수": 월수, "원본": h})
    if 문제:
        return {"error": "입력 확인 필요 — 계산하지 않음", "문제": 문제}
    본인 = [h for h in 정리 if h["소유자"] == "본인"]
    if not 본인:
        return {"error": "입력 확인 필요 — 계산하지 않음", "문제": ["본인 소유 주택이 없음 — 소유자='본인'인 임대 주택을 넣을 것"]}

    # 1) 주택 수(부부 합산, 국외 포함 — 제12조2호나목·시행령 제8조의2③; 소형주택 제외는 간주임대료에만)
    산입주택 = [h for h in 정리 if h["산입"]]
    주택수 = len(산입주택)
    고가기준 = _주택임대_고가기준[year]
    판단.append({"항목": "주택 수", "결론": f"{주택수}채(본인·배우자 합산, 국외주택 포함" +
                (", 공동소유 중 주택수_산입=False 제외" if len(산입주택) < len(정리) else "") + ")",
                "근거": "소득세법 제12조2호나목, 시행령 제8조의2③"})
    if 주택수 == 1 and 산입주택[0]["소유자"] == "본인":
        h = 산입주택[0]
        if h["소재"] == "국내" and h["기준시가"] <= 고가기준:
            판단.append({"항목": "1주택 비과세", "결론": f"1주택·국내·기준시가 {h['기준시가']:,}원 ≤ {고가기준:,}원 — 주택임대소득 전액 비과세",
                         "근거": "소득세법 제12조2호나목, 시행령 제8조의2⑤" + (" (2022.12.31. 개정 부칙 제11조: 2023 귀속부터 12억)" if year >= 2023 else "")})
            return _ok({"과세": False, "구분": "비과세(1주택)", "주택_수": 1, "총수입금액": 0, "비과세_수입금액": h["월세"]},
                       "소득세법 제12조2호나목", 판단=판단)
        사유 = "국외주택" if h["소재"] == "국외" else f"기준시가 {h['기준시가']:,}원 > {고가기준:,}원(고가주택)"
        판단.append({"항목": "1주택 비과세", "결론": f"1주택이지만 {사유} — 비과세 제외", "근거": "소득세법 제12조2호나목 괄호"})

    # 2) 총수입금액 = 월세 + 간주임대료(배우자 주택은 주택 수에만 — 보증금 0으로 넘김)
    간주입력 = [{"이름": h["이름"], "기준시가": h["기준시가"], "면적": h["면적"],
                "보증금": h["보증금"] if (h["소유자"] == "본인" and h["소재"] == "국내") else 0,
                **({"일수": int(h["일수"])} if h["일수"] else {})} for h in 산입주택]
    간주 = _주택간주임대료(year, 간주입력, financial_income, estimated, rate)
    if "error" in 간주:
        return 간주
    간주액 = 간주["결과"]["총수입금액_산입액"]
    간주별: dict[str, int] = {}
    if 간주액:
        적수합 = 간주["계산"]["적수_합계"]
        for row in 간주["계산"]["주택별"]:
            간주별[row["주택"]] = int(간주액 * row["적수"] / 적수합) if 적수합 else 0
        # 원 단위 절사 차이는 적수가 가장 큰 주택에 붙인다
        if 간주별:
            k = max(간주별, key=lambda n: next(r["적수"] for r in 간주["계산"]["주택별"] if r["주택"] == n))
            간주별[k] += 간주액 - sum(간주별.values())
    판단 += [{**d, "항목": "간주임대료 — " + d["항목"]} for d in 간주.get("판단", [])]
    주택별 = []
    for h in 본인:
        수입 = h["월세"] + 간주별.get(h["이름"], 0)
        주택별.append({"주택": h["이름"], "월세": h["월세"], "간주임대료": 간주별.get(h["이름"], 0), "수입금액": 수입, "_h": h})
    월세합 = sum(r["월세"] for r in 주택별)
    총수입 = 월세합 + 간주액
    결과: dict[str, Any] = {"주택_수": 주택수, "총수입금액": 총수입, "월세_합계": 월세합, "간주임대료": 간주액}
    계산 = {"주택별": [{k: v for k, v in r.items() if k != "_h"} for r in 주택별]}

    # ── 계산식 단계: 월세(주택별) + 간주임대료(적수 단계) → 총수입금액 → (분리과세) ──
    식입력: list[dict] = []
    식단계: list[dict] = []
    월세이름 = {}
    for k, r in enumerate(주택별, 1):
        식입력.append({"이름": f"월세{k}", "값": r["월세"], "설명": f"{r['주택']} 그 해 본인 귀속 월세·사용료"})
        월세이름[r["주택"]] = f"월세{k}"
    식단계.append({"이름": "월세_합계", "식": "+".join(월세이름.values()), "값": 월세합, "설명": "주택별 월세 합계"})
    if 간주.get("계산식"):
        식입력 += 간주["계산식_입력"]
        식단계 += 간주["계산식"]
    else:
        식입력.append({"이름": "간주임대료", "값": 간주액, "설명": "간주임대료 과세 요건 미충족(제25조①) — 0"})
    식단계.append({"이름": "총수입금액", "식": "월세_합계+간주임대료", "값": 총수입, "설명": "월세 + 간주임대료(시행령 제8조의2⑥)"})

    def 종합참고식() -> None:
        if "종합과세_주택임대분_세액_참고" not in 결과:
            return
        식입력.extend([{"이름": "다른과세표준", "값": other_taxable_base, "설명": "주택임대소득을 뺀 종합소득 과세표준(입력)"},
                     {"이름": "종합_주택임대소득금액", "값": comprehensive_rental_income, "설명": "종합과세 때 주택임대 사업소득금액(입력)"}])
        식단계.extend([
            {"이름": "종합_합산과세표준", "식": "다른과세표준+종합_주택임대소득금액", "값": other_taxable_base + comprehensive_rental_income, "설명": "합산 과세표준"},
            {"이름": "종합_합산세액", "식": I.income_tax_expr("종합_합산과세표준", year),
             "값": I.income_tax(other_taxable_base + comprehensive_rental_income, year), "설명": "기본세율(제55조①)"},
            {"이름": "종합_기존세액", "식": I.income_tax_expr("다른과세표준", year), "값": I.income_tax(other_taxable_base, year), "설명": "기본세율(제55조①)"},
            {"이름": "종합과세_주택임대분_세액_참고", "식": "종합_합산세액-종합_기존세액", "값": 결과["종합과세_주택임대분_세액_참고"],
             "설명": "종합과세 증가 세액(참고 — 소득공제·세액공제 차이 제외)"}])

    def 종합참고(결과: dict) -> None:
        if other_taxable_base is not None and comprehensive_rental_income is not None:
            종합 = I.income_tax(other_taxable_base + comprehensive_rental_income, year) - I.income_tax(other_taxable_base, year)
            결과["종합과세_주택임대분_세액_참고"] = 종합
            계산["종합과세_참고"] = (f"기본세율(과세표준 {other_taxable_base:,} + 주택임대 소득금액 {comprehensive_rental_income:,}) − "
                                 f"기본세율({other_taxable_base:,}) = {종합:,}원 — 소득공제·세액공제 차이는 반영하지 않은 참고값")

    # 3) 2천만원 기준 판정
    if 총수입 > 20_000_000:
        판단.append({"항목": "과세 방법", "결론": f"총수입금액 {총수입:,}원 > 2천만원 — 종합과세", "근거": "소득세법 제12조2호나목·제14조③7호"})
        결과.update({"과세": True, "구분": "종합과세"})
        종합참고(결과)
        종합참고식()
        return _ok(결과, "소득세법 제12조2호나목, 제14조③7호, 제25조①", 계산=계산, 판단=판단, 계산식_입력=식입력, 계산식=식단계)
    if year <= 2018:
        판단.append({"항목": "소액 비과세", "결론": f"총수입금액 {총수입:,}원 ≤ 2천만원 — 2018 귀속까지 비과세",
                     "근거": "소득세법 제12조2호나목 후단(2016.12.20. 개정: 2018.12.31. 이전에 끝나는 과세기간까지), 시행령 제8조의2⑥"})
        결과.update({"과세": False, "구분": "비과세(총수입금액 2천만원 이하)", "비과세_수입금액": 총수입})
        return _ok(결과, "소득세법 제12조2호나목", 계산=계산, 판단=판단, 계산식_입력=식입력, 계산식=식단계)

    # 4) 2019~ 분리과세 선택 가능(제14조③7호, 제64조의2; 법률 제12852호 부칙 제4조: 2019.1.1. 이후 발생 소득분부터)
    if other_income_amount is None:
        return {"error": "입력 확인 필요 — 계산하지 않음",
                "문제": [f"총수입금액 {총수입:,}원 ≤ 2천만원 — 분리과세 세액 계산에 other_income_amount(분리과세 주택임대소득을 뺀 종합소득금액) 필요"],
                "필요 입력": {"other_income_amount": "원"}}
    등록수입 = 0
    의무 = []
    for r in 주택별:
        h = r["_h"]
        if not h["등록임대"]:
            continue
        ok, why = _등록임대_판정(year, h["원본"], h["이름"])
        if ok is None:
            return {"error": "입력 확인 필요 — 계산하지 않음", "문제": [why]}
        판단.append({"항목": f"등록임대 — {h['이름']}", "결론": why + ("" if ok else " → 필요경비율 50%"),
                     "근거": "소득세법 시행령 제122조의2①②"})
        if not ok:
            continue
        r["등록기간_수입금액"] = r["수입금액"] * h["등록_월수"] // 12
        등록수입 += r["등록기간_수입금액"]
        신청일 = h["원본"].get("등록신청일")
        if 신청일:
            기간 = 10 if date.fromisoformat(신청일) >= date(2020, 8, 18) else 4
            의무.append({"주택": h["이름"], "등록신청일": 신청일, "의무임대기간_년": 기간})
        else:
            의무.append({"주택": h["이름"], "등록신청일": None, "의무임대기간_년": "2020.8.18. 전 신청 4년 / 이후 신청 10년 — 등록신청일 입력 시 확정"})
    일반수입 = 총수입 - 등록수입
    경비 = int(등록수입 * 0.6) + int(일반수입 * 0.5)
    공제 = 0
    if other_income_amount <= 20_000_000 and 총수입:
        공제 = int(4_000_000 * 등록수입 / 총수입) + int(2_000_000 * 일반수입 / 총수입)
    소득금액 = max(0, 총수입 - 경비 - 공제)
    세액 = int(소득금액 * 0.14)
    등록이름 = []
    for k, r in enumerate(주택별, 1):
        if "등록기간_수입금액" not in r:
            continue
        식입력 += [{"이름": f"간주배분{k}", "값": 간주별.get(r["주택"], 0),
                   "설명": f"{r['주택']} 간주임대료 배분액(적수 비율 안분, 원 미만 끝수는 적수 큰 주택)"},
                  {"이름": f"등록월수{k}", "값": r["_h"]["등록_월수"], "설명": f"{r['주택']} 그 해 등록임대로 임대한 월수"}]
        식단계.append({"이름": f"등록수입{k}", "식": f"ROUNDDOWN((월세{k}+간주배분{k})*등록월수{k}/12,0)", "값": r["등록기간_수입금액"],
                     "설명": f"{r['주택']} 등록임대 기간 수입금액(시행령 제122조의2 안분)"})
        등록이름.append(f"등록수입{k}")
    if not 등록이름:
        식입력.append({"이름": "등록임대없음", "값": 0, "설명": "등록임대주택(시행령 제122조의2①) 수입금액 없음"})
        등록이름.append("등록임대없음")
    식입력.append({"이름": "다른종합소득금액", "값": other_income_amount, "설명": "분리과세 주택임대소득을 뺀 종합소득금액(제64조의2② 추가 차감 요건)"})
    식단계 += [
        {"이름": "등록임대_수입금액", "식": "+".join(등록이름), "값": 등록수입, "설명": "등록임대주택 수입금액"},
        {"이름": "일반_수입금액", "식": "총수입금액-등록임대_수입금액", "값": 일반수입, "설명": "그 밖의 주택 수입금액"},
        {"이름": "필요경비", "식": "ROUNDDOWN(등록임대_수입금액*0.6,0)+ROUNDDOWN(일반_수입금액*0.5,0)", "값": 경비,
         "설명": "필요경비율 등록임대 60%·그 밖 50%(제64조의2②)"},
        {"이름": "추가차감", "식": "IF(다른종합소득금액<=20000000,ROUNDDOWN(4000000*등록임대_수입금액/MAX(총수입금액,1),0)"
                              "+ROUNDDOWN(2000000*일반_수입금액/MAX(총수입금액,1),0),0)", "값": 공제,
         "설명": "다른 종합소득금액 2천만원 이하면 400만(등록)·200만(일반)을 수입금액 비율로 안분"},
        {"이름": "분리과세_사업소득금액", "식": "MAX(0,총수입금액-필요경비-추가차감)", "값": 소득금액, "설명": "총수입금액 − 필요경비 − 추가 차감"},
        {"이름": "분리과세_세액", "식": "ROUNDDOWN(분리과세_사업소득금액*0.14,0)", "값": 세액, "설명": "14%(제64조의2①)"},
        {"이름": "분리과세_지방소득세", "식": "ROUNDDOWN(분리과세_세액/10,0)", "값": 세액 // 10, "설명": "소득세의 10%"},
    ]
    판단.append({"항목": "과세 방법", "결론": f"총수입금액 {총수입:,}원 ≤ 2천만원 — 분리과세(14%)와 종합과세 중 선택",
                 "근거": "소득세법 제14조③7호, 제64조의2①, 법률 제12852호 부칙 제4조(2019.1.1. 이후 발생 소득분부터)"})
    판단.append({"항목": "추가 차감", "결론": (f"다른 종합소득금액 {other_income_amount:,}원 ≤ 2천만원 — {공제:,}원 차감(등록·일반 수입금액 비율 안분)"
                                          if other_income_amount <= 20_000_000 else f"다른 종합소득금액 {other_income_amount:,}원 > 2천만원 — 추가 차감 없음"),
                 "근거": "소득세법 제64조의2②, 시행령 제122조의2(등록임대주택과 그 밖의 주택 수입금액 안분 계산식)"})
    if 의무:
        판단.append({"항목": "의무임대기간", "결론": "등록임대 60%·400만원을 적용받으면 의무임대기간(2020.8.18. 이후 등록 신청분 10년, 그 전 4년) 안에 임대하지 않을 때 차액 추징",
                     "근거": "소득세법 제64조의2③2호, 법률 제17757호 부칙 제13조②·제36조"})
    결과.update({"과세": True, "구분": "분리과세 선택 가능", "분리과세_사업소득금액": 소득금액, "분리과세_세액": 세액,
               "분리과세_지방소득세": 세액 // 10, "등록임대_수입금액": 등록수입})
    if 의무:
        결과["의무임대기간"] = 의무
    계산["주택별"] = [{k: v for k, v in r.items() if k != "_h"} for r in 주택별]
    계산["분리과세"] = (f"총수입금액 {총수입:,} − 필요경비(등록 {등록수입:,} × 60% + 일반 {일반수입:,} × 50% = {경비:,}) − 추가 차감 {공제:,} "
                     f"= {소득금액:,}원 × 14% = {세액:,}원")
    종합참고(결과)
    if "종합과세_주택임대분_세액_참고" in 결과:
        종합 = 결과["종합과세_주택임대분_세액_참고"]
        판단.append({"항목": "분리과세 vs 종합과세(참고)", "결론": f"분리과세 {세액:,}원 vs 종합과세 증가분 {종합:,}원 — " +
                     ("분리과세가 적음" if 세액 < 종합 else "종합과세가 적음" if 종합 < 세액 else "같음") + " (공제·감면 차이 제외)",
                     "근거": "소득세법 제64조의2①"})
    종합참고식()
    return _ok(결과, "소득세법 제14조③7호, 제64조의2①②, 시행령 제122조의2", 계산=계산, 판단=판단,
               산식=계산["분리과세"], 계산식_입력=식입력, 계산식=식단계)


# ── 결손금·이월결손금 공제(소득세법 제45조) ──────────────────────────────────────
@mcp.tool(annotations=CALC)
@_guard
def income_loss_offset(
    year: YEAR,
    incomes: Annotated[dict, Field(description=(
        "{부동산임대, 주거용임대, 근로, 연금, 기타, 이자, 배당} — 각 소득금액(원), 음수=결손금. "
        "부동산임대 = 주거용 건물 임대업을 뺀 부동산임대업(결손금은 부동산임대 소득에서만 공제, 제45조②). "
        "주거용임대 = 주거용 건물 임대업(결손금도 다른 종합소득에서 공제, 제45조② 단서) — 임대 결손금은 주거용분을 꼭 나눠 넣는다. "
        "사업소득은 business로 입력(사업소득은 incomes['사업']으로 넣지 않음)."
    ))],
    carryforwards: Annotated[list | None, Field(description=(
        "[{year, kind, amount}] — 이월결손금 목록. kind: '사업'|'부동산임대'. "
        "예: [{\"year\":2022, \"kind\":\"사업\", \"amount\":10000000}]"
    ))] = None,
    estimated: Annotated[bool, Field(description="추계 과세 여부(True면 그 연도에는 이월결손금을 공제하지 않고 소멸 없이 다음 연도로 넘김, 제45조④)")] = False,
    financial_flat_rate: Annotated[int, Field(description="이자·배당 중 원천징수세율 적용분 합계(원, 이자에서 먼저 차감)")] = 0,
    business: Annotated[dict | None, Field(description=(
        "{book_income, adjustments} — 사업소득 계산용. book_income: 장부상 사업 손익(원, 손실은 음수). "
        "adjustments: [{kind, amount}]. kind: 사업주_본인_급여·가사관련경비·기타_가산·기타_차감. "
        "예: {\"book_income\":-90000000, \"adjustments\":[{\"kind\":\"사업주_본인_급여\",\"amount\":30000000}]}"
    ))] = None,
    history: Annotated[list | None, Field(description=(
        "[{year, incomes, business?, financial_flat_rate?, estimated?}] — 대상 연도보다 앞선 연도들. "
        "year 오름차순으로 계산 후 이월결손금을 다음 연도로 넘김. "
        "history와 carryforwards 중 하나만 사용. "
        "예: [{\"year\":2022, \"incomes\":{\"근로\":15000000}, \"business\":{\"book_income\":-30000000, \"adjustments\":[]}}]"
    ))] = None,
) -> dict:
    """결손금·이월결손금 공제 후 종합소득금액. 언제: 사업소득 결손금이나 이월결손금을 다른 종합소득에서 공제할 때.
    입력: 귀속연도 → year, 소득 목록 → incomes(사업소득 제외), 이월결손금 → carryforwards(선택),
          추계 여부 → estimated, 금융소득 분리과세분 → financial_flat_rate, 사업소득 상세 → business(선택),
          이전 연도 이력 → history(선택, carryforwards와 상호 배제).
    결과: 공제 후 소득금액·종합소득금액·이월결손금(다음 연도)·공제 내역·소멸·이력_계산·근거."""
    return K.income_loss_offset(year=year, incomes=incomes, carryforwards=carryforwards,
                                estimated=estimated, financial_flat_rate=financial_flat_rate,
                                business=business, history=history)


# ── 법인 재해손실세액공제(법인세법 제58조, 시행령 제95조) ────────────────────────
@mcp.tool(annotations=CALC)
@_guard
def disaster_loss_credit(
    year: YEAR,
    assets_before: Annotated[int, Field(description="재해 전 자산총액(토지 제외 사업용자산 + 변상책임 있는 타인 자산, 장부가액, 원) — assets 사용 시 생략")] = 0,
    assets_lost: Annotated[int, Field(description="상실 자산가액(토지 제외, 원) — assets 사용 시 생략")] = 0,
    current_tax: Annotated[int, Field(description="재해 발생 사업연도 법인세 산출세액(원, 공제·감면 전)")] = 0,
    other_credits: Annotated[int, Field(description="공제대상 법인세에서 차감할 다른 감면·공제세액(원, 기존 호환용 — credit_items와 함께 입력 불가)")] = 0,
    unpaid_prior_tax: Annotated[int, Field(description="재해일 현재 미납·미부과 법인세 본세(원) — 가산세는 넣지 말고 penalty_items로 따로(두 번 들어감)")] = 0,
    penalty: Annotated[int, Field(description="가산세(원, 기존 호환용 — penalty_items와 함께 입력 불가)")] = 0,
    assets: Annotated[list | None, Field(description=(
        "[{구분, 장부가액, 상실액}] — assets_before/assets_lost 대신 사용. "
        "구분: '사업용'(토지 제외)|'토지'|'타인소유_변상책임있음'|'타인소유_변상책임없음'. "
        "예: [{\"구분\":\"사업용\",\"장부가액\":500000000,\"상실액\":140000000}, {\"구분\":\"타인소유_변상책임있음\",\"장부가액\":100000000,\"상실액\":100000000}]"
    ))] = None,
    credit_items: Annotated[list | None, Field(description=(
        "[{종류, 금액}] — 항목형 공제 입력. 종류: '조세특례제한법'(current_tax에서 차감)|'법인세법'(공제대상 법인세에서 빼지 않음). "
        "예: [{\"종류\":\"조세특례제한법\",\"금액\":30000000}]"
    ))] = None,
    penalty_items: Annotated[list | None, Field(description=(
        "[{근거, 금액}] — 항목형 가산세 입력. 근거: '법인세법 제75조의3'|'국세기본법 제47조의2'|'국세기본법 제47조의3'|'국세기본법 제47조의4'|'국세기본법 제47조의5'|'기타'(가산세 대상 아님). "
        "예: [{\"근거\":\"국세기본법 제47조의5\",\"금액\":6000000}]"
    ))] = None,
) -> dict:
    """법인 재해손실세액공제. 언제: 재해로 자산을 상실한 법인의 법인세 세액공제를 계산할 때.
    입력: 사업연도 → year, 자산(assets 또는 assets_before/assets_lost), 당기 법인세 → current_tax,
          공제대상 법인세에서 차감할 세액 → other_credits 또는 credit_items, 미납 법인세 → unpaid_prior_tax,
          가산세 → penalty 또는 penalty_items.
    결과: 자산상실비율·공제대상 법인세·재해손실세액공제·판정·자산계산과정(assets 사용 시)·판단(credit_items·penalty_items 사용 시)·근거.
    공제대상 법인세 = 미납 법인세 + max(당기 법인세 − 조세특례제한법 공제, 0) + 가산세(시행령 제95조③ 대상).
    법인세법 공제(credit_items '법인세법')는 공제대상 법인세에서 빼지 않음. '기타' 가산세도 제외."""
    return K.disaster_loss_credit(year=year, assets_before=assets_before, assets_lost=assets_lost,
                                  current_tax=current_tax, other_credits=other_credits, unpaid_prior_tax=unpaid_prior_tax,
                                  penalty=penalty, assets=assets, credit_items=credit_items, penalty_items=penalty_items)


# ── 일시상각충당금·압축기장충당금 ──────────────────────────────────
@mcp.tool(annotations=CALC)
@_guard
def asset_grant_reserve(
    year: YEAR,
    kind: Annotated[Literal["국고보조금", "공사부담금", "보험차익"], Field(description="재원 구분")],
    received: Annotated[int, Field(description="수령액(금전·현물 가액)")],
    assets: Annotated[list, Field(description="[{asset, cost, used?(그 자산에 쓴 보조금·공사부담금·보험차익 금액, 원), depreciable, depreciation_this_year(그해 세무상 손금으로 인정되는 감가상각비 — 회사계상액이 상각범위액을 넘으면 상각범위액, depreciation_adjustment로 먼저 계산), disposed?, disposed_ratio?, book_provision?, book_reversal?}]")],
    lost_book_value: Annotated[int | None, Field(description="멸실·손괴 보험대상자산 장부가액")] = None,
    receipt_year: Annotated[bool, Field(description="수령·설정 연도")] = True,
    planned_use: WON = 0,
    expired_unused: WON = 0,
    claim_by_return: Annotated[bool, Field(description="신고조정 손금산입")] = True,
) -> dict:
    """국고보조금·공사부담금·보험차익 일시상각충당금·압축기장충당금 세무조정. 언제: 국고보조금·공사부담금·보험차익으로 사업용자산을 취득·개량한 법인의 손금산입·충당금 상계·처분 계산.
입력: 사업연도 → year, 재원 구분 → kind, 수령액 → received, 자산별 내역 → assets(취득·개량가액·사용액·상각·처분·충당금 등).
결과: 손금산입 한도·충당금 설정금액·상계액·처분 익금·사용계획·미사용·기말 잔액·소득금액 순영향·세무조정·근거. 일시상각충당금·압축기장충당금 계산."""
    return C.asset_grant_reserve(
        year=year, kind=kind, received=received, assets=assets,
        lost_book_value=lost_book_value, receipt_year=receipt_year,
        planned_use=planned_use, expired_unused=expired_unused, claim_by_return=claim_by_return,
    )

# ── 재고자산 평가 세무조정 ──────────────────────────────────────
@mcp.tool(annotations=CALC)
@_guard
def inventory_valuation_adjustment(
    year: YEAR,
    items: Annotated[list, Field(description="재고자산 구분별 목록 (1개 이상). 각 항목: category, reported_method, used_method, book_value, valuations, change_report(선택), reported_lcm/used_lcm·name·prior_adjustment(선택)")],
    fy_start: Annotated[str | None, Field(description="사업연도 개시일(YYYY-MM-DD). 생략 시 year-01-01")] = None,
    fy_end: Annotated[str | None, Field(description="사업연도 종료일(YYYY-MM-DD). 생략 시 year-12-31")] = None,
    new_corp: Annotated[bool, Field(description="설립(수익사업 개시) 첫 사업연도 여부 — 최초 신고기한 적용")] = False,
) -> dict:
    """재고자산 평가 세무조정(세법상 평가방법 결정·평가차 유보). 언제: 법인세 세무조정에서 재고자산 기말평가액과 장부가액 차이를 유보로 조정할 때.

입력:
  year — 사업연도 개시 연도(2016~2026)
  fy_start/fy_end — 사업연도 개시일·종료일(YYYY-MM-DD). 생략 시 year-01-01/year-12-31
  new_corp — 설립(수익사업 개시) 첫 사업연도인지 (최초 신고기한: 종료일 달 말일 + 3개월)
  items — [{category: '제품·상품'|'반제품·재공품'|'원재료'|'저장품'|'매매목적부동산',
            name?(표시용), reported_method: 원가법 6종|'무신고',
            reported_lcm?(저가법 신고 여부), used_method: 원가법 6종,
            used_lcm?(저가법 평가 여부), change_report?: {method, lcm?, filed_on: YYYY-MM-DD},
            book_value: int≥0, valuations: {개별법·선입선출법·...·시가: int≥0},
            prior_adjustment?: int(전기 유보 잔액)}]

결과: 세무조정사항 목록, 구분별 판정·세법 방법·평가액·조정액, 익금산입·손금산입 합계, 순조정액, 계산(기한 등), 근거.
근거: 법인세법 제42조①2호·②, 시행령 제73조1호·제74조①~⑥."""
    from .engine.calc_cit import inventory_valuation_adjustment as _calc
    r = _calc(year=year, items=items, fy_start=fy_start, fy_end=fy_end, new_corp=new_corp)
    if "error" in r:
        return r
    return _ok(r["결과"], r["근거"], 계산=r["계산"], 판단=r.get("판단", []))

# ── 0.2 추가: 도구 공백 5건 (SPEC_gold_gaps2) ──────────────────────

def _rate_norm(r):
    """비율 입력 정규화: 1 초과면 ÷100, 아니면 그대로. 0~1 소수 또는 '%' 표기(예 '5%') 둘 다 허용."""
    if r is None: return None
    if isinstance(r, str):
        r = r.strip()
        if r.endswith("%"):
            r = r[:-1]
        r = float(r)
    else:
        r = float(r)
    return r / 100 if r > 1 else r


@mcp.tool(annotations=CALC)
@_guard
def corporate_deemed_rent_income(
    year: YEAR,
    rentals: Annotated[list, Field(description="[{구분:'주택'|'비주택', 보증금(원), 건설비(건물 취득·건설 누계, 원), 임대면적비율?(기본1=건물 전체, 소수), 일수?(기본 그 해 일수)}]")],
    financial_income: Annotated[int, Field(description="비주택 임대사업 금융수익 합계(수입이자·할인료·배당금 등, 원)", ge=0)] = 0,
    rate: Annotated[str | float | None, Field(description="정기예금이자율(소수 또는 '%' 표기 예 '5%', 1 초과면 ÷100, 생략 시 연도표)")]=None,
    debt_over_2x: Annotated[bool, Field(description="차입금적수 > 자기자본적수×2 여부(기본 True)")] = True,
    rental_main_business: Annotated[bool, Field(description="부동산임대업 주업 여부(기본 True)")] = True,
) -> dict:
    """법인 임대보증금 간주익금(조특법 제138조, 조특령 제132조). 언제: 법인이 부동산임대업에서 임대보증금을 받을 때 그중 일부를 간주익금으로 익금산입할지 계산한다.
입력: 사업연도(2016~2026) → year, 임대물건 목록 [{구분, 보증금, 건설비, 임대면적비율, 일수}] → rentals, 비주택 금융수익 합계(원) → financial_income, 정기예금이자율(소수·% 모두 허용, 생략 시 연도표) → rate, 차입금 과다 여부 → debt_over_2x, 부동산임대 주업 여부 → rental_main_business.
결과: 간주익금·보증금적수·건설비상당액적수·제외_주택보증금·세무조정, 판단(건물별 결론), 근거(조세특례제한법 제138조①·②, 시행령 제132조)."""
    import calendar as _cal
    rate = _rate_norm(rate)
    if rate is None:
        rate = V.DEEMED_RENT_RATE.get(year)
        if rate is None:
            return {"error": "입력 확인 필요 — 계산하지 않음",
                    "문제": [f"{year}년 정기예금이자율 연도표 없음"]}
    year_days = 366 if _cal.isleap(year) else 365

    if not debt_over_2x or not rental_main_business:
        return _ok({"간주익금": 0, "세무조정": []},
                   "조세특례제한법 제138조①",
                   계산={"보증금적수": 0, "건설비상당액적수": 0, "제외_주택보증금": 0},
                   판단=[{"항목": "적용요건", "결론": "적용 대상 아님 — 차입금 과다·부동산임대 주업 법인만 (조특법 제138조①)", "근거": "조특법 제138조①"}],
                   산식=f"적용요건 미충족(부채과다={debt_over_2x}, 주업={rental_main_business}) → 간주익금 0")

    보증금적수 = 0
    건설비상당액적수 = 0
    제외_주택보증금 = 0
    판단 = []

    for r in rentals:
        구분 = r.get("구분", "")
        보증금 = int(r.get("보증금", 0) or 0)
        건설비 = int(r.get("건설비", 0) or 0)
        _임대면적비율_raw = r.get("임대면적비율")
        if _임대면적비율_raw is None:
            임대면적비율 = 1
        else:
            임대면적비율 = _rate_norm(_임대면적비율_raw)
        일수 = int(r.get("일수", 0) or year_days)
        if 구분 == "주택":
            제외_주택보증금 += 보증금
            판단.append({"항목": f"{구분} 보증금 {보증금:,}원",
                         "결론": "제외 — 주택(부수토지 포함) 임대보증금은 간주익금 대상 아님 (조특법 제138조① 괄호)",
                         "근거": "조특법 제138조①"})
        elif 구분 == "비주택":
            보증금적수 += 보증금 * 일수
            건설비상당액적수 += 건설비 * 임대면적비율 * 일수
            판단.append({"항목": f"{구분} 보증금 {보증금:,}원 · 건설비 {건설비:,}원 · 비율 {임대면적비율} · 일수 {일수}",
                         "결론": f"적용 — 보증금적수 {보증금*일수:,}원, 건설비상당액적수 {건설비*임대면적비율*일수:,}원",
                         "근거": "조특법 제138조①, 조특령 제132조"})
        else:
            return {"error": "입력 확인 필요 — 계산하지 않음",
                    "문제": [f"rentals: 구분='{구분}'은(는) 허용값 아님 — '주택'|'비주택'"]}

    net_jeoksu = max(0, 보증금적수 - 건설비상당액적수)
    간주익금 = int(max(0, net_jeoksu / year_days * rate - financial_income))
    세무조정 = []
    if 간주익금 > 0:
        세무조정.append({"구분": "임대보증금 간주익금", "금액": 간주익금, "소득처분": "기타사외유출"})

    return _ok({"간주익금": 간주익금, "세무조정": 세무조정},
               "조세특례제한법 제138조①②, 시행령 제132조",
               계산={"보증금적수": 보증금적수, "건설비상당액적수": 건설비상당액적수, "제외_주택보증금": 제외_주택보증금},
               판단=판단,
               산식=f"max(0, ({보증금적수:,} − {건설비상당액적수:,}) ÷ {year_days} × {rate} − {financial_income:,}) = {간주익금:,}원")


@mcp.tool(annotations=CALC)
@_guard
def estimated_business_income(
    year: YEAR,
    industry_group: Annotated[Literal["가", "나", "다"], Field(description="업종군: '가'=농업·도소매·부동산매매 등, '나'=제조·숙박음식·건설 등, '다'=부동산임대·서비스 등")],
    prior_year_revenue: Annotated[int, Field(description="직전연도 수입금액(원)", ge=0)],
    revenues: Annotated[list, Field(description="[{구분:'매출'|'보조금'|'유형자산처분'|'기타', 금액(원)}]")],
    expenses: Annotated[list, Field(description="[{구분:'매입'|'임차료'|'인건비'|'기타'|'자산취득', 금액(원)}] — 증빙 입증 지출")],
    base_rate: Annotated[str | float, Field(description="기준경비율(소수 또는 '%' 표기 예 '20%')")],
    simple_rate: Annotated[str | float, Field(description="단순경비율(소수 또는 '%' 표기 예 '60%')")],
    new_business: Annotated[bool, Field(description="신규 사업 여부(기본 False)")] = False,
    multiplier: Annotated[str | float | None, Field(description="배율(예 2.8 또는 '280%'). 생략 시 귀속연도 기준 — 2016~2018년 간편 2.6·복식 3.2, 2020년~ 간편 2.8·복식 3.4. 2019년 귀속은 결정 시점에 따라 달라 직접 입력 필요")] = None,
) -> dict:
    """추계 사업소득금액(소득령 제143조, 제208조⑤). 언제: 장부가 없거나 불충분한 사업자의 사업소득금액을 추계로 계산할 때.
입력: 귀속연도(2016~2026) → year, 업종군(가·나·다) → industry_group, 직전연도 수입금액(원) → prior_year_revenue, 수입금액 목록 [{구분, 금액}] → revenues, 증빙 입증 지출 목록 [{구분, 금액}] → expenses, 기준경비율 → base_rate, 단순경비율 → simple_rate, 신규 사업 여부 → new_business, 배율(소수·% 허용, 생략 시 귀속연도 기본) → multiplier.
연도: 단순경비율 대상의 '해당 과세기간 수입금액' 요건은 2019년 귀속부터(영 제143조④, 2018.2.13. 개정), 2018년까지는 신규사업자에만 적용.
복식부기의무자의 사업용 유형자산 처분액 수입금액 산입은 2018년 귀속부터(법 제19조①20호, 2017.12.19. 신설).
배율(규칙 제67조): 2020.3.12.까지 간편 2.6·복식 3.2, 2020.3.13.부터 간편 2.8·복식 3.4 — 2019년 귀속은 multiplier 직접 입력.
결과: 사업소득금액·계산(수입금액·장부구분·적용방법·주요경비·기준경비율_소득·한도), 판단(수입금액 구성·제외·적용방법 사유·주요경비 제외), 근거(소득세법 시행령 제143조·제208조⑤)."""
    base_rate = _rate_norm(base_rate)
    simple_rate = _rate_norm(simple_rate)
    if isinstance(multiplier, str) and multiplier.strip().endswith("%"):
        multiplier = float(multiplier.strip()[:-1]) / 100
    elif multiplier is not None:
        multiplier = float(multiplier)

    # revenues·expenses 구분 검증
    _rev_kinds = {"매출", "보조금", "유형자산처분", "기타"}
    _exp_kinds = {"매입", "임차료", "인건비", "기타", "자산취득"}
    for i, r in enumerate(revenues or []):
        if r.get("구분") not in _rev_kinds:
            return {"error": "입력 확인 필요 — 계산하지 않음",
                    "문제": [f"revenues 항목 {i+1}: 구분='{r.get('구분')}' 허용값 아님 — {sorted(_rev_kinds)}"]}
    for i, e in enumerate(expenses or []):
        if e.get("구분") not in _exp_kinds:
            return {"error": "입력 확인 필요 — 계산하지 않음",
                    "문제": [f"expenses 항목 {i+1}: 구분='{e.get('구분')}' 허용값 아님 — {sorted(_exp_kinds)}"]}

    # 기준금액표
    기준표 = {"가": {"간편장부": 300_000_000, "단순경비율": 60_000_000},
              "나": {"간편장부": 150_000_000, "단순경비율": 36_000_000},
              "다": {"간편장부": 75_000_000,  "단순경비율": 24_000_000}}
    if industry_group not in 기준표:
        return {"error": "입력 확인 필요 — 계산하지 않음",
                "문제": [f"industry_group='{industry_group}'은(는) 허용값 아님 — '가'|'나'|'다'"]}
    기준 = 기준표[industry_group]

    # 장부구분
    if new_business:
        장부구분 = "간편장부대상자"
    elif prior_year_revenue < 기준["간편장부"]:
        장부구분 = "간편장부대상자"
    else:
        장부구분 = "복식부기의무자"

    # 배율(소득세법 시행규칙 제67조): 2020.3.12.까지 3.2(간편 2.6), 2020.3.13. 개정 후 3.4(간편 2.8).
    # 2019년 귀속은 2020.3.13. 전후 어느 시점에 결정·경정하느냐에 따라 달라져(개정 부칙 미확인) 직접 입력을 받는다.
    배율_근거 = "소득세법 시행규칙 제67조(입력값)"
    if multiplier is None:
        if year == 2019:
            pass   # 기준경비율 계산에 배율이 필요할 때 오류(아래)
        elif year <= 2018:
            multiplier = 2.6 if 장부구분 == "간편장부대상자" else 3.2
            배율_근거 = "소득세법 시행규칙 제67조(2020.3.13. 개정 전)"
        else:
            multiplier = 2.8 if 장부구분 == "간편장부대상자" else 3.4
            배율_근거 = "소득세법 시행규칙 제67조(2020.3.13. 개정 후)"

    # 수입금액
    매출 = sum(int(r.get("금액", 0) or 0) for r in revenues if r.get("구분") == "매출")
    보조금 = sum(int(r.get("금액", 0) or 0) for r in revenues if r.get("구분") == "보조금")
    기타 = sum(int(r.get("금액", 0) or 0) for r in revenues if r.get("구분") == "기타")
    유형자산처분 = sum(int(r.get("금액", 0) or 0) for r in revenues if r.get("구분") == "유형자산처분")

    수입금액 = 매출 + 보조금 + 기타
    수입금액_판단 = []
    if 장부구분 == "복식부기의무자" and year >= 2018:
        수입금액 += 유형자산처분
        수입금액_판단.append(f"복식부기의무자 — 유형자산처분 {유형자산처분:,}원 포함(법 제19조①20호)")
    elif 장부구분 == "복식부기의무자":
        수입금액_판단.append(f"{year}년 귀속 — 유형자산처분 {유형자산처분:,}원 제외(법 제19조①20호는 2018년 귀속부터)")
    else:
        수입금액_판단.append(f"간편장부대상자 — 유형자산처분 {유형자산처분:,}원 제외(법 제19조①20호는 복식부기의무자만)")

    # 적용방법(영 제143조④). 2019년 귀속부터: (신규 또는 직전 수입 기준 미달) 그리고 해당 과세기간 수입 < 제208조⑤2호 금액.
    # 2018년 귀속까지: 신규사업자만 해당 과세기간 수입 요건, 계속사업자는 직전 수입 기준만.
    직전_미달 = prior_year_revenue < 기준["단순경비율"]
    당해_미달 = 수입금액 < 기준["간편장부"]
    if year >= 2019:
        단순대상 = (new_business or 직전_미달) and 당해_미달
    else:
        단순대상 = (new_business and 당해_미달) or (not new_business and 직전_미달)
    적용방법 = "단순경비율" if 단순대상 else "기준경비율"

    판단 = []
    판단.append({"항목": "직전연도 수입금액", "결론": f"{prior_year_revenue:,}원 — {'단순경비율 기준 미만' if (new_business or prior_year_revenue < 기준['단순경비율']) else '단순경비율 기준 이상'} (기준 {기준['단순경비율']:,}원)",
                 "근거": "소득령 제143조④"})
    판단.append({"항목": "당해 수입금액", "결론": f"{수입금액:,}원 — {'간편장부 기준 미만' if 수입금액 < 기준['간편장부'] else '간편장부 기준 이상'} (기준 {기준['간편장부']:,}원)"
                 + ("" if (year >= 2019 or new_business) else " — 2018년 귀속까지 계속사업자는 이 요건 없음"),
                 "근거": "소득령 제143조④" + (" (2018.2.13. 개정, 2019년 귀속부터)" if year >= 2019 else " (2018.2.13. 개정 전)")})
    판단.append({"항목": "적용방법", "결론": 적용방법, "근거": "소득령 제143조③④"})

    if 적용방법 == "단순경비율":
        사업소득금액 = int(round(수입금액 * (1 - simple_rate), 6))
        return _ok({"사업소득금액": 사업소득금액},
                   "소득세법 시행령 제143조③",
                   계산={"수입금액": 수입금액, "장부구분": 장부구분, "적용방법": 적용방법, "주요경비": 0,
                         "기준경비율_소득": 0, "한도": None},
                   판단=판단 + [{"항목": "단순경비율 소득", "결론": f"{수입금액:,} × (1 − {simple_rate}) = {사업소득금액:,}원",
                                  "근거": "소득령 제143조③"}])

    # 기준경비율
    매입 = sum(int(r.get("금액", 0) or 0) for r in expenses if r.get("구분") == "매입")
    임차료 = sum(int(r.get("금액", 0) or 0) for r in expenses if r.get("구분") == "임차료")
    인건비 = sum(int(r.get("금액", 0) or 0) for r in expenses if r.get("구분") == "인건비")
    주요경비 = 매입 + 임차료 + 인건비

    주요경비_판단 = []
    기타_지출 = sum(int(r.get("금액", 0) or 0) for r in expenses if r.get("구분") == "기타")
    자산취득_지출 = sum(int(r.get("금액", 0) or 0) for r in expenses if r.get("구분") == "자산취득")
    if 기타_지출 > 0:
        주요경비_판단.append(f"기타 지출 {기타_지출:,}원 — 주요경비에서 제외")
    if 자산취득_지출 > 0:
        주요경비_판단.append(f"자산취득 지출 {자산취득_지출:,}원 — 주요경비에서 제외")

    적용기준경비율 = base_rate / 2 if 장부구분 == "복식부기의무자" else base_rate
    기준경비율_소득 = 수입금액 - 주요경비 - 수입금액 * 적용기준경비율
    기준경비율_소득 = max(0, int(round(기준경비율_소득, 6)))
    단순경비율_소득 = int(round(수입금액 * (1 - simple_rate), 6))
    if multiplier is None:
        return {"error": "입력 확인 필요 — 계산하지 않음",
                "문제": ["2019년 귀속은 추계 배율이 결정·경정 시점에 따라 간편 2.6·복식 3.2(2020.3.12.까지) 또는 "
                         "간편 2.8·복식 3.4(2020.3.13.부터, 소득세법 시행규칙 제67조) — multiplier를 직접 입력"]}
    한도 = int(round(단순경비율_소득 * multiplier, 6))
    사업소득금액 = min(기준경비율_소득, 한도)

    판단.append({"항목": "주요경비", "결론": f"매입 {매입:,} + 임차료 {임차료:,} + 인건비 {인건비:,} = {주요경비:,}원 (기타·자산취득 제외)",
                 "근거": "소득령 제143조③"})
    if 주요경비_판단:
        for j in 주요경비_판단:
            판단.append({"항목": "주요경비 제외", "결론": j, "근거": "소득령 제143조③"})
    판단.append({"항목": "기준경비율 소득", "결론": f"{수입금액:,} − {주요경비:,} − {수입금액:,}×{적용기준경비율} = {기준경비율_소득:,}원",
                 "근거": "소득령 제143조③"})
    판단.append({"항목": "단순경비율 소득 × 배율", "결론": f"{단순경비율_소득:,} × {multiplier} = {한도:,}원 (한도)",
                 "근거": f"소득령 제143조③ 단서, {배율_근거}"})
    판단.append({"항목": "사업소득금액", "결론": f"min({기준경비율_소득:,}, {한도:,}) = {사업소득금액:,}원",
                 "근거": "소득령 제143조③"})

    return _ok({"사업소득금액": 사업소득금액},
               "소득세법 시행령 제143조③④, 제208조⑤",
               계산={"수입금액": 수입금액, "장부구분": 장부구분, "적용방법": 적용방법, "주요경비": 주요경비,
                     "기준경비율_소득": 기준경비율_소득, "한도": 한도},
               판단=판단,
               산식=f"수입금액 {수입금액:,} − 주요경비 {주요경비:,} − {수입금액:,}×{적용기준경비율}={기준경비율_소득:,}, 단순경비율 소득×{multiplier}={한도:,} → min = {사업소득금액:,}원")


@mcp.tool(annotations=CALC)
@_guard
def vat_exempt_use_deemed_supply(
    use_date: Annotated[str, Field(description="면세사업에 쓰기 시작한 날(YYYY-MM-DD)")],
    mode: Annotated[Literal["전부", "일부"], Field(description="전용 범위: '전부'=전액 면세전용, '일부'=과세·면세 공통 사용")],
    assets: Annotated[list, Field(description="[{구분:'건물'|'구축물'|'기타'|'토지'|'비상각', 취득일(YYYY-MM-DD), 취득가액(원), 매입세액공제?(기본True, bool), 시가?(비상각시 필수, 원)}]")],
    supplies: Annotated[dict | None, Field(description="일부일 때 필수: {taxable: 과세 공급가액(원), exempt: 면세 공급가액(원)}")]=None,
) -> dict:
    """면세전용 간주공급 과세표준(부가법 제10조①, 부가령 제66조). 언제: 과세사업에 쓰던 재화를 면세사업에 전용할 때 간주공급 과세표준을 계산한다.
입력: 전용일(YYYY-MM-DD) → use_date, 전용 범위(전부/일부) → mode, 자산 목록 [{구분, 취득일, 취득가액, 매입세액공제, 시가}] → assets, 일부일 때 공급가액 {taxable, exempt} → supplies.
결과: 과세표준·계산(자산별 경과과세기간수·간주시가·제외사유, 면세비율)·판단, 근거(부가가치세법 제10조①, 시행령 제66조)."""
    if mode == "일부" and (supplies is None or not isinstance(supplies, dict) or "taxable" not in supplies or "exempt" not in supplies):
        return {"error": "입력 확인 필요 — 계산하지 않음",
                "문제": ["일부 전용 시 supplies={taxable, exempt} 필요"]}

    판단 = []
    자산별 = []
    간주시가_합 = 0

    for a in assets:
        구분 = a.get("구분", "")
        취득일 = str(a.get("취득일", ""))
        취득가액 = int(a.get("취득가액", 0) or 0)
        매입세액공제 = bool(a.get("매입세액공제", True))
        시가 = a.get("시가")

        if 구분 not in ("건물", "구축물", "기타", "토지", "비상각"):
            return {"error": "입력 확인 필요 — 계산하지 않음",
                    "문제": [f"assets: 구분='{구분}'은(는) 허용값 아님 — '건물'|'구축물'|'기타'|'토지'|'비상각'"]}

        # 토지 제외
        if 구분 == "토지":
            자산별.append({"구분": 구분, "제외사유": "토지 — 면세 재화(부가법 제10조① 괄호)",
                           "근거": "부가법 제10조①"})
            판단.append({"항목": f"{구분} {취득일} 취득가액 {취득가액:,}원", "결론": "제외 — 토지(면세 재화)",
                         "근거": "부가법 제10조①"})
            continue

        # 매입세액공제 받지 않은 자산 제외
        if not 매입세액공제:
            자산별.append({"구분": 구분, "제외사유": "매입세액공제 받지 않은 자산(부가법 제10조① 괄호)",
                           "근거": "부가법 제10조①"})
            판단.append({"항목": f"{구분} {취득일} 취득가액 {취득가액:,}원", "결론": "제외 — 매입세액공제 미수령 자산",
                         "근거": "부가법 제10조①"})
            continue

        # 비상각 재화
        if 구분 == "비상각":
            if 시가 is None:
                return {"error": "입력 확인 필요 — 계산하지 않음",
                        "문제": [f"assets: 비상각 자산은 시가 필수"]}
            간주시가 = int(시가)
            자산별.append({"구분": 구분, "경과과세기간수": 0, "간주시가": 간주시가})
            간주시가_합 += 간주시가
            판단.append({"항목": f"{구분} {취득일} 취득가액 {취득가액:,}원, 시가 {간주시가:,}원",
                         "결론": f"비상각 재화 — 시가 {간주시가:,}원 적용",
                         "근거": "부가령 제66조①"})
            continue

        # 감가상각자산
        체감률 = 0.05 if 구분 in ("건물", "구축물") else 0.25
        경과 = V.elapsed_periods(취득일, use_date)
        if 경과 < 0:
            return {"error": "입력 확인 필요 — 계산하지 않음",
                    "문제": [f"자산 '{구분}': 전용일이 취득일보다 앞섬"]}
        간주시가 = int(max(0, 취득가액 * (1 - 체감률 * 경과)))
        자산별.append({"구분": 구분, "경과과세기간수": 경과, "간주시가": 간주시가})
        간주시가_합 += 간주시가
        판단.append({"항목": f"{구분} {취득일} 취득, 취득가액 {취득가액:,}원",
                     "결론": f"경과과세기간수 {경과}기, 간주시가 {취득가액:,}×(1−{체감률}×{경과})={간주시가:,}원",
                     "근거": "부가령 제66조①"})

    if mode == "전부":
        과세표준 = 간주시가_합
        면세비율 = None
    else:
        taxable = int(supplies.get("taxable", 0) or 0)
        exempt = int(supplies.get("exempt", 0) or 0)
        total = taxable + exempt
        면세비율 = exempt / total if total > 0 else 0
        if 면세비율 < 0.05:
            과세표준 = 0
            판단.append({"항목": "면세비율", "결론": f"{면세비율*100:.1f}% < 5% → 과세표준 0 (부가령 제66조③ 단서)",
                         "근거": "부가령 제66조③"})
        else:
            과세표준 = int(간주시가_합 * 면세비율)
            판단.append({"항목": "면세비율", "결론": f"면세 공급가액 {exempt:,} / 총 {total:,} = {면세비율*100:.1f}% ≥ 5%",
                         "근거": "부가령 제66조③"})

    계산 = {"자산별": 자산별}
    if mode == "일부":
        계산["면세비율"] = 면세비율

    return _ok({"과세표준": 과세표준},
               "부가가치세법 제10조①, 시행령 제66조",
               계산=계산,
               판단=판단,
               산식=("전부: 간주시가 합 = " + f"{간주시가_합:,}원" if mode == "전부" else
                     f"간주시가 합 {간주시가_합:,} × 면세비율 {면세비율*100:.1f}% = {과세표준:,}원"))


@mcp.tool(annotations=CALC)
@_guard
def liquidation_income(
    assets_realized: Annotated[int | list, Field(description="환가 총액(원) 또는 [{항목, 금액(원)}]")],
    liabilities_paid: WON,
    capital: WON,
    surplus: WON,
    yubo: Annotated[int, Field(description="세무상 유보 잔액(원, △유보는 음수)", ge=None)] = 0,
    carryforward_loss: Annotated[int, Field(description="세무상 이월결손금 잔액(원)", ge=0)] = 0,
    capitalized_surplus_2y: Annotated[int, Field(description="해산등기일 전 2년 이내 자본금에 전입한 잉여금(원)", ge=0)] = 0,
    refund_tax: Annotated[int, Field(description="해산등기일 현재 환급될 법인세(원)", ge=0)] = 0,
) -> dict:
    """청산소득금액(법인세법 제79조). 언제: 법인이 해산·청산할 때 청산소득금액을 계산한다.
입력: 환가 총액(원) 또는 [{항목, 금액}] → assets_realized, 부채 상환액(원) → liabilities_paid, 자본금(원) → capital, 자본잉여금+이익잉여금(원) → surplus, 세무상 유보 잔액(원, △유보는 음수) → yubo, 이월결손금 잔액(원) → carryforward_loss, 2년 이내 자본전입 잉여금(원) → capitalized_surplus_2y, 환급 법인세(원) → refund_tax.
결과: 청산소득금액·계산(잔여재산가액·자기자본총액·결손금_상계·자본금·잉여금)·판단, 근거(법인세법 제79조)."""
    # 환가 총액
    if isinstance(assets_realized, list):
        환가총액 = sum(int(a.get("금액", 0) or 0) for a in assets_realized)
    else:
        환가총액 = int(assets_realized) if assets_realized else 0

    잔여재산가액 = 환가총액 - liabilities_paid

    # 2년 이내 자본전입 잉여금 조정: 전입하지 않은 것으로 봄
    자본금 = capital - int(capitalized_surplus_2y)
    잉여금 = int(surplus) + int(yubo) + int(capitalized_surplus_2y)

    판단 = []
    if int(capitalized_surplus_2y) > 0:
        판단.append({"항목": "2년 이내 자본전입 잉여금", "결론": f"{int(capitalized_surplus_2y):,}원 — 전입하지 않은 것으로 봄: 자본금 {capital:,}→{자본금:,}원, 잉여금 +{int(capitalized_surplus_2y):,}원 (법 제79조⑤)",
                     "근거": "법인세법 제79조⑤"})

    # 이월결손금 상계
    상계액 = min(int(carryforward_loss), 잉여금)
    잉여금_after = 잉여금 - 상계액
    if 상계액 > 0:
        판단.append({"항목": "이월결손금 상계", "결론": f"min({int(carryforward_loss):,}원, 잉여금 {잉여금:,}원) = {상계액:,}원 상계, 잔여 잉여금 {잉여금_after:,}원 (법 제79조④)",
                     "근거": "법인세법 제79조④"})
    elif int(carryforward_loss) > 잉여금:
        판단.append({"항목": "이월결손금", "결론": f"이월결손금 {int(carryforward_loss):,}원 > 잉여금 {잉여금:,}원 → 초과분 없는 것으로 봄 (법 제79조④)",
                     "근거": "법인세법 제79조④"})

    자기자본총액 = 자본금 + 잉여금_after + int(refund_tax)
    청산소득금액 = 잔여재산가액 - 자기자본총액

    if 청산소득금액 <= 0:
        판단.append({"항목": "청산소득금액", "결론": "청산소득 없음", "근거": "법인세법 제79조"})

    return _ok({"청산소득금액": int(청산소득금액)},
               "법인세법 제79조",
               계산={"잔여재산가액": 잔여재산가액, "자기자본총액": 자기자본총액,
                     "결손금_상계": 상계액, "자본금": 자본금, "잉여금": 잉여금_after},
               판단=판단,
               산식=f"잔여재산가액 {잔여재산가액:,} − 자기자본총액 {자기자본총액:,} = {청산소득금액:,}원")


@mcp.tool(annotations=CALC)
@_guard
def income_adjustment_total(
    net_income: WON,
    items: Annotated[list, Field(description="[{항목(이름), 종류, 금액(원), 방향?('가산'|'차감'|'없음'), 소득처분?}] — 종류: 법인세비용·법인지방소득세·벌금·과태료·가산세·환급금이자·자기주식처분이익(자본잉여금)·미수이자(원천징수대상)·비지정기부금·기업업무추진비 한도초과·감가상각비 한도초과·업무무관 지급이자·수입배당금 익금불산입·이월결손금 보전 자산수증이익·사업용 재산세·기부금 한도초과(별칭: 법인세·국세환급가산금·자기주식처분이익·미수이자·재산세 등). 사전에 없으면 방향을 주면 됨")],
    donation_excess: Annotated[int, Field(description="기부금 한도초과액(원), 차가감 이후 가산", ge=0)] = 0,
    donation_carryover_deduction: Annotated[int, Field(description="기부금 한도초과 이월액 손금산입(원)", ge=0)] = 0,
) -> dict:
    """차가감소득금액 집계(법인세 과세표준 및 세액조정계산서). 언제: 당기순이익에서 세법상 익금산입·손금산입을 가감해 차가감소득금액과 각 사업연도소득금액을 구할 때.
입력: 당기순이익(원) → net_income, 조정 항목 목록 [{항목, 종류, 금액, 방향, 소득처분}] → items, 기부금 한도초과액(원) → donation_excess, 기부금 한도초과 이월액 손금산입(원) → donation_carryover_deduction.
결과: 차가감소득금액·각사업연도소득금액·익금산입_손금불산입·손금산입_익금불산입·세무조정, 판단, 근거(법인세법 제18·19·21·23·24·25·28조 등)."""
    kinds = vocab.INCOME_ADJUSTMENT_KINDS
    # 흔히 쓰는 다른 이름 → 사전 종류 (Claude 2026-10-28)
    _별칭 = {"법인세": "법인세비용", "법인세등": "법인세비용", "국세환급가산금": "환급금이자", "지방세환급가산금": "환급금이자",
            "자기주식처분이익": "자기주식처분이익(자본잉여금)", "미수이자": "미수이자(원천징수대상)", "재산세": "사업용 재산세",
            "접대비 한도초과": "기업업무추진비 한도초과", "벌금": "벌금·과태료·가산세", "과태료": "벌금·과태료·가산세", "가산세": "벌금·과태료·가산세"}
    for it in items:
        if it.get("종류") in _별칭:
            it["종류"] = _별칭[it["종류"]]
    _모름 = [f"'{it.get('항목', '')}'(종류 '{it.get('종류', '')}')" for it in items
            if it.get("종류") not in kinds and it.get("방향") is None]
    if _모름:
        return {"error": "입력 확인 필요 — 계산하지 않음",
                "문제": [f"사전에 없는 종류 {len(_모름)}건: {', '.join(_모름)} — 종류를 아래 중에서 고르거나 방향('가산'|'차감'|'없음')을 주세요"],
                "허용 종류": list(kinds)}
    # 조정 이름: 가산 = 익금산입 또는 손금불산입, 차감 = 손금산입 또는 익금불산입
    _조정명 = {"법인세비용": "손금불산입", "법인지방소득세": "손금불산입", "벌금·과태료·가산세": "손금불산입", "비지정기부금": "손금불산입",
             "기업업무추진비 한도초과": "손금불산입", "감가상각비 한도초과": "손금불산입", "업무무관 지급이자": "손금불산입",
             "자기주식처분이익(자본잉여금)": "익금산입", "환급금이자": "익금불산입", "수입배당금 익금불산입": "익금불산입",
             "이월결손금 보전 자산수증이익": "익금불산입", "미수이자(원천징수대상)": "익금불산입"}
    익금산입 = 0
    손금산입 = 0
    세무조정 = []
    판단 = []

    # 기부금 한도초과 금액 추출 (items 내 종류='기부금 한도초과' 항목에서)
    extracted_donation_excess = 0
    for it in items:
        if it.get("종류") == "기부금 한도초과":
            extracted_donation_excess += int(it.get("금액", 0) or 0)
    donation_excess = int(donation_excess) + extracted_donation_excess

    for it in items:
        항목명 = it.get("항목", "")
        종류 = it.get("종류", "")
        금액 = int(it.get("금액", 0) or 0)
        방향_입력 = it.get("방향", None)
        소득처분_입력 = it.get("소득처분", None)

        if 종류 not in kinds:
            if 방향_입력 is None:
                return {"error": "입력 확인 필요 — 계산하지 않음",
                        "문제": [f"항목 '{항목명}': 종류 '{종류}' 사전에 없음, 방향을 주세요"]}
            # 방향은 있으나 사전에 없는 종류 — 소득처분은 입력값 또는 기본
            방향 = 방향_입력
            # 소득처분 기본값: 가산→기타사외유출, 차감→기타, 없음→—
            if 소득처분_입력:
                소득처분 = 소득처분_입력
            elif 방향 == "가산":
                소득처분 = "기타사외유출"
            elif 방향 == "차감":
                소득처분 = "기타"
            else:
                소득처분 = "—"
            근거 = "입력 방향 기준"
        else:
            kd = kinds[종류]
            방향 = 방향_입력 if 방향_입력 is not None else kd["방향"]
            소득처분 = 소득처분_입력 if 소득처분_입력 is not None else kd["소득처분"]
            근거 = kd["근거"]

        if 방향_입력 is not None and 방향_입력 not in ("가산", "차감", "없음"):
            return {"error": "입력 확인 필요 — 계산하지 않음",
                    "문제": [f"항목 '{항목명}': 방향 '{방향_입력}' 허용값 아님 — '가산'|'차감'|'없음'"]}

        # 기부금 한도초과 — 차가감 이후 가산
        if 종류 == "기부금 한도초과":
            판단.append({"항목": 항목명, "결론": f"차가감소득금액 이후 가산 — 금액 {금액:,}원 (법 제24조)",
                         "근거": 근거})
            continue

        if 방향 == "없음":
            판단.append({"항목": 항목명, "결론": f"조정 없음 — {금액:,}원 (손금 인정, 세무조정 대상 아님)",
                         "근거": 근거})
            continue

        if 방향 == "가산":
            익금산입 += 금액
            조정 = _조정명.get(종류, "익금산입·손금불산입")
            세무조정.append({"구분": 항목명, "금액": 금액, "소득처분": 소득처분, "조정": "익금산입", "조정구분": 조정, "종류": 종류})   # '조정'은 방향(엑셀 SUMIF용, 호환)
            판단.append({"항목": 항목명, "결론": f"{조정}(가산) +{금액:,}원, 소득처분 {소득처분}",
                         "근거": 근거})
        elif 방향 == "차감":
            손금산입 += 금액
            조정 = _조정명.get(종류, "손금산입·익금불산입")
            세무조정.append({"구분": 항목명, "금액": 금액, "소득처분": 소득처분, "조정": "손금산입", "조정구분": 조정, "종류": 종류})   # '조정'은 방향(엑셀 SUMIF용, 호환)
            판단.append({"항목": 항목명, "결론": f"{조정}(차감) −{금액:,}원, 소득처분 {소득처분}",
                         "근거": 근거})

    차가감소득금액 = net_income + 익금산입 - 손금산입
    각사업연도소득금액 = 차가감소득금액 + int(donation_excess) - int(donation_carryover_deduction)

    if int(donation_excess) > 0:
        세무조정.append({"구분": "기부금 한도초과", "금액": int(donation_excess), "소득처분": "기타사외유출", "조정": "익금산입"})
        판단.append({"항목": "기부금 한도초과", "결론": f"차가감 이후 가산 +{int(donation_excess):,}원 (법 제24조)",
                     "근거": "법 제24조"})

    # ── 계산식 단계 (SPEC_export_workbook_c) ──────────────────────────────
    계산식_입력: list[dict] = [
        {"이름": "당기순이익", "값": net_income, "설명": "당기순이익(세무조정 전)"},
    ]
    if int(donation_excess) > 0:
        계산식_입력.append({"이름": "기부금한도초과", "값": int(donation_excess),
                             "설명": "기부금 한도초과액 (법 제24조)"})
    if int(donation_carryover_deduction) > 0:
        계산식_입력.append({"이름": "기부금이월공제", "값": int(donation_carryover_deduction),
                             "설명": "기부금 한도초과 이월액 손금산입"})

    # 항목은 익금산입·손금산입으로 나누어 번호 부여 (방향은 kinds 기준 최종값)
    calc_items_add: list[dict] = []   # 익금산입·손금불산입 (가산)
    calc_items_sub: list[dict] = []   # 손금산입·익금불산입 (차감)
    for it in items:
        금액 = int(it.get("금액", 0) or 0)
        if 금액 == 0:
            continue
        종류 = it.get("종류", "")
        kd = kinds.get(종류, {})
        direction = it.get("방향") or kd.get("방향", "가산")
        항목명 = it.get("항목", "")
        if direction == "차감":
            calc_items_sub.append({"항목명": 항목명, "종류": 종류, "금액": 금액})
        else:
            calc_items_add.append({"항목명": 항목명, "종류": 종류, "금액": 금액})

    for idx, item in enumerate(calc_items_add, 1):
        계산식_입력.append({"이름": f"항목{idx}", "값": item["금액"],
                             "설명": f"{item['항목명']} — {item['종류']} (익금산입)"})
    가산_start = len(calc_items_add)
    for idx, item in enumerate(calc_items_sub, 1):
        계산식_입력.append({"이름": f"항목{가산_start + idx}", "값": item["금액"],
                             "설명": f"{item['항목명']} — {item['종류']} (손금산입)"})

    add_refs = "+".join(f"항목{i}" for i in range(1, len(calc_items_add) + 1)) or "0"
    sub_refs = "+".join(f"항목{i}" for i in range(가산_start + 1, 가산_start + len(calc_items_sub) + 1)) or "0"

    각사_식 = "차가감소득금액"
    if int(donation_excess) > 0:
        각사_식 += "+기부금한도초과"
    if int(donation_carryover_deduction) > 0:
        각사_식 += "-기부금이월공제"

    계산식: list[dict] = [
        {"이름": "익금산입_손금불산입", "식": add_refs,
         "값": 익금산입, "설명": "익금산입·손금불산입 항목 합계"},
        {"이름": "손금산입_익금불산입", "식": sub_refs,
         "값": 손금산입, "설명": "손금산입·익금불산입 항목 합계"},
        {"이름": "차가감소득금액", "식": "당기순이익+익금산입_손금불산입-손금산입_익금불산입",
         "값": 차가감소득금액, "설명": "당기순이익 + 익금산입 − 손금산입"},
        {"이름": "각사업연도소득금액", "식": 각사_식,
         "값": 각사업연도소득금액, "설명": "차가감소득금액 + 기부금한도초과 − 이월공제"},
    ]

    return _ok({"차가감소득금액": 차가감소득금액, "각사업연도소득금액": 각사업연도소득금액,
                "익금산입_손금불산입": 익금산입, "손금산입_익금불산입": 손금산입,
                "세무조정": 세무조정},
               "법인세법 제18·19·21·23·24·25·28조 등",
               판단=판단,
               산식=f"당기순이익 {net_income:,} + 익금산입 {익금산입:,} − 손금산입 {손금산입:,} = 차가감 {차가감소득금액:,}원, + 기부금한도초과 {donation_excess:,} − 이월공제 {donation_carryover_deduction:,} = 각사업연도 {각사업연도소득금액:,}원",
               계산식_입력=계산식_입력, 계산식=계산식)

# ── 모든 도구 공통 — 정의되지 않은 최상위 인자 거부 ─────────────
# MCP SDK의 arg_model은 기본적으로 extra='ignore'라 모르는 키가 조용히 버려짐.
# 각 도구의 arg_model을 extra='forbid'로 재창출해 밸리데이션 단계에서 거부한다.
# (Tool.run은 validate_arguments → arg_model.model_validate → model_dump_one_level 순으로 args를 처리)
from mcp.server.mcpserver.utilities.func_metadata import ArgModelBase
from pydantic import ConfigDict, Field, create_model

for _tool_name, _tool in list(mcp._tool_manager._tools.items()):
    _meta = _tool.fn_metadata
    _old_model = _meta.arg_model
    if _old_model is None:
        continue
    _new_fields = {}
    for _fname, _f in _old_model.model_fields.items():
        # FieldInfo를 그대로 재사용 — ge/le 등 유효성 제약 보존(필수 필드도 마찬가지)
        _new_fields[_fname] = (_f.annotation, _f)
    _new_model = create_model(
        f"{_old_model.__name__}_NoExtra",
        __base__=ArgModelBase,
        __config__=ConfigDict(extra='forbid'),
        **_new_fields,
    )
    _meta.arg_model = _new_model

# ── export_workbook: 계산 근거 엑셀 내보내기 ─────────────────────────────
_INVALID_NAME_CHARS = frozenset("/\\:*?\"<>|")


def _sanitize_name(s: str) -> str:
    return "".join("_" if c in _INVALID_NAME_CHARS else c for c in s)


def _unique_path(dir_path: str, base_name: str, ext: str = ".xlsx") -> str:
    """이미 있으면 (2), (3)… 접미사를 붙여 새 경로 반환."""
    safe = _sanitize_name(base_name)
    candidate = os.path.join(dir_path, safe + ext)
    if not os.path.exists(candidate):
        return candidate
    n = 2
    while True:
        candidate = os.path.join(dir_path, f"{safe} ({n}){ext}")
        if not os.path.exists(candidate):
            return candidate
        n += 1


def _flatten_args(args, prefix: str = "") -> list[tuple[str, Any]]:
    """중첩 dict/list를 경로|값 튜플 목록으로 펼침. 예: revenue.매출액, ledger[0].금액."""
    rows: list[tuple[str, Any]] = []
    if isinstance(args, dict):
        for k, v in args.items():
            path = f"{prefix}.{k}" if prefix else str(k)
            if isinstance(v, (dict, list)):
                rows.extend(_flatten_args(v, path))
            else:
                rows.append((path, v))
    elif isinstance(args, list):
        for i, v in enumerate(args):
            path = f"{prefix}[{i}]"
            if isinstance(v, (dict, list)):
                rows.extend(_flatten_args(v, path))
            else:
                rows.append((path, v))
    else:
        rows.append((prefix, args))
    return rows


def _numeric_items(d: dict) -> list[tuple[str, Any]]:
    """결과 dict에서 숫자·문자 값을 (키, 값) 튜플로 추출."""
    out: list[tuple[str, Any]] = []
    for k, v in d.items():
        if isinstance(v, (int, float)) or (isinstance(v, str) and v):
            out.append((str(k), v))
    return out


def _cell_value(v: Any) -> Any:
    """엑셀 셀에 쓸 수 있는 스칼라 값으로 변환. 목록·dict는 규칙에 따라 문자열로."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return v
    if isinstance(v, str):
        return v
    if isinstance(v, list):
        if _is_basis_list(v):
            return _format_basis(v)
        # 항목마다 한 겹 더 감싼 근거 목록인가? (earned_income_gross_pay)
        if _is_wrapped_basis_list(v):
            return "\n".join(_format_basis(inner) for inner in v if _is_basis_list(inner))
        return _json.dumps(v, ensure_ascii=False)
    if isinstance(v, dict):
        return _json.dumps(v, ensure_ascii=False)
    return str(v)


def _is_basis_list(v: list) -> bool:
    """v가 [법령, 조, 항, 요지] 목록의 목록인가? (근거 목록 판정)"""
    if not v or not isinstance(v, list):
        return False
    for item in v:
        if not isinstance(item, list):
            return False
        if not (3 <= len(item) <= 4):
            return False
        if not all(isinstance(x, str) for x in item):
            return False
        # 둘째 요소가 조 번호인가? (예: "제40조", "제81조①")
        if len(item) >= 2 and not _looks_like_article(item[1]):
            return False
    return True


def _looks_like_article(s: str) -> bool:
    """s가 조 번호처럼 보이는가? (예: "제40조", "제81조①")"""
    if not s.startswith("제"):
        return False
    조_idx = s.find("조")
    if 조_idx <= 1:
        return False
    return s[1:조_idx].isdigit()


def _is_wrapped_basis_list(v: list) -> bool:
    """v가 근거 목록들을 한 겹 더 감싼 목록인가? (earned_income_gross_pay 근거)"""
    if not v or not isinstance(v, list):
        return False
    for item in v:
        if not _is_basis_list(item):
            return False
    return True


def _format_basis(v: list) -> str:
    """근거 목록([법령, 조, 항, 요지] 목록)을 사람이 읽는 문장으로 변환.

    각 항목의 자리(0 법령·1 조·2 항·3 요지)로 먼저 읽고, 빈 값만 생략한다.
    빈 값을 먼저 걸러내면 요지가 항 자리에 붙는 문제(예: 제40조실지귀속)가 생긴다.
    """
    lines = []
    for item in v:
        if not item:
            continue
        law_name = item[0] if len(item) > 0 and item[0] else ""
        article = item[1] if len(item) > 1 and item[1] else ""
        clause = item[2] if len(item) > 2 and item[2] else ""
        gist = item[3] if len(item) > 3 and item[3] else ""
        if not law_name:
            continue

        # 법령 + 조항(+ 항)
        law_article = law_name
        if article:
            law_article += " " + article
        if clause:
            law_article += clause

        # 요지 추가
        if gist:
            lines.append(f"{law_article} — {gist}")
        else:
            lines.append(law_article)

    return "\n".join(lines)


def _format_basis_short(v: list) -> str:
    """근거 목록([법령, 조, 항, 요지] 목록)을 법조문 번호만 중복 없이 `, `로 이어 한 줄로 변환.

    요약 시트용: 요지 문장 없이 법령 조항만 추출하며, 중복 제거 후 한 줄로 연결한다.
    4번째 칸(요지)이 ``N호…``로 시작하면 ``N호``까지만 조항에 붙이고 나머지 문장은 버린다.

    항목별로 한 겹 더 감싼 근거 목록(wrapped basis list)도 같은 규칙으로 펼쳐 한 줄로 만든다.
    """
    import re as _re
    # 항목별로 한 겹 더 감싼 근거 목록이면 내부 항목만 평면화
    if _is_wrapped_basis_list(v):
        v = [inner for wrapper in v for inner in (wrapper if _is_basis_list(wrapper) else [])]
    seen: list[str] = []
    for item in v:
        if not item:
            continue
        law_name = item[0] if len(item) > 0 and item[0] else ""
        article = item[1] if len(item) > 1 and item[1] else ""
        clause = item[2] if len(item) > 2 and item[2] else ""
        gist = item[3] if len(item) > 3 and item[3] else ""
        if not law_name:
            continue
        law_article = law_name
        if article:
            law_article += " " + article
        if clause:
            law_article += clause
        # 요지가 'N호…'로 시작하면 호까지만 조항에 붙이고 나머지 문장은 버림.
        # 항이 비어 있으면 조와 호 사이 한 칸 띄움(예: "제17조의4 1호"), 항이 있으면 바로 붙임(예: "제20조①1호").
        if gist and _re.match(r"^\d+호", gist):
            ho = _re.match(r"^\d+호", gist).group(0)
            if clause:
                law_article += ho
            else:
                law_article += " " + ho
        if law_article not in seen:
            seen.append(law_article)
    return ", ".join(seen)


# ── 계산식 이름 치환 헬퍼 (SPEC_export_workbook_c) ─────────────────────────
import re as _re
from openpyxl.workbook.defined_name import DefinedName

_NAME_RE = _re.compile(r"[가-힣A-Za-z][가-힣A-Za-z0-9_]*")

_FUNCTIONS = frozenset({"MIN", "MAX", "IF", "ROUNDDOWN", "TRUE", "FALSE"})

_CELL_A1_RE = _re.compile(r"^[A-Za-z]{1,3}[0-9]+$")
_CELL_R1C1_RE = _re.compile(r"^[RrCc]$|^[Rr][0-9]*[Cc][0-9]*$")
_VALID_START_RE = _re.compile(r"^[가-힣A-Za-z_]$")
_SANITIZE_CHAIN = _re.compile(r"[^가-힣A-Za-z0-9_.]+")


def _sanitize_dname(raw: str, *, wb) -> str:
    """Excel 정의 이름 규칙에 맞게 이름 문자열을 정제한다.

    - 허용 문자(한글·영문·숫자·밑줄·마침표)만 남기고 나머지는 '_'로 치환.
    - 연속 '_'를 하나로 압축, 끝 '_' 제거.
    - 첫 글자가 한글·영문·밑줄이 아니면 앞에 '_'.
    - 셀 주소(A1, AB12, R, C, R1C1 등)로 읽히면 앞에 '_'.
    - 255자 초과 시 자름.
    - 이미 등록된 이름과 대소문자 무시로 겹치면 '_2', '_3'… 접미사.
    """
    s = _SANITIZE_CHAIN.sub("_", raw)
    s = _re.sub(r"_+", "_", s)
    s = s.rstrip("_")
    if not s:
        s = "_"
    if not _VALID_START_RE.match(s[0]):
        s = "_" + s
    if _CELL_A1_RE.match(s) or _CELL_R1C1_RE.match(s):
        s = "_" + s
    s = s[:255]
    key = s.lower()
    if key in wb.defined_names:
        base, ext = s, ""
        i = 2
        while f"{base}{ext}".lower() in wb.defined_names:
            ext = f"_{i}"
            i += 1
        s = f"{base}{ext}"
    return s


def _이름_치환(식: str, 이름_맵: dict[str, str]) -> str:
    """엑셀 수식 문법의 식에서 각 이름(식별자)을 정의 이름(호출N_이름)으로 치환한다.

    함수명(MIN/MAX/IF/ROUNDDOWN)과 TRUE/FALSE는 치환하지 않는다.
    괄호·쉼표·연산자·숫자는 그대로 보존하며, Excel이 연산자 우선순위를 해석한다.
    """
    tok: list[str] = []
    i, n = 0, len(식)
    while i < n:
        c = 식[i]
        if c in " \t":
            i += 1
            continue
        if c in "+-*/()":
            tok.append(c); i += 1; continue
        if c == ",":
            tok.append(","); i += 1; continue
        if c.isdigit() or (c == "." and i + 1 < n and 식[i + 1].isdigit()):
            j = i
            # 정수부
            while j < n and 식[j].isdigit():
                j += 1
            # 소수부
            if j < n and 식[j] == "." and j + 1 < n and 식[j + 1].isdigit():
                j += 1
                while j < n and 식[j].isdigit():
                    j += 1
            # 지수부: e/E 뒤에 숫자(또는 부호+숫자)가 있을 때만 포함
            if j < n and 식[j] in "eE":
                k = j + 1
                if k < n and 식[k] in "+-":
                    k += 1
                if k < n and 식[k].isdigit():
                    j = k
                    while j < n and 식[j].isdigit():
                        j += 1
                # e/E 뒤에 숫자(부호+숫자)가 없으면 거기서 숫자를 끝내고
                # e/E부터는 다음 반복에서 이름으로 읽는다
            tok.append(식[i:j]); i = j; continue
        m = _NAME_RE.match(식, i)
        if m:
            name = m.group(0)
            i = m.end()
            if name not in _FUNCTIONS:
                tok.append(이름_맵.get(name, name))
            else:
                tok.append(name)
        else:
            tok.append(c); i += 1
    return "".join(tok)


@mcp.tool(annotations=CALC)
def export_workbook(
    title: Annotated[str, Field(description="문서 제목 — 파일명 기본값")],
    calls: Annotated[list | str, Field(description="[{label, tool, args}] 1~30개 — JSON 문자열로 줘도 됨")],
    output_dir: Annotated[str, Field(description="내보낼 폴더 경로(기본: 현재 작업 폴더)")] = ".",
    filename: Annotated[str | None, Field(description="파일명 기본값(제목), 확장자 .xlsx 자동)")] = None,
    layout: Annotated[Literal["forms", "steps"], Field(description="'forms'(법정서식 틀 + 계산·입력가정·판단근거, 기본) 또는 'steps'(요약 + 계산 N)")] = "forms",
    company: Annotated[str, Field(description="법인명(서식 머리에 표시, 기본 빈 문자열)")] = "",
    period: Annotated[str, Field(description="사업연도 표기(예 '2025.01.01 ~ 2025.12.31', 기본 빈 문자열)")] = "",
    biz_no: Annotated[str, Field(description="사업자등록번호(서식 머리에 표시, 기본 빈 문자열)")] = "",
) -> dict:
    """계산 도구 결과를 엑셀(.xlsx)로 정리한다. 각 call은 서버가 직접 다시 실행하며, 결과·계산·세무조정·판단·산식을 탭으로 펼친다.
언제: 계산 결과(세무조정·소득처분·판단·산식)를 실무자가 검토할 수 있는 엑셀 파일로 남길 때.
입력:
  title — 문서 제목(파일명 기본값).
  calls — [{label, tool, args}] 목록(1~30개). tool 이름이 없거나 'export_workbook'이면 오류(재귀 금지).
  output_dir — 저장할 폴더(기본: '.'). 없으면 오류.
  filename — 파일명(기본: title, .xlsx 자동). '/ \\ : * ? " < > |'는 '_'로 대체.
  layout — 'forms'(기본: 서식 틀 시트 + 계산 N + 입력·가정 + 판단·근거) 또는 'steps'(요약 + 계산 N + 입력·가정 + 판단·근거).
  company — 법인명(서식 머리에 표시, 기본 '').
  period — 사업연도 표기(예 '2025.01.01 ~ 2025.12.31', 기본 '').
  biz_no — 사업자등록번호(서식 머리에 표시, 기본 '').

결과:
  결과.파일 — 절대경로.
  결과.시트 — 시트 이름 목록.
  결과.call_수 — calls 길이.
  결과.오류_call — 오류 발생한 call의 label 목록.
  근거 — '계산 도구 결과 정리(세무 자문 아님)'."""
    # ── calls가 문자열이면 JSON으로 풀기 ──
    if isinstance(calls, str):
        try:
            parsed = _json.loads(calls)
        except Exception:
            return {"error": "입력 확인 필요 — calls는 [{label, tool, args}] 목록(JSON 배열)", "안내": "calls가 올바른 JSON 배열이 아님"}
        if not isinstance(parsed, list):
            return {"error": "입력 확인 필요 — calls는 [{label, tool, args}] 목록(JSON 배열)", "안내": "calls가 JSON 객체 등으로 넘어옴 — 배열이어야 함"}
        calls = parsed

    # ── 입력 검증 ──
    if not calls:
        return {"error": "입력 확인 필요 — calls가 비어 있음", "안내": "1~30개 call을 주세요"}
    if len(calls) > 30:
        return {"error": "입력 확인 필요 — calls가 30개 초과", "안내": "최대 30개까지"}
    if output_dir is None or output_dir == "":
        return {"error": "입력 확인 필요 — output_dir 없음", "안내": "output_dir을 지정하세요"}
    if not os.path.isdir(output_dir):
        return {"error": "입력 확인 필요 — output_dir이 폴더가 아님", "안내": f"'{output_dir}'은(는) 폴더가 아님"}

    오류_call: list[str] = []
    결과들: list[dict] = []  # 각 call의 (label, tool, result_or_error, basis)

    # ── 각 call 실행 ──
    for call in calls:
        label = call.get("label", "")
        tool = call.get("tool", "")
        args = call.get("args", {})
        if isinstance(args, str):
            try:
                parsed_args = _json.loads(args)
            except Exception:
                오류_call.append(label)
                결과들.append({"label": label, "tool": tool,
                               "error": "입력 확인 필요 — call args가 올바른 JSON이 아님",
                               "basis": "", "result": None})
                continue
            if not isinstance(parsed_args, dict):
                오류_call.append(label)
                결과들.append({"label": label, "tool": tool,
                               "error": "입력 확인 필요 — call args가 JSON 객체(dict)가 아님",
                               "basis": "", "result": None})
                continue
            args = parsed_args
            call["args"] = args  # 뒷부분의 calls[i].get("args", …)도 dict를 받게 갱신

        if not tool or tool == "export_workbook":
            return {"error": "입력 확인 필요 — 도구 이름 없음 또는 export_workbook 재귀 호출 금지",
                    "안내": "calls에서 tool 이름을 확인하세요"}

        try:
            r = asyncio.run(mcp.call_tool(tool, args))
            raw = r.content[0].text if hasattr(r, "content") else r[0][0].text
            data = _json.loads(raw)
        except Exception as e:
            오류_call.append(label)
            결과들.append({"label": label, "tool": tool,
                           "error": f"입력 확인 필요 — 호출 오류: {e}",
                           "basis": "", "result": None})
            continue

        if "error" in data:
            오류_call.append(label)
            결과들.append({"label": label, "tool": tool,
                           "error": data.get("error", str(data)),
                           "basis": data.get("안내") or data.get("문제") and "\n".join(data["문제"]) or "",
                           "result": data,
                           "data": data})
            continue
        else:
            결과들.append({"label": label, "tool": tool,
                           "error": "",
                           "basis": data.get("근거", ""),
                           "result": data.get("결과", data),
                           "data": data})

    # ── 워크북 생성 ──
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side, numbers
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()

    # 번호 서식
    money_fmt = "#,##0"
    bold = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="D9E1F2")

    def _set_money(cell, val):
        cell.value = val
        cell.number_format = money_fmt

    # ── 폼 템플릿 로드 ──
    _FORM_TEMPLATES = _load_form_templates()

    시트이름들: list[str] = []

    if layout == "steps":
        # ── 1. 요약 시트 (steps 모드, 기존 동작) ──
        ws_sum = wb.active
        ws_sum.title = "요약"

        오늘 = date.today().isoformat()
        ws_sum["A1"] = "제목"; ws_sum["B1"] = title
        ws_sum["A2"] = "작성일"; ws_sum["B2"] = 오늘
        ws_sum["A3"] = "call 수"; ws_sum["B3"] = len(calls)
        for c in ("A1", "A2", "A3"):
            ws_sum[c].font = bold

        hdr_row = 4
        ws_sum.cell(hdr_row, 1, "label").font = bold
        ws_sum.cell(hdr_row, 2, "도구").font = bold
        첫데이터행 = hdr_row + 1

        for i, cd in enumerate(결과들):
            r = 첫데이터행 + i
            ws_sum.cell(r, 1, cd["label"])
            ws_sum.cell(r, 2, cd["tool"])
            c = 3
            if cd["error"]:
                ws_sum.cell(r, c, cd["error"]); c += 1
            else:
                for k, v in _numeric_items(cd["result"] or {}):
                    kc = ws_sum.cell(r, c, k); kc.font = bold
                    vc = ws_sum.cell(r, c + 1)
                    if isinstance(v, (int, float)):
                        _set_money(vc, v)
                    else:
                        vc.value = _cell_value(v)
                    c += 2
                if cd["basis"]:
                    if isinstance(cd["basis"], list) and (_is_basis_list(cd["basis"]) or _is_wrapped_basis_list(cd["basis"])):
                        bc = ws_sum.cell(r, c, _format_basis_short(cd["basis"]))
                    else:
                        bc = ws_sum.cell(r, c, _cell_value(cd["basis"]))
                    c += 1

        고지행 = 첫데이터행 + len(결과들) + 1
        ws_sum.cell(고지행, 1,
                    "이 파일은 계산 도구의 결과를 정리한 참고 자료이며 세무 자문이 아닙니다. 판단과 신고는 세무 전문가와 확인하세요.")
        ws_sum.cell(고지행, 1).font = Font(italic=True, color="C00000")
        ws_sum.merge_cells(start_row=고지행, start_column=1, end_row=고지행, end_column=8)

        ws_sum.column_dimensions["A"].width = 20
        ws_sum.column_dimensions["B"].width = 28
        for c in range(3, 20):
            ws_sum.column_dimensions[get_column_letter(c)].width = 16
        ws_sum.freeze_panes = "A5"

        시트이름들.append("요약")

        # ── 2. 계산 시트들 (steps 모드, 기존 동작) ──
        calc_idx = 0
        for cd in 결과들:
            if cd["error"] or not cd["result"]:
                continue
            calc_idx += 1
            sh_name = f"계산 {calc_idx}"
            ws = wb.create_sheet(sh_name)
            시트이름들.append(sh_name)
            res = cd["result"] or {}
            calc_dict = (cd.get("data") or {}).get("계산") if isinstance(cd.get("data"), dict) else None
            formula = (cd.get("data") or {}).get("계산식") if isinstance(cd.get("data"), dict) else None

            if formula:
                호출접두사 = f"호출{calc_idx}_"
                dname_map: dict[str, str] = {}
                for inp in (cd.get("data") or {}).get("계산식_입력", []) or []:
                    dname_map[inp["이름"]] = _sanitize_dname(f"{호출접두사}{inp['이름']}", wb=wb)
                for s in formula:
                    dname_map[s["이름"]] = _sanitize_dname(f"{호출접두사}{s['이름']}", wb=wb)

                행 = 1
                for cc, h in enumerate(["이름", "값", "설명"], 1):
                    c = ws.cell(행, cc, h); c.font = bold; c.fill = header_fill
                행 += 1
                for inp in (cd.get("data") or {}).get("계산식_입력", []) or []:
                    ws.cell(행, 1, inp["이름"])
                    vc = ws.cell(행, 2)
                    v = inp["값"]
                    if isinstance(v, bool):
                        vc.value = v
                    elif isinstance(v, (int, float)):
                        _set_money(vc, v)
                    else:
                        vc.value = v
                    _nm = dname_map[inp["이름"]]
                    wb.defined_names[_nm] = DefinedName(
                        name=_nm, attr_text=f"'{sh_name}'!$B${행}")
                    ws.cell(행, 3, inp.get("설명", ""))
                    행 += 1

                calc_hdr_row = 행
                for cc, h in enumerate(["단계", "수식", "도구 값", "확인", "설명"], 1):
                    c = ws.cell(행, cc, h); c.font = bold; c.fill = header_fill
                행 += 1
                for s in formula:
                    단계명 = s["이름"]
                    dname = dname_map[단계명]
                    ws.cell(행, 1, 단계명)
                    수식셀 = ws.cell(행, 2, f"={_이름_치환(s['식'], dname_map)}")
                    수식셀.number_format = money_fmt
                    wb.defined_names[dname] = DefinedName(
                        name=dname, attr_text=f"'{sh_name}'!$B${행}")
                    도구값 = s["값"]
                    vc = ws.cell(행, 3)
                    if isinstance(도구값, (int, float)):
                        _set_money(vc, 도구값)
                    else:
                        vc.value = 도구값
                    ws.cell(행, 4, f'=IF(B{행}=C{행},"일치","확인 필요")')
                    ws.cell(행, 5, s.get("설명", ""))
                    행 += 1
                next_table_row = 행

            행 = next_table_row if formula else 1
            ws.cell(행, 1, "항목").font = bold
            ws.cell(행, 2, "금액").font = bold
            ws.cell(행, 1).fill = header_fill
            ws.cell(행, 2).fill = header_fill
            행 += 1

            all_items: list[tuple[str, Any]] = []
            all_items.extend(_numeric_items(res))
            if isinstance(calc_dict, dict):
                all_items.extend(_numeric_items(calc_dict))

            항목맵: dict[str, Any] = {}
            for k, v in all_items:
                항목맵[k] = v

            for k, v in 항목맵.items():
                ws.cell(행, 1, k)
                if isinstance(v, (int, float)):
                    _set_money(ws.cell(행, 2), v)
                else:
                    ws.cell(행, 2, _cell_value(v))
                행 += 1

            정착_start = 행 + 1
            settlements = res.get("세무조정", [])
            합계_수식행 = None
            if settlements:
                행 = 정착_start
                ws.cell(행, 1, "구분").font = bold
                ws.cell(행, 2, "조정").font = bold
                ws.cell(행, 3, "금액").font = bold
                ws.cell(행, 4, "소득처분").font = bold
                for cc in range(1, 5):
                    ws.cell(행, cc).fill = header_fill
                행 += 1
                처분별: dict[str, int] = {}
                가산first_row = 행
                for s in settlements:
                    raw = s.get("금액", 0)
                    if isinstance(raw, (int, float)):
                        amt = raw; numeric = True
                    else:
                        amt = raw; numeric = False
                    raw_조정 = s.get("조정", "")
                    if raw_조정 in ("익금산입", "손금불산입"):
                        조정 = "가산"
                    elif raw_조정 in ("손금산입", "익금불산입"):
                        조정 = "차감"
                    elif raw_조정 == "":
                        처분_tmp = s.get("소득처분", "")
                        조정 = "차감" if 처분_tmp.startswith("△") else "가산"
                    else:
                        조정 = raw_조정
                    처분 = s.get("소득처분", "")
                    ws.cell(행, 1, s.get("구분", ""))
                    ws.cell(행, 2, 조정)
                    if numeric:
                        _set_money(ws.cell(행, 3), amt)
                    else:
                        ws.cell(행, 3, amt)
                    ws.cell(행, 4, 처분)
                    if numeric:
                        처분별[처분] = 처분별.get(처분, 0) + amt
                    행 += 1
                처분last_row = 행 - 1

                가산_sumif_cond = f"B{가산first_row}:B{처분last_row}"
                차감_sumif_cond = f"B{가산first_row}:B{처분last_row}"
                금액_sumif_addr = f"C{가산first_row}:C{처분last_row}"

                ws.cell(행, 1, "가산 합계").font = bold
                ws.cell(행, 3, f'=SUMIF({가산_sumif_cond},"가산",{금액_sumif_addr})')
                ws.cell(행, 3).number_format = money_fmt
                ws.cell(행, 3).font = bold
                가산_수식행 = 행; 행 += 1

                ws.cell(행, 1, "차감 합계").font = bold
                ws.cell(행, 3, f'=SUMIF({차감_sumif_cond},"차감",{금액_sumif_addr})')
                ws.cell(행, 3).number_format = money_fmt
                ws.cell(행, 3).font = bold
                차감_수식행 = 행; 행 += 1

                ws.cell(행, 1, "순액").font = bold
                ws.cell(행, 3, f"=C{가산_수식행}-C{차감_수식행}")
                ws.cell(행, 3).number_format = money_fmt
                ws.cell(행, 3).font = bold
                행 += 1

                ws.cell(행, 1, "금액 합계(방향 무관)").font = bold
                ws.cell(행, 3, f"=SUM({금액_sumif_addr})")
                ws.cell(행, 3).number_format = money_fmt
                ws.cell(행, 3).font = bold
                합계_수식행 = 행; 행 += 1

                ws.cell(행, 1, "소득처분별 합계").font = bold
                ws.cell(행, 2, "소득처분").font = bold
                ws.cell(행, 3, "금액").font = bold
                for cc in range(1, 4):
                    ws.cell(행, cc).fill = header_fill
                행 += 1
                for 처분, tot in 처분별.items():
                    ws.cell(행, 1, "")
                    ws.cell(행, 2, 처분)
                    sumif_cond = f"D{가산first_row}:D{처분last_row}"
                    ws.cell(행, 3, f'=SUMIF({sumif_cond},"{처분}",{금액_sumif_addr})')
                    ws.cell(행, 3).number_format = money_fmt
                    행 += 1

            if 합계_수식행 is not None:
                v_익금 = res.get("익금산입_손금불산입")
                v_손금 = res.get("손금산입_익금불산입")
                if isinstance(v_익금, (int, float)) and isinstance(v_손금, (int, float)):
                    ws.cell(행, 1, "확인 (익금산입_손금불산입)").font = bold
                    ws.cell(행, 3, f'=IF(C{가산_수식행}={v_익금},"일치","확인 필요")')
                    행 += 1
                    ws.cell(행, 1, "확인 (손금산입_익금불산입)").font = bold
                    ws.cell(행, 3, f'=IF(C{차감_수식행}={v_손금},"일치","확인 필요")')
                    행 += 1
                else:
                    도구합계 = None; 도구합계키 = None
                    for key in ("손금불산입_합계", "한도초과", "소득처분별_합계"):
                        v = res.get(key)
                        if isinstance(v, (int, float)):
                            도구합계 = v; 도구합계키 = key; break
                    if 도구합계 is not None:
                        ws.cell(행, 1, f"확인 ({도구합계키})").font = bold
                        ws.cell(행, 3, f'=IF(C{합계_수식행}={도구합계},"일치","확인 필요")')
                        행 += 1

            산식 = (cd.get("data") or {}).get("산식") or ""
            if 산식:
                행 += 1
                ws.cell(행, 1, "산식").font = bold
                ac = ws.cell(행, 2, _cell_value(산식))
                ac.alignment = Alignment(wrap_text=True)

            ws.column_dimensions["A"].width = 32
            ws.column_dimensions["B"].width = 20
            ws.column_dimensions["C"].width = 16
            ws.freeze_panes = "A2"

    else:  # layout == "forms"
        # ── 1. 틀 템플릿 목록 구축 (call 순서대로, 한 call의 모든 템플릿 포함) ──
        템플릿_목록: list[dict] = []
        _tpl_seen: dict[str, int] = {}
        _cd_by_tool: dict[str, list[int]] = {}  # tool → call 인덱스 목록
        _skipped_tpl: set[str] = set()  # when 조건으로 skipped된 템플릿 키
        for i, cd in enumerate(결과들):
            if cd["error"] or not cd["result"]:
                continue
            tool = cd.get("tool", "")
            tpl_list = _FORM_TEMPLATES.get(tool, [])
            if not tpl_list:
                continue
            _cd_by_tool.setdefault(tool, []).append(i)
            for tpl in tpl_list:
                tkey = tpl["key"]
                # when 조건이 있으면 계산식_입력/결과 값으로 평가하여 거짓이면 렌더하지 않음
                when = tpl.get("when")
                if when:
                    계산식_입력 = (cd.get("data") or {}).get("계산식_입력", [])
                    입력_값 = {inp["이름"]: inp["값"] for inp in 계산식_입력}
                    결과_값 = cd.get("result", {}) or {}
                    # [입력.이름] 구문을 입력_값["이름"]으로, [결과.키] 구문을 결과_값["키"]로 치환
                    # FALSE/TRUE는 Python bool로 변환
                    def _when_repl(m):
                        prefix = m.group(1)  # "입력.이름" 또는 "결과.키"
                        if prefix.startswith("입력."):
                            return f"입력_값[{prefix[3:]!r}]"
                        elif prefix.startswith("결과."):
                            return f"결과_값[{prefix[3:]!r}]"
                        return f"{prefix}"
                    when_py = _re.sub(r"\[([^]]+)\]", _when_repl, when)
                    when_py = when_py.replace("FALSE", "False").replace("TRUE", "True")
                    if not eval(when_py, {"__builtins__": {}}, {"입력_값": 입력_값, "결과_값": 결과_값}):
                        _skipped_tpl.add(tkey)
                        continue
                sheet_name = tpl["sheet"]
                _tpl_seen[tkey] = _tpl_seen.get(tkey, 0) + 1
                if _tpl_seen[tkey] > 1:
                    sheet_name = f"{tpl['sheet']} ({_tpl_seen[tkey]})"
                템플릿_목록.append({"key": tkey, "sheet": sheet_name, "form": tpl["form"], "tool": tool, "call_idx": i})

        # 템플릿_목록을 spec 순서(별지3, 별지15, 별지23갑, 별지23을, 별지20정률, 별지16의2, 별지50을,
        # 부가22, 부가근거63, 부가근거64, 종소배당, 종소근로)로 정렬
        _tpl_render_order = {"별지3": 1, "별지15": 2, "별지23갑": 3, "별지23을": 4,
                             "별지20정률": 5, "별지16의2": 6, "별지26갑": 6.1, "별지26을": 6.2, "별지50을": 7,
                             "부가22": 8, "부가근거63": 9, "부가근거64": 10, "종소배당": 11, "종소근로": 12,
                             "별지84": 13, "별지84세액": 14, "별지84부표1": 15}
        템플릿_목록.sort(key=lambda t: _tpl_render_order.get(t["key"], 99))

        # ── 1. 목록 시트 ──
        ws_list = wb.active
        ws_list.title = "목록"
        시트이름들.append("목록")
        행 = 1

        ws_list.cell(행, 1, "서식 시트").font = bold
        ws_list.cell(행, 2, "서식 표기").font = bold
        for cc in range(1, 3):
            ws_list.cell(행, cc).fill = header_fill
        행 += 1

        for tinfo in 템플릿_목록:
            ws_list.cell(행, 1, tinfo["sheet"])
            ws_list.cell(행, 2, tinfo["form"])
            행 += 1

        # 자료 원칙에 따른 줄: 법인세 도구 call이 하나라도 있으면
        has_corp_tool = any(
            cd.get("tool") in ("entertainment_adjustment", "income_adjustment_total",
                                "depreciation_adjustment", "dividend_received_deduction",
                                "nonbusiness_interest_disallowance")
            for cd in 결과들 if not cd.get("error")
        )
        if has_corp_tool:
            ws_list.cell(행, 1, "별지 제50호서식(갑)")
            ws_list.cell(행, 2, "자본금과 적립금조정명세서 — 자료 없음(재무상태표 자본 자료 필요)")
            행 += 1
            ws_list.cell(행, 1, "별지 제50호서식(병)")
            ws_list.cell(행, 2, "해당 없음(해외현지기업 출자)")
            행 += 1

        # 틀의 `목록`([[이름, 설명]]) — 서식 중 이번에 만들지 않은 부분(자료 없음 등)을 렌더된 틀마다 한 번
        _목록_본: set[str] = set()
        for tinfo in 템플릿_목록:
            _t = next((t for t in _FORM_TEMPLATES.get(tinfo["tool"], []) if t["key"] == tinfo["key"]), None)
            for 이름, 설명 in (_t or {}).get("목록", []):
                if 이름 in _목록_본:
                    continue
                _목록_본.add(이름)
                ws_list.cell(행, 1, 이름)
                ws_list.cell(행, 2, 설명)
                행 += 1

        # when 조건으로 만들지 않은 틀에 `목록_생략` 문구가 있으면 그 줄
        for _tool_tpls in _FORM_TEMPLATES.values():
            for _t in _tool_tpls:
                if _t["key"] in _skipped_tpl and _t.get("목록_생략") and _t["sheet"] not in _목록_본:
                    _목록_본.add(_t["sheet"])
                    ws_list.cell(행, 1, _t["sheet"])
                    ws_list.cell(행, 2, _t["목록_생략"])
                    행 += 1

        # 종소배당이 when 조건으로 skipped된 경우 목록 줄에 해당 없음 추가
        if "종소배당" in _skipped_tpl:
            ws_list.cell(행, 1, "종소 배당·이자소득명세서")
            ws_list.cell(행, 2, "❻ 배당소득명세서·❺ 이자소득명세서 — 해당 없음(금융소득 2천만원 이하 분리과세)")
            행 += 1

        ws_list.cell(행, 1,
                    "이 파일은 계산 도구의 결과를 정리한 참고 자료이며 세무 자문이 아닙니다. 판단과 신고는 세무 전문가와 확인하세요.")
        ws_list.cell(행, 1).font = Font(italic=True, color="C00000")
        ws_list.merge_cells(start_row=행, start_column=1, end_row=행, end_column=8)

        ws_list.column_dimensions["A"].width = 30
        ws_list.column_dimensions["B"].width = 55
        ws_list.freeze_panes = "A2"

        # ── 2. 입력·가정 시트 생성 (템플릿 렌더링 전에 정의된 이름 생성) ──
        ws_inp = wb.create_sheet("입력·가정")
        시트이름들.append("입력·가정")
        행 = 1
        ws_inp.cell(행, 1, "call").font = bold
        ws_inp.cell(행, 2, "경로").font = bold
        ws_inp.cell(행, 3, "값").font = bold
        for cc in range(1, 4):
            ws_inp.cell(행, cc).fill = header_fill
        행 += 1
        for i, cd in enumerate(결과들):
            if cd["error"]:
                continue
            label = cd["label"]
            args = (calls[i] or {}).get("args", {})
            for 경로, 값 in _flatten_args(args):
                ws_inp.cell(행, 1, label)
                ws_inp.cell(행, 2, 경로)
                vc = ws_inp.cell(행, 3)
                if isinstance(값, (int, float)):
                    _set_money(vc, 값)
                else:
                    vc.value = _cell_value(값)
                행 += 1
        # forms 모드: 계산식_입력·계산식 단계별 정의된 이름 생성 (입력·가정 시트 아래)
        행 += 1
        for i, cd in enumerate(결과들):
            if cd["error"]:
                continue
            call_idx = i
            호출접두사 = f"호출{call_idx + 1}_"
            for inp in (cd.get("data") or {}).get("계산식_입력", []) or []:
                ws_inp.cell(행, 1, 호출접두사 + inp["이름"])
                vc = ws_inp.cell(행, 2)
                v = inp["값"]
                if isinstance(v, bool):
                    vc.value = v
                elif isinstance(v, (int, float)):
                    _set_money(vc, v)
                else:
                    vc.value = v
                _nm = _sanitize_dname(f"{호출접두사}{inp['이름']}", wb=wb)
                wb.defined_names[_nm] = DefinedName(
                    name=_nm, attr_text=f"'입력·가정'!$B${행}")
                ws_inp.cell(행, 3, inp.get("설명", ""))
                행 += 1
            for s in (cd.get("data") or {}).get("계산식", []) or []:
                ws_inp.cell(행, 1, 호출접두사 + s["이름"])
                vc = ws_inp.cell(행, 2)
                v = s["값"]
                if isinstance(v, bool):
                    vc.value = v
                elif isinstance(v, (int, float)):
                    _set_money(vc, v)
                else:
                    vc.value = v
                _nm = _sanitize_dname(f"{호출접두사}{s['이름']}", wb=wb)
                wb.defined_names[_nm] = DefinedName(
                    name=_nm, attr_text=f"'입력·가정'!$B${행}")
                ws_inp.cell(행, 3, s.get("설명", ""))
                행 += 1
            # 결과 dict 항목도 정의된 이름으로 등록 (템플릿 [결과.경로] 참조용)
            res = cd.get("result", {}) or {}
            for k, v in res.items():
                if isinstance(v, (int, float)):
                    ws_inp.cell(행, 1, 호출접두사 + "결과_" + k)
                    vc = ws_inp.cell(행, 2)
                    _set_money(vc, v)
                    _nm = _sanitize_dname(f"{호출접두사}결과_{k}", wb=wb)
                    wb.defined_names[_nm] = DefinedName(
                        name=_nm, attr_text=f"'입력·가정'!$B${행}")
                    행 += 1
            # calc_dict(=계산 단계 값) 항목도 정의된 이름으로 등록 (템플릿 [계산.경로] 참조용)
            calc_dict = (cd.get("data") or {}).get("계산") if isinstance((cd.get("data") or {}), dict) else None
            if isinstance(calc_dict, dict):
                for k, v in calc_dict.items():
                    if isinstance(v, (int, float)) and not (isinstance(v, float) and v == float("inf")):
                        ws_inp.cell(행, 1, 호출접두사 + k)
                        vc = ws_inp.cell(행, 2)
                        _set_money(vc, v)
                        _nm = _sanitize_dname(f"{호출접두사}{k}", wb=wb)
                        wb.defined_names[_nm] = DefinedName(
                            name=_nm, attr_text=f"'입력·가정'!$B${행}")
                        ws_inp.cell(행, 3, "")
                        행 += 1

        # ── 3. 틀 시트 렌더링 ──
        sheet_name_map: dict[str, str] = {}
        slot_map_all: dict[str, str] = {}
        _tpl_by_key: dict[str, dict] = {}
        for tinfo in 템플릿_목록:
            sheet_name_map[tinfo["key"]] = tinfo["sheet"]
            if tinfo["key"] not in _tpl_by_key:
                for _tpl in _FORM_TEMPLATES.get(tinfo["tool"], []):
                    if _tpl["key"] == tinfo["key"]:
                        _tpl_by_key[tinfo["key"]] = _tpl
                        break
            tpl = _tpl_by_key.get(tinfo["key"])
            if tpl is None:
                continue
            for s in tpl.get("slots", []):
                sid = s["id"]
                cell = s["cell"]
                slot_map_all[sid] = cell

        # slot_map_cross: {(ref_key, ref_id): (sheet, cell)} — cross 참조 lookup용
        # 1단계: 모든 템플릿의 slots 수집 (tkey → {sid: (sheet, cell)})
        all_slots: dict[str, dict[str, tuple[str, str]]] = {}
        for tinfo in 템플릿_목록:
            tpl = _tpl_by_key[tinfo["key"]]
            all_slots[tinfo["key"]] = {}
            for s in tpl.get("slots", []):
                all_slots[tinfo["key"]][s["id"]] = (tinfo["sheet"], s["cell"])

        # tool: "*" (별지50을)는 유보 행이 1개 이상일 때 추가
        # (income_adjustment_total의 세무조정에서 유보 + depreciation 유보 행 포함)
        _has_유보 = any(
            s.get("소득처분") == "유보"
            for cd in 결과들 if not cd.get("error")
            for s in (cd.get("result") or {}).get("세무조정", [])
        ) or ("별지20정률" in all_slots)
        if _FORM_TEMPLATES.get("*", []) and _has_유보 and "별지50을" not in _tpl_seen:
            _tpl_seen["별지50을"] = 1
            별지50_tpl = _FORM_TEMPLATES["*"][0]
            sheet_name = 별지50_tpl["sheet"]
            adj_call_idx = next((i for i, cd in enumerate(결과들)
                                if cd.get("tool") == "income_adjustment_total" and not cd.get("error")), 0)
            tinfo_50 = {"key": "별지50을", "sheet": sheet_name, "form": 별지50_tpl["form"],
                       "tool": "*", "call_idx": adj_call_idx}
            템플릿_목록.append(tinfo_50)
            sheet_name_map["별지50을"] = sheet_name
            _tpl_by_key["별지50을"] = 별지50_tpl
            all_slots["별지50을"] = {}
            for s in 별지50_tpl.get("slots", []):
                all_slots["별지50을"][s["id"]] = (sheet_name, s["cell"])
            # 별지50을 추가 후 다시 정렬하여 올바른 순서 유지
            템플릿_목록.sort(key=lambda t: _tpl_render_order.get(t["key"], 99))

        # adj_slots 구축: adjustment → [(시트명, 셀, check_path, call_idx)]
        adj_slots: dict[str, list[tuple[str, str, str, int]]] = {}
        for tkey, slots in all_slots.items():
            tpl = _tpl_by_key[tkey]
            for s in tpl.get("slots", []):
                adj = s.get("adjustment")
                if adj:
                    sid = s["id"]
                    시트명, 셀 = all_slots[tkey][sid]
                    # call_idx 찾기: 템플릿_목록에서 해당 tinfo의 call_idx
                    call_idx = next((ti["call_idx"] for ti in 템플릿_목록 if ti["key"] == tkey), 0)
                    adj_slots.setdefault(adj, []).append((시트명, 셀, s.get("check", ""), call_idx))

        # 2단계: {틀key!id} cross 참조 파싱 (slot 수식 + table 열 수식)
        cross_refs: set[tuple[str, str]] = set()
        for tinfo in 템플릿_목록:
            tpl = _tpl_by_key[tinfo["key"]]
            for s in tpl.get("slots", []):
                f = s.get("f", "") or ""
                for m in _re.finditer(r"\{([^}]+)\}", f):
                    rest = m.group(1)
                    if "!" in rest:
                        ref_key, ref_id = rest.split("!", 1)
                        cross_refs.add((ref_key, ref_id))
            for tbl in tpl.get("tables", []):
                for col_spec in tbl.get("cols", {}).values():
                    if isinstance(col_spec, str):
                        for m in _re.finditer(r"\{([^}]+)\}", col_spec):
                            rest = m.group(1)
                            if "!" in rest:
                                ref_key, ref_id = rest.split("!", 1)
                                cross_refs.add((ref_key, ref_id))

        # 3단계: slot_map_cross 구성 (대상이 없는 cross 참조는 제외 — 렌더러에서 fallback 처리)
        slot_map_cross: dict[tuple[str, str], tuple[str, str]] = {}
        for ref_key, ref_id in cross_refs:
            if ref_key in all_slots and ref_id in all_slots[ref_key]:
                slot_map_cross[(ref_key, ref_id)] = all_slots[ref_key][ref_id]
            # 대상이 없으면 제외 (render_form_sheet에서 fallback 값 사용)

        # ── source_data 구축 (테이블 소스 + cross fallback용) ──
        source_data: dict[str, list[dict]] = {}
        # adjustment 소스: income_adjustment_total의 세무조정에서 가산·차감 분리
        adj_result = 결과들[3].get("result") if len(결과들) > 3 else {}
        adj_items = (adj_result or {}).get("세무조정", [])
        _adj_src_add: list[dict] = []
        _adj_src_sub: list[dict] = []
        for s in adj_items:
            조정 = s.get("조정", "")
            금액 = s.get("금액", 0)
            if isinstance(금액, (int, float)):
                if 조정 in ("익금산입", "손금불산입"):
                    처분 = s.get("소득처분", "")
                    코드 = "400" if 처분 == "유보" else ("100" if 처분 == "△유보" else "500")
                    _adj_src_add.append({"과목": s.get("구분", ""), "금액": 금액, "처분": 처분, "코드": 코드,
                                         "종류": s.get("종류", s.get("구분", ""))})
                elif 조정 in ("손금산입", "익금불산입"):
                    처분 = s.get("소득처분", "")
                    코드 = "100" if 처분 == "△유보" else "200"
                    _adj_src_sub.append({"과목": s.get("구분", ""), "금액": 금액, "처분": 처분, "코드": 코드,
                                         "종류": s.get("종류", s.get("구분", ""))})
        source_data["조정_가산"] = _adj_src_add
        source_data["조정_차감"] = _adj_src_sub
        # 피출자법인 소스: dividend_received_deduction 결과 → SPEC대로 {이름, 지분율(%)=출자비율×100}
        div_result = 결과들[2].get("result") if len(결과들) > 2 else {}
        _피출자 = []
        for inv in (div_result or {}).get("피출자법인별", []):
            _피출자.append({
                "이름": inv.get("이름", ""),
                "지분율(%)": inv.get("출자비율", 0) * 100,
            })
        source_data["피출자법인"] = _피출자
        # 유보 소스: depreciation_adjustment (별지20정률로 렌더된 경우) + income_adjustment_total 유보
        _유보_rows: list[dict] = []
        dep_result = 결과들[1].get("result") if len(결과들) > 1 else {}
        if "별지20정률" in all_slots:
            _유보_rows.append({
                "과목": "감가상각비 한도초과(기계장치)",
                "기초": "={별지20정률!⑪}",
                "감소": "={별지20정률!㉗}",
                "증가": "={별지20정률!㉖}",
                "비고": "",
            })
        for s in adj_items:
            처분 = s.get("소득처분", "")
            if 처분 == "유보":
                adj_key = s.get("종류", s.get("구분", ""))
                is_손금_유보 = s.get("조정") == "손금산입"
                if adj_key == "감가상각비 한도초과":
                    continue  # 위 depreciation 행에서 처리됨
                _유보_rows.append({
                    "과목": adj_key,
                    "기초": 0,
                    "감소": 0 if not is_손금_유보 else s.get("금액", 0),
                    "증가": s.get("금액", 0) if not is_손금_유보 else 0,
                    "비고": "",
                })
        source_data["유보"] = _유보_rows
        # 공통사용재화 소스: common_input_tax args common_asset_sales → {순번: 1..}
        for i, cd in enumerate(결과들):
            if cd.get("tool") == "common_input_tax" and not cd.get("error"):
                args = calls[i].get("args", {})
                common_asset_sales = args.get("common_asset_sales", [])
                _재화_rows = []
                for idx, sale in enumerate(common_asset_sales, 1):
                    _재화_rows.append({
                        "순번": idx,
                        "공급가액": int(sale.get("sale_price", 0)),
                    })
                source_data["공통사용재화"] = _재화_rows
                break

        # 배당명세·이자명세 소스: financial_income_aggregate 계산.항목별 → {코드, 일련번호, 항목, 금액}
        for i, cd in enumerate(결과들):
            if cd.get("tool") == "financial_income_aggregate" and not cd.get("error"):
                calc_dict = (cd.get("data") or {}).get("계산", {})
                항목별 = calc_dict.get("항목별", [])
                판단목록 = (cd.get("data") or {}).get("판단", [])
                # 판단목록에서 입력 순서 복구 (계산.항목별은 dict라 순서 없음)
                순서맵 = {h["항목"]: idx for idx, h in enumerate(판단목록)}
                _배당_rows: list[dict] = []
                _이자_rows: list[dict] = []
                # 코드 매핑: 배당 GU대상=21, GU비대상=22, 출자공동=28, 이자=13, 제외는 제외
                코드맵 = {"GU대상": "21", "GU비대상": "22", "출자공동": "28", "이자": "13", "제외": None}
                for item in 항목별:
                    구분 = item.get("구분", "")
                    코드 = 코드맵.get(구분)
                    if 코드 is None:
                        continue  # 제외는 넣지 않음
                    항목명 = item.get("항목", "").replace("_", " ")
                    금액 = item.get("금액", 0)
                    if 구분 in ("GU대상", "GU비대상", "출자공동"):
                        _배당_rows.append({"코드": 코드, "일련번호": 0, "항목": 항목명, "금액": 금액,
                                           "_순서": 순서맵.get(item.get("항목", ""), 999)})
                    elif 구분 == "이자":
                        _이자_rows.append({"코드": 코드, "일련번호": 0, "항목": 항목명, "금액": 금액,
                                            "_순서": 순서맵.get(item.get("항목", ""), 999)})
                # 코드 오름차순, 같은 코드 안에서는 입력 순서(_순서)로 정렬
                _배당_rows.sort(key=lambda r: (r["코드"], r["_순서"]))
                _이자_rows.sort(key=lambda r: (r["코드"], r["_순서"]))
                # 일련번호는 코드별 1부터
                last_code = None
                seq = 0
                for r in _배당_rows:
                    if r["코드"] != last_code:
                        last_code = r["코드"]
                        seq = 1
                    else:
                        seq += 1
                    r["일련번호"] = seq
                last_code = None
                seq = 0
                for r in _이자_rows:
                    if r["코드"] != last_code:
                        last_code = r["코드"]
                        seq = 1
                    else:
                        seq += 1
                    r["일련번호"] = seq
                # _순서 필드 제거
                for r in _배당_rows:
                    del r["_순서"]
                for r in _이자_rows:
                    del r["_순서"]
                source_data["배당명세"] = _배당_rows
                source_data["이자명세"] = _이자_rows
                break

        # connection: adjustment 소스 행의 금액을 렌더된 서식 slot 참조로 치환
        for src_name in ("조정_가산", "조정_차감"):
            rows = source_data.get(src_name, [])
            for row in rows:
                kind = row.get("종류", row.get("과목", ""))
                amount = row.get("금액", 0)
                if isinstance(amount, (int, float)) and kind in adj_slots:
                    for 시트명, 셀, check_path, call_idx in adj_slots[kind]:
                        cd = 결과들[call_idx]
                        # cd["data"]에는 nested "결과"가 있음 (템플릿 [결과.경로] 참조용)
                        res = cd.get("data") or {}
                        tv = tool_value(res, check_path)
                        if tv == amount:
                            row["금액"] = f"='{시트명}'!{셀}"
                            break

        for tinfo in 템플릿_목록:
            cd = 결과들[tinfo["call_idx"]]
            tkey = tinfo["key"]
            tpl = _tpl_by_key.get(tkey)
            if tpl is None:
                continue

            ws = wb.create_sheet(tinfo["sheet"])
            input_map: dict[str, Any] = {}
            for inp in (cd.get("data") or {}).get("계산식_입력", []) or []:
                input_map[inp["이름"]] = inp["값"]
            for s in (cd.get("data") or {}).get("계산식", []) or []:
                input_map[s["이름"]] = s["값"]

            # {{과세기간}} 구성 (부가22 템플릿용): year·period args로 "{year}년 제N기 ( ... )"
            과세기간 = ""
            if tkey == "부가22":
                for i, c in enumerate(calls):
                    if c.get("tool") == "common_input_tax" and not 결과들[i].get("error"):
                        args = c.get("args", {})
                        year = args.get("year", "")
                        period = args.get("period", "")
                        if period == "1기":
                            과세기간 = f"{year}년 제1기 ( 1월 1일 ~ 6월 30일 )"
                        elif period == "2기":
                            과세기간 = f"{year}년 제2기 ( 7월 1일 ~ 12월 31일 )"
                        break

            # 별지26 (갑)·(을): 표 source는 그 call 결과로(같은 도구 call이 여럿이면 call마다 따로)
            call_sources = source_data
            form_res = cd.get("data", cd.get("result", {}))
            if tinfo["tool"] == "nonbusiness_interest_disallowance":
                form_res = nonbusiness_form_data(form_res, (calls[tinfo["call_idx"]] or {}).get("args", {}))
                call_sources = {**source_data, **nonbusiness_form_sources(form_res)}
            # 교차 참조는 같은 call에서 만든 시트를 가리킨다(같은 도구 call이 둘 이상이면 ' (2)' 시트끼리)
            cross_this = dict(slot_map_cross)
            for (ref_key, ref_id) in slot_map_cross:
                same = next((t for t in 템플릿_목록 if t["key"] == ref_key and t["call_idx"] == tinfo["call_idx"]), None)
                if same is not None:
                    cross_this[(ref_key, ref_id)] = (same["sheet"], all_slots[ref_key][ref_id][1])
            render_form_sheet(ws, tpl, form_res, input_map,
                               sheet_name_map, slot_map_all, cross_this, tkey,
                               call_idx=tinfo["call_idx"], company=company, period=period, biz_no=biz_no,
                               source_data=call_sources, all_slots=all_slots,
                               label=cd.get("label", ""), 과세기간=과세기간)
            시트이름들.append(tinfo["sheet"])

        # ── 3b. 계산 N 시트들 (틀 없는 도구들, call 순서대로) ──
        calc_idx = 0
        for cd in 결과들:
            if cd["error"] or not cd["result"]:
                continue
            tool = cd.get("tool", "")
            if tool in _FORM_TEMPLATES:
                continue  # 틀 있는 도구는 이미 처리됨
            tpl_list = _FORM_TEMPLATES.get(tool, [])
            if tpl_list:
                continue  # 틀 있는 도구
            calc_idx += 1
            sh_name = f"계산 {calc_idx}"
            ws = wb.create_sheet(sh_name)
            시트이름들.append(sh_name)
            res = cd.get("result", {}) or {}
            calc_dict = (cd.get("data") or {}).get("계산") if isinstance(cd.get("data"), dict) else None
            formula = (cd.get("data") or {}).get("계산식") if isinstance(cd.get("data"), dict) else None

            # 계산 N 시트 내용 — steps 모드와 동일한 레이아웃:
            #   계산식_입력 표 → 계산식 단계 수식·도구 값·확인 열 → 결과·계산 키-값 →
            #   세무조정 표 → 확인 수식 → 산식
            next_table_row = 1
            if formula:
                호출접두사 = f"호출{calc_idx}_"
                dname_map: dict[str, str] = {}
                for inp in (cd.get("data") or {}).get("계산식_입력", []) or []:
                    dname_map[inp["이름"]] = _sanitize_dname(f"{호출접두사}{inp['이름']}", wb=wb)
                for s in formula:
                    dname_map[s["이름"]] = _sanitize_dname(f"{호출접두사}{s['이름']}", wb=wb)

                행 = 1
                for cc, h in enumerate(["이름", "값", "설명"], 1):
                    c = ws.cell(행, cc, h); c.font = bold; c.fill = header_fill
                행 += 1
                for inp in (cd.get("data") or {}).get("계산식_입력", []) or []:
                    ws.cell(행, 1, inp["이름"])
                    vc = ws.cell(행, 2)
                    v = inp["값"]
                    if isinstance(v, bool):
                        vc.value = v
                    elif isinstance(v, (int, float)):
                        _set_money(vc, v)
                    else:
                        vc.value = v
                    _nm = dname_map[inp["이름"]]
                    wb.defined_names[_nm] = DefinedName(
                        name=_nm, attr_text=f"'{sh_name}'!$B${행}")
                    ws.cell(행, 3, inp.get("설명", ""))
                    행 += 1

                for cc, h in enumerate(["단계", "수식", "도구 값", "확인", "설명"], 1):
                    c = ws.cell(행, cc, h); c.font = bold; c.fill = header_fill
                행 += 1
                for s in formula:
                    단계명 = s["이름"]
                    dname = dname_map[단계명]
                    ws.cell(행, 1, 단계명)
                    수식셀 = ws.cell(행, 2, f"={_이름_치환(s['식'], dname_map)}")
                    수식셀.number_format = money_fmt
                    wb.defined_names[dname] = DefinedName(
                        name=dname, attr_text=f"'{sh_name}'!$B${행}")
                    도구값 = s["값"]
                    vc = ws.cell(행, 3)
                    if isinstance(도구값, (int, float)):
                        _set_money(vc, 도구값)
                    else:
                        vc.value = 도구값
                    ws.cell(행, 4, f'=IF(B{행}=C{행},"일치","확인 필요")')
                    ws.cell(행, 5, s.get("설명", ""))
                    행 += 1
                next_table_row = 행

            행 = next_table_row
            ws.cell(행, 1, "항목").font = bold
            ws.cell(행, 2, "금액").font = bold
            ws.cell(행, 1).fill = header_fill
            ws.cell(행, 2).fill = header_fill
            행 += 1

            all_items: list[tuple[str, Any]] = []
            all_items.extend(_numeric_items(res))
            if isinstance(calc_dict, dict):
                all_items.extend(_numeric_items(calc_dict))

            항목맵: dict[str, Any] = {}
            for k, v in all_items:
                항목맵[k] = v

            for k, v in 항목맵.items():
                ws.cell(행, 1, k)
                if isinstance(v, (int, float)):
                    _set_money(ws.cell(행, 2), v)
                else:
                    ws.cell(행, 2, _cell_value(v))
                행 += 1

            정착_start = 행 + 1
            settlements = res.get("세무조정", [])
            합계_수식행 = None
            if settlements:
                행 = 정착_start
                ws.cell(행, 1, "구분").font = bold
                ws.cell(행, 2, "조정").font = bold
                ws.cell(행, 3, "금액").font = bold
                ws.cell(행, 4, "소득처분").font = bold
                for cc in range(1, 5):
                    ws.cell(행, cc).fill = header_fill
                행 += 1
                처분별: dict[str, int] = {}
                가산first_row = 행
                for s in settlements:
                    raw = s.get("금액", 0)
                    if isinstance(raw, (int, float)):
                        amt = raw; numeric = True
                    else:
                        amt = raw; numeric = False
                    raw_조정 = s.get("조정", "")
                    if raw_조정 in ("익금산입", "손금불산입"):
                        조정 = "가산"
                    elif raw_조정 in ("손금산입", "익금불산입"):
                        조정 = "차감"
                    elif raw_조정 == "":
                        처분_tmp = s.get("소득처분", "")
                        조정 = "차감" if 처분_tmp.startswith("△") else "가산"
                    else:
                        조정 = raw_조정
                    처분 = s.get("소득처분", "")
                    ws.cell(행, 1, s.get("구분", ""))
                    ws.cell(행, 2, 조정)
                    if numeric:
                        _set_money(ws.cell(행, 3), amt)
                    else:
                        ws.cell(행, 3, amt)
                    ws.cell(행, 4, 처분)
                    if numeric:
                        처분별[처분] = 처분별.get(처분, 0) + amt
                    행 += 1
                처분last_row = 행 - 1

                가산_sumif_cond = f"B{가산first_row}:B{처분last_row}"
                차감_sumif_cond = f"B{가산first_row}:B{처분last_row}"
                금액_sumif_addr = f"C{가산first_row}:C{처분last_row}"

                ws.cell(행, 1, "가산 합계").font = bold
                ws.cell(행, 3, f'=SUMIF({가산_sumif_cond},"가산",{금액_sumif_addr})')
                ws.cell(행, 3).number_format = money_fmt
                ws.cell(행, 3).font = bold
                가산_수식행 = 행; 행 += 1

                ws.cell(행, 1, "차감 합계").font = bold
                ws.cell(행, 3, f'=SUMIF({차감_sumif_cond},"차감",{금액_sumif_addr})')
                ws.cell(행, 3).number_format = money_fmt
                ws.cell(행, 3).font = bold
                차감_수식행 = 행; 행 += 1

                ws.cell(행, 1, "순액").font = bold
                ws.cell(행, 3, f"=C{가산_수식행}-C{차감_수식행}")
                ws.cell(행, 3).number_format = money_fmt
                ws.cell(행, 3).font = bold
                행 += 1

                ws.cell(행, 1, "금액 합계(방향 무관)").font = bold
                ws.cell(행, 3, f"=SUM({금액_sumif_addr})")
                ws.cell(행, 3).number_format = money_fmt
                ws.cell(행, 3).font = bold
                합계_수식행 = 행; 행 += 1

                ws.cell(행, 1, "소득처분별 합계").font = bold
                ws.cell(행, 2, "소득처분").font = bold
                ws.cell(행, 3, "금액").font = bold
                for cc in range(1, 4):
                    ws.cell(행, cc).fill = header_fill
                행 += 1
                for 처분, tot in 처분별.items():
                    ws.cell(행, 1, "")
                    ws.cell(행, 2, 처분)
                    sumif_cond = f"D{가산first_row}:D{처분last_row}"
                    ws.cell(행, 3, f'=SUMIF({sumif_cond},"{처분}",{금액_sumif_addr})')
                    ws.cell(행, 3).number_format = money_fmt
                    행 += 1

            if 합계_수식행 is not None:
                v_익금 = res.get("익금산입_손금불산입")
                v_손금 = res.get("손금산입_익금불산입")
                if isinstance(v_익금, (int, float)) and isinstance(v_손금, (int, float)):
                    ws.cell(행, 1, "확인 (익금산입_손금불산입)").font = bold
                    ws.cell(행, 3, f'=IF(C{가산_수식행}={v_익금},"일치","확인 필요")')
                    행 += 1
                    ws.cell(행, 1, "확인 (손금산입_익금불산입)").font = bold
                    ws.cell(행, 3, f'=IF(C{차감_수식행}={v_손금},"일치","확인 필요")')
                    행 += 1
                else:
                    도구합계 = None; 도구합계키 = None
                    for key in ("손금불산입_합계", "한도초과", "소득처분별_합계"):
                        v = res.get(key)
                        if isinstance(v, (int, float)):
                            도구합계 = v; 도구합계키 = key; break
                    if 도구합계 is not None:
                        ws.cell(행, 1, f"확인 ({도구합계키})").font = bold
                        ws.cell(행, 3, f'=IF(C{합계_수식행}={도구합계},"일치","확인 필요")')
                        행 += 1

            산식 = (cd.get("data") or {}).get("산식") or ""
            if 산식:
                행 += 1
                ws.cell(행, 1, "산식").font = bold
                ac = ws.cell(행, 2, _cell_value(산식))
                ac.alignment = Alignment(wrap_text=True)

            ws.column_dimensions["A"].width = 32
            ws.column_dimensions["B"].width = 20
            ws.column_dimensions["C"].width = 16
            ws.freeze_panes = "A2"

        # ── 4. 시트 순서 조정 (forms 모드: 입력·가정 시트를 틀+계산N 시트들 뒤로) ──
        if layout == "forms":
            sheets = wb._sheets
            inp_idx = None
            for i, s in enumerate(sheets):
                if s.title == "입력·가정":
                    inp_idx = i
                    break
            if inp_idx is not None:
                # 목록(1) + 틀 시트들 + 계산 N 시트들
                틀시트들_끝 = 1 + len(템플릿_목록) + calc_idx
                if inp_idx != 틀시트들_끝:
                    sheet = sheets.pop(inp_idx)
                    sheets.insert(틀시트들_끝, sheet)

    # ── 5. 입력·가정 시트 (steps 모드에서만 생성, forms 모드는 앞에서 생성) ──
    if layout != "forms":
        ws_inp = wb.create_sheet("입력·가정")
        시트이름들.append("입력·가정")
    행 = 1
    ws_inp.cell(행, 1, "call").font = bold
    ws_inp.cell(행, 2, "경로").font = bold
    ws_inp.cell(행, 3, "값").font = bold
    for cc in range(1, 4):
        ws_inp.cell(행, cc).fill = header_fill
    행 += 1
    for i, cd in enumerate(결과들):
        if cd["error"]:
            continue
        label = cd["label"]
        args = (calls[i] or {}).get("args", {})
        for 경로, 값 in _flatten_args(args):
            ws_inp.cell(행, 1, label)
            ws_inp.cell(행, 2, 경로)
            vc = ws_inp.cell(행, 3)
            if isinstance(값, (int, float)):
                _set_money(vc, 값)
            else:
                vc.value = _cell_value(값)
            행 += 1
    ws_inp.column_dimensions["A"].width = 20
    ws_inp.column_dimensions["B"].width = 36
    ws_inp.column_dimensions["C"].width = 18
    ws_inp.freeze_panes = "A2"

    # ── 6. 판단·근거 시트 (두 모드 공통) ──
    ws_jud = wb.create_sheet("판단·근거")
    시트이름들.append("판단·근거")
    행 = 1
    for cc, h in enumerate(["call", "항목", "결론", "금액", "근거", "이유"], 1):
        c = ws_jud.cell(행, cc, h); c.font = bold; c.fill = header_fill
    행 += 1
    for cd in 결과들:
        if cd["error"]:
            continue
        label = cd["label"]
        판단목록 = (cd.get("data") or {}).get("판단", [])
        if not 판단목록:
            continue
        if isinstance(판단목록, dict):      # religious_income 등 {항목: 결론} 형태
            판단목록 = [{"항목": k, "결론": v} for k, v in 판단목록.items()]
        elif isinstance(판단목록, str):
            판단목록 = [판단목록]
        for j in 판단목록:
            if not isinstance(j, dict):  # 판단이 문장 목록인 도구(예: nonbusiness_interest_disallowance)
                j = {"결론": j}
            ws_jud.cell(행, 1, label)
            ws_jud.cell(행, 2, j.get("항목", ""))
            con = ws_jud.cell(행, 3, _cell_value(j.get("결론", "")))
            amt = j.get("금액", 0)
            if isinstance(amt, (int, float)):
                _set_money(ws_jud.cell(행, 4), amt)
            else:
                ws_jud.cell(행, 4, amt)
            근거_val = j.get("근거", "")
            근거_cell = ws_jud.cell(행, 5, _cell_value(근거_val))
            if isinstance(근거_val, list) and _is_basis_list(근거_val):
                근거_cell.alignment = Alignment(wrap_text=True)
            이유_val = j.get("이유", "")
            이유_cell = ws_jud.cell(행, 6, _cell_value(이유_val))
            if isinstance(이유_val, list) and (_is_basis_list(이유_val) or _is_wrapped_basis_list(이유_val)):
                이유_cell.alignment = Alignment(wrap_text=True)
            행 += 1
    ws_jud.column_dimensions["A"].width = 20
    ws_jud.column_dimensions["B"].width = 16
    ws_jud.column_dimensions["C"].width = 40
    ws_jud.column_dimensions["D"].width = 16
    ws_jud.column_dimensions["E"].width = 30
    ws_jud.column_dimensions["F"].width = 36
    ws_jud.freeze_panes = "A2"
    # ── 저장 ──
    base = filename if filename else title
    fpath = _unique_path(output_dir, base)
    wb.save(fpath)
    return {
        "결과": {
            "파일": os.path.abspath(fpath),
            "시트": 시트이름들,
            "call_수": len(calls),
            "오류_call": 오류_call,
        },
        "근거": "계산 도구 결과 정리(세무 자문 아님)",
    }


# ── 국세 우선권에 따른 매각대금 배분(국세기본법 제35조·제36조) ──────────────────────
@mcp.tool(annotations=CALC)
@_guard
def tax_priority_distribution(
    sale_amount: Annotated[int, Field(description="매각대금(원) — 경매 매각대금 또는 공매 매각금액", ge=0)],
    sale_date: Annotated[str, Field(description="매각결정일·매각허가결정일 YYYY-MM-DD(2016-01-01~2026-12-31) — 2023.4.1.·2024.1.1. 개정 적용을 이 날짜로 판정")],
    claims: Annotated[list, Field(description=(
        "배분 받을 채권 목록. 각 항목 {종류, 금액, 이름?, ...}. 종류: 소액임차보증금(최우선변제 일정액) | 최우선임금채권(최종 3개월 임금·재해보상금·최종 3년 퇴직금) | "
        "당해세(세목=상속세·증여세·종합부동산세, 그 재산에 부과된 국세) | 국세 | 지방세(당해세=True면 세목=재산세·자동차세·지역자원시설세·지방교육세) | "
        "저당권 | 근저당권 | 전세권 | 질권 | 가등기담보 | 확정일자임차보증금 | 일반임금채권 | 일반채권. "
        "조세: 법정기일 YYYY-MM-DD(또는 확정방식=신고|고지|원천징수|인지세|제2차납세의무|양도담보|압류후확정 + 기준일), 압류일?(압류한 조세). "
        "담보: 설정일(확정일자 임차인은 대항요건·확정일자를 모두 갖춘 날), 순위?(같은 날 설정 시 접수 순서), 양도전_설정?(True=현 소유자 취득 전 설정, ①3호의2). "
        "일반채권: 가압류일?. 금액은 배분 요구액(원)."))],
    execution_cost: Annotated[int, Field(description="집행비용(원) — 강제집행·경매·공매 절차 비용(국세기본법 제35조①2호)", ge=0)] = 0,
    property_type: Annotated[Literal["주택", "상가건물", "기타"], Field(description="매각 재산 종류 — 소액임차 2분의 1 한도·주택 임차인 당해세 특례(제35조⑦) 판정용")] = "기타",
    prior_owner: Annotated[dict | None, Field(description=(
        "양도전_설정 담보가 있고 2023.4.1. 이후 매각일 때만: 직전 보유자 체납 정보 {체납: [{금액, 법정기일}], 보유전_전세권등: bool, 최초설정일?} "
        "또는 계산된 한도 {한도: 원}(시행령 제18조③)"))] = None) -> dict:
    """국세 우선권에 따른 매각대금 배분 순서·금액(경매·공매, 국세기본법 제35조·제36조). 언제: 체납자 재산이 경매·공매로 매각될 때 집행비용 → 최우선변제 → 당해세 → 법정기일과 담보 설정일 비교 → 임금 → 그 밖의 조세 → 일반채권 순으로 누가 얼마를 받는지 계산할 때.
    입력: 매각대금 → sale_amount, 매각(허가)결정일 → sale_date, 채권 목록 → claims, 집행비용 → execution_cost, 재산 종류 → property_type, 직전 보유자 체납(양도된 재산) → prior_owner.
    지원 연도: 2016.1.1.~2026.12.31. 매각결정분. 2023.4.1. 이후 매각결정분은 ①3호의2(양도 전 설정 담보 > 현 소유자 국세, 직전 보유자 체납액 한도 안 국세는 우선)·③ 후단(3호의2 재산은 종부세만 당해세 우선)·⑦(주택 임차인 당해세 순위 대신 변제), 2024.1.1. 이후는 한도 안 국세를 법정기일이 설정일보다 빠른 것으로 한정. 지방세 당해세·법정기일 비교(지방세기본법 제71조), 국세·지방세 상호간 압류선착(제36조·지방세기본법 제73조, 압류 없으면 동순위 안분).
    지원하지 않음(입력 확인 필요로 반환): 가압류채권과 후순위 담보권 안분·흡수배당, 조세 상호간 흡수배당(조세가 담보권 앞뒤로 나뉘고 뒤 조세가 덜 받는 경우), 담보권 사이 순환관계, 공과금, 납세담보(제37조), 3호·3호의2 담보 혼재·3호의2 담보 둘 이상, 일반임금채권과 전세권·확정일자 임차권의 순위, 제35조⑦ 대상 임차인 둘 이상. 소액임차 일정액·최우선임금 범위는 계산된 금액을 넣는다(2분의 1 한도만 적용).
    결과: 채권별 배분액·미배분액·단계별 근거, 배분 순서, 잔여(소유자 몫), 판단(적용한 개정·특례), 적용표."""
    try:
        return TP.distribute(sale_amount, execution_cost, sale_date, property_type, claims, prior_owner)
    except TP.입력확인 as e:
        return {"error": "입력 확인 필요 — 계산하지 않음", "문제": e.문제, **e.extra}


# ── 도구 찾기 ──────────────────────────────────────────────────────────────
# 도구가 70개를 넘어 목록만 보고는 맞는 도구를 놓치기 쉽다(10/28 점검: 있는 도구를 "없다"고 판단한 사례 3건).
# 사용자가 쓰는 말 → 도구. 별칭은 설명문에 없는 일상 표현·옛 용어·실무 용어를 잇는다.

# ── 주식등 양도소득: 대주주 판정·세율 (소득세법 제94조①3호·제104조①11호·12호, 시행령 제157조·제167조의8) ─────────
@mcp.tool(annotations=CALC)
@_guard
def stock_major_shareholder(transfer_date: Annotated[str, Field(description="양도일 YYYY-MM-DD(2016-01-01 ~ 2026-12-31) — 요건은 양도일 구간별로 다름")],
                            market: Annotated[str, Field(description="시장 — 코스피|코스닥|코넥스|비상장|K-OTC. 평소 말(유가증권·KOSDAQ·장외 등)도 됨")],
                            own_ratio: Annotated[float | None, Field(description="직전 사업연도 종료일 현재 본인 지분율 — 0~1 비율(0.015) 또는 %(1.5). 모르면 own_mcap만", ge=0)] = None,
                            own_mcap: Annotated[int | None, Field(description="직전 사업연도 종료일 현재 본인 보유 주식 시가총액(원) — 상장은 그날 최종시세가액 × 주식 수, 비상장은 상증법 평가액", ge=0)] = None,
                            related_ratio: Annotated[float | None, Field(description="특수관계인(기타주주) 지분율 합계 — 본인 제외. 구간에 따라 합산 여부·범위가 다르므로 아는 대로 넣으면 도구가 판단", ge=0)] = None,
                            related_mcap: Annotated[int | None, Field(description="특수관계인 보유 시가총액 합계(원), 본인 제외", ge=0)] = None,
                            is_max_group: Annotated[bool | None, Field(description="주주 1인과 특수관계인의 지분 합계가 그 법인 주주 중 최대인지(최대주주군). 2023.1.1.~ 양도분은 이것이 True일 때만 합산. 모르면 비움")] = None,
                            venture: Annotated[bool, Field(description="비상장·K-OTC만: 벤처기업 주식이면 True(시가총액 기준 40억)")] = False,
                            is_corporation: Annotated[bool, Field(description="주주가 법인이면 True(2025.2.28.~ 대주주 판정 대상 아님)")] = False,
                            post_acquisition_ratio: Annotated[float | None, Field(description="직전 사업연도 종료일 뒤 양도일 전에 주식을 더 취득했으면 취득 후 지분율(0~1 또는 %) — 지분율 기준에 이르면 취득일 이후 대주주(영 제157조⑨, 시가총액 기준은 적용 안 됨)", ge=0)] = None,
                            total_shares: Annotated[int | None, Field(description="주주 명부로 계산할 때: 직전 사업연도 종료일 현재 해당 법인 주식등 합계(보통주·우선주·자기주식 포함, 신주인수권증권·출자지분 포함). 이걸 주면 own_ratio 대신 명부에서 지분율을 계산", ge=1)] = None,
                            own_shares: Annotated[int | None, Field(description="명부: 본인 보유 주식 수(직전 사업연도 종료일 현재)", ge=0)] = None,
                            holders: Annotated[list | None, Field(description="명부: 다른 주주 [{이름, 관계, 주식수}] — 관계는 평소 말(배우자·아들·형·사촌·처남·장인·사위 …, '혈족 4촌'·'인척 3촌', 경영지배 관계 법인은 '경영지배'). 합산 여부·범위는 양도일 구간·최대주주군 여부로 도구가 판단해 주주별로 보여줌")] = None,
                            price: Annotated[float | None, Field(description="명부: 1주당 가액(원) — 상장은 직전 사업연도 종료일 최종시세가액, 비상장은 소득세법 시행령 제165조④ 기준시가(stock_basis_income_tax). 없으면 시가총액 기준은 판정하지 않음", ge=0)] = None,
                            lent_shares: Annotated[int, Field(description="명부: 본인이 대여 중인 주식 수(반환받는 조건) — 대여자 소유로 봄(영 제157조⑦)", ge=0)] = 0,
                            pef_shares: Annotated[int, Field(description="명부: 사모집합투자기구를 통해 보유한 주식 중 본인 투자비율로 안분한 주식 수(영 제157조⑧)", ge=0)] = 0) -> dict:
    """주식 양도소득세 대주주 판정 — 양도일·시장·직전 사업연도 종료일 현재 지분율·시가총액으로 대주주 여부(소득세법 시행령 제157조·제167조의8, 2016~2026 구간별 기준: 코스피 1%/25억→15억→10억→50억 등)와 특수관계인 합산 방식을 판정한다. 주주 명부(total_shares·own_shares·holders·price)를 주면 지분율·시가총액과 주주별 합산 여부("처남은 2023년부터 인척 2촌이라 최대주주군일 때만 합산" 등)까지 도구가 계산한다.
    언제: "나 대주주야?", 주식 양도세 세율·과세 여부를 정하기 전에. 결과의 판정(대주주/아님/확인 필요)을 stock_capital_gains_tax의 major_shareholder에 넣는다.
    입력: 양도일 → transfer_date, 시장 → market, 본인 지분율·시가총액 → own_ratio·own_mcap, 특수관계인 합계 → related_ratio·related_mcap, 최대주주군 여부 → is_max_group, 벤처 → venture.
    결과: 판정·이유·적용 기준(지분율·시가총액)·합산 방식, 판단(기준일·시가 산정·확인 필요), 근거(시행령 조항·개정령·시행일). 시가총액은 사용자가 계산해 넣는다(도구가 시세를 조회하지 않음)."""
    return ST.major_shareholder(transfer_date, market, own_ratio=own_ratio, own_mcap=own_mcap, related_ratio=related_ratio, related_mcap=related_mcap,
                                is_max_group=is_max_group, venture=venture, is_corporation=is_corporation, post_acquisition_ratio=post_acquisition_ratio,
                                total_shares=total_shares, own_shares=own_shares, holders=holders, price=price, lent_shares=lent_shares, pef_shares=pef_shares)


@mcp.tool(annotations=CALC)
@_guard
def stock_capital_gains_tax(transfer_date: Annotated[str, Field(description="양도일 YYYY-MM-DD(2016-01-01 ~ 2026-12-31)")],
                            market: Annotated[str, Field(description="시장 — 코스피|코스닥|코넥스|비상장|K-OTC|국외. 평소 말도 됨")],
                            major_shareholder: Annotated[bool | None, Field(description="대주주 여부(stock_major_shareholder 판정값). 국외주식은 불필요")] = None,
                            sme: Annotated[bool | None, Field(description="발행법인이 중소기업인지(중소기업기본법 제2조, 판정 기준일은 결과 판단에 표시)")] = None,
                            mid_sized: Annotated[bool, Field(description="K-OTC만: 중견기업이면 True(소액주주 과세 제외 판정용)")] = False,
                            otc_trade: Annotated[bool, Field(description="상장주식을 증권시장 밖(장외)에서 양도했으면 True — 대주주가 아니어도 과세")] = False,
                            transfer_price: Annotated[int | None, Field(description="양도가액(원)", ge=0)] = None,
                            acquisition_price: Annotated[int | None, Field(description="취득가액(원, 취득 수수료 포함)", ge=0)] = None,
                            expenses: Annotated[int, Field(description="양도 시 필요경비(증권거래세·수수료 등, 원)", ge=0)] = 0,
                            gain: Annotated[int | None, Field(description="양도차익(원)을 이미 알면 가액 대신 이것만")] = None,
                            acquisition_date: Annotated[str, Field(description="취득일 YYYY-MM-DD — 대주주이고 중소기업이 아니면 필요(1년 미만 보유 30%)")] = "",
                            prior_basic_deduction: Annotated[int, Field(description="같은 해 주식등 양도에서 이미 쓴 기본공제(원, 연 250만원 한도)", ge=0)] = 0) -> dict:
    """주식등 양도소득세 — 과세 대상 여부(상장주식은 대주주 양도·장외 양도만, K-OTC 중소·중견 소액주주 제외), 양도차익 → 기본공제(연 250만원) → 과세표준 → 세율(대주주 20%·3억 초과 25%, 1년 미만 비중소 30%, 소액주주 중소 10%·그 밖 20%, 국외 10%/20%; 연도별) → 산출세액·지방소득세·예정신고기한(반기 말일+2개월).
    언제: 주식을 팔았을 때 세금이 얼마인지, 과세 대상인지 물을 때. 대주주 여부는 먼저 stock_major_shareholder로 판정한다.
    입력: 양도일 → transfer_date, 시장 → market, 대주주 → major_shareholder, 중소기업 → sme, 양도가액·취득가액·필요경비 → transfer_price·acquisition_price·expenses(또는 gain), 취득일 → acquisition_date.
    지원하지 않음: 2019 이전 국외주식(제118조의2), 파생상품, 상장주식 장내 소액주주(과세 대상 아님으로 답함), 양도차손 통산 계산.
    결과: 과세대상·양도차익·기본공제·과세표준·적용세율·산출세액·지방소득세·예정신고기한, 세율 후보(경합 시 큰 세액), 판단, 근거."""
    return ST.stock_gains_tax(transfer_date, market, major_shareholder=major_shareholder, sme=sme, mid_sized=mid_sized, otc_trade=otc_trade,
                              transfer_price=transfer_price, acquisition_price=acquisition_price, expenses=expenses, gain=gain,
                              acquisition_date=acquisition_date, prior_basic_deduction=prior_basic_deduction)



@mcp.tool(annotations=ToolAnnotations(read_only_hint=True, destructive_hint=False, idempotent_hint=True, open_world_hint=True))
@_guard
def stock_price_lookup(ticker: Annotated[str, Field(description="6자리 종목코드(예: 005930). 종목명은 받지 않음")],
                       day: Annotated[str, Field(description="기준일 YYYY-MM-DD — close면 그날(휴장이면 직전 거래일) 종가, avg_2m이면 평가기준일")],
                       mode: Annotated[Literal["close", "avg_2m"], Field(description="close: 최종시세가액(양도세 대주주 시가총액용, 직전 사업연도 종료일) / avg_2m: 평가기준일 전후 각 2개월 종가 평균(상속·증여 시가)")] = "close") -> dict:
    """상장주식 시세 조회 — KRX 정보데이터시스템(pykrx)에서 종가를 가져와 (1) 직전 사업연도 종료일 최종시세가액(소득세법 시행령 제157조④1호, 휴장이면 직전거래일) 또는 (2) 상속·증여 평가 시가(평가기준일 전후 2개월 종가 평균, 상증법 제63조①1호가목)를 돌려준다.
    언제: 대주주 판정(stock_major_shareholder의 price·own_mcap)이나 상속·증여 주식 평가에 넣을 가액이 필요할 때.
    입력: 종목코드 → ticker, 기준일 → day, 용도 → mode(close|avg_2m).
    주의: 이 도구만 외부(KRX)로 조회하며 보내는 정보는 종목코드·날짜뿐. pykrx가 없으면 설치 안내를 돌려줌(pip install 'korean-tax-calc-mcp[stock]'). pykrx의 'KRX 로그인 실패' 경고는 종가 조회와 무관(로그인 불필요). 비상장주식은 상증법 보충적 평가 도구로.
    결과: 종가 또는 2개월 평균·거래일수·기간, 판단(직전거래일 사용 여부·기간 조정 확인 필요), 근거."""
    return SP.stock_price(ticker, day, mode=mode)



@mcp.tool(annotations=CALC)
@_guard
def unlisted_stock_valuation(valuation_date: Annotated[str, Field(description="평가기준일(상속개시일·증여일) YYYY-MM-DD, 2016~2026")],
                             shares: Annotated[int, Field(description="평가기준일 현재 발행주식총수", ge=1)],
                             incomes: Annotated[list | None, Field(description="최근 3개 사업연도 [{사업연도, 순손익액 | 각사업연도소득 + 가산:{…} + 차감:{…} | calls:[{tool, args}](세무조정 계산기 호출을 주면 서버가 실행해 각사업연도소득·가산·차감을 채움), 주식수?}] — 1년 전부터. 순손익액을 바로 알면 숫자만, 세무조정계산서에서 뽑으면 각사업연도소득과 가감 항목으로(영 제56조④, 도구가 산출 과정을 보여줌). 사업 개시 3년 미만이면 생략 가능")] = None,
                             net_asset: Annotated[int | dict | None, Field(description="평가기준일 현재 순자산가액(원), 또는 재무상태표 항목형 {자산_장부가액, 자산_평가차액?(상증법 평가액−장부가액), 부채_장부가액, 자산_가산:{지급받을 권리…}, 자산_차감:{선급비용, 영업권…}, 부채_가산:{미지급 법인세, 확정 배당, 퇴직급여추계액…}, 부채_차감:{충당금, 준비금…}} — 영 제55조·규칙 제17조의2 가감을 도구가 적용하고 산출 과정을 보여줌")] = None,
                             equity: Annotated[int | None, Field(description="영업권 계산용 평가기준일 자기자본(원). goodwill을 직접 주면 불필요")] = None,
                             goodwill: Annotated[int | None, Field(description="영업권 평가액(원)을 이미 알면 직접")] = None,
                             real_estate_ratio: Annotated[float | None, Field(description="자산총액 중 부동산·부동산권리 비율(0~1) — 0.5 이상이면 2:3 가중, 0.8 이상이면 순자산가치만(2018.2.13.~)")] = None,
                             stock_asset_ratio: Annotated[float | None, Field(description="자산총액 중 주식등 비율(0~1) — 0.8 이상이면 순자산가치만(2017.2.7.~)")] = None,
                             business_start: Annotated[str, Field(description="사업개시일 YYYY-MM-DD(개시 후 3년 미만 판정용)")] = "",
                             status: Annotated[str, Field(description="청산|휴폐업|존속3년(정관 존속기한 3년 이내) 해당 시")] = "",
                             is_max_shareholder: Annotated[bool, Field(description="평가 대상 주식이 최대주주등 보유분이면 True(할증)")] = False,
                             ratio_over_50: Annotated[bool | None, Field(description="2019.12.31.까지: 최대주주등 지분이 50% 초과면 True(30%), 아니면 False(20%)")] = None,
                             sme: Annotated[bool | None, Field(description="중소기업 여부(할증 제외 판정) — 최대주주면 필수")] = None,
                             mid_sized: Annotated[bool, Field(description="중견기업(직전 3개 사업연도 매출액 평균 5천억 미만)이면 True — 2023.1.1.~ 할증 제외")] = False,
                             premium_exclusion: Annotated[str | None, Field(description="할증 제외 사유가 있으면 적음(3년 연속 결손·전부 매각·사업개시 3년 내 영업이익 0 이하·청산 확정·명의신탁 등, 영 제53조⑧)")] = None,
                             capital_changes: Annotated[list | None, Field(description="평가기준일 전 3년 내 유상증자·감자 [{일자, 구분:'증자'|'감자', 주식수, 1주당금액}] — 각 연도 순손익액 가산·차감(영 제56조⑤)과 환산주식수(규칙 제17조의3⑤)를 도구가 계산")] = None,
                             estimated_profit: Annotated[dict | None, Field(description="추정이익으로 평가하려면 {1주당추정이익평균(둘 이상 평가기관 산출 평균), 사유(자산수증이익|합병|분할|업종 변경|휴업|처분손익|매출 3년 미만|고시), 신고기한, 신고일, 산정기준일, 평가서작성일} — 영 제56조② 요건 4개를 판정해 충족 시 3년 가중평균 대신 적용")] = None,
                             assets: Annotated[list | None, Field(description="자산별 평가로 평가차액을 만들려면 [{종류(토지|건물|주택|상장주식|비상장주식|예금|채권|매출채권|기타), 명칭?, 장부가액, 시가?|기준시가?|평가액?, 하한 예외?}] — 평가액>시가>기준시가>장부가액 순으로 고르고, 장부가액보다 낮으면 장부가액(영 제55조①), 합계를 자산_평가차액에 더함")] = None,
                             kind: Annotated[Literal["상속", "증여"], Field(description="상속·증여 구분(추정이익 요건 4호 연도 판정용)")] = "증여") -> dict:
    """비상장주식 상속·증여 평가(상증법 보충적 평가) — 순손익가치(3년 가중평균 3:2:1 ÷ 10%)와 순자산가치((순자산+영업권) ÷ 주식수)를 3:2(부동산 과다 2:3)로 가중평균, 하한 순자산가치×80%(2017.4~2018.3 70%), 순자산가치만 평가하는 사유(청산·개시 3년 미만·휴폐업·부동산 80%·주식 80%·존속 3년, 구간별), 최대주주 할증(2019까지 20/30%·중소 0, 2020~ 20%·중소 제외, 2023~ 중견 제외)까지 2016~2026 구간별로 계산한다.
    언제: 비상장주식을 상속·증여받았을 때 1주당 평가액, 또는 비상장법인 지분의 상증법 가액이 필요할 때. 양도세 대주주 시가총액·기준시가는 산식이 달라 stock_basis_income_tax를 쓴다.
    입력: 평가기준일 → valuation_date, 발행주식총수 → shares, 3년 순손익 → incomes, 순자산가액 → net_asset, 부동산 비율 → real_estate_ratio, 최대주주·중소기업 → is_max_shareholder·sme.
    재무제표 입력: incomes에 각사업연도소득+가산·차감 항목, net_asset에 자산·부채 장부가액+평가차액+가감 항목을 주면 순손익액·순자산가액을 도구가 만들고 줄별 근거를 보여준다.
    세무조정 연결: incomes 항목에 calls=[{tool, args}]를 주면 서버가 세무조정 계산기(income_adjustment_total·corporate_tax·entertainment_adjustment·donation_limit·dividend_received_deduction·business_car_expense·nonbusiness_interest_disallowance·depreciation_adjustment)를 실행해 각사업연도소득·가산·차감에 자동으로 넣는다(영 제56조④ 항목만).
    증감자(capital_changes)·추정이익(estimated_profit, 요건 판정)·자산별 평가(assets: 시가→기준시가→장부가액 선택, 장부가액 하한)까지 받는다.
    지원하지 않음: 평가심의위원회, 자산별 시가의 평가기간(상속 전후 6개월·증여 전 6개월~후 3개월) 판정(사용자가 기간 안 가액만 넣음), 국외 비상장.
    결과: 1주당 평가액(할증 후·전), 순손익가치·순자산가치·가중평균, 평가방법, 할증률, 판단(사유·하한·확인 필요), 근거."""
    연결 = []
    if incomes:
        new_items = []
        for it in incomes:
            if isinstance(it, dict) and it.get("calls"):
                it2, lines, probs = _link_income_calls(it)
                if probs:
                    return {"error": "입력 확인 필요 — 계산하지 않음", "문제": probs}
                연결.append({"사업연도": it2.get("사업연도") or it2.get("year"), "연결": lines}); it = it2
            new_items.append(it)
        incomes = new_items
    out = SV.unlisted_valuation(valuation_date, shares, incomes=incomes, net_asset=net_asset, equity=equity, goodwill=goodwill, real_estate_ratio=real_estate_ratio,
                                 stock_asset_ratio=stock_asset_ratio, business_start=business_start, status=status, is_max_shareholder=is_max_shareholder,
                                 ratio_over_50=ratio_over_50, sme=sme, mid_sized=mid_sized, premium_exclusion=premium_exclusion,
                                 capital_changes=capital_changes, estimated_profit=estimated_profit, assets=assets, kind=kind)
    if 연결 and isinstance(out, dict) and "계산" in out:
        out["계산"]["세무조정 연결"] = 연결
    return out


@mcp.tool(annotations=CALC)
@_guard
def stock_basis_income_tax(transfer_date: Annotated[str, Field(description="양도일(또는 대주주 판정 시 양도일) YYYY-MM-DD")],
                           shares: Annotated[int, Field(description="직전 사업연도 종료일 현재 발행주식총수", ge=1)],
                           prior_year_income: Annotated[int | None, Field(description="양도일이 속하는 사업연도의 직전 사업연도 순손익액(원)")] = None,
                           book_net_asset: Annotated[int | None, Field(description="직전 사업연도 종료일 현재 장부가액 순자산(원, 토지는 기준시가로 바꿔 계산한 값)")] = None,
                           real_estate_ratio: Annotated[float | None, Field(description="자산총액 중 부동산 비율(0~1), 0.5 이상이면 2:3")] = None,
                           stock_asset_ratio: Annotated[float | None, Field(description="자산총액 중 주식등 비율(0~1) — 2023.2.28.~ 0.8 이상이면 순자산가치만")] = None,
                           business_start: Annotated[str, Field(description="사업개시일 YYYY-MM-DD(개시 후 1년 미만이면 순자산가치만)")] = "",
                           status: Annotated[str, Field(description="청산|휴폐업|존속3년 해당 시")] = "") -> dict:
    """비상장주식 소득세법 기준시가(영 제165조④) — 양도세 기준시가와 대주주 시가총액 판정(영 제157조④2호)에 쓰는 1주당 가액. 상속·증여 평가와 달리 직전 사업연도 순손익만(÷10%), 순자산은 장부가액, 할증 없음, 하한 순자산가치×80%(2018.4.1.~).
    언제: 비상장주식 양도세에서 기준시가가 필요하거나, stock_major_shareholder의 price(1주당 가액)에 넣을 값이 필요할 때.
    입력: 양도일 → transfer_date, 발행주식총수 → shares, 직전 사업연도 순손익액 → prior_year_income, 장부 순자산 → book_net_asset.
    결과: 1주당 기준시가, 순손익가치·순자산가치, 평가방법, 판단(할증 없음·이자율 확인 필요), 근거."""
    return SV.income_tax_unlisted_basis(transfer_date, shares, prior_year_income=prior_year_income, book_net_asset=book_net_asset, real_estate_ratio=real_estate_ratio,
                                        business_start=business_start, status=status, stock_asset_ratio=stock_asset_ratio)



@mcp.tool(annotations=ToolAnnotations(read_only_hint=True, destructive_hint=False, idempotent_hint=True, open_world_hint=True))
@_guard
def official_property_price(address: Annotated[str, Field(description="지번 주소 또는 도로명 주소(예: '서울특별시 강남구 역삼동 737', '서울 강남구 테헤란로 7길 22'). pnu를 알면 생략")] = "",
                            pnu: Annotated[str, Field(description="필지번호 19자리(법정동코드 10 + 산 여부 1 + 본번 4 + 부번 4)를 알면 직접")] = "",
                            kind: Annotated[str, Field(description="토지|아파트(공동주택)|단독주택 — 평소 말(땅·대지·연립·다세대·다가구)도 됨")] = "토지",
                            year: Annotated[int | None, Field(description="공시 기준연도(1월 1일 기준). 비우면 조회되는 최신 연도. 평가기준일이 그 해 공시 전이면 전년도를 넣음", ge=1990, le=2030)] = None,
                            area_m2: Annotated[float | None, Field(description="토지만: 면적(㎡) — 주면 공시지가 × 면적 총액을 계산", ge=0)] = None,
                            dong: Annotated[str, Field(description="아파트만: 동(예: '101')")] = "",
                            ho: Annotated[str, Field(description="아파트만: 호(예: '1203')")] = "") -> dict:
    """부동산 공시가격 조회 — 주소(또는 PNU)로 개별공시지가(토지, 원/㎡)·공동주택가격(아파트, 세대별)·개별주택가격(단독)을 국토교통부 국가공간정보(브이월드)에서 가져온다. 상속·증여 보충적 평가(상증법 제61조①)와 양도세 기준시가(소득세법 제99조①1호)에 쓰는 값.
    언제: 비상장주식 평가(unlisted_stock_valuation의 assets 기준시가), 상속·증여재산 평가(inheritance_gift_valuation의 supplementary_value), 양도세 기준시가가 필요할 때. 시가(매매·감정 사례)가 있으면 시가가 우선(법 제60조) — 이 도구는 공시가격만.
    입력: 주소 → address(또는 pnu), 종류 → kind, 기준연도 → year, 토지 면적 → area_m2, 아파트 동·호 → dong·ho.
    주의: 이 도구는 외부(브이월드)로 주소·PNU·연도를 보냄. 인증키 VWORLD_KEY가 필요(vworld.kr 오픈API 인증키, 국가중점API 포함) — 없으면 안내와 직접 입력 대안을 돌려줌.
    결과: 기준연도·공시가격(토지는 ㎡당과 총액), 연도별 목록, 판단(필지 확인·평가 사용법·기준일), 근거."""
    return OP.official_price(address=address, pnu=pnu, kind=kind, year=year, area_m2=area_m2, dong=dong, ho=ho)



@mcp.tool(annotations=CALC)
@_guard
def building_standard_price(year: Annotated[int, Field(description="적용연도(양도·상속·증여일이 속한 해, 2007~2026). 새 고시 전이면 직전 연도", ge=2007, le=2026)],
                            area_m2: Annotated[float, Field(description="평가대상 건물 연면적(㎡, 집합건물은 전용+공용)", gt=0)],
                            structure: Annotated[str, Field(description="구조 — 철근콘크리트·철골·경량철골·연와(벽돌)·블록·목조·조립식패널·컨테이너 등 평소 말(건축물대장 표기)")],
                            use: Annotated[str, Field(description="용도 — 사무실·근린생활시설(상가)·공장·창고·병원·호텔·학원 등 평소 말(건축물대장 주용도). 주택·아파트·오피스텔은 이 고시 대상 아님(공시가격·고시가격)")],
                            land_price_per_m2: Annotated[float | None, Field(description="부속토지 ㎡당 개별공시지가(원) — official_property_price(토지)로 조회", ge=0)] = None,
                            built_year: Annotated[int | None, Field(description="신축연도(사용승인일 기준; 가사용·사실상 사용이 앞서면 그 해; 증축은 증축연도)", ge=1900, le=2026)] = None,
                            purpose: Annotated[Literal["양도세", "상증세"], Field(description="양도세 기준시가(조정률 미적용) / 상증세 보충적 평가(개별건물 특성 조정률 적용)")] = "양도세",
                            floors: Annotated[int | None, Field(description="상증세만: 최고층수(지하·옥탑 제외) — 규모 조정률", ge=1)] = None,
                            total_area_m2: Annotated[float | None, Field(description="상증세만: 건물 전체 연면적(㎡) — 규모 조정률(비우면 area_m2)", gt=0)] = None,
                            roof: Annotated[str, Field(description="상증세만, 구조지수 100 미만 건물: 지붕재료(슬래브·기와|패널·유리·슬레이트|함석·천막·초가)")] = "",
                            floor_use: Annotated[str, Field(description="상증세만: 평가 부분이 상가 1층·2층·지하1층·지하2층·부속주차장·기계실·옥탑이면 적음")] = "",
                            safety_grade: Annotated[str, Field(description="상증세만: 안전진단 등급 B|C|D|E(납세자 입증)")] = "",
                            remodel_count: Annotated[int, Field(description="상증세만: 일부개축 횟수(개축 부분만 해당)", ge=0)] = 0) -> dict:
    """건물 기준시가 — 국세청 「건물 기준시가 계산방법 고시」 산식(신축가격기준액 × 구조지수 × 용도지수 × 위치지수 × 경과연수별 잔가율 × (상증세만) 개별건물 특성 조정률, ㎡당 1,000원 미만 버림 × 면적). 양도세 기준시가(소득세법 제99조①1호나목)와 상속·증여 보충적 평가(상증법 제61조①2호)에 쓴다.
    언제: 일반 건물(상가·사무실·공장·창고 등)의 기준시가가 필요할 때 — 비상장주식 평가(unlisted_stock_valuation의 assets 기준시가), 상속·증여 신고, 양도세 환산취득가액. 주택·아파트는 공시가격(official_property_price), 오피스텔·일괄고시 상업용은 국세청 고시가격을 쓴다.
    입력: 적용연도 → year, 면적 → area_m2, 구조·용도 → structure·use, 부속토지 공시지가 → land_price_per_m2, 신축연도 → built_year, 상증세면 purpose='상증세'와 층수·연면적·지붕·층별·안전등급.
    지수표는 2026년 고시(제2025-39호) 원문. 2007~2025년은 그 해 신축가격기준액에 2026 지수표를 적용한 근사(결과에 표시).
    결과: 기준시가·㎡당 금액·지수·잔가율·조정률, 판단(산식·시가 우선·국세청 감정평가 가능성·범위), 근거."""
    return BL.building_standard_price(year, area_m2, structure, use, land_price_per_m2, built_year, purpose=purpose, floors=floors, total_area_m2=total_area_m2,
                                      roof=roof, floor_use=floor_use, safety_grade=safety_grade, remodel_count=remodel_count)



# ── 세무조정 계산기 결과 → 비상장주식 평가 순손익액 입력 연결(영 제56조④) ────────────────────────
_VAL_LINK = {   # 도구 → [(구분, 항목, 결과 키 또는 키 목록)]; 구분 '각사업연도소득'은 기준값
    "income_adjustment_total": [("각사업연도소득", "", "각사업연도소득금액")],
    "corporate_tax": [("차감", "법인세", "산출세액")],
    "entertainment_adjustment": [("차감", "접대비 한도초과", "손금불산입_합계")],
    "donation_limit": [("차감", "기부금 한도초과", "한도초과(기타사외유출)"), ("가산", "기부금이월", "이월분 손금산입")],
    "dividend_received_deduction": [("가산", "수입배당금", "익금불산입액")],
    "business_car_expense": [("차감", "업무용승용차 한도초과", ["유지비 손금불산입(상여)", "감가상각 사적사용 손금불산입(상여)", "감가상각 한도초과(유보)"]), ("가산", "27조의2 손금산입", "감가상각 미계상 손금산입(유보)")],
    "nonbusiness_interest_disallowance": [("차감", "지급이자 손금불산입", "손금불산입(기타사외유출)")],
    "depreciation_adjustment": [("차감", "감가상각 시인부족", "전기부인액_추인")],
}


def _link_income_calls(item: dict) -> tuple[dict, list[dict], list[str]]:
    """incomes 항목의 calls=[{tool, args}]를 서버가 실행해 각사업연도소득·가산·차감에 넣는다. (항목, 연결 줄, 문제)"""
    import json as _json
    it = dict(item); 가산 = dict(it.get("가산") or {}); 차감 = dict(it.get("차감") or {}); lines = []; problems = []
    for call in it.get("calls") or []:
        tool = call.get("tool", ""); args = call.get("args", {})
        if isinstance(args, str):
            try: args = _json.loads(args)
            except Exception: problems.append(f"{tool}: args가 JSON이 아님"); continue
        if tool not in _VAL_LINK:
            problems.append(f"{tool}: 순손익액 연결 미지원 도구 — 결과 숫자를 가산·차감에 직접 넣으세요(연결 지원: {', '.join(_VAL_LINK)})"); continue
        try:
            r = asyncio.run(mcp.call_tool(tool, args)); data = _json.loads(r.content[0].text if hasattr(r, "content") else r[0][0].text)
        except Exception as e:
            problems.append(f"{tool}: 호출 오류 {e}"); continue
        if "error" in data:
            problems.append(f"{tool}: {data['error']} {('— ' + ' / '.join(data.get('문제', []))) if data.get('문제') else ''}"); continue
        res = data.get("결과", {})
        for 구분, 항목, keys in _VAL_LINK[tool]:
            ks = keys if isinstance(keys, list) else [keys]
            vals = [res.get(k) for k in ks if isinstance(res.get(k), (int, float))]
            if not vals:
                problems.append(f"{tool}: 결과에 {ks} 없음"); continue
            amt = float(sum(vals))
            if 구분 == "각사업연도소득":
                it["각사업연도소득"] = amt; lines.append({"도구": tool, "구분": "각사업연도소득", "금액": amt, "근거": "법인세법 제14조"})
            elif 구분 == "가산":
                가산[항목] = 가산.get(항목, 0) + amt; lines.append({"도구": tool, "구분": "+ " + 항목, "금액": amt, "근거": "영 제56조④1호"})
            else:
                차감[항목] = 차감.get(항목, 0) + amt; lines.append({"도구": tool, "구분": "− " + 항목, "금액": amt, "근거": "영 제56조④2호"})
    if it.get("각사업연도소득") is None and it.get("순손익액") is None and it.get("calls"):
        problems.append("각사업연도소득이 없음 — income_adjustment_total call을 넣거나 각사업연도소득을 직접 주세요")
    it["가산"] = 가산; it["차감"] = 차감; it.pop("calls", None)
    return it, lines, problems


@mcp.tool(annotations=CALC)
@_guard
def individual_donation_expense(year: YEAR,
                                base_income: Annotated[int, Field(description="기준소득금액 — 기부금을 필요경비에 산입하기 전의 해당 과세기간 사업소득금액(원). 이자·배당 등 다른 종합소득은 넣지 않음", ge=0)],
                                carried_loss: Annotated[int, Field(description="소득세법 제45조 이월결손금(원)", ge=0)] = 0,
                                political: Annotated[int, Field(description="정치자금 기부금 총액(원) — 10만원까지 세액공제, 초과분 필요경비", ge=0)] = 0,
                                hometown: Annotated[int, Field(description="고향사랑 기부금 총액(원) — 10만원까지 세액공제, 초과분 필요경비", ge=0)] = 0,
                                special: Annotated[int, Field(description="특례기부금 금전 지출(원): 국가·지자체, 국방헌금, 이재민 구호금품, 사립학교·병원 시설비·교육비·장학금·연구비 등(법인세법 제24조②1호)", ge=0)] = 0,
                                general: Annotated[int, Field(description="일반기부금 금전 지출(원, 종교단체 제외): 사회복지법인·의료법인·공익법인 고유목적사업비 등", ge=0)] = 0,
                                religious: Annotated[int, Field(description="종교단체 일반기부금(원)", ge=0)] = 0,
                                carried_special: Annotated[int, Field(description="이월된 특례기부금 한도초과액(원, 10년)", ge=0)] = 0,
                                carried_general: Annotated[int, Field(description="이월된 일반기부금 한도초과액(원, 10년)", ge=0)] = 0,
                                volunteer: Annotated[dict | None, Field(description="특별재난지역 자원봉사용역 {hours: 본인 봉사시간, fuel: 부수 유류비, transport_fuel: 그중 이동용 유류비(제외), materials: 재료비} — 특례기부금")] = None,
                                in_kind: Annotated[list | None, Field(description="현물 기부 [{kind: 특례|일반|종교, book: 장부가액, market: 시가}] — 시가(시가가 장부가액보다 낮으면 장부가액)로 평가")] = None) -> dict:
    """개인사업자 기부금 필요경비 산입 한도(소득세법 제34조). 언제: 사업소득자의 종합소득세 신고에서 정치자금·고향사랑·특례·일반(종교) 기부금을 어디까지 필요경비로 넣고, 얼마가 세액공제·이월로 넘어가는지 계산할 때.
    순서(시행령 제81조④): 정치자금·고향사랑 10만원 초과분(이월결손금 뺀 소득금액 한도) → 특례(기준소득−이월결손금) → 일반(종교 있으면 10% + min(20%, 종교 외), 없으면 30%). 이월분은 같은 구분에서 먼저 산입(제79조④).
    현물은 시가(시가<장부가액이면 장부가액)(제81조③), 자원봉사용역은 봉사일수(8시간 단위 올림)×8만원 + 부수 유류비(이동분 제외)·재료비(제81조⑤).
    입력: 기준소득금액 → base_income(이자·배당 제외), 이월결손금 → carried_loss, 기부금 → political·hometown·special·general·religious, 이월 → carried_*, 자원봉사 → volunteer, 현물 → in_kind.
    결과: 필요경비 산입 합계·구분별 산입·한도초과 이월·기부금 산입 후 사업소득금액·정치·고향 세액공제 대상(10만원 이하), 판단, 근거. 세액공제(제59조의4) 금액 자체는 계산하지 않음 — 종합소득 기준은 별도."""
    return DI.individual_donation(year, base_income, carried_loss, political, hometown, special, general, religious, carried_special, carried_general, volunteer, in_kind)


_TOOL_ALIASES: dict[str, list[str]] = {
    "양도세": ["capital_gains_income", "capital_gains_tax"], "양도소득세": ["capital_gains_income", "capital_gains_tax"],
    "단기 양도": ["capital_gains_tax"], "미등기": ["capital_gains_tax"], "비사업용 토지": ["capital_gains_tax", "capital_gains_income"],
    "집 팔": ["capital_gains_income"], "아파트 매도": ["capital_gains_income"], "1세대1주택": ["capital_gains_income"],
    "다주택 중과": ["capital_gains_income", "capital_gains_tax"], "다주택자": ["capital_gains_income", "capital_gains_tax"],
    "조정대상지역": ["capital_gains_income", "capital_gains_tax"], "중과세율": ["capital_gains_tax", "capital_gains_income"],
    "고가주택": ["capital_gains_income"], "장기보유특별공제": ["capital_gains_income"], "부담부증여": ["capital_gains_income", "gift_tax"],
    "청산": ["liquidation_income"], "해산": ["liquidation_income"], "잔여재산": ["liquidation_income", "deemed_dividend_capital_reduction"],
    "보험차익": ["asset_grant_reserve"], "국고보조금": ["asset_grant_reserve"], "공사부담금": ["asset_grant_reserve"],
    "일시상각충당금": ["asset_grant_reserve"], "압축기장충당금": ["asset_grant_reserve"],
    "차가감소득금액": ["income_adjustment_total"], "각사업연도소득": ["income_adjustment_total", "corporate_tax"],
    "세무조정 합계": ["income_adjustment_total"], "소득금액조정합계표": ["income_adjustment_total"],
    "접대비": ["entertainment_adjustment"], "기업업무추진비": ["entertainment_adjustment"],
    "연말정산": ["wage_income_tax", "special_tax_credit", "personal_deduction", "earned_income_gross_pay"],
    "월정액급여": ["earned_income_gross_pay"], "비과세 근로소득": ["earned_income_gross_pay"], "식대": ["earned_income_gross_pay"],
    "가지급금": ["nonbusiness_interest_disallowance", "deemed_interest"], "업무무관": ["nonbusiness_interest_disallowance"],
    "지급이자 손금불산입": ["nonbusiness_interest_disallowance"], "채권자불분명": ["nonbusiness_interest_disallowance"],
    "건설자금이자": ["nonbusiness_interest_disallowance"], "가수금": ["nonbusiness_interest_disallowance"], "급료 가불": ["nonbusiness_interest_disallowance"],
    "기업구매자금": ["nonbusiness_interest_disallowance"], "미지급소득 대납": ["nonbusiness_interest_disallowance"],
    "인정이자": ["deemed_interest"], "간주임대료": ["housing_deemed_rent_income", "vat_deemed_rent", "corporate_deemed_rent_income"],
    "주택 임대 보증금": ["housing_deemed_rent_income"], "전세 보증금 소득세": ["housing_deemed_rent_income"], "주택임대": ["housing_rental_income", "housing_deemed_rent_income"],
    "주택임대소득": ["housing_rental_income"], "월세 소득세": ["housing_rental_income"], "분리과세 주택임대": ["housing_rental_income"],
    "1주택 비과세 임대": ["housing_rental_income"], "임대사업자 등록": ["housing_rental_income"], "2천만원 이하 임대": ["housing_rental_income"],
    "임대보증금": ["corporate_deemed_rent_income", "vat_deemed_rent"], "상속세": ["inheritance_tax", "inheritance_gift_valuation"],
    "증여세": ["gift_tax", "inheritance_gift_valuation"], "재산평가": ["inheritance_gift_valuation"], "시가": ["inheritance_gift_valuation"],
    "무상주": ["deemed_dividend_bonus_shares"], "감자": ["deemed_dividend_capital_reduction"], "의제배당": ["deemed_dividend_bonus_shares", "deemed_dividend_capital_reduction"],
    "금융소득": ["financial_income_aggregate"], "그로스업": ["financial_income_aggregate"], "gross-up": ["financial_income_aggregate"],
    "배당세액공제": ["dividend_tax_credit"], "비교산출세액": ["dividend_tax_credit"], "비교과세": ["dividend_tax_credit"], "금융소득 종합과세 산출세액": ["dividend_tax_credit", "financial_income_aggregate"],
    "결손금": ["income_loss_offset", "loss_carryforward_limit"], "이월결손금": ["loss_carryforward_limit", "income_loss_offset"],
    "연금저축": ["pension_account_withdrawal", "pension_income"], "연금수령한도": ["pension_account_withdrawal"], "IRP": ["pension_account_withdrawal"],
    "퇴직금": ["retirement_income_tax", "retirement_allowance_adjustment"], "퇴직소득": ["retirement_income_tax"],
    "감가상각": ["depreciation_adjustment"], "대손": ["bad_debt_adjustment", "bad_debt_allowance", "vat_bad_debt_credit"],
    "재고자산": ["inventory_valuation_adjustment"], "업무용승용차": ["business_car_expense"], "법인차": ["business_car_expense"],
    "기부금": ["donation_limit"], "최저한세": ["minimum_tax"], "외국납부세액": ["foreign_tax_credit"],
    "재해손실": ["disaster_loss_credit"], "부당행위": ["unfair_transaction", "deemed_interest"], "특수관계": ["related_judge", "kinship_check"],
    "친족": ["kinship_check"], "일감몰아주기": ["tunnelling_gift_nts", "tunnelling_gift", "dominant_shareholder"],
    "가산세": ["underreporting_penalty", "late_payment_penalty", "invoice_penalty"], "납부지연": ["late_payment_penalty"],
    "세금계산서": ["invoice_penalty", "vat_input_tax_deduction"], "공통매입세액": ["common_input_tax", "vat_common_input_allocation", "input_tax_recalculation"],
    "면세전용": ["vat_exempt_use_deemed_supply"], "간주공급": ["vat_deemed_supply_items", "vat_exempt_use_deemed_supply"],
    "의제매입": ["vat_deemed_input_credit"], "농산물": ["vat_deemed_input_credit"], "간이과세": ["simplified_vat_return", "vat_simplified_taxpayer"],
    "신용카드매출": ["vat_card_sales_credit"], "일괄공급": ["bundled_supply_split"], "토지 건물 안분": ["bundled_supply_split"],
    "공급시기": ["vat_supply_timing"], "장기할부": ["vat_supply_timing"], "중간지급조건부": ["vat_supply_timing"],
    "완성도기준": ["vat_supply_timing"], "선발급": ["vat_supply_timing"], "세금계산서 먼저": ["vat_supply_timing"],
    "예정신고 귀속": ["vat_supply_timing"], "과세기간 귀속": ["vat_supply_timing"],
    "비거주자": ["nonresident_withholding"], "원천징수": ["withholding_tax", "nonresident_withholding", "daily_worker_withholding"],
    "일용": ["daily_worker_withholding"], "추계": ["estimated_business_income"], "경비율": ["estimated_business_income"],
    "인적공제": ["personal_deduction"], "자녀세액공제": ["child_tax_credit"], "자녀공제": ["child_tax_credit", "personal_deduction", "inheritance_tax"],
    "출산공제": ["child_tax_credit"], "출생": ["child_tax_credit"], "입양": ["child_tax_credit"], "손자녀": ["child_tax_credit"],
    "다자녀": ["child_tax_credit"], "자녀 세액": ["child_tax_credit"], "의료비": ["special_tax_credit"], "교육비": ["special_tax_credit"], "보험료": ["special_tax_credit"],
    "기타소득": ["other_income_aggregate"], "종교인소득": ["religious_income"], "종교인": ["religious_income"], "목사": ["religious_income"], "사례비": ["religious_income"], "승려": ["religious_income"], "복권": ["other_income_aggregate"], "강연료": ["other_income_aggregate"],
    "연금소득": ["pension_income"], "제척기간": ["assessment_limitation"], "과소자본": ["thin_capitalization"],
    "수입배당금": ["dividend_received_deduction"], "엑셀": ["export_workbook"], "서식": ["export_workbook"], "신고서 검토": ["return_precheck", "precheck_schema"],
    "국세 우선": ["tax_priority_distribution"], "국세우선권": ["tax_priority_distribution"], "당해세": ["tax_priority_distribution"],
    "법정기일": ["tax_priority_distribution"], "배당표": ["tax_priority_distribution"], "배당 순위": ["tax_priority_distribution"], "매각대금 배분": ["tax_priority_distribution"],
    "경매": ["tax_priority_distribution"], "공매": ["tax_priority_distribution"], "매각대금": ["tax_priority_distribution"],
    "압류선착": ["tax_priority_distribution"], "저당권": ["tax_priority_distribution"], "전세권": ["tax_priority_distribution"],
    "소액임차": ["tax_priority_distribution"], "최우선변제": ["tax_priority_distribution"], "임금채권": ["tax_priority_distribution"],
    # 평소 말(처음 쓰는 사람 점검 2026-10-30): 연봉·월급·프리랜서·배당 같은 생활어
    "연봉": ["wage_income_tax", "earned_income_gross_pay"], "월급": ["wage_income_tax", "earned_income_gross_pay"],
    "직장인": ["wage_income_tax", "earned_income_gross_pay"], "근로자 세금": ["wage_income_tax"], "급여 세금": ["wage_income_tax"],
    "근로소득": ["wage_income_tax", "earned_income_gross_pay"], "총급여": ["wage_income_tax", "earned_income_gross_pay"],
    "프리랜서": ["estimated_business_income", "withholding_tax", "income_tax"], "3.3%": ["withholding_tax", "estimated_business_income"],
    "3.3": ["withholding_tax", "estimated_business_income"], "사업소득": ["estimated_business_income", "income_loss_offset", "income_tax"],
    "종합소득세": ["income_tax", "estimated_business_income", "financial_income_aggregate"], "종소세": ["income_tax", "estimated_business_income"],
    "배당": ["financial_income_aggregate", "dividend_tax_credit"], "이자소득": ["financial_income_aggregate"], "예금이자": ["financial_income_aggregate"],
    "종합과세": ["financial_income_aggregate", "dividend_tax_credit"], "2천만원 초과": ["financial_income_aggregate", "dividend_tax_credit"],
    "2천만 원 초과": ["financial_income_aggregate", "dividend_tax_credit"], "금융소득종합과세": ["financial_income_aggregate", "dividend_tax_credit"],
    "증여": ["gift_tax", "inheritance_gift_valuation"], "상속": ["inheritance_tax", "inheritance_gift_valuation"],
    "법인세": ["corporate_tax", "income_adjustment_total"], "퇴직소득세": ["retirement_income_tax"],
    "주식 양도": ["stock_capital_gains_tax", "stock_major_shareholder"], "주식 양도세": ["stock_capital_gains_tax", "stock_major_shareholder"], "주식 팔": ["stock_capital_gains_tax", "stock_major_shareholder"],
    "대주주": ["stock_major_shareholder", "stock_capital_gains_tax"], "종가": ["stock_price_lookup"], "공시지가": ["official_property_price"], "개인 기부금": ["individual_donation_expense"], "기부금 필요경비": ["individual_donation_expense"], "사업자 기부금 한도": ["individual_donation_expense"], "연금수령한도": ["pension_account_withdrawal"], "연금수령연차": ["pension_account_withdrawal"], "연금계좌 인출": ["pension_account_withdrawal"], "연금외수령": ["pension_account_withdrawal"], "건물 기준시가": ["building_standard_price"], "건물 평가": ["building_standard_price"], "상가 기준시가": ["building_standard_price"], "신축가격기준액": ["building_standard_price"], "공시가격": ["official_property_price"], "기준시가 조회": ["official_property_price"], "공동주택가격": ["official_property_price"], "개별주택가격": ["official_property_price"], "아파트 공시": ["official_property_price"], "비상장주식 평가": ["unlisted_stock_valuation", "stock_basis_income_tax"], "보충적 평가": ["unlisted_stock_valuation"], "순손익가치": ["unlisted_stock_valuation"], "순자산가치": ["unlisted_stock_valuation", "stock_basis_income_tax"], "최대주주 할증": ["unlisted_stock_valuation"], "기준시가 주식": ["stock_basis_income_tax"], "시세": ["stock_price_lookup"], "주가 조회": ["stock_price_lookup"], "시가총액": ["stock_price_lookup", "stock_major_shareholder"], "비상장주식": ["stock_capital_gains_tax", "stock_major_shareholder"], "장외 양도": ["stock_capital_gains_tax"], "소액주주": ["stock_major_shareholder", "stock_capital_gains_tax"],
    "양도": ["capital_gains_income", "capital_gains_tax"], "필요경비": ["capital_gains_income", "estimated_business_income", "other_income_aggregate"],
}

# 계산 도구가 없는 흔한 질문 — 있는 척 엉뚱한 도구를 고르지 않도록 먼저 걸러 근거 조문만 안내(처음 쓰는 사람 점검 2026-10-30)
_NOT_SUPPORTED: list[tuple[tuple[str, ...], str, list[tuple[str, str]]]] = [
    (("월세", "세액공제"), "월세액 세액공제(조세특례제한법 제95조의2)", [("조세특례제한법", "제95조의2")]),
    (("월세액",), "월세액 세액공제(조세특례제한법 제95조의2)", [("조세특례제한법", "제95조의2")]),
    (("연금", "세액공제"), "연금계좌 세액공제(소득세법 제59조의3) — 연금저축·IRP 납입액 공제", [("소득세법", "제59조의3")]),
    (("종합부동산세",), "종합부동산세(종합부동산세법)", [("종합부동산세법", "제7조"), ("종합부동산세법", "제9조")]),
    (("종부세",), "종합부동산세(종합부동산세법)", [("종합부동산세법", "제7조"), ("종합부동산세법", "제9조")]),
    (("신용카드", "소득공제"), "신용카드 등 사용금액 소득공제(조세특례제한법 제126조의2)", [("조세특례제한법", "제126조의2")]),
    (("근로장려금",), "근로장려금(조세특례제한법 제100조의2~)", [("조세특례제한법", "제100조의2"), ("조세특례제한법", "제100조의5")]),
    (("자녀장려금",), "자녀장려금(조세특례제한법 제100조의27~)", [("조세특례제한법", "제100조의27"), ("조세특례제한법", "제100조의29")]),
    (("취득세",), "취득세(지방세법 — 국세 아님)", [("지방세법", "제10조"), ("지방세법", "제11조")]),
    (("재산세",), "재산세(지방세법 — 국세 아님)", [("지방세법", "제110조"), ("지방세법", "제111조")]),
    (("자동차세",), "자동차세(지방세법 — 국세 아님)", [("지방세법", "제127조")]),
]
_LAW_MCP_HINT = {
    "이름": "korean-tax-mcp(세법 근거 MCP — 조문·부칙·해석·판례)",
    "레포": "https://github.com/UpstageAI/korean-tax-mcp",
    "설치": "uvx korean-tax-mcp — Claude Desktop이면 mcpServers에 {\"korean-tax\": {\"command\": \"uvx\", \"args\": [\"korean-tax-mcp\"]}} 추가(조문 조회는 무료 법제처 키 LAW_OC)",
    "안내": "이 서버가 연결돼 있지 않으면 사용자에게 위 설치를 권하고, 조문 원문 없이 답할 때는 그 사실을 밝히세요",
}

_NOT_SUPPORTED_EXEMPT = ("양도", "필요경비", "취득가액", "보유세")  # 양도 계산 질문에 섞인 '취득세·재산세'는 걸러내지 않음


def _bigrams(text: str) -> set[str]:
    t = "".join(ch for ch in text.lower() if not ch.isspace())
    return {t[i:i + 2] for i in range(len(t) - 1)}


@mcp.tool(annotations=CALC)
def find_tool(query: Annotated[str, Field(description="찾는 계산을 평소 말로. 예: '아파트 팔았을 때 양도세', '보험차익으로 산 기계', '접대비 한도', '해산한 법인 청산소득'")] = "",
              n: Annotated[int, Field(description="돌려줄 도구 수(1~10)", ge=1, le=10)] = 5,
              keyword: Annotated[str, Field(description="query와 같음(korean-tax-mcp find_article과 같은 이름으로 불러도 되게)")] = "") -> dict:
    """맞는 계산 도구 찾기 — 도구가 70개가 넘으니, 어떤 도구를 쓸지 모르겠으면 먼저 이걸 부른다.
    언제: 계산 문제를 받았는데 맞는 도구 이름이 바로 떠오르지 않을 때, 또는 '이 계산 도구는 없다'고 판단하기 전에 반드시.
    입력: 찾는 계산 → query(평소 말, 세목·쟁점·조문 무엇이든), 개수 → n.
    결과: 후보 도구 목록(이름·한 줄 설명·언제 쓰는지·일치 이유). 후보가 없으면 계산 도구가 없는 것이니 korean-tax-mcp로 조문을 찾는다."""
    tools = _tool_catalog()
    q = (query or keyword).strip()
    if not q:
        return {"error": "입력 확인 필요 — query(찾는 계산)를 넣으세요"}
    qn = q.replace(" ", "").lower()
    for words, what, arts in _NOT_SUPPORTED:
        if all(w.replace(" ", "").lower() in qn for w in words) and not (words[0] in ("취득세", "재산세", "자동차세") and any(e in qn for e in _NOT_SUPPORTED_EXEMPT)):
            return {"결과": {"후보": [], "없는 계산": what,
                            "근거 조회": [{"도구": "law_article", "law_name": ln, "article": ar} for ln, ar in arts],
                            "근거 서버": _LAW_MCP_HINT,
                            "안내": f"{what}은(는) 이 서버에 계산 도구가 없습니다. 위 '근거 조회'대로 korean-tax-mcp law_article을 불러 조문 원문을 읽고 그에 따라 직접 계산하되, "
                                    "계산 도구 없이 계산했다는 사실을 답에 밝히세요. (지방세는 korean-tax-mcp에도 없을 수 있음 — 그러면 법제처 원문으로)"},
                    "근거": "도구 목록 검색(세무 자문 아님)"}
    qb = _bigrams(q)
    scores: dict[str, float] = {}
    why: dict[str, list[str]] = {}
    matched = [alias for alias in _TOOL_ALIASES if alias.replace(" ", "").lower() in qn]
    # 긴 별칭이 맞으면 그 안에 든 짧은 별칭은 세지 않음('수입배당금'이 맞으면 '배당'은 무시)
    matched = [a for a in matched if not any(a != b and a.replace(" ", "").lower() in b.replace(" ", "").lower() for b in matched)]
    for alias in matched:
        for rank, name in enumerate(_TOOL_ALIASES[alias]):
            scores[name] = scores.get(name, 0) + 3.0 - rank * 0.5
            why.setdefault(name, []).append(f"'{alias}'")
    for name, desc in tools.items():
        if name.lower() in q.lower():
            scores[name] = scores.get(name, 0) + 5.0
            why.setdefault(name, []).append("도구 이름")
        words = {w for w in re.findall(r"[a-z]+", q.lower()) if len(w) > 2}
        name_words = set(name.split("_")) | set(re.findall(r"[a-z]+", desc.lower()))
        hit = words & name_words
        if hit:
            scores[name] = scores.get(name, 0) + 1.5 * len(hit)
            why.setdefault(name, []).append("영어 낱말 " + ", ".join(sorted(hit)))
        db = _bigrams(desc[:600])
        if qb and db:
            ov = len(qb & db) / len(qb)
            if ov > 0:
                scores[name] = scores.get(name, 0) + ov * 2.0
    # 별칭·이름 일치가 없고 글자 겹침만 약하면(점수 < 1.0) 맞는 도구가 없는 것으로 본다
    ranked = [k for k in sorted((k for k in scores if k in tools), key=lambda k: -scores[k]) if scores[k] >= 1.0][:n]
    if not ranked:
        return {"결과": {"후보": [], "안내": "맞는 계산 도구가 없습니다. 조문·기한·절차 질문이면 korean-tax-mcp(find_article·law_article)로 근거를 찾아 답하고, 계산 도구가 없다는 사실을 답에 밝히세요."},
                "근거": "도구 목록 검색(세무 자문 아님)"}
    out = []
    for name in ranked:
        desc = tools[name]
        lines = [ln.strip() for ln in desc.splitlines() if ln.strip()]
        언제 = next((ln for ln in lines if ln.startswith("언제")), "")
        out.append({"도구": name, "설명": lines[0][:160] if lines else "", "언제": 언제[:200],
                    "일치": ", ".join(why.get(name, [])) or "설명문 낱말 겹침", "점수": round(scores[name], 2)})
    return {"결과": {"후보": out,
                    "안내": "후보의 설명·입력 스키마를 확인한 뒤 호출하세요. 후보에 맞는 도구가 없으면 이 서버에 해당 계산이 없는 것이니 "
                            "조문 계산은 korean-tax-mcp(find_article·law_article)로 근거를 찾아 직접 계산하고 그 사실을 답에 밝히세요.",
                    "근거 조회": "계산 결과의 '근거'에 적힌 조문(예: 소득세법 제47조①)은 korean-tax-mcp law_article(law_name, article, as_of=귀속연도 말일)로 원문을 확인해 답에 함께 인용하세요. "
                                "사실관계가 조문의 요건에 맞는지(비과세·공제 요건, 적용 시기)는 도구가 아니라 조문으로 판단합니다.",
                    "근거 서버": _LAW_MCP_HINT},
            "근거": "도구 목록 검색(세무 자문 아님)"}


_CATALOG: dict[str, str] | None = None


def _tool_catalog() -> dict[str, str]:
    global _CATALOG
    if _CATALOG is None:
        cat = {}
        for t in mcp._tool_manager.list_tools():
            if t.name != "find_tool":
                cat[t.name] = (t.title or "") + "\n" + (t.description or "")
        _CATALOG = cat
    return _CATALOG


if __name__ == "__main__":
    main()
