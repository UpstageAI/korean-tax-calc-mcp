"""살아 있는 수식(계산식 단계) 정답 테스트. 작성 Claude 2026-10-24 (docs/SPEC_export_workbook_c.md)."""
import asyncio, json, re, shutil, subprocess
import openpyxl
import pytest
from korean_tax_calc_mcp.server import mcp


def c(tool, a):
    r = asyncio.run(mcp.call_tool(tool, a))
    return json.loads((r.content[0] if hasattr(r, "content") else r[0][0]).text)


ENT = {"year": 2025, "sme": True, "ledger": [
    {"항목": "일반 접대비", "금액": 52000000, "증빙": "적격"},
    {"항목": "체육활동 입장권", "금액": 20000000, "증빙": "적격", "성격": "문화"},
    {"항목": "직원단체(법인) 복리시설비", "금액": 4000000, "증빙": "적격", "성격": "직원단체_법인"},
    {"항목": "거래처 증정 제품", "금액": 9000000, "증빙": "적격", "성격": "현물", "시가": 10000000, "장부가액": 8000000}],
    "revenue": {"매출액": 15000000000, "특수관계매출": 3000000000}}
ADJ = {"net_income": 300000000, "items": [
    {"항목": "법인세비용", "종류": "법인세비용", "금액": 20000000},
    {"항목": "기업업무추진비 한도초과", "종류": "기업업무추진비 한도초과", "금액": 2280000},
    {"항목": "기계장치 감가상각비 한도초과", "종류": "감가상각비 한도초과", "금액": 5550000},
    {"항목": "㈜B 수입배당금 익금불산입", "종류": "수입배당금 익금불산입", "금액": 7200000}]}
NAME = re.compile(r"^[가-힣A-Za-z][가-힣A-Za-z0-9_]*$")


def _py(expr):
    e = re.sub(r"\bROUNDDOWN\(", "_RD(", expr)
    assert not re.search(r"(?<![<>])=(?!=)", e), f"= 비교 금지: {expr}"
    return e


def evaluate(r):
    env = {"MIN": min, "MAX": max, "IF": lambda c, a, b: a if c else b, "_RD": lambda x, n=0: int(x), "TRUE": True, "FALSE": False}
    for i in r["계산식_입력"]:
        assert NAME.match(i["이름"]), i["이름"]
        env[i["이름"]] = i["값"]
    out = {}
    for s in r["계산식"]:
        assert NAME.match(s["이름"]), s["이름"]
        names = set(re.findall(r"[가-힣A-Za-z][가-힣A-Za-z0-9_]*", s["식"])) - {"MIN", "MAX", "IF", "ROUNDDOWN", "TRUE", "FALSE"}
        assert names and names <= set(env), (s["이름"], s["식"], names - set(env))
        v = eval(_py(s["식"]), {"__builtins__": {}}, env)
        assert abs(v - s["값"]) < 1, (s["이름"], s["식"], v, s["값"])
        env[s["이름"]] = v
        out[s["이름"]] = s
    return out


def test_entertainment_formula_steps():
    r = c("entertainment_adjustment", ENT)
    steps = evaluate(r)
    assert {"기본한도", "수입금액한도", "문화한도", "한도", "시부인대상", "한도초과"} <= set(steps)
    assert steps["한도초과"]["값"] == r["결과"]["한도초과"] == 2280000
    assert steps["한도"]["값"] == r["결과"]["한도"]
    assert "시부인대상" in steps["한도초과"]["식"] and "한도" in steps["한도초과"]["식"]
    assert any("11000000" in str(i["값"]) or i["값"] == 11000000 for i in r["계산식_입력"])  # 현물 평가액은 판단 결과 입력


def test_income_adjustment_formula_steps():
    r = c("income_adjustment_total", ADJ)
    steps = evaluate(r)
    assert steps["차가감소득금액"]["값"] == r["결과"]["차가감소득금액"] == 320630000
    assert steps["익금산입_손금불산입"]["값"] == 27830000 and steps["손금산입_익금불산입"]["값"] == 7200000
    assert "당기순이익" in steps["차가감소득금액"]["식"]


def _formula_sheet(tmp_path):
    r = c("export_workbook", {"layout": "steps", "title": "수식", "calls": [{"label": "기업업무추진비", "tool": "entertainment_adjustment", "args": ENT},
                                                         {"label": "차가감", "tool": "income_adjustment_total", "args": ADJ}],
                               "output_dir": str(tmp_path)})
    return r["결과"]["파일"]


def test_workbook_has_live_formulas_and_names(tmp_path):
    path = _formula_sheet(tmp_path)
    wb = openpyxl.load_workbook(path)
    names = set(wb.defined_names.keys()) if hasattr(wb.defined_names, "keys") else {d.name for d in wb.defined_names.definedName}
    assert {"호출1_한도초과", "호출1_한도", "호출2_차가감소득금액", "호출2_당기순이익"} <= names
    ws = wb["계산 1"]
    fs = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert any("호출1_시부인대상" in f and "호출1_한도" in f for f in fs)        # 한도초과 수식이 이름 참조
    assert not any(re.fullmatch(r"=\d+", f) for f in fs)                         # 값만 붙인 수식 없음


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice 없음")
def test_formulas_recalculate_to_tool_values(tmp_path):
    path = _formula_sheet(tmp_path)
    out = tmp_path / "recalc"
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(out), path], check=True, capture_output=True, timeout=120)
    wb = openpyxl.load_workbook(next(out.glob("*.xlsx")), data_only=True)
    checks = [c.value for s in ("계산 1", "계산 2") for row in wb[s].iter_rows() for c in row if c.value in ("일치", "확인 필요")]
    assert len(checks) >= 10 and set(checks) == {"일치"}


@pytest.mark.parametrize("mod", [{"sme": False}, {"revenue": {"매출액": 80000000000, "특수관계매출": 3000000000}},
                                 {"revenue": {"매출액": 500000000}}, {"months": 6}, {"year": 2019},
                                 {"sme": False, "revenue": {"매출액": 120000000000, "특수관계매출": 70000000000}}])
def test_entertainment_formula_steps_hold_for_other_inputs(mod):
    r = c("entertainment_adjustment", {**ENT, **mod})
    steps = evaluate(r)
    assert steps["한도초과"]["값"] == r["결과"]["한도초과"] and steps["한도"]["값"] == r["결과"]["한도"]


def test_item_inputs_carry_original_names():
    r = c("entertainment_adjustment", ENT)
    descs = " ".join(i["설명"] for i in r["계산식_입력"] if i["이름"].startswith("항목"))
    assert "거래처 증정 제품" in descs and "체육활동 입장권" in descs
