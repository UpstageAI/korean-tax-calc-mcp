"""엑셀을 법정서식 순서·모양으로 — 정답 테스트. 작성 Claude 2026-10-25 (docs/SPEC_workbook_forms.md).
값은 가상 사례(가람정밀, 2025) 도구 결과를 손 검산한 것."""
import asyncio, json, pathlib, re, shutil, subprocess
import openpyxl
import pytest
from korean_tax_calc_mcp.server import mcp

CALLS = json.loads((pathlib.Path(__file__).parent / "data" / "workbook_forms_calls.json").read_text(encoding="utf-8"))
ORDER = ["목록", "별지3", "별지15", "별지23(갑)", "별지20(1)", "별지16의2", "별지50(을)", "부가 별지22", "부가 계산근거", "종소 금융·근로소득"]


def c(tool, a):
    r = asyncio.run(mcp.call_tool(tool, a))
    return json.loads((r.content[0] if hasattr(r, "content") else r[0][0]).text)


@pytest.fixture(scope="module")
def books(tmp_path_factory):
    d = tmp_path_factory.mktemp("forms")
    r = c("export_workbook", {"title": "가람정밀 서식", "calls": CALLS, "output_dir": str(d),
                              "company": "㈜가람정밀(가상)", "period": "2025.01.01 ~ 2025.12.31"})
    path = r["결과"]["파일"]
    raw = openpyxl.load_workbook(path)
    calc = None
    if shutil.which("soffice"):
        subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(d / "re"), path],
                       check=True, capture_output=True, timeout=180)
        calc = openpyxl.load_workbook(next((d / "re").glob("*.xlsx")), data_only=True)
    return raw, calc


def sheet(wb, prefix):
    hits = [n for n in wb.sheetnames if n.startswith(prefix)]
    assert hits, (prefix, wb.sheetnames)
    return wb[hits[0]]


def amount(ws, label, col=3):
    """A(번호) 또는 B(구분)가 label로 시작하는 첫 행의 col 값."""
    for row in ws.iter_rows():
        a = str(row[0].value or ""); b = str(row[1].value or "")
        if a.startswith(label) or b.startswith(label) or (a + " " + b).startswith(label):
            return row[col - 1].value
    raise AssertionError(f"{ws.title}: '{label}' 행 없음")


def test_sheet_order_follows_forms(books):
    raw, _ = books
    names = raw.sheetnames
    idx = [next(i for i, n in enumerate(names) if n.startswith(p)) for p in ORDER]
    assert idx == sorted(idx), names
    assert "요약" not in names and not any(re.fullmatch(r"계산 \d+", n) for n in names), names
    assert names[-2:] == ["입력·가정", "판단·근거"]


def test_form_headers(books):
    raw, _ = books
    ws = sheet(raw, "별지23(갑)")
    head = " ".join(str(x.value) for row in ws.iter_rows(max_row=3) for x in row if x.value)
    assert "별지 제23호서식(갑)" in head and "㈜가람정밀(가상)" in head and "2025.01.01" in head
    for p in ORDER[1:]:
        txt = " ".join(str(x.value) for row in sheet(raw, p).iter_rows() for x in row if x.value)
        assert not re.search(r"[a-z]+_[a-z_]+", txt), (p, "도구 영문 이름이 서식 시트에 보임")


def test_amount_cells_are_formulas(books):
    raw, _ = books
    ws = sheet(raw, "별지23(갑)")
    for lab in ("⑧", "⑮", "⑯"):
        v = amount(ws, lab)
        assert isinstance(v, str) and v.startswith("="), (lab, v)
    assert re.search(r"C\d+.*C\d+.*C\d+", amount(ws, "⑧"))          # ④+⑥+⑦ 셀 참조
    assert "별지15" in amount(sheet(raw, "별지3"), "102")             # 합계표에서 옮겨 옴


need_soffice = pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice 없음")


@need_soffice
@pytest.mark.parametrize("prefix, label, want", [
    ("별지3", "101", 300000000), ("별지3", "102", 27830000), ("별지3", "103", 7200000),
    ("별지3", "104", 320630000), ("별지3", "107", 320630000),
    ("별지23(갑)", "①", 87000000), ("별지23(갑)", "④", 36000000), ("별지23(갑)", "⑧", 78400000), ("별지23(갑)", "⑯", 2280000),
    ("별지23(갑)", "⑩", 6000000), ("별지23(갑)", "⑮", 84720000), ("별지23(갑)", "⑰", 84720000),
    ("별지20(1)", "⑤", 100000000), ("별지20(1)", "⑪", 4000000), ("별지20(1)", "⑫", 31500000), ("별지20(1)", "㉒", 9450000),
    ("별지20(1)", "㉓", 15000000), ("별지20(1)", "㉖", 5550000), ("별지20(1)", "㉘", 9550000),
])
def test_form_amounts_after_recalc(books, prefix, label, want):
    _, calc = books
    assert amount(sheet(calc, prefix), label) == want


@need_soffice
def test_adjustment_summary_table(books):
    _, calc = books
    ws = sheet(calc, "별지15")
    rows = [[x.value for x in r] for r in ws.iter_rows()]
    left = {r[0]: (r[1], r[2], r[3]) for r in rows if r[0] and isinstance(r[1], (int, float))}
    right = {r[4]: (r[5], r[6], r[7]) for r in rows if len(r) > 7 and r[4] and isinstance(r[5], (int, float))}
    assert left["법인세비용"][0] == 20000000 and str(left["법인세비용"][2]) == "500"
    assert left["기업업무추진비 한도초과"][0] == 2280000 and str(left["기업업무추진비 한도초과"][2]) == "500"
    assert left["감가상각비 한도초과"][0] == 5550000 and str(left["감가상각비 한도초과"][2]) == "400"
    assert any(k.startswith("수입배당금") and v[0] == 7200000 and str(v[2]) == "200" for k, v in right.items())
    assert left["합계"][0] == 27830000 and right["합계"][0] == 7200000


@need_soffice
def test_dividend_and_reserve_tables(books):
    _, calc = books
    rows = [[x.value for x in r] for r in sheet(calc, "별지16의2").iter_rows()]
    b = next(r for r in rows if r[0] and str(r[0]).startswith("B"))
    assert b[1:6] == [15000000, 80, 12000000, 4800000, 7200000]
    rows = [[x.value for x in r] for r in sheet(calc, "별지50(을)").iter_rows()]
    dep = next(r for r in rows if r[0] and "감가상각비" in str(r[0]))
    assert dep[1:5] == [4000000, 0, 5550000, 9550000]


@need_soffice
def test_vat_and_income_sheets(books):
    _, calc = books
    txt = [[x.value for x in r] for r in sheet(calc, "부가 계산근거").iter_rows()]
    flat = [v for r in txt for v in r]
    assert 10400000 in flat and 1040000 in flat                          # 공통사용재화 과세표준·세액
    assert 300000000 in flat and 200000000 in flat and 30000000 in flat  # 일괄양도 건물·토지·부가세
    flat = [v for r in sheet(calc, "종소 금융·근로소득").iter_rows() for v in [x.value for x in r]]
    assert 500000 in flat and 31500000 in flat                           # Gross-up 가산액·배당소득금액
    assert 138400000 in flat and 15518000 in flat and 122882000 in flat  # 총급여·근로소득공제·소득금액


@need_soffice
def test_all_checks_match(books):
    _, calc = books
    checks = [x.value for ws in calc.worksheets for r in ws.iter_rows() for x in r if x.value in ("일치", "확인 필요")]
    assert len(checks) >= 30 and set(checks) == {"일치"}


def test_unmapped_tool_keeps_calc_sheet(tmp_path):
    r = c("export_workbook", {"title": "기타", "output_dir": str(tmp_path),
                              "calls": [{"label": "청산소득", "tool": "liquidation_income",
                                         "args": {"assets_realized": 900000000, "liabilities_paid": 300000000, "capital": 200000000, "surplus": 100000000}}]})
    wb = openpyxl.load_workbook(r["결과"]["파일"])
    assert any(re.fullmatch(r"계산 \d+", n) for n in wb.sheetnames), wb.sheetnames
