"""PR 리뷰 반영(교차 참조 대체값·회색 칸) 정답 테스트. 작성 Claude 2026-10-25 (docs/SPEC_form_templates4.md)."""
import copy, json, pathlib
import openpyxl
import pytest
from korean_tax_calc_mcp.forms_workbook import cross_fallback_value, load_templates
from test_form_templates import CALLS, build, c

ENT = next(x for x in CALLS if x["tool"] == "entertainment_adjustment")


def test_cross_fallback_uses_target_slot_check():
    res = c("entertainment_adjustment", ENT["args"])
    tpls = load_templates()
    assert cross_fallback_value("별지23갑", "⑯", templates=tpls, results={"entertainment_adjustment": res}) == 2280000
    with pytest.raises(KeyError):
        cross_fallback_value("별지23갑", "②", templates=tpls, results={"entertainment_adjustment": res})   # check 없는 slot
    with pytest.raises(KeyError):
        cross_fallback_value("별지23갑", "⑯", templates=tpls, results={})                                # 도구 결과 없음


@pytest.mark.parametrize("variant", ["demo", "straight_line"])
def test_no_literal_none_in_cells(tmp_path, variant):
    calls = copy.deepcopy(CALLS)
    if variant == "straight_line":
        dep = next(x for x in calls if x["tool"] == "depreciation_adjustment")
        dep["args"].update(method="정액법", rate=0.1, booked_depreciation=12000000, prior_disallowed=0, expensed_capex=0)
    wb = openpyxl.load_workbook(build(tmp_path, calls))
    bad = [(ws.title, x.coordinate) for ws in wb.worksheets for r in ws.iter_rows() for x in r
           if isinstance(x.value, str) and ("None" in x.value)]
    assert not bad, bad


def test_cell_fill_applied(tmp_path):
    wb = openpyxl.load_workbook(build(tmp_path, CALLS))
    assert wb["별지15 소득금액조정합계표"]["C31"].fill.fgColor.rgb.endswith("808080")
    assert wb["별지20(1) 감가상각비(정률법)"]["E36"].fill.fgColor.rgb.endswith("BFBFBF")
    assert wb["별지3 과세표준및세액조정계산서"]["E12"].fill.fgColor.rgb.endswith("808080")
