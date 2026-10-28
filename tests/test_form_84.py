"""별지 제84호서식(양도소득과세표준 신고 및 납부계산서 제1쪽)·부표1(양도소득금액 계산명세서) 틀 테스트.

값은 가상 사례 손 검산(소득세법 제95조·제103조·제104조·제55조①, 시행령 제160조①).
"""
import copy
import shutil

import openpyxl
import pytest

from test_form_templates import build, recalc, c

S84, S84_2, S84T = "별지84 양도소득과세표준 신고", "별지84 양도소득과세표준 신고 (2)", "별지84 신고서(세율 적용)"
B1, B1_2 = "별지84부표1 양도소득금액 계산명세서", "별지84부표1 양도소득금액 계산명세서 (2)"
soffice = pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice 없음")

# 1세대1주택 고가주택: 2015.3.10. 취득 8억, 2025.6.20. 양도 20억, 양도비 2천만, 거주 10년
#   양도차익 = 20억 − 8억 − 2천만 = 11.8억(제97조)
#   과세대상 양도차익 = 11.8억 × (20억 − 12억)/20억 = 4.72억(영 제160조①, 2021.12.8. 이후 기준 12억)
#   장기보유특별공제 = 11.8억 × 80%(표2 보유 40%+거주 40%) × 0.4 = 3.776억(제95조②)
#   양도소득금액 9,440만 − 기본공제 250만(제103조①) = 과세표준 9,190만
#   산출세액 = 9,190만 × 35% − 1,544만(2023~ 제55조①) = 16,725,000
HIGH = dict(asset_type="주택", acquisition_date="2015-03-10", transfer_date="2025-06-20", transfer_price=2_000_000_000,
            acquisition_price=800_000_000, transfer_expenses=20_000_000, one_house=True, exempt_met=True,
            residence_years=10, apply_basic_deduction=True)
# 조정대상지역 3주택 이상 중과: 2019.3.10. 취득 5억, 2026.7.20. 양도 9억(한시 배제 기한 2026.5.9. 경과), 양도비 1천만
#   양도차익 3.9억, 장기보유특별공제 배제(제95조② 괄호), 과세표준 3.875억
#   산출세액 = 3.875억 × 40% − 2,594만 = 129,060,000 + 3.875억 × 30%(제104조⑦3호) 116,250,000 = 245,310,000
HEAVY = dict(asset_type="주택", acquisition_date="2019-03-10", transfer_date="2026-07-20", transfer_price=900_000_000,
             acquisition_price=500_000_000, transfer_expenses=10_000_000, one_house=False, houses_count=3,
             adjusted_area=True, apply_basic_deduction=True)
TAX = dict(base=91_900_000, asset_type="주택", acquisition_date="2015-03-10", transfer_date="2025-06-20")
CALLS = [{"label": "고가주택 양도", "tool": "capital_gains_income", "args": HIGH},
         {"label": "중과 주택 양도", "tool": "capital_gains_income", "args": HEAVY},
         {"label": "세율 적용", "tool": "capital_gains_tax", "args": TAX}]


@pytest.fixture(scope="module")
def wb(tmp_path_factory):
    d = tmp_path_factory.mktemp("f84")
    path = build(d, CALLS, period="")
    return openpyxl.load_workbook(path), (recalc(path, d) if shutil.which("soffice") else None)


def test_tool_form_values():
    r = c("capital_gains_income", HIGH)
    k = r["계산"]
    assert (k["세율구분_코드"], k["자산종류_코드"], k["장기보유특별공제_코드"]) == ("1-10", "2", "01")
    assert (k["양도가액"], k["취득가액"], k["기타필요경비"], k["고가주택_기준"]) == (2_000_000_000, 800_000_000, 20_000_000, 1_200_000_000)
    assert (k["세율"], k["누진공제"], k["가산세율"], k["비과세_양도차익"]) == (0.35, 15_440_000, 0.0, 708_000_000)
    assert {i["이름"] for i in r["계산식_입력"]} >= {"양도가액", "취득가액", "기타필요경비", "비과세_전액", "고가주택_기준", "공제율", "기본공제", "세율", "누진공제", "가산세율"}
    assert r["결과"]["산출세액"] == 16_725_000 and r["결과"]["양도소득_과세표준"] == 91_900_000   # 기존 결과 키 그대로
    h = c("capital_gains_income", HEAVY)["계산"]
    assert (h["세율구분_코드"], h["장기보유특별공제_코드"], h["세율"], h["누진공제"], h["가산세율"]) == ("1-49", "03", 0.4, 25_940_000, 0.3)
    assert "배제" in h["장기보유특별공제_배제"]


def test_rate_codes_short_term_heavy():
    # 2주택 중과인데 보유 1년 미만: 70%(제104조①3호) 7천만 > 기본세율 1,956만 + 20%p 2천만 → 코드 1-84(작성방법 ㉘)
    r = c("capital_gains_tax", dict(base=100_000_000, asset_type="주택", acquisition_date="2025-09-01",
                                    transfer_date="2026-07-01", houses_count=2, adjusted_area=True))
    assert r["결과"]["산출세액"] == 70_000_000
    assert (r["계산"]["세율구분_코드"], r["계산"]["세율"], r["계산"]["누진공제"], r["계산"]["가산세율"]) == ("1-84", 0.7, 0, 0.0)
    # 비사업용 토지 2년 이상: 기본세율 + 10%p(제104조①8호) → 1-11
    r = c("capital_gains_tax", dict(base=50_000_000, asset_type="토지", acquisition_date="2018-01-01",
                                    transfer_date="2025-01-01", non_business_land=True))
    assert r["계산"]["세율구분_코드"] == "1-11" and r["계산"]["가산세율"] == 0.1
    assert r["결과"]["산출세액"] == 50_000_000 * 15 // 100 - 1_260_000 + 5_000_000   # 15% 구간 누진공제 126만


def test_sheets_and_formulas(wb):
    raw, _ = wb
    n = raw.sheetnames
    assert [x for x in n if x.startswith("별지84")] == [S84, S84_2, S84T, B1, B1_2]
    assert not any(x.startswith("계산 ") for x in n)
    w = raw[S84]
    assert w["E7"].value == f"='{B1}'!E30"              # ④ = 부표1 ⑱
    assert raw[S84_2]["E7"].value == f"='{B1_2}'!E30"   # 둘째 call은 자기 부표1을 가리킴
    assert w["E10"].value == "=호출1_기본공제" and w["E11"].value == "=MAX(0,E7+E8-E9-E10)"
    assert w["E12"].value == "=호출1_세율+호출1_가산세율" and w["E12"].number_format == "0%"
    assert w["E13"].value == "=ROUNDDOWN(ROUND(E11*호출1_세율,4),0)-호출1_누진공제+ROUNDDOWN(ROUND(E11*호출1_가산세율,4),0)"
    assert w["E24"].value == "=E13-E14-E15-E16-E17-E18+E22-E23"
    assert (w["D11"].value, w["C11"].value) == ("=SUM(E11:G11)", "=D11+H11")
    assert w["E6"].value == "( 1-10 )" and raw[S84_2]["E6"].value == "( 1-49 )"
    assert "2025 년 귀속" in w["C3"].value
    b = raw[B1]
    assert b["E20"].value == "=호출1_양도가액" and b["E25"].value == "=MAX(0,E20-E21-E23-E24)"
    assert "호출1_고가주택_기준" in b["E27"].value and "호출1_공제율" in b["F28"].value
    assert b["E30"].value == "=MAX(0,E27-F28)"
    assert (b["E5"].value, b["E8"].value, b["E28"].value) == ("( 1-10 )", "( 2 )", "( 01 )")
    assert (b["E17"].value, b["F17"].value, b["E29"].value) == ("10년 이상", "11년 미만", "10년 이상")
    assert raw[B1_2]["E28"].value == "( 03 )" and raw[B1_2]["E17"].value == "년 이상"
    assert "배제" in raw[B1_2]["A41"].value
    t = raw[S84T]
    assert t["E11"].value == "=호출3_과세표준" and t["E7"].value is None
    assert raw[S84].print_area and raw[S84].page_setup.paperSize == 9


def test_list_omitted_parts(wb):
    raw, _ = wb
    txt = "\n".join(" ".join(str(x.value) for x in r if x.value) for r in raw["목록"].iter_rows())
    for s in ("①신고인·②양수인", "농어촌특별세", "부표1 ②소재지", "부표2·부표2의2·부표3", "④~⑦(세율 적용 시트)"):
        assert s in txt
    assert txt.count("자료 없음") >= 5


@soffice
def test_values_high_price_house(wb):
    _, calc = wb
    w, b = calc[S84], calc[B1]
    assert [b[a].value for a in ("E20", "E21", "E24", "E25", "E26", "E27", "F28", "E30")] == \
        [2_000_000_000, 800_000_000, 20_000_000, 1_180_000_000, 708_000_000, 472_000_000, 377_600_000, 94_400_000]
    assert [w[a].value for a in ("E7", "E10", "E11", "E13", "E24", "C13")] == \
        [94_400_000, 2_500_000, 91_900_000, 16_725_000, 16_725_000, 16_725_000]
    assert w["E12"].value == pytest.approx(0.35)


@soffice
def test_values_heavy(wb):
    _, calc = wb
    w, b = calc[S84_2], calc[B1_2]
    assert [b[a].value for a in ("E25", "E26", "E27", "F28", "E30")] == [390_000_000, 0, 390_000_000, 0, 390_000_000]
    assert [w[a].value for a in ("E7", "E11", "E13")] == [390_000_000, 387_500_000, 245_310_000]
    assert w["E12"].value == pytest.approx(0.70)
    assert calc[S84T]["E13"].value == 16_725_000


@soffice
def test_all_checks_match(wb):
    _, calc = wb
    checks = [x.value for s in (S84, S84_2, S84T, B1, B1_2) for r in calc[s].iter_rows() for x in r
              if x.value in ("일치", "확인 필요")]
    assert len(checks) == 4 + 4 + 1 + 5 + 5 and set(checks) == {"일치"}


def test_without_basic_deduction_skips_return_sheet(tmp_path):
    a = copy.deepcopy(HIGH)
    a["apply_basic_deduction"] = False
    raw = openpyxl.load_workbook(build(tmp_path, [{"label": "고가주택", "tool": "capital_gains_income", "args": a}]))
    assert B1 in raw.sheetnames and S84 not in raw.sheetnames
    txt = " ".join(str(x.value) for r in raw["목록"].iter_rows() for x in r if x.value)
    assert "apply_basic_deduction=True" in txt
