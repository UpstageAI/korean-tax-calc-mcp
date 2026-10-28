"""업무무관 지급이자 손금불산입 — 가지급금 상대방의 법인세법 특수관계인 판정 연결 (Claude 2026-10-28).

근거(조문 원문 대조):
- 법인세법 제28조①4호나목: 특수관계인에게 업무와 관련 없이 지급한 가지급금 등 → 업무무관 산식 적수.
- 특수관계인: 법 제2조제12호(2019.1.1. 시행) ← 영 제2조⑧(2024.1.1.~; 2019.2.12.~2023.12.31.은 제2조⑤, 그 전 제87조①).
  1호 사실상 영향력자와 그 친족, 2호 비소액주주등과 그 친족, 3호 임원·직원 등, 4·5호 지배 법인, 6호 30%×30%, 7호 계열회사.
손계산(영 제53조②, 대여일 산입·말일 산입):
- 차입금: 이자 10,000,000, 적수 36,500,000,000.
- 김배우 10,000,000 × 92일(2025-10-01~12-31) = 920,000,000 → 10,000,000 × 920,000,000 ÷ 36,500,000,000 = 252,054(절사).
- ㈜모회사 20,000,000 × 92일 = 1,840,000,000 → 합계 2,760,000,000 → 756,164(절사).
"""
import openpyxl
from test_form_templates import build, c

T = "nonbusiness_interest_disallowance"
CO = "㈜가람정밀"
LOANS = [{"interest": 10_000_000, "jeoksu": 36_500_000_000}]
REL = {
    "people": {CO: {"구분": "법인"}, "김대표": {"구분": "개인"}, "김배우": {"구분": "개인"}, "박거래": {"구분": "개인"},
               "㈜모회사": {"구분": "법인"}, "이직원": {"구분": "개인"}},
    "family": [["김대표", "김배우", "배우자"]],
    "stakes": [["김대표", CO, 0.60], ["㈜모회사", CO, 0.30]],
    "officers": {CO: ["김대표"]},
    "employees": {CO: ["이직원"]},
}


def run(adv, **kw):
    return c(T, {"loans": LOANS, "advances": adv, "year": 2025, "company": CO, "relations": REL, **kw})


def test_spouse_of_ceo_is_related_kin():
    # 대표이사(60% 주주)의 배우자 — 직원 아님, 지분 없음 → 영 제2조⑧2호(비소액주주등의 친족) → 기타특수관계인, 포함
    r = run([{"상대방": "김배우", "사유": "업무무관대여", "금액": 10_000_000, "지급일": "2025-10-01"}])
    row = r["계산"]["가지급금"][0]
    assert row["대상"] == "기타특수관계인" and row["포함액"] == 10_000_000 and row["적수"] == 920_000_000
    assert "제2조⑧2호" in row["판단"] and "친족" in row["판단"]
    assert row["근거"].startswith("법인세법 제2조제12호, 법인세법 시행령 제2조⑧2호")
    assert row["특수관계_판정"]["특수관계"] is True
    assert r["결과"]["손금불산입(기타사외유출)"] == 252_054


def test_unrelated_customer_excluded():
    r = run([{"상대방": "박거래", "사유": "업무무관대여", "금액": 10_000_000, "지급일": "2025-10-01"}])
    row = r["계산"]["가지급금"][0]
    assert row["대상"] == "비특수관계인" and row["포함액"] == 0 and row["적수"] == 0
    assert "해당 없음" in row["판단"] and "법인세법 제28조①4호나목" in row["근거"]
    assert r["결과"]["손금불산입(기타사외유출)"] == 0


def test_30pct_corporate_shareholder_is_related_corp():
    r = run([{"상대방": "㈜모회사", "사유": "업무무관대여", "금액": 20_000_000, "지급일": "2025-10-01"},
             {"상대방": "김배우", "사유": "업무무관대여", "금액": 10_000_000, "지급일": "2025-10-01"}])
    row = r["계산"]["가지급금"][0]
    assert row["대상"] == "특수관계법인" and "비소액주주" in row["판단"] and "30.00%" in row["판단"]
    assert r["결과"]["손금불산입(기타사외유출)"] == 10_000_000 * 2_760_000_000 // 36_500_000_000 == 756_164


def test_roles_from_officers_employees():
    r = run([{"상대방": "김대표", "사유": "업무무관대여", "금액": 1_000_000, "지급일": "2025-10-01"},
             {"상대방": "이직원", "사유": "급료가불", "금액": 3_000_000, "지급일": "2025-10-01", "월정급여": 2_000_000}])
    a, b = r["계산"]["가지급금"]
    assert a["대상"] == "임원" and a["특수관계_판정"]["역할"] == ["임원", "주주"]
    assert b["대상"] == "직원" and b["포함액"] == 1_000_000 and "제2조⑧3호" in b["판단"]


def test_conflict_is_error():
    r = run([{"대상": "비특수관계인", "상대방": "김배우", "사유": "업무무관대여", "금액": 10_000_000, "지급일": "2025-10-01"}])
    assert "error" in r and "입력 확인 필요" in r["error"] and "김배우" in r["문제"][0]
    r = run([{"대상": "직원", "상대방": "박거래", "금액": 1_000_000, "지급일": "2025-10-01"}])
    assert "error" in r


def test_consistent_given_target_kept():
    r = run([{"대상": "대표이사", "상대방": "김대표", "금액": 1_000_000, "지급일": "2025-10-01"}])
    row = r["계산"]["가지급금"][0]
    assert row["대상"] == "대표이사" and "경고" not in row["판단"]


def test_year_article_location():
    # 2018 사업연도: 영 제87조①, 법 제52조① / 2022: 영 제2조⑤
    r = c(T, {"loans": LOANS, "advances": [{"상대방": "김배우", "금액": 1_000_000, "지급일": "2018-10-01"}],
              "year": 2018, "company": CO, "relations": REL})
    row = r["계산"]["가지급금"][0]
    assert "제87조①2호" in row["판단"] and row["근거"].startswith("법인세법 제52조①")
    r = c(T, {"loans": LOANS, "advances": [{"상대방": "김배우", "금액": 1_000_000, "지급일": "2022-10-01"}],
              "year": 2022, "company": CO, "relations": REL})
    assert "제2조⑤2호" in r["계산"]["가지급금"][0]["판단"]


def test_missing_relations_or_unknown_party():
    r = c(T, {"loans": LOANS, "advances": [{"상대방": "김배우", "금액": 1, "지급일": "2025-10-01"}], "year": 2025})
    assert "error" in r and "relations" in r["문제"][0]
    r = run([{"상대방": "모르는사람", "금액": 1, "지급일": "2025-10-01"}])
    assert "error" in r


def test_old_inputs_unchanged():
    r = c(T, {"loans": LOANS, "advances": [{"대상": "대표이사", "금액": 10_000_000, "지급일": "2025-10-01"}], "year": 2025})
    row = r["계산"]["가지급금"][0]
    assert "특수관계_판정" not in row and r["결과"]["손금불산입(기타사외유출)"] == 252_054


def test_export_workbook_form_26(tmp_path):
    adv = [{"상대방": "김배우", "금액": 10_000_000, "지급일": "2025-10-01"},
           {"상대방": "박거래", "금액": 5_000_000, "지급일": "2025-10-01"}]
    path = build(tmp_path, [{"label": "지급이자", "tool": T,
                             "args": {"loans": LOANS, "advances": adv, "year": 2025, "company": CO, "relations": REL}}])
    wb = openpyxl.load_workbook(path)
    eul = wb["별지26(을) 업무무관 적수"]
    rows = [(eul.cell(r, 3).value, eul.cell(r, 4).value, eul.cell(r, 5).value, eul.cell(r, 8).value) for r in range(12, 14)]
    assert rows[0] == ("2025.10.01", "김배우 증가", 10_000_000, 92) and rows[1][0] in (None, "")
    from test_form_26 import Evaluator, GAP
    ev = Evaluator(wb)
    assert ev.v(GAP, "E8") == 920_000_000 and ev.v(GAP, "I8") == 252_054


def test_top_level_basis_and_judgment():
    r = run([{"상대방": "김배우", "금액": 10_000_000, "지급일": "2025-10-01"}])
    assert "특수관계인: 법인세법 제2조제12호, 법인세법 시행령 제2조⑧2호" in r["근거"]
    assert any("김배우 → 기타특수관계인" in x for x in r["판단"])
