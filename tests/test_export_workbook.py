"""계산 근거 엑셀 내보내기 export_workbook 정답 테스트. 작성 Claude 2026-10-24 (docs/SPEC_export_workbook.md)."""
import asyncio, json
import openpyxl
from korean_tax_calc_mcp.server import mcp


def c(tool, a):
    r = asyncio.run(mcp.call_tool(tool, a))
    return json.loads((r.content[0] if hasattr(r, "content") else r[0][0]).text)


ENT = {"year": 2025, "sme": True, "ledger": [{"항목": "일반", "금액": 95000000, "증빙": "적격"},
                                             {"항목": "관광", "금액": 6000000, "증빙": "적격", "성격": "문화"}],
       "revenue": {"매출액": 18000000000, "특수관계매출": 2000000000}}
ADJ = {"net_income": 100000000, "items": [{"항목": "법인세", "종류": "법인세비용", "금액": 8000000},
                                          {"항목": "환급금 이자", "종류": "환급금이자", "금액": 2000000}]}


def build(tmp_path, calls, **kw):
    return c("export_workbook", {"title": "제25기 세무조정 검토", "calls": calls, "output_dir": str(tmp_path), **kw})


def test_workbook_tabs_values_and_live_formulas(tmp_path):
    r = build(tmp_path, [{"label": "기업업무추진비", "tool": "entertainment_adjustment", "args": ENT},
                         {"label": "차가감소득금액", "tool": "income_adjustment_total", "args": ADJ}])
    path = r["결과"]["파일"]
    assert path.endswith(".xlsx") and r["결과"]["시트"][:1] == ["요약"]
    wb = openpyxl.load_workbook(path)  # 수식 그대로
    assert {"요약", "입력·가정", "판단·근거"} <= set(wb.sheetnames)
    assert any(s.startswith("계산") for s in wb.sheetnames)
    calc = wb[[s for s in wb.sheetnames if s.startswith("계산")][0]]
    formulas = [cell.value for row in calc.iter_rows() for cell in row if isinstance(cell.value, str) and cell.value.startswith("=")]
    assert any(f.upper().startswith("=SUM(") for f in formulas)
    vals = [cell.value for row in calc.iter_rows() for cell in row]
    assert 16600000 in vals and 84400000 in vals
    j = [cell.value for row in wb["판단·근거"].iter_rows() for cell in row]
    assert "법인세법 제25조④(문화접대비)" in j
    inp = [cell.value for row in wb["입력·가정"].iter_rows() for cell in row]
    assert 18000000000 in inp and any(isinstance(v, str) and "revenue" in v for v in inp)
    summ = [cell.value for row in wb["요약"].iter_rows() for cell in row]
    assert "차가감소득금액" in summ and 106000000 in summ
    assert any(isinstance(v, str) and "세무 자문이 아닙니다" in v for v in summ)


def test_workbook_records_tool_error_and_rejects_unknown_tool(tmp_path):
    r = build(tmp_path, [{"label": "잘못된 입력", "tool": "income_adjustment_total", "args": {"net_income": 1, "items": [{"항목": "가", "종류": "없는 종류", "금액": 1}]}}])
    wb = openpyxl.load_workbook(r["결과"]["파일"])
    assert any(isinstance(v, str) and "입력 확인 필요" in v for row in wb["요약"].iter_rows() for v in [c.value for c in row])
    r = build(tmp_path, [{"label": "x", "tool": "no_such_tool", "args": {}}])
    assert "입력 확인 필요" in json.dumps(r, ensure_ascii=False)


def test_workbook_no_overwrite_and_no_recursion(tmp_path):
    calls = [{"label": "차가감", "tool": "income_adjustment_total", "args": ADJ}]
    p1 = build(tmp_path, calls, filename="검토")["결과"]["파일"]
    p2 = build(tmp_path, calls, filename="검토")["결과"]["파일"]
    assert p1 != p2
    r = build(tmp_path, [{"label": "x", "tool": "export_workbook", "args": {}}])
    assert "입력 확인 필요" in json.dumps(r, ensure_ascii=False)
