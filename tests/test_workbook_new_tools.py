"""0.14.0 새 도구들의 살아 있는 수식(계산식_입력·계산식) 회귀 테스트.

각 도구 결과의 계산식 단계를 입력값으로 직접 계산해 도구 값과 맞는지(evaluate), export_workbook이 만든 엑셀의
수식 칸을 정의 이름만으로 다시 계산해 도구 값과 맞는지, 입력 칸을 바꾸면 도구를 다시 부른 값과 같아지는지 확인한다.
기대값은 조문으로 손계산(각 사례 주석).
"""
import copy, re, shutil, subprocess
import openpyxl
import pytest
from test_export_workbook_c import c, evaluate
from test_excel_defined_names import VALID, CELL_A1, CELL_R1C1, _names

# ── 사례(도구, args, {단계 이름: 손계산 값}) ─────────────────────────────────────

# 소득세법 제56조·제62조(2024 귀속): 이자 1,000만 + 상장배당 4,000만 → Gross-up = MIN(4,000만, 5,000만 − 2,000만) × 10% = 300만.
#  일반산출 = 기본세율(5,300만 − 2,000만 + 5,000만 = 8,300만) 1,416만 + 2,000만 × 14% 280만 = 1,696만
#  비교산출 = 5,000만 × 14% 700만 + 기본세율(5,000만) 624만 = 1,324만 → 공제 = MIN(300만, 372만) = 300만
DTC = {"year": 2024, "items": [{"kind": "예금이자", "amount": 10_000_000}, {"kind": "상장법인현금배당", "amount": 40_000_000}],
       "other_base": 50_000_000}
# 제62조2호 나목 단서(2025): 14% 500만·25% 300만, 출자공동사업자배당 400만, 과세표준 전체 6,000만 → 다른 과세표준 5,200만.
#  가목 = 70만 + 75만 = 145만, 나목 = MAX(기본세율(5,200만) 672만, 400만 × 14% 56만 + 기본세율(4,800만) 594만 = 650만) = 672만
DTC2 = {"year": 2025, "grossup": 0, "financial_14": 5_000_000, "financial_25": 3_000_000, "partner_dividend": 4_000_000,
        "taxable_base": 60_000_000}
# 제59조의2(2025): 8세 이상 3명 → 55만 + 40만, 넷째 출생(0세, ①인원 제외) → ③ 셋째 이상 70만
CTC = {"year": 2025, "children": [{"age": 10, "basic_deduction": True}, {"age": 12, "basic_deduction": True},
                                  {"age": 14, "basic_deduction": True},
                                  {"age": 0, "basic_deduction": True, "born_or_adopted_this_year": True, "order": 4}]}
# 2017: ① 2명 30만, ② 6세 이하 2명 → 1명 초과 15만
CTC2 = {"year": 2017, "children": [{"age": 3, "basic_deduction": True}, {"age": 5, "basic_deduction": True}]}
# 종교인소득(2025): 식사대 월 20만 × 12 = 240만, 자가운전(본인명의임차 2025~) MIN(300만, 240만), 보육 월 20만 × 12 = 240만,
#  종교활동비 500만 → 비과세 1,220만, 과세대상 3,780만 → 필요경비 1,600만 + 1,780만 × 50% = 2,490만 → 1,290만
REL = {"year": 2025, "amount": 50_000_000, "meal_allowance": 2_400_000, "meal_months": 12, "meal_provided": False,
       "vehicle_allowance": 3_000_000, "vehicle_months": 12, "vehicle": "본인명의임차", "religious_activity_fund": 5_000_000,
       "childcare_allowance": 3_600_000, "childcare_months": 12, "childcare_children": 2}
# 간주임대료(제25조①, 영 제53조③): A 4억×365, B 2억×200 → 적수 큰 A에서 3억 차감 → (1억×365 + 2억×200) = 765억
#  × 60% ÷ 365 × 3.1% = 3,898,356 − 금융수익 50만 = 3,398,356
HOUSES = [{"이름": "A", "기준시가": 500_000_000, "면적": 85, "보증금": 400_000_000},
          {"이름": "B", "기준시가": 400_000_000, "면적": 85, "보증금": 200_000_000, "일수": 200},
          {"이름": "C", "기준시가": 300_000_000, "면적": 60, "보증금": 0}]
HDR = {"year": 2025, "houses": HOUSES, "financial_income": 500_000, "rate": 0.031}
# 주택임대(제64조의2, 2025): 월세 600만 + 등록 B 900만(6개월 등록 → 450만). 일반 1,050만.
#  경비 450만×60% + 1,050만×50% = 795만, 추가 차감 400만×0.3 + 200만×0.7 = 260만 → 445만 × 14% = 623,000
HRI = {"year": 2025, "houses": [
    {"이름": "A", "기준시가": 500_000_000, "면적": 85, "월세": 6_000_000},
    {"이름": "B", "기준시가": 400_000_000, "면적": 85, "월세": 9_000_000, "등록임대": True, "임대유형": "장기일반",
     "아파트": False, "등록신청일": "2021-03-02", "등록_월수": 6}],
    "other_income_amount": 15_000_000, "other_taxable_base": 30_000_000, "comprehensive_rental_income": 7_000_000}
# 3주택 월세 1,800만 + 간주임대료 3,398,356 = 21,398,356 > 2천만 → 종합과세
HRI2 = {"year": 2025, "houses": [{**h, "월세": 6_000_000} for h in HOUSES], "financial_income": 500_000, "rate": 0.031}
# 제104조①(2024 양도): 토지 1년 이상 2년 미만 40% 4,000만 vs 비사업용 기본세율(1억) 1,956만 + 10% 1,000만 → 4,000만
CGT = {"base": 100_000_000, "asset_type": "토지", "acquisition_date": "2023-03-01", "transfer_date": "2024-06-01",
       "non_business_land": True}
# 제104조⑦(2021.8 양도, 3주택 조정지역): 기본세율(2억) 5,660만 + 30% 6,000만 = 1억1,660만
CGT2 = {"base": 200_000_000, "asset_type": "주택", "acquisition_date": "2015-03-01", "transfer_date": "2021-08-01",
        "houses_count": 3, "adjusted_area": True}
# 토지 10년 보유: 양도차익 3억, 장기보유 20% → 2억4,000만 − 기본공제 250만 = 2억3,750만 → 3,706만 + 8,750만 × 38% = 7,031만
CGI = {"asset_type": "토지", "acquisition_date": "2015-03-01", "transfer_date": "2025-06-01", "transfer_price": 500_000_000,
       "acquisition_price": 200_000_000, "apply_basic_deduction": True}
# 국세기본법 제35조: 2억9,700만 → 최우선 7,500 → 당해세 200 → 부가세 1,000 → 저당 1억 → 임금 3,000 → 소득세 2,000 → 대여금 6,000
M = 10_000
TPD = {"sale_amount": 30000 * M, "execution_cost": 300 * M, "sale_date": "2022-08-01", "property_type": "주택", "claims": [
    {"이름": "임차인", "종류": "소액임차보증금", "금액": 5500 * M},
    {"이름": "근로자", "종류": "최우선임금채권", "금액": 2000 * M},
    {"이름": "재산세", "종류": "지방세", "당해세": True, "세목": "재산세", "금액": 200 * M, "법정기일": "2021-07-16"},
    {"이름": "부가세", "종류": "국세", "금액": 1000 * M, "확정방식": "신고", "기준일": "2020-01-25"},
    {"이름": "은행", "종류": "저당권", "금액": 10000 * M, "설정일": "2020-05-01"},
    {"이름": "임금", "종류": "일반임금채권", "금액": 3000 * M},
    {"이름": "소득세", "종류": "국세", "금액": 2000 * M, "법정기일": "2021-05-31"},
    {"이름": "대여금", "종류": "일반채권", "금액": 10000 * M}]}
# ⑦(2024 매각): 종부세 1억 → 임차인 1억5천 → 잔여 0 (임차인이 종부세 몫 5천을 대신 받는 것은 채권별 값)
TPD2 = {"sale_amount": 25000 * M, "sale_date": "2024-03-01", "property_type": "주택", "claims": [
    {"이름": "임차인", "종류": "확정일자임차보증금", "금액": 20000 * M, "설정일": "2021-03-01"},
    {"이름": "종부세", "종류": "당해세", "세목": "종합부동산세", "금액": 10000 * M, "법정기일": "2023-12-15"}]}
# 공급시기(영 제28조③, 규칙 제17조): 장기할부 1,000만 × 3회 각 약정일 + 일반판매 인도일
VST = {"transactions": [
    {"이름": "기계", "종류": "장기할부", "인도일": "2025-02-10", "대가": [
        {"약정일": "2025-02-10", "금액": 10_000_000}, {"약정일": "2025-08-10", "금액": 10_000_000}, {"약정일": "2026-02-10", "금액": 10_000_000}]},
    {"이름": "상품", "종류": "일반판매", "인도일": "2025-05-20", "공급가액": 5_000_000}], "period": "2025-1기 확정"}
# 영 제17조①(2026): 월정액급여 = (급여 2,520만 + 자가운전 과세분 60만 + 식사대 240만 + 사회보험료 대납 180만) ÷ 12 = 250만 ≤ 260만
#  → 연장근로수당 비과세 MIN(400만, 240만), 과세 160만 → 총급여 3,520만
EGP = {"year": 2026, "items": [
    {"kind": "급여", "amount": 25_200_000},
    {"kind": "자가운전보조금", "amount": 3_000_000, "facts": {"months": 12, "vehicle": "본인소유", "business_use": True, "actual_travel_reimbursed": False}},
    {"kind": "식사대", "amount": 2_400_000, "facts": {"months": 12, "meal_provided": False}},
    {"kind": "연장근로수당", "amount": 4_000_000, "facts": {"eligible_occupation": True, "prior_year_gross": 30_000_000}},
    {"kind": "상여", "amount": 6_000_000},
    {"kind": "사회보험료_본인부담분_회사대납", "amount": 1_800_000}]}

CASES = [
    ("dividend_tax_credit", DTC, {"일반산출세액": 16_960_000, "비교산출세액": 13_240_000, "배당세액공제액": 3_000_000, "공제후세액": 13_960_000}),
    ("dividend_tax_credit", DTC2, {"비교_가목": 1_450_000, "비교_나목": 6_720_000, "비교산출세액": 8_170_000, "배당세액공제액": 0}),
    ("child_tax_credit", CTC, {"기본분": 950_000, "출산분": 700_000, "자녀세액공제": 1_650_000}),
    ("child_tax_credit", CTC2, {"기본분": 300_000, "추가분": 150_000, "자녀세액공제": 450_000}),
    ("religious_income", REL, {"비과세_합계": 12_200_000, "필요경비_표금액": 24_900_000, "기타소득금액": 12_900_000}),
    ("housing_deemed_rent_income", HDR, {"적수_합계": 76_500_000_000, "금융수익_차감_전": 3_898_356, "간주임대료": 3_398_356}),
    ("housing_rental_income", HRI, {"필요경비": 7_950_000, "추가차감": 2_600_000, "분리과세_세액": 623_000,
                                    "종합과세_주택임대분_세액_참고": 1_050_000}),
    ("housing_rental_income", HRI2, {"간주임대료": 3_398_356, "총수입금액": 21_398_356}),
    ("capital_gains_tax", CGT, {"기본세율세액": 19_560_000, "산출세액": 40_000_000, "지방소득세": 4_000_000}),
    ("capital_gains_tax", CGT2, {"기본세율세액": 56_600_000, "산출세액": 116_600_000}),
    ("capital_gains_income", CGI, {"과세표준": 237_500_000, "산출세액": 70_310_000, "지방소득세": 7_031_000}),
    ("tax_priority_distribution", TPD, {"배당할금액": 29700 * M, "잔액4": 11000 * M, "잔여": 0, "배분8": 6000 * M}),
    ("tax_priority_distribution", TPD2, {"단계1_배분": 10000 * M, "단계2_배분": 15000 * M, "잔여": 0}),
    ("vat_supply_timing", VST, {"거래_합계": 35_000_000, "조회기간_과세표준": 5_000_000}),
    ("earned_income_gross_pay", EGP, {"월정액급여": 2_500_000, "연장근로_비과세": 2_400_000, "총급여액": 35_200_000}),
]
# 단계 이름 → 결과 키(같은 값이어야 하는 것)
RESULT_KEY = {"간주임대료": "총수입금액_산입액"}


@pytest.mark.parametrize("tool, args, expect", CASES, ids=[f"{t}-{i}" for i, (t, _, _) in enumerate(CASES)])
def test_formula_steps_match_hand_values(tool, args, expect):
    r = c(tool, args)
    assert "error" not in r, r
    steps = evaluate(r)          # 모든 단계 식을 입력값으로 계산해 단계 값과 대조
    for name, want in expect.items():
        assert name in steps, f"{tool}: 단계 {name} 없음"
        assert steps[name]["값"] == want, (name, steps[name]["값"], want)
        key = RESULT_KEY.get(name, name)
        if key in r["결과"] and tool != "housing_rental_income":
            assert r["결과"][key] == want, (key, r["결과"][key])


def test_vat_period_sums_are_row_sums():
    steps = evaluate(c("vat_supply_timing", VST))
    assert steps["과세표준_2025_1기_예정"]["식"] == "거래1_공급1"
    assert steps["거래_합계"]["값"] == 35_000_000
    assert steps["조회기간_과세표준"]["식"] == "거래2_공급1"


def test_tax_priority_stage_balances_chain():
    steps = evaluate(c("tax_priority_distribution", TPD))
    for k in range(1, 8):
        assert steps[f"단계{k}_배분"]["식"].startswith("MIN(")
        assert steps[f"잔액{k}"]["식"] == f"{'배당할금액' if k == 1 else f'잔액{k - 1}'}-단계{k}_배분"


# ── 엑셀: 정의 이름만으로 수식 칸을 다시 계산 ─────────────────────────────────────

_FUNCS = {"MIN", "MAX", "IF", "ROUNDDOWN", "TRUE", "FALSE"}
_TOK = re.compile(r"[가-힣A-Za-z_][가-힣A-Za-z0-9_.]*")


class Book:
    """openpyxl로 연 통합 문서의 정의 이름 → 셀, 수식은 + - * / MIN MAX IF ROUNDDOWN 비교만 계산."""

    def __init__(self, path):
        self.wb = openpyxl.load_workbook(path)
        self.ref = {}
        for nm, dn in self.wb.defined_names.items():
            m = re.fullmatch(r"'?([^'!]+)'?!\$([A-Z]+)\$(\d+)", dn.attr_text)
            if m:
                self.ref[nm] = (m.group(1), f"{m.group(2)}{m.group(3)}")
        self.over = {}
        self.memo = {}

    def value(self, name):
        if name in self.over:
            return self.over[name]
        if name in self.memo:
            return self.memo[name]
        sh, cell = self.ref[name]
        v = self.wb[sh][cell].value
        if isinstance(v, str) and v.startswith("="):
            v = self.calc(v[1:])
        self.memo[name] = v
        return v

    def calc(self, f):
        def sub(m):
            t = m.group(0)
            return t if t in _FUNCS else f"V({t!r})"
        expr = _TOK.sub(sub, f).replace("ROUNDDOWN(", "RD(")
        assert not re.search(r"(?<![<>])=", expr), f
        env = {"V": self.value, "MIN": min, "MAX": max, "IF": lambda c_, a, b: a if c_ else b,
               "RD": lambda x, n=0: int(x), "TRUE": True, "FALSE": False}
        return eval(expr, {"__builtins__": {}}, env)


CALLS = [{"label": f"{t}-{i}", "tool": t, "args": a} for i, (t, a, _) in enumerate(CASES)]


def _book(tmp_path, layout="steps", calls=CALLS):
    r = c("export_workbook", {"layout": layout, "title": f"새 도구 수식 {layout}", "calls": calls, "output_dir": str(tmp_path)})
    assert "error" not in r, r
    return r["결과"]["파일"]


def _formula_rows(ws):
    """계산 시트의 단계 표: (단계, 수식, 도구 값) — 수식 칸이 '='로 시작하는 행."""
    out = []
    for row in ws.iter_rows(min_col=1, max_col=4):
        a, b, cc, d = (x.value for x in row)
        if isinstance(b, str) and b.startswith("=") and isinstance(d, str) and d.startswith("=IF(B"):
            out.append((a, b, cc))
    return out


def test_steps_workbook_formulas_recompute_to_tool_values(tmp_path):
    path = _book(tmp_path)
    bk = Book(path)
    for i, (tool, args, expect) in enumerate(CASES, 1):
        ws = bk.wb[f"계산 {i}"]
        rows = _formula_rows(ws)
        assert rows, (i, tool)
        for 단계, 수식, 도구값 in rows:
            got = bk.calc(수식[1:])
            assert abs(got - 도구값) < 1, (tool, 단계, 수식, got, 도구값)
        for name, want in expect.items():
            assert abs(bk.value(f"호출{i}_{name}") - want) < 1, (tool, name)


def test_forms_workbook_calc_sheets_have_formulas(tmp_path):
    path = _book(tmp_path, "forms")
    bk = Book(path)
    sheets = [ws for ws in bk.wb.worksheets if ws.title.startswith("계산 ")]
    assert len(sheets) >= len({t for t, _, _ in CASES} - {"earned_income_gross_pay"})
    for ws in sheets:
        rows = _formula_rows(ws)
        assert rows, ws.title
        for 단계, 수식, 도구값 in rows:
            assert abs(bk.calc(수식[1:]) - 도구값) < 1, (ws.title, 단계)


@pytest.mark.parametrize("layout", ["steps", "forms"])
def test_defined_names_valid_and_all_defined(tmp_path, layout):
    names, formulas = _names(_book(tmp_path, layout))
    user = [n for n, _, _ in names if not n.startswith("_xlnm.")]
    bad = [n for n in user if not VALID.match(n) or CELL_A1.match(n) or CELL_R1C1.match(n) or len(n) > 255]
    assert not bad, bad[:10]
    assert len({n.lower() for n in user}) == len(user)
    defined = {n.lower() for n in user}
    used = {t for f in formulas for t in re.findall(r"호출\d+_[가-힣A-Za-z0-9_.]+", re.sub(r'"[^"]*"', "", f))}
    assert used and not sorted(t for t in used if t.lower() not in defined)


# ── 입력 칸을 바꾸면 수식이 따라가 도구를 다시 부른 값과 같아진다 ─────────────────────

LIVE = [
    ("religious_income", REL, "받은금액총액", 60_000_000, {"amount": 60_000_000}, "기타소득금액"),
    ("dividend_tax_credit", DTC, "다른과세표준", 80_000_000, {"other_base": 80_000_000}, "공제후세액"),
    ("capital_gains_tax", CGT, "과세표준", 30_000_000, {"base": 30_000_000}, "산출세액"),
    ("housing_deemed_rent_income", HDR, "보증금2", 350_000_000,
     {"houses": [HOUSES[0], {**HOUSES[1], "보증금": 350_000_000}, HOUSES[2]]}, "총수입금액_산입액"),
    ("tax_priority_distribution", TPD, "매각대금", 25000 * M, {"sale_amount": 25000 * M}, "잔여"),
    ("vat_supply_timing", VST, "거래2_공급1", 7_000_000,
     {"transactions": [VST["transactions"][0], {**VST["transactions"][1], "공급가액": 7_000_000}]}, "조회기간_과세표준"),
    ("child_tax_credit", CTC, "대상인원", 4, None, "자녀세액공제"),
]
STEP_OF = {"총수입금액_산입액": "간주임대료"}


@pytest.mark.parametrize("tool, args, inp, new, new_args, key", LIVE, ids=[x[0] for x in LIVE])
def test_changing_input_cell_recomputes(tmp_path, tool, args, inp, new, new_args, key):
    path = _book(tmp_path, calls=[{"label": tool, "tool": tool, "args": args}])
    bk = Book(path)
    bk.over[f"호출1_{inp}"] = new
    got = bk.value(f"호출1_{STEP_OF.get(key, key)}")
    if new_args is None:      # 자녀 4명(8세 이상) → 55만 + 40만 × 2 = 135만 + 출산 70만
        assert got == 2_050_000
        return
    a = copy.deepcopy(args); a.update(new_args)
    want = c(tool, a)["결과"][key]
    assert got == want, (tool, got, want)


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice 없음")
def test_libreoffice_recalc_all_match(tmp_path):
    path = _book(tmp_path)
    out = tmp_path / "recalc"
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(out), path],
                   check=True, capture_output=True, timeout=180)
    wb = openpyxl.load_workbook(next(out.glob("*.xlsx")), data_only=True)
    for i in range(1, len(CASES) + 1):
        checks = [x.value for row in wb[f"계산 {i}"].iter_rows() for x in row if x.value in ("일치", "확인 필요")]
        assert checks and set(checks) == {"일치"}, (i, CASES[i - 1][0], checks)
