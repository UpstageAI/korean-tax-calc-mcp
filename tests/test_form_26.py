"""별지 제26호서식(갑)·(을) 업무무관부동산등에관련한차입금이자조정명세서 — 서식 틀 회귀 (Claude 2026-10-28).

서식 원문: 법인세법 시행규칙 별지 제26호서식(갑) <개정 2006.3.14>·(을) <개정 2021.10.28> (2026.7.1. 시행본 PDF 대조).
값은 조문으로 손계산(법인세법 제28조①1호·4호, 시행령 제53조②③·제55조):
- 채권자불분명 사채이자 1,500,000 → (갑) ⑬ 상단, 업무무관 산식 지급이자·차입금에서 뺌.
- 가지급금 등 적수 490,000,000 + 184,000,000 + 244,000,000 + 505,500,000 = 1,423,500,000 (을 ⑧ 계 = 갑 ④).
- ⑧ = 3,250,000 × min(1,423,500,000, 23,725,000,000) ÷ 23,725,000,000 = 195,000. 합계 1,500,000 + 195,000 = 1,695,000.
"""
import re, shutil
import openpyxl
import pytest
from test_form_templates import build, c, recalc
from test_new_l import ADV, LOANS, T

GAP, EUL = "별지26(갑) 업무무관 지급이자", "별지26(을) 업무무관 적수"


# ── 작은 수식 계산기(이 서식이 쓰는 함수만: SUM·N·MIN·MAX·IF·ROUND·ROUNDDOWN·ISNUMBER) ──
def _cell_rc(a):
    m = re.match(r"([A-Z]+)(\d+)$", a)
    col = 0
    for ch in m.group(1):
        col = col * 26 + ord(ch) - 64
    return int(m.group(2)), col


class Evaluator:
    def __init__(self, wb):
        self.wb, self.memo = wb, {}

    def v(self, sheet, addr):
        key = (sheet, addr)
        if key not in self.memo:
            x = self.wb[sheet][addr].value
            self.memo[key] = self.formula(sheet, x[1:]) if isinstance(x, str) and x.startswith("=") else x
        return self.memo[key]

    def rng(self, sheet, a, b):
        (r1, c1), (r2, c2) = _cell_rc(a), _cell_rc(b)
        return [self.v(sheet, f"{openpyxl.utils.get_column_letter(cc)}{rr}")
                for rr in range(r1, r2 + 1) for cc in range(c1, c2 + 1)]

    def formula(self, sheet, f):
        refs = []

        def ext(m):
            refs.append((m.group(1), m.group(2)))
            return f"__R{len(refs) - 1}__"
        f = re.sub(r"'([^']+)'!([A-Z]+\d+)", ext, f)
        f = re.sub(r"\b([A-Z]+\d+):([A-Z]+\d+)\b", lambda m: f"self.rng(SH,{m.group(1)!r},{m.group(2)!r})", f)
        f = re.sub(r"(?<![\w'])([A-Z]{1,2}\d+)\b", lambda m: f"self.v(SH,{m.group(1)!r})", f)
        f = re.sub(r"__R(\d+)__", lambda m: f"self.v({refs[int(m.group(1))][0]!r},{refs[int(m.group(1))][1]!r})", f)
        f = re.sub(r"(?<![<>=!])=(?!=)", "==", f)
        f = re.sub(r"\b(SUM|N|MIN|MAX|IF|ROUND|ROUNDDOWN|ISNUMBER)\(", lambda m: f"F_{m.group(1)}(", f)
        num = lambda x: isinstance(x, (int, float)) and not isinstance(x, bool)
        flat = lambda a: [y for x in a for y in (x if isinstance(x, list) else [x])]
        env = {"self": self, "SH": sheet,
               "F_SUM": lambda *a: sum(x for x in flat(a) if num(x)),
               "F_N": lambda x: x if num(x) else 0,
               "F_MIN": lambda *a: min(x for x in flat(a) if num(x)),
               "F_MAX": lambda *a: max(x for x in flat(a) if num(x)),
               "F_IF": lambda c_, a, b: a if c_ else b,
               "F_ROUND": lambda x, n: round(x, int(n)),
               "F_ROUNDDOWN": lambda x, n: int(x) if int(n) == 0 else int(x * 10 ** n) / 10 ** n,
               "F_ISNUMBER": num}
        return eval(f, env)


@pytest.fixture(scope="module")
def hand(tmp_path_factory):
    d = tmp_path_factory.mktemp("f26")
    path = build(d, [{"label": "지급이자", "tool": T, "args": {"loans": LOANS, "advances": ADV, "year": 2025}}])
    return path, openpyxl.load_workbook(path), d


def test_sheets_and_list(hand):
    _, wb, _ = hand
    assert wb.sheetnames == ["목록", GAP, EUL, "입력·가정", "판단·근거"]
    listed = [wb["목록"].cell(r, 2).value for r in range(2, 4)]
    assert listed == ["법인세법 시행규칙 별지 제26호서식(갑)", "법인세법 시행규칙 별지 제26호서식(을)"]


def test_gap_layout_and_formulas(hand):
    _, wb, _ = hand
    ws = wb[GAP]
    merged = {str(m) for m in ws.merged_cells.ranges}
    assert {"A6:B7", "C6:F6", "D10:E10", "F10:G10", "H10:I10", "A12:A13", "H12:H13", "I24:I25"} <= merged
    assert "(①×⑦÷⑥)" in ws["I6"].value and "(⑩-⑬-<16>)" in ws["H11"].value and ws["A26"].value == "합 계"
    assert ws.print_area and "A1:I27" in str(ws.print_area).replace("$", "")
    assert ws["A8"].value == "=H26"                                    # ① = <18> 합계(작성요령 1)
    assert ws["E8"].value.startswith(f"='{EUL}'!I28-'{EUL}'!I35")       # ④ = (을) ⑧ 계 − ⑨ 계(동일인 상계)
    assert ws["F8"].value == "=C8+D8+E8" and ws["G8"].value == "=I26" and ws["H8"].value == "=MIN(F8,G8)"
    assert "ROUNDDOWN(A8*H8/G8,0)" in ws["I8"].value                   # ⑧ = ①×⑦÷⑥, 원 미만 절사
    # 채권자불분명 차입금 행: ⑩ = ⑬ 상단 = 1,500,000, <18> = ⑩-⑬-<16>
    assert (ws["B12"].value, ws["D12"].value) == (1_500_000, 1_500_000)
    assert ws["H12"].value == "=N(B12)-N(D12)-N(D13)-N(F12)-N(F13)"
    assert (ws["B14"].value, ws["C14"].value) == (3_250_000, 23_725_000_000)
    assert ws["B16"].value in (None, "")                               # 기업구매자금대출(영 제53조④2호)은 적지 않음


def test_eul_ledger_rows(hand):
    _, wb, _ = hand
    ws = wb[EUL]
    rows = [(ws.cell(r, 3).value, ws.cell(r, 4).value, ws.cell(r, 5).value, ws.cell(r, 8).value) for r in range(12, 16)]
    assert rows == [("2025.05.01", "영업부장 증가", 2_000_000, 245), ("2025.07.01", "직원갑 증가", 1_000_000, 184),
                    ("2025.09.01", "직원을 증가", 2_000_000, 122), ("2025.11.22", "대표이사 증가", 12_637_500, 40)]
    assert ws["G12"].value == "=N(E12)-N(F12)" and ws["I12"].value == "=G12*H12"
    assert ws["I28"].value == "=SUM(I12:I27)" and ws["C28"].value == "계"


def test_values_and_checks(hand):
    _, wb, _ = hand
    ev = Evaluator(wb)
    g = lambda a: ev.v(GAP, a)
    assert [g(a) for a in ("A8", "C8", "D8", "E8", "F8", "G8", "H8", "I8")] == [
        3_250_000, 0, 0, 1_423_500_000, 1_423_500_000, 23_725_000_000, 1_423_500_000, 195_000]
    assert g("D26") == 1_500_000 and g("B26") == 4_750_000
    assert g("D26") + g("I8") == 1_695_000                              # 채권자불분명 + 업무무관 = 손금불산입 합계
    assert ev.v(EUL, "I28") == 1_423_500_000
    checks = [ev.v(s, x.coordinate) for s in (GAP, EUL) for row in wb[s].iter_rows() for x in row
              if isinstance(x.value, str) and '"일치"' in x.value]
    assert len(checks) == 11 and set(checks) == {"일치"}


def test_deposit_offset(tmp_path):
    # 대표이사 가수금 1,000,000(12/1~12/19, 19일) → 가수금 적수 19,000,000, 동일인 상계(영 제53조③ 후단)
    # ④ = 1,423,500,000 − 19,000,000 = 1,404,500,000, ⑧ = 3,250,000 × 1,404,500,000 ÷ 23,725,000,000 = 192,397(절사)
    dep = [{"이름": "대표이사", "금액": 1_000_000, "일자": "2025-12-01", "반환일": "2025-12-20"}]
    path = build(tmp_path, [{"label": "상계", "tool": T, "args": {"loans": LOANS, "advances": ADV, "deposits": dep, "year": 2025}}])
    wb = openpyxl.load_workbook(path)
    ws = wb[EUL]
    assert [ws.cell(29, cc).value for cc in (3, 4, 6, 8)] == ["2025.12.01", "대표이사 가수금 증가", 1_000_000, 19]
    assert ws["G29"].value == "=N(F29)-N(E29)" and ws["G30"].value == "=G29+N(F30)-N(E30)"
    ev = Evaluator(wb)
    assert ev.v(EUL, "I35") == 19_000_000
    assert ev.v(GAP, "E8") == 1_404_500_000 and ev.v(GAP, "I8") == 192_397
    checks = [ev.v(s, x.coordinate) for s in (GAP, EUL) for row in wb[s].iter_rows() for x in row
              if isinstance(x.value, str) and '"일치"' in x.value]
    assert checks and set(checks) == {"일치"}


def test_simple_mode_direct_jeoksu(tmp_path):
    # 적수 직접 입력: 30,000,000 × 876,000,000 ÷ 3,650,000,000 = 7,200,000 (영 제53조②)
    path = build(tmp_path, [{"label": "직접", "tool": T, "args": {"loans": [{"interest": 30_000_000, "jeoksu": 3_650_000_000}],
                                                                  "nonbusiness_jeoksu": 876_000_000}}])
    wb = openpyxl.load_workbook(path)
    assert wb[EUL]["D6"].value == "업무무관자산 적수(직접 입력)"
    ev = Evaluator(wb)
    assert ev.v(GAP, "C8") == 876_000_000 and ev.v(GAP, "I8") == 7_200_000
    checks = [ev.v(s, x.coordinate) for s in (GAP, EUL) for row in wb[s].iter_rows() for x in row
              if isinstance(x.value, str) and '"일치"' in x.value]
    assert checks and set(checks) == {"일치"}


def test_tool_ledger_matches_person_jeoksu():
    dep = [{"이름": "대표이사", "금액": 1_000_000, "일자": "2025-12-01", "반환일": "2025-12-20"}]
    r = c(T, {"loans": LOANS, "advances": ADV, "deposits": dep, "year": 2025})
    k = r["계산"]
    assert r["결과"]["손금불산입_합계"] == 1_500_000 + 192_397
    assert sum(x["적수"] for x in k["적수표"]["가지급금"]) == k["가지급금적수_합계"] == 1_423_500_000
    assert sum(x["적수"] for x in k["적수표"]["가수금"]) == k["가수금적수_합계"] == 19_000_000
    assert k["불분명이자_합계"] == 1_500_000 and k["가수금_상계초과적수"] == 0 and k["업무무관자산적수"] == 0


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice 없음")
def test_recalc_libreoffice(hand):
    path, _, d = hand
    calc = recalc(path, d)
    assert calc[GAP]["I8"].value == 195_000 and calc[GAP]["D26"].value == 1_500_000
    checks = [x.value for s in (GAP, EUL) for row in calc[s].iter_rows() for x in row if x.value in ("일치", "확인 필요")]
    assert len(checks) == 11 and set(checks) == {"일치"}
