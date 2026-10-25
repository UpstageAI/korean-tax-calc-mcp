"""법정서식 2단계(별지3·15·20(1)·16의2·50(을)) 정답 테스트. 작성 Claude 2026-10-25 (docs/SPEC_form_templates2.md).
값은 가상 사례(가람정밀, 2025) 손 검산."""
import copy, pathlib, re, shutil
import openpyxl
import pytest
from test_form_templates import CALLS, build, recalc, c

S3, S15, S20, S16, S50 = ("별지3 과세표준및세액조정계산서", "별지15 소득금액조정합계표", "별지20(1) 감가상각비(정률법)",
                          "별지16의2 수입배당금액명세서", "별지50(을) 자본금과적립금")
soffice = pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice 없음")


@pytest.fixture(scope="module")
def demo(tmp_path_factory):
    d = tmp_path_factory.mktemp("tpl2")
    path = build(d, CALLS)
    return openpyxl.load_workbook(path), (recalc(path, d) if shutil.which("soffice") else None)


def test_sheet_order(demo):
    raw, _ = demo
    n = raw.sheetnames
    assert n[:8] == ["목록", S3, S15, "별지23(갑) 기업업무추진비", "별지23(을) 기업업무추진비", S20, S16, S50], n
    assert [x for x in n if re.fullmatch(r"계산 \d+", x)] == ["계산 1", "계산 2", "계산 3", "계산 4"]
    assert n[-2:] == ["입력·가정", "판단·근거"]


def test_links_between_forms(demo):
    raw, _ = demo
    assert "'별지15 소득금액조정합계표'!B31" in raw[S3]["F6"].value and "'별지15 소득금액조정합계표'!F31" in raw[S3]["F7"].value
    ws = raw[S15]
    rows = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(8, 31) if ws.cell(r, 1).value}
    assert "'별지23(갑) 기업업무추진비'!D27" in str(rows["기업업무추진비 한도초과"])
    assert "'별지20(1) 감가상각비(정률법)'!E31" in str(rows["기계장치 감가상각비 한도초과"])
    assert rows["법인세비용"] == 20000000                                    # 연결할 서식 없음 → 숫자
    assert "'별지16의2 수입배당금액명세서'!F28" in str(ws["F8"].value)
    assert "'별지20(1) 감가상각비(정률법)'!E16" in str(raw[S50]["B8"].value)


@soffice
def test_values_after_recalc(demo):
    _, calc = demo
    assert [calc[S3].cell(r, 6).value for r in range(5, 13)] == [300000000, 27830000, 7200000, 320630000, 0, 0, 320630000, 320630000]
    ws = calc[S15]
    assert [ws.cell(r, 3).value for r in (8, 9, 10)] == ["기타사외유출", "기타사외유출", "유보"]
    assert [str(ws.cell(r, 4).value) for r in (8, 9, 10)] == ["500", "500", "400"]
    assert (ws["B31"].value, ws["F31"].value, ws["G8"].value, str(ws["H8"].value)) == (27830000, 7200000, "기타", "200")
    w = calc[S20]
    assert [w.cell(r, 5).value for r in (10, 11, 12, 17, 19, 27, 28, 29, 31, 32, 33)] == [
        100000000, 87500000, 12500000, 31500000, 9450000, 9450000, 15000000, 5550000, 5550000, 0, 9550000]
    assert w["E18"].value == 0.3 and w["D33"].value == 9550000
    w = calc[S16]
    assert [w.cell(20, c).value for c in range(1, 7)] == ["B", 15000000, 80, 12000000, 4800000, 7200000]
    assert w["F11"].value == 40 and w["F28"].value == 7200000
    assert [calc[S50].cell(8, c).value for c in range(2, 6)] == [4000000, 0, 5550000, 9550000]
    checks = [x.value for s in (S3, S15, S20, S16) for r in calc[s].iter_rows() for x in r if x.value in ("일치", "확인 필요")]
    assert len(checks) >= 8 and set(checks) == {"일치"}


def test_adjustment_items_carry_kind():
    adj = next(x for x in CALLS if x["tool"] == "income_adjustment_total")
    r = c("income_adjustment_total", adj["args"])
    assert [i["종류"] for i in r["결과"]["세무조정"]] == ["법인세비용", "기업업무추진비 한도초과", "감가상각비 한도초과", "수입배당금 익금불산입"]


@soffice
def test_reversal_and_two_investees(tmp_path):
    calls = copy.deepcopy(CALLS)
    dep = next(x for x in calls if x["tool"] == "depreciation_adjustment"); dep["args"]["booked_depreciation"] = 5000000
    div = next(x for x in calls if x["tool"] == "dividend_received_deduction")
    div["args"]["investees"].append({"name": "C", "listed": True, "ratio": 0.1, "dividend": 4000000, "book_value_jeoksu": 36500000000})
    adj = next(x for x in calls if x["tool"] == "income_adjustment_total")
    adj["args"]["items"] = [i for i in adj["args"]["items"] if i["종류"] not in ("감가상각비 한도초과", "수입배당금 익금불산입")] + [
        {"항목": "수입배당금 익금불산입", "종류": "수입배당금 익금불산입", "금액": 8100000}]
    path = build(tmp_path, calls)
    calc = recalc(path, tmp_path)
    w = calc[S20]
    assert (w["E29"].value, w["E31"].value, w["E32"].value, w["E33"].value) == (-4450000, 0, 4000000, 0)   # 시인부족 → 추인
    assert [calc[S50].cell(8, col).value for col in range(2, 6)] == [4000000, 4000000, 0, 0]
    s16 = calc[S16]
    assert [s16.cell(21, col).value for col in range(1, 7)] == ["C", 4000000, 30, 1200000, 300000, 900000]
    assert s16["F28"].value == 8100000


@soffice
def test_fallback_calc_sheets_keep_live_formulas(demo):
    raw, calc = demo
    for name in ("계산 1", "계산 2", "계산 3", "계산 4"):
        fs = [x.value for r in raw[name].iter_rows() for x in r if isinstance(x.value, str) and x.value.startswith("=")]
        assert len(fs) >= 3, (name, "계산식 단계 수식 없음")               # steps 모양(입력·계산 표 + 확인) 그대로
        checks = [x.value for r in calc[name].iter_rows() for x in r if x.value in ("일치", "확인 필요")]
        assert checks and set(checks) == {"일치"}, (name, checks)


@soffice
def test_two_checks_on_same_row(demo):
    _, calc = demo
    checks = [x.value for r in calc[S15].iter_rows() for x in r if x.value in ("일치", "확인 필요")]
    assert checks == ["일치", "일치"]                                        # ②합계(J열)·⑤합계(L열) 둘 다
