"""엑셀 법정서식 그대로(서식 틀 + 값 채우기) 1단계 정답 테스트. 작성 Claude 2026-10-25 (docs/SPEC_form_templates.md).
값은 가상 사례(가람정밀, 2025) 손 검산."""
import asyncio, json, pathlib, re, shutil, subprocess
import openpyxl
import pytest
from korean_tax_calc_mcp.server import mcp
from test_gold_gaps5 import ENT as MARKET_LEDGER, REV as MARKET_REV, VOUCHER

CALLS = json.loads((pathlib.Path(__file__).parent / "data" / "workbook_forms_calls.json").read_text(encoding="utf-8"))
GAP, EUL = "별지23(갑) 기업업무추진비", "별지23(을) 기업업무추진비"
soffice = pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice 없음")


def c(tool, a):
    r = asyncio.run(mcp.call_tool(tool, a))
    return json.loads((r.content[0] if hasattr(r, "content") else r[0][0]).text)


def build(d, calls, **kw):
    r = c("export_workbook", {"title": "서식", "calls": calls, "output_dir": str(d), "company": "㈜가람정밀(가상)",
                              "period": "2025.01.01 ~ 2025.12.31", **kw})
    assert "결과" in r, r
    return r["결과"]["파일"]


def recalc(path, d):
    out = pathlib.Path(d) / "re"
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(out), path], check=True, capture_output=True, timeout=180)
    return openpyxl.load_workbook(next(out.glob("*.xlsx")), data_only=True)


@pytest.fixture(scope="module")
def demo(tmp_path_factory):
    d = tmp_path_factory.mktemp("tpl")
    path = build(d, CALLS)
    return openpyxl.load_workbook(path), (recalc(path, d) if shutil.which("soffice") else None)


def test_sheet_order(demo):
    raw, _ = demo
    n = raw.sheetnames
    assert n[0] == "목록" and n.index(EUL) == n.index(GAP) + 1, n
    assert n[-2:] == ["입력·가정", "판단·근거"] and "요약" not in n
    assert sum(bool(re.fullmatch(r"계산 \d+", x)) for x in n) == 4          # 틀 없는 나머지 4개 call(2단계 기준)


def test_template_layout(demo):
    raw, _ = demo
    ws = raw[GAP]
    merged = {str(m) for m in ws.merged_cells.ranges}
    assert {"A9:A19", "B10:B13", "B14:B17", "A20:A21", "A22:A25", "C2:C3", "A27:C27"} <= merged
    assert ws.sheet_view.showGridLines is False
    assert ws.print_area and "A1:D29" in str(ws.print_area).replace("$", "")
    assert ws["A6"].font.name == "맑은 고딕"
    assert "㈜가람정밀(가상)" in ws["D2"].value and "2025.01.01" in ws["B2"].value and "2025.12.31" in ws["B2"].value
    assert "월수( 12 )" in ws["B9"].value
    assert ws["D6"].value.startswith("=") and "'별지23(을) 기업업무추진비'!F13" in ws["D6"].value
    assert re.search(r"D8.*D26", ws["D27"].value)                         # ⑯ = ③ − ⑮
    assert "호출1_" in ws["D9"].value                                     # ④는 입력·가정 이름 참조
    assert ws["D6"].border.left.style == "thin" and ws["A6"].border.left.style == "medium"


@soffice
def test_values_after_recalc(demo):
    _, calc = demo
    ws = calc[GAP]
    assert [ws.cell(r, 4).value for r in range(6, 29)] == [
        87000000, 0, 87000000, 36000000, 30000000, 10000000, 0, 40000000, 30000000, 4000000, 0, 34000000,
        600000, 70600000, 20000000, 14120000, 0, 0, 0, 0, 84720000, 2280000, 84720000]
    e = calc[EUL]
    assert (e["C7"].value, e["E7"].value, e["F7"].value) == (12000000000, 3000000000, 15000000000)
    assert e["F13"].value == 87000000 and e["F22"].value == 0
    checks = [x.value for r in ws.iter_rows() for x in r if x.value in ("일치", "확인 필요")]
    assert len(checks) >= 6 and set(checks) == {"일치"}


@soffice
@pytest.mark.parametrize("year, ledger, want", [
    (2025, MARKET_LEDGER, {22: 9000000, 23: 0, 24: 0, 25: 9000000, 27: 16600000}),
    (2024, MARKET_LEDGER, {22: 9000000, 25: 7840000, 27: 17760000}),
    (2026, VOUCHER, {22: 0, 23: 4000000, 24: 3000000, 25: 7000000}),
    (2022, MARKET_LEDGER, {22: 0, 25: 0}),
])
def test_market_rows_by_year(tmp_path, year, ledger, want):
    path = build(tmp_path, [{"label": "업추", "tool": "entertainment_adjustment",
                             "args": {"year": year, "sme": True, "ledger": ledger, "revenue": MARKET_REV}}])
    ws = recalc(path, tmp_path)[GAP]
    for row, v in want.items():
        assert ws.cell(row, 4).value == v, (year, row, ws.cell(row, 4).value)
    checks = [x.value for r in ws.iter_rows() for x in r if x.value in ("일치", "확인 필요")]
    assert set(checks) == {"일치"}, checks


def test_list_sheet(demo):
    raw, _ = demo
    txt = " ".join(str(x.value) for r in raw["목록"].iter_rows() for x in r if x.value)
    assert "별지 제23호서식(갑)" in txt and "별지 제23호서식(을)" in txt
    assert "별지 제50호서식(갑)" in txt and "자료 없음" in txt and "별지 제50호서식(병)" in txt and "해당 없음" in txt


def test_market_spend_in_calc():
    r = c("entertainment_adjustment", {"year": 2026, "sme": True, "ledger": VOUCHER, "revenue": MARKET_REV})
    assert (r["계산"]["전통시장_지출"], r["계산"]["지역사랑상품권_지출"], r["계산"]["온누리상품권_지출"]) == (0, 4000000, 3000000)
    r = c("entertainment_adjustment", {"year": 2024, "sme": True, "ledger": VOUCHER, "revenue": MARKET_REV})
    assert (r["계산"]["지역사랑상품권_지출"], r["계산"]["온누리상품권_지출"]) == (0, 0)


def test_steps_layout_unchanged(tmp_path):
    path = build(tmp_path, CALLS[:1], layout="steps")
    assert "요약" in openpyxl.load_workbook(path).sheetnames


def test_render_style(demo):
    raw, _ = demo
    ws = raw[GAP]
    for addr in ("D6", "D19", "D27"):
        assert ws[addr].number_format == "#,##0", (addr, ws[addr].number_format)
        assert ws[addr].alignment.horizontal == "right"
    assert ws.row_dimensions[9].height == 30 and ws.row_dimensions[25].height == 40   # row_heights
    assert ws.row_dimensions[6].height == 24                                            # default_row_height
    assert ws["A9"].alignment.vertical == "center" and ws["A9"].alignment.wrap_text
    assert ws["B9"].alignment.vertical == "center"
    assert "A1:C1" in {str(m) for m in ws.merged_cells.ranges}
    lst = " ".join(str(x.value) for r in raw["목록"].iter_rows() for x in r if x.value)
    assert "entertainment_adjustment" not in lst                                        # 목록에 도구 영문 이름 없음
